#!/usr/bin/env python3
"""校验七章终稿正文、Word、PDF、关键数值及样本参考矩阵。"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import re
import subprocess
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from docx import Document
from lxml import etree


FORBIDDEN = [
    ".py", ".csv", "results/runs", "capacity_action_delta_mva", "storage_modules",
    "cumulative_eac", "physical_clr", "ROBUST", "WEAKLY_SENSITIVE", "SENSITIVE",
    "证据链", "候选包", "artifact", "manifest", "pipeline", "fallback",
    "GitHub Actions", "case id", "task id",
]

EXPECTED_CHAPTER_TITLES = [
    "第一章 绪论",
    "第二章 文献综述与理论基础",
    "第三章 徐州典型地区“新能源—电网”运行特征分析",
    "第四章 基于实际工程的电网建设成本模型",
    "第五章 “一片一策”弹性容载比规划建议",
    "第六章 典型案例验证",
    "第七章 结论与展望",
]

EXPECTED_SECTIONS = [
    "1.1 研究背景与意义",
    "1.2 研究范围与对象",
    "1.3 研究目标与内容",
    "1.4 研究方法与技术路线",
    "2.1 国内外研究现状",
    "2.2 核心概念与理论基础",
    "3.1 徐州新能源发展与电力负荷概况",
    "3.2 主变负载率多工况统计分析",
    "3.3 源荷比、电量渗透率与容载比的关联规律",
    "3.4 “新能源渗透水平—电网运行状态”数据库的构建",
    "4.1 分布式新能源配套工程成本数据收集与处理",
    "4.2 成本驱动因子识别与参数化模型建立",
    "4.3 刚性容载比场景全寿命周期成本测算",
    "4.4 弹性容载比场景全寿命周期成本测算",
    "4.5 两场景经济技术对比与决策案例库形成",
    "5.1 历史负荷与分布式光伏接入特征分析典型片区负荷与分布式光伏预测",
    "5.2 含正向容载比与反向承载力协同约束的弹性规划评价指标体系设计",
    "5.3 基于层次分析法的指标赋权",
    "5.4 不同片区类型的容载比推荐区间及技术方案组合",
    "6.1 刚性规划方案复盘（以某高渗透率片区为例）",
    "6.2 含反向承载力校核的弹性容载比方案设计与对比",
    "6.3 效果评估与方法有效性验证",
]


def strip_math(text: str) -> str:
    return re.sub(r"\$\$.*?\$\$", "", text, flags=re.S)


def pdf_page_text(pdf: Path, page: int) -> str:
    result = subprocess.run(
        ["pdftotext", "-f", str(page), "-l", str(page), str(pdf), "-"],
        check=True, capture_output=True, text=True,
    )
    return result.stdout.strip()


def run_check(root: Path, markdown: Path, docx: Path, pdf: Path, output_dir: Path) -> dict:
    source = markdown.read_text(encoding="utf-8-sig")
    prose = strip_math(source)
    issues: list[str] = []

    # 1. 开题报告目录硬约束：一级、二级目录均不得自行改写。
    chapter_titles = re.findall(r"^# (第[一二三四五六七]+章[^\n]*)$", source, re.M)
    if chapter_titles != EXPECTED_CHAPTER_TITLES:
        issues.append(f"一级目录未严格遵循开题报告：{chapter_titles}")
    section_titles = re.findall(r"^## ([1-6]\.\d+[^\n]*)$", source, re.M)
    if section_titles != EXPECTED_SECTIONS:
        issues.append(f"二级目录未严格遵循开题报告：{section_titles}")

    for token in FORBIDDEN:
        if token.lower() in prose.lower():
            issues.append(f"正文含内部程序或治理表达：{token}")
    snake = sorted(set(re.findall(r"\b[a-z]+_[a-z][a-z0-9_]*\b", prose)))
    if snake:
        issues.append(f"正文含程序式字段：{snake}")

    # 2. 二次科学审查形成的核心口径。
    required_semantics = [
        "2022—2025年规划期累计在役等年成本",
        "存量容量豁免",
        "Rcap只约束规划期新增",
        "110 kV线路统计负载余度",
        "样本内一致性检查",
        "本报告不虚构具体AHP权重值",
        "不能形成具有外推能力的连续四维区间",
    ]
    for phrase in required_semantics:
        if phrase not in source:
            issues.append(f"缺少关键科学口径：{phrase}")

    forbidden_semantics = [
        "年化规划成本/万元·年⁻¹",
        "网络容量支撑裕度",
        "独立查询规则进行反向验证",
        "独立查询结果与正式优化结果一致",
        "推荐上限3.0",
    ]
    for phrase in forbidden_semantics:
        if phrase in source:
            issues.append(f"正文仍含过时或过强口径：{phrase}")

    if re.search(r"(?<!\$)\$[^$]+\$(?!\$)", source):
        issues.append("正文存在行内LaTeX，DOCX渲染可能显示源码")

    # 3. 图、表、公式与引用。
    figure_refs = re.findall(r"!\[(图\s*\d+-\d+[^\]]*)\]\(([^)]+)\)", source)
    expected_figure_numbers = ["1-1", "2-1", "3-1", "4-1", "4-2", "5-1"]
    actual_figure_numbers = [re.search(r"\d+-\d+", caption).group(0) for caption, _ in figure_refs]
    if actual_figure_numbers != expected_figure_numbers:
        issues.append(f"图编号异常：{actual_figure_numbers}")
    for _, relative in figure_refs:
        if not (markdown.parent / relative).exists():
            issues.append(f"图件不存在：{relative}")
    for number in expected_figure_numbers:
        if len(re.findall(rf"图\s*{re.escape(number)}", source)) < 2:
            issues.append(f"图{number}未被正文明确引用")

    captions = re.findall(r"^\*\*表\s*(\d+-\d+)", source, re.M)
    expected_tables = ["3-1", "3-2", "4-1", "4-2", "5-1", "5-2", "6-1", "6-2"]
    if captions != expected_tables:
        issues.append(f"表编号异常：{captions}")
    for number in expected_tables:
        if len(re.findall(rf"表\s*{re.escape(number)}", source)) < 2:
            issues.append(f"表{number}未被正文明确引用")

    equation_numbers = re.findall(r"^（(\d+-\d+)）$", source, re.M)
    expected_equations = (
        [f"2-{i}" for i in range(1, 6)]
        + [f"3-{i}" for i in range(1, 5)]
        + [f"4-{i}" for i in range(1, 4)]
        + ["5-1"]
    )
    if equation_numbers != expected_equations:
        issues.append(f"公式编号异常：{equation_numbers}")

    body_before_references = source.split("# 参考文献", 1)[0]
    cited: set[int] = set()
    for content in re.findall(r"\[([0-9,\-]+)\]", body_before_references):
        for part in content.split(","):
            if "-" in part:
                lo, hi = map(int, part.split("-", 1))
                cited.update(range(lo, hi + 1))
            else:
                cited.add(int(part))
    if cited != set(range(1, 21)):
        issues.append(f"参考文献正文引用覆盖异常：{sorted(cited)}")

    # 4. 正式冻结结果数值追溯。
    formal_path = root / "实验/研究/results/runs/real-2021-2025-v32-frozen/formal_matrix_110kv.csv"
    formal = pd.read_csv(formal_path).set_index("region_id")
    checks = [
        ("第四章", "表4-1", "QX-00001无上限实际物理容载比", formal.loc["QX-00001", "PATH_OPT_CLR_UNBOUNDED_clr_2025"], "2.3597466"),
        ("第四章", "表4-1", "QX-00001控制值2.0实际物理容载比", formal.loc["QX-00001", "PATH_OPT_CLR_LE_2_clr_2025"], "2.2679276"),
        ("第四章", "表4-1", "QX-00001无上限累计在役等年成本", formal.loc["QX-00001", "PATH_OPT_CLR_UNBOUNDED_cumulative_eac"], "188.1800"),
        ("第四章", "表4-1", "QX-00001控制值2.0累计在役等年成本", formal.loc["QX-00001", "PATH_OPT_CLR_LE_2_cumulative_eac"], "481.9566"),
        ("第四章", "表4-1", "QX-00005无上限实际物理容载比", formal.loc["QX-00005", "PATH_OPT_CLR_UNBOUNDED_clr_2025"], "2.1715720"),
        ("第四章", "表4-1", "QX-00005控制值2.0实际物理容载比", formal.loc["QX-00005", "PATH_OPT_CLR_LE_2_clr_2025"], "2.1192953"),
        ("第四章", "表4-1", "QX-00005无上限累计在役等年成本", formal.loc["QX-00005", "PATH_OPT_CLR_UNBOUNDED_cumulative_eac"], "1745.5283"),
        ("第四章", "表4-1", "QX-00005控制值2.0累计在役等年成本", formal.loc["QX-00005", "PATH_OPT_CLR_LE_2_cumulative_eac"], "3359.0522"),
        ("第五章", "表5-2", "QX-00001推荐下限", formal.loc["QX-00001", "rcap_robust_near_optimal_lower"], "2.50"),
        ("第五章", "表5-2", "QX-00005推荐下限", formal.loc["QX-00005", "rcap_robust_near_optimal_lower"], "2.30"),
    ]
    trace_rows = []
    for chapter, item, metric, official, report_value in checks:
        passed = report_value in source
        if not passed:
            issues.append(f"报告缺少关键值：{metric}={report_value}")
        trace_rows.append([
            chapter, item, metric, report_value,
            f"110 kV正式冻结结果：{official}", "通过" if passed else "不通过",
        ])

    if "2.50～" in source or "2.30～" in source:
        issues.append("正文存在可能的人为推荐上限")

    # 5. 样本参考规则只做样本内一致性，不允许产生人工上限或缺值误判。
    backtest_path = root / "实验/研究/reports/report_reconstruction/matrix_backtest.csv"
    backtest = pd.read_csv(backtest_path)
    if len(backtest) != 8 or not backtest["explains_formal_optimization"].all():
        issues.append("样本参考规则未合理解释全部已知片区")
    if int(backtest["matches_formal_optimization"].sum()) != 7:
        issues.append("样本回代直接一致片区数量应为7，QX-00007应为数据例外")
    if backtest["artificial_upper_bound"].any():
        issues.append("参考矩阵存在人为上限")

    matrix_module_path = root / "实验/研究/src/report_interval_matrix.py"
    module_spec = importlib.util.spec_from_file_location("report_interval_matrix_check", matrix_module_path)
    matrix_module = importlib.util.module_from_spec(module_spec)
    assert module_spec.loader is not None
    module_spec.loader.exec_module(matrix_module)
    complete_indicators = [0.973, 0.088, 0.1136, 0.0622]
    for missing_index in range(4):
        values = complete_indicators.copy()
        values[missing_index] = None
        result = matrix_module.query_interval_matrix(
            technical_feasible=True,
            source_load_scale_ratio=values[0],
            local_reverse_flow_ratio=values[1],
            network_capacity_support_margin=values[2],
            positive_peak_cagr_2021_2025=values[3],
        )
        if result[1] != "D":
            issues.append(f"第{missing_index + 1}项核心指标缺失时未进入数据不足分支")

    # 6. Word/PDF成品检查。
    word = Document(docx)
    with ZipFile(docx) as archive:
        document_bytes = archive.read("word/document.xml")
        document_xml = document_bytes.decode("utf-8")
    xml_root = etree.fromstring(document_bytes)
    math_count = len(xml_root.findall(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"))
    if math_count != 13:
        issues.append(f"Word原生公式数量应为13，实际为{math_count}")
    if document_xml.count("<w:drawing") != 6:
        issues.append("Word嵌入图件数量不是6")
    if len(word.tables) != 8:
        issues.append(f"Word表格数量应为8，实际为{len(word.tables)}")
    if "\\frac" in document_xml or "\\Delta" in document_xml or "$S_" in document_xml:
        issues.append("Word存在未转换公式源码")
    for number in expected_equations:
        if f"（{number}）" not in document_xml:
            issues.append(f"Word缺少公式编号（{number}）")
    for table in word.tables:
        if len(table.rows) > 1 and table.rows[0]._tr.get_or_add_trPr().find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader") is None:
            issues.append("Word表格缺少重复表头设置")
            break

    info = subprocess.run(["pdfinfo", str(pdf)], check=True, capture_output=True, text=True).stdout
    match = re.search(r"Pages:\s+(\d+)", info)
    page_count = int(match.group(1)) if match else 0
    if page_count < 30:
        issues.append(f"PDF页数异常：{page_count}")

    # 低文本页不直接判为空白：若页面含图题，则属于合法的图件主导页。
    sparse_pages: list[int] = []
    visual_only_pages: list[int] = []
    visual_page_text: dict[str, str] = {}
    for page in range(1, page_count + 1):
        text = pdf_page_text(pdf, page)
        compact = re.sub(r"\s+", "", text)
        if len(compact) < 20:
            if re.search(r"图\s*[1-7]-\d+", text):
                visual_only_pages.append(page)
                visual_page_text[str(page)] = text
            else:
                sparse_pages.append(page)
    if sparse_pages:
        issues.append(f"PDF存在疑似真实空白页：{sparse_pages}")

    full_pdf_text = subprocess.run(
        ["pdftotext", str(pdf), "-"], check=True, capture_output=True, text=True
    ).stdout
    if re.search(r"\\(?:Delta|frac|mathrm|left|right)", full_pdf_text):
        issues.append("PDF存在未编译公式源码")

    output_dir.mkdir(parents=True, exist_ok=True)
    trace_path = output_dir / "数值追溯表.csv"
    with trace_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["报告章节", "表/图", "指标", "报告值", "正式结果来源", "校验结果"])
        writer.writerows(trace_rows)

    result = {
        "passed": not issues,
        "issues": issues,
        "chapter_count": len(chapter_titles),
        "section_count": len(section_titles),
        "figure_count": len(figure_refs),
        "table_count": len(word.tables),
        "native_math_count": math_count,
        "pdf_pages": page_count,
        "visual_only_pages": visual_only_pages,
        "visual_only_page_text": visual_page_text,
        "matrix_backtest_regions": len(backtest),
        "trace_rows": len(trace_rows),
    }
    (output_dir / "自动检查结果.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--markdown", type=Path, required=True)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    result = run_check(
        args.root.resolve(), args.markdown.resolve(), args.docx.resolve(),
        args.pdf.resolve(), args.output_dir.resolve()
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
