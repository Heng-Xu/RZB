"""按导则原定义，以2025区县原始资料估计源荷比和电量渗透率。"""

from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from .baseline_2021 import read_csv
from .city_mapping_audit import write_csv
from .pizhou_station_scenarios import read_pizhou_station_series
from .regional_revision_run import OUTPUT
from .official_annual import STUDY_DIR


ROOT = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2"
PV_MONTHLY = ROOT / "逐月分县分布式光伏.xlsx"
PV_HOURLY = ROOT / "光伏8760小时数据.xlsx"
EQUIPMENT = ROOT / "2025设备负载统计表.xlsx"
CITY_HOURLY = STUDY_DIR / "rebuild_2026/source_audit/city_2025_aggregate_hourly.csv"
REGIONS = {"QX-00005": "邳州", "QX-00007": "市区"}


def monthly_pv_mw() -> dict:
    workbook = load_workbook(PV_MONTHLY, read_only=True, data_only=True)
    result = {}
    year = None
    try:
        for cells in workbook.active.iter_rows(values_only=True):
            label = str(cells[1] or "")
            if label.startswith("202") and label.endswith("年"):
                year = int(label[:4])
            if cells[0] in REGIONS and year in (2024, 2025):
                result[year, cells[0]] = [float(value) * 10 for value in cells[1:13]]
    finally:
        workbook.close()
    return result


def pv_hourly_factors() -> dict[datetime, float]:
    workbook = load_workbook(PV_HOURLY, read_only=True, data_only=True)
    try:
        return {row[1]: float(row[4]) for row in workbook.active.iter_rows(min_row=2, values_only=True)
                if isinstance(row[1], datetime) and row[1].year == 2025}
    finally:
        workbook.close()


def equipment_scenes() -> dict:
    workbook = load_workbook(EQUIPMENT, read_only=True, data_only=True)
    result = defaultdict(lambda: {"day": 0.0, "night": 0.0})
    try:
        for row in workbook["主变1"].iter_rows(min_row=3, values_only=True):
            if row[1] != 110 or row[2] not in REGIONS:
                continue
            result[row[2]]["day"] += float(row[6] or 0)
            result[row[2]]["night"] += float(row[8] or 0)
    finally:
        workbook.close()
    return result


def pizhou_net_energy_mwh() -> tuple[float, int, float]:
    sites = [values for (voltage, _), values in read_pizhou_station_series().items() if voltage == 110]
    common = set.intersection(*(set(series) for series in sites))
    observed = {hour: sum(series[hour] for series in sites) for hour in common}
    missing = [datetime(2025, 1, 1) + timedelta(hours=index) for index in range(8760)
               if datetime(2025, 1, 1) + timedelta(hours=index) not in observed]
    for hour in missing:
        observed[hour] = (observed[hour - timedelta(hours=1)] +
                          observed[hour + timedelta(hours=1)]) / 2
    return sum(observed.values()), len(missing), max(observed.values())


def city_net_energy_mwh() -> tuple[float, int, float]:
    rows = read_csv(CITY_HOURLY)
    values = [float(row["observed_mw"] or row["weekly_mean_mw"]) for row in rows]
    return sum(values), sum(row["data_status"] != "observed" for row in rows), max(values)


def run(output_dir: Path = OUTPUT) -> list[dict]:
    months = monthly_pv_mw()
    factors = pv_hourly_factors()
    scenes = equipment_scenes()
    series = {"QX-00005": pizhou_net_energy_mwh(),
              "QX-00007": city_net_energy_mwh()}
    rows = []
    for region, name in REGIONS.items():
        year_end_installed = months[2025, region][-1]
        august_installed = months[2025, region][7]
        noon_factor = factors[datetime(2025, 8, 20, 12)] * (1 / 3) + factors[datetime(2025, 8, 20, 13)] * (2 / 3)
        user_day = scenes[region]["day"] + noon_factor * august_installed
        user_night = scenes[region]["night"]
        user_max = max(user_day, user_night)
        generation = 0.0
        for hour, factor in factors.items():
            current_end = months[2025, region][hour.month - 1]
            prior_end = (months[2024, region][-1] if hour.month == 1 else
                         months[2025, region][hour.month - 2])
            generation += factor * (current_end + prior_end) / 2
        sample_net_energy, missing_hours, sample_peak = series[region]
        # 样本时序按官方县域2025年110 kV同步正向峰覆盖率外推；电量仍为估计。
        official_peak = scenes[region]["night"]
        coverage_factor = official_peak / sample_peak
        county_net_energy = sample_net_energy * coverage_factor
        user_energy = county_net_energy + generation
        rows.append({"区县": name, "区县代号": region, "年份": 2025,
                     "分布式电源装机渗透率（源荷比）": round(year_end_installed / user_max, 9),
                     "分布式电源电量渗透率": round(generation / user_energy, 9),
                     "分布式光伏年末装机（MW）": round(year_end_installed, 6),
                     "最大用电负荷典型场景估计（MW）": round(user_max, 6),
                     "夜间用户负荷候选（MW）": round(user_night, 6),
                     "日间用户负荷候选（MW）": round(user_day, 6),
                     "分布式光伏年发电量估计（MWh）": round(generation, 3),
                     "用户年用电量估计（MWh）": round(user_energy, 3),
                     "年度110kV样本净电量（MWh）": round(sample_net_energy, 3),
                     "样本覆盖外推系数": round(coverage_factor, 6),
                     "样本时序补值小时": missing_hours,
                     "统计状态": "按导则公式计算；最大用户负荷、发电量和用电量均为典型场景/时序估计",
                     "数据边界": "区县光伏全量+110kV站点净负荷样本外推；跨电压及其他分布式电源待核"})
    write_csv(rows, Path(output_dir) / "2025区县导则指标估计.csv")
    return rows


if __name__ == "__main__":
    for row in run():
        print(row["区县"], row["分布式电源装机渗透率（源荷比）"],
              row["分布式电源电量渗透率"])
