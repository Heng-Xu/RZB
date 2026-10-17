"""固定已有设备，对转供证据及运行自由度作逐项消融。"""
import copy
import json
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from ..static_joint_20261015 import model as base
from ..forecast_2035_20261017 import run as future

ROOT=base.ROOT
PUBLIC=ROOT/'研究报告/01研究推荐/当前版/转供机制审查_20261017'
SOURCE=ROOT/'研究报告/数据来源/当前版/转供机制审查_20261017'
REVIEW=ROOT/'研究报告/05_review/当前版/1017转供机制审查'
TRANSFER=ROOT/'实验/研究/data/tuomin/补充数据/徐州地区变电站全停全转情况分析汇总20250705.xlsx'
MODES={
    'original':'原运行自由度',
    'fault_only':'既有通道仅事故使用',
    'direct_fault':'既有仅事故＋事故仅直接邻接',
    'confirmed_only':'仅有路径依据的既有边＋事故直接邻接（资料下界对照）',
}


class RestrictedLP(base.OperatingLP):
    def bounds_for(self,prepared,failed):
        common,bounds,full,largest,bp,bn=super().bounds_for(prepared,failed)
        mode=self.inp.get('audit_mode','original')
        if mode!='original':
            bounds=[(0.,0.) if (failed<0 and arc[1]['existing']) or
                    (failed>=0 and mode in ['direct_fault','confirmed_only'] and failed not in (arc[5],arc[6]))
                    else bound for arc,bound in zip(self.arcs,bounds)]
        return common,bounds,full,largest,bp,bn

    def check(self,p,prepared,failed,with_deficit=True):
        if self.arcs:return super().check(p,prepared,failed,with_deficit)
        common,bounds,full,largest,bp,bn=self.bounds_for(prepared,failed)
        available=full.copy()
        if failed>=0:available[failed]-=largest[failed]
        deficit=np.maximum(np.maximum(p-available-bp,-p-available-bn),0.)
        return dict(verified=bool(deficit.max()<1e-7),deficit_mw=float(deficit.sum()),direct=True,
                    station_deficits={sid:float(v) for sid,v in zip(self.ids,deficit) if v>1e-7})


def source_audit():
    book=load_workbook(TRANSFER,read_only=True,data_only=True);lookup={}
    for nr,r in enumerate(book.worksheets[0].iter_rows(min_row=3,values_only=True),3):
        if r[1]=='110kV' and r[2]:lookup[r[2]]=dict(row=nr,feeder_count=r[4],external_tied_feeders=r[5],
            untied_or_internal=r[7],constraint_reason=r[8],J=r[9],K=r[10],unable_details=r[13])
    book.close();overview=[];station_rows=[]
    for region in ['city','pizhou']:
        inp=json.loads((base.RAW/region/'annual_2021/inputs.json').read_text())
        existing=[e for e in inp['network']['edges'] if e['existing']]
        synthetic=[e for e in existing if e['id'].startswith('EQ-')]
        neighbors=defaultdict(set);degree=Counter()
        for e in existing:
            for a,b in [(e['station_a'],e['station_b']),(e['station_b'],e['station_a'])]:
                neighbors[a].add(b);degree[a]+=1
        overview.append(dict(地区=base.REGIONS[region],研究站数=len(inp['stations']),既有边数=len(existing),
            其中标定假设边数=len(synthetic),标定站对数=len({tuple(sorted((e['station_a'],e['station_b']))) for e in synthetic}),
            原路径依据边数=len(existing)-len(synthetic),每站最多模型邻站数=max(map(len,neighbors.values())),
            每站最多标定或既有边数=max(degree.values()),
            原表转供率不低于百分之99站数=sum((v or 0.)>=.99 for v in inp['station_transfer_rates'].values()),
            逐站既有预算合计_MW=sum(inp['existing_transfer_budget_mw'].values())))
        for e in inp['station_transfer_evidence']:
            sid=e['station'];r=lookup.get(e['name'],{})
            feeders=[v for v in inp['network']['feeders'].values() if v['station']==sid and v.get('equivalent_existing')]
            station_rows.append(dict(地区=base.REGIONS[region],站号=sid,站名=e['name'],原表行=r.get('row',''),
                原表10kV线路条数=r.get('feeder_count',''),原表有站外联络线路条数=r.get('external_tied_feeders',''),
                原表可转供_MW=e['transferable_mw'],原表损失_MW=e['unserved_mw'],原表转供率=e['source_fraction'],
                模型邻站数=len(neighbors[sid]),模型既有边数=degree[sid],模型通用馈线数=len(feeders),
                模型通用馈线总热限_MW=sum(v['thermal_mw'] for v in feeders),
                原表不能全转原因=r.get('constraint_reason','') or '',原表内部联络或无联络情况=r.get('untied_or_internal','') or ''))
    base.old.write_csv(SOURCE/'既有网络证据审查.csv',overview)
    base.old.write_csv(SOURCE/'逐站原表与模型映射.csv',station_rows)
    return overview


def instances():
    for region in ['city','pizhou']:
        inp,scenes,_=base.prepare(region,'annual_2021')
        for case in ['rigid','elastic_free']:
            folder=base.RAW/region/'annual_2021'/('rigid_certified' if case=='rigid' else ('elastic_engineering' if region=='pizhou' else 'elastic_free_certified'))
            u,n,built=base.old.old.assets_from_csv(inp,folder,2025)
            yield region,case,2025,inp,scenes[2025],u,n,built,folder
    inp=json.loads((future.RAW/'inputs.json').read_text());scenes=np.load(future.RAW/'scenarios.npz')['2035']
    for case in ['rigid','elastic_free']:
        folder=future.RAW/(case+'_rolling');u,n,built=base.old.old.assets_from_csv(inp,folder,2035)
        yield 'city',case,2035,inp,scenes,u,n,built,folder


def run():
    for p in [PUBLIC,SOURCE,REVIEW]:p.mkdir(parents=True,exist_ok=True)
    overview=source_audit();summary=[];worst=[];sources=[]
    for region,case,y,inp,scenes,u,n,built,folder in instances():
        ids=list(inp['stations'])
        sources.append(dict(region=region,case=case,year=y,folder=str(folder.relative_to(ROOT)),
            stations_sha256=base.old.sha(folder/'stations.csv'),projects_sha256=base.old.sha(folder/'projects.csv')))
        for mode,label in MODES.items():
            active=copy.deepcopy(inp);active['audit_mode']=mode
            active['network']['edges']=[e for e in active['network']['edges'] if e['id'] in built or
                (e['existing'] and (mode!='confirmed_only' or not e['id'].startswith('EQ-')))]
            used={e[k] for e in active['network']['edges'] for k in ['a','b']}
            active['network']['feeders']={f:r for f,r in active['network']['feeders'].items() if f in used}
            lp=RestrictedLP(active);failures=[];counts=Counter();maximum=defaultdict(float)
            original_dispatch=[]
            for h,p in enumerate(scenes):
                ready=lp.prepared(p,y,u,n,built)
                for f in range(-1,len(ids)):
                    result=lp.check(p,ready,f);kind='正常' if f<0 else 'N-1';counts[kind]+=1
                    if not result['verified']:
                        maximum[kind]=max(maximum[kind],result['deficit_mw'])
                        failures.append(dict(地区=base.REGIONS[region],方案=case,年份=y,规则=label,截面=h,
                            工况=kind,故障站=ids[f] if f>=0 else '',约束松弛_MW=result['deficit_mw'],
                            缺口站=json.dumps(result.get('station_deficits',{}),ensure_ascii=False)))
            summary.append(dict(地区=base.REGIONS[region],方案=case,年份=y,规则=label,年末容量_MVA=float(u.sum()),
                储能_MWh=float(n.sum()*.215),累计新增联络_条=len(built),正常校核数=counts['正常'],
                正常失败数=sum(r['工况']=='正常' for r in failures),正常最大约束松弛_MW=maximum['正常'],
                N1校核数=counts['N-1'],N1失败数=sum(r['工况']=='N-1' for r in failures),N1最大约束松弛_MW=maximum['N-1']))
            worst.extend(sorted(failures,key=lambda r:r['约束松弛_MW'],reverse=True)[:10])
            print('TRANSFER_ABLATION',region,case,y,mode,'bad',len(failures),dict(maximum),flush=True)
    base.old.write_csv(PUBLIC/'固定设备转供规则对照.csv',summary)
    base.old.write_csv(REVIEW/'最差工况.csv',worst)
    base.old.dump(REVIEW/'机制审查.json',dict(verified=True,source_file=str(TRANSFER.relative_to(ROOT)),
        source_sha256=base.old.sha(TRANSFER),source_audit=overview,asset_sources=sources,results=summary,
        meaning='固定已交付设备后的运行自由度诊断；约束松弛为不满足程度，不是统计失负荷；仅已核实路径为资料下界对照，不声称其余实际通道不存在'))


if __name__=='__main__':run()
