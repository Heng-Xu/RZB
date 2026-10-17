"""生成转供机制审查说明和可检查的证据工作簿。"""
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from . import audit as a
from ..case_parameters_20261012.report import add_sheet, markdown


def main():
    raw=json.loads((a.REVIEW/'机制审查.json').read_text());assert raw['verified']
    probes=[]
    cases=[(region,case) for region in ['city','pizhou'] for case in ['rigid','elastic_free']]
    if (a.REVIEW/'单年机制筛查/pizhou/elastic_engineering/summary.json').exists():cases.append(('pizhou','elastic_engineering'))
    for region,case in cases:
            folder=a.REVIEW/'单年机制筛查'/region/case
            s=json.loads((folder/'summary.json').read_text());annual=s['annual'][0]
            proof=json.loads((folder/'independent_lp.json').read_text())
            assert proof['verified'] and all(r['verified'] for r in proof['records'])
            probes.append(dict(地区=a.base.REGIONS[region],方案={'rigid':'刚性','elastic_free':'弹性费用参考','elastic_engineering':'弹性工程优选'}[case],年份=2025,
                共同下限=s['common_floor'],容量_MVA=annual['capacity_mva'],容载比=annual['clr'],储能_MWh=annual['storage_mwh'],
                新增联络_条=annual['cumulative_lines'],费用现值_万元=s['cost_npv_10k'],相对间隙=s['solver']['mip_gap'],
                完整静态工况数=len(proof['records']),范围='2021设备起点的2025单年机制筛查，非逐年方案'))
    recommendations=[
        dict(项目='既有通道证据',现状='以全停全转总量标定168/76条假设边，并作为免费双向连接',调整='保留真实已建资产零新增投资；J、K只作既有转供能力总上界，逐站连接与可切换区段须另有证据，未闭合部分采用明示代理及敏感性',依据='原收资2025-07-05全停全转表；修订稿9.2.1.2',性质='数据映射修正'),
        dict(项目='正常供区',现状='各截面独立重配，既有事故能力同时承担日常优化',调整='既有免费能力默认用于事故备用；有明确依据的正常转供须登记年内固定运行方案，新建项目可形成有费用的固定供区调整',依据='修订稿5.2.1.2分区相对稳定、9.1.1逐年过渡；正常既有禁用是本次保守运行策略',性质='标准指导＋工程运行假设'),
        dict(项目='事故转供',现状='每个故障工况允许全地区健康站继续连锁重配',调整='保留已批准的正常方案，事故追加切换限于故障供区所在接线组及其可达相邻区段，健康供区不再全局重排',依据='修订稿9.2.1.1标准接线组、9.2.1.2实际容量及线路余量',性质='工程运行策略，非标准逐字强制'),
        dict(项目='区段切换',现状='站级连续MW可任意选择，光伏与负荷没有共同切换决策',调整='以明确线路区段为单元，同一次切换同时转移该区段负荷及光伏，全年截面共用固定分段和允许接收站；一段只能由一处供电',依据='修订稿9.2.1.1 b—e项',性质='建模结构修正；架空分段参数不可机械用于电缆'),
        dict(项目='共享与余量',现状='已有馈线和接收主变约束，但对象含标定的通用馈线',调整='逐路径绑定真实或明示代理馈线，约束源端区段、路径最弱热限、受端馈线余量、接收主变和本组切换能力，同一资源只计一次',依据='修订稿9.2.1.2、DL/T2041—2025 6.4',性质='标准指导＋路径建模'),
        dict(项目='两组公平比较',现状='上一轮刚性有分档下限，弹性取消下限',调整='同地区同年度使用相同容载比下限、源荷与运行约束；只在明确的刚性新增容量控制及弹性放宽条件上区别',依据='用户最新指示；修订稿6.2.2.1',性质='方案比较设定'),
        dict(项目='结果筛选',现状='仅费用寻优可能采用不同容量与联络替代',调整='保留费用参考，另外筛选增容替代候选：较高容量、较少储能及联络；若设置显式偏好约束或费用容差，单独登记',依据='用户最新结果要求',性质='工程优选目标，不称标准必然结论'),
    ]
    frames={'网络证据':raw['source_audit'],'固定设备规则对照':raw['results'],'同下限单年筛查':probes,
        '逐站原表与模型':a.base.old.old.base.previous.rows(a.SOURCE/'逐站原表与模型映射.csv'),'约束调整方向':recommendations}
    book=Workbook();book.remove(book.active)
    for name,rows in frames.items():
        a.base.old.write_csv(a.PUBLIC/(name+'.csv'),rows);add_sheet(book,name,rows)
        for row in book[name]:
            for cell in row:cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    path=a.PUBLIC/'转供机制与方案方向核对.xlsx';book.save(path)
    w=load_workbook(path,data_only=True);cells=0
    for name,rows in frames.items():
        actual=list(w[name].values);expected=[list(rows[0])]+[list(r.values()) for r in rows]
        assert len(actual)==len(expected)
        for aa,bb in zip(actual,expected):
            for x,y in zip(aa,bb):
                if isinstance(y,float):assert abs(x-y)<max(1e-8,abs(y)*1e-12)
                else:assert (x if x is not None else '')==(y if y is not None else '')
                cells+=1
    w.close()
    memo=a.PUBLIC/'转供调整判断.md'
    memo.write_text('\n'.join([
        '# 转供调整判断：面向两区县2021—2035逐年重算','',
        '本轮要达到的工程比选目标是：在相同容载比下限和供电安全条件下，以主变增容替代部分新增储能和10 kV联络。先核对转供机制，再做两地区全部年度重算。当前成果为机制诊断和2025单年筛查，尚未形成新的2021—2035全套逐年方案。','',
        '## 核心判断','',
        '既有转供能力具有实物基础，使用既有资产可以不计新增投资。现模型的主要问题是：把原表全停全转的逐站总量映射成可免费使用的站对网络，并允许它在正常及每个故障截面上独立调整供区，形成超出已核实线路与运行方式证据的优化自由度。',
        '市区168条既有模型边全部为标定假设，覆盖84个站对；邳州77条中76条为标定假设，另1条T01有局部路径依据。市区模型单站最多可连接12座邻站，邳州最多8座。这些站级计数不直接等同于单个标准接线组的电源数量，不能据此直接判定违反接线组条款，但说明当前尚未以接线组为基本校核单元。',
        '原表还提供10kV线路数、有站外联络的线路数、同站联络与专线情况、不能全转原因。原标定主要使用J、K总量，未将这些内容闭合为各接收站和具体线路区段的连接约束。','',
        '## 固定设备对照','',
        '保持上一版市区2035年两套设备不动，只关闭正常方式下的既有免费转供，原49个正常截面中均有13个不再通过。刚性最大约束松弛93.64 MW，弹性80.02 MW；这是约束不满足程度，不能当作预计失负荷。再禁止故障时健康站之间追加转接，N-1分别有377、472个工况不通过。',
        '这组对照表明上一版明显依赖正常免费转供及全区故障后重配。事故直接邻接是较保守的诊断规则，完整新模型应保留已批准正常方案，仅限制事故追加切换，避免把正常有费用的供区调整一并取消。',
        '邳州2025年的原方案在禁用正常免费转供、事故仅直接邻接后仍可行，说明两地区的响应不同；进一步只保留有明确路径依据的既有边时，原刚性、弹性分别有81、79个N-1工况不通过。仅确认路径的对照是资料覆盖条件下的保守边界，不能推断其余真实联络不存在。','',
        '## 同下限的单年机制筛查','',
        '从同一2021设备组合出发，只计算2025单年措施，两组下限一致：市区1.50、邳州1.70。正常禁止使用事故汇总授予的既有免费通道，付费新增项目仍可参与正常工况，事故只用故障站直接邻接通道。保留同一正常及N-1有效储能、逐站既有预算、馈线共享余量与主变接收限值。','',
        markdown(['地区','方案','容量/MVA','容载比','储能/MWh','新增联络/条'],[
            [r['地区'],r['方案'],f'{r["容量_MVA"]:.2f}',f'{r["容载比"]:.2f}',f'{r["储能_MWh"]+1e-10:.2f}',r['新增联络_条']] for r in probes]),'',
        '费用参考的四组单年方案均通过全部有限静态工况。市区弹性增加27.00 MVA并替代126.64 MWh储能；邳州弹性增加121.50 MVA、储能由201.24降至4.09 MWh。仅以费用为目标时，新增联络分别同为0条、4条。',
        '针对新增联络的方案选择，再在邳州弹性费用参考5%费用容差以内优先减少联络，取得2148.50 MVA、7.74 MWh、3条联络的可行组合，费用6171.93万元，较费用参考5926.91万元增加4.13%。相对刚性2027.00 MVA、201.24 MWh、4条联络，三项方向同时达到用户目标。该5%为工程优选条件，沿用既有比选方式，不称标准强制或纯费用最优。',
        '市区本次两方案新增联络均为0条，容量和储能已有响应。零新增联络不能再减少，原始2021配置相同的年度也应如实相同，不添加无必要工程制造差额。单年筛查尚不能保证2035年或所有年度也有相同排序，仍须完整逐年重算。','',
        '## 拟采用的转供约束','',
        '1. 逐站J、K作为既有事故转供总量上界，既有连接须落到有依据的接线组、相邻供区和区段。资料未闭合部分采用明确的代理结构与敏感性范围，原来按地区平均转供率放大免费网络的标定不再作为实物连接证据。',
        '2. 正常方式采用年内固定供区和登记的切换方案。新建联络可以形成永久供区调整并计工程费用。全停全转收资授予的事故备用能力，不再自动成为日常源荷互补能力。',
        '3. 故障方式以上述正常方案为基线，只允许故障所在接线组及有可达路径的相邻区段追加切换；健康接线组不全区连锁重排。',
        '4. 以可切换线路区段为单位分配供电来源，同一切换动作同时转移区段负荷和光伏。各截面使用相同的区段范围与允许连接，一段负荷只归属一个供电来源。',
        '5. 转供能力同时受源端可切区段、路径最弱热限、受端馈线余量、接收主变能力和故障可恢复范围约束，馈线与接收侧余量按同时占用计算。',
        '6. 两方案同地区同年度采用相同容载比下限和相同网络运行规则；保留费用参考，按用户增容替代目标另选工程候选。标准条款用于确定安全和结构边界，方案偏好单独登记。','',
        '## 标准原文与适用层级','',
        '2025年4月《配电网规划设计技术导则（修订稿）》PDF第45页第9.2.1.1条：架空线路联络、接线组电源和分段有具体建议；架空分段最大2000 kW、组电源不宜超过3座变电站等要求须施加于其对应对象，不能改成每座站仅允许两个邻站，也不能机械用于电缆网。PDF第46页第9.2.1.2条要求结合变压器和线路余量、分段及联络确定转供能力，并原则上不设置站间中压专用联络线。',
        '正常免费转供禁用、事故限制追加切换是本轮保守工程运行设定，标准没有给出“统一把全站转供率打五折”或“弹性容量必然更高”的规定。DL/T 2041—2025第6.4条区分主变分列与并列方式的反向负载率，不提供任意站间转供信用。',
        '标准原页图像、固定设备运行对照、单年原始解及输入哈希均保存在对应来源和核验目录。','']))
    proof=dict(verified=True,artifact_kind='model_mechanism_review',schemes=probes,workbook_cells=cells,static_pairs_fixed_assets=sum(r['正常校核数']+r['N1校核数'] for r in raw['results']),
        static_pairs_probes=sum(r['完整静态工况数'] for r in probes),probe_results=probes,
        limitation='机制诊断及2025单年筛查；未完成2021—2035逐年新方案；邳州工程优选已达三项方向，市区新增联络均为零',
        solution_sources=[dict(folder=str((a.REVIEW/'单年机制筛查'/region/case).relative_to(a.ROOT)),
            files={name:a.base.old.sha(a.REVIEW/'单年机制筛查'/region/case/name) for name in ['effective_inputs.json','summary.json','stations.csv','projects.csv','selected_pairs.json','chosen_solution.npz','independent_lp.json']})
            for region,case in cases],
        files=[dict(path=str(p.relative_to(a.ROOT)),sha256=a.base.old.sha(p)) for p in [path,memo,a.REVIEW/'机制审查.json']])
    a.base.old.dump(a.REVIEW/'方向核验.json',proof)
    print('TRANSFER_DIRECTION_VERIFIED',cells,proof['static_pairs_probes'],flush=True)


if __name__=='__main__':main()
