"""核验补充屋顶面积，并形成区县资源情景；不推定电网可接入容量。"""

import argparse
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .city_mapping_audit import write_csv
from .hourly_source_profile import OUTPUT_DIR


SOURCE = Path(__file__).resolve().parents[1] / "data/tuomin/补充数据/屋顶数据.xlsx"
URBAN_CORE_DISTRICTS = ("鼓楼区", "云龙区", "泉山区")
AREA_GROUPS = {
    "party_government": (5, 8),
    "public_building": (8, 11),
    "industrial_park": (11, 14),
    "other_commercial": (14, 17),
    "rural_residential": (17, 20),
    "urban_residential": (21, 24),
    "other": (24, 27),
}
KW_PER_USABLE_M2 = 0.10  # 徐州地铁招标案例：94,900 m² / 9.49 MWp。
USABLE_AREA_FRACTIONS = (0.20, 0.40, 0.60)  # 仅敏感性；0.20 为江苏 2009 年规划估算，非上限。


def _area(value: object, row_number: int, column_number: int) -> float:
    if value is None:
        return 0.0
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0:
        raise ValueError(f"屋顶表第 {row_number} 行第 {column_number} 列面积无效：{value!r}")
    return float(value)


def read_roof_records(source: Path = SOURCE) -> list[dict]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["输入-屋顶面积"]
        records = []
        seen_ids = set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=3, max_col=27, values_only=True), 3):
            city, district, town, village, area_id = (str(value).strip() if value is not None else "" for value in cells[:5])
            if not all((city, district, town, village, area_id)):
                raise ValueError(f"屋顶表第 {row_number} 行区域键缺失")
            if area_id in seen_ids:
                raise ValueError(f"屋顶表区域编号重复：{area_id}")
            seen_ids.add(area_id)
            groups = {
                name: sum(_area(cells[index], row_number, index + 1) for index in range(start, end))
                for name, (start, end) in AREA_GROUPS.items()
            }
            rural_subtotal = cells[20]
            if rural_subtotal is not None and abs(_area(rural_subtotal, row_number, 21) - groups["rural_residential"]) > 0.01:
                raise ValueError(f"屋顶表第 {row_number} 行农村居民小计不一致")
            records.append(
                {
                    "city": city,
                    "district": district,
                    "town": town,
                    "village": village,
                    "area_id": area_id,
                    **groups,
                    "total_m2": sum(groups.values()),
                }
            )
        return records
    finally:
        workbook.close()


def scope_areas(records: list[dict]) -> dict[str, float]:
    return {
        "QX-00005": sum(row["total_m2"] for row in records if row["city"] == "徐州市" and row["district"] == "邳州市"),
        "urban_core_proxy": sum(
            row["total_m2"] for row in records
            if row["city"] == "徐州市" and row["district"] in URBAN_CORE_DISTRICTS
        ),
    }


def capacity_scenario_mw(gross_area_m2: float, usable_fraction: float) -> float:
    return gross_area_m2 * usable_fraction * KW_PER_USABLE_M2 / 1000


def district_summaries(records: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for row in records:
        groups[(row["city"], row["district"])].append(row)
    result = []
    for (city, district), members in sorted(groups.items()):
        area = sum(row["total_m2"] for row in members)
        result.append(
            {
                "city": city,
                "district": district,
                "area_ids": len(members),
                "gross_roof_area_m2": round(area, 3),
                "party_government_m2": round(sum(row["party_government"] for row in members), 3),
                "public_building_m2": round(sum(row["public_building"] for row in members), 3),
                "industrial_park_m2": round(sum(row["industrial_park"] for row in members), 3),
                "other_commercial_m2": round(sum(row["other_commercial"] for row in members), 3),
                "rural_residential_m2": round(sum(row["rural_residential"] for row in members), 3),
                "urban_residential_m2": round(sum(row["urban_residential"] for row in members), 3),
                "other_m2": round(sum(row["other"] for row in members), 3),
            }
        )
    return result


def pizhou_village_match(records: list[dict], source: Path = SOURCE) -> dict:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        transformer_keys = {
            (str(cells[2]).strip(), str(cells[3]).strip())
            for cells in workbook["台区数据数据"].iter_rows(min_row=2, max_col=4, values_only=True)
            if str(cells[0]).strip() == "邳州市"
        }
    finally:
        workbook.close()
    roof = [row for row in records if row["city"] == "徐州市" and row["district"] == "邳州市"]
    matched = [row for row in roof if (row["town"], row["village"]) in transformer_keys]
    return {
        "roof_area_ids": len(roof),
        "matched_area_ids": len(matched),
        "matched_gross_area_m2": round(sum(row["total_m2"] for row in matched), 3),
        "unmatched_area_ids": len(roof) - len(matched),
        "unmatched_gross_area_m2": round(sum(row["total_m2"] for row in roof) - sum(row["total_m2"] for row in matched), 3),
    }


def sensitivity_rows(records: list[dict]) -> list[dict]:
    areas = scope_areas(records)
    rows = []
    for scope, area in areas.items():
        for fraction in USABLE_AREA_FRACTIONS:
            rows.append(
                {
                    "scope": scope,
                    "scope_definition": "邳州市行政区，110/35 kV 未分层" if scope == "QX-00005" else "鼓楼/云龙/泉山行政区代理，非29站供电边界",
                    "gross_roof_area_m2": round(area, 3),
                    "assumed_usable_area_fraction": fraction,
                    "kw_per_usable_m2": KW_PER_USABLE_M2,
                    "rooftop_nameplate_scenario_mw": round(capacity_scenario_mw(area, fraction), 3),
                    "status": "resource_sensitivity_not_installed_pv_or_grid_hosting_limit",
                }
            )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="核验社区屋顶面积和台区覆盖")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    records = read_roof_records()
    write_csv(district_summaries(records), args.output_dir / "rooftop_area_district_audit.csv")
    write_csv(sensitivity_rows(records), args.output_dir / "rooftop_resource_sensitivity.csv")
    match = pizhou_village_match(records)
    print("屋顶区域记录", len(records), "邳州镇村匹配", match["matched_area_ids"], "/", match["roof_area_ids"])
    for scope, area in scope_areas(records).items():
        print(scope, "屋顶面积 m²", round(area, 3))


if __name__ == "__main__":
    main()
