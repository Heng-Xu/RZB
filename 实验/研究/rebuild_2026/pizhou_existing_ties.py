"""六馈线既有跨站联络的 2025 年受端电流余量静态筛查。"""

import argparse
from collections import defaultdict
from hashlib import sha256
from math import sqrt
from pathlib import Path

from openpyxl import load_workbook

from .baseline_2021 import POWER_FACTOR, read_csv
from .city_mapping_audit import write_csv
from .feeder_2025 import OUTPUT_DIR, read_feeders
from .official_annual import STUDY_DIR


TOPOLOGY_SOURCE = STUDY_DIR / "data/tuomin/10kv_case/source/10kV案例建模前期准备_V2_新结构化表格与图纸.xlsx"
TRANSFORMER_PATH = OUTPUT_DIR / "transformer_only_path_station_2021_2025.csv"
FEEDER_ALIAS_IDS = {
    "dnan": "PZXL-00092",
    "dxi": "PZXL-00097",
    "dzhen": "PZXL-00099",
    "hdong": "PZXL-00154",
    "hpao": "PZXL-00161",
    "hzhen": "PZXL-00173",
}
NOMINAL_LINE_KV = 10.0


def active_mw_at_current(current_a: float) -> float:
    return sqrt(3) * NOMINAL_LINE_KV * current_a * POWER_FACTOR / 1000


def read_cross_station_ties(source: Path = TOPOLOGY_SOURCE) -> list[dict]:
    source = Path(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows = []
    try:
        for source_row, cells in enumerate(workbook["02_联络关系"].iter_rows(min_row=2, values_only=True), start=2):
            if cells[2] != "跨站":
                continue
            rows.append({
                "tie_id": cells[0],
                "from_feeder_id": FEEDER_ALIAS_IDS[cells[4]],
                "to_feeder_id": FEEDER_ALIAS_IDS[cells[8]],
                "from_endpoint": cells[6],
                "to_endpoint": cells[10],
                "drawing_evidence": cells[11],
                "source_normal_state": cells[12],
                "topology_note": cells[14],
                "source_file": str(source.resolve().relative_to(STUDY_DIR)),
                "source_sheet": "02_联络关系",
                "source_row": source_row,
                "source_sha256": source_hash,
            })
    finally:
        workbook.close()
    return rows


def build_feeder_headroom() -> list[dict]:
    rows = []
    for feeder in read_feeders():
        allowed = float(feeder["max_allowed_current_a"])
        observed = float(feeder["reported_2025_max_current_a"])
        raw_power = float(feeder["reported_2025_max_active_power_mw"])
        current_power = active_mw_at_current(observed)
        rows.append({
            "feeder_id": feeder["feeder_id"],
            "station_id": feeder["station_id"],
            "year": 2025,
            "allowed_current_a": allowed,
            "observed_max_current_a": observed,
            "unused_current_a": round(max(allowed - observed, 0), 6),
            "assumed_power_factor": POWER_FACTOR,
            "reported_max_active_power_mw": raw_power,
            "current_equivalent_active_power_mw": round(current_power, 6),
            "source_power_status": "zero_active_power_conflicts_with_current" if raw_power == 0 and observed > 0 else "reported_nonzero",
            "receiving_current_headroom_mw": round(active_mw_at_current(max(allowed - observed, 0)), 6),
            "calculation_scope": "2025_source_feeder_current_screen_not_downstream_path_rating",
            "source_file": feeder["source_file"],
            "source_row": feeder["source_row"],
            "source_sha256": feeder["source_sha256"],
        })
    return rows


def build_existing_tie_screen(
    feeders: list[dict],
    ties: list[dict],
    station_path: Path = TRANSFORMER_PATH,
) -> list[dict]:
    by_feeder = {row["feeder_id"]: row for row in feeders}
    stations = {
        row["model_station_id"]: row
        for row in read_csv(station_path)
        if row["study_region_id"] == "QX-00005"
        and int(row["voltage_kv"]) == 110
        and int(row["year"]) == 2025
        and row["reverse_variant"] == "night_central"
    }
    rows = []
    for tie in ties:
        for source_feeder_id, receiver_feeder_id in (
            (tie["from_feeder_id"], tie["to_feeder_id"]),
            (tie["to_feeder_id"], tie["from_feeder_id"]),
        ):
            source = by_feeder[source_feeder_id]
            receiver = by_feeder[receiver_feeder_id]
            station = stations[receiver["station_id"]]
            capacity = float(station["selected_capacity_mva"])
            forward_headroom = max(POWER_FACTOR * capacity - float(station["estimated_station_forward_peak_mw"]), 0)
            reverse_headroom = max(0.8 * POWER_FACTOR * capacity - float(station["reverse_screen_mw"]), 0)
            feeder_headroom = receiver["receiving_current_headroom_mw"]
            donor_max_magnitude = source["current_equivalent_active_power_mw"]
            rows.append({
                "tie_id": tie["tie_id"],
                "year": 2025,
                "direction_from_station_id": source["station_id"],
                "direction_to_station_id": receiver["station_id"],
                "donor_feeder_id": source_feeder_id,
                "receiver_feeder_id": receiver_feeder_id,
                "donor_annual_current_power_upper_proxy_mw": donor_max_magnitude,
                "receiver_feeder_headroom_mw": feeder_headroom,
                "receiver_station_forward_headroom_mw": round(forward_headroom, 6),
                "receiver_station_reverse_headroom_mw": round(reverse_headroom, 6),
                "forward_import_screen_mw": round(min(donor_max_magnitude, feeder_headroom, forward_headroom), 6),
                "reverse_import_screen_mw": round(min(donor_max_magnitude, feeder_headroom, reverse_headroom), 6),
                "incremental_tie_investment_10k_cny": 0,
                "source_normal_state": tie["source_normal_state"],
                "screen_status": "source_end_static_upper_screen_shared_feeder_and_station_limits_not_summed",
                "topology_source_row": tie["source_row"],
            })
    return rows


def check_simultaneous_transfers(
    moves: list[dict],
    ties: list[dict],
    feeders: list[dict],
    station_net_load_mw: dict[str, float],
    station_capacity_mva: dict[str, float],
) -> dict:
    """正值转移用电、负值转移反送；只查静态源端和站级边界。"""
    tie_by_id = {row["tie_id"]: row for row in ties}
    feeder_by_id = {row["feeder_id"]: row for row in feeders}
    post_net = dict(station_net_load_mw)
    donor_use = defaultdict(float)
    receiver_use = defaultdict(float)
    used_ties = set()
    violations = []
    for move in moves:
        tie_id = move["tie_id"]
        source_id = move["donor_feeder_id"]
        receiver_id = move["receiver_feeder_id"]
        shift = float(move["net_load_shift_to_receiver_mw"])
        tie = tie_by_id[tie_id]
        if {source_id, receiver_id} != {tie["from_feeder_id"], tie["to_feeder_id"]} or source_id == receiver_id:
            raise ValueError(f"{tie_id} 不属于该跨站馈线对")
        if tie_id in used_ties:
            raise ValueError(f"{tie_id} 在同一静态场景不能同时取两个方向")
        used_ties.add(tie_id)
        donor = feeder_by_id[source_id]
        receiver = feeder_by_id[receiver_id]
        donor_use[source_id] += abs(shift)
        receiver_use[receiver_id] += abs(shift)
        post_net[donor["station_id"]] -= shift
        post_net[receiver["station_id"]] += shift
    for donor_id, used_mw in donor_use.items():
        if used_mw > feeder_by_id[donor_id]["current_equivalent_active_power_mw"] + 1e-9:
            violations.append(f"{donor_id}:shared_donor_annual_current_upper_proxy")
    for receiver_id, used_mw in receiver_use.items():
        if used_mw > feeder_by_id[receiver_id]["receiving_current_headroom_mw"] + 1e-9:
            violations.append(f"{receiver_id}:shared_receiving_headroom")
    for station_id, net_mw in post_net.items():
        capacity = station_capacity_mva[station_id]
        if net_mw > POWER_FACTOR * capacity + 1e-9:
            violations.append(f"{station_id}:forward_station_capacity")
        if -net_mw > 0.8 * POWER_FACTOR * capacity + 1e-9:
            violations.append(f"{station_id}:reverse_aggregate_screen")
    return {
        "feasible_under_static_screen": not violations,
        "violations": violations,
        "post_station_net_load_mw": post_net,
        "donor_feeder_usage_mw": dict(donor_use),
        "receiver_feeder_usage_mw": dict(receiver_use),
        "district_net_load_before_mw": sum(station_net_load_mw.values()),
        "district_net_load_after_mw": sum(post_net.values()),
        "technical_scope": "source_end_and_station_static_screen_not_downstream_or_guide_certification",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="核对邳州两站六馈线既有联络的受端开放容量")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    feeders = build_feeder_headroom()
    ties = read_cross_station_ties()
    screen = build_existing_tie_screen(feeders, ties)
    write_csv(feeders, args.output_dir / "pizhou_six_feeder_receiving_headroom_2025.csv")
    write_csv(ties, args.output_dir / "pizhou_existing_cross_station_ties.csv")
    write_csv(screen, args.output_dir / "pizhou_existing_tie_direction_screen_2025.csv")
    print(f"已核对 {len(feeders)} 条馈线、{len(ties)} 条跨站既有联络、{len(screen)} 个双向静态受端筛查。")


if __name__ == "__main__":
    main()
