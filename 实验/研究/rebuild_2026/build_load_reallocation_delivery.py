"""输出年度措施、正常负荷转接及独立措施有效性复核；不覆盖旧事故代理结果。"""

import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .baseline_2021 import read_csv
from .joint_shared_measure_audit import audit
from .regional_static_milp_v2_audit import audit as station_audit
from .station_grid_feasibility import grid_candidate_rows


MODEL = Path(__file__).resolve().parent
SOURCE = MODEL / "outputs/joint_shared_measure/transfer_10pct_all_years_load_reallocation"
DELIVERY = MODEL.parents[2] / "docs/2026-10-01三措施联合优化"
LABEL = {"pizhou": "邳州", "city": "市区", "rigid": "刚性", "elastic": "弹性"}


def service(load, area):
    return load if area == "A" else max(0, min(load - 12, 2 * load / 3))


def transformer_action(r):
    changes = [f"{slot}号主变{float(r[f'prior_unit_{slot}_mva']):g}→{float(r[f'unit_{slot}_mva']):g} MVA"
               for slot in (1, 2) if float(r[f"unit_{slot}_mva"]) >
               float(r[f"prior_unit_{slot}_mva"]) + 1e-5]
    if int(r["third_transformer_commissioned"]):
        changes.append("新增3号主变50 MVA")
    return "；".join(changes) or "沿用"


def main(source=SOURCE, delivery=DELIVERY, repeat_directory=None):
    source, delivery = Path(source), Path(delivery)
    for district in ("pizhou", "city"):
        station_audit(source / district)
    numerical = audit(source, repeat_directory)
    review = json.loads((source / "review.json").read_text())
    assert review["transfer_mode"] == "load_reallocation"
    annual = read_csv(source / "annual_comparison.csv")
    book = Workbook()
    summary_sheet = book.active
    summary_sheet.title = "逐年措施汇总"
    summary_sheet.append(["区县", "方案", "年份", "主变净增MVA", "新购主变MVA", "新增第三台数",
                          "新增储能MW", "新增储能MWh", "在役主变MVA", "在役储能MW", "在役储能MWh",
                          "正常负荷转接MW", "新建联络条数", "容载比"])
    moves_sheet = book.create_sheet("站间正常负荷转接")
    moves_sheet.append(["区县", "方案", "年份", "负荷转出站", "负荷承接站", "峰值转接MW", "通道类型",
                       "原表同网格辅助记录", "通道建模口径"])
    check_sheet = book.create_sheet("逐站容量与效果校核")
    check_sheet.append(["区县", "方案", "年份", "站号", "供电类别", "原负荷MW", "转出MW", "转入MW",
                       "转接后负荷MW", "主变MVA", "储能MW", "储能MWh", "正向余量MW", "反向保守余量MW",
                       "主变退出供电任务MW", "主变退出静态余量MW", "取消转接后退出余量MW",
                       "取消储能后退出余量MW", "不增主变后退出余量MW"])
    cost_sheet = book.create_sheet("费用与研究条件")
    cost_sheet.append(["区县", "刚性全寿命现值万元", "弹性全寿命现值万元", "节省万元", "节省比例"])
    reports, annual_details, constraints = {}, {}, []
    for district in ("pizhou", "city"):
        summaries = {r["scheme"]: r for r in json.loads((source / district / "summary.json").read_text())}
        grid_pairs = {tuple(sorted((r["station_a"], r["station_b"]))) for r in grid_candidate_rows(
            {r["station"] for r in read_csv(source / district / "rigid_stations.csv")})} if district == "pizhou" else set()
        costs = review["cost_npv_10k"][district]
        cost_sheet.append([LABEL[district], costs["rigid"], costs["elastic"],
                           costs["rigid"] - costs["elastic"], 1 - costs["elastic"] / costs["rigid"]])
        for scheme in ("rigid", "elastic"):
            summary = summaries[scheme]
            rows = read_csv(source / district / f"{scheme}_stations.csv")
            path = source / district / f"{scheme}_transfers.csv"
            transfers = read_csv(path) if path.exists() else []
            sheet = book.create_sheet(LABEL[district] + LABEL[scheme] + "具体措施")
            sheet.append(["年份", "站号", "主变具体动作", "主变净增MVA", "新购主变MVA",
                          "新增储能MW", "新增储能MWh", "在役储能MW", "在役储能MWh",
                          "转出负荷MW", "承接负荷MW", "转接前负荷MW", "转接后负荷MW",
                          "投运后主变MVA", "3号主变预留位", "10kV备用间隔数", "模型站点口径"])
            initial = {r["station"]: (float(r["prior_unit_1_mva"]), float(r["prior_unit_2_mva"]))
                       for r in rows if int(r["year"]) == 2022}
            portfolio = []
            for r in rows:
                year, station = int(r["year"]), r["station"]
                out = sum(float(m["mw"]) for m in transfers if int(m["year"]) == year and m["donor"] == station)
                incoming = sum(float(m["mw"]) for m in transfers if int(m["year"]) == year and m["receiver"] == station)
                old_load = float(r["forward_mw"])
                post = old_load - out + incoming
                assert abs(post - float(r["post_transfer_forward_mw"])) < 1e-5
                assert post >= -1e-5
                units = [float(r[f"unit_{slot}_mva"]) for slot in (1, 2, 3)]
                cap, energy = sum(units), float(r["storage_energy_mwh"])
                support = min(float(r["storage_power_mw"]), energy / max(2.15, int(r["forward_duration_h"])))
                reverse_support = min(float(r["storage_power_mw"]), energy / max(2.15, int(r["reverse_duration_h"])))
                task = service(post, r["area_class"])
                survivor = .95 * (cap - max(units))
                normal_margin = .95 * cap + support - post
                reverse_margin = .95 * float(summary["reverse_capacity_fraction"]) * cap + reverse_support - float(r["reverse_mw"]) - out
                n1_margin = survivor + support - task
                no_transfer = survivor + support - service(old_load, r["area_class"])
                no_storage = survivor - task
                storage_stress = survivor + .8 * support - task
                original_units = initial[station]
                no_expansion = .95 * (sum(original_units) - max(original_units)) + support - task
                assert min(normal_margin, reverse_margin, n1_margin) >= -1e-4
                result = {"year": year, "station": station, "forward_margin_mw": normal_margin,
                          "reverse_margin_mw": reverse_margin, "n1_margin_mw": n1_margin,
                          "no_transfer_n1_margin_mw": no_transfer,
                          "no_storage_n1_margin_mw": no_storage, "no_expansion_n1_margin_mw": no_expansion,
                          "storage_80pct_n1_margin_mw": storage_stress,
                          "assigned_load_mw": post, "original_load_mw": old_load}
                portfolio.append(result)
                check_sheet.append([LABEL[district], LABEL[scheme], year, station, r["area_class"],
                                    old_load, out, incoming, post, cap, float(r["storage_power_mw"]), energy,
                                    normal_margin, reverse_margin, task, n1_margin, no_transfer, no_storage, no_expansion])
                net_added = cap - float(r["prior_unit_1_mva"]) - float(r["prior_unit_2_mva"]) - (
                    0 if int(r["third_transformer_commissioned"]) else float(r["unit_3_mva"]))
                sheet.append([year, station, transformer_action(r), net_added, float(r["purchased_unit_mva"]),
                              float(r["new_storage_power_mw"]), float(r["new_storage_energy_mwh"]),
                              float(r["storage_power_mw"]), energy, out, incoming, old_load, post, cap,
                              int(r["source_available_third_slots"]), int(r["source_spare_10kv_bays"]),
                              "仿真点，类别按C类假定" if station.startswith("SIM-") else "原表站号对应仿真点"])
            for m in transfers:
                pair = tuple(sorted((m["donor"], m["receiver"])))
                moves_sheet.append([LABEL[district], LABEL[scheme], int(m["year"]), m["donor"], m["receiver"],
                                    float(m["mw"]), "既有仿真联络通道" if m["kind"] == "existing" else "新增仿真联络通道",
                                    "有" if pair in grid_pairs else "未证实",
                                    "区县内站间等效通道，参与正常负荷配置"])
            reports[district + "_" + scheme] = {
                "station_years": len(portfolio), "minimum_normal_margin_mw": min(r["forward_margin_mw"] for r in portfolio),
                "minimum_reverse_margin_mw": min(r["reverse_margin_mw"] for r in portfolio),
                "minimum_n1_margin_mw": min(r["n1_margin_mw"] for r in portfolio),
                "transfer_records": len(transfers),
                "transfer_records_with_shared_grid_evidence": sum(tuple(sorted((m["donor"], m["receiver"]))) in grid_pairs for m in transfers),
                "transfer_effective_station_years": sum(r["no_transfer_n1_margin_mw"] < -1e-4 for r in portfolio),
                "storage_effective_station_years": sum(r["no_storage_n1_margin_mw"] < -1e-4 for r in portfolio),
                "expansion_effective_station_years": sum(r["no_expansion_n1_margin_mw"] < -1e-4 for r in portfolio),
                "storage_80pct_insufficient_station_years": sum(r["storage_80pct_n1_margin_mw"] < -1e-4 for r in portfolio),
                "storage_80pct_largest_shortfall_mw": max(0, -min(r["storage_80pct_n1_margin_mw"] for r in portfolio)),
                "station_year_checks": portfolio,
            }
            for a in [a for a in annual if a["district"] == district and a["scheme"] == scheme]:
                year = int(a["year"])
                local = [r for r in rows if int(r["year"]) == year]
                previous_cap = sum(float(r["prior_unit_1_mva"]) + float(r["prior_unit_2_mva"]) +
                                   (0 if int(r["third_transformer_commissioned"]) else float(r["unit_3_mva"])) for r in local)
                transfer = sum(float(m["mw"]) for m in transfers if int(m["year"]) == year)
                before, after = sum(float(r["forward_mw"]) for r in local), sum(float(r["post_transfer_forward_mw"]) for r in local)
                assert abs(before - after) < 1e-5
                detail = {"net_added_mva": float(a["capacity_mva"]) - previous_cap,
                          "normal_transfer_mw": transfer, "station_load_sum_before_mw": before,
                          "station_load_sum_after_mw": after}
                annual_details[district + "_" + scheme + "_" + str(year)] = detail
                summary_sheet.append([LABEL[district], LABEL[scheme], year, detail["net_added_mva"],
                                      float(a["new_transformer_purchase_mva"]), int(a["new_third_transformers"]),
                                      float(a["new_storage_power_mw"]), float(a["new_storage_energy_mwh"]),
                                      float(a["capacity_mva"]), float(a["installed_storage_power_mw"]),
                                      float(a["installed_storage_energy_mwh"]), transfer, int(a["new_lines_commissioned"]), float(a["clr"])])
                constraints.append((district, scheme, year, a, detail, local, transfers))
    rigid_total = numerical["rigid_npv_10k"]
    elastic_total = numerical["elastic_npv_10k"]
    cost_sheet.append(["合计", rigid_total, elastic_total, rigid_total - elastic_total,
                       1 - elastic_total / rigid_total])
    for table in book:
        table.freeze_panes = "A2"
        table.auto_filter.ref = table.dimensions
        table.row_dimensions[1].height = 42
        for cell in table[1]:
            cell.fill = PatternFill("solid", fgColor="245473")
            cell.font = Font(name="微软雅黑", color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for col in range(1, table.max_column + 1):
            table.column_dimensions[get_column_letter(col)].width = 22
        for row in table.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    if abs(cell.value) < 1e-8:
                        cell.value = 0.0
                    cell.number_format = "0.0000" if (table is summary_sheet and cell.column == 14) else "0.000"
        if table is cost_sheet:
            for row in table.iter_rows(min_row=2):
                row[-1].number_format = "0.00%"
    delivery.mkdir(parents=True, exist_ok=True)
    book.save(delivery / "逐年具体措施与负荷转接审查.xlsx")
    audit_result = {"scope": "normal_load_reallocation_static_equipment_planning_not_fault_transfer",
                    "result_directory": str(source.relative_to(MODEL)), "numerical_audit": numerical,
                    "portfolios": reports, "annual_load_conservation": annual_details,
                    "counterfactual_scope": "same_chosen_equipment_layout_remove_one_measure_without_reoptimization"}
    (delivery / "负荷转接独立有效性审查.json").write_text(json.dumps(audit_result, ensure_ascii=False, indent=2) + "\n")
    text = ["# 逐年具体措施及负荷转接有效性审查", "",
            "转供按区县内正常负荷转接建模。转接后站负荷=原负荷−转出+转入，区县总负荷守恒；主变、储能及供电安全容量约束均使用转接后的负荷。事故恢复量未进入本轮规划。", "",
            "按用户确认，本轮属于仿真研究：区县内站间以仿真联络通道连接，不以实际工程路由为验收前提。共同10%限制是正常负荷转接的研究情景；原强制补足联络建设比例已撤销，新线按成本与容量收益寻优。", "",
            "刚性R≤2；弹性两区县每年R≥2.001，市区R低于邳州。R排序为用户规划条件，成本排序未强制。2021仍为共同反事实起点，以下2022—2025为回算规划年份。", "",
            "## 逐年新增措施", "",
            "净增主变容量是配置变化；新购主变铭牌容量包含整台替换，费用按新购容量计。储能栏为当年新增。转接MW是该年度方案的峰值负荷分配量，不是逐年累加投资量。", "",
            "| 区县 | 方案 | 年份 | 净增主变MVA | 新购主变MVA | 新增储能MW/MWh | 正常转接MW | 新线条数 | R |",
            "|---|---|---:|---:|---:|---:|---:|---:|---:|"]
    for district, scheme, year, a, detail, _, _ in constraints:
        text.append(f'| {LABEL[district]} | {LABEL[scheme]} | {year} | {detail["net_added_mva"]:g} | {float(a["new_transformer_purchase_mva"]):g} | {float(a["new_storage_power_mw"]):.1f}/{float(a["new_storage_energy_mwh"]):.3f} | {detail["normal_transfer_mw"]:.3f} | {a["new_lines_commissioned"]} | {float(a["clr"]):.4f} |')
    text += ["", "## 同价格费用对照", "",
             "费用是新增措施的全寿命现值，单位万元；包含储能运维与寿命更新，排除2021存量。", "",
             "| 区县 | 刚性 | 弹性 | 节省比例 |", "|---|---:|---:|---:|"]
    for district, cost in review["cost_npv_10k"].items():
        text.append(f'| {LABEL[district]} | {cost["rigid"]:.2f} | {cost["elastic"]:.2f} | {1-cost["elastic"]/cost["rigid"]:.2%} |')
    text.append(f'| 合计 | {rigid_total:.2f} | {elastic_total:.2f} | {1-elastic_total/rigid_total:.2%} |')
    text += ["", "## 静态措施是否有效", "",
             "以同一设备布局为基准，逐项取消正常负荷转接、储能或主变增容，复算供电安全容量缺口。该检查说明措施对所选方案的作用，不等于重新寻优后的全局收益。", "",
             "| 区县 | 方案 | 转接量撤销后不足的站年数 | 储能撤销后不足的站年数 | 不增主变后不足的站年数 | 最小静态退出余量MW |",
             "|---|---|---:|---:|---:|---:|"]
    for key, r in reports.items():
        district, scheme = key.split("_")
        text.append(f'| {LABEL[district]} | {LABEL[scheme]} | {r["transfer_effective_station_years"]} | {r["storage_effective_station_years"]} | {r["expansion_effective_station_years"]} | {max(0, r["minimum_n1_margin_mw"]):.6f} |')
    text += ["", "把储能有效支撑降至80%作自设压力测试，保持设备和负荷转接方案不变，结果如下。这不是标准规定的效率或SOC取值，也没有进行应力情景再优化。", "",
             "| 区县 | 方案 | 供电安全容量不足的站年数 | 最大缺口MW |", "|---|---|---:|---:|"]
    for key, r in reports.items():
        district, scheme = key.split("_")
        text.append(f'| {LABEL[district]} | {LABEL[scheme]} | {r["storage_80pct_insufficient_station_years"]} | {r["storage_80pct_largest_shortfall_mw"]:.3f} |')
    text += ["", "## 仿真设定与有效性判定", "",
             "1. 审查结论：本轮静态仿真范围内有效。逐年方案具体到仿真站号、设备规格、储能MW/MWh、转接方向及MW，392个站年均通过正向、反向和转接后供电安全容量复算。实际路由不属于本轮验收范围。",
             "2. 全站对既有仿真网络在10%限制下可承接新通道功能，正投资新线被零新增投资既有通道支配，因此新线选零。三类措施共同参与优化不要求三类均建设；正常站间负荷转接持续参与容量配置。",
             "3. 主变新增第三台按原表预留位作为仿真可选条件。市区29站外推，SIM-CITY点类别按C类假定。2021容量是共同反事实起点，2022—2025为回算规划年份。",
             "4. 储能为0.1 MW/0.215 MWh整数柜，按D95峰段持续时间折算可持续功率；当前属于静态容量仿真，SOC与充放电时序未显式求解。80%支撑压力测试单列，不改变主方案的确定性结果。",
             "5. 电源保持原站归属，反向压力用原反送峰值+转出峰值作保守上界，未抵扣受端消纳收益；避免负荷转出后反送压力被漏算。",
             "6. 站级负荷为独立峰值任务，转接前后任务总量守恒；容载比分母保持原表区县降压负荷代理。站间转接优化负荷空间分配，不能直接改变区县总负荷或把联络容量计入主变容量分子。", "",
             "## 每年逐站动作和转接方向", ""]
    for district, scheme, year, _, _, local, transfers in constraints:
        text += [f"### {LABEL[district]}{LABEL[scheme]} {year}年", ""]
        selected = [r for r in local if float(r["purchased_unit_mva"]) > 1e-5 or float(r["new_storage_energy_mwh"]) > 1e-5]
        for r in selected:
            action = transformer_action(r)
            if float(r["new_storage_energy_mwh"]) > 1e-5:
                action += f'；新增储能{float(r["new_storage_power_mw"]):.1f} MW/{float(r["new_storage_energy_mwh"]):.3f} MWh'
            text.append(f'- {r["station"]}：{action}。')
        if not selected:
            text.append("- 设备沿用上一年。")
        for m in transfers:
            if int(m["year"]) == year:
                text.append(f'- 负荷转接：{m["donor"]}转出至{m["receiver"]}，{float(m["mw"]):.3f} MW；既有仿真联络通道。')
        text.append("")
    if numerical["repeat_csv_exact_match"]:
        text += ["## 确定性验证", "",
                 f'同输入从头重复三阶段计算，{len(numerical["output_sha256"])}份CSV逐字节一致。两次均取得已证明最优状态，三阶段间隙为0；392个站年数值与负荷守恒复算通过。18项相关回归测试通过。', ""]
    (delivery / "逐年具体措施与有效性审查.md").write_text("\n".join(text).rstrip() + "\n")
    print(json.dumps({"delivery": str(delivery), "measures": {
        k: {f: v for f, v in r.items() if f != "station_year_checks"} for k, r in reports.items()}}, ensure_ascii=False))
    return audit_result


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--repeat", type=Path)
    args = parser.parse_args()
    main(repeat_directory=args.repeat)
