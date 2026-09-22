"""摘录本地工程投资原表中与本轮两种电压相关的直接案例。"""

import argparse
import csv
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from .official_annual import STUDY_DIR


SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/江苏徐州邢楼110千伏变电站主变扩建等工程建设规模及投资汇总表(1).xlsx"
TARGET = Path(__file__).resolve().parent / "source_audit/transformer_project_cost_references.csv"
CASES = {
    14: (110, "third_transformer_expansion", 50, 0),
    16: (110, "third_transformer_expansion", 50, 0),
    43: (35, "transformer_replacement", 20, 10),
    45: (35, "transformer_replacement", 20, 10),
}


def read_cost_references(source: Path = SOURCE) -> list[dict]:
    source = Path(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    source_label = str(source.resolve().relative_to(STUDY_DIR))
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows = []
    try:
        sheet = workbook["Sheet1"]
        for source_row, (voltage_kv, measure_type, purchased_mva, replaced_mva) in CASES.items():
            cells = next(sheet.iter_rows(min_row=source_row, max_row=source_row, values_only=True))
            rows.append(
                {
                    "voltage_kv": voltage_kv,
                    "project_name": cells[1],
                    "project_scope_raw": cells[2],
                    "measure_type": measure_type,
                    "purchased_transformer_mva": purchased_mva,
                    "replaced_old_transformer_mva": replaced_mva,
                    "project_capacity_increment_mva": purchased_mva - replaced_mva,
                    "construction_cost_10k_cny": cells[9],
                    "equipment_purchase_cost_10k_cny": cells[10],
                    "installation_cost_10k_cny": cells[11],
                    "other_cost_10k_cny": cells[12],
                    "contingency_cost_10k_cny": cells[14],
                    "static_total_10k_cny": cells[15],
                    "dynamic_total_10k_cny": cells[16],
                    "pricing_role": "project_case_reference_not_generic_unit_price",
                    "source_file": source_label,
                    "source_sheet": sheet.title,
                    "source_row": source_row,
                    "source_sha256": source_hash,
                }
            )
    finally:
        workbook.close()
    return rows


def write_cost_references(source: Path = SOURCE, target: Path = TARGET) -> int:
    rows = read_cost_references(source)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="摘录邳州 110/35 kV 主变工程投资案例")
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=TARGET)
    args = parser.parse_args()
    print(f"已写入 {write_cost_references(args.source, args.output)} 条项目投资案例")


if __name__ == "__main__":
    main()
