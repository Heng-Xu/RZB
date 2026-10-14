"""只扩展输入映射；求解、设备、馈线和正常/N-1约束复用锁定模型。"""
from functools import lru_cache
from pathlib import Path
from types import FunctionType
import hashlib
import numpy as np
from openpyxl import load_workbook
from ..matrix_20261012 import model as locked

ROOT = Path(__file__).resolve().parents[4]
PV_SOURCE = ROOT/'实验/研究/data/tuomin/电网建模数据_Agent整合版_V1.2/光伏8760小时数据.xlsx'
LOAD_HOURS = 4400.0
INPUTS = ['eta', 'energy_penetration', 'growth_rate', 'cable_fraction',
          'initial_transfer_fraction', 'h95_h', 'diversity', 'line_price', 'storage_price']


@lru_cache(maxsize=1)
def solar_shape():
    book = load_workbook(PV_SOURCE, read_only=True, data_only=True)
    values = np.array([r[4] for r in book['光伏8760'].iter_rows(min_row=2, values_only=True)], float)
    book.close()
    assert len(values) == 8760 and np.isfinite(values).all() and values.min() >= 0
    # 场景假设：08:00—18:00为等效日间窗口，窗口外出力为零。
    daylight = np.isin(np.arange(8760) % 24, np.arange(8, 18))
    values[~daylight] = 0
    values /= values.max()
    return values, daylight


def peak_blocks(h):
    # 五个错开的冬季夜峰，分别用于站级局部高峰；全部保持区域总峰值。
    return [np.arange(day*24+18, day*24+18+h) for day in [10, 20, 30, 40, 50]]


def _curve(eta, amplitude, h):
    phi, daylight = solar_shape()
    load = np.where(daylight, .30, 0.)
    blocks = peak_blocks(h)
    peak_hours = np.concatenate(blocks)
    other = ~daylight.copy()
    other[peak_hours] = False
    load[peak_hours] = 1.
    load[other] = (LOAD_HOURS-load.sum())/other.sum()
    assert 0 <= load[other][0] < .94 and abs(load.sum()-LOAD_HOURS) < 1e-8
    power = eta*amplitude
    maximum = max(1., power-.30)
    generation = np.minimum(power*phi, .30+.949*maximum)
    # 年内唯一日间极值平台；其余时刻严格低于95%绝对峰值。
    day_block = np.arange(171*24+9, 171*24+9+h)
    generation[day_block] = power
    assert daylight[day_block].all() and np.all(generation[~daylight] == 0)
    return load, generation, blocks


@lru_cache(maxsize=8192)
def profile(eta, energy, h):
    eta, energy, h = float(eta), float(energy), int(h)
    assert eta >= 0 and 0 <= energy <= 1 and 1 <= h <= 8
    if eta == 0:
        if energy != 0: raise ValueError('零装机对应零发电量')
        amplitude = 0.
    else:
        if energy <= 0 or not .12-1e-10 <= energy/eta <= .40+1e-10:
            raise ValueError('电量与装机组合超出声明的利用水平域')
        _, upper, _ = _curve(eta, 1., h)
        if energy > upper.sum()/LOAD_HOURS+1e-10:
            raise ValueError('给定装机和出力上限不能达到所设年发电量')
        lo, hi = 0., 1.
        for _ in range(42):
            mid = (lo+hi)/2
            _, generation, _ = _curve(eta, mid, h)
            if generation.sum()/LOAD_HOURS < energy: lo = mid
            else: hi = mid
        amplitude = (lo+hi)/2
    load, generation, blocks = _curve(eta, amplitude, h)
    net = load-generation
    mask = abs(net) >= .95*abs(net).max()
    edges = np.diff(np.r_[False, mask, False].astype(int))
    longest = int(np.max(np.flatnonzero(edges == -1)-np.flatnonzero(edges == 1)))
    assert longest == h
    assert abs(generation.sum()/load.sum()-energy) < 1e-10
    assert abs(net.max()-1) < 1e-12 and generation.max() <= eta+1e-10
    return dict(amplitude=amplitude, energy_penetration=energy, h95_h=longest,
                cumulative_h95_h=int(mask.sum()), load_equivalent_hours=float(load.sum()),
                generation_equivalent_hours=float(generation.sum()/eta) if eta else 0.,
                forward_peak_per_unit=float(net.max()), reverse_peak_per_unit=float(max(-net.min(),0)),
                curve_sha256=hashlib.sha256(np.column_stack([load,generation]).astype('<f8').tobytes()).hexdigest())


def make_input(parameters):
    p = dict(parameters)
    p['support_duration_h'] = p['h95_h']
    certificate = profile(p['eta'], p['energy_penetration'], p['h95_h'])
    inp = locked.make_input(p)
    inp['generation'][-1] *= certificate['amplitude']/.75
    inp['net'] = inp['gross']-inp['generation']
    # 对全部新增联络项目使用相同价格倍率，固定技术能力与接收约束。
    inp['line_cost'] *= p.get('line_price', 1.)
    inp['profile_certificate'] = certificate
    return inp


def rebound(function):
    return FunctionType(function.__code__, dict(function.__globals__, make_input=make_input),
                        function.__name__, function.__defaults__, function.__closure__)


_solve = rebound(locked.solve)
independent_lp_check = rebound(locked.independent_lp_check)


def indicators(inp):
    net = inp['net']; peak = inp['peak']; p = inp['p']
    region = net.sum(axis=1)
    ct = inp['transformer_cost_per_purchased_mva']/p['pf']
    cl = inp['line_cost']/7.49106117
    cb = inp['storage_unit_cost']*p['storage_price']/inp['storage_effective']
    return dict(forward_peak_mw=float(max(region.max(),0)),
                reverse_peak_mw=float(max(-region.min(),0)),
                initial_max_loading=float(abs(net).max(axis=0).max()/(100*p['pf'])),
                transformer_line_ratio=ct/cl, transformer_storage_ratio=ct/cb,
                transformer_unit_cost=ct, line_unit_cost=cl, storage_unit_cost=cb)


def solve(parameters):
    result = _solve(parameters)
    inp = make_input(parameters)
    result.update(indicators(inp))
    result['profile_certificate'] = inp['profile_certificate']
    result['extension_version'] = 'indicator_extension_20261014'
    return result
