"""共享零储能可行候选并输出可审查的完整年度措施。"""
import copy
import json
import shutil

import numpy as np
from scipy.sparse import coo_matrix
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .run import OUT, ROOT, CASES, rows, dump, write_csv, sha, build, load_scenes, network, actual_limits

LABELS={'rigid_storage':'刚性·允许储能','elastic_storage':'弹性·允许储能',
        'rigid_no_storage':'刚性·无储能对照','elastic_no_storage':'弹性·无储能对照'}


def share_zero_storage():
    network.limits=actual_limits
    for region in ('city','pizhou'):
        root=OUT/region;inp=json.loads((root/'inputs.json').read_text());scenes=load_scenes(root,inp)
        for mode in ('rigid','elastic'):
            source=root/f'{mode}_no_storage';target=root/f'{mode}_storage'
            s=json.loads((source/'summary.json').read_text());t=json.loads((target/'summary.json').read_text())
            if 'annual' not in s or ('annual' in t and t['cost_npv_10k']<=s['cost_npv_10k']+.001):continue
            chosen=source/'chosen_solution.npz'
            if not chosen.exists():chosen=source/'static_solution.npz'
            z=np.load(chosen)['solution'];m,meta=build(inp,scenes,f'{mode}_storage')
            rr,cc,vv=zip(*m.entries);a=coo_matrix((vv,(rr,cc)),shape=(len(m.constraint_lower),len(m.costs))).tocsr()
            assert len(z)==len(m.costs)
            az=a@z
            error=max(0.,float(np.max(np.asarray(m.constraint_lower)-az)),float(np.max(az-np.asarray(m.constraint_upper))),
                      float(np.max(np.asarray(m.lower_bounds)-z)),float(np.max(z-np.asarray(m.upper_bounds))))
            assert error<1e-5
            dump(target/'independent_search_summary.json',t)
            merged=copy.deepcopy(s);merged['case']=f'{mode}_storage'
            merged['solver']=copy.deepcopy(t['solver'])
            lb=merged['solver'].get('dual_bound')
            merged['solver'].update(objective=s['cost_npv_10k'],feasible_incumbent=True,accepted=False,
                                    status='采用更便宜的已核查零储能候选',
                                    mip_gap=max(0.,(s['cost_npv_10k']-lb)/s['cost_npv_10k']) if lb is not None else None)
            merged['zero_storage_reuse']={'source':str(source.relative_to(ROOT)),'target_matrix_max_violation':error}
            for filename in ('years.csv','stations.csv','projects.csv'):
                shutil.copyfile(source/filename,target/filename)
            np.savez_compressed(target/'chosen_solution.npz',solution=z)
            dump(target/'summary.json',merged)


def strongest_bounds():
    for region in ('city','pizhou'):
        for case in CASES:
            out=OUT/region/case;p=out/'summary.json';s=json.loads(p.read_text())
            if 'annual' not in s:continue
            bounds=[s['solver'].get('dual_bound')]
            for f in out.glob('static*_solver.json'):
                bounds.append(json.loads(f.read_text()).get('dual_bound'))
            valid=[v for v in bounds if v is not None]
            if valid:
                s['solver']['dual_bound']=max(valid)
                s['solver']['mip_gap']=max(0.,(s['cost_npv_10k']-max(valid))/s['cost_npv_10k'])
                s['solver']['bound_basis']='本输入历次寻优的最强有效费用下界'
                dump(p,s)


def markdown_table(headers, data):
    return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+
                     ['| '+' | '.join(str(v) for v in r)+' |' for r in data])


def report():
    status=[];annual=[];units=[];lines=[];stationyears=[];costs=[];inputs={};summaries={};sources={}
    for region in ('city','pizhou'):
        root=OUT/region;inp=json.loads((root/'inputs.json').read_text());inputs[region]=inp
        sources[str((root/'inputs.json').relative_to(ROOT))]=sha(root/'inputs.json')
        sources[str((root/'静态工况.csv').relative_to(ROOT))]=sha(root/'静态工况.csv')
        for case in CASES:
            p=root/case;s=json.loads((p/'summary.json').read_text());summaries[region,case]=s
            sources[str((p/'summary.json').relative_to(ROOT))]=sha(p/'summary.json')
            solver=s['solver'];ok='annual' in s
            status.append({'地区':inp['region_name'],'方案':LABELS[case],
                           '状态':'已取得静态可行候选' if ok else ('本模型内无解' if solver.get('proven_infeasible') else '尚未取得可行解'),
                           '费用现值_万元':s.get('cost_npv_10k'),'费用下界_万元':solver.get('dual_bound'),
                           '相对间隙':solver.get('mip_gap'),'说明':s['scope']})
            if not ok:continue
            sr=rows(p/'stations.csv');pr=rows(p/'projects.csv')
            old={sid:list(d['initial_units'])+[0.] for sid,d in inp['stations'].items()};oldn={sid:0. for sid in old}
            for a in s['annual']:
                y=a['year']; prefix={'地区':inp['region_name'],'方案':LABELS[case],'年份':y}
                measures=[]
                for r in sr:
                    if int(r['year'])!=y:continue
                    sid=r['station'];uu=[float(r[f'unit_{k}_mva']) for k in (1,2,3)]
                    n=float(r['storage_mwh']);delta=sum(uu)-sum(old[sid]);bought=float(r['purchased_transformer_mva'])
                    stationyears.append({**prefix,'站点':sid,'主变1_MVA':uu[0],'主变2_MVA':uu[1],'主变3_MVA':uu[2],
                                         '容量_MVA':sum(uu),'容量净增_MVA':delta,'购置容量_MVA':bought,
                                         '储能累计_MWh':n,'储能新增_MWh':n-oldn[sid]})
                    if abs(delta)>.001 or abs(n-oldn[sid])>.001:
                        before='+'.join(f'{v:g}' for v in old[sid] if v);after='+'.join(f'{v:g}' for v in uu if v)
                        units.append({**prefix,'站点':sid,'调整前主变_MVA':before,'调整后主变_MVA':after,
                                      '容量净增_MVA':delta,'购置容量_MVA':bought,'新增储能_MWh':n-oldn[sid]})
                        if delta>.001:measures.append(f'{sid}：{before}→{after}')
                    old[sid]=uu;oldn[sid]=n
                yrprojects=[r for r in pr if int(r['year'])==y]
                assert len(yrprojects)==a['new_lines']
                for r in yrprojects:
                    lines.append({**prefix,'项目':r['id'],'源站':r['station_a'],'邻站':r['station_b'],
                                  '源端馈线':r['a'],'接收馈线':r['b'],'能力_MW':float(r['capacity_mw']),
                                  '新增间隔站':r.get('new_bay_station',''),'候选依据':r['basis']})
                annual.append({**prefix,'实际最大负荷_MW':a['peak_mw'],'实际统计容量_MVA':a['actual_capacity_mva'],
                               '方案容量_MVA':a['capacity_mva'],'方案容载比':a['clr'],
                               '容量净增_MVA':a['net_capacity_increment_mva'],'购置容量_MVA':a['purchased_transformer_mva'],
                               '新增第三台_台':a['new_third_count'],'新建联络_项':a['new_lines'],
                               '其中新出线_项':sum(bool(r.get('new_bay_station')) for r in yrprojects),
                               '储能新增_MWh':a['new_storage_mwh'],'储能累计_MWh':a['storage_mwh'],
                               '新增措施费用现值_万元':a['added_lifecycle_npv_10k'],
                               '累计费用现值_万元':a['cumulative_lifecycle_npv_10k'],
                               '被动超限':'是' if a['passive_exceedance'] else '否',
                               '主变措施':'；'.join(measures) or '保持'})
                costs.append({**prefix,'主变费用现值_万元':a['transformer_npv_10k'],
                              '联络费用现值_万元':a['line_npv_10k'],'储能费用现值_万元':a['storage_npv_10k'],
                              '当年建设投资_万元':a['construction_capex_10k'],
                              '新增措施费用现值_万元':a['added_lifecycle_npv_10k']})
    comparison=[]
    for region in ('city','pizhou'):
        r=summaries[region,'rigid_storage'];e=summaries[region,'elastic_storage']
        if 'annual' not in r or 'annual' not in e:continue
        lb=r['solver'].get('dual_bound')
        comparison.append({'地区':inputs[region]['region_name'],'刚性费用现值_万元':r['cost_npv_10k'],
                           '弹性费用现值_万元':e['cost_npv_10k'],'弹性相对刚性费用变化':e['cost_npv_10k']/r['cost_npv_10k']-1,
                           '费用排序是否由界证明':bool(lb is not None and e['cost_npv_10k']<lb),
                           '刚性联络合计':sum(a['new_lines'] for a in r['annual']),
                           '弹性联络合计':sum(a['new_lines'] for a in e['annual']),
                           '刚性期末储能_MWh':r['annual'][-1]['storage_mwh'],
                           '弹性期末储能_MWh':e['annual'][-1]['storage_mwh']})
    targets=[]
    for y in inputs['city']['years']:
        city=summaries['city','elastic_storage'];pz=summaries['pizhou','elastic_storage']
        if 'annual' in city and 'annual' in pz:
            ca=next(a for a in city['annual'] if a['year']==y);pa=next(a for a in pz['annual'] if a['year']==y)
            targets.append({'年份':y,'市区弹性容载比':ca['clr'],'邳州弹性容载比':pa['clr'],
                            '市区小于邳州_仅观察':ca['clr']<pa['clr'],
                            '两地弹性均突破2':min(ca['clr'],pa['clr'])>2.})
    assumptions=[
        {'事项':'核心数据','本轮处理':'近5年容载比.xlsx，市区Sheet1第9行、邳州第19行，容量和降压负荷原单位乘10；不以增长率替代逐年统计。'},
        {'事项':'共同起点','本轮处理':'市区2900、邳州1463.5 MVA保持；为规划仿真资产，实际容量另列，不替换或累计叠加。'},
        {'事项':'刚性被动超限','本轮处理':'C_y≤max(C_(y-1),2P_y)，容量非减；仅保留已有容量，不给超限新增扩容额度。'},
        {'事项':'弹性方案','本轮处理':'下限2.001及原物理候选上界保持；费用最小；地区排序按用户最新确认仅作观察，不约束、不据此判定方案是否通过；不强制逐年上升。'},
        {'事项':'静态范围','本轮处理':'只校核归档区县峰值和最低净负荷两个工况及各站最大单主变N-1；不新增逐时/SOC仿真，不宣称覆盖全年或各站局部峰值。'},
        {'事项':'反向工况','本轮处理':'沿用净负荷模板，按实际年度降压峰值同倍率标定；历史反送不是实测，新能源发展变化未单独恢复。'},
        {'事项':'市区网络','本轮处理':'与邳州使用同一候选边/共享馈线/N-1/备用间隔模型。上级电源交集只用于候选筛选假设；馈线按已用间隔等分峰值，热限取原新线参考值。'},
        {'事项':'市区存量转供','本轮处理':'无已核实10kV存量边，不按原40%给予免费全区能力；市区全部选中边计费。该保守边界需后续存量网络敏感性，不能当实际投资预算。'},
        {'事项':'邳州网络','本轮处理':'原T01及279条付费候选保持；共享馈线负荷块及接收余量不能重复使用。'},
        {'事项':'储能','本轮处理':'沿用地区正向持续时间市区6h、邳州2h；两地统一有效电量系数0.76。市区由旧1.0同步为邳州设置。'},
        {'事项':'候选线路费用','本轮处理':'原3km及综合单价假设保持；无测绘路线。'},
        {'事项':'旧结果比较','本轮处理':'原结果保留；本次变化含实际年度负荷、被动超限、静态工况简化、市区网络同步，不能只归因于更正负荷。'},
        {'事项':'求解精度','本轮处理':'有解为通过矩阵及独立网络核查的候选；间隙单列，未达1%不称精确最优。允许储能组可选已核查的更便宜无储能候选。'},
    ]
    sheets={'方案状态':status,'两地费用比较':comparison,'逐年方案措施':annual,'逐年费用分项':costs,
            '主变及储能变化':units,'联络项目明细':lines,'逐站年度容量':stationyears,
            '原表容量负荷':rows(OUT/'原表容量负荷.csv'),'负荷输入差异':rows(OUT/'负荷输入差异.csv'),
            '地区排序核对':targets,'无储能必要条件':rows(OUT/'无储能必要条件.csv'),'模型与假设':assumptions}
    wb=Workbook();wb.remove(wb.active)
    for name,data in sheets.items():
        if not data:continue
        write_csv(OUT/f'{name}.csv',data)
        ws=wb.create_sheet(name);ws.append(list(data[0]));
        for r in data:ws.append([r.get(k) for k in data[0]])
        ws.freeze_panes='D2' if '年份' in data[0] else 'A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='25465F');c.font=Font(name='方正仿宋_GBK',color='FFFFFF',bold=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font=Font(name='方正仿宋_GBK',size=11);c.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(c.value,float):c.number_format='0.00'
        for j,key in enumerate(data[0],1):
            ws.column_dimensions[get_column_letter(j)].width=60 if key in ('说明','本轮处理','主变措施','候选依据') else 21
            if key in ('相对间隙','弹性相对刚性费用变化'):
                for c in list(ws.columns)[j-1][1:]:c.number_format='0%'
    path=OUT/'两地静态逐年方案对比.xlsx';wb.save(path)
    check=load_workbook(path,read_only=True,data_only=True)
    assert check['逐年方案措施'].max_row==len(annual)+1;check.close()
    for region in ('city','pizhou'):
        for case in CASES:
            s=summaries[region,case]
            if 'annual' not in s:continue
            assert abs(sum(r['added_lifecycle_npv_10k'] for r in s['annual'])-s['cost_npv_10k'])<.02
            assert all(r['max_network_deficit_mw']<=.001 for r in s['static_network_checks'])
    text=['# 两地实际年度负荷静态方案对比','',
          '本次为独立试算，原方案及正式报告保持。核心年度统计来自原容载比表，仿真起点保持。两地采用同一网络约束结构；市区馈线及候选关系含规划假设。', '',
          markdown_table(['地区','刚性费用/万元','弹性费用/万元','费用变化','是否由上下界证明更低'],
                         [[r['地区'],f"{r['刚性费用现值_万元']:.2f}",f"{r['弹性费用现值_万元']:.2f}",f"{100*r['弹性相对刚性费用变化']:.0f}%",'是' if r['费用排序是否由界证明'] else '否，仅为当前候选比较'] for r in comparison])]
    for region in ('city','pizhou'):
        name=inputs[region]['region_name'];text.extend(['',f'## {name}逐年方案','',
            markdown_table(['方案','年','实际负荷/MW','方案容量/MVA','容载比','容量净增/MVA','新联络/项','新增储能/MWh','年度措施现值/万元'],
                [[r['方案'],r['年份'],f"{r['实际最大负荷_MW']:.2f}",f"{r['方案容量_MVA']:.1f}",f"{r['方案容载比']:.3f}",
                  f"{r['容量净增_MVA']:.1f}",r['新建联络_项'],f"{r['储能新增_MWh']:.3f}",f"{r['新增措施费用现值_万元']:.2f}"]
                 for r in annual if r['地区']==name and r['方案'] in ('刚性·允许储能','弹性·允许储能')])])
    text.extend(['','## 状态与适用范围','',markdown_table(['地区','方案','状态','求解间隙'],
        [[r['地区'],r['方案'],r['状态'],('<1%' if 0<r['相对间隙']<.01 else f"{100*r['相对间隙']:.0f}%") if r['相对间隙'] is not None else '—'] for r in status]),'',
        '地区容载比关系逐年列示，仅供观察。用户已确认不强制市区小于邳州，不为匹配排序追加容量。','',
        markdown_table(['年份','市区弹性','邳州弹性','市区小于邳州'],
                       [[r['年份'],f"{r['市区弹性容载比']:.3f}",f"{r['邳州弹性容载比']:.3f}",'是（观察）' if r['市区小于邳州_仅观察'] else '否（观察）'] for r in targets]),''])
    text.extend(f"- {r['事项']}：{r['本轮处理']}" for r in assumptions)
    text.extend(['','刚性无储能若无解，只表示本轮数据、候选和控制规则下无解。市区无储能结果还受假设馈线限制，不能据此认定现实工程必须配置储能。',
                 '逐站主变调整、储能新增、具体候选联络及完整费用分项见同目录工作簿。容量净增与主变购置容量分别列示，购置费用不能按净增容量替代。'])
    (OUT/'两地静态方案说明.md').write_text('\n'.join(text)+'\n')
    dump(OUT/'结果汇总.json',{'status':status,'comparison':comparison,'annual':annual,'targets':targets,'assumptions':assumptions})
    dump(OUT/'来源哈希.json',sources)
    dump(OUT/'核验.json',{'annual_rows':len(annual),'station_rows':len(stationyears),'project_rows':len(lines),
                         'cost_recomputed':True,'static_network_certificates_passed':True,'workbook_reopened':True,
                         'model_tests':'test_rules.py：析取约束和实际年度标定已验证'})
    print(json.dumps(comparison,ensure_ascii=False,indent=2))


if __name__=='__main__':
    strongest_bounds();share_zero_storage();report()
