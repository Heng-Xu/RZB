"""两地共同资产、费用和网络约束；仅使用已归档静态工况与年度统计。

python -m 实验.研究.rebuild_2026.static_actual_20261010.run --prepare
python -m 实验.研究.rebuild_2026.static_actual_20261010.run --region city --case elastic_storage
"""
import argparse
import copy
import csv
import itertools
import json
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from ..official_annual import read_official_annual, DEFAULT_SOURCE
from ..pilot_20261009.run import ROOT, PF, CASES, dump, write_csv, sha
from ..synchronous_20261009 import model as assets_model
from ..network_case_20261010 import run as network
from ..network_case_20261010.solver import solve
from ..two_district_ordered_guide_run import guide_range

OUT = ROOT / '研究报告/试算/20261010_实际年度负荷静态试算'
AUDIT = ROOT / '实验/研究/rebuild_2026/source_audit'
SEEDS = Path(__file__).resolve().parent / 'frozen_inputs'


def rows(path):
    with Path(path).open(encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))


def city_network(inp):
    """以同上级电源筛候选；明确是规划假设，不认定已有10kV连接。

    已用间隔数作为出线数代理；单线分担负荷取本站静态峰值等分。
    同额定新线的7.491MW作为缺失馈线热限的试算值。无已核实存量边。
    """
    source = DEFAULT_SOURCE.parent / '高压网架结构.xlsx'
    w = load_workbook(source, read_only=True, data_only=True)
    upstream = {r[1]: {str(v) for v in r[5:8] if str(v).startswith('BDZ-')}
                for r in w.worksheets[0].iter_rows(min_row=3, values_only=True)
                if r[1] in inp['stations']}
    w.close()
    feeders = {}; pools = {}; evidence = []
    scale = inp['annual_observed_peak_mw']['2025'] / inp['initial_peak_mw']
    for sid, s in inp['stations'].items():
        count = max(0, s['metadata']['total_10kv_bays'] - s['metadata']['spare_10kv_bays'])
        pools[sid] = []
        for j in range(count):
            f = f'ASSUMED-{sid}-{j+1:02d}'; pools[sid].append(f)
            feeders[f] = dict(id=f, station=sid, grid='city_upstream_candidate',
                              thermal_mw=inp['line_increment_mw'],
                              maximum_load_proxy_2025_mw=s['base_forward_mw'] * scale / count,
                              source_row='assumed_equal_load_per_occupied_bay', source_reported_active_mw=None)
        evidence.append(dict(station=sid, occupied_bay_proxy=count,
                             upstream='|'.join(sorted(upstream.get(sid, set()))),
                             candidate_status='assumed_from_shared_upstream' if sid in upstream else 'no_candidate_evidence'))
    edges = []; cursor = {sid: 0 for sid in pools}
    for a, b in itertools.combinations(inp['stations'], 2):
        shared = upstream.get(a, set()) & upstream.get(b, set())
        if not shared or not pools[a] or not pools[b]:
            continue
        width = min(4, len(pools[a]), len(pools[b]))
        # 各候选站对轮流分配源端馈线，避免所有走廊人为挤在同四条代理馈线上。
        endpoints = {s: [pools[s][(cursor[s]+j) % len(pools[s])] for j in range(width)] for s in (a,b)}
        for s in (a,b): cursor[s] += width
        for j in range(width):
            edges.append(dict(id=f'CITY-TIE-{len(edges)+1:04d}', a=endpoints[a][j], b=endpoints[b][j],
                              station_a=a, station_b=b, capacity_mw=inp['line_increment_mw'],
                              existing=False, grid='|'.join(sorted(shared)), source_section_fraction=1.,
                              bidirectional=True, route_length_km_assumption=3., new_bay_station='',
                              basis='规划假设：同上级电源筛候选，既有馈线负荷等分；非已核实10kV线路'))
        for donor, receiver in ((a, b), (b, a)):
            if inp['stations'][receiver]['metadata']['spare_10kv_bays'] <= 0:
                continue
            for f in endpoints[donor][:2]:
                eid = f'CITY-FEEDER-{len(edges)+1:04d}'; nf = 'SIM-' + eid
                feeders[nf] = dict(id=nf, station=receiver, grid='city_upstream_candidate',
                                   thermal_mw=inp['line_increment_mw'], maximum_load_proxy_2025_mw=0.,
                                   source_row='planning_new_feeder', source_reported_active_mw=0.)
                edges.append(dict(id=eid, a=f, b=nf, station_a=donor, station_b=receiver,
                                  capacity_mw=inp['line_increment_mw'], existing=False,
                                  grid='|'.join(sorted(shared)), source_section_fraction=1.,
                                  bidirectional=False, route_length_km_assumption=3., new_bay_station=receiver,
                                  basis='规划假设：既有源端馈线接邻站新出线，占用受站备用间隔'))
    used = {e[k] for e in edges for k in ('a', 'b')}
    inp['network'] = dict(feeders={f: feeders[f] for f in sorted(used)}, edges=edges,
                          source_file=str(source.relative_to(ROOT)), source_sha256=sha(source),
                          existing_credit='无已核实市区既有10kV边，不按40%授予免费全区转供',
                          evidence='上级电源及间隔为源表；馈线负荷分配、热限、10kV可连接性为假设')
    inp['initial_transfer_fraction_usage'] = '40%保留为原方案参数记录；本轮转供由候选馈线、余量和投运决定，不生成免费能力'
    return evidence


def prepare():
    official = [r for r in read_official_annual(DEFAULT_SOURCE) if r['voltage_kv'] == 110]
    OUT.mkdir(parents=True, exist_ok=True)
    write_csv(OUT / '原表容量负荷.csv', official)
    audit = []
    for region in ('city', 'pizhou'):
        out = OUT / region; out.mkdir(exist_ok=True)
        inp = json.loads((SEEDS / f'{region}.json').read_text())
        actual = {str(r['year']): r for r in official if r['region_id'] == inp['region_id']}
        inp['annual_observed_peak_mw'] = {y: r['reported_downward_load_mw'] for y, r in actual.items()}
        inp['annual_observed_capacity_mva'] = {y: r['capacity_mva'] for y, r in actual.items()}
        inp['official_source_sha256'] = sha(DEFAULT_SOURCE)
        inp['storage_usable_energy_fraction'] = .76
        inp['scope'] = '原表年度负荷+归档正向峰值及低净负荷静态工况；各站单台主变N-1；无时序仿真'
        inp['rigid_rule'] = 'C_y<=max(C_(y-1),2*P_y)，C非减；仅允许保留以前已建容量形成的被动超限'
        inp['simulation_capacity_note'] = '保持原共同仿真起点；不等同原表实际存量，不以实际历史扩建叠加仿真措施'
        inp['storage_basis'] = '静态有效功率min(额定功率,额定电量×0.76/原方案持续时间)，不声称通过充电调度'
        if region == 'city':
            write_csv(out / '市区候选依据.csv', city_network(inp))
        source = AUDIT / f'{region}_2025_station_values_at_district_scenes.csv'
        selected = [r for r in rows(source) if r['voltage_kv'] == '110']
        key = 'model_station_id' if region == 'city' else 'station_id'
        kinds = sorted({r['scenario_kind'] for r in selected})
        forward = 'district_forward_peak'
        source_forward = sum(float(r['station_net_load_mw']) for r in selected if r['scenario_kind'] == forward)
        data = {kind: {r[key]: float(r['station_net_load_mw']) for r in selected if r['scenario_kind'] == kind} for kind in kinds}
        assert all(set(v) == set(inp['stations']) for v in data.values())
        # 静态分配模板：相同年度比例作用于两种典型净负荷工况，非实测历史反送曲线。
        scen = []
        for y in inp['years']:
            factor = inp['annual_observed_peak_mw'][str(y)] / source_forward
            for kind in kinds:
                for sid in inp['stations']:
                    scen.append(dict(year=y, scenario=kind, station=sid, power_mw=data[kind][sid] * factor,
                                     basis='归档静态分配模板按实际年度降压峰值校准；反向同倍率为沿用的净负荷假设'))
            old = inp['initial_peak_mw'] * (1 + inp['growth_rate']) ** (y - 2021)
            a = actual[str(y)]
            audit.append(dict(region=inp['region_name'], year=y, actual_capacity_mva=a['capacity_mva'],
                              actual_peak_mw=a['reported_downward_load_mw'], source_clr=a['reported_clr'],
                              old_growth_peak_mw=old, difference_mw=old-a['reported_downward_load_mw'],
                              preserved_initial_capacity_mva=inp['initial_capacity_mva']))
        inp['static_source_file'] = str(source.relative_to(ROOT)); inp['static_source_sha256'] = sha(source)
        inp['static_scene_limit'] = '仅两个区县典型工况，不覆盖所有站局部峰值或全年运行；费用变化同时受工况简化影响'
        dump(out / 'inputs.json', inp); write_csv(out / '静态工况.csv', scen)
        write_csv(out / '候选联络.csv', inp['network']['edges'])
        print('PREPARED', region, len(inp['network']['edges']), 'edges', flush=True)
    write_csv(OUT / '负荷输入差异.csv', audit)


def actual_limits(inp, p, y):
    ids = list(inp['stations']); idx = {s: j for j, s in enumerate(ids)}
    annual = inp['annual_observed_peak_mw'][str(y)] / inp['annual_observed_peak_mw']['2025']
    scale2025 = inp['annual_observed_peak_mw']['2025'] / inp['initial_peak_mw']
    donor = {}; receive = {}
    for f, r in inp['network']['feeders'].items():
        peak = inp['stations'][r['station']]['base_forward_mw'] * scale2025
        donor[f] = r['maximum_load_proxy_2025_mw'] * max(0., float(p[idx[r['station']]])) / peak
        receive[f] = max(0., r['thermal_mw'] - r['maximum_load_proxy_2025_mw'] * annual)
    return donor, receive


def capacity_terms(inp, meta, y):
    c = {v: r for (s, slot, yy, r), v in meta['x'].items() if yy == y}
    c.update({v: 50 for (s, yy), v in meta['third'].items() if yy == y})
    return c


def retain_existing_upper(m, current, previous, initial, threshold, big_m):
    """精确析取：不扩容或新增后仍<=2P，非历史峰值容量额度结转。"""
    branch = m.variable(0)
    increase = dict(current)
    for v, c in previous.items(): increase[v] = increase.get(v, 0) - c
    increase[branch] = -big_m
    m.constraint(increase, upper=0 if previous else initial)
    m.constraint({**current, branch: big_m}, upper=threshold + big_m)


def build(inp, scenes, case):
    # 构建原有资产和明确候选网络；不创建旧的逐时运行约束。
    m, meta = network.build(inp, np.zeros((1, len(inp['stations']))), [], case)
    original_rows = {}
    for row, col, value in m.entries:
        original_rows.setdefault(row, {})[col] = value
    caps = {}; flowmeta = []
    for y in inp['years']:
        cap = capacity_terms(inp, meta, y)
        matches = [r for r, terms in original_rows.items() if terms == cap]
        assert len(matches) == 1, matches
        row = matches[0]; peak = inp['annual_observed_peak_mw'][str(y)]
        m.constraint_lower[row] = (2.001 if case.startswith('elastic') else guide_range(inp['growth_rate'])[0]) * peak
        m.constraint_upper[row] = inp['physical_capacity_upper_mva']
        meta['peaks'][y] = peak
        if case.startswith('rigid'):
            previous = capacity_terms(inp, meta, y-1) if y > inp['years'][0] else {}
            retain_existing_upper(m, cap, previous, inp['initial_capacity_mva'], 2*peak, inp['physical_capacity_upper_mva'])
        for sid, s in inp['stations'].items():
            # 上限仅限定搜索范围，涵盖实际年度非单调负荷。
            if not case.endswith('no_storage'):
                maxp = max(abs(float(p[list(inp['stations']).index(sid)])) for p in scenes[y])
                m.upper_bounds[meta['modules'][sid,y]] = math.ceil(maxp * max(s['forward_duration_h'], s['reverse_duration_h'], 2.15) / (.215*.76)) + 1
            full = {v:r for (ss, slot, yy, r),v in meta['x'].items() if ss == sid and yy == y}
            if (sid,y) in meta['third']: full[meta['third'][sid,y]] = 50
            for f in range(-1,len(inp['stations'])):
                variants = [full]
                if f == list(inp['stations']).index(sid):
                    c = {v:r for (ss,slot,yy,r),v in meta['x'].items() if ss==sid and yy==y and slot!=1}
                    if (sid,y) in meta['third']: c[meta['third'][sid,y]] = 50
                    variants = [c]
                    if s['allow_third'] and max(s['initial_units']) < 50:
                        c3 = {v:r for (ss,slot,yy,r),v in meta['x'].items() if ss==sid and yy==y}
                        one = m.variable(0); m.lower_bounds[one] = 1
                        c3[one] = 50; c3[meta['third'][sid,y]] = -50; variants.append(c3)
                caps[sid,y,f] = variants
        for h,p in enumerate(scenes[y]):
            for f in range(-1,len(inp['stations'])):
                start = len(flowmeta)
                network.add_network_scene(m, inp, p, y, f, caps, meta['modules'], meta['edgevars'], flowmeta)
                flowmeta[start:] = [(yy,h,ff,e,d,v) for yy,ff,e,d,v in flowmeta[start:]]
    meta['network_flows'] = flowmeta
    return m,meta


def load_scenes(root, inp):
    rr = rows(root / '静态工况.csv'); scenes = {}
    for y in inp['years']:
        scenes[y] = np.array([[float(next(r['power_mw'] for r in rr if int(r['year']) == y and r['scenario'] == k and r['station'] == sid))
                               for sid in inp['stations']] for k in sorted({r['scenario'] for r in rr})])
        assert abs(scenes[y][0].sum() - inp['annual_observed_peak_mw'][str(y)]) < 1e-5
    return scenes


def run(region, case, seconds=180):
    root = OUT / region; out = root / case; out.mkdir(exist_ok=True)
    inp = json.loads((root/'inputs.json').read_text()); scenes = load_scenes(root,inp)
    network.limits = actual_limits  # 仅本进程替换年度缩放；历史模块文件不改。
    m,meta = build(inp,scenes,case)
    stage = 'static'; warm = root/'no_warm_start'
    if (out/'stations.csv').exists():
        warm = out
        attempt = 2
        while (out/f'static_{attempt:02}_solver.json').exists(): attempt += 1
        stage = f'static_{attempt:02}'
        old = json.loads((out/'summary.json').read_text())
        assert old['input_sha256'] == sha(root/'inputs.json'), '输入变更应使用新的结果目录'
        dump(out/f'before_{stage}_summary.json',old)
        m.constraint({i:c for i,c in enumerate(m.costs) if c},upper=old['cost_npv_10k']+1e-5)
    warm_inp = copy.deepcopy(inp); warm_inp['storage_usable_energy_fraction']=1.
    z,result = solve(m,out,stage,seconds,.01,warm_inp,meta,warm)
    if warm == out and result.get('feasible_incumbent'):
        bounds=[v for v in (result.get('dual_bound'),old['solver'].get('dual_bound')) if v is not None]
        if bounds:
            result['dual_bound']=max(bounds)
            result['mip_gap']=max(0.,(result['objective']-result['dual_bound'])/result['objective'])
            result['bound_basis']='历次相同输入求解的最强有效下界'
    summary = dict(region=region,case=case,solver=result,scope=inp['scope'])
    if z is None or not result.get('feasible_incumbent'):
        if warm == out:
            print('RETAIN previous verified incumbent',region,case,flush=True);return
        dump(out/'summary.json',summary); return
    checks = []
    for y in inp['years']:
        units,n,ends,_ = assets_model.assets(inp,z,meta,y)
        built = {e for (e,yy),v in meta['edgevars'].items() if yy==y and z[v]>.5}
        check,violations = network.check_year(inp,scenes[y],y,units,n,built)
        assert not violations, check
        check['static_scenario_contingency_pairs'] = check.pop('total_hour_contingency_pairs')
        checks.append(check)
    summary,stations,projects = network.extract(inp,None,case,z,meta,result,[],checks)
    summary.update(region=region,scope=inp['scope'],static_scene_limit=inp['static_scene_limit'],
                   static_network_checks=summary.pop('full_hour_network_checks'),input_sha256=sha(root/'inputs.json'))
    np.savez_compressed(out/'chosen_solution.npz',solution=z)
    previous = inp['initial_capacity_mva']
    for a in summary['annual']:
        y = a['year']; c = a['capacity_mva']; p = a['peak_mw']
        a.update(net_capacity_increment_mva=c-previous,actual_capacity_mva=inp['annual_observed_capacity_mva'][str(y)],
                 actual_clr=inp['annual_observed_capacity_mva'][str(y)]/p,
                 passive_exceedance=case.startswith('rigid') and c > 2*p+1e-5)
        if case.startswith('rigid'): assert c <= max(previous,2*p)+1e-5
        else: assert c >= 2.001*p-1e-5
        assert c >= previous-1e-5
        previous = c
    dump(out/'summary.json',summary); write_csv(out/'years.csv',summary['annual'])
    write_csv(out/'stations.csv',stations); write_csv(out/'projects.csv',projects)
    print('COMPLETE',region,case,summary['cost_npv_10k'],flush=True)


if __name__ == '__main__':
    p=argparse.ArgumentParser();p.add_argument('--prepare',action='store_true')
    p.add_argument('--region',choices=['city','pizhou']);p.add_argument('--case',choices=CASES)
    p.add_argument('--seconds',type=int,default=180);a=p.parse_args()
    if a.prepare: prepare()
    if a.region and a.case: run(a.region,a.case,a.seconds)
