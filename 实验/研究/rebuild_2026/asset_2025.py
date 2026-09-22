"""从 2025 年原始设备明细读取三层主变，核对年度统计容量。"""

import argparse
import csv
from collections import defaultdict
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from .official_annual import DEFAULT_TARGET as ANNUAL_CSV
from .official_annual import LAYERS, STUDY_DIR


SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/110（35）kv设备明细.xlsx"
OUTPUT_DIR = Path(__file__).resolve().parent / "source_audit"
SHEETS = {
    110: ("110千伏变电站1", 14, 15, 17),
    35: ("35千伏变电站1", 12, 13, 15),
}


def read_assets(source: Path = SOURCE) -> list[dict]:
    """逐设备行保留来源；开关站的零容量占位行不算在役主变。"""
    source = Path(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    source_label = str(source.resolve().relative_to(STUDY_DIR))
    workbook = load_workbook(source, read_only=True, data_only=True)
    assets = []
    try:
        for voltage_kv, (sheet_name, id_col, capacity_col, year_col) in SHEETS.items():
            sheet = workbook[sheet_name]
            for source_row, cells in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
                region_id, station_id = cells[1], cells[3]
                if (region_id, voltage_kv) not in LAYERS or not station_id:
                    continue
                capacity = float(cells[capacity_col] or 0)
                assets.append(
                    {
                        "region_id": region_id,
                        "voltage_kv": voltage_kv,
                        "station_id": station_id,
                        "transformer_id": cells[id_col],
                        "capacity_mva": capacity,
                        "station_commission_year": cells[8],
                        "transformer_commission_year": cells[year_col],
                        "asset_kind": "transformer" if capacity > 0 else "zero_capacity_placeholder",
                        "source_file": source_label,
                        "source_sheet": sheet_name,
                        "source_row": source_row,
                        "source_sha256": source_hash,
                    }
                )
    finally:
        workbook.close()
    return assets


def reconcile_capacity(assets: list[dict], annual_csv: Path = ANNUAL_CSV) -> list[dict]:
    with Path(annual_csv).open(encoding="utf-8-sig", newline="") as handle:
        annual = {
            (row["region_id"], int(row["voltage_kv"])): row
            for row in csv.DictReader(handle)
            if int(row["year"]) == 2025
        }
    by_layer = defaultdict(list)
    for asset in assets:
        by_layer[(asset["region_id"], asset["voltage_kv"])].append(asset)
    results = []
    for key in sorted(LAYERS):
        rows = by_layer[key]
        live = [row for row in rows if row["asset_kind"] == "transformer"]
        asset_capacity = sum(row["capacity_mva"] for row in live)
        reported_capacity = float(annual[key]["capacity_mva"])
        results.append(
            {
                "region_id": key[0],
                "voltage_kv": key[1],
                "year": 2025,
                "asset_station_count": len({row["station_id"] for row in live}),
                "asset_transformer_count": len(live),
                "zero_capacity_placeholder_rows": len(rows) - len(live),
                "asset_capacity_mva": asset_capacity,
                "official_capacity_mva": reported_capacity,
                "asset_minus_official_mva": round(asset_capacity - reported_capacity, 9),
                "capacity_status": "match" if abs(asset_capacity - reported_capacity) < 1e-6 else "unresolved_difference",
                "official_source_row": annual[key]["source_row"],
            }
        )
    return results


def write_csv(rows: list[dict], target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="读取 2025 主变设备并核对年度统计容量")
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--annual", type=Path, default=ANNUAL_CSV)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    assets = read_assets(args.source)
    checks = reconcile_capacity(assets, args.annual)
    write_csv(assets, args.output_dir / "asset_transformers_2025.csv")
    write_csv(checks, args.output_dir / "asset_capacity_reconciliation_2025.csv")
    print(f"已写入 {len(assets)} 条设备行、{len(checks)} 条容量核对记录")


if __name__ == "__main__":
    main()
