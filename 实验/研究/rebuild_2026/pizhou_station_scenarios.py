"""按已复核的逐列映射构造邳州 110/35 kV 逐站静态场景。"""

import argparse
import csv
from collections import defaultdict
from datetime import datetime
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from .city_mapping_audit import write_csv
from .city_scenarios import scenario_metrics
from .pizhou_scenarios import AUDIT_DIR, SOURCE


MAPPING_EVIDENCE = AUDIT_DIR / "pizhou_2025_mapping_evidence.csv"
AGGREGATE_SCENARIOS = AUDIT_DIR / "pizhou_2025_scenarios_evidence.csv"


def read_pizhou_station_series(
    hourly_source: Path = SOURCE,
    mapping_source: Path = MAPPING_EVIDENCE,
) -> dict[tuple[int, str], dict[datetime, float]]:
    with Path(mapping_source).open(encoding="utf-8-sig", newline="") as handle:
        mapping = [row for row in csv.DictReader(handle) if row["station_status"] == "supported_for_station_aggregation"]
    columns = defaultdict(list)
    for row in mapping:
        columns[(int(row["resolved_voltage_kv"]), row["resolved_station_id"])].append(int(row["source_column"]) - 1)
    workbook = load_workbook(hourly_source, read_only=True, data_only=True)
    series = {key: {} for key in columns}
    try:
        for cells in workbook["Sheet3"].iter_rows(min_row=3, values_only=True):
            raw_time = str(cells[0] or "")
            if not raw_time.startswith("2025-"):
                continue
            hour = datetime.fromisoformat(raw_time)
            for key, indices in columns.items():
                values = [cells[index] for index in indices]
                if all(value not in (None, "") for value in values):
                    series[key][hour] = round(sum(float(value) for value in values), 6)
    finally:
        workbook.close()
    return series


def build_station_scenarios(
    hourly_source: Path = SOURCE,
    mapping_source: Path = MAPPING_EVIDENCE,
    aggregate_source: Path = AGGREGATE_SCENARIOS,
) -> tuple[list[dict], list[dict]]:
    series = read_pizhou_station_series(hourly_source, mapping_source)
    with Path(aggregate_source).open(encoding="utf-8-sig", newline="") as handle:
        aggregate = {int(row["voltage_kv"]): row for row in csv.DictReader(handle)}
    source_hash = sha256(Path(hourly_source).read_bytes()).hexdigest()
    mapping_hash = sha256(Path(mapping_source).read_bytes()).hexdigest()
    station_rows = [
        {
            "region_id": "QX-00005",
            "voltage_kv": voltage,
            "year": 2025,
            "station_id": station_id,
            "observed_hours": len(values),
            **scenario_metrics(values),
            "hourly_source_sha256": source_hash,
            "mapping_source_sha256": mapping_hash,
        }
        for (voltage, station_id), values in sorted(series.items())
    ]
    at_district_scenes = [
        {
            "region_id": "QX-00005",
            "voltage_kv": voltage,
            "year": 2025,
            "scenario_kind": kind,
            "scenario_time": time.isoformat(sep=" "),
            "station_id": station_id,
            "station_net_load_mw": values[time],
            "district_net_load_mw": float(aggregate[voltage]["forward_peak_mw"])
            if kind == "district_forward_peak" else -float(aggregate[voltage]["reverse_peak_mw"]),
            "data_status": "observed",
        }
        for (voltage, station_id), values in sorted(series.items())
        for kind, time in (
            ("district_forward_peak", datetime.fromisoformat(aggregate[voltage]["forward_peak_time"])),
            ("district_reverse_peak", datetime.fromisoformat(aggregate[voltage]["reverse_peak_time"])),
        )
    ]
    return station_rows, at_district_scenes


def main() -> None:
    parser = argparse.ArgumentParser(description="输出邳州逐站静态场景")
    parser.add_argument("--output-dir", type=Path, default=AUDIT_DIR)
    args = parser.parse_args()
    station_rows, at_district_scenes = build_station_scenarios()
    write_csv(station_rows, args.output_dir / "pizhou_2025_station_scenarios_evidence.csv")
    write_csv(at_district_scenes, args.output_dir / "pizhou_2025_station_values_at_district_scenes.csv")
    print(f"邳州逐站场景 {len(station_rows)} 站，片区峰值时刻站值 {len(at_district_scenes)} 条")


if __name__ == "__main__":
    main()
