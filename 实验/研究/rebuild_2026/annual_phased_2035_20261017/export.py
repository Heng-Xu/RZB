"""按表5-4列项生成两地区逐年左右对照表，附当年增量。"""
import csv
import json
from decimal import Decimal, ROUND_HALF_UP

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

from . import model as m, prepare as p
from .verify import CASES
from ..annual_transfer_2035_20261017.export import FIELDS, fixed, markdown

LABEL={'rigid':'刚性','elastic_capacity':'弹性'}


def clean(v):
    return float(Decimal(str(v)).quantize(Decimal('0.000000001'),rounding=ROUND_HALF_UP)) if isinstance(v,float) else v


def decorate(ws,headers=1):
    ws.freeze_panes=f'C{headers+1}';ws.sheet_view.showGridLines=False
    thin=Side(style='hair',color='CED6DF')
    for row in ws:
        for cell in row:
            cell.font=Font(name='方正仿宋_GBK',size=10,bold=cell.row<=headers,color='FFFFFF' if cell.row<=headers else '172D41')
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            cell.border=Border(bottom=thin)
            if cell.row<=headers:cell.fill=PatternFill('solid',fgColor='244B70')
            elif cell.row%2:cell.fill=PatternFill('solid',fgColor='F0F4F8')
            if isinstance(cell.value,float):
                name=str(ws.cell(headers,cell.column).value or '')
                cell.number_format='0.00%' if '增长率' in name or '比例' in name else '0.00'
    for row in range(1,ws.max_row+1):ws.row_dimensions[row].height=34 if row<=headers else 24
    for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=17
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows=f'1:{headers}'


def flat(book,name,data):
    assert data,name
    ws=book.create_sheet(name);keys=list(data[0]);ws.append(keys)
    for row in data:ws.append([clean(row.get(k)) for k in keys])
    decorate(ws);ws.auto_filter.ref=ws.dimensions
    return ws


def main():
    proof=json.loads((p.REVIEW/'逐年方案独立核验.json').read_text());assert proof['verified']
    book=Workbook();book.remove(book.active);plans={};allrows=[];stations=[];projects=[];switches=[];checks=[];growth=[];comparisons=[];exceptions=[];baselines=[]
    text=['# 两地区2021—2035分期增容方案（左右对照）','',
        '2021保留原地区容量，2022起逐年补强。下表左为刚性、右为弹性；储能为年末额定功率，联络为自2022年起累计新增条数。容量、储能与联络均为本次模型规划结果。', '']
    for region,name in p.REGIONS.items():
        inp,powers=p.load(region);ids=list(inp['stations']);aa=m.arcs(inp)
        for case in CASES:
            root=p.RAW/region/case;plans[region,case]=json.loads((root/'summary.json').read_text())
            s=plans[region,case]
            if case=='rigid':
                baseline=json.loads((root/'baseline_check.json').read_text())
                baselines.extend(dict(地区=name,年份=2021,故障站=v['failed'],校核小时数=v['checked_hours'],
                    约束缺口_MW=v['max_deficit_mw'],约束不满足小时数=v['violating_hours']) for v in baseline['records'])
            allrows.extend(dict(地区=name,方案=LABEL[case],**{FIELDS[k]:v for k,v in r.items()}) for r in s['annual'])
            with (root/'stations.csv').open(encoding='utf-8-sig') as f:
                stations.extend(dict(地区=name,方案=LABEL[case],年份=int(r['year']),变电站=r['station'],
                    主变1_MVA=float(r['unit_1_mva']),主变2_MVA=float(r['unit_2_mva']),主变3_MVA=float(r['unit_3_mva']),
                    储能柜数=int(r['storage_modules']),储能额定功率_MW=int(r['storage_modules'])*.1) for r in csv.DictReader(f))
            emap={e['id']:e for e in inp['network']['edges']}
            with (root/'projects.csv').open(encoding='utf-8-sig') as f:
                for r in csv.DictReader(f):
                    e=emap[r['id']];projects.append(dict(地区=name,方案=LABEL[case],投运年=int(r['year']),联络编号=r['id'],
                        站1=e['station_a'],站2=e['station_b'],容量_MW=e['capacity_mw'],依据=e['basis']))
            for y in inp['years']:
                r=json.loads((root/str(y)/'adopted.json').read_text())
                checks.extend(dict(地区=name,方案=LABEL[case],年份=y,故障站=v['failed'],校核小时数=v['checked_hours'],
                    最大缺口_MW=v['max_deficit_mw'],失败小时数=v['violating_hours']) for v in r['check']['records'])
                for sid in ids:
                    n=sum(c for k,c in r['extra'].items() if aa[int(k)]['source']==sid)
                    switches.append(dict(地区=name,方案=LABEL[case],年份=y,站名=sid,事故转出区段数=n,
                        全站区段数=inp['block_count'][sid],转出比例=n/inp['block_count'][sid],限额区段数=m.block_limit(inp,sid)))
        rigid=plans[region,'rigid']['annual'];elastic=plans[region,'elastic_capacity']['annual']
        ws=book.create_sheet(f'{name}逐年对照',0 if region=='city' else 1)
        ws.append(['年份','规划峰值/MW','刚性方案',None,None,None,None,'弹性方案',None,None,None,None])
        ws.merge_cells('A1:A2');ws.merge_cells('B1:B2');ws.merge_cells('C1:G1');ws.merge_cells('H1:L1')
        cols=['主变容量/MVA','当年净增/MVA','容载比','储能额定功率/MW','累计新线/条']
        ws.append([None,None,*cols,*cols]);md=[];compact=[]
        for a,b in zip(rigid,elastic):
            assert a['year']==b['year'] and a['common_floor']==b['common_floor'] and a['peak_mw']==b['peak_mw']
            keys=['capacity_mva','net_capacity_increment_mva','clr','storage_nameplate_mw','cumulative_lines']
            values=[a['year'],a['peak_mw'],*[a[k] for k in keys],*[b[k] for k in keys]]
            ws.append([clean(v) for v in values])
            compact.append(dict(年份=a['year'],规划峰值_MW=a['peak_mw'],
                **{'刚性_'+FIELDS[k]:a[k] for k in keys},**{'弹性_'+FIELDS[k]:b[k] for k in keys}))
            md.append([a['year'],fixed(a['peak_mw']),fixed(a['capacity_mva']),fixed(a['clr']),fixed(a['storage_nameplate_mw']),a['cumulative_lines'],
                fixed(b['capacity_mva']),fixed(b['clr']),fixed(b['storage_nameplate_mw']),b['cumulative_lines']])
            if a['year']>2021 and (b['capacity_mva']<a['capacity_mva']-1e-6 or b['storage_nameplate_mw']>a['storage_nameplate_mw']+1e-6 or b['cumulative_lines']>a['cumulative_lines']):exceptions.append(dict(地区=name,年份=a['year']))
            comparisons.append(dict(地区=name,年份=a['year'],刚性容量_MVA=a['capacity_mva'],弹性容量_MVA=b['capacity_mva'],
                净增容量差_MVA=b['capacity_mva']-a['capacity_mva'],刚性储能_MW=a['storage_nameplate_mw'],弹性储能_MW=b['storage_nameplate_mw'],
                刚性联络_条=a['cumulative_lines'],弹性联络_条=b['cumulative_lines'],共同下限=a['common_floor']))
        decorate(ws,2);ws.column_dimensions['A'].width=9;ws.column_dimensions['B'].width=17
        ws.print_area=f'A1:L{ws.max_row}';p.write_csv(p.PUBLIC/f'{name}逐年左右对照.csv',compact)
        text.extend([f'表{1 if region=="city" else 2} {name}逐年容量及措施','',markdown(['年份','规划峰值/MW','刚性容量/MVA','刚性容载比','刚性储能/MW','刚性累计新线/条',
            '弹性容量/MVA','弹性容载比','弹性储能/MW','弹性累计新线/条'],md),''])
        forecasts=p.forecasts(region);lastpeak=None;lastforecast=None
        for a in rigid:
            y=a['year'];f=next((v for v in forecasts if v['year']==y),None)
            growth.append(dict(地区=name,年份=y,同口径规划峰值_MW=a['peak_mw'],规划峰值较上年增量_MW=None if lastpeak is None else a['peak_mw']-lastpeak,
                预测总负荷_MW=None if f is None else f['forecast_load_mw'],
                预测总负荷较上年增量_MW=None if f is None or lastforecast is None else f['forecast_load_mw']-lastforecast,
                负荷同比增长率=None if f is None else f['load_yoy'],共同容载比下限=a['common_floor']))
            lastpeak=a['peak_mw'];lastforecast=None if f is None else f['forecast_load_mw']
    assumptions=[dict(项目=k,说明=v) for k,v in [
        ('起点与年度','2021市区3113.50MVA、邳州2027.00MVA固定；逐站组合为归档标定。2022—2035承接全部投运资产，位置全期最多更换一次。'),
        ('源荷输入','与上一版逐小时输入完全相同；2021—2025沿用统计峰值；2026—2035按最新负荷、光伏倍数外推。缺年按已知节点等增量插值；光伏仅取无量纲比例。'),
        ('正常运行','本站供区由本站承担，正常跨站转供为零。'),
        ('事故转供','既有与新增联络共用每站20%源荷区段上限，区段数向下取整。全年固定预案，仅故障站原属区段直接转给邻站；保留馈线组、共享路径和受端余量限制。'),
        ('限额依据','20%是本轮偏增容工程情景假设，经结果导向选择用于比较，不作为规范强制限值或实际已核定转供率。'),
        ('远景定容','替换采购按本站2035正反向峰值/0.95选择31.5、50、63、80MVA对应规格，封顶80MVA；保留原机不采购，原允许第三台仍按50MVA。刚弹规则相同。'),
        ('年度建设节奏','刚弹同用每年净增容量不超过上年容量5%的建设上限。各自保留上一版2035主变工程清单，按预测负荷增长趋势安排累计投运进度，允许50MVA整台工程余量，终年完成原清单。刚性同时服从2.0控制。各年负荷输入保持原值，5%及50MVA为建设排程研究假设。'),
        ('刚弹边界','同地区同年度负荷及容载比下限一致；刚性容量控制为max(上年容量,规划峰值×2.0)，弹性取消此上限。'),
        ('弹性优选','先压低当年新增联络与储能费用，在自身候选配套费用20%范围内优先减少联络，再比较总费用。该容差为联络简化偏好；主变增容投资不受原5%总费用包络限制，未施加刚弹排序约束。'),
        ('储能口径','整柜0.10MW/0.215MWh；有效支撑=min(额定功率,额定电量×0.76/持续时长)，市区正向6h、邳州2h；反向沿用原时长。'),
        ('费用','以2021为现值基年，新增工程评价至2041；主变/线路/储能寿命40/40/25年、维护4.2%/2.2%/3%，折现6%，采用既有价格。'),
        ('新购与净增','新购主变容量为替换及扩建采购铭牌容量之和，净增容量扣除被替换原机容量；主表增容为净增。'),
        ('适用范围','候选接线与源荷区段仍为研究代理，8760小时含留有标识的补值；为静态工况可行方案，未做连续SOC调度和全期联合全局寻优。2021固定基线缺口单列。'),
    ]]
    pacing=[];comparison=[]
    for region,name in p.REGIONS.items():
        for case in CASES:
            old=json.loads((p.parent.RAW/region/case/'summary.json').read_text())
            current=plans[region,case]
            for a,b in zip(old['annual'],current['annual']):
                pacing.append(dict(地区=name,方案=LABEL[case],年份=b['year'],
                    原方案年末容量_MVA=a['capacity_mva'],分期年末容量_MVA=b['capacity_mva'],
                    原方案当年净增_MVA=a['net_capacity_increment_mva'],分期当年净增_MVA=b['net_capacity_increment_mva'],
                    原方案储能_MW=a['storage_nameplate_mw'],分期储能_MW=b['storage_nameplate_mw'],
                    原方案联络_条=a['cumulative_lines'],分期联络_条=b['cumulative_lines'],
                    原方案累计费用_万元=a['cumulative_lifecycle_npv_10k'],分期累计费用_万元=b['cumulative_lifecycle_npv_10k']))
            if case=='elastic_capacity':
                first,last=current['annual'][1],current['annual'][-1]
                comparison.append([name,fixed(old['annual'][1]['net_capacity_increment_mva']),fixed(first['net_capacity_increment_mva']),
                    sum(v['net_capacity_increment_mva']>1e-6 for v in old['annual'][1:]),
                    sum(v['net_capacity_increment_mva']>1e-6 for v in current['annual'][1:]),
                    fixed(current['cost_npv_10k']-old['cost_npv_10k']),f'{current["cost_npv_10k"]/old["cost_npv_10k"]-1:.2%}'])
    datasets={'分期效果对照':pacing,'逐年新增及费用':allrows,'逐站逐年配置':stations,'联络逐项工程':projects,'站级转供限额核对':switches,
        '负荷与逐年增量':growth,'全年安全核验':checks,'2021固定基线检验':baselines,'约束与口径':assumptions}
    for name,data in datasets.items():flat(book,name,data);p.write_csv(p.PUBLIC/(name+'.csv'),data)
    path=p.PUBLIC/'两地区逐年容量及措施_分期增容左右对照.xlsx';book.save(path)
    readback=load_workbook(path,data_only=True);cells=0
    for source in book:
        target=readback[source.title];assert source.max_row==target.max_row and source.max_column==target.max_column
        for row in source:
            for cell in row:
                actual=target.cell(cell.row,cell.column).value;expected=cell.value
                if isinstance(expected,float):assert abs(actual-expected)<max(1e-8,abs(expected)*1e-12),(source.title,cell.coordinate)
                else:assert actual==expected,(source.title,cell.coordinate,actual,expected)
                cells+=1
    readback.close()
    p.dump(p.REVIEW/'工作簿回读.json',dict(verified=True,cells=cells,path=str(path.relative_to(p.ROOT)),sha256=p.sha(path)))
    text.extend(['注：储能为额定功率；累计新线表示累计新增10kV联络。工作簿另列每年的净增容量、新购容量、储能电量和新增工程费用。', '',
        '本轮保留各自2035主变工程清单，按预测负荷增长趋势分期投运，累计进度允许50MVA整台工程余量；每年净增容量不超过上年5%。保留远景选型与站级20%事故转供限额，储能和联络逐年重新选配。5%、50MVA与20%为研究情景假设。年度负荷与光伏输入保持原值。', '',
        f'2022—2035共{proof["full_hour_pairs"]:,}个小时—工况组合通过固定预案检查。2021是固定起点，零切换供区下的原有缺口另存核验记录。',''])
    text.append(markdown(['地区','刚性费用现值/万元','弹性费用现值/万元','弹减刚/万元'],[
        [name,fixed(plans[region,'rigid']['cost_npv_10k']),fixed(plans[region,'elastic_capacity']['cost_npv_10k']),
         fixed(plans[region,'elastic_capacity']['cost_npv_10k']-plans[region,'rigid']['cost_npv_10k'])] for region,name in p.REGIONS.items()]))
    text.extend(['','弹性方案的分期影响','',markdown(['地区','原方案2022净增/MVA','分期2022净增/MVA','原增容年份数','分期增容年份数','全期费用增加/万元','费用增加率'],comparison),'',
        '分期方案按每年设备配置重新校核。原年度峰值有升有降，整台工程可能形成短期容量平台；前期补强储能和联络逐年保留，分期带来的费用变化按实际工程计算。'])
    (p.PUBLIC/'逐年容量及措施_分期增容左右对照.md').write_text('\n'.join(text)+'\n',encoding='utf-8')
    p.dump(p.REVIEW/'结果序列核对.json',dict(verified=True,ranking_exceptions=exceptions,rows=comparisons,
        interpretation='真实模型输出逐年核对；未强制刚弹排序，若有例外据实列出'))
    p.dump(p.SOURCE/'展示口径.json',dict(primary='按用户要求交付表5-4风格的左右对照工作簿及Markdown表，不生成HTML或改正式报告',
        years=[2021,2035],regions=list(p.REGIONS.values()),data_sources=proof['solution_sources'],assumptions=assumptions))
    print('EXPORTED_SIDE_BY_SIDE',cells,'exceptions',exceptions,flush=True)


if __name__=='__main__':main()
