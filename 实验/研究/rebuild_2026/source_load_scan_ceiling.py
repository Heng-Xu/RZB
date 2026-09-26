"""复算源荷压力代理与研究性容载比扫描覆盖上限；不作承载力判定。"""

import argparse
import csv
import math
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .city_mapping_audit import write_csv
from .hourly_source_profile import OUTPUT_DIR
from .pizhou_station_scenarios import read_pizhou_station_series
from .pv_source_audit import HOURLY_SOURCE, SOURCE as MONTHLY_SOURCE
from .rooftop_area_audit import capacity_scenario_mw, read_roof_records, scope_areas


CITY_HOURLY = OUTPUT_DIR / "city_2025_aggregate_hourly.csv"
REFERENCE_R_GROSS = 2.07  # 李阳彤等（2026）光伏算例上缘；不是标准限值。
PROVINCIAL_GROWTH_STRESS = 2.0  # 江苏 2030 年分布式光伏较 2025 年翻番的政策目标；非区县配额。


def month_end_capacity_proxy(source: Path = MONTHLY_SOURCE) -> dict[str, float]:
    """逐月表无单位；10 MW/原表单位仅作跨表量级核对后的研究假设。"""
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"]
        return {
            "QX-00005": float(sheet.cell(31, 13).value) * 10,
            "QX-00007": float(sheet.cell(29, 13).value) * 10,
        }
    finally:
        workbook.close()


def monthly_capacity_proxy(source: Path = MONTHLY_SOURCE) -> dict[str, dict[int, float]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"]
        return {
            region: {month: float(sheet.cell(row, month + 1).value) * 10 for month in range(1, 13)}
            for region, row in (("QX-00005", 31), ("QX-00007", 29))
        }
    finally:
        workbook.close()


def pv_shape(source: Path = HOURLY_SOURCE) -> dict[datetime, float]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        return {
            cells[1]: float(cells[4])
            for cells in workbook["光伏8760"].iter_rows(min_row=2, values_only=True)
        }
    finally:
        workbook.close()


def pizhou_net_series() -> dict[datetime, float]:
    layer_stations = read_pizhou_station_series()
    common_hours = set.intersection(*(set(values) for values in layer_stations.values()))
    return {hour: sum(values[hour] for values in layer_stations.values()) for hour in common_hours}


def city_net_series(source: Path = CITY_HOURLY) -> dict[datetime, float]:
    with source.open(encoding="utf-8-sig", newline="") as handle:
        return {
            datetime.fromisoformat(row["hour"]): float(row["observed_mw"])
            for row in csv.DictReader(handle)
            if row["data_status"] == "observed" and row["observed_mw"]
        }


def scan_ceiling(gross_reference_r: float, gross_proxy_peak: float, net_peak: float) -> float:
    """仅保证扫描覆盖；向上取到 0.5 档，不是推荐值或技术限值。"""
    return math.ceil(2 * gross_reference_r * gross_proxy_peak / net_peak - 1e-12) / 2


def build_research_rows() -> list[dict]:
    monthly = monthly_capacity_proxy()
    installed = month_end_capacity_proxy()
    shape = pv_shape()
    rooftop_areas = scope_areas(read_roof_records())
    rows = []
    for region, net in (("QX-00005", pizhou_net_series()), ("QX-00007", city_net_series())):
        common_hours = set(net) & set(shape)
        net_hour = max(common_hours, key=lambda hour: net[hour])
        gross_hour = max(common_hours, key=lambda hour: net[hour] + monthly[region][hour.month] * shape[hour])
        net_peak = net[net_hour]
        gross_proxy_peak = net[gross_hour] + monthly[region][gross_hour.month] * shape[gross_hour]
        ratio = installed[region] / gross_proxy_peak
        roof_scope = "QX-00005" if region == "QX-00005" else "urban_core_proxy"
        roof_area = rooftop_areas[roof_scope]
        roof_pv_20 = capacity_scenario_mw(roof_area, 0.20)
        roof_pv_40 = capacity_scenario_mw(roof_area, 0.40)
        rows.append(
            {
                "region_id": region,
                "voltage_scope": "110_and_35_combined" if region == "QX-00005" else "110_sample_only",
                "observed_common_hours": len(common_hours),
                "net_peak_time": net_hour.isoformat(sep=" "),
                "net_peak_mw": round(net_peak, 6),
                "gross_proxy_peak_time": gross_hour.isoformat(sep=" "),
                "gross_proxy_peak_mw": round(gross_proxy_peak, 6),
                "pv_2025_dec_capacity_proxy_mw": round(installed[region], 6),
                "pv_to_gross_load_ratio_proxy": round(ratio, 6),
                "policy_double_pv_ratio_stress": round(PROVINCIAL_GROWTH_STRESS * ratio, 6),
                "rooftop_area_scope": roof_scope,
                "rooftop_gross_area_m2": round(roof_area, 3),
                "rooftop_pv_20pct_usable_scenario_mw": round(roof_pv_20, 3),
                "rooftop_pv_40pct_usable_scenario_mw": round(roof_pv_40, 3),
                "rooftop_20pct_nameplate_to_gross_load_proxy": round(roof_pv_20 / gross_proxy_peak, 6),
                "rooftop_40pct_nameplate_to_gross_load_proxy": round(roof_pv_40 / gross_proxy_peak, 6),
                "equivalent_roof_fraction_for_2x_installed_proxy": round(
                    PROVINCIAL_GROWTH_STRESS * installed[region] / capacity_scenario_mw(roof_area, 1.0), 6
                ),
                "rooftop_scope_warning": (
                    "administrative_pizhou_roofs_not_split_to_110_or_35"
                    if region == "QX-00005"
                    else "urban_core_three_districts_not_29_station_supply_boundary"
                ),
                "literature_gross_denominator_example_upper": REFERENCE_R_GROSS,
                "gross_to_net_peak_factor_proxy": round(gross_proxy_peak / net_peak, 6),
                "translated_research_reference": round(REFERENCE_R_GROSS * gross_proxy_peak / net_peak, 6),
                "capacity_load_scan_ceiling": scan_ceiling(REFERENCE_R_GROSS, gross_proxy_peak, net_peak),
                "status": "research_scan_coverage_only_not_standard_or_guide_certification",
            }
        )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="复算研究性源荷压力和容载比扫描覆盖上限")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    rows = build_research_rows()
    write_csv(rows, args.output_dir / "source_load_scan_ceiling_research.csv")
    for row in rows:
        print(row["region_id"], "源荷压力", row["policy_double_pv_ratio_stress"], "研究扫描上限", row["capacity_load_scan_ceiling"])


if __name__ == "__main__":
    main()
