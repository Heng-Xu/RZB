"""将已核验联合规划输出整理为逐年对照表和可读计算说明。"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .baseline_2021 import read_csv
from .joint_shared_measure_audit import audit


ROOT = Path(__file__).resolve().parents[3]
MODEL = Path(__file__).resolve().parent
SOURCE = MODEL / "outputs/joint_shared_measure/transfer_10pct_all_years"
DELIVERY = ROOT / "docs/2026-10-01三措施联合优化"
LABELS = {"pizhou": "邳州", "city": "市区", "rigid": "刚性", "elastic": "弹性"}


def main():
    check = audit(SOURCE)
    review = json.loads((SOURCE / "review.json").read_text())
    rows = read_csv(SOURCE / "annual_comparison.csv")
    DELIVERY.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "逐年对照"
    headers = ["区县", "方案", "年份", "负荷MW", "在役主变MVA", "容载比",
               "新增购置主变MVA", "新增储能MW", "新增储能MWh",
               "在役储能MW", "在役储能MWh", "新联络条数", "既有正常转供MW", "新线正常转供MW"]
    sheet.append(headers)
    for r in rows:
        sheet.append([LABELS[r["district"]], LABELS[r["scheme"]], int(r["year"])] +
                     [float(r[key]) for key in (
                         "net_peak_proxy_mw", "capacity_mva", "clr",
                         "new_transformer_purchase_mva", "new_storage_power_mw",
                         "new_storage_energy_mwh", "installed_storage_power_mw",
                         "installed_storage_energy_mwh", "new_lines_commissioned",
                         "existing_transfer_mw", "new_line_transfer_mw")])
    cost_sheet = book.create_sheet("费用对照")
    cost_sheet.append(["区县", "刚性全寿命现值万元", "弹性全寿命现值万元", "节省万元", "节省比例"])
    for label, cost in review["cost_npv_10k"].items():
        saving = cost["rigid"] - cost["elastic"]
        cost_sheet.append([LABELS[label], cost["rigid"], cost["elastic"], saving, saving / cost["rigid"]])
    rigid = check["rigid_npv_10k"]
    elastic = check["elastic_npv_10k"]
    cost_sheet.append(["合计", rigid, elastic, rigid - elastic, (rigid - elastic) / rigid])
    for district in ("pizhou", "city"):
        for scheme in ("rigid", "elastic"):
            station_sheet = book.create_sheet(LABELS[district] + LABELS[scheme] + "逐站措施")
            stations = read_csv(SOURCE / district / f"{scheme}_stations.csv")
            station_sheet.append(list(stations[0]))
            for r in stations:
                station_sheet.append(list(r.values()))
    for table in book:
        table.freeze_panes = "A2"
        table.auto_filter.ref = table.dimensions
        for cell in table[1]:
            cell.fill = PatternFill("solid", fgColor="245473")
            cell.font = Font(name="微软雅黑", color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        table.row_dimensions[1].height = 34
        for index in range(1, table.max_column + 1):
            table.column_dimensions[get_column_letter(index)].width = 19
        for line in table.iter_rows(min_row=2):
            for cell in line:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000" if table.title == "逐年对照" and cell.column == 6 else "0.00"
        if table.title == "费用对照":
            for line in table.iter_rows(min_row=2):
                line[-1].number_format = "0.00%"
    book.save(DELIVERY / "两区县逐年三措施优化结果.xlsx")
    text = [
        "# 2026-10-01 两区县三措施联合优化结果", "",
        "回滚状态提交为3011fdc，模型锚点为9月29日eb417e0。恢复同措施集后，同步优化邳州、市区的刚性与弹性四条路径。2021年同区县共用反事实起点，2022—2025年按原表年度负荷规划。", "",
        "刚性年度R≤2.0；弹性两区县每年R≥2.001，且市区每年低于邳州至少0.001、弹性每年高于本区县刚性至少0.001。这三项R关系是用户指定规划条件。费用高低没有加入强制排序约束，按同一案例价格和同一供电需求计算。", "",
        "三类措施均保留：主变购置与增容、配置储能、10 kV联络（既有转供和新建）。正常转供与单一事故恢复的实际使用量均限制为对应站负荷的10%，两方案采用相同限制。该10%是响应减少转供依赖的研究情景，不是导则规定或实测转移能力；原邳州既有能力28.2879%、市区50%不改写。新线投资仍按网络建设条件计入。", "",
        "储能整数柜数由峰值功率、持续时间和成本寻优，站级上界按供电任务推导，替代旧版固定5 MW的探索上界。第三台主变限于原表可识别预留位置；两区县初始设备与年度负荷继续保持原来源。", "",
        "| 区县 | 刚性现值／万元 | 弹性现值／万元 | 节省／万元 | 节省比例 |",
        "| --- | ---: | ---: | ---: | ---: |"]
    for label, cost in review["cost_npv_10k"].items():
        saving = cost["rigid"] - cost["elastic"]
        text.append(f'| {LABELS[label]} | {cost["rigid"]:.2f} | {cost["elastic"]:.2f} | {saving:.2f} | {saving / cost["rigid"]:.2%} |')
    text += [f"| 合计 | {rigid:.2f} | {elastic:.2f} | {rigid-elastic:.2f} | {(rigid-elastic)/rigid:.2%} |", "",
             "费用为2022—2025年新增措施的全寿命折现估算，排除2021存量。替换主变按新购铭牌容量计费；储能计入购置、运维和寿命期更新。两区县起点不同属于地区差异，同一区县刚性与弹性起点一致。", "",
             "| 区县 | 年份 | 刚性R | 弹性R | 刚性容量MVA | 弹性容量MVA |",
             "| --- | ---: | ---: | ---: | ---: | ---: |"]
    by_key = {(r["district"], r["scheme"], int(r["year"])): r for r in rows}
    for label in ("pizhou", "city"):
        for year in (2022, 2023, 2024, 2025):
            a, b = by_key[label, "rigid", year], by_key[label, "elastic", year]
            text.append(f'| {LABELS[label]} | {year} | {float(a["clr"]):.4f} | {float(b["clr"]):.4f} | {float(a["capacity_mva"]):.1f} | {float(b["capacity_mva"]):.1f} |')
    text += ["", "第一阶段最小化四条路径新增措施的总全寿命现值；第二阶段在总费用最多相差0.001万元的同等成本解中最大化弹性各年R之和；第三阶段固定选定设备布局，最小化该布局的连续转供量。固定单线程、种子0和同一版本求解器，保存输入、代码哈希与费用最优断点。", "",
             "在10%使用上限下，既有全站对等效通道可以承接新通道的实际转供；联络建设仍满足原网络目标条件。因此在取得最低费用解后固定同价线路布局，缩小后续设备寻优规模。这一等价处理依赖当前完整站对代理网络，工程拓扑变更后需重新验证。", "",
             "逐站正反向承载、整数储能、主变投运位置、转供上限、静态事故供电、年度容量不减、原表分母和新增费用均已独立复算。容量增量、转供与储能的完整明细见同目录工作簿。", "",
             "该结果属于规定条件下的静态规划解。市区负荷由29站样本外推，部分站点类别为研究假定；3 km新联络为等效通道成本场景，尚不具备实际路由与完整站址费用。每年超过2.0是规划要求，不能归因于实测承载需要自然产生的唯一解。", "",
             "复现：在实验/研究下运行 python -m rebuild_2026.joint_shared_measure --case transfer_10pct --breakthrough all_years；保留最优断点后可加 --resume。数值审查为 python -m rebuild_2026.joint_shared_measure_audit <输出目录>。"]
    (DELIVERY / "计算结果与约束说明.md").write_text("\n".join(text) + "\n")
    print(json.dumps({"delivery": str(DELIVERY), "rigid_10k": rigid, "elastic_10k": elastic}, ensure_ascii=False))


if __name__ == "__main__":
    main()
