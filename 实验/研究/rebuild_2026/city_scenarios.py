"""从市区 29 座 110 kV 样本站提取同步场景，并核对共同缺测的影响。"""

import argparse
import csv
from datetime import datetime, timedelta
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from .city_mapping_audit import write_csv
from .hourly_source_profile import CITY_SOURCE, EXPECTED_2025, OUTPUT_DIR


MAPPING_AUDIT = OUTPUT_DIR / "city_2025_mapping_audit.csv"
HOURS = sorted(EXPECTED_2025)
WEEK = timedelta(days=7)


def read_city_series(city_source: Path, mapping_source: Path) -> tuple[dict[str, dict[datetime, float]], list[dict]]:
    with Path(mapping_source).open(encoding="utf-8-sig", newline="") as handle:
        selected = [row for row in csv.DictReader(handle) if row["in_study_cohort"] == "True"]
    series = {}
    with ZipFile(city_source) as archive:
        for row in selected:
            workbook = load_workbook(BytesIO(archive.read(row["archive_member"])), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                sheet.reset_dimensions()  # 原文件错误声明为 A1:A1
                values = {}
                for cells in sheet.iter_rows(values_only=True):
                    raw_time = str(cells[0] or "")
                    if raw_time.startswith("2025-"):
                        values[datetime.fromisoformat(raw_time)] = float(cells[1])
            finally:
                workbook.close()
            series[row["model_station_id"]] = values
    return series, selected


def weekly_analogs(hour: datetime, observed: dict[datetime, float]) -> list[float]:
    values = []
    for direction in (-1, 1):
        for weeks in range(1, 7):
            neighbor = hour + direction * weeks * WEEK
            if neighbor in observed:
                values.append(observed[neighbor])
                break
    if not values:
        raise ValueError(f"缺测小时前后六周均没有同一时刻参照：{hour}")
    return values


def fill_weekly(observed: dict[datetime, float], choice: str) -> dict[datetime, float]:
    filled = dict(observed)
    for hour in HOURS:
        if hour in observed:
            continue
        analogs = weekly_analogs(hour, observed)
        filled[hour] = {
            "weekly_low": min(analogs),
            "weekly_mean": sum(analogs) / len(analogs),
            "weekly_high": max(analogs),
        }[choice]
    return filled


def longest_run(hours: list[datetime]) -> int:
    longest = current = 0
    previous = None
    for hour in hours:
        current = current + 1 if previous is not None and hour - previous == timedelta(hours=1) else 1
        longest = max(longest, current)
        previous = hour
    return longest


def scenario_metrics(values: dict[datetime, float]) -> dict:
    ordered = sorted(values.items())
    forward_time, forward_value = max(ordered, key=lambda item: item[1])
    minimum_time, minimum_value = min(ordered, key=lambda item: item[1])
    forward_peak = max(forward_value, 0.0)
    reverse_peak = max(-minimum_value, 0.0)
    forward_hours = [hour for hour, value in ordered if forward_peak > 0 and value >= 0.95 * forward_peak]
    reverse_hours = [hour for hour, value in ordered if reverse_peak > 0 and -value >= 0.95 * reverse_peak]
    return {
        "forward_peak_time": forward_time.isoformat(sep=" ") if forward_peak else "",
        "forward_peak_mw": round(forward_peak, 6),
        "reverse_peak_time": minimum_time.isoformat(sep=" ") if reverse_peak else "",
        "reverse_peak_mw": round(reverse_peak, 6),
        "minimum_net_load_time": minimum_time.isoformat(sep=" "),
        "minimum_net_load_mw": round(minimum_value, 6),
        "forward_h95_hours": len(forward_hours),
        "reverse_h95_hours": len(reverse_hours),
        "forward_d95_max_run_hours": longest_run(forward_hours),
        "reverse_d95_max_run_hours": longest_run(reverse_hours),
    }


def build_city_scenarios(
    city_source: Path = CITY_SOURCE,
    mapping_source: Path = MAPPING_AUDIT,
) -> tuple[list[dict], list[dict], list[dict]]:
    station_series, stations = read_city_series(city_source, mapping_source)
    model_ids = [row["model_station_id"] for row in stations]
    if len(model_ids) != len(set(model_ids)):
        raise ValueError("市区研究样本有重复的模型站标识")
    common_hours = set.intersection(*(set(values) for values in station_series.values()))
    source_hash = sha256(Path(city_source).read_bytes()).hexdigest()
    mapping_hash = sha256(Path(mapping_source).read_bytes()).hexdigest()
    summary = []
    aggregate_variants = {}
    for variant in ("observed", "weekly_low", "weekly_mean", "weekly_high"):
        chosen = {
            station_id: {hour: values[hour] for hour in common_hours} if variant == "observed" else fill_weekly(values, variant)
            for station_id, values in station_series.items()
        }
        hours = sorted(common_hours) if variant == "observed" else HOURS
        aggregate = {hour: round(sum(chosen[station_id][hour] for station_id in model_ids), 6) for hour in hours}
        aggregate_variants[variant] = aggregate
        for station in stations:
            station_id = station["model_station_id"]
            summary.append(
                {
                    "region_id": "QX-00007",
                    "voltage_kv": 110,
                    "year": 2025,
                    "model_station_id": station_id,
                    "station_identity_basis": station["station_identity_basis"],
                    "variant": variant,
                    "scenario_status": "observed_incomplete" if variant == "observed" else "planning_estimate" if variant == "weekly_mean" else "sensitivity_only",
                    "observed_hours": len(common_hours),
                    "imputed_hours": 0 if variant == "observed" else len(HOURS) - len(common_hours),
                    **scenario_metrics(chosen[station_id]),
                    "city_source_sha256": source_hash,
                    "mapping_audit_sha256": mapping_hash,
                }
            )
        summary.append(
            {
                "region_id": "QX-00007",
                "voltage_kv": 110,
                "year": 2025,
                "model_station_id": "__DISTRICT__",
                "station_identity_basis": "synchronous_sum_of_29_stations",
                "variant": variant,
                "scenario_status": "observed_incomplete" if variant == "observed" else "planning_estimate" if variant == "weekly_mean" else "sensitivity_only",
                "observed_hours": len(common_hours),
                "imputed_hours": 0 if variant == "observed" else len(HOURS) - len(common_hours),
                **scenario_metrics(aggregate),
                "city_source_sha256": source_hash,
                "mapping_audit_sha256": mapping_hash,
            }
        )
    hourly = [
        {
            "hour": hour.isoformat(sep=" "),
            "data_status": "observed" if hour in common_hours else "weekly_analog_estimate",
            "observed_mw": aggregate_variants["observed"].get(hour, ""),
            "weekly_low_mw": aggregate_variants["weekly_low"][hour],
            "weekly_mean_mw": aggregate_variants["weekly_mean"][hour],
            "weekly_high_mw": aggregate_variants["weekly_high"][hour],
        }
        for hour in HOURS
    ]
    district_observed = next(row for row in summary if row["model_station_id"] == "__DISTRICT__" and row["variant"] == "observed")
    scene_times = {
        "district_forward_peak": datetime.fromisoformat(district_observed["forward_peak_time"]),
        "district_minimum_net_load": datetime.fromisoformat(district_observed["minimum_net_load_time"]),
    }
    station_scenes = [
        {
            "region_id": "QX-00007",
            "voltage_kv": 110,
            "year": 2025,
            "scenario_kind": kind,
            "scenario_time": time.isoformat(sep=" "),
            "model_station_id": station["model_station_id"],
            "station_identity_basis": station["station_identity_basis"],
            "station_net_load_mw": round(station_series[station["model_station_id"]][time], 6),
            "district_net_load_mw": aggregate_variants["observed"][time],
            "data_status": "observed",
        }
        for kind, time in scene_times.items()
        for station in stations
    ]
    return summary, hourly, station_scenes


def main() -> None:
    parser = argparse.ArgumentParser(description="计算市区 110 kV 同步静态场景与缺测敏感性")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    summary, hourly, station_scenes = build_city_scenarios()
    write_csv(summary, args.output_dir / "city_2025_scenario_sensitivity.csv")
    write_csv(hourly, args.output_dir / "city_2025_aggregate_hourly.csv")
    write_csv(station_scenes, args.output_dir / "city_2025_station_values_at_district_scenes.csv")
    observed = next(row for row in summary if row["model_station_id"] == "__DISTRICT__" and row["variant"] == "observed")
    print(f"市区 110 kV 同步有效小时 {observed['observed_hours']}；已写入 {len(summary)} 条场景敏感性记录")


if __name__ == "__main__":
    main()
