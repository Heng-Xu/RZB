"""回读实际Excel主表生成预览，并记录截图来源。"""
import argparse
from html import escape

from openpyxl import load_workbook

from . import prepare as p
from .export import FILENAME


def run(collect_only=False):
    book = load_workbook(p.PUBLIC/FILENAME,data_only=True)
    images = []
    for name in p.REGIONS.values():
        ws = book[f'{name}逐年对照']
        skip, spans = set(), {}
        for group in ws.merged_cells.ranges:
            spans[group.min_row,group.min_col] = (group.max_row-group.min_row+1,group.max_col-group.min_col+1)
            skip.update((r,c) for r in range(group.min_row,group.max_row+1)
                        for c in range(group.min_col,group.max_col+1) if (r,c)!=(group.min_row,group.min_col))
        rows = []
        for row in ws:
            cells = []
            for cell in row:
                if (cell.row,cell.column) in skip:
                    continue
                rs,cs = spans.get((cell.row,cell.column),(1,1))
                value = cell.value
                if isinstance(value,(int,float)) and '0.00' in cell.number_format:
                    value = f'{value:.2f}'
                value = '' if value is None else str(value)
                tag = 'th' if cell.row<=2 else 'td'
                cells.append(f'<{tag} rowspan="{rs}" colspan="{cs}">{escape(value)}</{tag}>')
            rows.append('<tr>'+''.join(cells)+'</tr>')
        html = '''<!doctype html><meta charset="utf-8"><style>
        *{box-sizing:border-box}body{margin:35px;font-family:"方正仿宋_GBK","Noto Sans CJK SC",serif;color:#172d41}
        h1{font-size:28px;margin:0 0 14px}p{font-size:19px;line-height:1.5;margin:8px 0 18px}
        table{border-collapse:collapse;width:100%;table-layout:fixed}th,td{text-align:center;border:1px solid #ced6df;padding:12px 6px;font-size:19px;height:44px;line-height:1.35}
        th{background:#244b70;color:white}tr:nth-child(odd) td{background:#f0f4f8}td:first-child,th:first-child{width:6%}
        </style>'''
        html += f'<h1>{name}2021—2035逐年容量及措施</h1><p>四年更换间隔，按当年新增工程总费用比较。刚性在左，弹性在右；储能为年末额定功率，新线为2022年起累计新增联络。</p>'
        html += '<table>'+''.join(rows)+'</table><p>数据为模型回算与规划结果。新增量另表列示；既有转供按收资MW封顶，通道及受端余量共同约束。</p>'
        page = p.REVIEW/f'{name}主表预览.html'
        if not collect_only:
            page.write_text(html)
        else:
            assert page.read_text() == html
            output = p.REVIEW/f'{name}主表预览.png'
            assert output.exists()
            images.append(dict(sheet=ws.title,rows=ws.max_row,columns=ws.max_column,
                               image=str(output.relative_to(p.ROOT)),sha256=p.sha(output)))
    book.close()
    if collect_only:
        p.dump(p.REVIEW/'主表渲染.json',dict(workbook_sha256=p.sha(p.PUBLIC/FILENAME),rendered=images))


if __name__=='__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--collect-only',action='store_true')
    run(parser.parse_args().collect_only)
