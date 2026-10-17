"""逐年左右对照、前后变化、费用分项及后期扩容能力检验。"""
import csv
import json

from openpyxl import Workbook, load_workbook

from ..annual_bounded_transfer_2035_20261017.export import clean, decorate, flat
from ..annual_transfer_2035_20261017.export import FIELDS, fixed, markdown
from . import model as m, prepare as p

LABEL = {'rigid':'刚性','elastic':'弹性'}
FILENAME = '两地区逐年容量及措施_四年间隔总费用左右对照.xlsx'


def csvrows(path):
    with path.open(encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))


def main():
    proof = json.loads((p.REVIEW/'逐年方案独立核验.json').read_text())
    probes = json.loads((p.REVIEW/'后期能力检验.json').read_text())
    assert proof['verified'] and probes['verified']
    book = Workbook()
    book.remove(book.active)
    plans, annual, stations, projects, switches, limits, checks, baselines = {}, [], [], [], [], [], [], []
    solvers, comparisons, changes, late, costs, replacements = [], [], [], [], [], []
    previous_refs = []
    text = ['# 两地区2021—2035逐年容量及措施','',
        '2021容量固定为市区3113.50 MVA、邳州2027.00 MVA，2022—2035逐年以上年设备和当年负荷、光伏求解。'
        '同一原主变位置两次更换年份之差至少为四年，按当年新增主变、储能和联络的费用现值直接比较。表中为模型回算和规划结果。','']
    for region, name in p.REGIONS.items():
        inp, powers = p.load(region)
        ids, aa = list(inp['stations']), m.arcs(inp)
        names = {r['station']:r['name'] or r['station'] for r in inp['station_transfer_evidence']}
        for case in p.CASES:
            root = p.RAW/region/case
            summary = json.loads((root/'summary.json').read_text())
            plans[region,case] = summary
            annual.extend(dict(地区=name,方案=LABEL[case],**{FIELDS[k]:v for k,v in a.items()}) for a in summary['annual'])
            before = {sid:inp['stations'][sid]['initial_units']+[0.] for sid in ids}
            before_n = {sid:0 for sid in ids}
            for r in csvrows(root/'stations.csv'):
                sid, n = r['station'], int(r['storage_modules'])
                u = [float(r[f'unit_{k}_mva']) for k in (1,2,3)]
                stations.append(dict(地区=name,方案=LABEL[case],年份=int(r['year']),变电站=sid,站名=names[sid],
                    主变1_MVA=u[0],主变2_MVA=u[1],主变3_MVA=u[2],年末容量_MVA=sum(u),
                    当年净增容量_MVA=sum(u)-sum(before[sid]),年末储能柜数=n,
                    当年新增储能_MW=(n-before_n[sid])*.1,年末储能额定功率_MW=n*.1,年末储能电量_MWh=n*.215))
                before[sid], before_n[sid] = u, n
            edges = {e['id']:e for e in inp['network']['edges']}
            for r in csvrows(root/'projects.csv'):
                e = edges[r['id']]
                projects.append(dict(地区=name,方案=LABEL[case],投运年=int(r['year']),联络编号=r['id'],
                    站1=e['station_a'],站2=e['station_b'],容量_MW=e['capacity_mw'],依据=e['basis']))
            for r in csvrows(root/'replacements.csv'):
                replacements.append(dict(地区=name,方案=LABEL[case],站名=names[r['station']],变电站=r['station'],
                    主变位置=int(r['slot']),更换年份=int(r['year']),上次更换年份=int(r['previous_year']),
                    原容量_MVA=float(r['old_mva']),新购容量_MVA=float(r['new_mva'])))
            if case=='rigid':
                base = json.loads((root/'baseline_check.json').read_text())
                baselines.extend(dict(地区=name,年份=2021,故障站=r['failed'],核验小时数=r['checked_hours'],
                    约束不满足小时数=r['violating_hours'],最大缺口_MW=r['max_deficit_mw']) for r in base['records'])
            for y in inp['years']:
                r = json.loads((root/str(y)/'adopted.json').read_text())
                s = r['solver']
                solvers.append(dict(地区=name,方案=LABEL[case],年份=y,新增工程费用现值_万元=s['objective'],
                    年度条件下界_万元=s['dual_bound'],距条件下界比例=s['mip_gap'],求解状态=s['status']))
                checks.extend(dict(地区=name,方案=LABEL[case],年份=y,故障站=v['failed'],校核小时数=v['checked_hours'],
                    最大缺口_MW=v['max_deficit_mw'],不满足小时数=v['violating_hours'],最小余量_MW=v['minimum_margin_mw']) for v in r['check']['records'])
                for field, mode in (('normal','正常固定'),('extra','故障站追加')):
                    for key, count in sorted(r[field].items(),key=lambda kv:int(kv[0])):
                        a = aa[int(key)]
                        switches.append(dict(地区=name,方案=LABEL[case],年份=y,方式=mode,原属站=a['source'],接收站=a['target'],
                            路径=a['id'],源馈线组=a['a'],受端馈线=a['b'],区段数=count,
                            峰值预约上界_MW=count*float(abs(powers[y][:,a['sa']]).max())/inp['block_count'][a['source']],
                            路径性质='既有能力代理' if a['existing'] else '新增付费候选'))
                for j, sid in enumerate(ids):
                    existing = sum(c for field in ('normal','extra') for key,c in r[field].items()
                                   if aa[int(key)]['source']==sid and aa[int(key)]['existing'])
                    paid = sum(c for field in ('normal','extra') for key,c in r[field].items()
                               if aa[int(key)]['source']==sid and not aa[int(key)]['existing'])
                    limits.append(dict(地区=name,方案=LABEL[case],年份=y,变电站=sid,站名=names[sid],
                        既有转供预约上界_MW=existing*float(abs(powers[y][:,j]).max())/inp['block_count'][sid],
                        原表既有能力_MW=inp['existing_transfer_budget_by_year_mw'][str(y)][sid],
                        新增转出区段数=paid,新增转出区段上限=int(inp['block_count'][sid]*inp['new_transfer_fraction']+1e-9)))
            oldcase = 'rigid' if case=='rigid' else 'elastic_capacity'
            oldpath = p.parent.RAW/region/oldcase/'summary.json'
            prior = json.loads(oldpath.read_text())
            previous_refs.append(dict(path=str(oldpath.relative_to(p.ROOT)),sha256=p.sha(oldpath)))
            for index, (old, new) in enumerate(zip(prior['annual'],summary['annual'])):
                olddelta = 0. if index==0 else old['storage_nameplate_mw']-prior['annual'][index-1]['storage_nameplate_mw']
                newdelta = 0. if index==0 else new['storage_nameplate_mw']-summary['annual'][index-1]['storage_nameplate_mw']
                changes.append(dict(地区=name,方案=LABEL[case],年份=new['year'],
                    前版容量_MVA=old['capacity_mva'],本轮容量_MVA=new['capacity_mva'],
                    前版当年新增储能_MW=olddelta,本轮当年新增储能_MW=newdelta,
                    前版年末储能_MW=old['storage_nameplate_mw'],本轮年末储能_MW=new['storage_nameplate_mw'],
                    前版累计新线_条=old['cumulative_lines'],本轮累计新线_条=new['cumulative_lines'],
                    前版累计费用现值_万元=old['cumulative_lifecycle_npv_10k'],本轮累计费用现值_万元=new['cumulative_lifecycle_npv_10k']))
            costs.append(dict(地区=name,方案=LABEL[case],主变费用现值_万元=sum(a['transformer_npv_10k'] for a in summary['annual']),
                储能费用现值_万元=sum(a['storage_npv_10k'] for a in summary['annual']),
                联络费用现值_万元=sum(a['line_npv_10k'] for a in summary['annual']),合计费用现值_万元=summary['cost_npv_10k'],
                前版合计费用现值_万元=prior['cost_npv_10k'],费用变化_万元=summary['cost_npv_10k']-prior['cost_npv_10k']))
        rigid, elastic = plans[region,'rigid']['annual'], plans[region,'elastic']['annual']
        ws = book.create_sheet(f'{name}逐年对照',0 if region=='city' else 1)
        ws.append(['年份','规划峰值/MW','刚性方案',None,None,None,None,'弹性方案',None,None,None,None])
        for area in ('A1:A2','B1:B2','C1:G1','H1:L1'):
            ws.merge_cells(area)
        cols = ['年末容量/MVA','当年净增/MVA','容载比','年末储能/MW','累计新线/条']
        ws.append([None,None,*cols,*cols])
        md, compact = [], []
        for i, (a,b) in enumerate(zip(rigid,elastic)):
            assert a['year']==b['year'] and a['peak_mw']==b['peak_mw'] and a['common_floor']==b['common_floor']
            keys = ['capacity_mva','net_capacity_increment_mva','clr','storage_nameplate_mw','cumulative_lines']
            values = [a['year'],a['peak_mw'],*[a[k] for k in keys],*[b[k] for k in keys]]
            ws.append([clean(v) for v in values])
            md.append([a['year'],fixed(a['peak_mw']),*[int(a[k]) if k=='cumulative_lines' else fixed(a[k]) for k in keys],
                       *[int(b[k]) if k=='cumulative_lines' else fixed(b[k]) for k in keys]])
            compact.append(dict(年份=a['year'],规划峰值_MW=a['peak_mw'],
                **{'刚性_'+FIELDS[k]:a[k] for k in keys},**{'弹性_'+FIELDS[k]:b[k] for k in keys}))
            comparisons.append(dict(地区=name,年份=a['year'],容量差_弹减刚_MVA=b['capacity_mva']-a['capacity_mva'],
                储能差_弹减刚_MW=b['storage_nameplate_mw']-a['storage_nameplate_mw'],
                联络差_弹减刚_条=b['cumulative_lines']-a['cumulative_lines'],
                费用差_弹减刚_万元=b['cumulative_lifecycle_npv_10k']-a['cumulative_lifecycle_npv_10k']))
            if a['year']>=2031:
                da = a['storage_nameplate_mw']-rigid[i-1]['storage_nameplate_mw']
                db = b['storage_nameplate_mw']-elastic[i-1]['storage_nameplate_mw']
                late.append(dict(地区=name,年份=a['year'],刚性当年净增容量_MVA=a['net_capacity_increment_mva'],
                    弹性当年净增容量_MVA=b['net_capacity_increment_mva'],刚性当年新增储能_MW=da,
                    弹性当年新增储能_MW=db,刚性当年新增联络_条=a['new_lines'],弹性当年新增联络_条=b['new_lines'],
                    弹性当年费用现值_万元=b['added_lifecycle_npv_10k'],
                    弹性实际只新增主变=bool(b['net_capacity_increment_mva']>1e-7 and db<1e-7 and b['new_lines']==0)))
        decorate(ws,2)
        ws.column_dimensions['A'].width = 9
        ws.print_area = f'A1:L{ws.max_row}'
        for row in ws.iter_rows(min_row=3):
            for c in row:
                c.number_format = '0' if c.column in (1,7,12) else '0.00'
        p.write_csv(p.PUBLIC/f'{name}逐年左右对照.csv',compact)
        text.extend([f'表 {name}逐年容量及措施','',markdown(
            ['年份','峰值/MW','刚性容量/MVA','刚性净增/MVA','刚性容载比','刚性储能/MW','刚性累计新线/条',
             '弹性容量/MVA','弹性净增/MVA','弹性容载比','弹性储能/MW','弹性累计新线/条'],md),''])
    assumptions = [dict(项目=k,说明=v) for k,v in [
        ('容量及新增量','主表储能为年末额定功率，联络为2022起累计新增条数。当年新增措施另表列示。新购主变容量为更换与扩建采购铭牌容量之和，净增容量扣除同期被替换容量。'),
        ('更换间隔','任意连续四个规划年度内同一原主变位置最多更换一次，新一次更换年份与上次之差至少为4；研究工程假设。基期设备既往更换年份缺失，从2022年记录本模型更换事件。'),
        ('年度求解','2022—2035仅依当年源荷和上年资产选工程，已投运主变、储能及新线逐年保留。新购31.5、50、63、80 MVA，允许扩建位置第三台整台50 MVA。'),
        ('第三台资料边界','46座有设备明细对应站点中，9座原表可扩容量超过50 MVA，分为63 MVA五座、80 MVA一座、100 MVA三座；均有1个预留位置。现有锁定扩建造价实例只对应50 MVA，本轮第三台仍限50 MVA。可行性结论仅适用于此候选范围。'),
        ('刚弹差别','共同源荷、容载比下限、安全和工程约束。刚性容量上限为max(上年容量,当年峰值×2.0)，弹性取消该新增控制；两者均按新增工程总费用直接寻优。'),
        ('转供证据','46座站沿用原Excel J列MW固定能力上限，其余3座站按零。既有正常校准区段和事故追加共同占用，正常区段可回切。2025资料用于早期回算，逐线投运及真实拓扑资料有缺口。'),
        ('网络约束','保留路径、共享源馈线组、受端馈线与接收站主变余量。新增付费路径最多转出全站20%源荷区段，向下取整；EQ组和路径为研究代理。'),
        ('储能校核','0.10 MW/0.215 MWh整柜，可用电量系数0.76，有效支撑为额定功率与可用电量/持续时长的较小值。正向市区6小时、邳州2小时，反向沿用逐站时长。静态校核不含连续SOC调度。'),
        ('后期能力检验','2031—2035各年继承本轮上年资产，固定已有储能，主变取当年候选范围内允许最大规格。分别固定已有新线和开放全部候选新线，核对停止新增储能的可行性。为单年反事实，非连续无储能新增链；未检验超过50 MVA的第三台工程。'),
        ('费用及界限','沿用锁定寿命、维护、折现及2041评价期。费用采用具体可行配置，年度条件下界和间隙逐年列示；不将其合计作为全期最优下界。'),
        ('对前版比较','更换限制与选择顺序同时变化，前后差额为两项变更的合并影响。前版原始结果完整保留，未替换正式报告采用数据。'),
        ('基期及源荷','2021固定地区容量，邳州基期事故缺口单列且未加入虚拟设备。历史为源荷模板回算，8760小时含补值，未来沿用原预测；年度容载比分母为原口径规划峰值。'),
    ]]
    data = {'后期当年新增措施':late,'前版与本轮逐年比较':changes,'逐年新增及费用':annual,'费用分项汇总':costs,
        '后期只扩容能力检验':probes['rows'],'逐站逐年配置':stations,'主变更换履历':replacements,
        '联络逐项工程':projects,'固定区段切换预案':switches,'站级转供上限核对':limits,
        '年度费用下界与间隙':solvers,'全年安全核验':checks,'2021固定基线检验':baselines,
        '刚弹完整序列比较':comparisons,'定义与假设':assumptions}
    for name, rows in data.items():
        if not rows:
            continue
        ws = flat(book,name,rows)
        p.write_csv(p.PUBLIC/(name+'.csv'),rows)
        if name=='定义与假设':
            ws.column_dimensions['A'].width = 22
            ws.column_dimensions['B'].width = 118
            for row in range(2,ws.max_row+1):
                ws.row_dimensions[row].height = max(42,20*((len(ws.cell(row,2).value)+49)//50))
        if name=='后期只扩容能力检验':
            ws.column_dimensions['H'].width = 80
    path = p.PUBLIC/FILENAME
    book.save(path)
    rb = load_workbook(path,data_only=True)
    cells = 0
    for ws in book:
        actual = rb[ws.title]
        assert actual.max_row==ws.max_row and actual.max_column==ws.max_column
        for row in ws:
            for c in row:
                value = actual.cell(c.row,c.column).value
                if isinstance(c.value,float):
                    assert abs(value-c.value)<max(1e-8,abs(c.value)*1e-12)
                else:
                    assert (value if value is not None else '') == (c.value if c.value is not None else '')
                cells += 1
    rb.close()
    p.dump(p.REVIEW/'工作簿回读.json',dict(verified=True,cells=cells,path=str(path.relative_to(p.ROOT)),sha256=p.sha(path)))
    text.extend(['储能按年末额定功率列示；当年新增与累计规模分别统计。','',
        '表 后期弹性方案当年新增措施','',markdown(['地区','年份','净增容量/MVA','新增储能/MW','新增联络/条'],[
            [r['地区'],r['年份'],fixed(r['弹性当年净增容量_MVA']),fixed(r['弹性当年新增储能_MW']),r['弹性当年新增联络_条']] for r in late]),'',
        '表 全期新增工程费用现值','',markdown(['地区','方案','主变/万元','储能/万元','联络/万元','合计/万元','较前版变化/万元'],[
            [r['地区'],r['方案'],*[fixed(r[k]) for k in ('主变费用现值_万元','储能费用现值_万元','联络费用现值_万元','合计费用现值_万元','费用变化_万元')]] for r in costs]),'',
        f'四套方案2022—2035合计{proof["full_hour_pairs"]:,}个小时—工况组合通过全年静态安全复核。'
        '邳州2021固定设备的事故缺口另列，不作为2022年在役储能。','',
        '两次更换间隔、费用分项及年度条件下界已逐项核对。2025收资用于早期回算、EQ拓扑和源荷区段为研究代理，储能按静态时长校核；当前费用为逐年条件可行方案费用。'])
    text.extend(['','本轮变化与后期扩容条件',''])
    for region,name in p.REGIONS.items():
        actual = plans[region,'elastic']['annual'][-1]
        previous = json.loads((p.parent.RAW/region/'elastic_capacity/summary.json').read_text())
        old = previous['annual'][-1]
        cost = next(r for r in costs if r['地区']==name and r['方案']=='弹性')
        delta = {kind:sum(a[kind+'_npv_10k'] for a in plans[region,'elastic']['annual'])
                 -sum(a[kind+'_npv_10k'] for a in previous['annual']) for kind in ('transformer','storage','line')}
        text.extend([f'{name}弹性方案2035年净增容量由前版{old["net_capacity_increment_mva"]:.2f} MVA变为{actual["net_capacity_increment_mva"]:.2f} MVA，'
            f'当年新增储能由{old["new_storage_mwh"]/2.15:.2f} MW变为{actual["new_storage_mwh"]/2.15:.2f} MW。'
            f'年末储能由{old["storage_nameplate_mw"]:.2f} MW变为{actual["storage_nameplate_mw"]:.2f} MW，累计新增联络由{old["cumulative_lines"]}条变为{actual["cumulative_lines"]}条。'
            f'全期新增工程费用现值为{cost["合计费用现值_万元"]:.2f}万元，较前版变化{cost["费用变化_万元"]:+.2f}万元。'
            f'差额由主变{delta["transformer"]:+.2f}万元、储能{delta["storage"]:+.2f}万元和联络{delta["line"]:+.2f}万元共同构成。',''])
        repeat = [r for r in replacements if r['地区']==name and r['方案']=='弹性' and r['上次更换年份']>0]
        if repeat:
            events = '、'.join(f'{r["站名"]}{r["更换年份"]}年（上次{r["上次更换年份"]}年）' for r in repeat)
            text.extend([f'{name}发生再次更换的原主变位置共{len(repeat)}个事件：{events}。'
                '再次更换按新购整台铭牌容量计费，原设备退出容量从净增容量中扣除。',''])
        categories = []
        for mode,title in (('只新增主变','只扩容'),('新增主变及联络','扩容与联络共同补强')):
            years = [str(r['年份']) for r in probes['rows'] if r['地区']==name and r['检验']==mode and r['结论']=='可行']
            categories.append(title+'可行年份：'+('、'.join(years) if years else '无'))
        text.extend([name+'2031—2035年停止新增储能的单年能力检验，'+'；'.join(categories)+'。'
            '各年均继承本轮上年资产，第三台仍限50 MVA，可行配置只说明能力，费用需另行比较。',''])
    text.extend(['设备原表另有9座站预留扩建容量超过50 MVA，现有扩建造价实例只对应50 MVA。'
        '超过50 MVA第三台的费用资料尚未闭合，本轮后期不可行性结论限于当前候选范围。',''])
    (p.PUBLIC/'逐年容量及措施_四年间隔总费用左右对照.md').write_text('\n'.join(text).rstrip()+'\n')
    p.dump(p.REVIEW/'结果序列核对.json',dict(verified=True,rows=comparisons,late=late,costs=costs))
    p.dump(p.SOURCE/'展示口径及前版依据.json',dict(workbook=str(path.relative_to(p.ROOT)),previous_results=previous_refs,assumptions=assumptions))
    print('EXPORTED_INTERVAL',cells,flush=True)


if __name__ == '__main__':
    main()
