"""2025 年原始逐时数据的轻量核对；不推断站点或电压映射。"""

import argparse
import csv
from datetime import datetime, timedelta
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


STUDY_DIR = Path(__file__).resolve().parents[1]
SOURCE_DIR = STUDY_DIR / "data/tuomin"
PIZHOU_SOURCE = SOURCE_DIR / "电网建模数据_Agent整合版_V1.2/邳州主变负载率.xlsx"
CITY_SOURCE = SOURCE_DIR / "补充数据/徐州市区脱敏.zip"
OUTPUT_DIR = Path(__file__).resolve().parent / "source_audit"
EXPECTED_2025 = {datetime(2025, 1, 1) + timedelta(hours=i) for i in range(8760)}


def _metrics(times: list[datetime], values: list[float | None]) -> dict:
    unique_times = set(times)
    missing_times = EXPECTED_2025 - unique_times
    longest_missing_run = 0
    current_run = 0
    for hour in sorted(EXPECTED_2025):
        current_run = current_run + 1 if hour in missing_times else 0
        longest_missing_run = max(longest_missing_run, current_run)
    numeric = [value for value in values if value is not None]
    forward_peak = max((value for value in numeric if value > 0), default=0.0)
    reverse_peak = max((-value for value in numeric if value < 0), default=0.0)
    return {
        "source_year": 2025,
        "source_rows": len(times),
        "unique_hours": len(unique_times),
        "missing_hours": len(missing_times),
        "longest_missing_run_hours": longest_missing_run,
        "duplicate_hours": len(times) - len(unique_times),
        "missing_values": len(values) - len(numeric),
        "first_missing_value_time": next(
            (time.isoformat(sep=" ") for time, value in zip(times, values) if value is None),
            "",
        ),
        "positive_hours": sum(value > 0 for value in numeric),
        "negative_hours": sum(value < 0 for value in numeric),
        "zero_hours": sum(value == 0 for value in numeric),
        "forward_peak_mw": round(forward_peak, 9),
        "reverse_peak_mw": round(reverse_peak, 9),
        "forward_h95_observed": sum(
            value is not None and value >= 0.95 * forward_peak
            for value in values
        ) if forward_peak else 0,
        "reverse_h95_observed": sum(
            value is not None and -value >= 0.95 * reverse_peak
            for value in values
        ) if reverse_peak else 0,
        "all_zero": bool(numeric) and all(value == 0 for value in numeric),
    }


def profile_pizhou_columns(source: Path) -> list[dict]:
    """Sheet3 中仅取 2025 年；列号是原始标识，不凭重复表头合并主变。"""
    source = Path(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet3"]
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator)
        next(iterator)  # 空白行
        times: list[datetime] = []
        columns: list[list[float | None]] = [[] for _ in headers[1:]]
        for cells in iterator:
            raw_time = str(cells[0] or "")
            if not raw_time.startswith("2025-"):
                continue
            times.append(datetime.fromisoformat(raw_time))
            for index, values in enumerate(columns, start=1):
                raw = cells[index]
                values.append(float(raw) if raw not in (None, "") else None)
        rows = []
        for index, values in enumerate(columns, start=1):
            rows.append(
                {
                    "source_file": str(source.resolve().relative_to(STUDY_DIR)),
                    "source_sheet": sheet.title,
                    "source_column": index + 1,
                    "header_station_id": headers[index],
                    "voltage_assignment": "unverified",
                    **_metrics(times, values),
                    "source_sha256": source_hash,
                }
            )
        return rows
    finally:
        workbook.close()


def profile_city_archive(source: Path) -> list[dict]:
    """逐个读取市区压缩包；修复工作簿错误的 A1:A1 维度元数据。"""
    source = Path(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    rows = []
    with ZipFile(source) as archive:
        for member in sorted(name for name in archive.namelist() if name.endswith(".xlsx")):
            workbook = load_workbook(BytesIO(archive.read(member)), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                sheet.reset_dimensions()
                times: list[datetime] = []
                values: list[float | None] = []
                series_hash = sha256()
                for cells in sheet.iter_rows(values_only=True):
                    raw_time = str(cells[0] or "")
                    if not raw_time.startswith("2025-"):
                        continue
                    time = datetime.fromisoformat(raw_time)
                    raw_value = cells[1] if len(cells) > 1 else None
                    value = float(raw_value) if raw_value not in (None, "") else None
                    times.append(time)
                    values.append(value)
                    series_hash.update(f"{time.isoformat()}|{value}\n".encode())
                rows.append(
                    {
                        "source_archive": str(source.resolve().relative_to(STUDY_DIR)),
                        "archive_member": member,
                        "source_sheet_index": 1,
                        "voltage_assignment": "unverified",
                        **_metrics(times, values),
                        "series_sha256": series_hash.hexdigest(),
                        "archive_sha256": source_hash,
                    }
                )
            finally:
                workbook.close()
    first_member_by_hash = {}
    for row in rows:
        digest = row["series_sha256"]
        row["duplicate_of"] = first_member_by_hash.get(digest, "")
        first_member_by_hash.setdefault(digest, row["archive_member"])
    return rows


def _write_csv(rows: list[dict], target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="核对 2025 年邳州与市区逐时源数据")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    pizhou = profile_pizhou_columns(PIZHOU_SOURCE)
    city = profile_city_archive(CITY_SOURCE)
    _write_csv(pizhou, args.output_dir / "pizhou_2025_column_profile.csv")
    _write_csv(city, args.output_dir / "city_2025_series_profile.csv")
    print(f"邳州 {len(pizhou)} 列；市区 {len(city)} 个文件，{sum(bool(r['duplicate_of']) for r in city)} 个重复序列")


if __name__ == "__main__":
    main()
