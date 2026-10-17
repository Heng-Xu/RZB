"""独立复算年度资产与费用，输出可供人工逐站核对的工作簿。"""
import json
import math
from decimal import Decimal, ROUND_HALF_UP

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from scipy.sparse import coo_matrix

from . import run as m
from ..case_parameters_20261012.report import add_sheet, markdown


def export():
    inp=json.loads((m.RAW/'inputs.json').read_text());ids=list(inp['stations'])
    archive=np.load(m.RAW/'scenarios.npz');scenes={y:archive[str(y)] for y in inp['years']}
    annual=[];stations=[];projects=[];costs=[];checks=[];summaries={};states={};support=[];folders={}
    for case,label in m.LABEL.items():
        options=[p for p in m.RAW.glob(case+'_*/summary.json') if (p.parent/'independent_lp.json').exists()]
        folder=min(options,key=lambda p:json.loads(p.read_text())['cost_npv_10k']).parent
        folders[case]=folder
        s=json.loads((folder/'summary.json').read_text());summaries[case]=s
        assert s['input_sha256']==m.sha(m.RAW/'inputs.json') and s['forecast_sha256']==m.sha(m.FORECAST)
        proof=json.loads((folder/'independent_lp.json').read_text());assert proof['verified']
        assert len(proof['records'])==sum(len(scenes[y])*30 for y in inp['years'])
        assert all(r['verified'] for r in proof['records'])
        selected={tuple(r) for r in json.loads((folder/'selected_pairs.json').read_text())}
        native,meta=m.base.build(inp,scenes,case,selected);z=np.load(folder/'chosen_solution.npz')['solution']
        rr,cc,vv=zip(*native.entries)
        lhs=coo_matrix((vv,(rr,cc)),shape=(len(native.constraint_lower),len(z))).tocsr()@z
        integers=np.array(native.integrality,dtype=bool)
        violation=max(0.,float(np.max(np.array(native.constraint_lower)-lhs)),float(np.max(lhs-np.array(native.constraint_upper))),
            float(np.max(np.array(native.lower_bounds)-z)),float(np.max(z-np.array(native.upper_bounds))),float(np.max(abs(z[integers]-np.rint(z[integers])))))
        assert violation<1e-5
        old=np.array([r['initial_units']+[0.] for r in inp['stations'].values()]);oldn=np.zeros(len(ids));oldbuilt=set()
        repeat=np.zeros((len(ids),2),int);total=0.;states[case]={}
        for a in s['annual']:
            y=a['year'];u,n,built=m.base.old.old.assets_from_csv(inp,folder,y)
            zu,zn,_,_=m.base.old.old.base.assets.assets(inp,z,meta,y)
            assert np.array_equal(u,zu) and np.array_equal(n,zn)
            assert np.all(u>=old) and np.all(n>=oldn) and built>=oldbuilt
            assert {eid for (eid,yy),v in meta['edgevars'].items() if yy==y and z[v]>.5}==built
            if case=='rigid':assert u.sum()<=max(old.sum(),2*a['peak_mw'])+1e-5
            bought=np.where(u[:,:2]>old[:,:2]+1e-8,u[:,:2],0.).sum()
            removed=np.where(u[:,:2]>old[:,:2]+1e-8,old[:,:2],0.).sum()
            thirds=(u[:,2]-old[:,2]).sum()/50.
            capex=dict(transformer=bought*inp['transformer_price_10k_per_purchased_mva']+thirds*inp['third_50mva_capex_10k'],
                storage=(n-oldn).sum()*inp['storage_module_capex_10k'],line=len(built-oldbuilt)*inp['line_capex_10k'])
            npvs={k:v*inp['cost_factors'][k][str(y)] for k,v in capex.items()};cost=sum(npvs.values());total+=cost
            assert abs(cost-a['added_lifecycle_npv_10k'])<1e-5
            for k,v in npvs.items():assert abs(v-a[k+'_npv_10k'])<1e-5
            assert abs(u.sum()-a['capacity_mva'])<1e-8 and abs(n.sum()*.215-a['storage_mwh'])<1e-8
            assert abs(u.sum()/a['peak_mw']-a['clr'])<1e-9
            repeat+=(u[:,:2]>old[:,:2]+1e-8)
            assert abs(bought+50*thirds-removed-(u.sum()-old.sum()))<1e-8
            annual.append(dict(方案=label,年份=y,规划负荷_MW=a['peak_mw'],年末主变_MVA=a['capacity_mva'],容载比=a['clr'],
                当年净增主变_MVA=u.sum()-old.sum(),当年新购主变_MVA=bought+50*thirds,
                当年新增第三台_台=int(thirds),年末储能_MW=a['storage_nameplate_mw'],年末储能_MWh=a['storage_mwh'],
                当年新增储能_MWh=a['new_storage_mwh'],当年新增联络_条=a['new_lines'],累计新增联络_条=a['cumulative_lines'],
                当年建设投资_万元=a['construction_capex_10k'],当年费用现值_万元=cost,累计费用现值_万元=total))
            for j,sid in enumerate(ids):
                stations.append(dict(方案=label,年份=y,站号=sid,站名=next((e['name'] for e in inp['station_transfer_evidence'] if e['station']==sid),sid) or sid,
                    年初配置_MVA='+'.join(f'{v:g}' for v in old[j] if v),年末配置_MVA='+'.join(f'{v:g}' for v in u[j] if v),
                    年末容量_MVA=float(u[j].sum()),当年净增容量_MVA=float((u[j]-old[j]).sum()),
                    当年新购容量_MVA=float(np.where(u[j,:2]>old[j,:2]+1e-8,u[j,:2],0.).sum()+u[j,2]-old[j,2]),
                    年末储能_MW=float(.1*n[j]),年末储能_MWh=float(.215*n[j]),当年新增储能_MWh=float(.215*(n[j]-oldn[j])),储能柜数=int(n[j])))
            states[case][y]=(u,n,built);old,oldn,oldbuilt=u,n,built
        assert repeat.max()<=1 and abs(total-s['cost_npv_10k'])<1e-4 and abs(float(np.dot(native.costs,z))-total)<1e-4
        for e in m.base.old.old.base.previous.rows(folder/'projects.csv'):
            projects.append(dict(方案=label,投运年份=int(e['year']),项目=e['id'],站A=e['station_a'],站B=e['station_b'],馈线A=e['a'],馈线B=e['b'],
                能力_MW=float(e['capacity_mw']),建设投资_万元=inp['line_capex_10k']))
        for kind,cn in [('transformer','主变'),('storage','储能'),('line','联络')]:
            costs.append(dict(方案=label,费用项=cn,费用现值_万元=sum(a[kind+'_npv_10k'] for a in s['annual'])))
        checks.append(dict(方案=label,费用现值_万元=total,费用下界_万元=s['solver']['dual_bound'],相对间隙=s['solver']['mip_gap'],
            独立静态校核工况数=len(proof['records']),最大原模型违反量=violation,重复更换次数最大值=int(repeat.max()),输入SHA256=s['input_sha256']))
    comparisons=[]
    for y in inp['years']:
        a,b=[next(r for r in annual if r['方案']==m.LABEL[c] and r['年份']==y) for c in m.LABEL]
        comparisons.append(dict(年份=y,规划负荷_MW=a['规划负荷_MW'],刚性主变_MVA=a['年末主变_MVA'],弹性主变_MVA=b['年末主变_MVA'],
            刚性容载比=a['容载比'],弹性容载比=b['容载比'],刚性储能_MWh=a['年末储能_MWh'],弹性储能_MWh=b['年末储能_MWh'],
            刚性累计联络_条=a['累计新增联络_条'],弹性累计联络_条=b['累计新增联络_条'],
            刚性累计费用_万元=a['累计费用现值_万元'],弹性累计费用_万元=b['累计费用现值_万元']))
    # 固定2035资产，对容量差异最大的站，独立求最少转供并手算剩余主变与储能支撑。
    delta=states['elastic_free'][2035][0].sum(axis=1)-states['rigid'][2035][0].sum(axis=1)
    targets=sorted([j for j,sid in enumerate(ids) if sid.startswith('BDZ-')],key=lambda j:abs(delta[j]),reverse=True)[:5]
    for case,label in m.LABEL.items():
        lp=m.base.OperatingLP(inp);u,n,built=states[case][2035]
        for j in targets:
            for h in sorted({int(np.argmax(scenes[2035][:,j])),int(np.argmin(scenes[2035][:,j]))}):
                p=scenes[2035][h];ready=lp.prepared(p,2035,u,n,built);dispatch=lp.dispatch(p,ready,j)
                outgoing=np.zeros(len(ids));incoming=np.zeros(len(ids))
                for k,value in dispatch['flow']:
                    arc=lp.arcs[k];outgoing[arc[5]]+=value;incoming[arc[6]]+=value
                served=p-outgoing+incoming
                common,bounds,full,largest,bp,bn=lp.bounds_for(ready,j)
                positive=full.copy();positive[j]-=largest[j]
                reverse=positive+bn;positive=positive+bp
                assert np.max(served-positive)<1e-5 and np.max(-served-reverse)<1e-5
                support.append(dict(方案=label,年份=2035,站号=ids[j],站名=next(e['name'] for e in inp['station_transfer_evidence'] if e['station']==ids[j]),
                    截面=h,基准小时=inp['forecast_snapshot_roles']['2035'][h]['source_hour'],
                    原净负荷_MW=float(p[j]),主变组合_MVA='+'.join(f'{v:g}' for v in u[j] if v),
                    N1剩余主变_MW=float(full[j]-largest[j]),转出_MW=float(outgoing[j]),转入_MW=float(incoming[j]),
                    转供后净负荷_MW=float(served[j]),储能正向有效功率_MW=float(bp[j]),储能反向有效功率_MW=float(bn[j]),
                    所需放电_MW=float(max(0.,served[j]-(full[j]-largest[j]))),
                    所需吸收_MW=float(max(0.,-served[j]-(full[j]-largest[j])))))
    rates=[]
    for r in json.loads((m.RAW/'forecast_annual.json').read_text()):
        rates.append(dict(年份=r['year'],市区总负荷_MW=r['city_load_mw'],徐州市光伏装机指数=r['citywide_pv_index'],
            负荷相对2025倍率=r['load_factor'],光伏相对2025倍率=r['pv_factor'],
            负荷同比增长率=r['load_yoy'],光伏同比增长率=r['pv_yoy'],负荷是否插值=r['load_interpolated'],光伏是否插值=r['pv_interpolated']))
    notes=[dict(事项=k,说明=v) for k,v in inp['forecast_basis'].items()]
    notes.extend([dict(事项='逐年决策',说明='采用逐年滚动方案，上年已投运主变、储能和联络固定，求当年措施。完成后以完整十年联合模型固定设备复核，不将2035配置反填各年。'),
        dict(事项='方案差别',说明='两方案共用源荷、设备与网络。刚性保留C_y≤max(C_(y-1),2P_y)及1.60分档下限，弹性取消容载比上下限，以相同费用目标寻优。本次为弹性经济参考，不进行5%以内简化联络的第二轮工程优选。'),
        dict(事项='差异解释',说明='2032年1.60×年度负荷首次超过3113.50 MVA，刚性下限成为地区增容的直接要求。2031年两组总容量相同，且刚性容载比上下限均未触及，该年储能与联络的小幅布局差别包含限时寻优容差影响。'),
        dict(事项='容量含义',说明='2021地区实际容量3113.50 MVA保持，逐站规格继承冻结标定；2025年末两方案仍为3113.50 MVA，2026接续。并非重新拟合2025实际逐台设备。'),
        dict(事项='人工核对范围',说明='已逐年复算采购与净增容量、储能柜数、联络投运和费用；关键工况核对2035年容量差异最大的5站正向及反向截面。运行校核覆盖所列有限静态截面及全部单站最大主变N-1。')])
    frames={'逐年对比':comparisons,'逐年措施':annual,'逐站设备':stations,'新增联络':projects,'费用分项':costs,
            '增长率':rates,'关键工况':support,'求解与核验':checks,'计算说明':notes}
    book=Workbook();book.remove(book.active)
    for name,rows in frames.items():
        m.write_csv(m.PUBLIC/(name+'.csv'),rows);add_sheet(book,name,rows)
        if not rows:continue
        ws=book[name]
        for row in ws:
            for cell in row:cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        for j,key in enumerate(rows[0],1):
            if any(k in key for k in ['增长率','相对间隙']):
                for row in ws.iter_rows(min_row=2,min_col=j,max_col=j):row[0].number_format='0.00%'
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
    workbook=m.PUBLIC/'市区2026—2035刚弹对比.xlsx';book.save(workbook)
    reopened=load_workbook(workbook,data_only=True);cell_count=0
    for name,rows in frames.items():
        if not rows:assert reopened[name].max_row==1;continue
        actual=list(reopened[name].values);expected=[list(rows[0])]+[list(r.values()) for r in rows]
        assert len(actual)==len(expected)
        for aa,bb in zip(actual,expected):
            for a,b in zip(aa,bb):
                if isinstance(b,float):assert math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-8)
                else:assert (a if a is not None else '')==(b if b is not None else '')
                cell_count+=1
    reopened.close()
    a,b=[summaries[c] for c in m.LABEL]
    def presentation(row):
        return {k:float(Decimal(str(round(float(v),12))).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
                if isinstance(v,(float,np.floating)) else v for k,v in row.items()}
    end_a,end_b=[presentation(s['annual'][-1]) for s in [a,b]]
    saving=a['cost_npv_10k']-b['cost_npv_10k'];first=next((r['年份'] for r in comparisons if any(abs(r[x]-r[y])>1e-6 for x,y in [
        ('刚性主变_MVA','弹性主变_MVA'),('刚性储能_MWh','弹性储能_MWh'),('刚性累计联络_条','弹性累计联络_条')])),None)
    differences={kind:sum(r[kind+'_npv_10k'] for r in a['annual'])-sum(r[kind+'_npv_10k'] for r in b['annual']) for kind in ['transformer','storage','line']}
    comparisons=[presentation(r) for r in comparisons]
    text=['# 市区不压缩起点容量：2026—2035年前推试算','',
        f'市区2025年末3113.50 MVA直接接续。2035年弹性经济参考相对刚性方案，新增措施费用现值减少{saving:.2f}万元（{saving/a["cost_npv_10k"]:.2%}）。两方案最早在{first}年出现配置差别。','',
        '模型按年滚动：固定上年已投运设备，决定当年新增措施，并对当年全部静态截面校核。完成后，再用原完整十年模型固定各年设备逐项复核。','',
        markdown(['2035年末及十年费用','刚性','弹性经济参考'],[
            ['主变/MVA',f'{end_a["capacity_mva"]:.2f}',f'{end_b["capacity_mva"]:.2f}'],
            ['容载比',f'{end_a["clr"]:.2f}',f'{end_b["clr"]:.2f}'],
            ['储能/MW',f'{end_a["storage_nameplate_mw"]:.2f}',f'{end_b["storage_nameplate_mw"]:.2f}'],
            ['储能/MWh',f'{end_a["storage_mwh"]:.2f}',f'{end_b["storage_mwh"]:.2f}'],
            ['累计新增10kV联络/条',end_a['cumulative_lines'],end_b['cumulative_lines']],
            ['2026—2035措施费用现值/万元',f'{a["cost_npv_10k"]:.2f}',f'{b["cost_npv_10k"]:.2f}']]),'',
        '费用按原2021现值基年折算，包含本轮新增措施投运后的2041年以前运维。既有设备不重复计投资。',
        f'本次为逐年滚动可行方案，距十年联合模型费用下界的相对间隙为刚性{a["solver"]["mip_gap"]:.2%}、弹性{b["solver"]["mip_gap"]:.2%}，费用差额用于比较本次具体方案，尚未证明十年全局最优。','',
        '## 增长率与输入','',
        '最新市区总负荷由2025年的2850.00 MW增至2030年的3440.00 MW、2035年的4150.00 MW，两个五年区间的复合增长率分别为'
        f'{(3440/2850)**.2-1:.2%}、{(4150/3440)**.2-1:.2%}。实际逐年输入保留预测表已有节点，2026—2028及2034按相邻节点等增量插值。',
        '徐州市光伏装机指数由2025年709.00增至2030年1107.00、2035年1840.00；相应五年复合增长率为'
        f'{(1107/709)**.2-1:.2%}、{(1840/1107)**.2-1:.2%}。2026—2030保留给定值，2031—2034每年增加146.60。原截图光伏单位未标明，本次只取相对2025年的倍率。','',
        '110kV原模型2025年统计负荷1528.82 MW乘市区负荷倍率，2035年为2226.18 MW。市区光伏采用徐州市相同增长比例，空间分配沿用2025模板。净负荷按“负荷代理×负荷倍率−光伏代理×光伏倍率”分别计算。',
        '源荷分解采用原归档night_central日间负荷代理：08:00—18:00将站级负荷代理与原同期净负荷中较大者作为负荷分量，差额作为光伏分量；夜间原净功率保留在非光伏分量。该分解逐点还原2025基准，是快速外推假设。','',
        '## 逐年对比','',
        markdown(['年份','规划负荷/MW','刚/弹主变/MVA','刚/弹容载比','刚/弹储能/MWh','刚/弹累计新联络/条'],[
            [r['年份'],f'{r["规划负荷_MW"]:.2f}',f'{r["刚性主变_MVA"]:.2f} / {r["弹性主变_MVA"]:.2f}',
             f'{r["刚性容载比"]:.2f} / {r["弹性容载比"]:.2f}',f'{r["刚性储能_MWh"]:.2f} / {r["弹性储能_MWh"]:.2f}',
             f'{r["刚性累计联络_条"]} / {r["弹性累计联络_条"]}'] for r in comparisons]),'',
        '2032年，1.60×1998.19 MW＝3197.11 MVA，首次超过起点3113.50 MVA，刚性方案因此需要提高地区总容量。2035年的下限为3561.88 MVA，刚性选配3562.50 MVA，弹性保留3143.50 MVA，并用更多联络满足具体站点的正常与N-1要求。两方案2035年容量均低于2.0上限，本次容量差异主要来自1.60下限。',
        '2031年两方案地区总容量相同，刚性容载比上下限均未触及，该年的储能与联络布局差别包含限时寻优容差影响。',
        f'费用分项核对：弹性主变费用减少{differences["transformer"]:.2f}万元、储能费用减少{differences["storage"]:.2f}万元，联络费用增加{-differences["line"]:.2f}万元，合计减少{saving:.2f}万元。','',
        '## 核对与范围','',
        '正常运行和N-1共同允许储能及联络替代，逐站既有转供预算、共享馈线余量、接收侧能力、正常反向80%与最大一台主变退出规则保持。2026—2035每个主变位置最多更换一次，两方案共同执行。刚性保留2.0新增容量控制及原分档下限，弹性取消容载比上下限。本次交付为费用目标下的弹性经济参考。',
        f'两方案共核对{sum(r["独立静态校核工况数"] for r in checks)}个有限静态工况，全部通过。逐年设备和费用另行复算，工作簿{cell_count}个单元格回读一致；关键工况给出5座差异站的主变、储能及实际所需转供量。',
        '本次为增长假设下的前推方案，沿用既有候选网络与补齐数据。储能仍按原静态功率和持续时间约束计算；研究报告及已确认成果保持原文件。','']
    memo=m.PUBLIC/'方案结果.md';memo.write_text('\n'.join(text))
    m.dump(m.REVIEW/'核验记录.json',dict(verified=True,schemes=checks,workbook_cells=cell_count,
        support_rows=len(support),manual_review='逐年检查采购容量与净增容量关系、容量/负荷比值、储能整数柜数及累计联络；2035关键5站逐项复算N-1剩余主变、正反向支撑及转供后净功率。',
        artifacts=[dict(path=str(p.relative_to(m.ROOT)),sha256=m.sha(p)) for p in [workbook,memo,m.RAW/'inputs.json']],
        solution_sources=[dict(case=c,folder=str(folders[c].relative_to(m.ROOT)),
            files={name:m.sha(folders[c]/name) for name in ['summary.json','stations.csv','projects.csv','selected_pairs.json','chosen_solution.npz','independent_lp.json','rolling_history.json'] if (folders[c]/name).exists()}) for c in m.LABEL],
        scope='本次独立试算仅交付Markdown、XLSX及CSV，不修改正式研究报告，无Word/PDF交付'))
    print('EXPORTED',workbook,'cells',cell_count,'checks',checks,flush=True)


if __name__=='__main__':export()
