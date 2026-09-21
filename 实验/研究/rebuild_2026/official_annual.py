"""从原始年度统计表读取本轮正式建模层的容量和降压负荷。"""

import argparse
import csv
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook


LAYERS = {("QX-00005", 110), ("QX-00005", 35), ("QX-00007", 110)}
STUDY_DIR = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/近5年容载比.xlsx"
DEFAULT_TARGET = Path(__file__).resolve().parent / "source_audit/official_annual.csv"


def read_official_annual(source: Path) -> list[dict]:
    """读取 2021—2025 年原表值；万千伏安/万千瓦乘 10 转为 MVA/MW。"""
    source = Path(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    try:
        source_label = str(source.resolve().relative_to(STUDY_DIR))
    except ValueError:
        source_label = str(source)
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows = []
    region_id = None
    try:
        sheet = workbook["Sheet1"]
        for source_row, cells in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
            if cells[0] is not None:
                value = str(cells[0]).strip()
                region_id = value if value.startswith("QX-") else None
            voltage_text = str(cells[1] or "").strip()
            if voltage_text.startswith("110"):
                voltage_kv = 110
            elif voltage_text.startswith("35"):
                voltage_kv = 35
            else:
                continue
            if (region_id, voltage_kv) not in LAYERS:
                continue
            for year in range(2021, 2026):
                offset = 2 + 3 * (year - 2021)
                rows.append(
                    {
                        "region_id": region_id,
                        "voltage_kv": voltage_kv,
                        "year": year,
                        "capacity_mva": round(float(cells[offset]) * 10, 9),
                        "reported_downward_load_mw": round(float(cells[offset + 1]) * 10, 9),
                        "reported_clr": float(cells[offset + 2]),
                        "source_file": source_label,
                        "source_sheet": sheet.title,
                        "source_row": source_row,
                        "source_sha256": source_hash,
                    }
                )
    finally:
        workbook.close()
    return rows


def write_official_annual(source: Path, target: Path) -> int:
    """将原表正式建模层写入独立 CSV，返回记录数。"""
    rows = read_official_annual(source)
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="读取 2021—2025 年原始年度容量和降压负荷")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_TARGET)
    args = parser.parse_args()
    count = write_official_annual(args.source, args.output)
    print(f"已写入 {count} 条：{args.output}")


if __name__ == "__main__":
    main()
