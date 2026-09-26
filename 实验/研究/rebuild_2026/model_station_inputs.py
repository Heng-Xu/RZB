"""合成两片区 2025 年站级场景与设备容量来源，不代替 2021 年仿真起点。"""

import argparse
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .asset_2025 import SOURCE as ASSET_SOURCE
from .asset_2025 import read_assets
from .city_mapping_audit import write_csv
from .hourly_source_profile import OUTPUT_DIR


CITY_SCENARIOS = OUTPUT_DIR / "city_2025_scenario_sensitivity.csv"
CITY_MAPPING = OUTPUT_DIR / "city_2025_mapping_audit.csv"
PIZHOU_SCENARIOS = OUTPUT_DIR / "pizhou_2025_station_scenarios_evidence.csv"
EXTERNAL_CITY_IDS = {"BDZ-00055", "BDZ-00183"}


def read_external_city_assets(source: Path = ASSET_SOURCE) -> dict[str, tuple[str, float]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    capacity = defaultdict(float)
    try:
        for cells in workbook["110千伏变电站1"].iter_rows(min_row=4, values_only=True):
            if cells[3] in EXTERNAL_CITY_IDS:
                capacity[(cells[3], cells[1])] += float(cells[15] or 0)
    finally:
        workbook.close()
    return {station_id: (region_id, value) for (station_id, region_id), value in capacity.items()}


def read_csv(source: Path) -> list[dict]:
    with Path(source).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def build_model_station_inputs(
    city_scenarios: Path = CITY_SCENARIOS,
    city_mapping: Path = CITY_MAPPING,
    pizhou_scenarios: Path = PIZHOU_SCENARIOS,
) -> list[dict]:
    capacity = defaultdict(float)
    for asset in read_assets():
        capacity[(asset["region_id"], asset["voltage_kv"], asset["station_id"])] += asset["capacity_mva"]
    external = read_external_city_assets()
    rows = []
    for scene in read_csv(pizhou_scenarios):
        voltage = int(scene["voltage_kv"])
        station_id = scene["station_id"]
        rows.append(
            {
                "study_region_id": "QX-00005",
                "voltage_kv": voltage,
                "year": 2025,
                "model_station_id": station_id,
                "source_station_id": station_id,
                "source_asset_region_id": "QX-00005",
                "source_capacity_mva_2025": capacity[("QX-00005", voltage, station_id)],
                "capacity_basis": "source_equipment_snapshot",
                "scenario_basis": "observed_mapping_supported_one_hour_gap",
                "forward_peak_mw": scene["forward_peak_mw"],
                "reverse_peak_mw": scene["reverse_peak_mw"],
                "forward_h95_hours": scene["forward_h95_hours"],
                "reverse_h95_hours": scene["reverse_h95_hours"],
            }
        )
    city_by_id = {
        row["model_station_id"]: row
        for row in read_csv(city_mapping) if row["in_study_cohort"] == "True"
    }
    for scene in read_csv(city_scenarios):
        station_id = scene["model_station_id"]
        if scene["variant"] != "weekly_mean" or station_id == "__DISTRICT__":
            continue
        mapping = city_by_id[station_id]
        source_station_id = mapping["station_id"]
        if source_station_id in external:
            asset_region, asset_capacity = external[source_station_id]
            basis = "source_equipment_other_region_user_confirmed_city"
        elif source_station_id:
            asset_region = "QX-00007"
            asset_capacity = capacity[("QX-00007", 110, source_station_id)]
            basis = "source_equipment_snapshot"
        else:
            asset_region = ""
            asset_capacity = ""
            basis = "simulation_capacity_to_select"
        rows.append(
            {
                "study_region_id": "QX-00007",
                "voltage_kv": 110,
                "year": 2025,
                "model_station_id": station_id,
                "source_station_id": source_station_id,
                "source_asset_region_id": asset_region,
                "source_capacity_mva_2025": asset_capacity,
                "capacity_basis": basis,
                "scenario_basis": "weekly_mean_estimate_421_common_missing_hours",
                "forward_peak_mw": scene["forward_peak_mw"],
                "reverse_peak_mw": scene["reverse_peak_mw"],
                "forward_h95_hours": scene["forward_h95_hours"],
                "reverse_h95_hours": scene["reverse_h95_hours"],
            }
        )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="合成两片区站级场景与设备容量来源")
    parser.add_argument("--output", type=Path, default=OUTPUT_DIR / "model_station_inputs_2025.csv")
    args = parser.parse_args()
    rows = build_model_station_inputs()
    write_csv(rows, args.output)
    print(f"已输出 {len(rows)} 个站—电压模型输入；其中 {sum(row['capacity_basis'] == 'simulation_capacity_to_select' for row in rows)} 座仿真站容量待选")


if __name__ == "__main__":
    main()
