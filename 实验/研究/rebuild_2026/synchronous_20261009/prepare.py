import copy
import json
import subprocess
from datetime import datetime, timedelta
from io import BytesIO
from zipfile import ZipFile

import numpy as np
from openpyxl import load_workbook

from ..pilot_20261009.run import ROOT, OUT as PREVIOUS, dump, sha, write_csv, peak_events, reconstruct_baseline
from ..hourly_source_profile import CITY_SOURCE, PIZHOU_SOURCE, OUTPUT_DIR
from ..baseline_2021 import read_csv
from ..joint_lifecycle_optimizer import load_inputs
from ..station_grid_feasibility import station_metadata
from ..pizhou_transfer_fraction import pizhou_sample_fraction
from ..current_economic_assumptions import line_capex

OUT=ROOT/'研究报告/试算/20261009_两地同期重建'
REGIONS={'city':('市区','QX-00007',CITY_SOURCE),'pizhou':('邳州','QX-00005',PIZHOU_SOURCE)}


def prepare(region):
    name,rid,source=REGIONS[region]
    out=OUT/region;out.mkdir(parents=True,exist_ok=True)
    b,sc,d,p,c,prices=load_inputs()
    keys=sorted(k for k in b if k[:2]==(rid,110));ids=[k[2] for k in keys]
    hours=[datetime(2025,1,1)+timedelta(hours=i) for i in range(8760)]
    position={h:i for i,h in enumerate(hours)}
    raw=np.full((8760,len(ids)),np.nan)
    if region=='city':
        mapping_path=PREVIOUS/'city_2025_mapping_audit.csv'
        mapping=[r for r in read_csv(mapping_path) if r['in_study_cohort']=='True']
        assert {r['model_station_id'] for r in mapping}==set(ids)
        assert all(r['city_source_sha256']==sha(source) for r in mapping)
        with ZipFile(source) as z:
            for r in mapping:
                w=load_workbook(BytesIO(z.read(r['archive_member'])),read_only=True,data_only=True)
                s=w.active;s.reset_dimensions();col=ids.index(r['model_station_id'])
                for cells in s.iter_rows(values_only=True):
                    h=str(cells[0] or '')
                    if h.startswith('2025-') and cells[1] is not None:
                        row=position[datetime.fromisoformat(h)]
                        assert np.isnan(raw[row,col]);raw[row,col]=float(cells[1])
                w.close()
    else:
        mapping_path=out/'pizhou_2025_mapping_evidence.csv'
        mapping_path.write_bytes(subprocess.check_output(['git','show','d7a3c38^:实验/研究/rebuild_2026/source_audit/pizhou_2025_mapping_evidence.csv']))
        mapping=[r for r in read_csv(mapping_path) if r['resolved_voltage_kv']=='110' and r['station_status']=='supported_for_station_aggregation']
        assert {r['resolved_station_id'] for r in mapping}==set(ids)
        assert all(r['hourly_source_sha256']==sha(source) for r in mapping)
        columns={sid:[int(r['source_column'])-1 for r in mapping if r['resolved_station_id']==sid] for sid in ids}
        w=load_workbook(source,read_only=True,data_only=True)
        for cells in w['Sheet3'].iter_rows(min_row=3,values_only=True):
            h=str(cells[0] or '')
            if not h.startswith('2025-'):continue
            row=position[datetime.fromisoformat(h)]
            for j,sid in enumerate(ids):
                vals=[cells[c] for c in columns[sid]]
                if all(v not in (None,'') for v in vals):
                    assert np.isnan(raw[row,j]);raw[row,j]=sum(float(v) for v in vals)
        w.close()
    observed=np.isfinite(raw).all(axis=1)
    filled=raw.copy()
    for row,col in zip(*np.where(np.isnan(raw))):
        candidates=[]
        if region=='pizhou':
            assert 0<row<8759 and np.isfinite(raw[[row-1,row+1],col]).all()
            candidates=raw[[row-1,row+1],col]
        else:
            for direction in (-1,1):
                for weeks in range(1,7):
                    near=row+direction*168*weeks
                    if 0<=near<8760 and np.isfinite(raw[near,col]):
                        candidates.append(raw[near,col]);break
        assert len(candidates)>0;filled[row,col]=np.mean(candidates)
    annual={int(r['year']):r for r in read_csv(OUTPUT_DIR/'official_annual.csv') if r['region_id']==rid and r['voltage_kv']=='110'}
    observed_peaks={y:float(r['reported_downward_load_mw']) for y,r in annual.items()}
    total=filled.sum(axis=1);factor=observed_peaks[2025]/total.max()
    matrix=filled*factor*observed_peaks[2021]/observed_peaks[2025]
    events=peak_events(hours,total);event=max(events,key=lambda e:e['hours'])
    original=peak_events([h for h,v in zip(hours,observed) if v],total[observed])
    duration=event['hours']
    inputs=copy.deepcopy(json.loads((PREVIOUS/'pilot_inputs.json').read_text()))
    inputs.update(region=region,region_name=name,region_id=rid,initial_peak_mw=observed_peaks[2021],
                  annual_observed_peak_mw=observed_peaks,growth_rate=(observed_peaks[2025]/observed_peaks[2021])**.25-1,
                  forward_duration_h=duration,years=[2022,2023,2024,2025],
                  scope='同期站级功率矢量；每个场景仅一台主变N-1，其余站正常；完整8760小时静态分离复核')
    if region=='city':
        stations=inputs['stations']
        assert sorted(stations)==ids
    else:
        physical=station_metadata(set(ids));stations={}
        for k in keys:
            sid=k[2];s=b[k];md=physical[sid]
            stations[sid]={'old_units':[float(s[f'simulation_unit_{i}_mva']) for i in (1,2)],
                           'asset_capacity_mva':float(s['source_capacity_mva_2025'] or 0),
                           'allow_third':md['available_third_slots']>=1 and md['available_third_mva']>=50,
                           'metadata':md}
        reconstruct_baseline(stations,2*observed_peaks[2021],out/'baseline_solver')
        sample=pizhou_sample_fraction();dump(out/'初始转供率样本依据.json',sample)
        inputs['initial_transfer_fraction']=sample['initial_fraction']
        inputs['line_cable_fraction']=.1
        inputs['line_capex_10k']=line_capex(.1)
        inputs['line_cost_coefficient']=1+.1*(180/24-1)
    for j,sid in enumerate(ids):
        stations[sid].update(base_forward_mw=float(max(0,matrix[:,j].max())),
                             base_reverse_mw=float(max(0,-matrix[:,j].min())),forward_duration_h=duration,
                             reverse_duration_h=int(d[(rid,110,sid)]['reverse_d95_max_run_hours']))
    inputs['stations']=stations
    inputs['initial_capacity_mva']=sum(sum(s['initial_units']) for s in stations.values())
    inputs['initial_clr']=inputs['initial_capacity_mva']/inputs['initial_peak_mw']
    inputs['physical_capacity_upper_mva']=sum(sum(max(v,80) for v in s['initial_units'])+50*int(s['allow_third']) for s in stations.values())
    inputs['n1_basis']='每个故障场景仅一个站的最大一台主变N-1，其他站按正常容量；转接可随时刻与故障变化，统一使用固定既有与已建新增能力。正反向均校核。'
    inputs['line_basis']='双端数量及源端/总能力预算的完整站群等效模型；无实线路路径、电压、潮流及接收馈线热限验证。'
    inputs['baseline_treatment']='四组共用原重构资产，不虚构免费基期补强；必要补强在首个规划年计费。另算零增长成本基准，不能当作工程逐项成本拆分。'
    inputs['source_hashes']={str(x.relative_to(ROOT)):sha(x) for x in (source,mapping_path,OUTPUT_DIR/'official_annual.csv',PREVIOUS/'pilot_inputs.json')}
    stats={'station_count':len(ids),'observed_hours':int(observed.sum()),'filled_hours':int((~observed).sum()),
           'sample_peak_mw':float(total.max()),'official_peak_mw':observed_peaks[2025],
           'calibration_factor':float(factor),'longest_event':event,
           'observed_longest_hours':max(e['hours'] for e in original),
           'longest_imputed_hours':int(sum(not observed[position[datetime.fromisoformat(event['start'])]+i] for i in range(duration))),
           'source_scope':'原始样本按区县官方年度峰值统一比例标定，非全域逐时实测',
           'imputation':'市区邻周同刻平均；邳州单小时缺口相邻时刻线性补值'}
    np.savez_compressed(out/'hourly.npz',power_base_mw=matrix,station_ids=np.array(ids),observed=observed,
                        hours=np.array([h.isoformat(sep=' ') for h in hours]))
    dump(out/'inputs.json',inputs);dump(out/'峰段与数据核查.json',stats)
    write_csv(out/'区县同期负载率.csv',[{'hour':h.isoformat(sep=' '),'observed':int(observed[i]),
                   'estimated_mw':float(total[i]*factor),'loading_fraction':float(total[i]*factor/(.95*float(annual[2025]['capacity_mva'])))} for i,h in enumerate(hours)])
    print(name,'PREPARED',len(ids),'stations','duration',duration,'capacity',inputs['initial_capacity_mva'],'growth',inputs['growth_rate'],flush=True)


if __name__=='__main__':
    raise SystemExit('该历史独立入口已归档；请使用 实验.研究.rebuild_2026.static_actual_20261010。公共函数仍保留供现行模型导入。')
    import argparse
    parser=argparse.ArgumentParser();parser.add_argument('region',choices=REGIONS)
    prepare(parser.parse_args().region)
