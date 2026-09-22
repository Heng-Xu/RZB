"""邳州 2025 年逐时列到站级资产的候选映射；不作正式映射审批。"""

import argparse
import csv
from collections import Counter
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook


STUDY_DIR = Path(__file__).resolve().parents[1]
ASSET_SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/110（35）kv设备明细.xlsx"
PROFILE_SOURCE = Path(__file__).resolve().parent / "source_audit/pizhou_2025_column_profile.csv"
OUTPUT = Path(__file__).resolve().parent / "source_audit/pizhou_2025_mapping_candidates.csv"

# 仅由原表极值与 2025 年设备清单推断；详见 2025-hourly-findings.md。
INFERRED_STATIONS = {12: "BDZ-00290", 13: "BDZ-00290", 28: "BDZ-00247", 29: "BDZ-00247"}


def build_mapping_candidates(asset_source: Path, profile_source: Path) -> list[dict]:
    asset_source = Path(asset_source)
    asset_hash = sha256(asset_source.read_bytes()).hexdigest()
    workbook = load_workbook(asset_source, read_only=True, data_only=True)
    try:
        asset_voltage = {}
        for voltage, sheet_name in (
            (110, "110千伏变电站1"),
            (35, "35千伏变电站1"),
        ):
            for cells in workbook[sheet_name].iter_rows(min_row=4, values_only=True):
                if cells[1] == "QX-00005" and cells[3]:
                    asset_voltage[cells[3]] = voltage
    finally:
        workbook.close()

    with Path(profile_source).open(encoding="utf-8-sig", newline="") as handle:
        profiles = list(csv.DictReader(handle))
    station_counts = Counter(
        INFERRED_STATIONS.get(int(row["source_column"]), row["header_station_id"])
        for row in profiles
    )
    rows = []
    for profile in profiles:
        column = int(profile["source_column"])
        station = INFERRED_STATIONS.get(column, profile["header_station_id"])
        rows.append(
            {
                "source_column": column,
                "raw_header_station_id": profile["header_station_id"],
                "candidate_station_id": station,
                "candidate_voltage_kv": asset_voltage[station],
                "candidate_station_columns": station_counts[station],
                "mapping_basis": "source_extrema_inference" if column in INFERRED_STATIONS else "header_and_asset",
                "mapping_status": "provisional",
                "data_status": "sparse_all_zero" if profile["all_zero"] == "True" else "one_missing_value",
                "hourly_source_sha256": profile["source_sha256"],
                "asset_source_sha256": asset_hash,
            }
        )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="生成邳州主变时序列的候选站级映射")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    rows = build_mapping_candidates(ASSET_SOURCE, PROFILE_SOURCE)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    print(f"已写入 {len(rows)} 条候选映射：{args.output}")


if __name__ == "__main__":
    main()
