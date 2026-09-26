"""摘录分县逐月光伏序列的覆盖与同月比例；原表未注明数值单位。"""

import argparse
import re
from collections import defaultdict
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from .city_mapping_audit import write_csv
from .hourly_source_profile import OUTPUT_DIR
from .official_annual import STUDY_DIR


SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/逐月分县分布式光伏.xlsx"
ASSET_SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/光伏装机.xlsx"
HOURLY_SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/光伏8760小时数据.xlsx"
REGIONS = ("QX-00005", "QX-00007")


def read_monthly_pv(source: Path = SOURCE) -> tuple[list[dict], list[dict]]:
    source = Path(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    raw = []
    try:
        sheet = workbook["Sheet1"]
        year = None
        month_columns = {}
        for source_row, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
            marker = cells[1] if len(cells) > 1 else None
            if isinstance(marker, str) and re.fullmatch(r"20\d{2}年", marker):
                year = int(marker[:4])
                month_columns = {}
                continue
            if cells[0] == "区县代号":
                month_columns = {
                    col: int(label[:-1])
                    for col, label in enumerate(cells)
                    if isinstance(label, str) and re.fullmatch(r"(?:[1-9]|1[0-2])月", label)
                }
                continue
            if cells[0] not in REGIONS or year is None:
                continue
            for col, month in month_columns.items():
                value = cells[col]
                if not isinstance(value, (int, float)):
                    continue
                raw.append(
                    {
                        "region_id": cells[0],
                        "year": year,
                        "month": month,
                        "source_value": float(value),
                        "source_value_unit": "not_stated_in_workbook",
                        "source_file": str(source.resolve().relative_to(STUDY_DIR)),
                        "source_sheet": sheet.title,
                        "source_row": source_row,
                        "source_sha256": source_hash,
                    }
                )
    finally:
        workbook.close()
    reference = {(row["region_id"], row["month"]): row["source_value"] for row in raw if row["year"] == 2025}
    for row in raw:
        denominator = reference.get((row["region_id"], row["month"]))
        row["relative_to_2025_same_month"] = round(row["source_value"] / denominator, 9) if denominator else ""
    counts = defaultdict(set)
    for row in raw:
        counts[(row["region_id"], row["year"])].add(row["month"])
    coverage = []
    for region in REGIONS:
        for calendar_year in range(2021, 2027):
            months = sorted(counts[(region, calendar_year)])
            coverage.append(
                {
                    "region_id": region,
                    "year": calendar_year,
                    "month_count": len(months),
                    "months_present": ",".join(map(str, months)),
                    "coverage_status": "full_12_months" if len(months) == 12 else ("missing" if not months else "partial"),
                    "station_allocation_available": "no",
                    "physical_unit_stated": "no",
                }
            )
    return raw, coverage


def read_asset_pv_coverage(source: Path = ASSET_SOURCE) -> list[dict]:
    source = Path(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    grouped = defaultdict(list)
    try:
        sheet = workbook["主变（跨区合并）1"]
        for cells in sheet.iter_rows(min_row=2, values_only=True):
            grouped[(str(cells[1]), str(cells[3]))].append(cells)
    finally:
        workbook.close()
    result = []
    for region in REGIONS:
        for voltage in (35, 110):
            rows = grouped[(region, f"{voltage}kV")]
            years = sorted({str(row[4])[:4] for row in rows if row[4]})
            result.append(
                {
                    "region_id": region,
                    "voltage_kv": voltage,
                    "row_count": len(rows),
                    "station_count": len({str(row[2]) for row in rows}),
                    "typical_time_years": ",".join(years),
                    "role_for_2021_2025_model": "not_same_year_or_no_region_rows",
                    "source_file": str(source.resolve().relative_to(STUDY_DIR)),
                    "source_sheet": "主变（跨区合并）1",
                    "source_sha256": source_hash,
                }
            )
    return result


def read_hourly_pv_metadata(source: Path = HOURLY_SOURCE) -> list[dict]:
    source = Path(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        notes = dict(workbook["说明"].values)
        data_rows = workbook["光伏8760"].max_row - 1
    finally:
        workbook.close()
    return [
        {
            "source_file": str(source.resolve().relative_to(STUDY_DIR)),
            "source_sha256": source_hash,
            "generator_name": notes["光伏机组名称"],
            "generator_capacity_mw": notes["装机容量（MW）"],
            "original_year_scenario": notes["源年份/场景"],
            "hour_count": data_rows,
            "calendar_basis": notes["时间说明"],
            "role_for_study": "system_generator_profile_not_observed_2025_station_distributed_pv",
        }
    ]


def main() -> None:
    parser = argparse.ArgumentParser(description="核对分县逐月光伏的年份覆盖和同比例")
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--asset-source", type=Path, default=ASSET_SOURCE)
    parser.add_argument("--hourly-source", type=Path, default=HOURLY_SOURCE)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    monthly, coverage = read_monthly_pv(args.source)
    asset_coverage = read_asset_pv_coverage(args.asset_source)
    hourly_metadata = read_hourly_pv_metadata(args.hourly_source)
    write_csv(monthly, args.output_dir / "pv_monthly_region_audit.csv")
    write_csv(coverage, args.output_dir / "pv_monthly_source_coverage.csv")
    write_csv(asset_coverage, args.output_dir / "pv_asset_source_coverage.csv")
    write_csv(hourly_metadata, args.output_dir / "pv_hourly_proxy_metadata.csv")
    print(f"已核对 {len(monthly)} 个分县月份值、{len(coverage)} 个片区年度覆盖项、{len(asset_coverage)} 个站级光伏源覆盖项和逐时代理元数据；逐月表单位未注明。")


if __name__ == "__main__":
    main()
