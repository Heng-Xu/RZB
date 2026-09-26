"""试算邳州 2025 年同步场景；支持候选映射和已复核逐列映射。"""

import argparse
import csv
from datetime import datetime, timedelta
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook


STUDY_DIR = Path(__file__).resolve().parents[1]
SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/邳州主变负载率.xlsx"
AUDIT_DIR = Path(__file__).resolve().parent / "source_audit"


def _longest_run(times: list[datetime]) -> int:
    longest = current = 0
    previous = None
    for time in times:
        current = current + 1 if previous is not None and time - previous == timedelta(hours=1) else 1
        longest = max(longest, current)
        previous = time
    return longest


def build_provisional_scenarios(
    hourly_source: Path,
    mapping_source: Path,
    annual_source: Path,
) -> list[dict]:
    with Path(mapping_source).open(encoding="utf-8-sig", newline="") as handle:
        mapping = list(csv.DictReader(handle))
    verified_mapping = "resolved_voltage_kv" in mapping[0]
    columns = {
        voltage: [
            int(row["source_column"]) - 1
            for row in mapping
            if (int(row["resolved_voltage_kv"]) if verified_mapping else int(row["candidate_voltage_kv"])) == voltage
            and (row["station_status"] == "supported_for_station_aggregation" if verified_mapping else row["candidate_station_id"] != "BDZ-00056")
        ]
        for voltage in (110, 35)
    }
    with Path(annual_source).open(encoding="utf-8-sig", newline="") as handle:
        official = {
            int(row["voltage_kv"]): float(row["reported_downward_load_mw"])
            for row in csv.DictReader(handle)
            if row["region_id"] == "QX-00005" and int(row["year"]) == 2025
        }
    workbook = load_workbook(hourly_source, read_only=True, data_only=True)
    series: dict[int, list[tuple[datetime, float]]] = {110: [], 35: []}
    try:
        for cells in workbook["Sheet3"].iter_rows(min_row=3, values_only=True):
            raw_time = str(cells[0] or "")
            if not raw_time.startswith("2025-"):
                continue
            time = datetime.fromisoformat(raw_time)
            for voltage in (110, 35):
                values = [cells[index] for index in columns[voltage]]
                if all(value not in (None, "") for value in values):
                    series[voltage].append(
                        (time, round(sum(float(value) for value in values), 6))
                    )
    finally:
        workbook.close()

    source_hash = sha256(Path(hourly_source).read_bytes()).hexdigest()
    mapping_hash = sha256(Path(mapping_source).read_bytes()).hexdigest()
    rows = []
    for voltage, observed in series.items():
        by_time = dict(observed)
        missing = sorted(
            datetime(2025, 1, 1) + timedelta(hours=hour)
            for hour in range(8760)
            if datetime(2025, 1, 1) + timedelta(hours=hour) not in by_time
        )
        gap_time = missing[0]
        gap_previous = by_time[gap_time - timedelta(hours=1)]
        gap_next = by_time[gap_time + timedelta(hours=1)]
        gap_estimate = round((gap_previous + gap_next) / 2, 6)
        filled = observed + [(gap_time, gap_estimate)]
        forward = max(observed, key=lambda item: item[1])
        reverse = min(observed, key=lambda item: item[1])
        filled_forward = max(filled, key=lambda item: item[1])
        filled_reverse = min(filled, key=lambda item: item[1])
        forward_h95_times = [
            time for time, power in observed if power >= 0.95 * forward[1]
        ]
        reverse_h95_times = [
            time for time, power in observed if power <= 0.95 * reverse[1]
        ]
        filled_forward_h95 = sum(power >= 0.95 * filled_forward[1] for _, power in filled)
        filled_reverse_h95 = sum(power <= 0.95 * filled_reverse[1] for _, power in filled)
        rows.append(
            {
                "region_id": "QX-00005",
                "voltage_kv": voltage,
                "year": 2025,
                "scenario_status": "mapping_supported_gap_sensitivity" if verified_mapping else "provisional",
                "valid_hours": len(observed),
                "missing_aggregate_hours": 8760 - len(observed),
                "missing_hour": gap_time.isoformat(sep=" "),
                "missing_hour_previous_mw": gap_previous,
                "missing_hour_next_mw": gap_next,
                "missing_hour_linear_estimate_mw": gap_estimate,
                "forward_h95_linear_fill": filled_forward_h95,
                "reverse_h95_linear_fill": filled_reverse_h95,
                "linear_fill_changes_peak": filled_forward[1] != forward[1] or filled_reverse[1] != reverse[1],
                "linear_fill_changes_h95": filled_forward_h95 != len(forward_h95_times) or filled_reverse_h95 != len(reverse_h95_times),
                "excluded_sparse_station": "BDZ-00056" if voltage == 110 else "",
                "forward_peak_time": forward[0].isoformat(sep=" "),
                "forward_peak_mw": forward[1],
                "reverse_peak_time": reverse[0].isoformat(sep=" "),
                "reverse_peak_mw": -reverse[1],
                "forward_h95_observed": len(forward_h95_times),
                "reverse_h95_observed": len(reverse_h95_times),
                "forward_d95_max_run_hours": _longest_run(forward_h95_times),
                "reverse_d95_max_run_hours": _longest_run(reverse_h95_times),
                "official_2025_downward_peak_mw": official[voltage],
                "hourly_minus_official_peak_mw": round(forward[1] - official[voltage], 6),
                "hourly_source_sha256": source_hash,
                "mapping_source_sha256": mapping_hash,
            }
        )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="试算邳州 2025 年同步静态场景")
    parser.add_argument(
        "--output",
        type=Path,
        default=AUDIT_DIR / "pizhou_2025_scenarios_evidence.csv",
    )
    args = parser.parse_args()
    rows = build_provisional_scenarios(
        SOURCE,
        AUDIT_DIR / "pizhou_2025_mapping_evidence.csv",
        AUDIT_DIR / "official_annual.csv",
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    print(f"已写入 {len(rows)} 条已复核映射场景试算：{args.output}")


if __name__ == "__main__":
    main()
