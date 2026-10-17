"""交付两地区逐年左右对照、年度新增与费用、安全及来源核对。"""
import csv
import json

from openpyxl import Workbook, load_workbook

from . import model as m, prepare as p
from .verify import CASES
from ..annual_bounded_transfer_2035_20261017.export import decorate,flat,clean
from ..annual_transfer_2035_20261017.export import FIELDS,fixed,markdown

LABEL={'rigid':'刚性','elastic_capacity':'弹性','elastic_engineering':'弹性经济对照'}


def csvrows(path):
    with path.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))


def main():
    proof=json.loads((p.REVIEW/'逐年方案独立核验.json').read_text());assert proof['verified']
    p.PUBLIC.mkdir(exist_ok=True)
    book=Workbook();book.remove(book.active)
    plans={};annual=[];stations=[];projects=[];switches=[];limits=[];checks=[];baselines=[];solver=[];comparison=[];exceptions=[];phases=[];growth=[]
    text=['# 两地区2021—2035逐年需求驱动方案','',
        '2021保留市区3113.50MVA、邳州2027.00MVA。2022—2035分别以上年在役设备及当年负荷、光伏求解，主变、储能及新增联络投运后逐年承接。左为刚性，右为弹性。表中为模型回算及规划结果。','']
    for region,name in p.REGIONS.items():
        inp,powers=p.load(region);aa=m.arcs(inp);ids=list(inp['stations'])
        for case in CASES:
            root=p.RAW/region/case;s=json.loads((root/'summary.json').read_text());plans[region,case]=s
            annual.extend(dict(地区=name,方案=LABEL[case],**{FIELDS[k]:v for k,v in row.items()}) for row in s['annual'])
            if case=='rigid':
                baseline=json.loads((root/'baseline_check.json').read_text())
                baselines.extend(dict(地区=name,年份=2021,方式='基期固定设备及校准供区',故障站=v['failed'],
                    校核小时数=v['checked_hours'],约束不满足小时数=v['violating_hours'],最大缺口_MW=v['max_deficit_mw']) for v in baseline['records'])
            before={sid:v['initial_units']+[0.] for sid,v in inp['stations'].items()}
            for row in csvrows(root/'stations.csv'):
                sid=row['station'];u=[float(row[f'unit_{k}_mva']) for k in (1,2,3)];n=int(row['storage_modules'])
                stations.append(dict(地区=name,方案=LABEL[case],年份=int(row['year']),变电站=sid,站名=inp['stations'][sid].get('name',sid),
                    主变1_MVA=u[0],主变2_MVA=u[1],主变3_MVA=u[2],年末容量_MVA=sum(u),
                    当年净增容量_MVA=sum(u)-sum(before[sid]),年末储能柜数=n,储能额定功率_MW=n*.1,储能电量_MWh=n*.215))
                before[sid]=u
            edges={e['id']:e for e in inp['network']['edges']}
            for row in csvrows(root/'projects.csv'):
                e=edges[row['id']]
                projects.append(dict(地区=name,方案=LABEL[case],投运年=int(row['year']),联络编号=e['id'],
                    站1=e['station_a'],站2=e['station_b'],容量_MW=e['capacity_mw'],依据=e['basis']))
            for y in inp['years']:
                r=json.loads((root/str(y)/'adopted.json').read_text())
                checks.extend(dict(地区=name,方案=LABEL[case],年份=y,故障站=v['failed'],校核小时数=v['checked_hours'],
                    最大缺口_MW=v['max_deficit_mw'],失败小时数=v['violating_hours'],最小余量_MW=v['minimum_margin_mw']) for v in r['check']['records'])
                actual=r['annual']['added_lifecycle_npv_10k']
                ref=r.get('economic_reference')
                if ref is None:
                    search=r['solver']['conditional_search']
                    usable=search.get('stage','').startswith('cost_') and not search.get('fallback')
                    ref=dict(objective=None,dual_bound=search['dual_bound'] if usable else 0.,status=search['status'])
                    gap=max(0.,(actual-ref['dual_bound'])/max(1.,actual))
                    for record in r['history']:
                        search=record['solver'];phase=record['phase']
                        phases.append(dict(地区=name,年份=y,阶段=phase,目标含义='新增储能及联络费用_万元' if phase=='support' else '新增工程总费用_万元' if phase=='cost' else '联络数量及配套费用缩放项',
                            阶段目标值=search['objective'],阶段下界=search['dual_bound'],阶段间隙比例=search['mip_gap'],
                            全年校核通过=record['verified'],求解状态=search['status']))
                else:gap=r['solver']['mip_gap']
                solver.append(dict(地区=name,方案=LABEL[case],年份=y,年度新增费用现值_万元=actual,
                    年度经济参考_万元=ref['objective'],条件费用下界_万元=ref['dual_bound'],
                    相对参考费用增量=None if ref['objective'] is None else max(0.,actual-ref['objective']),距条件下界比例=gap,
                    经济求解状态=ref['status'],采用状态=r['solver']['status'],
                    下界范围='本年配套费用和联络上限固定后的总费用' if case=='elastic_capacity' else '固定上年资产后的本年总费用'))
                for field,mode in (('normal','正常固定'),('extra','故障站追加')):
                    for key,count in sorted(r[field].items(),key=lambda x:int(x[0])):
                        a=aa[int(key)]
                        switches.append(dict(地区=name,方案=LABEL[case],年份=y,方式=mode,原属站=a['source'],接收站=a['target'],
                            路径=a['id'],源馈线组=a['a'],受端馈线=a['b'],区段数=count,全站区段数=inp['block_count'][a['source']],
                            净功率峰值上界_MW=count*float(abs(powers[y][:,a['sa']]).max())/inp['block_count'][a['source']],
                            路径性质='既有能力代理' if a['existing'] else '新增付费候选'))
                for j,sid in enumerate(ids):
                    existing=sum(c for field in ('normal','extra') for key,c in r[field].items() if aa[int(key)]['source']==sid and aa[int(key)]['existing'])
                    paid=sum(c for field in ('normal','extra') for key,c in r[field].items() if aa[int(key)]['source']==sid and not aa[int(key)]['existing'])
                    used=existing*float(abs(powers[y][:,j]).max())/inp['block_count'][sid]
                    limits.append(dict(地区=name,方案=LABEL[case],年份=y,变电站=sid,既有转供预约峰值上界_MW=used,
                        原表既有能力_MW=inp['existing_transfer_budget_by_year_mw'][str(y)][sid],新增转出区段数=paid,
                        新增转出区段上限=int(inp['block_count'][sid]*inp['new_transfer_fraction']+1e-9)))
        rigid=plans[region,'rigid']['annual'];elastic=plans[region,'elastic_capacity']['annual']
        forecast={r['year']:r for r in p.forecasts(region)};last=None
        for a in rigid:
            f=forecast.get(a['year'],{})
            growth.append(dict(地区=name,年份=a['year'],规划峰值_MW=a['peak_mw'],
                规划峰值较上年增量_MW=None if last is None else a['peak_mw']-last,
                预测总负荷_MW=f.get('forecast_load_mw'),全市光伏装机指数=f.get('pv_index'),
                负荷同比增长率=f.get('load_yoy'),共同容载比下限=a['common_floor']))
            last=a['peak_mw']
        ws=book.create_sheet(f'{name}逐年对照',0 if region=='city' else 1)
        ws.append(['年份','规划峰值/MW','刚性方案',None,None,None,None,'弹性方案',None,None,None,None])
        for area in ('A1:A2','B1:B2','C1:G1','H1:L1'):ws.merge_cells(area)
        columns=['年末容量/MVA','当年净增/MVA','容载比','储能额定功率/MW','累计新线/条']
        ws.append([None,None,*columns,*columns]);md=[];compact=[]
        for a,b in zip(rigid,elastic):
            assert a['year']==b['year'] and a['peak_mw']==b['peak_mw'] and a['common_floor']==b['common_floor']
            keys=['capacity_mva','net_capacity_increment_mva','clr','storage_nameplate_mw','cumulative_lines']
            ws.append([clean(v) for v in [a['year'],a['peak_mw'],*[a[k] for k in keys],*[b[k] for k in keys]]])
            compact.append(dict(年份=a['year'],规划峰值_MW=a['peak_mw'],**{'刚性_'+FIELDS[k]:a[k] for k in keys},**{'弹性_'+FIELDS[k]:b[k] for k in keys}))
            md.append([a['year'],fixed(a['peak_mw']),*[int(a[k]) if k=='cumulative_lines' else fixed(a[k]) for k in keys],
                *[int(b[k]) if k=='cumulative_lines' else fixed(b[k]) for k in keys]])
            issues=[]
            if b['capacity_mva']<=a['capacity_mva']+1e-6:issues.append('容量未高于刚性')
            if b['storage_nameplate_mw']>=a['storage_nameplate_mw']-1e-6:issues.append('储能未少于刚性')
            if b['cumulative_lines']>=a['cumulative_lines']:issues.append('新增联络未少于刚性')
            if a['year']>2021 and issues:exceptions.append(dict(地区=name,年份=a['year'],情况='、'.join(issues)))
            comparison.append(dict(地区=name,年份=a['year'],容量差_弹减刚_MVA=b['capacity_mva']-a['capacity_mva'],
                储能差_弹减刚_MW=b['storage_nameplate_mw']-a['storage_nameplate_mw'],联络差_弹减刚_条=b['cumulative_lines']-a['cumulative_lines'],
                费用差_弹减刚_万元=b['cumulative_lifecycle_npv_10k']-a['cumulative_lifecycle_npv_10k']))
        decorate(ws,2);ws.column_dimensions['A'].width=9;ws.print_area=f'A1:L{ws.max_row}'
        for row in ws.iter_rows(min_row=3):
            for cell in row:cell.number_format='0' if cell.column in (1,7,12) else '0.00'
        p.write_csv(p.PUBLIC/f'{name}逐年左右对照.csv',compact)
        text.extend([f'表 {name}逐年容量及措施','',markdown(['年份','规划峰值/MW','刚性容量/MVA','刚性净增/MVA','刚性容载比','刚性储能/MW','刚性累计新线/条',
            '弹性容量/MVA','弹性净增/MVA','弹性容载比','弹性储能/MW','弹性累计新线/条'],md),''])
    assumptions=[dict(项目=k,说明=v) for k,v in [
        ('年度求解','2022—2035逐年固定上年采用的全部在役设备，只用当年负荷、光伏及正常/N1条件选择当年工程。新购额定容量31.5、50、63、80MVA，允许位置的第三台按50MVA整台投运；同一原主变位置全期最多更换一次。'),
        ('容量控制','同地区同年度容载比下限一致；刚性容量上限为max(上年在役容量,当年规划峰值×2.0)，弹性取消该上限。每年投入由模型决定，可形成容量平台。'),
        ('既有能力','46座有原表对应的站点沿用原Excel J列MW能力上限，所有年份固定封顶；3个模拟站按零。2025能力用于早期年度回算，早期逐线投运资料缺失。'),
        ('共享与路径','既有能力为每站总预算，正常固定供区与故障追加共同占用；每条路径、源馈线组、受端馈线及接收站主变余量同时约束。EQ站对为研究代理，实际逐线拓扑尚待资料闭合。'),
        ('正常及事故','2021校准的既有正常转供区段可回切原属站，既有通道的正常转供上限保持校准值；新增付费联络可形成年度固定供区调整。事故仅从故障站原属源荷区段向直接邻站追加切换，全年共用预案。邳州基期事故缺口单列，2022起补强。'),
        ('新增转供','新增付费路径正常与事故合计最多转出全站20%源荷区段，向下取整；既有转供另受原表MW上限。20%为研究情景假设。EQ聚合端口按能力与基期区段峰值估计代表组数，每组5段、3端点、2邻站。'),
        ('历史源荷','2021—2024使用全市光伏相对2025的装机比例回算，原年度区域正向峰值保持；2025逐点还原，2026—2035沿用原预测。历史曲线为源荷模板回算，含保留标识的补值。'),
        ('储能','0.10MW/0.215MWh整柜，SOC窗口与单向效率合成0.76，支撑取min(额定功率,额定电量×0.76/时长)；市区正向6h、邳州2h。全年静态校核沿用时长条件。'),
        ('弹性选择','主表沿用配套费用优先：各年先减少新增储能与联络费用，在该配套费用候选20%范围内优先减少联络，再比较总费用。主变按本年需求选择。另保留总费用参考5%容差的弹性经济对照。阶段下界及间隙分别记录，无刚弹排序约束。'),
        ('费用','2021现值基年，评价至2041，次年起计运维；寿命40/40/25年，运维4.2%/2.2%/3%，折现6%。只计新增采购；新购主变容量扣被替换容量得到净增容量。'),
        ('适用范围','回算和规划方案受候选网络、等比例源荷区段及早期空间模板假设影响。全年8760小时正常及各站最大一台主变故障静态校核通过，储能未做连续SOC调度；年度条件下界不构成全期联合最优证明。'),
    ]]
    costcheck=[]
    for region,name in p.REGIONS.items():
        for case in CASES:
            s=plans[region,case];last=s['annual'][-1]
            costcheck.append(dict(地区=name,方案=LABEL[case],主变费用现值_万元=sum(r['transformer_npv_10k'] for r in s['annual']),
                储能费用现值_万元=sum(r['storage_npv_10k'] for r in s['annual']),联络费用现值_万元=sum(r['line_npv_10k'] for r in s['annual']),
                合计费用现值_万元=s['cost_npv_10k'],终期容量_MVA=last['capacity_mva']))
    datasets={'逐年新增及费用':annual,'费用分项汇总':costcheck,'逐站逐年配置':stations,'联络逐项工程':projects,
        '固定区段切换预案':switches,'站级转供上限核对':limits,'年度费用下界与间隙':solver,'弹性优选阶段':phases,'全年安全核验':checks,
        '2021固定基线检验':baselines,'刚弹完整序列比较':comparison,'负荷与逐年增量':growth,'定义与假设':assumptions}
    for name,data in datasets.items():
        if data:
            ws=flat(book,name,data);p.write_csv(p.PUBLIC/(name+'.csv'),data)
            if name=='定义与假设':
                ws.column_dimensions['A'].width=22;ws.column_dimensions['B'].width=118
                for row in range(2,ws.max_row+1):ws.row_dimensions[row].height=max(42,20*((len(ws.cell(row,2).value)+49)//50))
            if name=='年度费用下界与间隙':
                ws.column_dimensions['I'].width=26;ws.column_dimensions['J'].width=65
                for row in range(2,ws.max_row+1):ws.row_dimensions[row].height=42
            if name=='联络逐项工程':
                ws.column_dimensions['H'].width=65
                for row in range(2,ws.max_row+1):ws.row_dimensions[row].height=42
    path=p.PUBLIC/'两地区逐年容量及措施_需求驱动左右对照.xlsx';book.save(path)
    rb=load_workbook(path,data_only=True);cells=0
    for ws in book:
        actual=rb[ws.title];assert actual.max_row==ws.max_row and actual.max_column==ws.max_column
        for row in ws:
            for c in row:
                value=actual.cell(c.row,c.column).value
                if isinstance(c.value,float):assert abs(value-c.value)<max(1e-8,abs(c.value)*1e-12)
                else:assert (value or '')==(c.value or '')
                cells+=1
    rb.close()
    p.dump(p.REVIEW/'工作簿回读.json',dict(verified=True,cells=cells,path=str(path.relative_to(p.ROOT)),sha256=p.sha(path)))
    text.extend(['注：储能为年末额定功率，累计新线为2022年起累计新增10kV联络。主变净增容量已扣除被替换容量。','',
        markdown(['地区','方案','主变费用/万元','储能费用/万元','联络费用/万元','合计费用现值/万元'],[
            [r['地区'],r['方案'],*[fixed(r[k]) for k in ('主变费用现值_万元','储能费用现值_万元','联络费用现值_万元','合计费用现值_万元')]] for r in costcheck]),'',
        f'主表四套方案2022—2035共{proof["primary_hour_pairs"]:,}个小时—工况组合通过静态安全复核，含弹性经济对照共{proof["full_hour_pairs"]:,}项。邳州2021固定设备仍有事故缺口，单列于基线检验页；基期不加入补强设备。','',
        '源荷历史处理及转供模型已调整，当前结果与上一版固定远景清单分期排程属于不同模型条件。容量平台、刚弹相同及排序例外在完整序列中据实保留。', '',
        '适用条件：2025转供收资作为早期回算上限，EQ连接与源荷区段为研究代理；储能采用静态时长校核。费用为逐年条件可行方案费用，求解间隙见工作簿。'])
    (p.PUBLIC/'逐年容量及措施_需求驱动左右对照.md').write_text('\n'.join(text)+'\n')
    p.dump(p.REVIEW/'结果序列核对.json',dict(verified=True,ranking_exceptions=exceptions,rows=comparison))
    p.dump(p.SOURCE/'展示口径.json',dict(primary=str(path.relative_to(p.ROOT)),source_manifest=proof['solution_sources'],assumptions=assumptions))
    print('EXPORTED_DEMAND',cells,'ranking_exceptions',len(exceptions),flush=True)


if __name__=='__main__':main()
