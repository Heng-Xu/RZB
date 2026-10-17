"""从交付工作簿回读主表，以浏览器渲染核查对齐、表头和全部年度。"""
from html import escape
import argparse
import json
from pathlib import Path
import subprocess

from openpyxl import load_workbook

from . import prepare as p


def run(prepare_only=False,collect_only=False):
    source=p.PUBLIC/'两地区逐年容量及措施_需求驱动左右对照.xlsx'
    book=load_workbook(source,data_only=True);rendered=[]
    for name in ('市区','邳州'):
        ws=book[f'{name}逐年对照'];covered=set();starts={}
        for group in ws.merged_cells.ranges:
            starts[group.min_row,group.min_col]=(group.max_row-group.min_row+1,group.max_col-group.min_col+1)
            covered.update((r,c) for r in range(group.min_row,group.max_row+1) for c in range(group.min_col,group.max_col+1)
                if (r,c)!=(group.min_row,group.min_col))
        rows=[]
        for row in ws:
            cells=[]
            for cell in row:
                if (cell.row,cell.column) in covered:continue
                rs,cs=starts.get((cell.row,cell.column),(1,1));value=cell.value
                if isinstance(value,(int,float)) and '0.00' in cell.number_format:
                    value=f'{value:.2%}' if '%' in cell.number_format else f'{value:.2f}'
                value='' if value is None else str(value)
                tag='th' if cell.row<=2 else 'td'
                cells.append(f'<{tag} rowspan="{rs}" colspan="{cs}">{escape(value)}</{tag}>')
            rows.append('<tr>'+''.join(cells)+'</tr>')
        html='''<!doctype html><meta charset="utf-8"><style>
        *{box-sizing:border-box}body{margin:35px;font-family:"方正仿宋_GBK","Noto Sans CJK SC",serif;color:#172d41}
        h1{font-size:28px;margin:0 0 14px}p{font-size:19px;line-height:1.5;margin:8px 0 18px}
        table{border-collapse:collapse;width:100%;table-layout:fixed}th,td{text-align:center;border:1px solid #ced6df;padding:12px 6px;font-size:19px;height:44px;line-height:1.35}
        th{background:#244b70;color:white}tr:nth-child(odd) td{background:#f0f4f8}td:first-child,th:first-child{width:6%}
        </style>'''
        html+=f'<h1>{name}2021—2035逐年容量及措施</h1><p>以上年在役设备和当年需求逐年求解。刚性在左，弹性在右；储能为年末额定功率，新线为自2022年起累计新增联络。</p>'
        html+='<table>'+''.join(rows)+'</table><p>数据为研究模型回算及规划结果。既有转供按原收资MW封顶，路径及受端余量同时约束；2021基期事故缺口单列。</p>'
        page=p.REVIEW/f'{name}主表预览.html';page.write_text(html)
        output=p.REVIEW/f'{name}主表预览.png'
        if prepare_only:continue
        if not collect_only:
            subprocess.run(['google-chrome','--headless','--no-sandbox','--disable-gpu','--disable-background-networking','--hide-scrollbars',
                '--user-data-dir=/tmp/annual-demand-preview-chrome','--window-size=2200,1300',f'--screenshot={output}',page.as_uri()],
                check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        rendered.append(dict(sheet=ws.title,rows=ws.max_row,columns=ws.max_column,image=str(output.relative_to(p.ROOT)),sha256=p.sha(output)))
    book.close()
    if prepare_only:return
    p.dump(p.REVIEW/'主表渲染.json',dict(workbook_sha256=p.sha(source),rendered=rendered))
    print('MAIN_TABLES_RENDERED',flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--prepare-only',action='store_true');parser.add_argument('--collect-only',action='store_true')
    a=parser.parse_args();run(a.prepare_only,a.collect_only)
