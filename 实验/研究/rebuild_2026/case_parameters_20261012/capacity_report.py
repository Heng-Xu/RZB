"""导出刚性、弹性及下限对照，并明确首年补强和验证边界。"""
import json
from decimal import Decimal,ROUND_HALF_UP

import numpy as np
from openpyxl import Workbook,load_workbook

from .capacity_run import OUT,RESULT,REVIEW,CASES,previous,load_scenes,dump,write_csv,sha,ROOT
from .capacity_audit import FlowLP
from .report import add_sheet,markdown

LABEL={'rigid_storage':'刚性','elastic_floor_2':'弹性_有2.0下限','elastic_no_floor':'弹性_无2.0下限'}


def two(value):
    return str(Decimal(str(value)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))


def baseline_audit(inp,scenes):
    lp=FlowLP(inp);units=np.array([s['initial_units']+[0.] for s in inp['stations'].values()]);n=np.zeros(len(units))
    results=[]
    for year in (2021,2022):
        power=scenes[2022]*inp['annual_observed_peak_mw'][str(year)]/inp['annual_observed_peak_mw']['2022']
        normal=0;faults=0;worst=0.;critical=None
        for h,p in enumerate(power):
            ready=lp.prepared(p,year,units,n,set())
            for f in range(-1,len(units)):
                r=lp.check(p,ready,f)
                if not r['verified']:
                    if f<0:normal+=1
                    else:faults+=1
                    if r['deficit_mw']>worst:worst=r['deficit_mw'];critical=[h,lp.ids[f] if f>=0 else 'normal']
        results.append(dict(地区=inp['region_name'],方案='压缩设备_未补强',年份=year,容量_MVA=float(units.sum()),
            容载比=float(units.sum()/inp['annual_observed_peak_mw'][str(year)]),正常失败工况数=normal,
            N1失败工况数=faults,最大承载缺口_MW=worst,总检查工况数=len(power)*(len(units)+1),
            最差工况=json.dumps(critical,ensure_ascii=False)))
    return results


def report():
    verify=json.loads((REVIEW/'矩阵费用及独立承载力核验.json').read_text());assert verify['verified']
    RESULT.mkdir(parents=True,exist_ok=True)
    params=[];annual=[];costs=[];stations=[];projects=[];first=[];baseline=[];data={}
    for region in ('city','pizhou'):
        root=OUT/region;inp=json.loads((root/'inputs.json').read_text());name=inp['region_name'];data[region]={}
        baseline.extend(baseline_audit(inp,load_scenes(root,inp)))
        params.append(dict(地区=name,收资参考转供率=inp['initial_transfer_fraction'],参考年份=2025,
            电缆长度占比=inp['line_cable_fraction'],新增10kV联络单项造价_万元=inp['line_capex_10k'],
            收资条件标定能力=inp['transfer_reference_audit']['weighted_rate'],
            压缩起点重算可实现能力=inp['compressed_initial_transfer_audit']['weighted_rate'],
            标定参考设备容量_MVA=sum(inp['transfer_calibration_capacity_mva'].values()),
            共同压缩起点_MVA=inp['initial_capacity_mva'],静态同期截面数=inp['capacity_snapshot_count'],
            正向储能持续时间_h=inp['forward_duration_h'],既有通道处理='在2025参考设备下标定后固定，接收容量按各年决策设备校核'))
        for case in CASES:
            out=root/case;s=json.loads((out/'summary.json').read_text());data[region][case]=s;sol=s['solver']
            checked=verify['regions'][region][case]
            assert checked['verified'] and checked['input_sha256']==sha(root/'inputs.json')
            assert checked['solution_sha256']==sha(out/'chosen_solution.npz')
            assert abs(checked['cost_recomputed_10k']-s['cost_npv_10k'])<1e-5
            costs.append(dict(地区=name,方案=LABEL[case],全周期费用现值_万元=s['cost_npv_10k'],
                主变费用现值_万元=sum(a['transformer_npv_10k'] for a in s['annual']),
                联络费用现值_万元=sum(a['line_npv_10k'] for a in s['annual']),
                储能费用现值_万元=sum(a['storage_npv_10k'] for a in s['annual']),
                费用下界_万元=sol['dual_bound'],相对间隙=sol['mip_gap'],
                **{'2025年末容量_MVA':s['annual'][-1]['capacity_mva'],
                   '2025容载比':s['annual'][-1]['clr'],'2025累计新增联络_项':s['annual'][-1]['cumulative_lines'],
                   '2025年末储能_MWh':s['annual'][-1]['storage_mwh']}))
            for a in s['annual']:
                annual.append(dict(地区=name,方案=LABEL[case],年份=a['year'],地区峰值负荷_MW=a['peak_mw'],
                    年末容量_MVA=a['capacity_mva'],容载比=a['clr'],当年净增容量_MVA=a['net_capacity_increment_mva'],
                    当年购入主变容量_MVA=a['purchased_transformer_mva'],当年新增联络_项=a['new_lines'],
                    年末累计新增联络_项=a['cumulative_lines'],当年新增储能_MWh=a['new_storage_mwh'],
                    年末储能_MW=a['storage_nameplate_mw'],年末储能_MWh=a['storage_mwh'],
                    当年建设投资_万元=a['construction_capex_10k'],当年费用现值_万元=a['added_lifecycle_npv_10k'],
                    累计费用现值_万元=a['cumulative_lifecycle_npv_10k'],刚性被动超限=a['passive_exceedance']))
            a=s['annual'][0];check=s['static_network_checks'][0]
            first.append(dict(地区=name,方案=LABEL[case],年份=2022,起点容量_MVA=inp['initial_capacity_mva'],
                首年净增容量_MVA=a['net_capacity_increment_mva'],首年末容量_MVA=a['capacity_mva'],首年容载比=a['clr'],
                首年新增储能_MWh=a['new_storage_mwh'],首年新增联络_项=a['new_lines'],
                检查工况数=check['checked_pairs'],失败工况数=check['violating_pairs'],最大承载缺口_MW=check['max_deficit_mw']))
            old={sid:asset['initial_units']+[0.] for sid,asset in inp['stations'].items()};oldn={s:0 for s in inp['stations']}
            for r in previous.rows(out/'stations.csv'):
                sid=r['station'];units=[float(r[f'unit_{k}_mva']) for k in (1,2,3)];n=int(r['storage_modules'])
                stations.append(dict(地区=name,方案=LABEL[case],年份=int(r['year']),站点=sid,
                    调整前主变_MVA='/'.join(f'{v:g}' for v in old[sid] if v),年末主变_MVA='/'.join(f'{v:g}' for v in units if v),
                    年末容量_MVA=sum(units),当年净增容量_MVA=sum(units)-sum(old[sid]),
                    当年购入主变容量_MVA=float(r['purchased_transformer_mva']),当年新增储能_MWh=.215*(n-oldn[sid]),
                    年末储能_MW=.1*n,年末储能_MWh=.215*n))
                old[sid]=units;oldn[sid]=n
            for r in previous.rows(out/'projects.csv'):
                projects.append(dict(地区=name,方案=LABEL[case],投运年份=int(r['year']),项目=r['id'],站A=r['station_a'],站B=r['station_b'],
                    馈线A=r['a'],馈线B=r['b'],单项能力_MW=float(r['capacity_mw']),占用备用间隔站=r['new_bay_station'],
                    单项建设投资_万元=inp['line_capex_10k'],候选依据=r['basis']))
    assumptions=[dict(事项='收资转供能力',说明='全部电压等级ΣJ/(ΣJ+ΣK)，作为2025收资设备条件下的地区参考指标；保留完整计算精度。'),
        dict(事项='既有网络',说明='聚合通道限定在原候选站对，先按2025设备和同期负荷标定，再冻结通道、热限与既有预算。压缩设备后重新校核接收站余量。'),
        dict(事项='实际与回算',说明='地区年度容量负荷为实际统计；共同压缩设备和寻优措施为仿真回算。市区样本29站（含3座原模拟站）、邳州20站，与地区全站统计范围分别说明。'),
        dict(事项='静态截面',说明='市区45、邳州26个同期截面，包括地区正向峰值、最低净负荷及各站正向峰值、最低净负荷。按原年度统计峰值统一缩放，逐站峰值保持真实同期关系。'),
        dict(事项='数据补齐',说明='沿用现有站级8760序列及补值标记；市区421小时含补值，邳州1小时含补值。原始样本和补值日志保留，有限静态截面不等同全年运行校验。'),
        dict(事项='正常及N-1',说明='各工况校核正反向功率；最大单台主变故障仅作用于一个站，其余站按正常容量。共享馈线余量、既有通道预算、接收站容量和新出线备用间隔同时限制转供。'),
        dict(事项='刚性规则',说明='容量非减；C_y≤max(C_(y−1),2P_y)。新增后容量受当年2.0上限限制，负荷回落时允许已有容量被动超限；沿用原增长档下限。'),
        dict(事项='弹性规则',说明='无下限弹性由费用目标、设备和安全约束决定容量；有下限弹性仅另加C_y/P_y≥2.0。两组均保留物理设备上界。'),
        dict(事项='线路独立计价',说明='市区电缆占比80%，邳州按403条去重线路长度计算10.76390458%。单项3km×[24×(1-q)+180×q]+6万元，长度及端部价沿用原假设。'),
        dict(事项='其余参数',说明='原设备候选、主变/线路/储能价格、持续时间（市区6h、邳州2h）、能量系数0.76、寿命、维护费率、折现率及评价期保留。'),
        dict(事项='结果使用',说明='工程应用需用实际既有联络端点和接收余量核实聚合通道。市区零新增投资是当前等效网络和静态样本下的费用寻优结果。'),
        dict(事项='费用精度',说明='本版比较通过校核的可行方案费用；求解下界和间隙独立列示。设备购入整台容量与净增容量分列。')]
    frames={'地区参数':params,'首年补强与验收':first,'压缩设备直接校核':baseline,'方案费用':costs,
            '逐年方案':annual,'逐站设备':stations,'新增联络':projects,'模型说明':assumptions}
    book=Workbook();book.remove(book.active)
    for name,rows in frames.items():write_csv(RESULT/f'{name}.csv',rows);add_sheet(book,name,rows)
    for row in book['地区参数'].iter_rows(min_row=2):
        for j in (2,4,6,7):row[j-1].number_format='0.00000000'
    path=RESULT/'两地区刚性与弹性方案_承载力复核.xlsx';book.save(path)
    w=load_workbook(path,data_only=True);count=0
    for name,rows in frames.items():
        actual=list(w[name].iter_rows(values_only=True));expected=[list(rows[0])]+[[r.get(k) for k in rows[0]] for r in rows]
        assert len(actual)==len(expected)
        for a,b in zip(actual,expected,strict=True):
            for x,y in zip(a,b,strict=True):
                if isinstance(y,float):assert abs(x-y)<max(1e-8,abs(y)*1e-12)
                else:assert (x if x is not None else '')==(y if y is not None else '')
                count+=1
    w.close()
    old_source=ROOT/'研究报告/试算/20261012_案例参数修正_下限对照/pizhou/no_floor/summary.json'
    first_trial=json.loads(old_source.read_text())['annual'][0]
    old_audit=json.loads((ROOT/'研究报告/05_review/当前版/1012案例参数修正/首年承载力审查.json').read_text())
    failed=next(r['失败工况数'] for r in old_audit['summary'] if r['地区']=='邳州' and r['方案']=='首版无下限首年方案_2022')
    elastic=data['pizhou']['elastic_no_floor'];rigid=data['pizhou']['rigid_storage']
    ea=elastic['annual'][0];ra=rigid['annual'][0];last=elastic['annual'][-1]
    savings=rigid['cost_npv_10k']-elastic['cost_npv_10k']
    text=['# 两地区刚性与弹性方案及首年承载力复核','',
        '本版补齐各站局部极值同期工况，并将既有通道先在2025收资设备条件下标定、再固定用于压缩设备方案。主结果为刚性和无下限弹性，另保留弹性2.0下限对照。','',
        '## 首年压缩设备需要先通过承载校核','',
        markdown(['地区','直接保留压缩设备的年份','容量/MVA','容载比','正常失败工况','N−1失败工况','最大静态缺口/MW'],
          [[r['地区'],r['年份'],f"{r['容量_MVA']:.2f}",f"{r['容载比']:.2f}",r['正常失败工况数'],r['N1失败工况数'],f"{r['最大承载缺口_MW']:.2f}"] for r in baseline]),'',
        '邳州压缩设备需要补强后进入首个规划年。以下首年方案均由正常与N−1约束选出，补强在2022年计费。','',
        markdown(['地区','方案','首年末容量/MVA','首年容载比','首年净增容量/MVA','首年新增储能/MWh','首年校核'],
          [[r['地区'],r['方案'],f"{r['首年末容量_MVA']:.2f}",f"{r['首年容载比']:.2f}",f"{r['首年净增容量_MVA']:.2f}",f"{r['首年新增储能_MWh']:.2f}",f"{r['检查工况数']}工况通过"] for r in first]),'',
        '## 四年费用及设备结果','',
        markdown(['地区','方案','费用现值/万元','2025容量/MVA','2025容载比','累计新增联络/项','2025储能/MWh'],
          [[r['地区'],r['方案'],f"{r['全周期费用现值_万元']:.2f}",f"{r['2025年末容量_MVA']:.2f}",f"{r['2025容载比']:.2f}",r['2025累计新增联络_项'],two(r['2025年末储能_MWh'])] for r in costs]),'',
        '## 逐年措施','',
        markdown(['地区','方案','年份','年末容量/MVA','容载比','当年净增容量/MVA','当年新增联络/项','当年新增储能/MWh'],
          [[r['地区'],r['方案'],r['年份'],f"{r['年末容量_MVA']:.2f}",f"{r['容载比']:.2f}",f"{r['当年净增容量_MVA']:.2f}",r['当年新增联络_项'],two(r['当年新增储能_MWh'])] for r in annual]),'',
        '## 如何解释容载比','',
        '容载比是地区容量与年度峰值负荷之比。安全约束在每个站、每个工况上校核正常及故障后正反向承载力，同时检查转供路径、馈线共享余量和接收站容量。因此，容载比超过2.0的方案仍可能存在局部缺口；满足已建安全约束的方案也可能低于2.0。','',
        f"本轮首版邳州首年容量{first_trial['capacity_mva']:.2f} MVA、容载比{first_trial['clr']:.2f}，通过两类地区工况，但遗漏的局部工况有{failed}项未通过。补齐工况后，无下限弹性方案首年自动选择{ea['capacity_mva']:.2f} MVA、容载比{ea['clr']:.2f}。该方案到2025年容载比为{last['clr']:.2f}，并配置联络和储能满足所检工况。",'',
        f"邳州刚性首年容量{ra['capacity_mva']:.2f} MVA、容载比{ra['clr']:.2f}，需要配置{ra['storage_mwh']:.2f} MWh储能；无下限弹性首年多配置{ea['capacity_mva']-ra['capacity_mva']:.2f} MVA主变容量并减少储能，四年可行方案费用减少{savings:.2f}万元（{savings/rigid['cost_npv_10k']:.0%}）。市区在当前较高既有转供能力条件下，两主方案保持相同配置。",'',
        '实际年度统计反映现实投运设备；本模型的压缩起点与回算建设措施属于另一组设备配置。实际容载比大于2.0用于对照设备与负荷范围，工程推荐还需核实等效既有通道。','',
        '## 参数与适用条件','']+[r['事项']+'：'+r['说明']+'\n' for r in assumptions]+['',
        markdown(['地区','方案','费用下界/万元','相对间隙'],[[r['地区'],r['方案'],f"{r['费用下界_万元']:.2f}",f"{r['相对间隙']*100:.2f}%"] for r in costs]),'',
        f"重新复建约束矩阵、独立LP及直接零转供证书共检查{verify['checked_pairs']}个静态工况，设备与三项费用复算通过。逐站规格、购入容量及联络端点见工作簿。正式六章与Word/PDF、已锁定推广矩阵保持原版。"]
    (RESULT/'结果与承载力说明.md').write_text('\n'.join(text)+'\n',encoding='utf-8')
    dump(REVIEW/'导出核验.json',dict(verified=True,workbook=str(path.relative_to(ROOT)),sha256=sha(path),checked_cells=count,
        baseline_audit=baseline,first_year=first,checked_pairs=verify['checked_pairs']))
    print('EXPORTED_CAPACITY',path,'cells',count,flush=True)


if __name__=='__main__':report()
