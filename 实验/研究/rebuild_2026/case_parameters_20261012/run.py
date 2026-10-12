"""保留实际年度案例资产、负荷与费用，仅修正转供、电缆及下限。

既有等效通道的站对限制在原候选图内；新增项目沿用原馈线端点。
地区全部电压等级转供统计用于共同起点网络的情景标定，不作历史复原。
"""
import argparse
import copy
import csv
import json
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix

from ..current_economic_assumptions import line_capex
from ..network_case_20261010 import run as network
from ..network_case_20261010.solver import solve
from ..pilot_20261009.run import ROOT, PF, dump, write_csv, sha
from ..static_actual_20261010 import run as previous
from ..synchronous_20261009 import model as assets_model

OUT = ROOT / '研究报告/试算/20261012_案例参数修正_下限对照'
SOURCE = ROOT / '研究报告/数据来源/当前版/案例参数修正_20261012'
REVIEW = ROOT / '研究报告/05_review/当前版/1012案例参数修正'
RESULT = ROOT / '研究报告/01研究推荐/当前版/案例参数修正_20261012'
TRANSFER_FILE = ROOT / '实验/研究/data/tuomin/补充数据/徐州地区变电站全停全转情况分析汇总20250705.xlsx'
CABLE_FILE = ROOT / '实验/研究/data/tuomin/电网建模数据_Agent整合版_V1.2/邳州10kV线路基础数据表.xlsx'
CASES = ('floor_2', 'no_floor')


def transfer_evidence():
    w = load_workbook(TRANSFER_FILE, read_only=True, data_only=True)
    records = []
    for nr, row in enumerate(w.worksheets[0].iter_rows(min_row=3, values_only=True), 3):
        if row[0] not in ('徐州市区', '邳州') or not row[2]:
            continue
        if not all(isinstance(row[k], (int, float)) for k in (9, 10)):
            continue
        assert row[9] >= 0 and row[10] >= 0
        records.append(dict(地区=row[0], 电压等级=row[1], 变电站=row[2], 原表行号=nr,
                            可转供负荷_MW=row[9], 损失负荷_MW=row[10]))
    w.close()
    stats = []
    for name in ('徐州市区', '邳州'):
        for scope in ('全部电压等级', '仅110 kV'):
            chosen = [r for r in records if r['地区'] == name and
                      (scope == '全部电压等级' or r['电压等级'] == '110kV')]
            supply = sum(r['可转供负荷_MW'] for r in chosen)
            loss = sum(r['损失负荷_MW'] for r in chosen)
            stats.append(dict(地区=name, 统计范围=scope, 站点数=len(chosen),
                              可转供负荷_MW=supply, 损失负荷_MW=loss, 转供率=supply/(supply+loss)))
    return records, stats


def cable_evidence():
    w = load_workbook(CABLE_FILE, read_only=True, data_only=True)
    records = {}
    for nr, r in enumerate(w.worksheets[0].iter_rows(min_row=2, values_only=True), 2):
        if not r[1] or not all(isinstance(r[k], (float, int)) for k in (13, 15)):
            continue
        item = dict(线路=r[1], 原表行号=nr, 总长_km=r[13], 电缆长度_km=r[15])
        if r[1] in records:
            assert all(records[r[1]][k] == item[k] for k in ('总长_km', '电缆长度_km'))
        else:
            records[r[1]] = item
    w.close()
    total = sum(r['总长_km'] for r in records.values())
    cable = sum(r['电缆长度_km'] for r in records.values())
    assert len(records) == 403 and 0 < cable < total
    return list(records.values()), cable/total


def equivalent_network(inp, scale):
    """原候选站对上的两馈线组聚合通道；保留已闭合T01路径。"""
    p = copy.deepcopy(inp)
    ids = list(p['stations'])
    pairs = sorted({tuple(sorted((e['station_a'], e['station_b'])))
                    for e in p['network']['edges'] if not e['id'].startswith('EQ-')})
    connected = {s for pair in pairs for s in pair}
    for sid in sorted(connected):
        base = p['transfer_calibration_forward_mw'][sid]
        for group in range(2):
            fid = f'EQ-{sid}-{group}'
            p['network']['feeders'][fid] = dict(id=fid, station=sid, grid='calibrated_existing',
                thermal_mw=base*(1+scale)/2,
                maximum_load_proxy_2025_mw=base*p['annual_observed_peak_mw']['2025']/p['annual_observed_peak_mw'][str(p.get('transfer_calibration_year',2021))]/2,
                source_row='aggregate_transfer_calibration', source_reported_active_mw=None,
                equivalent_existing=True)
    for n, (a, b) in enumerate(pairs):
        for group in range(2):
            p['network']['edges'].append(dict(id=f'EQ-{n:03d}-{group}',
                a=f'EQ-{a}-{group}', b=f'EQ-{b}-{1-group}', station_a=a, station_b=b,
                capacity_mw=scale*max(p['transfer_calibration_forward_mw'][a], p['transfer_calibration_forward_mw'][b])/2,
                existing=True, grid='calibrated_existing', source_section_fraction=1.,
                bidirectional=True, route_length_km_assumption=0., new_bay_station='',
                basis='原候选站对范围内按地区转供率标定的既有聚合通道；路径为研究假设'))
    p['existing_transfer_budget_mw'] = {s:scale*p['transfer_calibration_forward_mw'][s] for s in ids}
    return p


def limits(inp,p,y):
    donor,receive=previous.actual_limits(inp,p,y)
    index={s:j for j,s in enumerate(inp['stations'])}
    annual=inp['annual_observed_peak_mw'][str(y)]/inp['annual_observed_peak_mw'][str(inp.get('transfer_calibration_year',2021))]
    for f,r in inp['network']['feeders'].items():
        if r.get('equivalent_existing'):
            sid=r['station']
            donor[f]=max(0.,float(p[index[sid]]))/2
            receive[f]=max(0.,r['thermal_mw']-inp['transfer_calibration_forward_mw'][sid]*annual/2)
    return donor,receive


def initial_transfer_audit(inp):
    """逐个送端全停的星形LP，接收站按共同起点正常主变余量。"""
    ids = list(inp['stations'])
    base = np.array([inp['transfer_calibration_forward_mw'][s] for s in ids])
    raw=inp.get('transfer_calibration_net_mw',inp['transfer_calibration_forward_mw'])
    assert abs(base.sum()-sum(max(v,0.) for v in raw.values()))<1e-5
    donor, receive = limits(inp, base, inp.get('transfer_calibration_year',2021))
    capacity=inp.get('transfer_calibration_capacity_mva',{s:sum(inp['stations'][s]['initial_units']) for s in ids})
    spare = {s:max(0., PF*capacity[s]-base[j]) for j,s in enumerate(ids)}
    values = []
    for sid in ids:
        arcs = []
        for e in inp['network']['edges']:
            if not e['existing']:
                continue
            for direction in range(2 if e['bidirectional'] else 1):
                a,b = (e['a'],e['b']) if direction == 0 else (e['b'],e['a'])
                if inp['network']['feeders'][a]['station'] != sid:
                    continue
                cap = min(e['capacity_mw'], donor[a]*e['source_section_fraction'])
                arcs.append((a,b,inp['network']['feeders'][b]['station'],cap))
        if not arcs:
            values.append(0.); continue
        rows = []; rhs = []
        for field, bounds_by_key in ((0,donor), (1,receive), (2,spare)):
            for key in sorted({a[field] for a in arcs}):
                rows.append([float(a[field] == key) for a in arcs]); rhs.append(bounds_by_key[key])
        rows.append(np.ones(len(arcs))); rhs.append(min(base[ids.index(sid)],inp['existing_transfer_budget_mw'][sid]))
        r = linprog(-np.ones(len(arcs)), A_ub=np.array(rows), b_ub=np.array(rhs),
                    bounds=[(0,a[3]) for a in arcs], method='highs')
        assert r.success, r.message
        values.append(max(0.,float(-r.fun)))
    return dict(reference_year=inp.get('transfer_calibration_year',2021),station_ids=ids, station_transfer_mw=values,
                station_rates=np.divide(values,base,out=np.zeros(len(base)),where=base>0).tolist(), weighted_rate=sum(values)/sum(base),
                basis='各站分别作为唯一送端；在指定参考年设备及负荷下，接收馈线和健康站主变余量共同限制')


def calibrate(inp, target):
    low, high = 0., 8.
    maximum = initial_transfer_audit(equivalent_network(inp, high))
    if maximum['weighted_rate'] < target-1e-8:
        raise ValueError(f"原候选站对最大可标定转供率{maximum['weighted_rate']:.6f}低于{target:.6f}")
    for _ in range(32):
        mid = (low+high)/2
        if initial_transfer_audit(equivalent_network(inp,mid))['weighted_rate'] < target:
            low = mid
        else:
            high = mid
    p = equivalent_network(inp,(low+high)/2)
    audit = initial_transfer_audit(p)
    assert abs(audit['weighted_rate']-target) < 1e-7
    p['initial_transfer_fraction'] = target
    p['initial_network_scale'] = (low+high)/2
    p['initial_transfer_audit'] = audit
    p['initial_transfer_fraction_usage'] = '共同起点既有等效网络标定；全部电压等级地区统计作为研究输入'
    return p


def prepare(region_names=('city','pizhou')):
    SOURCE.mkdir(parents=True,exist_ok=True)
    rr, stats = transfer_evidence()
    cc, cable = cable_evidence()
    write_csv(SOURCE/'转供原表逐站提取.csv',rr)
    write_csv(SOURCE/'转供统计范围.csv',stats)
    write_csv(SOURCE/'邳州电缆长度提取.csv',cc)
    rates = {('city' if r['地区']=='徐州市区' else 'pizhou'):r['转供率']
             for r in stats if r['统计范围']=='全部电压等级'}
    provenance = dict(转供原表=dict(path=str(TRANSFER_FILE.relative_to(ROOT)),sha256=sha(TRANSFER_FILE)),
                      邳州线路原表=dict(path=str(CABLE_FILE.relative_to(ROOT)),sha256=sha(CABLE_FILE)),
                      市区电缆比例=dict(value=.8,source='用户2026年10月12日提供约80%，无最新逐线长度明细'),
                      转供口径='Σ站外最大可供负荷/(Σ站外最大可供负荷+Σ损失负荷)，全部电压等级',
                      基期假设='2025年转供统计用于共同仿真起点；保持年度实际负荷与原模拟设备',
                      下限对照='两组均允许突破2.0上限；仅有下限组要求C/P≥2.0，无下限组取消该条件',
                      网络假设='标定通道限制在原候选站对；新项目保留原馈线端点及间隔约束；三座模拟站保持无网络')
    if (SOURCE/'来源及参数.json').exists():
        provenance=json.loads((SOURCE/'来源及参数.json').read_text()) | provenance
    for region in region_names:
        root = OUT/region
        if (root/'inputs.json').exists():
            raise FileExistsError(f'输入已存在，禁止重新准备覆盖：{root}')
        root.mkdir(parents=True,exist_ok=True)
        old = json.loads((previous.OUT/region/'inputs.json').read_text())
        seed=copy.deepcopy(old)
        forward=previous.load_scenes(previous.OUT/region,old)[2025][0]
        seed['transfer_calibration_net_mw']={s:float(forward[j]*old['initial_peak_mw']/old['annual_observed_peak_mw']['2025'])
                                           for j,s in enumerate(old['stations'])}
        seed['transfer_calibration_forward_mw']={s:max(v,0.) for s,v in seed['transfer_calibration_net_mw'].items()}
        p = calibrate(seed, rates[region])
        p['line_cable_fraction'] = .8 if region=='city' else cable
        p['line_capex_10k'] = line_capex(p['line_cable_fraction'])
        p['line_cost_coefficient'] = 1+p['line_cable_fraction']*(180/24-1)
        p['scope'] = '原实际年度静态案例参数修正：地区转供标定、电缆造价；2.0下限有无对照'
        p['transfer_calibration_assumption'] = provenance['基期假设']
        p['line_basis'] = '新增10kV联络独立计费：3km×[24×(1-q)+180×q]+6万元/项；原长度和端部价假设'
        p['source_input_sha256'] = sha(previous.OUT/region/'inputs.json')
        p['source_scene_sha256'] = sha(previous.OUT/region/'静态工况.csv')
        p['paired_floor_values'] = [2.,0.]
        p['parameter_changes'] = {k:dict(before=old[k],after=p[k]) for k in
                                  ('initial_transfer_fraction','line_cable_fraction','line_capex_10k')}
        dump(root/'inputs.json',p)
        (root/'静态工况.csv').write_bytes((previous.OUT/region/'静态工况.csv').read_bytes())
        write_csv(root/'网络及候选.csv',p['network']['edges'])
        provenance[region] = dict(原输入SHA256=p['source_input_sha256'],静态工况SHA256=p['source_scene_sha256'],
                                  参数修正=p['parameter_changes'],转供标定=p['initial_transfer_audit'])
        print('PREPARED',region,'transfer',p['initial_transfer_audit']['weighted_rate'],
              'cable',p['line_cable_fraction'],'line_cost',p['line_capex_10k'],flush=True)
    dump(SOURCE/'来源及参数.json',provenance)


def build(inp,scenes,floor):
    saved_limits=network.limits
    network.limits = limits
    try:
        m, meta = previous.build(inp,scenes,'elastic_storage')
    finally:
        network.limits=saved_limits
    # 两组只改变同一容量行的下界；既有资产非减、物理上界等保持。
    original = defaultdict(dict)
    for row,col,value in m.entries:
        original[row][col] = value
    for y in inp['years']:
        cap = previous.capacity_terms(inp,meta,y)
        match = [row for row,terms in original.items() if terms==cap]
        assert len(match)==1
        m.constraint_lower[match[0]] = floor*inp['annual_observed_peak_mw'][str(y)]
    grouped = defaultdict(list)
    for y,h,failed,e,d,v in meta['network_flows']:
        grouped[y,h,failed].append((e,d,v))
    edge_map = {e['id']:e for e in inp['network']['edges']}
    for (y,h,failed),flows in grouped.items():
        existing = defaultdict(dict)
        for eid,d,v in flows:
            e = edge_map[eid]
            if e['existing']:
                a = e['a'] if d==0 else e['b']
                existing[inp['network']['feeders'][a]['station']][v] = 1
        for sid,terms in existing.items():
            m.constraint(terms,upper=inp['existing_transfer_budget_mw'][sid])
    return m,meta


def physical_check(inp,scenes,z,meta):
    """从流量及设备直接复算共享余量、站级N-1、负荷守恒。"""
    ids = list(inp['stations']); index = {s:j for j,s in enumerate(ids)}
    edge_map = {e['id']:e for e in inp['network']['edges']}
    grouped = defaultdict(list)
    for y,h,f,e,d,v in meta['network_flows']:
        grouped[y,h,f].append((e,d,float(z[v])))
    checks = []; dispatch = []
    for y in inp['years']:
        units,n,_,_ = assets_model.assets(inp,z,meta,y)
        bp,bn = network.storage_power(inp,n)
        built = {e for (e,yy),v in meta['edgevars'].items() if yy==y and z[v]>.5}
        worst = 0.
        for h,p in enumerate(scenes[y]):
            donor,receive = limits(inp,p,y)
            for failed in range(-1,len(ids)):
                outgoing=defaultdict(float); incoming=defaultdict(float); total=defaultdict(float)
                existing=defaultdict(float); station_out=np.zeros(len(ids)); netout=np.zeros(len(ids))
                residual = []
                for eid,d,value in grouped[y,h,failed]:
                    e = edge_map[eid]; a,b = (e['a'],e['b']) if d==0 else (e['b'],e['a'])
                    sa=inp['network']['feeders'][a]['station']; sb=inp['network']['feeders'][b]['station']
                    outgoing[a]+=value; incoming[b]+=value; total[eid]+=value
                    station_out[index[sa]]+=value; netout[index[sa]]+=value; netout[index[sb]]-=value
                    residual.append(-value)
                    if e['existing']:
                        existing[sa]+=value
                        residual.append(value-donor[a]*e['source_section_fraction'])
                residual.extend(total[e['id']]-(e['capacity_mw'] if e['existing'] or e['id'] in built else 0)
                                for e in inp['network']['edges'])
                residual.extend(outgoing[f]-donor[f] for f in donor)
                residual.extend(incoming[f]-receive[f] for f in receive)
                residual.extend(existing[s]-inp['existing_transfer_budget_mw'][s] for s in ids)
                residual.extend(station_out-np.maximum(p,0))
                available=PF*units.sum(axis=1)
                if failed>=0:available[failed]-=PF*units[failed].max()
                served=p-netout
                residual.extend(served-available-bp); residual.extend(-served-available-bn)
                residual.append(abs(served.sum()-p.sum()))
                violation=max(0.,float(max(residual))); worst=max(worst,violation)
                for j,sid in enumerate(ids):
                    dispatch.append(dict(year=y,scenario=h,failed_station=ids[failed] if failed>=0 else 'normal',
                        station=sid,original_net_mw=float(p[j]),net_transfer_out_mw=float(netout[j]),
                        post_transfer_mw=float(served[j]),forward_margin_mw=float(available[j]+bp[j]-served[j]),
                        reverse_margin_mw=float(available[j]+bn[j]+served[j])))
        for sid,s in inp['stations'].items():
            used=sum(e['id'] in built and e.get('new_bay_station')==sid for e in inp['network']['edges'])
            worst=max(worst,float(used-s['metadata']['spare_10kv_bays']))
        assert worst<1e-5, (y,worst)
        checks.append(dict(year=y,checked_pairs=len(scenes[y])*(len(ids)+1),max_violation_mw=worst,verified=True))
    return checks,dispatch


def run(region,case,seconds=120):
    root=OUT/region; out=root/case; out.mkdir(exist_ok=True)
    if (out/'summary.json').exists():raise FileExistsError(f'已有结果：{out}')
    inp=json.loads((root/'inputs.json').read_text()); scenes=previous.load_scenes(root,inp)
    floor=2. if case=='floor_2' else 0.
    m,meta=build(inp,scenes,floor)
    warm=previous.OUT/region/'elastic_storage'
    z,result=solve(m,out,'solve',seconds,.01,dict(inp,storage_usable_energy_fraction=1.),meta,warm)
    summary=dict(region=region,case=case,clr_floor=floor,solver=result,input_sha256=sha(root/'inputs.json'))
    if z is None or not result.get('feasible_incumbent'):
        dump(out/'summary.json',summary); return
    checks,dispatch=physical_check(inp,scenes,z,meta)
    summary,stations,projects=network.extract(inp,None,'elastic_storage',z,meta,result,[],checks)
    summary.update(region=region,case=case,clr_floor=floor,input_sha256=sha(root/'inputs.json'),
                   scope=inp['scope'],static_network_checks=summary.pop('full_hour_network_checks'))
    old=inp['initial_capacity_mva']
    for r in summary['annual']:
        r['net_capacity_increment_mva']=r['capacity_mva']-old
        assert r['capacity_mva']>=old-1e-5 and r['clr']>=floor-1e-7
        old=r['capacity_mva']
    np.savez_compressed(out/'chosen_solution.npz',solution=z)
    dump(out/'summary.json',summary);write_csv(out/'years.csv',summary['annual'])
    write_csv(out/'stations.csv',stations);write_csv(out/'projects.csv',projects)
    write_csv(out/'dispatch.csv',dispatch)
    print('COMPLETE',region,case,summary['cost_npv_10k'],flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--prepare',action='store_true')
    parser.add_argument('--region',choices=('city','pizhou'))
    parser.add_argument('--case',choices=CASES)
    parser.add_argument('--seconds',type=int,default=120)
    args=parser.parse_args()
    if args.prepare:prepare((args.region,) if args.region else ('city','pizhou'))
    if args.region and args.case:run(args.region,args.case,args.seconds)
