"""固定原方案计量既有转供，并独立比较逐站上限和正常固定供区约束。"""
import argparse
import copy
import json
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix

from . import capacity_run as base
from .capacity_audit import FlowLP
from .report import add_sheet, markdown
from .run import TRANSFER_FILE, PF
from . import report_version

ROOT=base.ROOT
OUT=ROOT/'研究报告/试算/20261012_既有转供约束诊断'
RESULT=ROOT/'研究报告/01研究推荐/当前版/既有转供约束诊断_20261012'
REVIEW=ROOT/'研究报告/05_review/当前版/1012既有转供约束诊断'
SOURCE=ROOT/'研究报告/数据来源/当前版/既有转供逐站约束_20261012'
MAPPING=ROOT/'实验/研究/data/tuomin/变电站对应.md'
LABEL={'rigid_storage':'刚性','elastic_no_floor':'弹性无下限','elastic_floor_2':'弹性有2.0下限'}
POLICIES=('station_caps','normal_fixed','station_caps_normal_fixed')
dump=base.dump;write_csv=base.write_csv;sha=base.sha


def station_evidence(inp):
    names={}
    for line in MAPPING.read_text().splitlines()[1:]:
        fields=line.split()
        if len(fields)>=3:names[fields[-1]]=fields[1]
    book=load_workbook(TRANSFER_FILE,read_only=True,data_only=True)
    lookup=defaultdict(list)
    for nr,r in enumerate(book.worksheets[0].iter_rows(min_row=3,values_only=True),3):
        if r[1]=='110kV' and r[2] and all(isinstance(r[k],(float,int)) for k in (9,10)):
            lookup[r[2]].append(dict(name=r[2],source_region=r[0],source_row=nr,
                                    transferable_mw=float(r[9]),unserved_mw=float(r[10])))
    book.close();records=[]
    for sid in inp['stations']:
        matched=lookup.get(names.get(sid),[]);assert len(matched)<=1
        item=dict(station=sid,name=names.get(sid,''),matched=bool(matched),
                  original_budget_mw=inp['existing_transfer_budget_mw'][sid])
        if matched:item.update(matched[0])
        else:item.update(source_region='',source_row='',transferable_mw=0.,unserved_mw=0.)
        total=item['transferable_mw']+item['unserved_mw']
        item['source_fraction']=item['transferable_mw']/total if total else None
        item['capped_budget_mw']=min(item['original_budget_mw'],item['transferable_mw'])
        records.append(item)
    return records


def constrained_input(inp,policy):
    p=copy.deepcopy(inp);p['transfer_constraint_policy']=policy
    p['normal_supply_area_fixed']='normal_fixed' in policy
    p['station_transfer_evidence']=station_evidence(inp)
    if 'station_caps' in policy:
        p['existing_transfer_budget_mw']={r['station']:r['capped_budget_mw'] for r in p['station_transfer_evidence']}
    p['constraint_diagnostic_note']='原91%/67%保留为收资地区背景；逐站MW上限直接约束既有通道，不为匹配地区均值重新放大网络。'
    return p


class DiagnosticLP(FlowLP):
    def bounds_for(self,prepared,failed):
        common,bounds,full,largest,bp,bn=prepared
        if failed<0 and self.inp.get('normal_supply_area_fixed'):
            bounds=[(0.,0.)]*len(bounds)
        return common,bounds,full,largest,bp,bn

    def check(self,p,prepared,failed,with_deficit=True):
        return super().check(p,self.bounds_for(prepared,failed),failed,with_deficit)

    def dispatch(self,p,prepared,failed):
        common,bounds,full,largest,bp,bn=self.bounds_for(prepared,failed)
        available=full.copy()
        if failed>=0:available[failed]-=largest[failed]
        if np.all(p<=available+bp+1e-7) and np.all(-p<=available+bn+1e-7):
            return dict(total_mw=0.,existing_mw=0.,aggregate_mw=0.,confirmed_mw=0.,new_mw=0.,
                        net_shift_mw=0.,flow=[],max_violation_mw=0.,direct=True)
        rhs=np.r_[common,available+bp-p,available+bn+p]
        r=linprog(np.ones(len(self.arcs)),A_ub=self.matrix,b_ub=rhs,bounds=bounds,method='highs')
        assert r.success,r.message
        violation=max(0.,float((self.matrix@r.x-rhs).max()),float(-r.x.min()),
                      max(float(v-hi) for v,(_,hi) in zip(r.x,bounds,strict=True)))
        assert violation<1e-5
        totals=defaultdict(float);net=np.zeros(len(self.ids));flows=[]
        for j,amount in enumerate(r.x):
            if amount<1e-7:continue
            k,e,d,a,b,sa,sb=self.arcs[j];amount=float(amount)
            category='aggregate' if e['id'].startswith('EQ-') else ('confirmed' if e['existing'] else 'new')
            totals[category]+=amount;net[sa]+=amount;net[sb]-=amount;flows.append([j,amount])
        assert abs(sum(totals.values())-r.fun)<1e-5
        return dict(total_mw=float(r.fun),existing_mw=totals['aggregate']+totals['confirmed'],
                    aggregate_mw=totals['aggregate'],confirmed_mw=totals['confirmed'],new_mw=totals['new'],
                    net_shift_mw=float(np.maximum(net,0).sum()),flow=flows,max_violation_mw=violation,direct=False)


def assets_from_csv(inp,folder,y):
    rows={r['station']:r for r in base.previous.rows(folder/'stations.csv') if int(r['year'])==y}
    units=np.array([[float(rows[s][f'unit_{k}_mva']) for k in (1,2,3)] for s in inp['stations']])
    n=np.array([int(rows[s]['storage_modules']) for s in inp['stations']])
    built={r['id'] for r in base.previous.rows(folder/'projects.csv') if int(r['year'])<=y}
    return units,n,built


def audit():
    OUT.mkdir(parents=True,exist_ok=True);REVIEW.mkdir(parents=True,exist_ok=True);SOURCE.mkdir(parents=True,exist_ok=True)
    yearly=[];scenario_rows=[];edge_rows=[];raw_rows=[];evidence=[];sources=[];count=0;worst=0.
    for region in ('city','pizhou'):
        root=base.OUT/region;inp=json.loads((root/'inputs.json').read_text());scenes=base.load_scenes(root,inp)
        lp=DiagnosticLP(inp);cache={};ids=list(inp['stations']);name=inp['region_name']
        evidence.extend(dict(地区=name,**r) for r in station_evidence(inp))
        sources.append(dict(region=region,input_sha256=sha(root/'inputs.json'),snapshot_sha256=sha(root/'static_base.npz'),
                            solutions={case:sha(root/case/'chosen_solution.npz') for case in base.CASES}))
        for case in base.CASES:
            folder=root/case;summary=json.loads((folder/'summary.json').read_text())
            selected={tuple(r) for r in json.loads((folder/'selected_pairs.json').read_text())}
            m,meta=base.build(inp,scenes,case,selected);z=np.load(folder/'chosen_solution.npz')['solution']
            original=defaultdict(lambda:defaultdict(float));edge_map={e['id']:e for e in inp['network']['edges']}
            for y,h,f,e,d,v in meta['network_flows']:
                original[y,h,f]['total']+=float(z[v])
                if edge_map[e]['existing']:original[y,h,f]['existing']+=float(z[v])
            for y in inp['years']:
                units,n,built=assets_from_csv(inp,folder,y);rows=[];dispatches={}
                for h,p in enumerate(scenes[y]):
                    prepared=lp.prepared(p,y,units,n,built)
                    for f in range(-1,len(ids)):
                        key=(y,h,f,units.tobytes(),n.tobytes(),tuple(sorted(built)))
                        if key not in cache:cache[key]=lp.dispatch(p,prepared,f)
                        r=cache[key];dispatches[h,f]=r;count+=1;worst=max(worst,r['max_violation_mw'])
                        row=dict(地区=name,方案=LABEL[case],年份=y,截面=h,原始小时=inp['station_snapshot_roles'][h]['source_hour'],
                                 工况='正常' if f<0 else 'N-1',故障站=ids[f] if f>=0 else '',
                                 既有通道转供_MW=r['existing_mw'],其中等效通道_MW=r['aggregate_mw'],
                                 其中已核实通道_MW=r['confirmed_mw'],新增通道转供_MW=r['new_mw'],
                                 总转供_MW=r['total_mw'],站净转移_MW=r['net_shift_mw'])
                        rows.append(row);scenario_rows.append(row)
                        if (y,h,f) in original:
                            raw_rows.append(dict(地区=name,方案=LABEL[case],年份=y,截面=h,故障站=row['故障站'],
                                原解既有转供_MW=original[y,h,f]['existing'],最小总转供调度中既有转供_MW=r['existing_mw'],
                                原解总转供_MW=original[y,h,f]['total'],最小总转供_MW=r['total_mw']))
                normal=max((r for r in rows if r['工况']=='正常'),key=lambda r:r['既有通道转供_MW'])
                fault=max((r for r in rows if r['工况']=='N-1'),key=lambda r:r['既有通道转供_MW'])
                yearly.append(dict(地区=name,方案=LABEL[case],年份=y,
                    正常最大既有转供_MW=normal['既有通道转供_MW'],正常对应新增转供_MW=normal['新增通道转供_MW'],正常截面=normal['截面'],
                    N1最大既有转供_MW=fault['既有通道转供_MW'],N1对应新增转供_MW=fault['新增通道转供_MW'],
                    N1截面=fault['截面'],N1故障站=fault['故障站'],
                    正常使用既有通道截面数=sum(r['工况']=='正常' and r['既有通道转供_MW']>1e-6 for r in rows),
                    N1使用既有通道工况数=sum(r['工况']=='N-1' and r['既有通道转供_MW']>1e-6 for r in rows),
                    检查工况数=len(rows),原方案费用现值_万元=summary['cost_npv_10k']))
                for row in (normal,fault):
                    h=row['截面'];f=-1 if row['工况']=='正常' else ids.index(row['故障站'])
                    for j,amount in dispatches[h,f]['flow']:
                        _,e,d,a,b,sa,sb=lp.arcs[j]
                        edge_rows.append(dict(地区=name,方案=LABEL[case],年份=y,工况=row['工况'],截面=h,故障站=row['故障站'],
                            通道=e['id'],类型='等效既有' if e['id'].startswith('EQ-') else ('已核实既有' if e['existing'] else '新增'),
                            转出站=ids[sa],接收站=ids[sb],转出馈线=a,接收馈线=b,转供_MW=amount,
                            通道容量_MW=e['capacity_mw'],送端既有预算_MW=inp['existing_transfer_budget_mw'][ids[sa]]))
            print('AUDITED_TRANSFER',region,case,flush=True)
    for name,rows in [('逐年既有转供量',yearly),('逐工况最小转供',scenario_rows),('代表工况通道明细',edge_rows),
                      ('原始解与最小转供对照',raw_rows),('逐站收资上限',evidence)]:
        write_csv(OUT/f'{name}.csv',rows)
    dump(REVIEW/'转供量独立复核.json',dict(verified=True,checked_pairs=count,max_violation_mw=worst,sources=sources,
        transfer_file_sha256=sha(TRANSFER_FILE),mapping_sha256=sha(MAPPING),
        dispatch_basis='固定投资设备后，以通道转供MW之和最小为目标；剔除原经济目标未约束的多余转供。各故障互斥，工况MW不累加为年电量。'))
    write_csv(SOURCE/'逐站既有转供能力及对应.csv',evidence)
    dump(SOURCE/'来源与口径.json',dict(转供原表=str(TRANSFER_FILE.relative_to(ROOT)),转供原表SHA256=sha(TRANSFER_FILE),
        站名对应=str(MAPPING.relative_to(ROOT)),站名对应SHA256=sha(MAPPING),输入=sources,
        约束='Σ既有通道转出_i≤min(旧标定预算_i,原表站外最大可供负荷J_i)；未匹配站不新增信用；固定既有能力，新增改造另计费用。',
        样本='市区26座实体站含铜山夹河条目及3模拟站；邳州20站。地区汇总率只保留为背景，不重新标定补足差额。'))
    print('TRANSFER_COMPLETE',count,flush=True)


def audit_capped_baseline():
    summaries=[];failures=[];count=0
    for region in ('city','pizhou'):
        source=base.OUT/region;original=json.loads((source/'inputs.json').read_text());inp=constrained_input(original,'station_caps')
        scenes=base.load_scenes(source,original);lp=DiagnosticLP(inp);cache={}
        for case in base.CASES:
            for y in inp['years']:
                units,n,built=assets_from_csv(inp,source/case,y);bad=[]
                for h,p in enumerate(scenes[y]):
                    prepared=lp.prepared(p,y,units,n,built)
                    for f in range(-1,len(inp['stations'])):
                        key=(y,h,f,units.tobytes(),n.tobytes(),tuple(sorted(built)))
                        if key not in cache:cache[key]=lp.check(p,prepared,f)
                        r=cache[key];count+=1
                        if not r['verified']:
                            row=dict(地区=inp['region_name'],原方案=LABEL[case],年份=y,截面=h,
                                故障站=lp.ids[f] if f>=0 else '正常',承载缺口_MW=r['deficit_mw'],
                                缺口站=json.dumps(r.get('station_deficits',{}),ensure_ascii=False))
                            bad.append(row);failures.append(row)
                summaries.append(dict(地区=inp['region_name'],原方案=LABEL[case],年份=y,
                    检查工况数=len(scenes[y])*(len(inp['stations'])+1),失败工况数=len(bad),
                    最大承载缺口_MW=max((r['承载缺口_MW'] for r in bad),default=0.)))
            print('AUDITED_STATION_CAPS',region,case,flush=True)
    write_csv(OUT/'原设备增加逐站上限校核.csv',summaries);write_csv(OUT/'原设备逐站上限缺口.csv',failures)
    dump(REVIEW/'逐站上限影响审查.json',dict(verified=True,checked_pairs=count,summary=summaries,
        meaning='保持原设备，只将既有转出限额绑定逐站原表J列；缺口为静态承载约束松弛，并非预计失负荷统计。'))


def build(inp,scenes,case,selected):
    m,meta=base.build(inp,scenes,case,selected)
    if inp.get('normal_supply_area_fixed'):
        groups=defaultdict(dict)
        for y,h,f,e,d,v in meta['network_flows']:
            if f<0:groups[y,h][v]=1.
        for terms in groups.values():m.constraint(terms,upper=0.)
    return m,meta


def repair_warm(region,policy,source_folder,target):
    """仅生成可行求解起点：用独立LP缺口补充静态储能，不作为推荐方案。"""
    source_folder=Path(source_folder);target=Path(target);target.mkdir(parents=True,exist_ok=True)
    original=json.loads((base.OUT/region/'inputs.json').read_text());inp=constrained_input(original,policy)
    scenes=base.load_scenes(base.OUT/region,original);lp=DiagnosticLP(inp)
    rows=base.previous.rows(source_folder/'stations.csv');lookup={(int(r['year']),r['station']):r for r in rows}
    extras=defaultdict(int);initial_failures=0
    for y in inp['years']:
        units,n,built=assets_from_csv(inp,source_folder,y)
        for p in scenes[y]:
            prepared=lp.prepared(p,y,units,n,built)
            for f in range(-1,len(inp['stations'])):
                r=lp.check(p,prepared,f)
                if not r['verified']:
                    initial_failures+=1
                    for sid,deficit in r['station_deficits'].items():
                        s=inp['stations'][sid]
                        coefficient=min(.1,.215*inp['storage_usable_energy_fraction']/max(1,s['forward_duration_h']),
                                        .215*inp['storage_usable_energy_fraction']/max(1,s['reverse_duration_h']))
                        extras[y,sid]=max(extras[y,sid],int(np.ceil((deficit+1e-6)/coefficient)))
    for sid in inp['stations']:
        previous_n=0
        for y in inp['years']:
            r=lookup[y,sid];n=max(previous_n,int(r['storage_modules'])+extras[y,sid]);previous_n=n
            r['storage_modules']=n;r['storage_mwh']=n*.215
    write_csv(target/'stations.csv',rows);(target/'projects.csv').write_bytes((source_folder/'projects.csv').read_bytes())
    checked=0
    for y in inp['years']:
        units,n,built=assets_from_csv(inp,target,y)
        for p in scenes[y]:
            prepared=lp.prepared(p,y,units,n,built)
            for f in range(-1,len(inp['stations'])):
                assert lp.check(p,prepared,f)['verified'];checked+=1
    dump(target/'warm_proof.json',dict(verified=True,checked_pairs=checked,initial_failures=initial_failures,
        source_folder=str(source_folder.relative_to(ROOT)),meaning='独立LP缺口修复的可行起点，后续经济寻优决定是否保留此储能。'))
    print('REPAIRED_WARM',region,checked,flush=True)
    return target


def solve_trial(region,case,policy,seconds,attempt=None,warm_folder=None,selected_folder=None):
    root=OUT/'约束对照'/region/policy;root.mkdir(parents=True,exist_ok=True)
    source=base.OUT/region;original=json.loads((source/'inputs.json').read_text());inp=constrained_input(original,policy)
    if (root/'inputs.json').exists():assert json.loads((root/'inputs.json').read_text())==inp
    else:dump(root/'inputs.json',inp)
    out=root/(case+('_'+attempt if attempt else ''));out.mkdir(exist_ok=True)
    if (out/'summary.json').exists():raise FileExistsError(out)
    scenes=base.load_scenes(source,original);lp=DiagnosticLP(inp);selected=set()
    for y in inp['years']:
        for h,r in enumerate(inp['station_snapshot_roles']):
            if '地区正向峰值' in r['roles'] or '地区最低净负荷' in r['roles']:
                selected.update((y,h,f) for f in range(-1,len(inp['stations'])))
        if inp.get('normal_supply_area_fixed'):
            selected.update((y,h,-1) for h in range(len(scenes[y])))
    if selected_folder:selected.update(tuple(r) for r in json.loads((Path(selected_folder)/'selected_pairs.json').read_text()))
    # 将原设备在逐站上限下已知的不足工况直接纳入首轮，减少无效候选。
    shortage_file=OUT/'原设备逐站上限缺口.csv'
    if shortage_file.exists():
        ids=list(inp['stations'])
        for row in base.previous.rows(shortage_file):
            if row['地区']==inp['region_name'] and row['原方案']==LABEL[case]:
                selected.add((int(row['年份']),int(row['截面']),ids.index(row['故障站']) if row['故障站']!='正常' else -1))
    warm=Path(warm_folder) if warm_folder else source/case;history=[];bounds=[]
    for iteration in range(1,11):
        m,meta=build(inp,scenes,case,selected)
        z,result=base.parameters.solve(m,out,f'iteration_{iteration:02}',seconds,.01,
            dict(inp,storage_usable_energy_fraction=1.),meta,warm)
        if result.get('dual_bound') is not None:bounds.append(result['dual_bound'])
        if z is None or not result.get('feasible_incumbent'):
            dump(out/'failed_search.json',dict(solver=result,history=history));return
        checks,violations,detail=base.check_all(inp,scenes,z,meta,lp)
        history.append(dict(iteration=iteration,selected_pairs=len(selected),solver=result,checks=checks))
        summary,stations,projects=base.network.extract(inp,None,case,z,meta,result,history,checks)
        candidate=out/f'candidate_{iteration:02}';candidate.mkdir(exist_ok=True)
        write_csv(candidate/'stations.csv',stations);write_csv(candidate/'projects.csv',projects);dump(candidate/'summary.json',summary);warm=candidate
        print('TRANSFER_TRIAL_CHECK',region,case,policy,iteration,'failures',len(violations),flush=True)
        if violations:
            added={(y,h,f) for _,y,h,f in sorted(violations,reverse=True)[:40]};assert added-selected
            warm=repair_warm(region,policy,candidate,candidate/'repaired_warm')
            selected.update(added);continue
        result['dual_bound']=max(bounds);result['mip_gap']=max(0.,(result['objective']-result['dual_bound'])/max(1.,abs(result['objective'])))
        summary.update(region=region,case=case,policy=policy,solver=result,clr_floor=base.floor_value(inp,case),
                       result_folder=str(out.relative_to(OUT)),
                       input_sha256=sha(root/'inputs.json'),source_input_sha256=sha(source/'inputs.json'),
                       static_network_checks=summary.pop('full_hour_network_checks'))
        old=inp['initial_capacity_mva']
        for a in summary['annual']:
            a['net_capacity_increment_mva']=a['capacity_mva']-old;old=a['capacity_mva']
        np.savez_compressed(out/'chosen_solution.npz',solution=z)
        dump(out/'summary.json',summary);write_csv(out/'years.csv',summary['annual'])
        write_csv(out/'stations.csv',stations);write_csv(out/'projects.csv',projects)
        dump(out/'selected_pairs.json',sorted(selected));dump(out/'independent_lp.json',dict(verified=True,checks=checks,records=detail))
        print('TRANSFER_TRIAL_COMPLETE',region,case,policy,summary['cost_npv_10k'],flush=True)
        return
    raise RuntimeError('工况分离未收敛，不得使用结果')


def verify_trials(check_export=True):
    records=[];best={};lower_bounds=defaultdict(list)
    for path in sorted((OUT/'约束对照').glob('*/*/*/summary.json')):
        s=json.loads(path.read_text());key=(s['region'],s['case'],s['policy'])
        assert s['input_sha256']==sha(path.parent.parent/'inputs.json')
        lower_bounds[key].append(float(s['solver']['dual_bound']))
        if key not in best or s['cost_npv_10k']<best[key][0]:best[key]=(s['cost_npv_10k'],path)
    expected={(region,case,policy) for region in ('city','pizhou') for case in base.CASES
              for policy in ('station_caps','station_caps_normal_fixed')}
    assert set(best)==expected,('缺少或多出对照组',expected-set(best),set(best)-expected)
    for _,path in best.values():
        out=path.parent;root=out.parent;inp=json.loads((root/'inputs.json').read_text());s=json.loads(path.read_text())
        source=base.OUT/s['region'];original=json.loads((source/'inputs.json').read_text())
        assert constrained_input(original,s['policy'])==inp
        assert s['input_sha256']==sha(root/'inputs.json') and s['source_input_sha256']==sha(source/'inputs.json')
        scenes=base.load_scenes(source,original)
        selected={tuple(r) for r in json.loads((out/'selected_pairs.json').read_text())}
        m,meta=build(inp,scenes,s['case'],selected);z=np.load(out/'chosen_solution.npz')['solution']
        rr,cc,vv=zip(*m.entries);matrix=coo_matrix((vv,(rr,cc)),shape=(len(m.constraint_lower),len(m.costs))).tocsr()
        az=matrix@z;integer=np.array(m.integrality,dtype=bool)
        violation=max(0.,float(np.max(np.array(m.constraint_lower)-az)),float(np.max(az-np.array(m.constraint_upper))),
                      float(np.max(np.array(m.lower_bounds)-z)),float(np.max(z-np.array(m.upper_bounds))),
                      float(np.max(abs(z[integer]-np.round(z[integer])))))
        assert violation<1e-5
        assert abs(np.dot(m.costs,z)-s['cost_npv_10k'])<1e-5
        checks,failures,details=base.check_all(inp,scenes,z,meta,DiagnosticLP(inp));assert not failures
        oldunits=np.array([r['initial_units']+[0.] for r in inp['stations'].values()]);oldn=np.zeros(len(oldunits));cost=0.
        oldcap=oldunits.sum();oldbuilt=set();edge_map={e['id']:e for e in inp['network']['edges']}
        for a in s['annual']:
            y=a['year'];units,n,built=assets_from_csv(inp,out,y)
            uz,nz,_,_=base.assets.assets(inp,z,meta,y)
            assert np.allclose(units,uz) and np.allclose(n,nz)
            assert np.all(units>=oldunits) and np.all(n>=oldn) and oldbuilt<=built
            assert abs(units.sum()-a['capacity_mva'])<1e-6 and abs(units.sum()/a['peak_mw']-a['clr'])<1e-9
            assert a['clr']>=base.floor_value(inp,s['case'])-1e-6
            if s['case']=='rigid_storage':assert units.sum()<=max(oldcap,2*a['peak_mw'])+1e-6
            assert all(not edge_map[e]['existing'] for e in built)
            bought=float(np.where(units[:,:2]>oldunits[:,:2],units[:,:2],0.).sum())
            third=float((units[:,2]-oldunits[:,2]).sum()/50.)
            raw=dict(transformer=bought*inp['transformer_price_10k_per_purchased_mva']+third*inp['third_50mva_capex_10k'],
                     storage=float((n-oldn).sum())*inp['storage_module_capex_10k'],line=len(built-oldbuilt)*inp['line_capex_10k'])
            for category,value in raw.items():
                npv=value*inp['cost_factors'][category][str(y)]
                assert abs(npv-a[category+'_npv_10k'])<1e-5;cost+=npv
            oldunits=units;oldn=n;oldcap=units.sum();oldbuilt=built
        assert abs(cost-s['cost_npv_10k'])<1e-5
        lower=max(lower_bounds[s['region'],s['case'],s['policy']]);assert lower<=cost+1e-5
        record=dict(region=s['region'],case=s['case'],policy=s['policy'],verified=True,
                    result_folder=str(out.relative_to(OUT)),
                    checked_pairs=sum(c['checked_pairs'] for c in checks),matrix_violation=violation,
                    cost_recomputed_10k=cost,input_sha256=sha(root/'inputs.json'),solution_sha256=sha(out/'chosen_solution.npz'),
                    dual_bound=lower,mip_gap=max(0.,(cost-lower)/max(1.,abs(cost))))
        records.append(record);print('VERIFIED_TRANSFER_TRIAL',s['policy'],s['case'],flush=True)
    assert records
    report_version.verify(records)
    dump(REVIEW/'约束对照独立核验.json',dict(verified=True,records=records))
    exported=REVIEW/'导出核验.json'
    if check_export and exported.exists():
        delivery=json.loads(exported.read_text())
        keys=('region','case','policy','input_sha256','solution_sha256')
        expected={tuple(r[k] for k in keys) for r in records}
        assert expected=={tuple(r[k] for k in keys) for r in delivery['trial_sources']},'对照解已变化，需重新--export'
        for filename,digest in delivery['result_files'].items():assert sha(RESULT/filename)==digest,filename
    return records


def trial_transfers(records):
    rows=[];edge_rows=[];scenario_rows=[]
    for c in records:
        out=OUT/c['result_folder'];inp=json.loads((out.parent/'inputs.json').read_text())
        scenes=base.load_scenes(base.OUT/c['region'],inp);lp=DiagnosticLP(inp)
        for y in inp['years']:
            units,n,built=assets_from_csv(inp,out,y);normal=[];faults=[];dispatches={}
            for h,p in enumerate(scenes[y]):
                prepared=lp.prepared(p,y,units,n,built)
                for f in range(-1,len(inp['stations'])):
                    r=lp.dispatch(p,prepared,f);dispatches[h,f]=r
                    (normal if f<0 else faults).append((r['existing_mw'],h,f,r))
                    scenario_rows.append(dict(地区=inp['region_name'],方案=LABEL[c['case']],约束=c['policy'],年份=y,
                        截面=h,原始小时=inp['station_snapshot_roles'][h]['source_hour'],
                        工况='正常' if f<0 else 'N-1',故障站=lp.ids[f] if f>=0 else '',
                        既有转供_MW=r['existing_mw'],其中等效既有_MW=r['aggregate_mw'],
                        其中已核实既有_MW=r['confirmed_mw'],新增转供_MW=r['new_mw'],
                        总转供_MW=r['total_mw'],站净转移_MW=r['net_shift_mw']))
            nr=max(normal,key=lambda v:v[0]);fr=max(faults,key=lambda v:v[0])
            rows.append(dict(地区=inp['region_name'],方案=LABEL[c['case']],约束=c['policy'],年份=y,
                正常最大既有转供_MW=nr[0],正常对应新增转供_MW=nr[3]['new_mw'],正常截面=nr[1],
                N1最大既有转供_MW=fr[0],N1对应新增转供_MW=fr[3]['new_mw'],N1截面=fr[1],N1故障站=lp.ids[fr[2]],
                正常最大总转供_MW=max(v[3]['total_mw'] for v in normal),
                正常最大新增转供_MW=max(v[3]['new_mw'] for v in normal),
                N1最大总转供_MW=max(v[3]['total_mw'] for v in faults),
                N1最大新增转供_MW=max(v[3]['new_mw'] for v in faults),
                检查工况数=len(normal)+len(faults)))
            for _,h,f,r in (nr,fr):
                for j,amount in r['flow']:
                    _,e,d,a,b,sa,sb=lp.arcs[j]
                    edge_rows.append(dict(地区=inp['region_name'],方案=LABEL[c['case']],约束=c['policy'],年份=y,
                        工况='正常' if f<0 else 'N-1',截面=h,故障站=lp.ids[f] if f>=0 else '',
                        通道=e['id'],类型='等效既有' if e['id'].startswith('EQ-') else ('已核实既有' if e['existing'] else '新增'),
                        转出站=lp.ids[sa],接收站=lp.ids[sb],转供_MW=amount,
                        通道容量_MW=e['capacity_mw'],送端既有上限_MW=inp['existing_transfer_budget_mw'][lp.ids[sa]]))
        print('AUDITED_TRIAL_TRANSFER',c['policy'],c['case'],flush=True)
    write_csv(OUT/'对照逐年转供量.csv',rows);write_csv(OUT/'对照代表工况通道.csv',edge_rows)
    write_csv(OUT/'对照逐工况最小转供.csv',scenario_rows)
    return rows,edge_rows


def export():
    audit_result=json.loads((REVIEW/'转供量独立复核.json').read_text());assert audit_result['verified']
    assert audit_result['transfer_file_sha256']==sha(TRANSFER_FILE) and audit_result['mapping_sha256']==sha(MAPPING)
    for source in audit_result['sources']:
        root=base.OUT/source['region']
        assert source['input_sha256']==sha(root/'inputs.json') and source['snapshot_sha256']==sha(root/'static_base.npz')
        assert source['solutions']=={case:sha(root/case/'chosen_solution.npz') for case in base.CASES}
    trial_checks=json.loads((REVIEW/'约束对照独立核验.json').read_text());assert trial_checks['verified']
    report_version.verify(trial_checks['records'])
    trial_amounts,trial_edges=trial_transfers(trial_checks['records'])
    RESULT.mkdir(parents=True,exist_ok=True)
    annual=base.previous.rows(OUT/'逐年既有转供量.csv');edges=base.previous.rows(OUT/'代表工况通道明细.csv')
    evidence=base.previous.rows(OUT/'逐站收资上限.csv');raw=base.previous.rows(OUT/'原始解与最小转供对照.csv')
    # 从CSV回读时保留原始数值类型，避免工作簿将数值当文本。
    def typed(rows):
        result=[]
        for row in rows:
            item={}
            for key,value in row.items():
                if value in ('True','False'):item[key]=value=='True'
                elif value=='':item[key]=''
                else:
                    try:item[key]=float(value) if any(c in value for c in '.eE') else int(value)
                    except ValueError:item[key]=value
            result.append(item)
        return result
    annual,edges,evidence,raw=map(typed,(annual,edges,evidence,raw))
    costs=[];years=[];stations=[];projects=[]
    station_names={(r['地区'],r['station']):r['name'] for r in evidence}
    for region in ('city',):
        for case in ('rigid_storage','elastic_no_floor'):
            folder=base.OUT/region/case;s=json.loads((folder/'summary.json').read_text())
            costs.append(dict(地区='市区',方案=LABEL[case],约束='上版地区标定（逐站校核未通过）',费用现值_万元=s['cost_npv_10k'],
                              费用下界_万元=s['solver']['dual_bound'],相对间隙=s['solver']['mip_gap'],
                              首年容量_MVA=s['annual'][0]['capacity_mva'],首年储能_MWh=s['annual'][0]['storage_mwh'],
                              末年容量_MVA=s['annual'][-1]['capacity_mva'],末年储能_MWh=s['annual'][-1]['storage_mwh'],
                              首年容载比=s['annual'][0]['clr'],末年容载比=s['annual'][-1]['clr'],
                              累计新增联络_项=s['annual'][-1]['cumulative_lines']))
    policy_names={'station_caps':'逐站收资MW上限','normal_fixed':'正常供区固定','station_caps_normal_fixed':'逐站MW上限＋正常供区固定'}
    for c in trial_checks['records']:
        folder=OUT/c['result_folder'];root=folder.parent
        assert c['input_sha256']==sha(root/'inputs.json') and c['solution_sha256']==sha(folder/'chosen_solution.npz')
        s=json.loads((folder/'summary.json').read_text());assert abs(c['cost_recomputed_10k']-s['cost_npv_10k'])<1e-5
        costs.append(dict(地区='市区' if c['region']=='city' else '邳州',方案=LABEL[c['case']],约束=policy_names[c['policy']],
            费用现值_万元=s['cost_npv_10k'],费用下界_万元=c['dual_bound'],相对间隙=c['mip_gap'],
            首年容量_MVA=s['annual'][0]['capacity_mva'],首年储能_MWh=s['annual'][0]['storage_mwh'],
            末年容量_MVA=s['annual'][-1]['capacity_mva'],末年储能_MWh=s['annual'][-1]['storage_mwh'],
            首年容载比=s['annual'][0]['clr'],末年容载比=s['annual'][-1]['clr'],
            累计新增联络_项=s['annual'][-1]['cumulative_lines']))
        for a in s['annual']:
            years.append(dict(地区=costs[-1]['地区'],方案=LABEL[c['case']],约束=policy_names[c['policy']],年份=a['year'],
                年度负荷_MW=a['peak_mw'],年末容量_MVA=a['capacity_mva'],容载比=a['clr'],当年净增容量_MVA=a['net_capacity_increment_mva'],
                当年购入主变容量_MVA=a['purchased_transformer_mva'],当年新增第三台主变_台=a['new_third_count'],
                当年新增储能_MWh=a['new_storage_mwh'],年末储能_MWh=a['storage_mwh'],
                当年新增联络_项=a['new_lines'],累计新增联络_项=a['cumulative_lines'],
                当年建设投资_万元=a['construction_capex_10k'],当年主变费用现值_万元=a['transformer_npv_10k'],
                当年联络费用现值_万元=a['line_npv_10k'],当年储能费用现值_万元=a['storage_npv_10k'],
                当年费用现值_万元=a['added_lifecycle_npv_10k'],累计费用现值_万元=a['cumulative_lifecycle_npv_10k']))
        stations.extend(dict(地区=costs[-1]['地区'],方案=LABEL[c['case']],约束=policy_names[c['policy']],
                             站名=station_names.get((costs[-1]['地区'],r['station']),r['station']),**r)
                        for r in typed(base.previous.rows(folder/'stations.csv')))
        projects.extend(dict(地区=costs[-1]['地区'],方案=LABEL[c['case']],约束=policy_names[c['policy']],
                             端点站名A=station_names.get((costs[-1]['地区'],r['station_a']),r['station_a']),
                             端点站名B=station_names.get((costs[-1]['地区'],r['station_b']),r['station_b']),**r)
                        for r in typed(base.previous.rows(folder/'projects.csv')))
    purpose=report_version.purpose
    for rows in (costs,years,stations,projects):
        for r in rows:
            r['用途']='历史定位' if '上版地区标定' in r['约束'] else (
                '报告主方案' if r['约束']==policy_names['station_caps'] and r['方案'] in ('刚性','弹性无下限') else '保留对照')
    trial_amounts=[dict(r,用途=purpose(r['约束'],next(k for k,v in LABEL.items() if v==r['方案'])),约束=policy_names[r['约束']]) for r in trial_amounts]
    trial_edges=[dict(r,用途=purpose(r['约束'],next(k for k,v in LABEL.items() if v==r['方案'])),约束=policy_names[r['约束']]) for r in trial_edges]
    main=[]
    for c in report_version.selected(trial_checks['records']):
        main.append(dict(版本=report_version.VERSION,用途='报告主方案',地区='市区' if c['region']=='city' else '邳州',
            运行规则='正常允许调整供区',方案=LABEL[c['case']],费用现值_万元=c['cost_recomputed_10k'],
            费用下界_万元=c['dual_bound'],相对间隙=c['mip_gap'],原始结果目录=str((OUT/c['result_folder']).relative_to(ROOT)),
            模型入口='实验/研究/rebuild_2026/case_parameters_20261012/transfer_diagnostic.py',
            输入SHA256=c['input_sha256'],方案解SHA256=c['solution_sha256']))
    suggestions=[
        dict(优先级=1,约束='既有转供逐站MW上限',表达='Σ既有转出_i≤min(上版标定预算_i,J_i)',
             依据='原表逐站站外最大可供负荷J列，匹配110kV站名和站号；不为凑地区均值重新放大。',状态='已用于两地全部12组；市区26站含铜山夹河，邳州20站可匹配，3模拟站无既有通道。'),
        dict(优先级=2,约束='正常供区固定与故障转供分开',表达='固定供区对照：正常工况f=0；故障工况允许既有及已建通道转供',
             依据='2025年4月修订资料8.1.2相对独立供区；9.2.1.2转供取决于主变、线路裕度及分段联络。固定供区是本轮对照假设，条文未要求所有正常调整一律为零。',状态='按用户确认保留两地全部三类方案的两种运行规则对照，共12组。'),
        dict(优先级=3,约束='正常接线跨工况一致',表达='每年选定正常供区/开关状态，正常样本共用；故障后在允许切换范围内重构',
             依据='当前各静态工况独立选转供，可能把应急能力当成可随时调整的正常供区。需实际可切负荷块与常开点。',状态='建议下一步替代简单正常零转供假设；数据补齐后实施。'),
        dict(优先级=4,约束='既有和新增共用实际馈线余量',表达='同一馈线、负荷块的既有和新增转供合计受唯一热限与可切负荷约束',
             依据='当前站级接收容量、既有预算、馈线共享限制已有；等效馈线组与实际新建端点之间仍缺少逐线对应。',状态='补实际既有端点和分段后实施，不能把每个规划候选站对直接认定为存量线路。'),
        dict(优先级=5,约束='主变正常运行储备与储能站址上限',表达='按批准标准设置正常负载率上限，储能MW/MWh按场地和接入容量限制',
             依据='阈值需要地区规划标准和站址资料；当前储能上界是搜索范围，未取得场地限额。',状态='可作参数敏感性；不凭期望费用自行指定80%或90%。'),
        dict(优先级=6,约束='线路故障或关键站全停压力测试',表达='按适用供电等级和保障要求另设故障集',
             依据='当前N-1为单台最大主变故障；全停全转收资是另一停运方式。2025修订资料9.2.1.2仅在特定保障需求下考虑全停全转。',状态='需明确对象及实际线路对应，不作为全地区默认强制约束。')]
    notes=[dict(事项='报告采用版本',说明='用户已选定正常允许调整供区的弹性无下限、刚性，两地共4组作为后续报告数据和模型版本；其余8组保留对照。'),
        dict(事项='零费用含义',说明='上版0是共同压缩起点后新增设备投资及其生命周期费用为0；已建网络资产成本未重复计入。'),
        dict(事项='转供量口径',说明='固定原投资设备，最小化各通道转供MW之和；表中列代表可行调度，故障工况互斥，不累加为全年电量。'),
        dict(事项='年度方案口径',说明='2022—2025方案为共同压缩起点的规划回算，既有网络按2025收资能力固定；转供量为模型调度，不是实际逐年投运及运行记录。'),
        dict(事项='既有通道分类',说明='等效既有通道与已核实T01分别列示；既有等效站对为研究假设，地区全停全转率不唯一确定真实线路连接。'),
        dict(事项='费用目标与转供输出',说明='投资模型费用目标不惩罚转供，原解可存在多余转供；转供量另以固定设备的最小总转供LP重算，投资费用不改。'),
        dict(事项='调度非唯一',说明='最小总转供量确定时，既有与新增通道分配仍可能多解；明细是其中一个最小总量方案，不声称为实测调度。'),
        dict(事项='对照范围',说明='两地六组原方案均补出转供量；两地均按逐站上限，比较正常允许调整与正常供区固定两种规则，各含刚性、无下限弹性、有2.0下限弹性，共12组。'),
        dict(事项='资料版本',说明='约束依据使用2025年4月《配电网规划设计技术导则》修订资料，保留其真实版本；DL/T 5729—2023为另一已发布标准。')]
    frames={'报告采用版本':main,
            '报告主方案逐年措施':[r for r in years if r['用途']=='报告主方案'],
            '报告主方案转供量':[r for r in trial_amounts if r['用途']=='报告主方案'],
            '市区刚性原方案转供':[r for r in annual if r['地区']=='市区' and r['方案']=='刚性'],
            '两地六组转供':annual,'代表工况通道':edges,'逐站转供收资':evidence,'原解与最小调度':raw,
            '原方案加逐站上限校核':typed(base.previous.rows(OUT/'原设备增加逐站上限校核.csv')),
            '对照逐年转供':trial_amounts,'对照代表通道':trial_edges,
            '约束方案比选':costs,'对照逐年措施':years,'对照逐站设备':stations,'对照新增联络项目':projects,
            '约束建议':suggestions,'使用说明':notes}
    book=Workbook();book.remove(book.active)
    for name,rows in frames.items():
        write_csv(RESULT/f'{name}.csv',rows);add_sheet(book,name,rows)
        headers=list(rows[0])
        for j,key in enumerate(headers,1):
            if key=='相对间隙' or key.endswith('fraction'):
                for col in book[name].iter_cols(min_col=j,max_col=j,min_row=2):
                    for cell in col:cell.number_format='0.00%' if key=='相对间隙' else '0%'
    path=RESULT/'既有转供量与约束调整对照.xlsx';book.save(path);loaded=load_workbook(path,data_only=True);cells=0
    for name,rows in frames.items():
        actual=list(loaded[name].iter_rows(values_only=True));expected=[list(rows[0])]+[[r.get(k) for k in rows[0]] for r in rows]
        assert len(actual)==len(expected)
        for a,b in zip(actual,expected,strict=True):
            for x,y in zip(a,b,strict=True):
                if isinstance(y,float):assert abs(x-y)<max(1e-8,abs(y)*1e-12)
                else:assert (x if x is not None else '')==(y if y is not None else '')
                cells+=1
    loaded.close()
    city=frames['市区刚性原方案转供'];comparison=frames['约束方案比选']
    text=['# 既有转供量及约束调整讨论','',
        '报告采用版本：正常允许调整供区的弹性无下限和刚性，覆盖市区、邳州共4组；其余8组保留对照。版本号为`'+report_version.VERSION+'`，原始结果、输入和模型依赖的哈希见报告采用版本锁定清单。','',
        '上版市区零新增投资依赖正常和故障两类工况均可调用既有等效通道。本轮先固定设备计算必要转供量，再用逐站原始MW上限和正常固定供区分别比较。','',
        '2022—2025设备配置为共同压缩起点的规划回算，既有网络按2025收资能力固定；以下转供量为模型调度。两种规则的地区电缆长度占比及联络独立计价相同：市区80%、452.40万元/项；邳州11%、128.38万元/项，完整长度比例保留在输入。','',
        markdown(['年份','正常最大既有转供/MW','N−1最大既有转供/MW','正常使用转供的样本数'],
            [[r['年份'],f"{r['正常最大既有转供_MW']:.2f}",f"{r['N1最大既有转供_MW']:.2f}",r['正常使用既有通道截面数']] for r in city]),'',
        '各故障工况互斥；以上是固定原设备后最小总转供调度的工况最大值，不能相加为全年转供电量。市区既有转供全部来自等效通道，尚无已核实逐线既有边。','',
        markdown(['地区','方案','约束','费用现值/万元','首年容量/MVA','首年储能/MWh','累计新增联络/项','间隙'],
            [[r['地区'],r['方案'],r['约束'],f"{r['费用现值_万元']:.2f}",f"{r['首年容量_MVA']:.2f}",f"{r['首年储能_MWh']:.2f}",r['累计新增联络_项'],f"{100*r['相对间隙']:.2f}%"] for r in comparison]),'',
        '上版地区标定的两行零费用仅用于问题定位，其设备在逐站上限下有承载缺口。其余12组为已取得并通过本版承载约束的具体可行方案费用。有正间隙的组保留费用下界，费用比较不将其视为已证全局最优。逐年新增、年末累计、容载比及逐站设备见工作簿；新增联络另列端点和实施年份。','',
        '有2.0下限组逐年满足容载比≥2.0；无下限组由站级承载约束、设备离散规格及费用共同决定，允许部分年份低于2.0。刚性组继续采用新增容量的2.0控制规则，负荷回落造成的存量被动超限保留。容载比为地区总容量与原年度统计峰值之比，逐站正常及N−1校核同时执行。','',
        '建议首先把逐站收资MW上限加入既有转供约束。正常供区固定可作为收紧运行假设的对照；后续有常开点和可切负荷块时，应进一步建立跨正常工况一致的接线和供区决策。将所有正常转供设为零是本轮对照假设，修订资料的相对独立供区要求本身并未规定全部正常调整都为零。','',
        '现有站级接收容量、馈线共享余量和备用间隔约束继续保留。进一步完善的重点是等效既有通道与实际馈线的对应，防止同一余量在既有和新增通道中重复授予。主变储备、储能站址以及关键故障集可继续补充，但需要相应规划标准或工程数据确定参数。','',
        '逐站匹配覆盖市区26座实体站、邳州20站；市区样本中的夹河站使用原表铜山条目，3座模拟站保持无既有能力。地区91%/67%作为收资背景保留，逐站约束后不为匹配地区均值重新扩大通道。','',
        '依据：2025年4月《配电网规划设计技术导则》修订资料8.1.2、9.2.1.2；源文件在参考文献/前沿调研/pdfs/D8_政策标准。原始转供表及站名对应的路径和哈希见核验。','',
        '两地六组原方案22,752个工况的最小转供复算通过；两地12组新增约束对照另经矩阵、费用及45,504个静态工况独立核验。正式报告、锁定推广矩阵和上版六组原始结果保留。']
    (RESULT/'约束调整与结果说明.md').write_text('\n'.join(text)+'\n',encoding='utf-8')
    (RESULT/'报告采用版本说明.md').write_text('# 报告采用的数据和模型版本\n\n版本号：`'+report_version.VERSION+'`。\n\n用户确认采用正常允许调整供区的弹性无下限与刚性，市区、邳州共4组。工作簿首张“报告采用版本”登记费用、求解间隙、原始目录及输入和解哈希；“报告主方案逐年措施”“报告主方案转供量”给出可直接用于撰写的16条年度记录。其余8组继续保留并标记为“保留对照”。\n\n'+markdown(['地区','方案','费用现值/万元','间隙'],[[r['地区'],r['方案'],f"{r['费用现值_万元']:.2f}",f"{r['相对间隙']*100:.2f}%"] for r in main])+'\n\n模型入口为`实验/研究/rebuild_2026/case_parameters_20261012/transfer_diagnostic.py`；锁定清单为`研究报告/05_review/当前版/1012既有转供约束诊断/报告采用版本锁定清单.json`。正文主方案数据从本版引用，2.0下限和正常供区固定组用于对照讨论。正式六章、Word/PDF尚未切换。\n',encoding='utf-8')
    dump(REVIEW/'导出核验.json',dict(verified=True,workbook=str(path.relative_to(ROOT)),sha256=sha(path),checked_cells=cells,
        original_checked_pairs=audit_result['checked_pairs'],trial_checked_pairs=sum(r['checked_pairs'] for r in trial_checks['records']),
        trial_sources=[{k:r[k] for k in ('region','case','policy','input_sha256','solution_sha256')} for r in trial_checks['records']],
        result_files={p.name:sha(p) for p in sorted(RESULT.iterdir()) if p.suffix in ('.csv','.xlsx') or p.name in ('约束调整与结果说明.md','报告采用版本说明.md')}))
    print('EXPORTED_TRANSFER',cells,flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--audit',action='store_true');parser.add_argument('--region',choices=('city','pizhou'))
    parser.add_argument('--verify',action='store_true');parser.add_argument('--export',action='store_true')
    parser.add_argument('--audit-caps',action='store_true')
    parser.add_argument('--lock-report-version',action='store_true')
    parser.add_argument('--case',choices=base.CASES);parser.add_argument('--policy',choices=POLICIES)
    parser.add_argument('--seconds',type=int,default=120)
    parser.add_argument('--attempt');parser.add_argument('--warm');parser.add_argument('--selected-from')
    args=parser.parse_args()
    if args.audit:audit()
    if args.audit_caps:audit_capped_baseline()
    if args.region and args.case and args.policy:solve_trial(args.region,args.case,args.policy,args.seconds,args.attempt,args.warm,args.selected_from)
    if args.verify:verify_trials()
    if args.lock_report_version:report_version.freeze()
    if args.export:
        verify_trials(check_export=False)
        export()
