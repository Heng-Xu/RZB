"""从四组原始求解记录导出本轮参数修正结果，不改正式稿件。"""
import json

from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter

from .run import ROOT,OUT,SOURCE,RESULT,REVIEW,CASES,previous,write_csv,dump,sha

LABEL={'floor_2':'有2.0下限','no_floor':'无容载比下限'}


def markdown(headers,rows):
    return '\n'.join(['|'+'|'.join(headers)+'|','|'+'|'.join(['---']*len(headers))+'|']+
                     ['|'+'|'.join(map(str,r))+'|' for r in rows])


def add_sheet(book,name,rows):
    sheet=book.create_sheet(name)
    if not rows:return
    headers=list(rows[0]);sheet.append(headers)
    for row in rows:sheet.append([row.get(k) for k in headers])
    for cell in sheet[1]:
        cell.font=Font(name='方正仿宋_GBK',bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='27546B')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name='方正仿宋_GBK',size=11)
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            if isinstance(cell.value,float):cell.number_format='0.00'
    for j,key in enumerate(headers,1):
        sheet.column_dimensions[get_column_letter(j)].width=min(42,max(15,len(key)*1.7))
    sheet.freeze_panes='D2';sheet.auto_filter.ref=sheet.dimensions


def report():
    assert json.loads((REVIEW/'独立核验.json').read_text())['verified']
    RESULT.mkdir(parents=True,exist_ok=True)
    parameters=[];annual=[];station_changes=[];projects=[];comparison=[];solver_rows=[]
    data={}
    for region in ('city','pizhou'):
        root=OUT/region;inp=json.loads((root/'inputs.json').read_text());name=inp['region_name']
        parameters.append(dict(地区=name,初始转供率=inp['initial_transfer_fraction'],
            模型标定转供率=inp['initial_transfer_audit']['weighted_rate'],电缆长度占比=inp['line_cable_fraction'],
            新增10kV联络单项造价_万元=inp['line_capex_10k'],单项研究长度_km=3.,
            共同起点容量_MVA=inp['initial_capacity_mva'],正向储能持续时间_h=inp['forward_duration_h'],
            转供统计范围='全部电压等级',既有网络='原候选站对内聚合通道标定，存量不计新增投资'))
        data[region]={}
        for case in CASES:
            out=root/case;s=json.loads((out/'summary.json').read_text());data[region][case]=s
            for a in s['annual']:
                annual.append(dict(地区=name,方案=LABEL[case],年份=a['year'],年度负荷_MW=a['peak_mw'],
                    年末容量_MVA=a['capacity_mva'],容载比=a['clr'],当年净增容量_MVA=a['net_capacity_increment_mva'],
                    当年购入主变容量_MVA=a['purchased_transformer_mva'],当年新增第三台主变_台=a['new_third_count'],
                    当年新增联络_项=a['new_lines'],年末累计新增联络_项=a['cumulative_lines'],
                    当年新增储能_MWh=a['new_storage_mwh'],年末储能_MW=a['storage_nameplate_mw'],年末储能_MWh=a['storage_mwh'],
                    当年建设投资_万元=a['construction_capex_10k'],当年主变费用现值_万元=a['transformer_npv_10k'],
                    当年联络费用现值_万元=a['line_npv_10k'],当年储能费用现值_万元=a['storage_npv_10k'],
                    当年费用现值_万元=a['added_lifecycle_npv_10k'],累计费用现值_万元=a['cumulative_lifecycle_npv_10k']))
            old={sid:asset['initial_units']+[0.] for sid,asset in inp['stations'].items()}
            oldn={sid:0 for sid in inp['stations']}
            for row in previous.rows(out/'stations.csv'):
                sid=row['station'];units=[float(row[f'unit_{k}_mva']) for k in (1,2,3)];n=int(row['storage_modules'])
                station_changes.append(dict(地区=name,方案=LABEL[case],年份=int(row['year']),站点=sid,
                    调整前主变_MVA='/'.join(f'{v:g}' for v in old[sid] if v),
                    年末主变_MVA='/'.join(f'{v:g}' for v in units if v),年末容量_MVA=sum(units),
                    当年净增容量_MVA=sum(units)-sum(old[sid]),当年购入主变容量_MVA=float(row['purchased_transformer_mva']),
                    当年新增储能_MWh=(n-oldn[sid])*.215,年末储能_MW=n*.1,年末储能_MWh=n*.215))
                old[sid]=units;oldn[sid]=n
            for row in previous.rows(out/'projects.csv'):
                projects.append(dict(地区=name,方案=LABEL[case],投运年份=int(row['year']),项目=row['id'],
                    送端站=row['station_a'],受端站=row['station_b'],馈线A=row['a'],馈线B=row['b'],
                    单项能力_MW=float(row['capacity_mw']),新出线占用间隔站=row['new_bay_station'],
                    单项建设投资_万元=inp['line_capex_10k'],候选依据=row['basis']))
            sol=s['solver']
            solver_rows.append(dict(地区=name,方案=LABEL[case],可行方案费用现值_万元=s['cost_npv_10k'],
                费用下界_万元=sol['dual_bound'],相对间隙=sol['mip_gap'],求解状态=sol['status'],
                求解秒数=sol['elapsed_seconds'],独立LP工况数=sum(c['checked_pairs'] for c in s['static_network_checks']),
                输入SHA256=s['input_sha256']))
        a,b=(data[region][case] for case in CASES)
        totals=lambda s,k:sum(r[k] for r in s['annual'])
        comparison.append(dict(地区=name,有下限费用现值_万元=a['cost_npv_10k'],无下限费用现值_万元=b['cost_npv_10k'],
            费用减少_万元=a['cost_npv_10k']-b['cost_npv_10k'],费用降幅=(a['cost_npv_10k']-b['cost_npv_10k'])/a['cost_npv_10k'],
            有下限2025容量_MVA=a['annual'][-1]['capacity_mva'],无下限2025容量_MVA=b['annual'][-1]['capacity_mva'],
            有下限2025容载比=a['annual'][-1]['clr'],无下限2025容载比=b['annual'][-1]['clr'],
            主变费用减少_万元=totals(a,'transformer_npv_10k')-totals(b,'transformer_npv_10k'),
            联络费用减少_万元=totals(a,'line_npv_10k')-totals(b,'line_npv_10k'),
            储能费用减少_万元=totals(a,'storage_npv_10k')-totals(b,'storage_npv_10k')))
    notes=[dict(事项='对照边界',说明='两组均允许容载比超过2.0；有下限组要求各年C/P≥2.0，无下限组取消这一行下界。不是刚性2.0上限对比。'),
           dict(事项='既有转供能力',说明='新转供率映射为共同起点已存在的聚合通道能力，存量不计新增投资；新增联络独立计费。'),
           dict(事项='网络假设',说明='既有通道限定在原候选站对内，用共同起点设备正常余量标定各站分别全停的负荷加权能力。完整既有接线未收齐，通道为研究等值。'),
           dict(事项='同期负荷',说明='转供标定采用原静态工况按2021地区净峰值缩放的同期分配，以正向站级负荷加权；净反送站的可转出负荷取0。优化的两类工况和年度负荷均保留原值。'),
           dict(事项='时间口径',说明='2025年转供统计用于共同仿真起点的情景参数；2022—2025年使用实际年度负荷，方案为回算仿真结果。'),
           dict(事项='转供统计',说明='采用全部电压等级地区收资：ΣJ/(ΣJ+ΣK)。仅110kV统计另列来源表；使用完整精度，非逐站百分比算术平均。'),
           dict(事项='保留参数',说明='设备候选、正常及主变N-1、原馈线及新增候选、备用间隔、储能能量系数和持续时间、寿命、折现率及评价期保持原案例。'),
           dict(事项='联络造价',说明='3km×[24×(1-q)+180×q]+6万元/项。长度3km及端部6万元保持原研究假设；既有馈线互联与邻站新出线按原间隔条件区分。'),
           dict(事项='容量统计',说明='当年净增容量为投运后容量减上年末容量；购入主变容量包含更换设备的整台额定容量，两者分别列示。'),
           dict(事项='求解记录',说明='保存费用下界、间隙和可行解；非零间隙方案用于可行费用比较。独立LP覆盖两类静态工况及各站最大单台主变N-1。')]
    frames={'参数修正':parameters,'方案费用对照':comparison,'逐年方案':annual,'逐站设备':station_changes,
            '新增联络项目':projects,'求解与核验':solver_rows,'使用说明':notes}
    book=Workbook();book.remove(book.active)
    for name,rows in frames.items():
        write_csv(RESULT/f'{name}.csv',rows);add_sheet(book,name,rows)
    for sheet in [book['参数修正']]:
        for j in (2,3,4):
            for row in range(2,sheet.max_row+1):sheet.cell(row,j).number_format='0.00000000'
    path=RESULT/'两地区参数修正_2.0下限有无对照.xlsx';book.save(path)
    checked=load_workbook(path,data_only=True)
    cells=0
    for name,rows in frames.items():
        expected=[list(rows[0])]+[[r.get(k) for k in rows[0]] for r in rows]
        actual=list(checked[name].iter_rows(values_only=True))
        assert len(actual)==len(expected)
        for a,b in zip(actual,expected,strict=True):
            for x,y in zip(a,b,strict=True):
                if isinstance(y,float):assert abs(x-y)<max(1e-8,abs(y)*1e-12)
                else:assert (x if x is not None else '')==(y if y is not None else '')
                cells+=1
    checked.close()
    text=['# 两地区参数修正及2.0下限对照（首版）','',
          '本版把地区转供率作为既有网络的基准能力，新增10 kV联络另行计费。两组采用相同设备、负荷、网络和费用参数，仅改变容载比2.0下限。两组均保留物理设备上界并允许容载比超过2.0。','',
          markdown(['地区','既有转供率（原表加权）','电缆长度占比','新增10 kV联络单项造价/万元'],
                   [[r['地区'],f"{r['初始转供率']*100:.2f}%",f"{r['电缆长度占比']*100:.2f}%",f"{r['新增10kV联络单项造价_万元']:.2f}"] for r in parameters]),'',
          '## 费用和容量对照','',
          markdown(['地区','有下限费用/万元','无下限费用/万元','费用减少/万元','降幅','有/无下限2025年容载比'],
                   [[r['地区'],f"{r['有下限费用现值_万元']:.2f}",f"{r['无下限费用现值_万元']:.2f}",f"{r['费用减少_万元']:.2f}",f"{r['费用降幅']:.0%}",f"{r['有下限2025容载比']:.2f} / {r['无下限2025容载比']:.2f}"] for r in comparison]),'',
          '费用为原评价口径下新增措施的生命周期现值。分项差额如下，正值表示取消下限后费用减少。','',
          markdown(['地区','主变费用减少/万元','联络费用减少/万元','储能费用减少/万元'],
                   [[r['地区']]+[f"{r[k]:.2f}" for k in ('主变费用减少_万元','联络费用减少_万元','储能费用减少_万元')] for r in comparison]),'',
          '## 逐年措施','',
          markdown(['地区','方案','年份','年末容量/MVA','容载比','当年净增容量/MVA','当年新增联络/项','年末累计联络/项','当年新增储能/MWh'],
                   [[r['地区'],r['方案'],r['年份'],f"{r['年末容量_MVA']:.2f}",f"{r['容载比']:.2f}",f"{r['当年净增容量_MVA']:.2f}",r['当年新增联络_项'],r['年末累计新增联络_项'],f"{r['当年新增储能_MWh']:.2f}"] for r in annual]),'',
          '## 结果解释','']
    for region in ('city','pizhou'):
        a,b=(data[region][case] for case in CASES);name=a['annual'][0] and ('市区' if region=='city' else '邳州')
        for case,s in ((CASES[0],a),(CASES[1],b)):
            actions=[]
            for r in s['annual']:
                bits=[]
                if r['net_capacity_increment_mva']:bits.append(f"净增主变容量{r['net_capacity_increment_mva']:.2f} MVA（购入{r['purchased_transformer_mva']:.2f} MVA）")
                if r['new_lines']:bits.append(f"新增联络{r['new_lines']}项")
                if r['new_storage_mwh']:bits.append(f"新增储能{r['new_storage_mwh']:.2f} MWh")
                if bits:actions.append(f"{r['year']}年"+'、'.join(bits))
            text.extend([f"{name}{LABEL[case]}："+('；'.join(actions) if actions else '四年均保留共同起点设备，既有网络通过本版静态工况校核，新增投资为0。')+'',''])
    text.extend(['市区下限对比反映地区容量指标对扩容的直接影响。邳州取消下限后，由较少的主变投资配合新增联络和少量储能满足静态安全条件。逐站更换规格、逐年购入容量和联络端点见工作簿。','',
                 '## 计算口径及核验','']+[r['事项']+'：'+r['说明']+'\n' for r in notes]+['',
                 markdown(['地区','方案','费用下界/万元','相对间隙','独立LP工况数'],
                          [[r['地区'],r['方案'],f"{r['费用下界_万元']:.2f}",f"{r['相对间隙']*100:.2f}%",r['独立LP工况数']] for r in solver_rows]),'',
                 '本版为独立试算，正式六章、已锁定推广矩阵和Word/PDF未修改。模型及复现入口：`实验/研究/rebuild_2026/case_parameters_20261012/README.md`。'])
    (RESULT/'结果说明.md').write_text('\n'.join(text),encoding='utf-8')
    dump(REVIEW/'导出核验.json',dict(verified=True,workbook=str(path.relative_to(ROOT)),sha256=sha(path),
                                   sheets=len(frames),checked_cells=cells,
                                   raw_inputs={str((OUT/r/'inputs.json').relative_to(ROOT)):sha(OUT/r/'inputs.json') for r in ('city','pizhou')}))
    print('EXPORTED',path,'checked cells',cells,flush=True)


if __name__=='__main__':report()
