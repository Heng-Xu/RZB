"""以 2025 年主变年度极值复核邳州逐时列的站与主变对应。"""

import argparse
import csv
from collections import defaultdict
from datetime import datetime, timedelta
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from .pizhou_mapping_candidates import INFERRED_STATIONS
from .official_annual import STUDY_DIR


STATS_SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/2025设备负载统计表.xlsx"
HOURLY_SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/邳州主变负载率.xlsx"
PROFILE_SOURCE = Path(__file__).resolve().parent / "source_audit/pizhou_2025_column_profile.csv"
TARGET = Path(__file__).resolve().parent / "source_audit/pizhou_2025_mapping_evidence.csv"


def read_transformer_stats(source: Path = STATS_SOURCE) -> dict[str, list[dict]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    stations = defaultdict(list)
    try:
        sheet = workbook["主变1"]
        for source_row, cells in enumerate(sheet.iter_rows(min_row=3, max_row=522, values_only=True), 3):
            if cells[2] != "QX-00005" or cells[1] not in (110, 35):
                continue
            stations[cells[3]].append(
                {
                    "transformer_id": cells[4],
                    "voltage_kv": int(cells[1]),
                    "capacity_mva": float(cells[5]),
                    "annual_max_mw": float(cells[14]) if cells[14] is not None else None,
                    "annual_max_time": cells[16],
                    "annual_min_mw": float(cells[17]) if cells[17] is not None else None,
                    "stats_source_row": source_row,
                }
            )
    finally:
        workbook.close()
    return dict(stations)


def read_peak_hour_values(source: Path, stats: dict[str, list[dict]]) -> dict[datetime, tuple]:
    times = set()
    for units in stats.values():
        for unit in units:
            raw = unit["annual_max_time"]
            if isinstance(raw, str) and raw.startswith("2025-"):
                hour = datetime.fromisoformat(raw).replace(minute=0, second=0)
                times.update((hour, hour + timedelta(hours=1)))
    workbook = load_workbook(source, read_only=True, data_only=True)
    values = {}
    try:
        for cells in workbook["Sheet3"].iter_rows(min_row=3, values_only=True):
            raw = str(cells[0] or "")
            if raw.startswith("2025-"):
                time = datetime.fromisoformat(raw)
                if time in times:
                    values[time] = cells
    finally:
        workbook.close()
    return values


def match_evidence(
    profile_source: Path = PROFILE_SOURCE,
    stats_source: Path = STATS_SOURCE,
    hourly_source: Path = HOURLY_SOURCE,
) -> list[dict]:
    with Path(profile_source).open(encoding="utf-8-sig", newline="") as handle:
        profiles = list(csv.DictReader(handle))
    stats = read_transformer_stats(stats_source)
    peak_hours = read_peak_hour_values(hourly_source, stats)
    grouped = defaultdict(list)
    for profile in profiles:
        column = int(profile["source_column"])
        station = INFERRED_STATIONS.get(column, profile["header_station_id"])
        grouped[station].append(profile)

    def discrepancy(profile: dict, unit: dict) -> float:
        return abs(float(profile["forward_peak_mw"]) - unit["annual_max_mw"]) + abs(
            float(profile["observed_min_mw"]) - unit["annual_min_mw"]
        )

    def peak_time_discrepancy(profile: dict, unit: dict) -> float:
        hour = datetime.fromisoformat(unit["annual_max_time"]).replace(minute=0, second=0)
        column = int(profile["source_column"]) - 1
        near = [
            float(peak_hours[time][column])
            for time in (hour, hour + timedelta(hours=1))
            if time in peak_hours and peak_hours[time][column] not in (None, "")
        ]
        return abs(max(near) - unit["annual_max_mw"])

    stats_hash = sha256(Path(stats_source).read_bytes()).hexdigest()
    rows = []
    for station in sorted(grouped):
        columns = sorted(grouped[station], key=lambda row: int(row["source_column"]))
        units = sorted(stats[station], key=lambda row: row["transformer_id"])
        if all(unit["annual_max_mw"] is not None for unit in units):
            direct = discrepancy(columns[0], units[0]) + discrepancy(columns[1], units[1])
            swapped = discrepancy(columns[0], units[1]) + discrepancy(columns[1], units[0])
            assigned = units if direct <= swapped else units[::-1]
            score = min(direct, swapped)
            alternative = max(direct, swapped)
            peak_time_score = peak_time_margin = None
            unit_status = "supported_by_extrema"
            if alternative - score < 2:
                direct_time = peak_time_discrepancy(columns[0], units[0]) + peak_time_discrepancy(columns[1], units[1])
                swapped_time = peak_time_discrepancy(columns[0], units[1]) + peak_time_discrepancy(columns[1], units[0])
                assigned = units if direct_time <= swapped_time else units[::-1]
                peak_time_score = min(direct_time, swapped_time)
                peak_time_margin = abs(direct_time - swapped_time)
                unit_status = "supported_by_peak_timestamp" if peak_time_margin >= 2 else "ambiguous_between_two_units"
            station_status = "supported_for_station_aggregation" if all(
                float(profile["forward_peak_mw"]) <= unit["annual_max_mw"] + 0.01
                and float(profile["observed_min_mw"]) >= unit["annual_min_mw"] - 0.01
                for profile, unit in zip(columns, assigned)
            ) else "extrema_conflict"
        else:
            assigned = [None, None]
            score = alternative = None
            peak_time_score = peak_time_margin = None
            unit_status = "unresolved_new_station_all_zero"
            station_status = "late_station_sparse_all_zero"
        for profile, unit in zip(columns, assigned):
            column = int(profile["source_column"])
            rows.append(
                {
                    "source_column": column,
                    "raw_header_station_id": profile["header_station_id"],
                    "resolved_station_id": station,
                    "resolved_voltage_kv": units[0]["voltage_kv"],
                    "resolved_transformer_id": unit["transformer_id"] if unit else "",
                    "transformer_capacity_mva": unit["capacity_mva"] if unit else "",
                    "hourly_max_mw": profile["forward_peak_mw"],
                    "hourly_min_mw": profile["observed_min_mw"],
                    "annual_max_mw": unit["annual_max_mw"] if unit else "",
                    "annual_min_mw": unit["annual_min_mw"] if unit else "",
                    "pair_score_mw": round(score, 6) if score is not None else "",
                    "alternative_pair_score_mw": round(alternative, 6) if alternative is not None else "",
                    "score_margin_mw": round(alternative - score, 6) if score is not None else "",
                    "peak_time_score_mw": round(peak_time_score, 6) if peak_time_score is not None else "",
                    "peak_time_margin_mw": round(peak_time_margin, 6) if peak_time_margin is not None else "",
                    "station_basis": "corrected_by_extrema_and_asset_census" if column in INFERRED_STATIONS else "raw_header_and_stats_station",
                    "station_status": station_status,
                    "unit_status": unit_status,
                    "annual_max_time": unit["annual_max_time"] if unit else "",
                    "stats_source_sheet": "主变1",
                    "stats_source_row": unit["stats_source_row"] if unit else "",
                    "stats_source_sha256": stats_hash,
                    "hourly_source_sha256": profile["source_sha256"],
                }
            )
    return sorted(rows, key=lambda row: row["source_column"])


def main() -> None:
    parser = argparse.ArgumentParser(description="核对邳州逐时列到主变的一对一映射")
    parser.add_argument("--output", type=Path, default=TARGET)
    args = parser.parse_args()
    rows = match_evidence()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    print(f"已写入 {len(rows)} 条映射证据")


if __name__ == "__main__":
    main()
