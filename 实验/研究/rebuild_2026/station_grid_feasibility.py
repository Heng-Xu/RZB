"""原表中的站点类别、扩建余量与 10 kV 供电网格；不推测地理距离。"""

from collections import defaultdict
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook

from .city_mapping_audit import write_csv


SOURCE_DIR = Path(__file__).resolve().parents[1] / "data/tuomin/电网建模数据_Agent整合版_V1.2"
ASSETS = SOURCE_DIR / "110（35）kv设备明细.xlsx"
FEEDERS = SOURCE_DIR / "邳州10kV线路基础数据表.xlsx"
AREA_RATINGS = {"A": (50.0, 63.0), "B": (40.0, 50.0, 63.0),
                "C": (31.5, 40.0, 50.0)}


def station_metadata(station_ids: set[str]) -> dict[str, dict]:
    sheet = load_workbook(ASSETS, read_only=True, data_only=True)["110千伏变电站1"]
    result = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
        station_id = row[3]
        if station_id not in station_ids or row[2] != 1:
            continue
        area = str(row[4]).strip()
        if area not in AREA_RATINGS:
            raise ValueError(f"{station_id} 供电区域 {area!r} 尚未建立主变候选")
        result[station_id] = {
            "station": station_id,
            "area_class": area,
            "available_third_slots": int(row[10] or 0),
            "available_third_mva": float(row[11] or 0),
            "total_10kv_bays": int(row[12] or 0),
            "spare_10kv_bays": int(row[13] or 0),
            "asset_source_sheet": sheet.title,
            "asset_source_row": row_number,
        }
    # 市区三份站级负荷原始文件没有对应设备明细行；按 C 类作为规划候选假定，
    # 不把 2025 年间隔或扩建余量假定为已核实设备事实。
    for station_id in station_ids - set(result):
        if station_id.startswith("SIM-CITY-"):
            result[station_id] = {
                "station": station_id, "area_class": "C",
                "available_third_slots": 0, "available_third_mva": 0.0,
                "total_10kv_bays": 0, "spare_10kv_bays": 0,
                "asset_source_sheet": "no_equipment_row_city_load_file_mapping",
                "asset_source_row": "",
                "area_class_status": "research_assumption_C_not_source_verified",
            }
    if set(result) != station_ids:
        raise ValueError(f"站点原表缺 {sorted(station_ids - set(result))}")
    return result


def grid_candidate_rows(station_ids: set[str]) -> list[dict]:
    sheet = load_workbook(FEEDERS, read_only=True, data_only=True).worksheets[0]
    grids = defaultdict(set)
    feeders = defaultdict(set)
    source_rows = defaultdict(set)
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        station_id, grid = row[12], row[4]
        if station_id not in station_ids or not isinstance(grid, str) or not grid.startswith("JS-XZ-PZ-"):
            continue
        grids[station_id].add(grid)
        feeders[station_id, grid].add(str(row[1]))
        source_rows[station_id, grid].add(row_number)
    metadata = station_metadata(station_ids)
    candidates = []
    for a, b in combinations(sorted(station_ids), 2):
        for grid in sorted(grids[a] & grids[b]):
            candidates.append({
                "station_a": a, "station_b": b, "shared_grid": grid,
                "station_a_feeders": "|".join(sorted(feeders[a, grid])),
                "station_b_feeders": "|".join(sorted(feeders[b, grid])),
                "station_a_source_rows": "|".join(map(str, sorted(source_rows[a, grid]))),
                "station_b_source_rows": "|".join(map(str, sorted(source_rows[b, grid]))),
                "station_a_spare_bays": metadata[a]["spare_10kv_bays"],
                "station_b_spare_bays": metadata[b]["spare_10kv_bays"],
                "both_have_spare_bay": int(
                    metadata[a]["spare_10kv_bays"] > 0 and metadata[b]["spare_10kv_bays"] > 0),
                "distance_km": "",
                "distance_status": "not_in_source_no_station_coordinates_or_route",
                "source_file": str(FEEDERS),
                "source_sheet": sheet.title,
            })
    return candidates


def write_screening_files(station_ids: set[str], directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    write_csv(list(station_metadata(station_ids).values()), directory / "station_physical_metadata.csv")
    write_csv(grid_candidate_rows(station_ids), directory / "shared_grid_station_pairs.csv")
