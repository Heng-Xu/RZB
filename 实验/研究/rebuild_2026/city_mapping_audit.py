"""以市区逐时文件为样本起点，核对站名、电压和设备表边界。"""

import argparse
import csv
from collections import defaultdict
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from .asset_2025 import read_assets
from .hourly_source_profile import CITY_SOURCE, OUTPUT_DIR
from .official_annual import STUDY_DIR


MAPPING_SOURCE = STUDY_DIR / "data/tuomin/变电站对应.md"
SUPPLEMENT_SOURCE = Path(__file__).resolve().parent / "city_2025_mapping_supplement.csv"
PROFILE_SOURCE = OUTPUT_DIR / "city_2025_series_profile.csv"
STATS_SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/2025设备负载统计表.xlsx"
ALIASES = {"西郊变新": "西郊变", "矿大变": "矿大"}


def read_name_mapping(source: Path = MAPPING_SOURCE) -> dict[str, list[dict]]:
    by_name = defaultdict(list)
    with Path(source).open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle, delimiter="\t"):
            if len(row) < 3 or not row[2].startswith("BDZ-"):
                continue
            full_name, name, station_id = (value.strip() for value in row[:3])
            voltage = next((value for value in (110, 220, 35) if full_name.startswith(f"{value}千伏")), None)
            by_name[name].append({"station_id": station_id, "voltage_kv": voltage})
    return dict(by_name)


def read_city_supplement(source: Path = SUPPLEMENT_SOURCE) -> dict[str, dict]:
    with Path(source).open(encoding="utf-8-sig", newline="") as handle:
        return {row["archive_member"]: row for row in csv.DictReader(handle)}


def read_city_station_extrema(source: Path = STATS_SOURCE) -> dict[str, tuple[float, float]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    result = {}
    try:
        for row in workbook["变电站1"].iter_rows(min_row=2, max_row=550, values_only=True):
            if row[2] == "QX-00007" and row[1] == 110 and row[3]:
                result[row[3]] = (float(row[12]), float(row[15]))
    finally:
        workbook.close()
    return result


def audit_city_mapping(
    mapping_source: Path = MAPPING_SOURCE,
    city_source: Path = CITY_SOURCE,
    profile_source: Path = PROFILE_SOURCE,
    supplement_source: Path = SUPPLEMENT_SOURCE,
) -> tuple[list[dict], list[dict]]:
    names = read_name_mapping(mapping_source)
    supplements = read_city_supplement(supplement_source)
    with Path(profile_source).open(encoding="utf-8-sig", newline="") as handle:
        profiles = list(csv.DictReader(handle))
    assets = read_assets()
    # 设备表每台主变占一行，站容量需按站合计。
    capacity = defaultdict(float)
    for row in assets:
        if row["region_id"] == "QX-00007" and row["voltage_kv"] == 110:
            capacity[row["station_id"]] += row["capacity_mva"]
    city_assets = {station: value for station, value in capacity.items() if value > 0}
    extrema = read_city_station_extrema()
    mapping_hash = sha256(Path(mapping_source).read_bytes()).hexdigest()
    supplement_hash = sha256(Path(supplement_source).read_bytes()).hexdigest()
    rows = []
    with ZipFile(city_source) as archive:
        for profile in profiles:
            member = profile["archive_member"]
            workbook = load_workbook(BytesIO(archive.read(member)), read_only=True, data_only=True)
            try:
                raw_name = workbook.active.title.strip()
            finally:
                workbook.close()
            supplement = supplements.get(member)
            normalized = supplement["station_name"] if supplement else ALIASES.get(raw_name, raw_name)
            candidates = names.get(normalized, [])
            eligible = [item for item in candidates if item["station_id"] in city_assets and item["voltage_kv"] == 110]
            match = eligible[0] if len(eligible) == 1 else None
            other_110 = [item for item in candidates if item["station_id"] not in city_assets and item["voltage_kv"] == 110]
            if supplement and supplement["station_id"] and not any(
                item["station_id"] == supplement["station_id"] and item["voltage_kv"] == 110
                for item in candidates
            ):
                raise ValueError(f"补充映射的站码与原对应表不一致：{member}")
            if profile["duplicate_of"]:
                status = "duplicate_series"
            elif match:
                status = "matched_city_110"
            elif candidates and all(item["voltage_kv"] == 220 for item in candidates):
                status = "known_220_out_of_scope"
            elif supplement and len(other_110) == 1:
                status = "confirmed_city_110_source_region_conflict"
            elif supplement and not candidates:
                status = "confirmed_city_110_no_station_id"
            elif candidates:
                status = "known_but_not_city_110"
            else:
                status = "name_absent_from_mapping"
            station_id = match["station_id"] if match else other_110[0]["station_id"] if status == "confirmed_city_110_source_region_conflict" else ""
            in_cohort = status in {"matched_city_110", "confirmed_city_110_source_region_conflict", "confirmed_city_110_no_station_id"}
            simulation_station_id = supplement["simulation_station_id"] if supplement else ""
            model_station_id = station_id or simulation_station_id
            if in_cohort and not model_station_id:
                raise ValueError(f"市区样本站缺少 BDZ 或仿真站标识：{member}")
            voltage_kv = 110 if in_cohort or (status == "duplicate_series" and match) else 220 if status == "known_220_out_of_scope" else ""
            voltage_evidence = (
                "mapping" if match or status in {"confirmed_city_110_source_region_conflict", "known_220_out_of_scope"}
                else "user_confirmation" if status == "confirmed_city_110_no_station_id" else ""
            )
            annual_max, annual_min = extrema.get(station_id, (None, None))
            hourly_max = float(profile["forward_peak_mw"])
            hourly_min = float(profile["observed_min_mw"])
            rows.append(
                {
                    "archive_member": member,
                    "matching_status": status,
                    "in_study_cohort": in_cohort,
                    "study_region_id": supplement["study_region_id"] if supplement else "QX-00007" if match else "",
                    "source_asset_region_id": supplement["source_asset_region_id"] if supplement else "QX-00007" if match else "",
                    "station_id": station_id,
                    "simulation_station_id": simulation_station_id,
                    "model_station_id": model_station_id,
                    "station_identity_basis": "source_bdz" if station_id else "simulation" if simulation_station_id else "",
                    "voltage_kv": voltage_kv,
                    "voltage_evidence": voltage_evidence,
                    "known_candidate_ids": ";".join(item["station_id"] for item in candidates),
                    "known_candidate_voltages": ";".join(str(item["voltage_kv"] or "") for item in candidates),
                    "name_alias_used": raw_name != normalized,
                    "user_scope_confirmation": bool(supplement),
                    "duplicate_of": profile["duplicate_of"],
                    "hourly_max_mw": hourly_max,
                    "hourly_min_mw": hourly_min,
                    "annual_max_mw": annual_max if annual_max is not None else "",
                    "annual_min_mw": annual_min if annual_min is not None else "",
                    "extrema_within_annual": (
                        hourly_max <= annual_max + 0.1 and hourly_min >= annual_min - 0.1
                    ) if match else "",
                    "mapping_source_sha256": mapping_hash,
                    "mapping_supplement_sha256": supplement_hash,
                    "city_source_sha256": profile["archive_sha256"],
                }
            )
    matched = {row["station_id"]: row["archive_member"] for row in rows if row["matching_status"] == "matched_city_110"}
    by_id = {
        item["station_id"]: name
        for name, items in names.items() for item in items
    }
    coverage = [
        {
            "station_id": station,
            "voltage_kv": 110,
            "asset_capacity_mva": city_assets[station],
            "archive_member": matched.get(station, ""),
            "coverage_status": "in_hourly_cohort" if station in matched else "outside_hourly_cohort",
            "name_present_in_mapping": station in by_id,
            "mapping_source_sha256": mapping_hash,
        }
        for station in sorted(city_assets)
    ]
    return rows, coverage


def write_csv(rows: list[dict], target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="以市区时序为样本起点核对 110 kV 站名与设备边界")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    rows, coverage = audit_city_mapping()
    write_csv(rows, args.output_dir / "city_2025_mapping_audit.csv")
    write_csv(coverage, args.output_dir / "city_2025_asset_coverage.csv")
    print(f"已核对 {len(rows)} 个时序文件、{sum(row['in_study_cohort'] for row in rows)} 个市区 110 kV 样本站序列；设备表参考站 {len(coverage)} 座")


if __name__ == "__main__":
    main()
