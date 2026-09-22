"""直接读取邳州六馈线原始台账和 7z 中的杆塔档距明细。"""

import argparse
import csv
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from subprocess import check_output

from openpyxl import load_workbook

from .official_annual import STUDY_DIR


FEEDER_SOURCE = STUDY_DIR / "data/tuomin/电网建模数据_Agent整合版_V1.2/邳州10kV线路基础数据表.xlsx"
SPAN_ARCHIVE = STUDY_DIR / "data/tuomin/10kv_case/source/10kv.7z"
OUTPUT_DIR = Path(__file__).resolve().parent / "source_audit"
FEEDER_MEMBERS = {
    "PZXL-00092": "10kv/邳州_10kVdnan线明细.xlsx",
    "PZXL-00097": "10kv/邳州_10kV墩西线明细.xlsx",
    "PZXL-00099": "10kv/邳州_10kV墩振线明细.xlsx",
    "PZXL-00154": "10kv/邳州_10kV河东线明细.xlsx",
    "PZXL-00161": "10kv/邳州_10kV河炮线明细.xlsx",
    "PZXL-00173": "10kv/邳州_10kV河镇线明细.xlsx",
}


def read_feeders(source: Path = FEEDER_SOURCE) -> list[dict]:
    source = Path(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    feeders = []
    try:
        sheet = workbook["邳州10千伏线路基础数据"]
        for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            feeder_id = row[1]
            if feeder_id not in FEEDER_MEMBERS:
                continue
            feeders.append(
                {
                    "feeder_id": feeder_id,
                    "station_id": row[12],
                    "voltage_kv": 10,
                    "max_allowed_current_a": row[7],
                    "reported_2025_max_current_a": row[8],
                    "reported_2025_max_active_power_mw": row[9],
                    "reported_2025_max_loading_pct": row[10],
                    "reported_total_length_km": row[13],
                    "reported_overhead_length_km": row[14],
                    "reported_cable_length_km": row[15],
                    "source_file": str(source.resolve().relative_to(STUDY_DIR)),
                    "source_sheet": sheet.title,
                    "source_row": source_row,
                    "source_sha256": source_hash,
                }
            )
    finally:
        workbook.close()
    return sorted(feeders, key=lambda row: row["feeder_id"])


def read_spans(archive: Path = SPAN_ARCHIVE) -> tuple[list[dict], list[dict]]:
    archive = Path(archive)
    archive_label = str(archive.resolve().relative_to(STUDY_DIR))
    archive_hash = sha256(archive.read_bytes()).hexdigest()
    spans, summaries = [], []
    for feeder_id, member in FEEDER_MEMBERS.items():
        raw = check_output(["7z", "x", "-so", str(archive), member])
        member_hash = sha256(raw).hexdigest()
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            sheet = workbook["档距段明细"]
            member_spans = []
            for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not row[0]:
                    continue
                member_spans.append(
                    {
                        "feeder_id": feeder_id,
                        "span_name": row[0],
                        "from_pole": row[1],
                        "to_pole": row[2],
                        "length_m": float(row[3]),
                        "conductor_model": row[4],
                        "source_archive": archive_label,
                        "source_member": member,
                        "source_sheet": sheet.title,
                        "source_row": source_row,
                    }
                )
        finally:
            workbook.close()
        spans.extend(member_spans)
        summaries.append(
            {
                "feeder_id": feeder_id,
                "span_count": len(member_spans),
                "span_length_sum_km": round(sum(row["length_m"] for row in member_spans) / 1000, 9),
                "length_50m_count": sum(row["length_m"] == 50 for row in member_spans),
                "length_60m_count": sum(row["length_m"] == 60 for row in member_spans),
                "other_length_count": sum(row["length_m"] not in (50, 60) for row in member_spans),
                "source_archive": archive_label,
                "archive_sha256": archive_hash,
                "source_member": member,
                "member_sha256": member_hash,
            }
        )
    return spans, summaries


def write_csv(rows: list[dict], target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="读取邳州六馈线原始边界与逐段档距")
    parser.add_argument("--feeder-source", type=Path, default=FEEDER_SOURCE)
    parser.add_argument("--span-archive", type=Path, default=SPAN_ARCHIVE)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    feeders = read_feeders(args.feeder_source)
    spans, summaries = read_spans(args.span_archive)
    write_csv(feeders, args.output_dir / "pizhou_six_feeders_2025.csv")
    write_csv(spans, args.output_dir / "pizhou_six_feeder_spans.csv")
    write_csv(summaries, args.output_dir / "pizhou_six_feeder_span_summary.csv")
    print(f"已写入 {len(feeders)} 条馈线、{len(spans)} 条档距段")


if __name__ == "__main__":
    main()
