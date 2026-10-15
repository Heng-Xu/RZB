"""逐年设备、费用、原始输入及新推广矩阵端点的独立复核。"""
import argparse
import hashlib
import json
import math
from concurrent.futures import ProcessPoolExecutor, as_completed

import numpy as np
import pandas as pd

from . import cases, matrix, export
from ..official_annual import read_official_annual, DEFAULT_SOURCE


def endpoint_signature(row):
    keys=['mode',*matrix.old.FIELDS,'units','modules','built','cost']
    return hashlib.sha256(json.dumps({k:row[k] for k in keys},sort_keys=True).encode()).hexdigest()


def endpoint_check(row):
    p={k:row[k] for k in matrix.old.FIELDS};p['normal_reverse_limit_fraction']=.8
    r=matrix.independent_lp_check(p,row['units'],row['modules'],set(row['built']))
    assert r['verified'],(p,r)
    return dict(raw_line=row['raw_line'],signature=endpoint_signature(row),mode=row['mode'],parameters=p,**r)


def verify_cases(selected):
    official={(r['region_id'],r['year']):r for r in read_official_annual(DEFAULT_SOURCE) if r['voltage_kv']==110}
    checked=[]
    for chosen in selected:
        folder=cases.ROOT/chosen['folder'];root=folder.parent
        inp=json.loads((root/'inputs.json').read_text());summary=json.loads((folder/'summary.json').read_text())
        assert cases.sha(root/'inputs.json')==chosen['input_sha256']
        assert cases.sha(folder/'chosen_solution.npz')==chosen['solution_sha256']
        for y,value in inp['annual_observed_peak_mw'].items():
            assert abs(value-official[inp['region_id'],int(y)]['reported_downward_load_mw'])<1e-8
        if chosen['baseline']=='annual_2021':assert inp['initial_capacity_mva']==official[inp['region_id'],2021]['capacity_mva']
        scenes=cases.old.base.load_scenes(cases.old.base.OUT/chosen['region'],inp);lp=cases.StationLP(inp)
        oldunits=np.array([s['initial_units']+[0.] for s in inp['stations'].values()]);oldn=np.zeros(len(oldunits));oldbuilt=set()
        total=0.;pairs=0;maxviolation=0.
        target={r['year']:r for r in summary.get('target_annual',[])}
        for a in summary['annual']:
            y=a['year'];units,n,built=cases.old.assets_from_csv(inp,folder,y)
            assert np.all(units>=oldunits-1e-7) and np.all(n>=oldn) and built.issuperset(oldbuilt)
            assert abs(units.sum()-a['capacity_mva'])<1e-7
            assert abs(units.sum()/a['peak_mw']-a['clr'])<1e-10
            assert abs(n.sum()*.215-a['storage_mwh'])<1e-7
            assert len(built)==a['cumulative_lines']
            bought=np.where(units[:,:2]>oldunits[:,:2]+1e-7,units[:,:2],0.).sum()
            third=(units[:,2]-oldunits[:,2]).sum()/50.
            capex=dict(transformer=float(bought*inp['transformer_price_10k_per_purchased_mva']+third*inp['third_50mva_capex_10k']),
                       storage=float((n-oldn).sum()*inp['storage_module_capex_10k']),line=float(len(built-oldbuilt)*inp['line_capex_10k']))
            costs={k:value*inp['cost_factors'][k][str(y)] for k,value in capex.items()}
            for k,v in costs.items():assert abs(v-a[k+'_npv_10k'])<.01,(chosen,k,v,a[k+'_npv_10k'])
            total+=sum(costs.values())
            if chosen['case']=='rigid':assert a['capacity_mva']<=max(oldunits.sum(),2*a['peak_mw'])+1e-7
            if target:
                t=target[y];assert a['capacity_mva']>=t['capacity_mva']+.5-1e-7
                assert a['storage_mwh']<=t['storage_mwh']+1e-7
                limit=max(0,t['cumulative_lines']-int(y==inp['years'][-1] and t['cumulative_lines']>0))
                assert len(built)<=limit
            for h,p in enumerate(scenes[y]):
                ready=lp.prepared(p,y,units,n,built)
                for f in range(-1,len(inp['stations'])):
                    r=lp.check(p,ready,f);pairs+=1
                    assert r['verified'],(chosen,y,h,f,r)
                    maxviolation=max(maxviolation,r.get('max_violation_mw',0.))
            oldunits,oldn,oldbuilt=units,n,built
        assert abs(total-summary['cost_npv_10k'])<.02
        assert summary['solver']['dual_bound']<=total+.01
        checked.append(dict(**chosen,checked_pairs=pairs,max_violation_mw=maxviolation,cost_recomputed=total,verified=True))
        print('VERIFIED_CASE',chosen['region'],chosen['baseline'],chosen['case'],pairs,flush=True)
    return checked


def verify_matrix(workers,with_lp):
    raw,rows=export.collect_matrix(matrix.OUT/'transfer_grid_procurement.jsonl')
    assert len(rows)==5832
    valid,pairs,adjacent,cells,public=export.matrix_data(rows)
    assert len(cells)<=864 and set(cells['模型'])=={export.MODE['stock'],export.MODE['greenfield']}
    assert all(r['status']=='Infeasible' for r in rows if not r['verified']), '尚有未求得有效结论的节点'
    endpoints=set(cells['下界原始行'])|set(cells['上界原始行'])
    chosen=[r for r in rows if r['raw_line'] in endpoints]
    maxcosterr=0.;maxgap=0.
    for r in rows:
        if not r['verified']:continue
        allowed=set(matrix.CATALOG) if r['mode']=='greenfield' else {50.,63.,80.}
        assert all(v in allowed for unit in r['units'] for v in unit[:2]),'新购规格越界'
        inp=matrix.old.make_input({k:r[k] for k in matrix.old.FIELDS});units=np.array(r['units']);n=np.array(r['modules'])
        purchased=units[:,:2].sum() if r['mode']=='greenfield' else np.where(units[:,:2]>50,units[:,:2],0.).sum()
        rawcost=dict(transformer=float(purchased*inp['transformer_cost_per_purchased_mva']+units[:,2].sum()/50*inp['third_cost']),
            storage=float(n.sum()*inp['storage_unit_cost']),line=float(len(r['built'])*inp['line_cost']))
        cost=sum(v*inp['factors'][k] for k,v in rawcost.items())
        assert abs(cost-r['cost'])<1e-4
        assert r['cost_lower']<=cost+1e-4
        assert r['physical_check']['max_violation_mw']<1e-5
        assert abs(units.sum()/inp['peak']-r['clr'])<1e-10
        maxcosterr=max(maxcosterr,abs(cost-r['cost']));maxgap=max(maxgap,r['relative_gap'])
    proofs=[]
    if with_lp:
        with ProcessPoolExecutor(max_workers=workers) as pool:
            futures=[pool.submit(endpoint_check,r) for r in chosen]
            for i,f in enumerate(as_completed(futures),1):
                proofs.append(f.result())
                if i%50==0:print('VERIFIED_MATRIX_ENDPOINT',i,'/',len(chosen),flush=True)
        cases.dump(cases.REVIEW/'矩阵端点独立LP.json',dict(verified=True,endpoint_count=len(proofs),records=proofs))
    else:
        proof_file=cases.REVIEW/'矩阵端点独立LP.json'
        assert proof_file.exists(),'须先执行--lp，生成当前矩阵的端点独立证书'
        saved=json.loads(proof_file.read_text())
        assert saved['verified']
        assert {r['raw_line']:r['signature'] for r in saved['records']}=={r['raw_line']:endpoint_signature(r) for r in chosen}
        for r in saved['records']:assert r['verified'] and r.get('max_violation_mw',0.)<1e-5
        proofs=saved['records']
    return dict(raw_records=len(raw),unique_nodes=len(rows),valid_nodes=len(valid),cells=len(cells),
        designed_cells=864,endpoint_count=len(chosen),independent_lp_count=len(proofs),max_cost_error=maxcosterr,max_relative_gap=maxgap,verified=True)


def verify_workbooks():
    from openpyxl import load_workbook
    tables,_=export.cases_data();tables['逐年转供']=pd.read_csv(cases.PUBLIC/'逐年转供.csv')
    _,rows=export.collect_matrix(matrix.OUT/'transfer_grid_procurement.jsonl')
    _,pairs,adjacent,cells,public=export.matrix_data(rows)
    pilot=export.records(matrix.OUT/'pilot_procurement.jsonl')
    pilot_rows=pd.DataFrame([dict(模型=export.MODE[r['mode']],转供率=r['initial_transfer_fraction'],支撑时间_h=r['support_duration_h'],
        主变_MVA=r.get('capacity_mva'),容载比=r.get('clr'),储能_MWh=r.get('storage_mwh'),新增联络_条=r.get('line_count'),
        费用现值_万元=r.get('cost'),间隙=r.get('relative_gap')) for r in pilot])
    sheets={'原句完整序列':pilot_rows,'首尾影响':pairs,'相邻节点':adjacent,'全部单元':cells,'全部节点':public}
    sheets.update({'新_'+k:v for k,v in export.matrix_frames(cells,'greenfield').items()})
    sheets.update({'旧_'+k:v for k,v in export.matrix_frames(cells,'stock').items()})
    def normalized(v):
        if isinstance(v,(list,dict,tuple)):return json.dumps(v,ensure_ascii=False)
        if v is None or v=='' or isinstance(v,(float,np.floating)) and not math.isfinite(v):return None
        return v
    records=[]
    for filename,frames in [('区县案例逐年方案.xlsx',tables),('推广转供率复核矩阵.xlsx',sheets)]:
        book=load_workbook(cases.PUBLIC/filename,read_only=True,data_only=True);count=0
        assert set(book.sheetnames)==set(frames)|({'说明'} if filename.startswith('推广') else set())
        for title,frame in frames.items():
            sheet=book[title];values=list(sheet.values)
            assert len(values)==len(frame)+1 and list(values[0])==list(frame.columns),(filename,title,'表头或行数')
            for i,(actual,expected) in enumerate(zip(values[1:],frame.itertuples(index=False,name=None)),2):
                for j,(x,y) in enumerate(zip(actual,expected),1):
                    x,y=normalized(x),normalized(y)
                    ok=math.isclose(x,y,rel_tol=1e-10,abs_tol=1e-8) if isinstance(x,(int,float)) and isinstance(y,(int,float)) else x==y
                    assert ok,(filename,title,i,j,x,y)
                    count+=1
        records.append(dict(file=filename,sha256=cases.sha(cases.PUBLIC/filename),sheets=len(book.sheetnames),data_cells=count,verified=True))
        book.close()
    cases.dump(cases.REVIEW/'工作簿逐格核验.json',dict(verified=True,records=records))
    return records


def main():
    p=argparse.ArgumentParser();p.add_argument('--lp',action='store_true');p.add_argument('--workers',type=int,default=8)
    a=p.parse_args();stats=json.loads((cases.REVIEW/'统计核验.json').read_text())
    assert stats['status']=='complete'
    case_checks=verify_cases(stats['selected_cases']);matrix_check=verify_matrix(a.workers,a.lp)
    workbook_checks=verify_workbooks()
    cases.dump(cases.REVIEW/'最终模型核验.json',dict(verified=True,cases=case_checks,matrix=matrix_check,
        workbooks=workbook_checks,
        normative_sources=json.loads((cases.SOURCE/'规范核读定位.json').read_text())))
    print('VERIFIED_REVIEW_20261015',sum(r['checked_pairs'] for r in case_checks),matrix_check,flush=True)


if __name__=='__main__':main()
