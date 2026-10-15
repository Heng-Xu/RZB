"""导出会后建模复核的逐年结果、转供完整序列和有限节点矩阵。"""
import argparse
import hashlib
import itertools
import json
import math
import os
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import cases, matrix

ROOT = cases.ROOT; OUT = cases.PUBLIC; REVIEW = cases.REVIEW; SOURCE = cases.SOURCE
MODE = {'stock': '存量改造', 'stock_allcost': '500MVA下限且全部计价', 'greenfield': '主变重新选配'}
CASE = {'rigid': '刚性', 'elastic_free': '弹性自由寻优', 'elastic_target': '弹性指定替代对照'}
BASE = {'compressed': '共同压缩起点', 'annual_2021': '保留2021原表地区容量'}
FIELDS = matrix.old.FIELDS


def dump(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def records(path):
    return [json.loads(s) for s in path.read_text().splitlines() if s.strip()]


def collect_matrix(path):
    raw = records(path); grouped = {}; statuses={}
    for line, r in enumerate(raw, 1):
        key = (r['mode'], *(r[k] for k in FIELDS))
        statuses.setdefault(key,set()).add('feasible' if r['verified'] else r['status'])
        assert not {'feasible','Infeasible'}.issubset(statuses[key]), ('同一节点出现矛盾可行性结论',key)
        r = dict(r, raw_line=line)
        previous = grouped.get(key)
        rank = lambda x: (not x['verified'], x.get('cost', np.inf), x.get('relative_gap') or 0.)
        if previous is None or rank(r) < rank(previous): grouped[key] = r
    return raw, list(grouped.values())


def workbook(path, sheets):
    w = Workbook(); w.remove(w.active)
    for name, data in sheets.items():
        s = w.create_sheet(name[:31]); frame = data if isinstance(data, pd.DataFrame) else pd.DataFrame(data)
        if not len(frame.columns): continue
        s.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            s.append([json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict, tuple)) else
                      None if (v is None or isinstance(v, (float, np.floating)) and not math.isfinite(v)) else v for v in row])
        s.freeze_panes = 'A2'; s.auto_filter.ref = s.dimensions
        for c in s[1]:
            c.font = Font(name='方正仿宋_GBK', bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='244C6A'); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name='方正仿宋_GBK', size=11)
                c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                if isinstance(c.value, float): c.number_format = '0.00'
                label = str(frame.columns[c.column - 1])
                if isinstance(c.value, (float, int)) and any(k in label for k in ('转供率','增长率','年增长率','电缆长度占比','降幅','占比')):
                    c.number_format = '0%'
                if isinstance(c.value, (float, int)) and ('间隙' in label): c.number_format = '0.00%'
        for i, col in enumerate(frame.columns, 1): s.column_dimensions[get_column_letter(i)].width = min(42, max(14, len(str(col)) * 2))
        s.row_dimensions[1].height = 36
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.orientation = 'landscape'; s.page_setup.paperSize = s.PAPERSIZE_A3
        s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
        s.print_title_rows = '1:1'
    w.save(path)


def cases_data():
    annual = []; stations = []; projects = []; fees = []; checks = []; inputs = []; initial = []; chosen = []
    for region in ('pizhou', 'city'):
        for baseline in ('compressed', 'annual_2021'):
            root = cases.OUT / region / (baseline + '_reverse80')
            inp = json.loads((root / 'inputs.json').read_text())
            for sid, rate in inp['station_transfer_rates'].items():
                evidence = next(r for r in cases.old.station_evidence(inp) if r['station'] == sid)
                for y in inp['years']:
                    inputs.append(dict(地区=inp['region_name'], 起点=BASE[baseline], 站号=sid, 年份=y,
                        原表J_MW=evidence['transferable_mw'], 原表K_MW=evidence['unserved_mw'], 逐站转供率=rate,
                        模型既有转供上限_MW=inp['existing_transfer_budget_by_year_mw'][str(y)][sid],
                        原表行号=evidence['source_row'], 匹配=evidence['matched']))
            if (root / 'initial_capacity_fit.json').exists():
                initial.extend(dict(地区=inp['region_name'], **r) for r in records_array(root / 'initial_capacity_fit.json'))
            for case in CASE:
                folder = cases.best_folder(root, case)
                if folder is None:
                    if case == 'elastic_target': continue
                    raise FileNotFoundError(root / case / 'summary.json')
                r = json.loads((folder / 'summary.json').read_text()); solver = r['solver']
                label = dict(地区=inp['region_name'], 起点=BASE[baseline], 方案=CASE[case])
                chosen.append(dict(region=region, baseline=baseline, case=case, folder=str(folder.relative_to(ROOT)),
                    input_sha256=cases.sha(root / 'inputs.json'), solution_sha256=cases.sha(folder / 'chosen_solution.npz'),
                    cost=r['cost_npv_10k'], lower=solver['dual_bound'], gap=solver['mip_gap']))
                fees.append(dict(**label, 费用现值_万元=r['cost_npv_10k'], 费用下界_万元=solver['dual_bound'],
                    相对间隙=solver['mip_gap'], 比较条件=r['scenario_constraint'], 原始目录=str(folder.relative_to(ROOT))))
                checks.extend(dict(**label, **item) for item in r['checks'])
                oldcap = inp['initial_capacity_mva']
                for a in r['annual']:
                    annual.append(dict(**label, 年份=a['year'], 当年负荷_MW=a['peak_mw'], 年末主变_MVA=a['capacity_mva'],
                        容载比=a['clr'], 当年净增容量_MVA=a['capacity_mva']-oldcap,
                        当年新购主变_MVA=a['purchased_transformer_mva'], 新增第三台主变_台=a['new_third_count'],
                        当年新增储能_MWh=a['new_storage_mwh'], 年末储能_MW=a['storage_nameplate_mw'], 年末储能_MWh=a['storage_mwh'],
                        当年新增联络_条=a['new_lines'], 累计新增联络_条=a['cumulative_lines'],
                        当年主变费用现值_万元=a['transformer_npv_10k'], 当年储能费用现值_万元=a['storage_npv_10k'],
                        当年联络费用现值_万元=a['line_npv_10k'], 当年费用现值_万元=a['added_lifecycle_npv_10k']))
                    oldcap = a['capacity_mva']
                stations.extend(dict(**label, **item) for item in cases.old.base.previous.rows(folder / 'stations.csv'))
                projects.extend(dict(**label, **item) for item in cases.old.base.previous.rows(folder / 'projects.csv'))
    return dict(逐年方案=pd.DataFrame(annual), 费用及精度=pd.DataFrame(fees), 逐站设备=pd.DataFrame(stations),
                联络项目=pd.DataFrame(projects), 逐站转供依据=pd.DataFrame(inputs), 起点容量标定=pd.DataFrame(initial),
                承载力核验=pd.DataFrame(checks)), chosen


def records_array(path): return json.loads(path.read_text())


def transfer_data(selected):
    yearly=[];detail=[];source_key=[(r['folder'],r['solution_sha256'],r['input_sha256']) for r in selected]
    cache=REVIEW/'转供调度核验.json';csvpath=OUT/'逐年转供.csv'
    saved_keys=set();cached_year=None;cached_detail=None
    if cache.exists() and csvpath.exists():
        saved=json.loads(cache.read_text())
        if saved['source_key']==[list(r) for r in source_key] and saved['model_sha256']==cases.sha(Path(cases.__file__)):
            return pd.read_csv(csvpath)
        detail_path=cases.OUT.parent/'案例逐工况最小转供.csv'
        if saved['model_sha256']==cases.sha(Path(cases.__file__)) and detail_path.exists():
            saved_keys={tuple(r) for r in saved['source_key']};cached_year=pd.read_csv(csvpath);cached_detail=pd.read_csv(detail_path)
    for chosen in selected:
        folder=ROOT/chosen['folder'];inp=json.loads((folder.parent/'inputs.json').read_text())
        signature=(chosen['folder'],chosen['solution_sha256'],chosen['input_sha256'])
        if signature in saved_keys:
            def selected_rows(frame):
                return frame[(frame['地区'].eq(inp['region_name']))&(frame['起点'].eq(BASE[chosen['baseline']]))&(frame['方案'].eq(CASE[chosen['case']]))].to_dict('records')
            yearly.extend(selected_rows(cached_year));detail.extend(selected_rows(cached_detail));continue
        scenes=cases.old.base.load_scenes(cases.old.base.OUT/chosen['region'],inp);lp=cases.StationLP(inp)
        for y in inp['years']:
            units,n,built=cases.old.assets_from_csv(inp,folder,y);rows=[]
            for h,p in enumerate(scenes[y]):
                ready=lp.prepared(p,y,units,n,built)
                for f in range(-1,len(inp['stations'])):
                    r=lp.dispatch(p,ready,f)
                    row=dict(地区=inp['region_name'],起点=BASE[chosen['baseline']],方案=CASE[chosen['case']],年份=y,
                        截面=h,故障站=lp.ids[f] if f>=0 else '正常',工况='正常' if f<0 else 'N-1',
                        最小总转供_MW=r['total_mw'],所用既有转供_MW=r['existing_mw'],所用新增转供_MW=r['new_mw'],
                        约束最大偏差_MW=r['max_violation_mw'])
                    rows.append(row);detail.append(row)
            normals=[r for r in rows if r['工况']=='正常'];faults=[r for r in rows if r['工况']=='N-1']
            normal=max(normals,key=lambda r:r['最小总转供_MW']);fault=max(faults,key=lambda r:r['最小总转供_MW'])
            yearly.append(dict(地区=inp['region_name'],起点=BASE[chosen['baseline']],方案=CASE[chosen['case']],年份=y,
                正常最大总转供_MW=normal['最小总转供_MW'],正常对应既有_MW=normal['所用既有转供_MW'],正常对应新增_MW=normal['所用新增转供_MW'],
                N1最大总转供_MW=fault['最小总转供_MW'],N1对应既有_MW=fault['所用既有转供_MW'],N1对应新增_MW=fault['所用新增转供_MW'],
                正常最大工况=normal['截面'],N1最大工况=fault['截面'],N1故障站=fault['故障站'],检查工况数=len(rows)))
        print('TRANSFER_EXPORTED',chosen['region'],chosen['baseline'],chosen['case'],flush=True)
    pd.DataFrame(detail).to_csv(cases.OUT.parent/'案例逐工况最小转供.csv',index=False,encoding='utf-8-sig')
    d=pd.DataFrame(yearly);d.to_csv(csvpath,index=False,encoding='utf-8-sig')
    dump(cache,dict(source_key=source_key,model_sha256=cases.sha(Path(cases.__file__)),checked_pairs=len(detail),verified=True,
        objective='固定设备逐工况最小化总转供MW；既有与新增为所选最小总转供调度的分量，各互斥故障不累加成年电量'))
    return d


def matrix_data(rows):
    valid = pd.DataFrame([r for r in rows if r['verified']]); samples = []; cells = []; pairs = []
    fixed = ['mode', 'eta', 'growth_rate', 'cable_fraction', 'support_duration_h', 'diversity']
    for key, group in valid.groupby(fixed):
        group = group.sort_values('initial_transfer_fraction')
        by = {round(r['initial_transfer_fraction'], 8): r for r in group.to_dict('records')}
        if .1 in by and .9 in by:
            lo, hi = by[.1], by[.9]
            pairs.append(dict(zip(fixed, key)) | dict(低转供容量_MVA=lo['capacity_mva'], 高转供容量_MVA=hi['capacity_mva'],
                容量减少_MVA=lo['capacity_mva']-hi['capacity_mva'], 容载比减少=lo['clr']-hi['clr'],
                容量降幅=(lo['capacity_mva']-hi['capacity_mva'])/lo['capacity_mva'],
                费用减少_万元=lo['cost']-hi['cost'], 有效节点数=len(group)))
        for a, b in zip(group.to_dict('records'), group.to_dict('records')[1:]):
            samples.append(dict(zip(fixed, key)) | dict(转供率起=a['initial_transfer_fraction'], 转供率止=b['initial_transfer_fraction'],
                容量变化_MVA=b['capacity_mva']-a['capacity_mva'], 容载比变化=b['clr']-a['clr'], 费用变化_万元=b['cost']-a['cost']))
    for mode in ('stock', 'greenfield'):
        d = valid[valid['mode'].eq(mode)].copy()
        d['转供分档'] = d.initial_transfer_fraction.map(lambda x: 0 if x <= .3+1e-9 else 1 if x <= .5+1e-9 else 2 if x <= .7+1e-9 else 3)
        for (q, kd, eta, growth, rb), g in d.groupby(['cable_fraction', 'diversity', 'eta', 'growth_rate', '转供分档']):
            lower = g.loc[g.clr.idxmin()]; upper = g.loc[g.clr.idxmax()]
            cells.append(dict(模型=MODE[mode], 电缆长度占比=q, 负荷分散系数=kd, 装机渗透率节点=eta,
                年增长率节点=growth, 转供分档=rb, 转供区间=['≤30%', '(30%,50%]', '(50%,70%]', '>70%'][rb],
                容载比下界=float(lower.clr), 容载比上界=float(upper.clr), 有效节点数=len(g),
                下界原始行=int(lower.raw_line), 上界原始行=int(upper.raw_line),
                计算条件='2/4/6h有限节点包络；同一电缆、分散、装机、增长节点；未插值'))
    public = []
    for r in rows:
        item = dict(模型=MODE[r['mode']], **{matrix.old.FIELDS[i]:r[matrix.old.FIELDS[i]] for i in range(6)},
                    状态=r['status'], 已验证=r['verified'], 原始行=r['raw_line'], 求解间隙=r.get('relative_gap'))
        if r['verified']:
            item.update(主变容量_MVA=r['capacity_mva'], 容载比=r['clr'], 储能_MW=r['storage_mw'],
                储能_MWh=r['storage_mwh'], 新增联络_条=r['line_count'], 费用现值_万元=r['cost'], 费用下界_万元=r['cost_lower'])
        public.append(item)
    return valid, pd.DataFrame(pairs), pd.DataFrame(samples), pd.DataFrame(cells), pd.DataFrame(public)


def matrix_frames(cells, mode):
    frames = {}
    for (q, kd), d in cells[cells['模型'].eq(MODE[mode])].groupby(['电缆长度占比', '负荷分散系数']):
        data = []
        for (growth, eta), g in d.groupby(['年增长率节点', '装机渗透率节点']):
            row = {'年增长率节点':growth, '装机渗透率节点':eta}
            for b in range(4):
                selected = g[g['转供分档'].eq(b)]
                if selected.empty: row[['≤30%', '(30%,50%]', '(50%,70%]', '>70%'][b]] = '无有效节点'
                else:
                    r = selected.iloc[0]
                    row[r['转供区间']] = f"{math.floor((r['容载比下界']+1e-10)*100)/100:.2f}～{math.ceil((r['容载比上界']-1e-10)*100)/100:.2f}"
            data.append(row)
        frames[f'电缆{q:.0%}_分散{kd:.2f}'] = pd.DataFrame(data)
    return frames


def plots(annual, valid, pilot):
    os.environ.setdefault('MPLCONFIGDIR','/tmp/xuzhou110kv-review-matplotlib')
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.font_manager import fontManager, FontProperties
    fontManager.addfont('/usr/share/fonts/wps-office/FZFSK.TTF')
    chinese = FontProperties(fname='/usr/share/fonts/wps-office/FZFSK.TTF').get_name()
    plt.rcParams.update({'font.family':['Times New Roman',chinese], 'axes.unicode_minus':False, 'font.size':11})
    figures = OUT / '图件'; figures.mkdir(exist_ok=True)
    d = pd.DataFrame([r for r in pilot if r['verified'] and r['support_duration_h']==4.])
    fig, axes = plt.subplots(1, 2, figsize=(11,4.4))
    for mode, g in d.groupby('mode'):
        g=g.sort_values('initial_transfer_fraction')
        axes[0].plot(g.initial_transfer_fraction*100,g.capacity_mva,marker='o',label=MODE[mode])
        axes[1].plot(g.initial_transfer_fraction*100,g.clr,marker='o',label=MODE[mode])
    for ax in axes: ax.set_xlabel('初始转供率/%'); ax.grid(alpha=.2); ax.legend(fontsize=9)
    axes[0].set_ylabel('主变容量/MVA');axes[1].set_ylabel('容载比')
    fig.tight_layout();fig.savefig(figures/'原句条件转供完整序列.png',dpi=220);fig.savefig(figures/'原句条件转供完整序列.pdf');plt.close(fig)
    for region in ('邳州','市区'):
        fig, axes=plt.subplots(1,3,figsize=(13,4.4))
        for case,g in annual[(annual['地区'].eq(region))&(annual['起点'].eq(BASE['compressed']))].groupby('方案'):
            g=g.sort_values('年份')
            for ax,col in zip(axes,['容载比','年末储能_MWh','累计新增联络_条']):ax.plot(g['年份'],g[col],marker='o',label=case)
        for ax,label in zip(axes,['容载比','年末储能/MWh','累计新增联络/条']):
            ax.set_xticks([2022,2023,2024,2025]);ax.set_ylabel(label);ax.grid(alpha=.2);ax.legend(fontsize=8)
        fig.tight_layout();fig.savefig(figures/f'{region}逐年方案.png',dpi=220);fig.savefig(figures/f'{region}逐年方案.pdf');plt.close(fig)


def main():
    p=argparse.ArgumentParser();p.add_argument('--partial',action='store_true');p.add_argument('--transfers',action='store_true');a=p.parse_args()
    for folder in (OUT,REVIEW,SOURCE):folder.mkdir(parents=True,exist_ok=True)
    raw, rows=collect_matrix(matrix.OUT/'transfer_grid_procurement.jsonl')
    if not a.partial: assert len(rows)==5832, f'矩阵计算未完成：{len(rows)}/5832'
    tables,chosen=cases_data()
    if a.transfers:tables['逐年转供']=transfer_data(chosen)
    for name,frame in tables.items():frame.to_csv(OUT/(name+'.csv'),index=False,encoding='utf-8-sig')
    workbook(OUT/'区县案例逐年方案.xlsx',tables)
    valid,pairs,adjacent,cells,public=matrix_data(rows)
    for name,frame in [('推广节点',public),('转供首尾比较',pairs),('转供相邻节点',adjacent),('转供分档单元',cells)]:
        frame.to_csv(OUT/(name+'.csv'),index=False,encoding='utf-8-sig')
    pilot=records(matrix.OUT/'pilot_procurement.jsonl')
    pilot_rows=pd.DataFrame([dict(模型=MODE[r['mode']],转供率=r['initial_transfer_fraction'],支撑时间_h=r['support_duration_h'],
        主变_MVA=r.get('capacity_mva'),容载比=r.get('clr'),储能_MWh=r.get('storage_mwh'),新增联络_条=r.get('line_count'),
        费用现值_万元=r.get('cost'),间隙=r.get('relative_gap')) for r in pilot])
    sheets={'说明':pd.DataFrame([dict(事项='有限节点设计',说明='3个装机×4个增长×3个电缆×3个支撑时长×3个分散×9个转供×2模型，共5832；原有有效库保留'),
        dict(事项='用途',说明='存量改造与主变重新选配分开；表内区间仅为已计算节点包络，不补造无效节点'),
        dict(事项='分类',说明='增长0/4/5.5/8%，装机0.5/1.5/2.5，电缆10/30/50%，分散1/1.10/1.15；支撑2/4/6h')]),
        '原句完整序列':pilot_rows,'首尾影响':pairs,'相邻节点':adjacent,'全部单元':cells,'全部节点':public}
    sheets.update({'新_'+k:v for k,v in matrix_frames(cells,'greenfield').items()})
    sheets.update({'旧_'+k:v for k,v in matrix_frames(cells,'stock').items()})
    workbook(OUT/'推广转供率复核矩阵.xlsx',sheets)
    metrics={}
    for mode,g in pairs.groupby('mode'):
        adj=adjacent[adjacent['mode'].eq(mode)]
        sub=valid[valid['mode'].eq(mode)].copy(); fixed=[k for k in FIELDS if k!='initial_transfer_fraction']
        x=sub.initial_transfer_fraction-sub.groupby(fixed).initial_transfer_fraction.transform('mean')
        y=sub.clr-sub.groupby(fixed).clr.transform('mean')
        metrics[mode]=dict(首尾均有效组数=len(g),首尾容量下降组数=int((g['容量减少_MVA']>1e-7).sum()),
            平均首尾容量降幅=float(g['容量降幅'].mean()),平均首尾容载比下降=float(g['容载比减少'].mean()),
            相邻不变占比=float((adj['容量变化_MVA'].abs()<1e-7).mean()),相邻上升占比=float((adj['容量变化_MVA']>1e-7).mean()),
            转供率与容载比组内Pearson=float((x@y)/np.sqrt((x@x)*(y@y))))
    dump(REVIEW/'统计核验.json',dict(status='partial' if a.partial else 'complete',raw_records=len(raw),unique_nodes=len(rows),valid_nodes=len(valid),
        infeasible_nodes=sum(r['status']=='Infeasible' for r in rows),other_unverified_nodes=sum(not r['verified'] and r['status']!='Infeasible' for r in rows),
        cells=len(cells),metrics=metrics,selected_cases=chosen))
    plots(tables['逐年方案'],valid,pilot)
    print(json.dumps(dict(nodes=len(rows),valid=len(valid),cells=len(cells),metrics=metrics),ensure_ascii=False),flush=True)


if __name__=='__main__':main()
