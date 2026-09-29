"""从现行模型结果和来源登记生成可复核的研究报告附录。"""

import csv
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
OUT = ROOT / '实验/研究/rebuild_2026/outputs'
SRC = ROOT / '研究报告/数据来源/2026-09-28_10kV线路与成本依据'
REGIONS = [('邳州', 'ordered_guide_pizhou_n1'), ('市区', 'ordered_guide_city_n1')]
SCHEMES = [('刚性', 'rigid'), ('弹性', 'elastic')]


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def n(value, decimals=3):
    return f'{float(value):.{decimals}f}'


def table(headers, rows):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |',
                      '| ' + ' | '.join(['---'] * len(headers)) + ' |'] +
                     ['| ' + ' | '.join(map(str, r)) + ' |' for r in rows]) + '\n\n'


def save(name, text):
    (HERE / name).write_text(text.rstrip() + '\n', encoding='utf-8')


def appendix_a():
    s = '# 附录 A 原始资料定位、指标换算和参数登记\n\n'
    s += '本附录给出正文数值的原始定位、计算式和证据性质。原始收资表及公开资料的文件哈希由资料登记表保存。模型结果的计算值不得称为原表实测。\n\n'
    manifest = read_csv(SRC / 'source_manifest.csv')
    s += '## A.1 来源索引\n\n'
    s += table(['来源代号', '原始定位', '使用限制'], [
        (r['source_id'], r['original_locator_and_value'], r['use_and_limitation'])
        for r in manifest])
    s += '来源文件的 SHA-256 与原始网址见研究报告配套的《source_manifest.csv》；原始资料可按来源代号检索。正式导则为 DL/T 5729—2023，2025 年文件仅作修订稿研究指标参考。\n\n'
    s += '## A.2 区县年度统计表逐格换算\n\n'
    s += '《近5年容载比》Sheet1 原表负荷单位为万千瓦，乘 10 后为 MW；原表容量单位为万千伏安，乘 10 后为 MVA。各年列按 D/G/J/M/P（负荷）、C/F/I/L/O（容量）排列。以下各值用于核对本轮负荷代理与历史容量背景。\n\n'
    import openpyxl
    book = openpyxl.load_workbook(SRC / '项目收资与已存档依据/近5年容载比.xlsx', read_only=True, data_only=True)
    ws = book['Sheet1']
    rows = []
    for name, row in [('邳州', 19), ('市区', 9), ('丰县', 34)]:
        for year, capcol, loadcol in zip(range(2021, 2026), 'CFILO', 'DGJMP'):
            cap = ws[f'{capcol}{row}'].value
            load = ws[f'{loadcol}{row}'].value
            rows.append((name, year, f'Sheet1!{capcol}{row}', n(cap, 4), n(cap*10, 3),
                         f'Sheet1!{loadcol}{row}', n(load, 4), n(load*10, 3)))
    s += table(['区县', '年', '容量原格', '原表万 kVA', '换算 MVA', '负荷原格', '原表万 kW', '换算 MW'], rows)
    s += '2021 年仿真容量另按离散主变规格重设：邳州 1463.5 MVA、市区 2462.0 MVA。它们分别除以同年 D19、D9 换算的负荷得到 1.999945、1.698200；原表容量只作为历史背景。2022 年起容量不减。\n\n'
    s += '## A.3 增长率、光伏和补充指标\n\n'
    s += '年度增长率为当年负荷除上年负荷再减一。2022—2025 年同比率按 1、2、3、4 加权，再除以 10；邳州为 8.0742%，市区为 1.8000%。这是回顾性分档计算，不是导则规定的权重或未来预测。\n\n'
    s += '分布式光伏装机取《逐月分县分布式光伏》相应区县年度 12 月单元格，万千瓦乘 10 换算成 MW。2021—2022 年缺少原表记录。源荷比的分母为用户最大用电负荷；电量渗透率的分母为用户年用电量，不以《近5年容载比》的降压负荷直接替代。\n\n'
    x = read_csv(SRC / '丰县2025补充指标估计.csv')[0]
    s += table(['丰县补充量', '估计值', '计算或限制'], [
        ('2025 年光伏装机 MW', n(x['年末分布式光伏装机MW_原表'], 3), x['分布式光伏原表格']),
        ('2025 年降压负荷 MW', n(x['110kV降压负荷MW_原表'], 3), x['负荷与容量原表格']),
        ('源荷比', n(x['源荷比_估计'], 5), '年末装机/最大用户负荷估计'),
        ('光伏年发电量 MWh', n(x['分布式光伏年发电量MWh_月装机曲线估计'], 0), '月末装机与项目光伏逐时曲线积分'),
        ('用户年用电量 MWh', n(x['用户年用电量MWh_借用负荷率估计'], 0), '用户最大负荷×8760×借用负荷率'),
        ('电量渗透率', n(float(x['电量渗透率_估计'])*100, 2)+'%', '发电量/用户用电量；借用负荷率'),
        ('单台主变负载率中位数', n(x['110kV主变负载率原表中位数百分数'], 2)+'%', x['负载率原表格']),
    ])
    s += '丰县年用电量缺同范围原始统计，本轮按邳州和市区估计负荷率均值 0.502489 计算约 2397002 MWh；光伏发电量用月末装机和项目光伏逐时曲线积分估计约 720020 MWh，故电量渗透率约 30.04%。该数只作描述性对照。\n\n'
    s += '## A.4 10 kV 转供、主变、储能和线路费用\n\n'
    s += table(['参数', '基准值及计算', '原始定位与证据性质'], [
        ('邳州样本可转移比例', '9.983428 / 35.292234 = 28.2879%', '六馈线及三处联络的静态组合上界；区县外推'),
        ('市区既有站间比例', '50%', 'DL/T 5729—2023 第 7.1.2 条两类分区建议范围的公共边界；研究假定'),
        ('等效新线容量', '7.491 MW/条', '邳州墩振线 23—24 杆区段重建上界；2025 静态筛查'),
        ('等效新线投资', '3×(0.9×24+0.1×120)+2×3=106.8 万元/条', '华容县规划报告表 7-1；3 km、90/10 混合和两开关为情景假定'),
        ('主变购置', '(1151+1335)/(50+100)=16.573333 万元/MVA', '徐州两项工程原表 Sheet1 第 21、36 行；含配套工程'),
        ('新设第三台主变', '(1206+1240)/2=1223 万元/50 MVA', '徐州两项工程静态总投资均值'),
        ('储能规格', '0.1 MW/0.215 MWh 每柜', '苏州采购规格；浏阳 1 MW/至少 2.15 MWh 成交价交叉核对'),
        ('储能 10 柜包', '210.456262 万元', '浏阳中标公告 2,104,562.62 元；按柜数线性外推'),
        ('主变/线路价格比', '(16.573333/0.95)/(106.8/7.491)=1.223643', '同有功 MW 的购置单价比较，不含站址费用'),
        ('主变/储能价格比', '(16.573333/0.95)/210.456262=0.082894', '储能按 1 MW/2.15 MWh 成套功率；持续时长改变有效能力'),
    ])
    s += '成本比例只反映价格锚点。设备的作用场景、持续能力和年度更新不同，应在完整路径目标函数中比较。评价期 2022—2041 年，折现率 6%；主变与线路寿命 20 年、储能 10 年；固定运维分别为投资额 1%、1%、3%，从投运次年计入。寿命和运维率均为统一研究情景假定。\n\n'
    s += '## A.5 缩写与单位\n\n'
    s += table(['缩写', '中文含义', '本报告用法'], [
        ('kV', '千伏', '电压等级'), ('MW', '兆瓦', '有功功率'), ('MVA', '兆伏安', '主变铭牌容量'),
        ('MWh', '兆瓦时', '储能额定能量及年度电量'), ('PV', '光伏', '分布式光伏出力或装机'),
        ('N−1', '单一元件停运情景', '本研究只作静态事故供电量筛查'),
        ('MILP', '混合整数线性规划', '离散设备与连续转供同时优化'),
        ('SHA-256', '文件摘要算法', '用于原始资料文件完整性登记')])
    save('附录A 原始资料与指标复算.md', s)


def appendix_b():
    s = '# 附录 B 两区县逐站年度容量与措施清单\n\n'
    s += '下列站号为项目脱敏编码。每一行均为该站该年的模型结果，主变购买量按新设备铭牌计算；站级正、反向压力为静态典型场景，不是全年同期实测峰。2021 年初态为统一反事实基准，表从 2022 年列起。\n\n'
    count = 0
    for region, folder in REGIONS:
        for label, scheme in SCHEMES:
            records = read_csv(OUT / folder / f'{scheme}_stations.csv')
            flow_records = read_csv(OUT / folder / f'{scheme}_outage_flows.csv')
            recovered = {}
            for flow in flow_records:
                key = (flow['failed_station'], flow['year'])
                recovered[key] = recovered.get(key, 0.0) + float(flow['recoverable_mw'])
            by_station = {}
            for r in records:
                by_station.setdefault(r['station'], []).append(r)
            s += f'## B.{1+REGIONS.index((region,folder))*2+SCHEMES.index((label,scheme))} {region}{label}方案\n\n'
            for station, sr in sorted(by_station.items()):
                sr.sort(key=lambda r: int(r['year']))
                s += f'### {station}（{sr[0]["area_class"]} 类研究分区）\n\n'
                s += table(['年', '在役主变 MVA', '本年购置 MVA', '在役储能 MW/MWh',
                            '本年新增储能 MW/MWh', '正向/反向压力 MW'], [
                    (r['year'], n(r['capacity_mva'],1), n(r['purchased_unit_mva'],1),
                     f'{n(r["storage_power_mw"],2)}/{n(r["storage_energy_mwh"],3)}',
                     f'{n(r["new_storage_power_mw"],2)}/{n(r["new_storage_energy_mwh"],3)}',
                     f'{n(r["forward_mw"],3)}/{n(r["reverse_mw"],3)}') for r in sr])
                s += table(['年', '主变一/二/三 MVA', '新设第三台', '储能柜数',
                            '购置主变投资 万元', '储能投资 万元', '事故静态接收量合计 MW'], [
                    (r['year'], '/'.join(n(r[f'unit_{k}_mva'],1) for k in (1,2,3)),
                     '是' if r['third_transformer_commissioned']=='1' else '否',
                     r['storage_modules'], n(r['transformer_capex_10k'],2),
                     n(r['storage_capex_10k'],2),
                     n(recovered.get((station,r['year']),0),3)) for r in sr])
                purchased = [r for r in sr if float(r['purchased_unit_mva'])>0]
                storage = [r for r in sr if int(r['new_storage_modules'])>0]
                s += ('本站购置主变年份：' + ('、'.join(r['year'] for r in purchased) if purchased else '无') +
                      '；新增储能年份：' + ('、'.join(r['year'] for r in storage) if storage else '无') +
                      '。事故接收量是各非零静态转移记录之和，按故障站汇总；不能直接视作实际停电后的转供电量。\n\n')
                count += len(sr)
    s += f'本附录共列 {count} 个站年记录。容量净增与购置主变铭牌量的差异来自旧设备替换；费用仍按购置量和对应年度折现因子计算。\n'
    save('附录B 逐站年度措施.md', s)


def appendix_c():
    s = '# 附录 C 条件参数矩阵与价格敏感性全表\n\n'
    s += '本矩阵以邳州 20 站、市区 29 站两种固定结构模板外推。增长率和源荷比均为假定值；源荷比通过同比例缩放反向功率代理进入模型，不能代替逐时光伏发电量。每格重新求四年路径。经济推荐仅指当前目标函数、价格和静态约束下费用较低的方案。\n\n'
    data = read_csv(OUT / 'generic_parameter_matrix_v2/conditional_recommendation_matrix.csv')
    for region, prefix in [('邳州模板', 'PZ-'), ('市区模板', 'CITY-')]:
        ss = [r for r in data if r['archetype'].startswith(prefix)]
        s += f'## C.{1 if prefix=="PZ-" else 2} {region}\n\n'
        s += table(['年增长', '源荷比', '刚性费用 万元', '弹性费用 万元', '推荐',
                    '刚性新线/储能 MWh', '弹性新线/储能 MWh'], [
            (f'{float(r["annual_load_growth_rate_assumed"])*100:.0f}%', n(r['source_load_ratio_assumed'],1),
             n(r['rigid_npv_10k'],2), n(r['elastic_npv_10k'],2),
             '弹性' if r['economically_preferred']=='elastic' else '刚性',
             f'{r["rigid_new_lines"]}/{n(r["rigid_new_storage_mwh"],3)}',
             f'{r["elastic_new_lines"]}/{n(r["elastic_new_storage_mwh"],3)}') for r in ss])
        s += table(['年增长', '源荷比', '刚性 2022/23/24/25 容载比', '弹性 2022/23/24/25 容载比'], [
            (f'{float(r["annual_load_growth_rate_assumed"])*100:.0f}%', n(r['source_load_ratio_assumed'],1),
             '/'.join(n(r[f'rigid_clr_{y}'],4) for y in range(2022,2026)),
             '/'.join(n(r[f'elastic_clr_{y}'],4) for y in range(2022,2026))) for r in ss])
    s += '18 组中只有邳州模板、零增长、源荷比 1.8 一格选择弹性。其余单元保持刚性，是当前参数和严格弹性比值顺序条件共同作用的结果。不能把本矩阵解释为导则上限或任意区县的无条件推荐。\n\n'
    s += '## C.3 两个实际价格比例的再优化结果\n\n'
    price = read_csv(OUT / 'ordered_guide_cost_sensitivity.csv')
    labels = {'baseline':'基准','line':'线路单价','storage':'储能单价'}
    s += table(['区县编码', '方案', '变动价格', '乘数', '主变/线价比', '主变/储能价比',
                '费用现值 万元', '新线条数', '新增储能柜'], [
        (r['region_id'], '刚性' if r['scheme']=='rigid' else '弹性',
         labels[r['changed_component']], n(r['price_scale'],1),
         n(r['transformer_to_line_price_ratio'],4),n(r['transformer_to_storage_price_ratio'],5),
         n(r['lifecycle_npv_10k'],2),r['new_lines'],r['new_storage_cabinets']) for r in price])
    save('附录C 条件矩阵与敏感性.md', s)


def appendix_d():
    s = '# 附录 D 新建等效联络与事故静态转供明细\n\n'
    s += '本附录保留求解模型的等效站对与故障站转移功率，便于复算第 5 章线路条数。站对是“全区县任意站可候选”的计算变量；3 km 为统一费用情景，不是测绘距离。事故转移功率是静态供电量模型结果，不等于实际操作记录或完整 N−1 校核。\n\n'
    for ri,(region,folder) in enumerate(REGIONS):
        for si,(label,scheme) in enumerate(SCHEMES):
            s += f'## D.{ri*2+si+1} {region}{label}\n\n'
            lines = read_csv(OUT / folder / f'{scheme}_new_lines.csv')
            s += f'共投运 {len(lines)} 条等效新线。\n\n'
            s += table(['投运年', '站对 A', '站对 B', '统一长度 km', '单条容量 MW', '初始投资 万元'], [
                (r['commissioning_year'],r['station_a'],r['station_b'],
                 n(r['planned_length_km'],1), n(r['screen_capacity_mw_2025'],3),
                 n(r['construction_capex_10k'],1)) for r in lines])
            flows = read_csv(OUT / folder / f'{scheme}_outage_flows.csv')
            for year in range(2022,2026):
                yr = [r for r in flows if int(r['year'])==year]
                s += f'### {year} 年事故静态转移\n\n'
                s += table(['故障站', '接收站', '通道性质', '可恢复功率 MW'], [
                    (r['failed_station'],r['receiver_station'],
                     '既有能力' if r['kind']=='existing_county_proxy' else '新建等效线',
                     n(r['recoverable_mw'],3)) for r in yr])
                s += f'本年列示 {len(yr)} 条非零静态转移记录；同一故障站的接收功率按各接收站记录求和。\n\n'
    save('附录D 联络与转供明细.md', s)


if __name__ == '__main__':
    appendix_a(); appendix_b(); appendix_c(); appendix_d()
    print('已生成 A—D 四项附录')
