"""核验原件未变、来源图、逐项取值、工作簿和PDF；只写本阶段核查证据。"""
from pathlib import Path
import csv,hashlib,json,math,re,zipfile
import openpyxl,pymupdf,nbformat

ROOT=Path(__file__).resolve().parents[3];OUT=ROOT/'研究报告/00审查';E=OUT/'证据'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p):return json.loads(p.read_text())
def dump(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
def same(a,b):
    if a is None or b is None:return a is None and b is None
    try:return math.isclose(float(a),float(b),rel_tol=1e-12,abs_tol=1e-9)
    except (TypeError,ValueError):return a==b

def main():
    checks=0
    def check(ok,context):
        nonlocal checks
        checks+=1
        if not ok:raise AssertionError(context)
    snapshot=read(E/'输入与原结果快照.json')
    for r in snapshot:check(sha(ROOT/r['path'])==r['sha256'],r['path'])
    original=read(ROOT/'docs/FREEZE-MANIFEST-2026-09-27.json')
    for r in original['files']:check(sha(ROOT/r['path'])==r['sha256'],r['path'])
    data=read(OUT/'data_source.json');records=data['records'];byid={r['id']:r for r in records}
    check(len(byid)==len(records),'重复数据编号')
    nodes={n['path']:n for n in data['source_graph']}
    graph=read(E/'数据来源图.json')['nodes'];check(graph==data['source_graph'],'来源图与JSON不一致')
    for n in nodes.values():
        check(sha(ROOT/n['path'])==n['sha256'],n['path'])
        for p in n['parents']:check(p in nodes,'未注册父节点 '+p)
    colors={}
    def visit(p):
        check(colors.get(p)!=1,'来源图循环 '+p)
        if colors.get(p)==2:return
        colors[p]=1
        for parent in nodes[p]['parents']:visit(parent)
        colors[p]=2
    for p in nodes:visit(p)
    caches={};workbooks={}
    for r in records:
        for k in ['name','variable','unit','source','source_location','formula','chapter']:
            check(bool(r.get(k)),r['id']+':'+k)
        check(r['source_sha256']==nodes[r['source']]['sha256'],r['id']+'来源哈希')
        check('QX-' not in r['name'],r['id']+'数据名称地区编码')
        check(r['value'] is None or not isinstance(r['value'],(int,float)) or math.isfinite(r['value']),r['id']+'非有限值')
        src=ROOT/r['source'];loc=r['source_location']
        match=re.fullmatch(r'CSV物理行(\d+)；字段(.+)',loc)
        if match:
            if src not in caches:
                with src.open(encoding='utf-8-sig',newline='') as f:caches[src]=list(csv.DictReader(f))
            raw=caches[src][int(match[1])-2][match[2]]
            check(raw==r['raw_value'] if isinstance(r['raw_value'],str) else same(raw,r['raw_value']),r['id']+'原值不同')
            check(same(None if raw in ('','None') else raw,r['value']),r['id']+'CSV值不同')
        elif loc.startswith('JSON $.'):
            if src not in caches:caches[src]=read(src)
            m=re.fullmatch(r'JSON \$\.(\w+)\[(\d+)\]\.(\w+)',loc);check(bool(m),r['id']+'JSON坐标')
            check(same(caches[src][m[1]][int(m[2])][m[3]],r['value']),r['id']+'JSON值不同')
        elif src.suffix=='.xlsx' and re.fullmatch(r'Sheet1![A-Z]+\d+',loc):
            if src not in workbooks:workbooks[src]=openpyxl.load_workbook(src,read_only=True,data_only=True)
            raw=workbooks[src]['Sheet1'][loc.split('!')[1]].value
            check(same(raw,r['raw_value']),r['id']+'Excel原值不同')
            factor=10 if r['unit'] in ('MW','MVA') else 1
            check(same(float(raw)*factor,r['value']),r['id']+'Excel换算不同')
    for w in workbooks.values():w.close()
    workbook=openpyxl.load_workbook(OUT/'最终数据字典.xlsx',read_only=True,data_only=True)
    groups=[('报告准入数据','admitted_record_ids'),('技术附录明细','technical_appendix_record_ids'),('历史汇总限制引用','historical_evidence_record_ids')]
    allseen=set()
    for sheet,key in groups:
        rows=list(workbook[sheet].iter_rows(min_row=2,values_only=True))
        check(len(rows)==len(data[key]),sheet+'行数')
        check([r[0] for r in rows]==data[key],sheet+'编号顺序')
        for row in rows:
            r=byid[row[0]];allseen.add(row[0]);check(same(row[3],r['value']),r['id']+'字典数值')
            check(row[5]==r['source'] and row[6]==r['source_location'] and row[7]==r['formula'] and row[8]==r['chapter'],r['id']+'字典来源或章节')
    check(allseen==set(byid),'字典遗漏登记记录');workbook.close()
    document=pymupdf.open(OUT/'数据来源审查.pdf');text=''.join(p.get_text() for p in document)
    for rid in data['admitted_record_ids']:check(rid in text,rid+'PDF遗漏或无法检索')
    check('QX-' not in text,'PDF出现地区编码')
    check('\ufffd' not in text,'PDF替换字符')
    pages=len(document)
    bounds=[]
    for i,p in enumerate(document):
        for b in p.get_text('blocks'):
            if b[0]<0 or b[1]<0 or b[2]>p.rect.width+.1 or b[3]>p.rect.height+.1:bounds.append(i+1)
    check(not bounds,'PDF文字越界 '+str(bounds))
    document.close()
    for name in ['data_source.json','数据来源审查.pdf']:check(sha(OUT/name)==sha(ROOT/'研究报告/数据来源'/name),'来源目录副本不同')
    nb=nbformat.read(OUT/'审计工具/阶段0只读核查.ipynb',as_version=4);nbformat.validate(nb)
    check(all(c.execution_count and all(o.output_type!='error' for o in c.outputs) for c in nb.cells if c.cell_type=='code'),'笔记本未执行或报错')
    cleanup=read(E/'废弃结果清理.json');archive=ROOT/cleanup['archive']
    check(sha(archive)==cleanup['archive_sha256'],'归档哈希不同')
    with zipfile.ZipFile(archive) as z:
        check(z.testzip() is None,'归档CRC失败')
        for r in cleanup['entries']:check(hashlib.sha256(z.read(r['path'])).hexdigest()==r['sha256'],r['path']+'归档原件不同')
    for folder in cleanup['removed_directories']:check(not (ROOT/folder).exists(),'未清理 '+folder)
    result={'status':'PASS','checks':checks,'records':len(records),'original_snapshot_files':len(snapshot),'original_manifest_files':len(original['files']),'pdf_pages':pages,'pdf_admitted_ids_complete':True,'source_graph_acyclic':True,'archived_files':cleanup['archived_files'],'scope':'原件哈希、取值定位、文件血缘、字典与JSON一致性、PDF检索与越界、归档完整性；不能代替工程可实施性核查'}
    dump(E/'交付核验.json',result);print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':main()
