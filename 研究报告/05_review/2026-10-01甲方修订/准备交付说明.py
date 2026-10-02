from pathlib import Path
import json,re,hashlib
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
HERE=Path(__file__).parent;R=HERE.parents[1];ROOT=R.parent;D=R/'01研究推荐/2026-10-01甲方修订';A=R/'02图表/2026-10-01甲方修订';S=R/'数据来源/2026-10-01甲方修订';S.mkdir(exist_ok=True)
T=json.loads((A/'计算表.json').read_text());rows=[]
for line in T['monthly_indicators'].splitlines():
 cols=[t.strip() for t in line.strip('|').split('|')]
 if cols[0].isdigit():rows.append([float(x) for x in cols])
font_manager.fontManager.addfont('/usr/share/fonts/wps-office/FZFSK.TTF');font_manager.fontManager.addfont('/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf')
plt.rcParams.update({'font.family':['Times New Roman',font_manager.FontProperties(fname='/usr/share/fonts/wps-office/FZFSK.TTF').get_name()],'font.size':11,'axes.unicode_minus':False})
fig,axs=plt.subplots(1,2,figsize=(7,3.7))
for ax,j,label in [(axs[0],2,'容载比（月度分析）'),(axs[1],3,'主变负载率估计 / %')]:
 ax.scatter([r[1] for r in rows],[r[j] for r in rows],color='#29465b');ax.set_xlabel('月度分布式电源电量渗透率 / %');ax.set_ylabel(label)
 for r in rows:ax.annotate(str(int(r[0])),(r[1],r[j]),fontsize=9,xytext=(3,2),textcoords='offset points')
plt.tight_layout();fig.savefig(A/'monthly_correlation.png',dpi=300,bbox_inches='tight');plt.close(fig)
mapping=dict(zip(['year','station','area_class','source_area_class','area_class_status','source_available_third_slots','source_spare_10kv_bays','new_third_without_reserved_slot','prior_unit_1_mva','prior_unit_2_mva','unit_1_mva','unit_2_mva','unit_3_mva','third_transformer_commissioned','capacity_mva','purchased_unit_mva','transformer_capex_10k','storage_power_mw','storage_energy_mwh','storage_modules','new_storage_power_mw','new_storage_energy_mwh','new_storage_modules','storage_capex_10k','forward_mw','reverse_mw','forward_duration_h','reverse_duration_h','normal_load_transferred_out_mw','normal_load_transferred_in_mw','post_transfer_forward_mw','post_transfer_reverse_upper_mw','transfer_purpose','transfer_base_load_mw','existing_transfer_fraction_scenario','target_transfer_fraction_scenario','transfer_fraction_ceiling','incident_new_line_units','existing_transfer_capacity_mw','new_line_fraction_increment','effective_transfer_fraction','effective_transfer_capacity_mw','credited_new_transfer_capacity_mw','actual_outgoing_fraction','transfer_base_status'],['仿真年','站号','规划供电区域类别','样本供电区域类别','类别设定方式','第三台预留位置数','备用10kV间隔数','无预留新增第三台标记','前期1号主变MVA','前期2号主变MVA','1号主变MVA','2号主变MVA','3号主变MVA','新增第三台状态','主变总容量MVA','主变采购容量MVA','主变初始投资万元','在役储能MW','在役储能MWh','在役储能柜数','新增储能MW','新增储能MWh','新增储能柜数','储能初始投资万元','最大正向净负荷MW','最大反向净负荷MW','正向95%峰值持续时间h','反向95%峰值持续时间h','转出负荷MW','承接负荷MW','转接后正向负荷MW','反向净负荷校核上界MW','转接用途','转供能力基准负荷MW','初始站间转供率','规划目标转供率','转供率研究上限','关联新增项目端数','既有转供能力MW','新增项目提升比例','有效转供率','有效转供能力MW','实际使用新增能力MW','实际转出比例','转供能力基准说明']))
wb=load_workbook(D/'统一起点逐年方案与复核.xlsx')
if '字段说明' in wb:del wb['字段说明']
ws=wb.create_sheet('字段说明');ws.append(['计算字段','报告参数名称']);[ws.append([a,b]) for a,b in mapping.items()]
for sh in wb.worksheets[:4]:
 assert len(mapping)==sh.max_column
 if sh.cell(1,1).value=='year':
  for cell in sh[1]:cell.value=mapping[cell.value]
  for i in range(2,sh.max_row+1):sh.cell(i,1,int(sh.cell(i,1).value)-2021)
for w in [wb,load_workbook(D/'通用容载比推荐矩阵.xlsx')]:
 for sh in w:
  for cell in sh[1]:cell.font=Font(name='宋体',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='29465B');cell.alignment=Alignment(wrap_text=True,vertical='center')
  sh.row_dimensions[1].height=32;sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
  for col in sh.columns:
   if sh.title!='使用说明':sh.column_dimensions[col[0].column_letter].width=22
  for row in sh.iter_rows(min_row=2):
   for cell in row:cell.font=Font(name='宋体',size=10);cell.alignment=Alignment(vertical='center')
 path=D/('统一起点逐年方案与复核.xlsx' if w is wb else '通用容载比推荐矩阵.xlsx');w.save(path)
register=json.loads((R/'数据来源/2026-10-01本轮来源说明/本轮原始文件登记.json').read_text())
(S/'原始文件登记.json').write_text(json.dumps(register,ensure_ascii=False,indent=2))
data='''# 数据来源与参数说明

本说明对应2026年10月甲方修订稿，研究对象为邳州、市区110 kV公用变电设备。历史统计、统一起点四年仿真和一般站群参数外推分别列示。

## 1 原始资料及使用方法

年度容量和降压负荷采用《近5年容载比.xlsx》，按区县与110 kV筛选，统一换算为MVA、MW。降压负荷用于估计年网供最大负荷，2021—2025年同比率按1∶2∶3∶4权重估计年负荷平均增长率，邳州8.0742%、市区1.8000%。

设备规格、供电类别、第三台位置及年最大负载率来自《2025设备负载统计表.xlsx》和设备明细。邳州20站、市区29站构成仿真结构，市区含已映射与补充站，规划供电类别按A类设定。站级净负荷采用邳州主变时序及市区脱敏资料，最大正向、反向净负荷和95%峰值持续时间分别提取；缺测按报告所列方法处理。

月末光伏装机采用《逐月分县分布式光伏.xlsx》。原表未注明数值单位，研究采用原值乘10作为MW的单位假定，2021—2022年按同月份几何关系回推，2023—2025年使用原记录。年度光伏电量按相邻月末平均装机和8760小时典型出力系数积分；用户最大负荷、电量按报告中的样本覆盖与本地发电关系估计。

主变费用采用徐州工程建设规模及投资汇总表：邢楼、墩集第三台项目位于第14、16行，王庄、鲁庙更换项目位于第21、36行。替换系数16.573333万元/MVA、第三台50 MVA费用1223万元/项。线路采用公开配网规划参考造价与3 km、90%架空/10%电缆、两台开关假定，形成106.8万元/项目。储能采用浏阳1 MW/2.15 MWh成套成交价折算21.0456262万元/柜，苏州单柜最高限价作为对照。

## 2 统一起点四年规划

采用模型离散设备初态：邳州主变1463.5 MVA、市区2462 MVA。基期年网供最大负荷分别标定为731.75 MW、1231 MW，容载比均为2.0；历史实际容量在正文第三章独立列示。后续四个仿真年按年负荷平均增长率推演，站级正向、反向负荷按2025年分布同步缩放。计算文件中的2022—2025年为仿真第1—4年现金流时点标签。

市区初始站间转供率假定50%；邳州六馈线及跨站联络样本最大可转负荷9.983428 MW、源端负荷35.292234 MW，测算28.2879%，统一应用于县域各站。既有能力零新增建设投资，新建项目提供7.491061 MW能力增量，按双端关联、共享新增能力预算和106.8万元/项目计费。两区县内任意不同站对均可成为等效新建项目候选。

正常正向、反向及单台主变退出容量逐站校核，退出后满足全部转接负荷。设备状态逐年非减，区县转出与承接守恒，电源保留原站。刚性容载比不超过2.0，弹性2.001～2.6，地区和方案比较条件见第五章。费用、容量偏好、固定布局转接量依次求解，费用阶段间隙0.99995%、容量阶段0.49944%。392个站年经独立复算通过。

## 3 通用矩阵外推

以A类/B、C类两种五站结构为一般模板，每站两台50/40 MVA，基期容载比2.0。装机渗透率0～3、年负荷平均增长率0～10%、95%峰值持续时间1～6 h，两种负荷分布、五组交叉价格形成7000个目标年组合。目标年采用弹性2.001～2.6范围，逐点静态寻优和35000条站级容量、守恒、费用独立复算通过。

电量渗透率按装机渗透率×1200/(8760×0.5)换算；光伏日间系数0.75、用户日间负荷系数0.30。主变与线路成本比0.97～1.64、主变与储能成本比0.047～0.119，后者按3 h参考尺度列示。矩阵按已计算点包络向外取整0.05；增长率第二、三档为左开右闭。最终推荐表仅列参数及容载比，计算点工作表保存具体措施。

## 4 费用与规范依据

正式规范采用DL/T 5729—2023、DL/T 2041—2025；项目留存2025年规划导则修订稿用于渗透率指标的研究定义。新增措施现金流以2021年末为基准，评价截止2041年，折现率6%；主变与线路寿命20年、运维1%，储能寿命10年、运维3%，均为统一研究参数。目标采用新增建设、固定运维与寿命更新支出，无期末残值及处置收益。

## 5 复算文件

配套《统一起点逐年方案与复核.xlsx》保存四组逐站结果、年度区域方案及复核；《通用容载比推荐矩阵.xlsx》保存推荐区间、7000个计算点和使用条件。完整JSON、CSV、站级外推记录及求解日志在研究报告/01研究推荐/2026-10-01甲方修订中保留。原始文件路径、字节数及SHA256如下。

'''
for record in register:data+=f'- {record["文件"]}；{record["字节数"]}字节；SHA256：{record["SHA256"]}\n\n'
(S/'数据来源与参数说明.md').write_text(data)
print('图表、配套工作簿及数据说明已整理')
