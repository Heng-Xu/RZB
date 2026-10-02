from pathlib import Path
import json,re
from openpyxl import load_workbook
HERE=Path(__file__).parent;R=HERE.parents[1];D=R/'01研究推荐/2026-10-01甲方修订';A=R/'02图表/2026-10-01甲方修订';TEXT=R/'03_MD/2026-10-01甲方修订'
names=[('容量／月正向峰值比','容载比（月度分析）'),('容量／月正向峰','容载比（月度分析）'),('绝对有功压力比','主变负载率估计值'),('绝对有功压力/%','主变负载率估计/%'),('绝对有功压力 / %','主变负载率估计 / %'),('绝对有功压力','主变负载率估计值')]
for p in sorted(TEXT.glob('0[1-6] *.md')):
 s=p.read_text()
 for old,new in names:s=s.replace(old,new)
 if p.name.startswith('03'):
  s=s.replace('相关系数采用12个月的样本计算','主变负载率按净负荷最大绝对值除以0.95倍同范围主变容量估计，0.95为功率因数假定。相关系数采用12个月的样本计算')
 if p.name.startswith('05'):
  s=s.replace('已有设备的铭牌规格和第三台预留位置进入候选集合','样本设备规格、供电区域类别和第三台预留位置进入候选集合')
  key='## 5.3 约束条件与求解流程'
  s=s.replace(key,'以两站负荷转接说明计算关系。送端原正向负荷80 MW，初始转供率28.2879%，既有能力为22.6303 MW；新建一个项目后，能力增至30.1214 MW。若转出25 MW，新增能力实际使用2.3697 MW，原站正向负荷降至55 MW；原反向净负荷10 MW时，校核上界增至35 MW。受端原负荷40 MW，承接后为65 MW，按65 MW校核正常和单台退出容量。两站负荷合计仍为120 MW，新建项目投资106.8万元计取一次。\n\n'+key)
  s=s.replace('表中仅列参数与推荐容载比。','表中仅列参数与推荐容载比。增长率采用0%～2%、大于2%至7%、大于7%至10%三个分档；区间下限2.00为0.05步长向外取整的显示值，逐点弹性计算执行2.001下限。')
 p.write_text(s)
matrix=json.loads((D/'通用推荐矩阵.json').read_text())
for row in matrix:
 row['年负荷平均增长率']=row['年负荷平均增长率'].replace('2%～7%','＞2%～7%').replace('7%～10%','＞7%～10%')
(D/'通用推荐矩阵.json').write_text(json.dumps(matrix,ensure_ascii=False,indent=2))
tables=json.loads((A/'计算表.json').read_text())
for k,s in tables.items():
 for old,new in names:s=s.replace(old,new)
 s=s.replace('原表容载比','容载比')
 if k=='general_matrix':s=s.replace('2%～7%','＞2%～7%').replace('7%～10%','＞7%～10%').replace('| 装机渗透率 |','| 分布式电源装机渗透率 |').replace('主变/线路成本比','主变与线路成本比').replace('主变/储能成本比','主变与储能成本比')
 tables[k]=s
(A/'计算表.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2))
wb=load_workbook(D/'通用容载比推荐矩阵.xlsx');ws=wb['通用推荐区间']
for i,row in enumerate(matrix,2):ws.cell(i,2,row['年负荷平均增长率'])
if '使用说明' in wb:del wb['使用说明']
ws=wb.create_sheet('使用说明');ws.append(['项目','条件'])
for row in [['推荐性质','第4年目标年的静态参数外推；用于前期比选范围'],['统计对象','A类、B/C类五站结构；各站初始两台50/40 MVA'],['初始容载比','2.0；目标年弹性2.001～2.6'],['初始转供率','50%或28.2879%'],['负荷与光伏空间分布','两种负荷权重；独立光伏权重'],['源荷构造','日间光伏系数0.75、用户负荷系数0.30'],['价格范围','线路0.75～1.25倍；储能0.50～1.25倍；四角与中心'],['电量渗透率换算','装机渗透率×1200/(8760×0.5)'],['区间端点','增长率第二、三档为左开右闭；装机及持续时间共享边界'],['显示步长','0.05向外取整；2.00显示下限对应逐点2.001约束'],['技术约束','正常、全负荷单台退出、反向容量、转供与守恒'],['复核','7000点、35000条站级记录通过；费用与容量独立复算']]:ws.append(row)
ws.column_dimensions['A'].width=24;ws.column_dimensions['B'].width=80;ws.freeze_panes='A2';wb.save(D/'通用容载比推荐矩阵.xlsx')
wb=load_workbook(D/'统一起点逐年方案与复核.xlsx')
for sheet,head in [('逐年区域方案',['地区','方案','仿真年','年网供最大负荷MW','主变容量MVA','净增MVA','采购MVA','新增储能MW/MWh','新建联络项','转接MW','容载比']),('容量费用复核',['地区','方案','站年','最小容量余量MW','费用偏差万元','状态'])]:
 ws=wb[sheet];ws.insert_rows(1);[ws.cell(1,j,v) for j,v in enumerate(head,1)];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
wb.save(D/'统一起点逐年方案与复核.xlsx')
# 保存下一次生成时的相同术语与边界。
for fn in ['准备现状图表.py','更新方案图表.py']:
 p=HERE/fn;s=p.read_text()
 for old,new in names:s=s.replace(old,new)
 p.write_text(s)
p=HERE/'通用矩阵外推.py';s=p.read_text().replace("f'{glo*100:g}%～{ghi*100:g}%'","(('＞' if glo>0 else '')+f'{glo*100:g}%～{ghi*100:g}%')");p.write_text(s)
print('边界标签、术语及工作簿说明已同步')
