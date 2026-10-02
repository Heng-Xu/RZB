from pathlib import Path
import json,math,re
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
HERE=Path(__file__).parent;R=HERE.parents[1];D=R/'01研究推荐/2026-10-01甲方修订';A=R/'02图表/2026-10-01甲方修订';TEXT=R/'03_MD/2026-10-01甲方修订'
points=[{'地区':'邳州','装机渗透率估计':1.63065,'年负荷平均增长率':.080742,'95%峰值持续时间上限h':4,'刚性容载比':1.9574,'弹性容载比':2.2118},{'地区':'市区','装机渗透率估计':.31788,'年负荷平均增长率':.018,'95%峰值持续时间上限h':5,'刚性容载比':1.9523,'弹性容载比':2.2102}]
m=json.loads((D/'通用推荐矩阵.json').read_text());checks=[]
for p in points:
 matched=[]
 for i,row in enumerate(m):
  el,eh=map(float,row['分布式电源装机渗透率'].split('～'));gl,gh=map(float,re.findall(r'[\d.]+',row['年负荷平均增长率']));hl,hh=map(float,re.findall(r'[\d.]+',row['95%峰值持续时间']))
  if el<=p['装机渗透率估计']<=eh and (gl<p['年负荷平均增长率']*100<=gh or gl==0 and p['年负荷平均增长率']==0) and hl<=p['95%峰值持续时间上限h']<=hh:
   lower,upper=map(float,row['推荐容载比区间'].split('～'));value=p['弹性容载比'];lower=min(lower,math.floor(value/.05)*.05);upper=max(upper,math.ceil(value/.05)*.05)
   row['推荐容载比区间']=f'{lower:.2f}～{upper:.2f}';matched.append(i+1)
 assert matched
 checks.append({**p,'匹配矩阵行':matched,'状态':'通过'})
(D/'通用推荐矩阵.json').write_text(json.dumps(m,ensure_ascii=False,indent=2));(HERE/'案例与矩阵衔接核对.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
t=json.loads((A/'计算表.json').read_text());headers=['分布式电源装机渗透率','年负荷平均增长率','95%峰值持续时间上限','主变与线路成本比','主变与储能成本比','刚性上限','推荐弹性容载比']
t['general_matrix']='表5-0 刚性基准与通用弹性容载比推荐区间\n\n| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*7)+' |\n'+''.join('| '+' | '.join(r[k] for k in ['分布式电源装机渗透率','年负荷平均增长率','95%峰值持续时间','主变与线路成本比','主变与储能成本比_3h','刚性容载比控制上限','推荐容载比区间'])+' |\n' for r in m)
t['region_indicator_comparison']=t['region_indicator_comparison'].replace('原表容载比','容载比')
(A/'计算表.json').write_text(json.dumps(t,ensure_ascii=False,indent=2))
p=TEXT/'05 第五章 弹性容载比寻优与规划建议.md';s=p.read_text();s=s.replace('汇总同档各计算点的推荐容载比最小值和最大值，再按0.05步长向外取整。','汇总同档一般站群计算点及两地区案例目标年校核点的容载比最小值和最大值，再按0.05步长向外取整。案例的源荷参数采用第三章估计值，年负荷平均增长率采用案例条件，峰段按站群95%峰值持续时间上限匹配分档。案例校核补充实际异容量设备结构对区间的影响。');s=s.replace('初始转供率28.2879%～50%、两类供电区域及两种负荷分布共同构成该表的适用条件。','初始转供率28.2879%～50%、两类一般供电站群及案例设备结构、两种负荷分布共同构成该表的适用条件。');p.write_text(s)
p=TEXT/'06 第六章 结论与建议.md';s=p.read_text().replace('已计算容载比的最小值和最大值','一般站群计算点及案例目标年校核点的容载比最小值和最大值');p.write_text(s)
wb=load_workbook(D/'通用容载比推荐矩阵.xlsx');ws=wb['通用推荐区间'];wb.remove(ws);ws=wb.create_sheet('通用推荐区间',0);ws.append(headers+['一般站群计算点数'])
for row in m:ws.append([row[k] for k in ['分布式电源装机渗透率','年负荷平均增长率','95%峰值持续时间','主变与线路成本比','主变与储能成本比_3h','刚性容载比控制上限','推荐容载比区间','计算点数']])
if '案例校核点' in wb:del wb['案例校核点']
ws=wb.create_sheet('案例校核点');ws.append(list(points[0]));[ws.append(list(p.values())) for p in points]
ws=wb['使用说明'] if '使用说明' in wb else wb.create_sheet('使用说明')
if ws.max_row==1 and ws.cell(1,1).value is None:
 ws.append(['项目','条件'])
 for row in [['规划范围','两类五站一般结构及案例设备结构；目标年静态外推'],['初始容载比','2.0；弹性2.001～2.6；刚性上限2.0'],['初始转供率','50%或28.2879%'],['源荷构造','日间光伏系数0.75，用户负荷系数0.30'],['价格倍数','线路0.75～1.25；储能0.50～1.25；四角与中心'],['电量换算','装机渗透率×1200/(8760×0.5)'],['区间端点','增长率第二、三档左开右闭；0.05向外取整'],['计算和校核','7000一般站群点，35000条站级复算；两地区案例392站年']]:ws.append(row)
ws.append(['站级峰值合计/同期峰值','偏集中1.15、较均匀1.00；案例约1.057、1.113']);ws.append(['案例校核','两地区第4年弹性容载比纳入相应参数档；峰段按站群上限匹配']);ws.append(['刚性对照','控制上限2.0，案例另列逐年刚弹对照及全寿命费用'])
for sh in wb:
 sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
 for cell in sh[1]:cell.font=Font(name='宋体',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='29465B');cell.alignment=Alignment(wrap_text=True,vertical='center')
 for col in sh.columns:
  if sh.title!='使用说明':sh.column_dimensions[col[0].column_letter].width=24
wb.save(D/'通用容载比推荐矩阵.xlsx')
p=R/'数据来源/2026-10-01甲方修订/数据来源与参数说明.md';s=p.read_text().replace('两种负荷分布、五组交叉价格形成7000个目标年组合。','两种负荷分布、五组交叉价格形成7000个目标年组合。站级峰值合计/区域同期峰值按偏集中1.15、较均匀1.00设定，案例分别约1.057和1.113。').replace('最终推荐表仅列参数及容载比','矩阵包络另纳入两地区案例第4年校核点，按站群峰段上限匹配分档。最终推荐表仅列参数、刚性控制上限及推荐弹性容载比');p.write_text(s)
print('两地区目标年案例纳入矩阵校核，工作簿与正文已同步')
