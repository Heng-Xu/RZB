import csv,json,sys,re,importlib.util
from pathlib import Path
from collections import defaultdict
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook
R=Path(__file__).resolve().parents[2];ROOT=R.parent;HERE=Path(__file__).parent
A=R/'02图表/2026-10-01甲方修订';D=R/'01研究推荐/2026-10-01甲方修订';BASE=D/'统一起点基准'
font_manager.fontManager.addfont('/usr/share/fonts/wps-office/FZFSK.TTF')
for f in ['Times_New_Roman.ttf','Times_New_Roman_Bold.ttf']:
 font_manager.fontManager.addfont('/usr/share/fonts/truetype/msttcorefonts/'+f)
F=font_manager.FontProperties(fname='/usr/share/fonts/wps-office/FZFSK.TTF').get_name()
plt.rcParams.update({'font.family':['Times New Roman',F],'font.size':11,'axes.unicode_minus':False,'figure.dpi':180})
C=['#29465b','#9b4b39','#778f68']
T=json.loads((A/'计算表.json').read_text());old=json.loads((R/'02图表/2026-10-01重构/计算表.json').read_text())
for k in ['storage_cashflow_example','project_npv_examples']:T[k]=old[k]
SECTION={};AUDIT=[]
def rows(p):
 with Path(p).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def num(v,n=3):return f'{float(v):.{n}f}'
def table(key,chapter,title,headers,data):
 T[key]=f'表{chapter}-0 '+title+'\n\n| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(map(str,row))+' |\n' for row in data)
def save(name):
 plt.tight_layout();plt.savefig(A/(name+'.png'),dpi=300,bbox_inches='tight');plt.close()
def cash(y,life,om):
 return np.array([1/1.06**(y-2021),sum(om/1.06**(t-2021) for t in range(y+1,2042)),sum(1/1.06**(t-2021) for t in range(y+life,2041,life))])
REG={'pizhou':'邳州','city':'市区'};SC={'rigid':'刚性','elastic':'弹性'}
def main():
 review=json.loads((BASE/'review.json').read_text())
 portfolio={};components=[];comparison=[];measures=[];transfers=[];base=[];resultrows=[]
 wb=Workbook();wb.remove(wb.active)
 for c,name in REG.items():
  summaries={r['scheme']:r for r in json.loads((BASE/c/'summary.json').read_text())}
  for s,slabel in SC.items():
   summ=summaries[s];yrs=rows(BASE/c/(s+'_years.csv'));st=rows(BASE/c/(s+'_stations.csv'));moves=rows(BASE/c/(s+'_transfers.csv'))
   portfolio[c,s]=(summ,yrs,st)
   assert abs(summ['baseline_2021_clr']-2)<1e-8
   before=summ['baseline_2021_capacity_mva'];annual=[]
   component=np.zeros((3,3));minmargin=1e9
   for yr in yrs:
    y=int(yr['year']);a=[r for r in st if int(r['year'])==y];v=lambda k:float(yr[k])
    assert abs(sum(float(r['capacity_mva']) for r in a)-v('capacity_mva'))<1e-4
    assert abs(v('clr')-v('capacity_mva')/v('net_peak_proxy_mw'))<1e-7
    if s=='rigid':assert v('clr')<=2+1e-6
    else:assert 2.001-1e-6<=v('clr')<=2.6+1e-6
    annual.append([y-2021,num(v('net_peak_proxy_mw')),num(v('capacity_mva'),1),num(v('capacity_mva')-before,1),num(v('new_transformer_purchase_mva'),1),num(v('new_storage_power_mw'),1)+'/'+num(v('new_storage_energy_mwh')),int(yr['new_lines_commissioned']),num(v('existing_transfer_mw')),num(v('clr'),4)])
    before=v('capacity_mva')
    out=sum(float(r['normal_load_transferred_out_mw']) for r in a);inc=sum(float(r['normal_load_transferred_in_mw']) for r in a)
    assert abs(out-inc)<1e-4
    for j,(typ,life,om) in enumerate([('transformer',20,.01),('storage',10,.03),('line',20,.01)]):component[j]+=v(typ+'_capex_10k')*cash(y,life,om)
    existing_use=0;new_use=0
    for r in a:
     q=float(r['normal_load_transferred_out_mw']);existing=float(r['existing_transfer_capacity_mw'])
     existing_use+=min(q,existing);new_use+=max(0,q-existing)
     u=[float(r[k]) for k in ['unit_1_mva','unit_2_mva','unit_3_mva']]
     e=float(r['storage_energy_mwh']);fp=float(r['post_transfer_forward_mw']);rp=float(r['post_transfer_reverse_upper_mw'])
     b=e/max(2.15,float(r['forward_duration_h']));br=e/max(2.15,float(r['reverse_duration_h']))
     margins=[.95*sum(u)+b-fp,.95*sum(u)+br-rp,.95*(sum(u)-max(u))+b-fp]
     assert min(margins)>=-1e-4
     minmargin=min(minmargin,min(margins))
     assert q<=float(r['effective_transfer_capacity_mw'])+1e-4
     assert all(new+1e-5>=old for new,old in zip(u[:2],[float(r['prior_unit_1_mva']),float(r['prior_unit_2_mva'])]))
     assert abs(e-.215*int(r['storage_modules']))<1e-4
    assert new_use<=int(yr['new_lines_in_service'])*7.49106117+1e-4
    transfers.append([name,slabel,y-2021,num(existing_use),num(new_use),int(yr['new_lines_in_service']),num(out)])
   assert abs(component.sum()-summ['objective_npv_10k'])<.002
   AUDIT.append({'地区':name,'方案':slabel,'站年':len(st),'最小容量余量MW':minmargin,'费用偏差万元':float(component.sum()-summ['objective_npv_10k']),'状态':'通过'})
   table(c+'_'+s+'_years',5,name+slabel+'方案逐年措施及容载比',['仿真年','年网供最大负荷/MW','主变容量/MVA','净增/MVA','采购/MVA','新增储能MW/MWh','新建联络/项','转接/MW','容载比'],annual)
   componentrows=[name,slabel,*[num(v,2) for v in component.sum(axis=1)],num(component.sum(),2)]
   components.append(componentrows)
   last=yrs[-1]
   measures.append([name,slabel,num(sum(float(r['new_transformer_purchase_mva']) for r in yrs),1),sum(int(r['new_third_transformers']) for r in yrs),num(last['installed_storage_power_mw'],1)+'/'+num(last['installed_storage_energy_mwh']),int(last['new_lines_in_service'])])
   resultrows.extend([name,slabel,*r] for r in annual)
   ws=wb.create_sheet(name+slabel+'逐站');ws.append(list(st[0]));[ws.append(list(r.values())) for r in st];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  cr=summaries['rigid']['objective_npv_10k'];ce=summaries['elastic']['objective_npv_10k']
  comparison.append([name,num(cr,2),num(ce,2),num(cr-ce,2),num((1-ce/cr)*100,2)+'%'])
  rvals='、'.join(num(r['clr'],4) for r in portfolio[c,'rigid'][1]);evals='、'.join(num(r['clr'],4) for r in portfolio[c,'elastic'][1])
  SECTION[c+'_results']=f'{name}刚性方案第1—4年容载比分别为{rvals}，弹性方案分别为{evals}。刚性全寿命费用现值为{cr:.2f}万元，弹性为{ce:.2f}万元。各年结果由相应主变、储能和转接配置形成。'
  b=summaries['rigid'];base.append([name,num(b['baseline_2021_capacity_mva'],1),num(b['baseline_2021_load_mw']),num(b['baseline_2021_clr'],1),'8.0742%' if c=='pizhou' else '1.8000%','28.2879%' if c=='pizhou' else '50%'])
 cr=sum(portfolio[c,'rigid'][0]['objective_npv_10k'] for c in REG);ce=sum(portfolio[c,'elastic'][0]['objective_npv_10k'] for c in REG)
 SECTION['cost_results']=f'两地区刚性方案新增措施全寿命费用现值合计{cr:.2f}万元，弹性合计{ce:.2f}万元。弹性较刚性费用{("降低" if ce<cr else "增加")}{abs(1-ce/cr)*100:.2f}%。费用差由所选主变、储能和联络工程的投资、运维及更新共同形成。'
 table('baseline_states',5,'统一容载比2.0的仿真起点',['地区','基期主变容量/MVA','基期年网供最大负荷/MW','容载比','年负荷平均增长率','初始转供率'],base)
 table('cost_comparison',5,'刚性与弹性方案全寿命费用现值比较',['地区','刚性/万元','弹性/万元','差额/万元','弹性费用降幅'],comparison)
 table('cost_components',5,'全寿命费用按工程类型分解',['地区','方案','主变/万元','储能/万元','线路/万元','合计/万元'],components)
 table('measure_totals',5,'全期采购和第4年在役状态',['地区','方案','主变采购/MVA','第三台/项','储能MW/MWh','新增联络/项'],measures)
 table('transfer_results',5,'既有与新增转供能力的实际利用',['地区','方案','仿真年','既有能力使用/MW','新增能力使用/MW','在役项目/项','总转接/MW'],transfers)
 table('base_audit',5,'392个站年的独立复算',['地区','方案','核对站年','最小容量余量/MW','费用偏差/万元','结果'],[[r['地区'],r['方案'],r['站年'],num(r['最小容量余量MW'],6),num(r['费用偏差万元'],6),r['状态']] for r in AUDIT])
 hist=review['solve_history']
 table('solver_stages',5,'分阶段求解精度',['阶段','计算目标','最优性间隙','耗时/s'],[[i+1,['全寿命费用','弹性容载比偏好','固定布局转接量'][i],num(r['mip_gap']*100,4)+'%',num(r['elapsed_seconds'],1)] for i,r in enumerate(hist)])
 table('base_cash_components',4,'方案初始投资与费用现值',['地区','方案','初始投资/万元','现值/万元'],[[REG[c],SC[s],num(sum(sum(float(r[k+'_capex_10k']) for k in ['transformer','storage','line']) for r in yrs),2),num(summ['objective_npv_10k'],2)] for (c,s),(summ,yrs,_) in portfolio.items()])
 table('indicator_roles',5,'参数定义及寻优作用',['参数','采用口径','模型作用'],[
 ['容载比','公用主变总容量/年网供最大负荷[4]','规划结果'],
 ['分布式电源装机渗透率','装机容量/用户最大负荷[69]','构造源荷情景'],
 ['分布式电源电量渗透率','发电量/用户用电量[69]','源荷背景'],
 ['最大正向、反向净负荷','站级方向极值','正常与退出容量校核'],
 ['95%峰值持续时间','最长连续达到本方向峰值95%的时段','储能能量折算'],
 ['年负荷平均增长率','历史同比率递增权重平均值','后续负荷推演'],
 ['主变与线路成本比','相同有效MW初始投资比','联络费用敏感性'],
 ['主变与储能成本比','3 h参考有效MW初始投资比','储能费用敏感性']])
 # 第三章：原始装机序列与实际采用的仿真净负荷分别绘制。
 metrics=json.loads((R/'02图表/2026-10-01重构/指标复算.json').read_text())
 table('pv_capacity',3,'年末分布式光伏装机容量',['地区','2021年/MW','2022年/MW','2023年/MW','2024年/MW','2025年/MW'],[[name]+[num(metrics['model_pv'][rid][str(y)][-1]) for y in range(2021,2026)] for rid,name in [('QX-00005','邳州'),('QX-00007','市区')]])
 fig,axes=plt.subplots(2,1,figsize=(6.7,4.8),sharex=True)
 for k,(rid,name) in enumerate([('QX-00005','邳州'),('QX-00007','市区')]):
  vals=np.array([metrics['model_pv'][rid][str(y)] for y in range(2021,2026)]).flatten();x=np.arange(60);ax=axes[k]
  ax.plot(x[:25],vals[:25],ls='--',color=C[k],label='2021—2022年回推');ax.plot(x[24:],vals[24:],color=C[k],label='2023—2025年月末记录');ax.set_ylabel(name+'装机容量 / MW');ax.legend(fontsize=10);ax.grid(alpha=.2)
 axes[-1].set_xticks([0,12,24,36,48,59],['2021.1','2022.1','2023.1','2024.1','2025.1','2025.12']);save('pv_monthly')
 pressure=[];rankings=[];dominance=[]
 for c,name in REG.items():
  st=portfolio[c,'elastic'][2]
  for y in range(2022,2026):
   a=[r for r in st if int(r['year'])==y]
   pressure.append([name,y-2021,num(sum(float(r['forward_mw']) for r in a)),num(sum(float(r['reverse_mw']) for r in a)),sum(float(r['reverse_mw'])>0 for r in a),num(max(max(float(r['forward_mw']),float(r['reverse_mw'])) for r in a)),len(a)])
  last=[r for r in st if int(r['year'])==2025]
  ranked=sorted(last,key=lambda r:max(float(r['forward_mw']),float(r['reverse_mw'])),reverse=True)[:5]
  for r in ranked:rankings.append([name,r['station'],num(r['forward_mw']),num(r['reverse_mw']),r['forward_duration_h'],r['reverse_duration_h']])
  dominance.append([name,len(last),sum(float(r['reverse_mw'])>float(r['forward_mw']) for r in last),sum(float(r['reverse_mw'])>0 for r in last)])
  ids=[r['station'] for r in last];matrix=np.array([[max(float(r['forward_mw']),float(r['reverse_mw'])) for r in st if r['station']==sid] for sid in ids])
  fig,ax=plt.subplots(figsize=(6.4,5.8 if c=='pizhou' else 7.6));im=ax.imshow(matrix,aspect='auto',cmap='YlOrRd');ax.set_xticks(range(4),['第1年','第2年','第3年','第4年']);ax.set_yticks(range(len(ids)),ids,fontsize=9);fig.colorbar(im,ax=ax,label='净负荷最大绝对值 / MW');save(c+'_modeled_peaks')
 table('annual_model_pressure',3,'仿真各年站级正反向净负荷',['地区','仿真年','各站最大正向净负荷合计/MW','各站最大反向净负荷合计/MW','反向站数','站级最大绝对值/MW','变电站数量'],pressure)
 table('modeled_peak_rankings',3,'第4年净负荷较大站的运行特征',['地区','站号','正向/MW','反向/MW','正向持续/h','反向持续/h'],rankings)
 table('direction_dominance',3,'第4年站级净负荷方向统计',['地区','变电站数量','反向较大站数','存在反向净负荷站数'],dominance)
 fig,axs=plt.subplots(1,2,figsize=(7,3.5))
 for c,name in REG.items():
  data=[r for r in pressure if r[0]==name]
  for ax,j in zip(axs,[2,3]):ax.plot(range(1,5),[float(r[j]) for r in data],'o-',label=name,color=C[list(REG).index(c)])
 for ax,label in zip(axs,['最大正向净负荷合计 / MW','最大反向净负荷合计 / MW']):ax.set_ylabel(label);ax.set_xticks(range(1,5));ax.legend();ax.grid(alpha=.2)
 save('annual_model_pressure')
 plt.figure(figsize=(6.6,3.8))
 for (c,s),(_,yrs,_) in portfolio.items():plt.plot(range(1,5),[float(r['clr']) for r in yrs],marker='o',ls='-' if s=='elastic' else '--',color=C[list(REG).index(c)],label=REG[c]+SC[s])
 plt.axhline(2,color='gray',lw=.8);plt.xticks(range(1,5));plt.xlabel('仿真年');plt.ylabel('容载比');plt.legend(ncol=2);plt.grid(alpha=.2);save('annual_clr')
 labels=[r[0]+r[1] for r in components];plt.figure(figsize=(6.7,3.8));bottom=np.zeros(4)
 for j,lab in enumerate(['主变','储能','线路']):
  vals=[float(r[j+2]) for r in components];plt.bar(labels,vals,bottom=bottom,label=lab,color=C[j]);bottom+=vals
 plt.ylabel('全寿命费用现值 / 万元');plt.legend();save('lifecycle_components')
 fig,axs=plt.subplots(1,2,figsize=(7,3.8))
 axs[0].bar(labels,[float(r[2]) for r in measures],color=C[0]);axs[0].set_ylabel('主变采购容量 / MVA')
 axs[1].bar(labels,[float(r[4].split('/')[1]) for r in measures],color=C[1]);axs[1].set_ylabel('第4年储能能量 / MWh')
 for ax in axs:ax.tick_params(axis='x',labelrotation=30)
 save('measure_comparison')
 matrix=json.loads((D/'通用推荐矩阵.json').read_text())
 table('general_matrix',5,'通用容载比推荐区间',['分布式电源装机渗透率','年负荷平均增长率','95%峰值持续时间','主变与线路成本比','主变与储能成本比','推荐容载比'],[[r[k] for k in ['分布式电源装机渗透率','年负荷平均增长率','95%峰值持续时间','主变与线路成本比','主变与储能成本比_3h','推荐容载比区间']] for r in matrix])
 table('general_assumptions',5,'通用外推参数及站群假定',['参数','取值或范围','设定方式'],[
 ['供电区域','A类、B/C类','两类一般站群'],
 ['初始主变','5站；各2×50或2×40 MVA','等容量基期'],
 ['初始容载比','2.0','两类站群一致'],
 ['负荷分布','0.10/0.15/0.20/0.25/0.30；0.16/0.18/0.20/0.22/0.24','权重之和为1'],
 ['站级峰值合计/区域同期峰值','偏集中1.15；较均匀1.00','案例1.057和1.113支持外推量级'],
 ['光伏分布','0.08/0.14/0.18/0.25/0.35','权重之和为1'],
 ['装机渗透率','0、0.3、0.8、1、1.3、1.8、2、2.3、2.8、3','参数假定'],
 ['年负荷平均增长率','0、1%、2%、4%、7%、8%、10%','参数假定'],
 ['95%峰值持续时间','1、2、3、4、6 h','正反向采用相同场景时长'],
 ['初始站间转供率','A类50%；B/C类28.2879%','案例参数外推'],
 ['线路、储能价格倍数','(0.75,0.50)/(1,1)/(1.25,1.25)/(0.75,1.25)/(1.25,0.50)','四角及中心'],
 ['光伏利用小时、用户年负荷率','1200 h；0.5','电量指标换算'],
 ['目标年、费用时点','第4年；按2025年投运折现','目标年静态规划外推']])
 external=json.loads((D/'外推逐点结果.json').read_text())
 sensitive=defaultdict(set)
 for r in external:sensitive[r['装机渗透率'],r['年负荷平均增长率'],r['95%峰值持续时间h'],r['供电区域'],r['负荷分布']].add(round(r['推荐容载比'],7))
 SECTION['matrix_results']=f'{len(external)}个外推计算点均满足所设正常容量、单台主变退出、转供能力和负荷守恒约束，最大最优性间隙为{max(r["最优性间隙"] for r in external)*100:.4f}%。在{len(sensitive)}组固定技术条件中，{sum(len(v)>1 for v in sensitive.values())}组随价格组合变化发生容载比选择变化。其余组合主要保持容量规格或通过储能、转接量调整费用。'
 for name,items in [('逐年区域方案',resultrows),('容量费用复核',[[r[k] for k in r] for r in AUDIT])]:
  ws=wb.create_sheet(name)
  ws.append(['地区','方案','仿真年','年网供最大负荷/MW','主变总容量/MVA','主变净增容量/MVA','主变采购容量/MVA','新增储能/MW与MWh','新增联络/项','转接负荷/MW','容载比'] if name=='逐年区域方案' else ['地区','方案','站年','最小容量余量/MW','费用偏差/万元','状态'])
  for r in items:ws.append(r)
 wb.save(D/'统一起点逐年方案与复核.xlsx')
 (A/'计算表.json').write_text(json.dumps(T,ensure_ascii=False,indent=2))
 (A/'生成段落.json').write_text(json.dumps(SECTION,ensure_ascii=False,indent=2))
 (HERE/'基准独立复核.json').write_text(json.dumps(AUDIT,ensure_ascii=False,indent=2))
 (HERE/'结果摘要.json').write_text(json.dumps({'刚性合计万元':cr,'弹性合计万元':ce,'费用变化比例':1-ce/cr,'基准':base,'矩阵':matrix,'求解信息':review},ensure_ascii=False,indent=2))
 print('基准392站年复算通过，图表与动态结果已生成')
if __name__=='__main__':main()

