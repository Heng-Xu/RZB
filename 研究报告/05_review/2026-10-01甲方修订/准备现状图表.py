import csv,json,sys,calendar
from collections import defaultdict
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import load_workbook,Workbook
ROOT=Path(__file__).resolve().parents[3];sys.path.insert(0,str(ROOT))
from 实验.研究.rebuild_2026.pizhou_station_scenarios import read_pizhou_station_series
from 实验.研究.rebuild_2026.joint_lifecycle_optimizer import cost_factors
BASE=ROOT/'实验/研究/rebuild_2026/outputs/joint_shared_measure/station_rate_all_years_load_reallocation_full_service_city_A'
RAW=ROOT/'实验/研究/data/tuomin/电网建模数据_Agent整合版_V1.2'
OUT=ROOT/'研究报告/02图表/2026-10-01甲方修订';OUT.mkdir(parents=True,exist_ok=True)
DATA=ROOT/'研究报告/01研究推荐/2026-10-01条件矩阵'
import subprocess
fontpath='/usr/share/fonts/wps-office/FZFSK.TTF'
font_manager.fontManager.addfont(fontpath)
font_manager.fontManager.addfont('/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf')
font_manager.fontManager.addfont('/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman_Bold.ttf')
plt.rcParams.update({'font.family':['Times New Roman',font_manager.FontProperties(fname=fontpath).get_name()],'font.size':10,'axes.unicode_minus':False,'figure.dpi':160})
TABLES={};METRICS={};COLORS=['#29465b','#9b4b39','#778f68','#8b7ba5']
def rows(p):
 with Path(p).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def table(key,title,headers,data):
 TABLES[key]=title+'\n\n| '+' | '.join(map(str,headers))+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(map(str,r))+' |\n' for r in data)+'\n'
def save(name):
 plt.tight_layout();plt.savefig(OUT/(name+'.png'),dpi=240,bbox_inches='tight');plt.close()
def num(x,n=3):return f'{float(x):.{n}f}'
def source_tables():
 table('source_scope','表3-1 原始资料范围与使用口径', ['资料','统计范围','计算用途','处理性质'],[
 ['年度容量和降压负荷[76]','区县、110 kV、2021—2025年','现状、年网供最大负荷及增长','原表换算'],
 ['月末光伏装机[77]','区县、2023—2025年','装机与发电量估计','原表换算'],
 ['设备负载统计[78]','逐台主变、2025年','规格、类别及设备极值','逐台原表'],
 ['站级净负荷[79]','邳州20站、市区样本','方向峰值与持续时间','映射及缺测标记'],
 ['光伏小时系数[80]','统一典型出力曲线','月度发电量估计','典型曲线外推'],
 ['工程投资及采购[70—75]','各项目工程范围','费用系数','原表标定与情景外推']])
 wb=load_workbook(RAW/'近5年容载比.xlsx',read_only=True,data_only=True);s=wb.active
 official=[];annual={}
 for name,rn in [('邳州',19),('市区',9)]:
  cells=list(s.iter_rows(min_row=rn,max_row=rn,values_only=True))[0]
  annual[name]={}
  for y,j in zip(range(2021,2026),[2,5,8,11,14]):
   cap=float(cells[j])*10;load=float(cells[j+1])*10;annual[name][y]=(cap,load)
   official.append([name,y,num(cap,1),num(load),num(cap/load,4)])
 table('official_annual','表3-2 邳州和市区年度原表容量与降压负荷[76]', ['地区','年份','容量/MVA','年网供最大负荷/MW','容载比'],official)
 METRICS['annual']=annual
 gr=[]
 for n in ['邳州','市区']:
  gs=[annual[n][y][1]/annual[n][y-1][1]-1 for y in range(2022,2026)]
  g=sum(i*x for i,x in enumerate(gs,1))/10
  METRICS[n+'_growth']=g
  for y,x,w in zip(range(2022,2026),gs,range(1,5)):gr.append([n,y,num(x*100,4)+'%',w])
  gr.append([n,'年负荷平均增长率',num(g*100,4)+'%','Σ权重=10'])
 table('growth','表3-3 年负荷平均增长率计算（式3-1）',['地区','年份','同比增长率','权重'],gr)
 plt.figure(figsize=(6.5,3.6))
 for n,c in zip(['邳州','市区'],COLORS):plt.plot(range(2021,2026),[annual[n][y][1] for y in range(2021,2026)],'o-',label=n,color=c)
 plt.xticks(range(2021,2026));plt.ylabel('年网供最大负荷 / MW');plt.legend();plt.grid(alpha=.2);save('annual_load')
 w=load_workbook(RAW/'逐月分县分布式光伏.xlsx',read_only=True,data_only=True);m={};year=None
 names={'QX-00005':'邳州','QX-00007':'市区'}
 for rr in w.active.iter_rows(values_only=True):
  if str(rr[1]).startswith('202') and str(rr[1]).endswith('年'):year=int(str(rr[1])[:4])
  if rr[0] in names and year in (2023,2024,2025) and all(x is not None for x in rr[1:13]):m[year,names[rr[0]]]=[float(x)*10 for x in rr[1:13]]
 table('pv_capacity','表3-4 年末分布式光伏装机[77]',['地区','2023年/MW','2024年/MW','2025年/MW'],[[n]+[num(m[y,n][-1]) for y in range(2023,2026)] for n in ['邳州','市区']])
 plt.figure(figsize=(6.5,3.6))
 for n,c in zip(['邳州','市区'],COLORS):plt.plot(range(1,13),m[2025,n],'o-',color=c,label=n)
 plt.xticks(range(1,13));plt.xlabel('月份');plt.ylabel('月末装机 / MW');plt.legend();plt.grid(alpha=.2);save('pv_monthly')
 wb=load_workbook(RAW/'2025设备负载统计表.xlsx',read_only=True,data_only=True);equip=defaultdict(list)
 for rr in wb['主变1'].iter_rows(min_row=3,values_only=True):
  if rr[1]==110 and rr[2] in names and rr[15] is not None:equip[names[rr[2]]].append(float(rr[15]))
 table('equipment_loading','表3-6 逐台110 kV主变年最大负载率统计[78]',['地区','样本台数','均值/%','中位数/%','≥80%台数','最大值/%'],[[n,len(equip[n]),num(np.mean(equip[n]),2),num(np.median(equip[n]),2),sum(x>=80 for x in equip[n]),num(max(equip[n]),2)] for n in ['邳州','市区']])
 METRICS['loading']={n:{'mean':float(np.mean(a)),'median':float(np.median(a)),'n':len(a)} for n,a in equip.items()}
 plt.figure(figsize=(6,3.6));plt.boxplot([equip[n] for n in ['邳州','市区']],tick_labels=['邳州','市区']);plt.axhline(80,color=COLORS[1],ls='--',lw=.8);plt.ylabel('逐台最大负载率 / %');save('loading_distribution')
 ind=rows(ROOT/'实验/研究/rebuild_2026/source_audit/capacity_release_simulation/regional_revision_2026_09_28/2025区县导则指标估计.csv')
 comparison=[]
 for rr in ind:
  n=rr['区县'];comparison.append([n,num(rr['分布式电源装机渗透率（源荷比）'],4),num(float(rr['分布式电源电量渗透率'])*100,2)+'%',num(annual[n][2025][0]/annual[n][2025][1],4),num(np.median(equip[n]),2)+'%','典型场景及样本外推'])
 table('region_indicator_comparison','表3-7 邳州和市区2025年指标对照[76—80]',['地区','装机渗透率估计','电量渗透率估计','原表容载比','负载率中位数','电量性质'],comparison)
 a=[];h=[]
 for region,n in [('pizhou','邳州'),('city','市区')]:
  sr=[rr for rr in rows(BASE/region/'elastic_stations.csv') if rr['year']=='2025']
  for rr in sr:
   f=float(rr['forward_mw']);rev=float(rr['reverse_mw']);a.append([n,rr['station'],num(f),num(rev),num(max(f,rev)),rr['forward_duration_h'],rr['reverse_duration_h']]);h.append((n,int(rr['forward_duration_h']),int(rr['reverse_duration_h'])))
 table('station_pressure','表3-8 2025年站级正反向净负荷与持续时间[78—79]',['地区','变电站编号','正向/MW','反向/MW','净负荷最大绝对值/MW','正向H95/h','反向H95/h'],a)
 plt.figure(figsize=(6.5,3.8))
 for k,n in enumerate(['邳州','市区']):
  vals=[v for v in h if v[0]==n];plt.scatter([v[1] for v in vals],[v[2] for v in vals],label=n,alpha=.65,color=COLORS[k])
 plt.xlabel('正向H95 / h');plt.ylabel('反向H95 / h');plt.legend();plt.grid(alpha=.2);save('duration_distribution')
 # 同样本月度分析；固定2025年末容量是研究量，非月度资产事实。
 series=[v for (kv,sid),v in read_pizhou_station_series().items() if kv==110]
 common=set.intersection(*(set(v) for v in series));net={t:sum(v[t] for v in series) for t in common}
 from datetime import datetime,timedelta
 for h in range(8760):
  t=datetime(2025,1,1)+timedelta(hours=h)
  if t not in net:net[t]=(net[t-timedelta(hours=1)]+net[t+timedelta(hours=1)])/2
 wh=load_workbook(RAW/'光伏8760小时数据.xlsx',read_only=True,data_only=True)
 pv=defaultdict(float)
 for rr in wh.active.iter_rows(min_row=2,values_only=True):
  t=rr[1];month=t.month;prev=m[2024,'邳州'][-1] if month==1 else m[2025,'邳州'][month-2]
  pv[month]+=float(rr[4])*(prev+m[2025,'邳州'][month-1])/2
 mr=[]
 for month in range(1,13):
  v=[x for t,x in net.items() if t.month==month];e=sum(v)*1.005033;u=e+pv[month];pk=max(v)*1.005033;abs_pk=max(abs(x) for x in v)*1.005033
  mr.append([month,pv[month]/u,2139.5/pk,100*abs_pk/(.95*2139.5),pv[month],u,pk])
 cr=[['电量渗透率—容载比（月度分析）',num(np.corrcoef([r[1] for r in mr],[r[2] for r in mr])[0,1],4),12,'固定容量、月度估计、样本内'],['电量渗透率—主变负载率估计值',num(np.corrcoef([r[1] for r in mr],[r[3] for r in mr])[0,1],4),12,'功率因数假定、非逐台实测']]
 table('monthly_correlations','表3-9 邳州12个月的描述性相关系数（式3-8）',['关联对象','Pearson系数','样本量','解释范围'],cr)
 table('monthly_indicators','表3-10 邳州月度指标计算明细[77—80]',['月份','电量渗透率/%','容载比（月度分析）','主变负载率估计/%','光伏电量/MWh','用户电量/MWh'],[[r[0],num(r[1]*100,2),num(r[2],3),num(r[3],2),num(r[4],0),num(r[5],0)] for r in mr])
 METRICS['monthly_correlations']=cr
 fig,axs=plt.subplots(1,2,figsize=(7,3.3))
 for ax,j,label in [(axs[0],2,'容载比（月度分析）'),(axs[1],3,'主变负载率估计 / %')]:
  ax.scatter([r[1]*100 for r in mr],[r[j] for r in mr],color=COLORS[0]);ax.set_xlabel('月电量渗透率 / %');ax.set_ylabel(label)
  for r in mr:ax.annotate(str(r[0]),(r[1]*100,r[j]),fontsize=7,xytext=(3,2),textcoords='offset points')
 save('monthly_correlation')
def cost_tables():
 table('economic_assumptions','表4-1 统一经济评价假定',['参数','取值','依据性质'],[['基准年／截止年','2021／2041','研究评价边界'],['折现率','6%','统一实价评价假定'],['主变／线路寿命','20／20年','研究寿命假定'],['储能寿命','10年','更新周期假定'],['主变／线路固定运维','初始投资1%/年','统一研究参数'],['储能固定运维','初始投资3%/年','统一研究参数'],['旧设备回收／期末残值','未计','无同范围可靠依据'],['储能充电购电／损耗电费','未计','未取得统一运行资料']])
 project=rows(ROOT/'实验/研究/rebuild_2026/source_audit/transformer_project_cost_references.csv')
 table('transformer_projects','表4-2 徐州110 kV主变工程价格依据[70]',['项目','新购/MVA','退出/MVA','净增/MVA','静态投资/万元'],[[r['project_name'].strip(),r['purchased_transformer_mva'],r['replaced_old_transformer_mva'],r['project_capacity_increment_mva'],r['static_total_10k_cny']] for r in project if r['voltage_kv']=='110'])
 table('line_cost_components','表4-3 等效联络线参考投资构成[71]',['项目','工程量','参考单价','费用/万元'],[['架空','2.7 km','24万元/km','64.8'],['电缆电气','0.3 km','120万元/km','36.0'],['智能开关','2台','3万元/台','6.0'],['合计','3 km及2台开关','线路电气与开关费用','106.8']])
 table('storage_anchors','表4-4 储能采购价格依据[72][73]',['案例','公告规模','保证规模','价格/万元','价格性质'],[['苏州','100 kW/215 kWh','同公告','27.2','储能子项最高限价'],['浏阳','1 MW/2.2 MWh','最低1 MW/2.15 MWh','210.456262','储能项目中标总价']])
 ds=[1,2,2.15,3,4,6,8,12]
 table('storage_duration_cost','表4-5 10柜包理想持续支撑与初始费用（式4-6）',['持续时间/h','有效功率/MW','单位有效MW费用/万元','主变／储能成本比'],[[d,num(min(1,2.15/d),6),num(210.456262/min(1,2.15/d),3),num((16.573333/.95)/(210.456262/min(1,2.15/d)),5)] for d in ds])
 plt.figure(figsize=(6.5,3.5));x=np.linspace(1,12,100);plt.plot(x,210.456262/np.minimum(1,2.15/x),color=COLORS[0]);plt.xlabel('持续时间 / h');plt.ylabel('初始投资 / 万元 / 有效MW');plt.grid(alpha=.2);save('storage_effective_cost')
 f=cost_factors();table('lifecycle_factors','表4-6 单位初始投资的2021年末现值因子（式4-10）',['投运年','主变','线路','储能'],[[y,num(f['transformer'][y],6),num(f['line'][y],6),num(f['storage'][y],6)] for y in range(2022,2026)])
def model_tables():
 # 正文的年度图表采用实际进入本轮模型的数据，原表行列追溯另存。
 wb=load_workbook(RAW/'逐月分县分布式光伏.xlsx',read_only=True,data_only=True);monthly={};year=None
 for rr in wb.active.iter_rows(values_only=True):
  if str(rr[1]).startswith('202') and str(rr[1]).endswith('年'):year=int(str(rr[1])[:4])
  if rr[0] in ['QX-00005','QX-00007'] and year in (2023,2024,2025) and all(x is not None for x in rr[1:13]):monthly[rr[0],year]=[float(x)*10 for x in rr[1:13]]
 for rid in ['QX-00005','QX-00007']:
  for y in [2021,2022]:monthly[rid,y]=[a*(a/b)**(2023-y) for a,b in zip(monthly[rid,2023],monthly[rid,2024])]
 pvs=[];model=[];pressure=[]
 for c,rid,name in [('pizhou','QX-00005','邳州'),('city','QX-00007','市区')]:
  summary=json.loads((BASE/c/'summary.json').read_text())[0];annual=METRICS['annual'][name];st=rows(BASE/c/'elastic_stations.csv')
  pvs.append([name]+[num(monthly[rid,y][-1]) for y in range(2021,2026)])
  model.append([name,2021,num(summary['baseline_2021_capacity_mva'],1),num(summary['baseline_2021_load_mw']),num(summary['baseline_2021_clr'],4),num(monthly[rid,2021][-1]),'共同反事实初态／光伏回推'])
  for y in range(2022,2026):
   a=[r for r in st if r['year']==str(y)];pk=annual[y][1]
   model.append([name,y,'寻优确定',num(pk),'寻优确定',num(monthly[rid,y][-1]),'光伏回推' if y==2022 else '光伏原表；压力模板换算'])
   pressure.append([name,y,num(sum(float(r['forward_mw']) for r in a)),num(sum(float(r['reverse_mw']) for r in a)),sum(float(r['reverse_mw'])>0 for r in a),num(max(max(float(r['forward_mw']),float(r['reverse_mw'])) for r in a)),len(a)])
 table('official_annual','表3-2 本轮年度规划输入及共同初态（式3-1、式3-9）',['地区','年份','初态容量/MVA','年网供最大负荷/MW','初态比值','年末光伏/MW','建模性质'],model)
 table('pv_capacity','表3-4 进入历史压力重建的年末光伏规模（式3-9）',['地区','2021年/MW','2022年/MW','2023年/MW','2024年/MW','2025年/MW'],pvs)
 table('annual_model_pressure','表3-11 基准模型年度站级压力汇总（式3-10）',['地区','年份','正向任务之和/MW','反送任务之和/MW','反送站数','最大站净负荷最大绝对值/MW','模型站数'],pressure)
 fig,axes=plt.subplots(2,1,figsize=(6.7,4.6),sharex=True)
 for ax,(rid,name,col) in zip(axes,[('QX-00005','邳州',COLORS[0]),('QX-00007','市区',COLORS[1])]):
  vals=np.array([monthly[rid,y] for y in range(2021,2026)]).reshape(-1);x=np.arange(60)
  ax.plot(x[:25],vals[:25],ls='--',color=col,label='2021—2022年几何回推');ax.plot(x[24:],vals[24:],color=col,label='2023—2025年月末原表');ax.set_ylabel(name+'装机 / MW');ax.axvline(23.5,color='gray',ls=':',lw=.7);ax.legend(fontsize=8);ax.grid(alpha=.2)
 axes[-1].set_xticks([0,12,24,36,48,59],['2021.1','2022.1','2023.1','2024.1','2025.1','2025.12']);save('pv_monthly')
 fig,axes=plt.subplots(1,2,figsize=(7,3.4))
 for name,col in [('邳州',COLORS[0]),('市区',COLORS[1])]:
  a=[r for r in pressure if r[0]==name];axes[0].plot([r[1] for r in a],[float(r[2]) for r in a],marker='o',color=col,label=name);axes[1].plot([r[1] for r in a],[float(r[3]) for r in a],marker='o',color=col,label=name)
 for ax,label in zip(axes,['独立正向峰任务之和 / MW','反送筛查任务之和 / MW']):ax.set_ylabel(label);ax.set_xticks(range(2022,2026));ax.legend();ax.grid(alpha=.2)
 save('annual_model_pressure')
 METRICS['model_pv']={rid:{str(y):monthly[rid,y] for y in range(2021,2026)} for rid in ['QX-00005','QX-00007']}
def main():
 source_tables();cost_tables()
 (OUT/'计算表.json').write_text(json.dumps(TABLES,ensure_ascii=False,indent=2))
 (OUT/'指标复算.json').write_text(json.dumps(METRICS,ensure_ascii=False,indent=2))
 print('tables',len(TABLES),'figures',len(list(OUT.glob('*.png'))));print(json.dumps(METRICS['monthly_correlations'],ensure_ascii=False))
if __name__=='__main__':main()
