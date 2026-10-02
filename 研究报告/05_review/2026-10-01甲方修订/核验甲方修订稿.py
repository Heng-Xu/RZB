"""交付文档、案例结果、通用矩阵及配套工作簿的一致性验收。"""
from pathlib import Path
import json,re,zipfile,hashlib
from collections import Counter
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
import pymupdf
HERE=Path(__file__).resolve().parent
R=HERE.parents[1]
ITER='2026-10-01甲方修订'
TEXT=R/'03_MD'/ITER
DATA=R/'01研究推荐'/ITER
ASSET=R/'02图表'/ITER
OUT=R/'04_word'/ITER

def read(p):return json.loads(p.read_text())
def norm(s):return re.sub(r'\s+','',s)
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def sequential(items,kind):
 grouped={}
 for c,n in items:grouped.setdefault(c,[]).append(int(n))
 for c,ns in grouped.items():assert ns==list(range(1,len(ns)+1)),(kind,c,ns)
 return sum(map(len,grouped.values()))

def main():
 source=(TEXT/'研究报告装配稿.md').read_text()
 assert len(re.findall(r'^# 第[一二三四五六]章 ',source,re.M))==6
 assert not re.search(r'^#{1,4} 4\.7(?:\s|\.)',source,re.M)
 assert source.index('## 5.2 三类方案')<source.index('## 5.4 寻优结果')<source.index('## 5.6 参数拓展')<source.index('## 5.7 规划建议')
 for term in ['丰县','共同反事实','规划分母代理','建模绝对峰','最大站绝对峰','模型站数','正向/反送任务之和','分母代理','四条路径','[[TABLE:','[[SECTION:','[[FIG:','为本项目本研究']:
  assert term not in source,term
 refs=re.findall(r'^\[(\d+)\]\s+',source,re.M)
 assert list(map(int,refs))==list(range(1,81)),refs
 tables=sequential(re.findall(r'^表(\d+)-(\d+)\s',source,re.M),'表')
 figures=sequential(re.findall(r'^图(\d+)-(\d+)\s',source,re.M),'图')
 eqs=read(ASSET/'公式/公式清单.json')
 assert sequential([x['number'].split('-') for x in eqs],'公式')==38
 mathsource=list(re.finditer(r'\$\$([\s\S]*?)\$\$|\$([^$\n]+)\$',source))
 displays=sum(m[1] is not None for m in mathsource)
 doc=Document(OUT/'研究报告_WPS兼容稿.docx')
 objects=doc._element.xpath('.//m:oMath')
 assert len(objects)==len(mathsource)==247
 assert len(doc._element.xpath('.//m:oMathPara'))==displays==38
 assert len(doc.inline_shapes)==figures==13
 assert len(doc.tables)==tables+displays==72
 for tab in doc.tables:
  if not tab._tbl.xpath('.//m:oMathPara'):continue
  assert len(tab.rows)==1 and len(tab.columns)==3
  assert tab.cell(0,0).width==tab.cell(0,2).width
  assert len(tab.cell(0,1)._tc.xpath('.//m:oMathPara'))==1
  para=tab.cell(0,1).paragraphs[0]
  assert para.alignment==1
  assert para._p.xpath('.//m:jc')[0].get(qn('m:val'))=='center'
  assert re.fullmatch(r'（\d+-\d+）',tab.cell(0,2).text)
 docx_paths=re.findall(r'!\[\]\(([^)]+\.png)\)',source)
 assert len(docx_paths)==13 and all(Path(p).is_file() for p in docx_paths)
 font=read(HERE/'图件字体说明.json')
 assert font['中文字体']=='方正仿宋_GBK'
 assert doc.styles['Normal'].element.xpath('.//w:rFonts')[0].get(qn('w:eastAsia'))==font['中文字体']
 # 目录以实际排版的正文页核定，同时检查Word缓存。
 pdf=pymupdf.open(OUT/'研究报告_WPS兼容稿.pdf')
 pages=[norm(p.get_text()) for p in pdf]
 cache=read(HERE/'目录页码核对.json')
 assert len(cache)==44
 for item in cache:
  actual=[i+1 for i,t in enumerate(pages) if i>=4 and norm(item['标题']) in t]
  assert actual and actual[0]==item['页码'],(item,actual[:1])
 indices={x['锚点']:str(x['页码']) for x in cache}
 fields=0
 for field in doc._element.xpath('.//w:fldSimple'):
  match=re.search(r'PAGEREF\s+(\w+)',field.get(qn('w:instr'),''))
  if not match:continue
  assert ''.join(x.text or '' for x in field.iter(qn('w:t')))==indices[match[1]],(match[1],''.join(x.text or '' for x in field.iter(qn('w:t'))))
  fields+=1
 assert fields==44
 with zipfile.ZipFile(OUT/'研究报告_WPS兼容稿.docx') as z:
  ns={'ep':'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'}
  tree=etree.fromstring(z.read('docProps/app.xml'))
  assert int(tree.find('ep:Pages',ns).text)==len(pdf)
 assert not any('¿' in p.get_text() for p in pdf),'PDF数学解析错误符号'
 for page in pdf:
  for info in page.get_image_info():
   rect=pymupdf.Rect(info['bbox'])
   assert rect.x0>=-1 and rect.x1<=page.rect.width+1 and rect.y0>=-1 and rect.y1<=page.rect.height+1,('图形越出页面',page.number,rect)
 # 技术、费用及求解精度采用独立复算记录。
 review=read(DATA/'统一起点基准/review.json')
 base=read(HERE/'基准独立复核.json')
 external=read(HERE/'外推独立复核.json')
 stats=read(DATA/'外推核验.json')
 assert sum(x['站年'] for x in base)==392 and all(x['状态']=='通过' for x in base)
 assert external['状态']=='通过' and external['计算组合']==7000 and external['站级复核']==35000
 assert stats['全部约束核对']=='通过' and stats['最大最优性间隙']<=.005
 assert review['primary_cost_relative_gap']<=.01 and review['capacity_preference_relative_gap']<=.005
 assert review['cost_order_imposed'] is False
 assert review['same_measure_set'] and review['same_start_per_district']
 costs=review['cost_npv_10k']
 rigid=sum(x['rigid'] for x in costs.values());elastic=sum(x['elastic'] for x in costs.values())
 assert all(x['elastic']<x['rigid'] for x in costs.values())
 assert '82.89%' in source and abs((1-elastic/rigid)*100-82.89)<.005
 assert '推荐采用弹性方案' in source
 matrix=read(DATA/'通用推荐矩阵.json')
 assert len(matrix)==18
 allowed={'分布式电源装机渗透率','年负荷平均增长率','95%峰值持续时间','主变与线路成本比','主变与储能成本比_3h','推荐容载比区间','计算点数','刚性容载比控制上限'}
 assert all(set(x)==allowed and float(x['刚性容载比控制上限'])==2 for x in matrix)
 for row in matrix:
  low,high=map(float,row['推荐容载比区间'].split('～'))
  assert 2<=low<=high<=2.6
 cases=read(HERE/'案例与矩阵衔接核对.json')
 for case in cases:
  assert case['状态']=='通过'
  for index in case['匹配矩阵行']:
   lower,upper=map(float,matrix[index-1]['推荐容载比区间'].split('～'))
   assert lower<=case['弹性容载比']<=upper
 workbook=load_workbook(DATA/'通用容载比推荐矩阵.xlsx',read_only=True)
 ws=workbook['通用推荐区间'];assert ws.max_row==19 and ws.max_column==8
 headers=[c.value for c in ws[1]]
 assert not any(any(t in h for t in ['柜数','措施','项目数']) for h in headers)
 assert '刚性上限' in headers and '推荐弹性容载比' in headers
 for row,expected in zip(ws.iter_rows(min_row=2,values_only=True),matrix):
  assert row[5]=='2.0' and row[6]==expected['推荐容载比区间']
 assert workbook['外推计算点'].max_row==7001
 workbook=load_workbook(DATA/'统一起点逐年方案与复核.xlsx',read_only=True)
 assert workbook['逐年区域方案'].max_row==17 and workbook['容量费用复核'].max_row==5
 assert sum(workbook[x].max_row-1 for x in ['邳州刚性逐站','邳州弹性逐站','市区刚性逐站','市区弹性逐站'])==392
 # 源文件说明、摘要及结论与模型保持相同范围。
 notes=(R/'数据来源'/ITER/'数据来源与参数说明.md').read_text()
 assert '原值乘10' in notes and '7000' in notes and '1.15' in notes
 style=read(HERE/'语体与术语审查.json');assert style['状态']=='通过'
 result={'状态':'通过','章节':6,'参考文献':80,'PDF页数':len(pdf),'真实数据表':tables,'图件':figures,'独立居中公式':displays,'行内原生公式':len(objects)-displays,'原生数学对象':len(objects),'目录页码核对':fields,
 '基准独立复核站年':392,'一般站群外推组合':7000,'外推独立复核站级记录':35000,'通用矩阵行数':18,'案例校核点':cases,
 '刚性费用现值万元':rigid,'弹性费用现值万元':elastic,'弹性费用降低比例':1-elastic/rigid,'费用排序作为约束':False,
 '费用阶段最优性间隙':review['primary_cost_relative_gap'],'容量偏好阶段最优性间隙':review['capacity_preference_relative_gap'],'外推最大最优性间隙':stats['最大最优性间隙'],
 '公式处理':'交付Word全部为原生OMML，行内可编辑；独立公式在等宽侧栏间居中，编号单独右置。PDF由同一LaTeX源生成数学图形后在本地排版导出；排版副本不替代原生Word。',
 '版面复核':'已查看内容提要、第五章起点、增容/储能公式、负荷转接示例、刚弹对照表、通用矩阵及规划建议页面；图形均未越出PDF页面。',
 '审查范围':'数据和模型复算、文档结构与本地版面核对；交付供电科院审核。',
 '字体':font,'语体审核':style['技能'],'文件SHA256':{p.name:digest(p) for p in [OUT/'研究报告_WPS兼容稿.docx',OUT/'研究报告_WPS兼容稿.pdf',DATA/'通用容载比推荐矩阵.xlsx',DATA/'统一起点逐年方案与复核.xlsx']}}
 (HERE/'甲方修订稿最终核验.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
 print(json.dumps({k:result[k] for k in ['状态','PDF页数','真实数据表','图件','原生数学对象','目录页码核对','基准独立复核站年','外推独立复核站级记录','通用矩阵行数']},ensure_ascii=False))
if __name__=='__main__':main()
