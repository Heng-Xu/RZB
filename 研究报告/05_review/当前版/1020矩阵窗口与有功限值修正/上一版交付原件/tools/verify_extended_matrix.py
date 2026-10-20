"""核验扩展任务完整性、窗口依据、费用及独立转供LP。"""
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path
from multiprocessing import get_context
import argparse
import json
import math
import sys

import numpy as np
import pandas as pd

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'tools'))
from extend_matrix_scenarios import TRIAL,PUBLIC,SOURCE,REVIEW,tasks,write,model
from locked_parameter_results import key,public_frame,sha,independent_initial_check
from matrix_planning_windows import records,export
from archive_extended_matrix_raw import raw_path,open_raw


def check_lp(row):
    result=model.independent_lp_check(row,row['units'],row['modules'],set(row['built']))
    assert result['verified'],(key(row),result)
    initial=independent_initial_check(model.make_input(row))
    assert math.isclose(initial['weighted_rate'],row['initial_transfer_fraction'],abs_tol=1e-8)
    return dict(节点=key(row),独立转供LP=result,初始转供率=initial['weighted_rate'])


def verify(lp=False,workers=8):
    design=json.loads((TRIAL/'试算设计.json').read_text())
    assert sha(Path(model.__file__))==design['模型SHA256']
    library=raw_path()
    rows=[json.loads(line) for line in open_raw(library)]
    assert len(rows)==design['任务数']==127578
    assert len({key(r) for r in rows})==len(rows)
    assert {key(r) for r in rows}=={key(r) for r in tasks()}
    good=[r for r in rows if r['verified']]
    assert all(r['verified'] or r['status']=='Infeasible' for r in rows)
    actual=pd.read_csv(PUBLIC/'统一有效样本.csv')
    expected=public_frame(pd.DataFrame(good))
    fields=list(expected.columns[:6])
    pd.testing.assert_frame_equal(expected.sort_values(fields).reset_index(drop=True),actual.reset_index(drop=True),
        check_dtype=False,check_exact=False,rtol=1e-12,atol=1e-9)
    # 由原设备重新求容量、能量及三类费用，独立于求解目标矩阵。
    from 实验.研究.rebuild_2026.current_economic_assumptions import cost_factors,line_capex
    factors={k:cost_factors()[k][2025] for k in ['transformer','storage','line']}
    for r in good:
        units=np.array(r['units']);modules=sum(r['modules'])
        assert units.sum()==r['capacity_mva']
        assert math.isclose(units.sum()/r['peak_mw'],r['clr'],abs_tol=1e-12)
        assert math.isclose(modules*.215,r['storage_mwh'],abs_tol=1e-9)
        costs=dict(transformer=(np.where(units[:,:2]>50,units[:,:2],0).sum()*16.57333333333333+
                    units[:,2].sum()/50*1223)*factors['transformer'],
                   storage=modules*21.0456262*r['storage_price']*factors['storage'],
                   line=len(r['built'])*line_capex(r['cable_fraction'])*factors['line'])
        assert abs(sum(costs.values())-r['cost'])<1e-4
        for k,value in costs.items():assert abs(value-r['cost_components'][k])<1e-5
        assert r['physical_check']['max_violation_mw']<1e-5
    if lp:cells,sequence=export(actual,PUBLIC)
    else:
        cells=records(actual)
        sequence=pd.read_csv(PUBLIC/'负荷分散系数11点完整序列.csv')
        assert cells==json.loads((PUBLIC/'窗口与原始节点依据.json').read_text())
    # 独立按明示数值条件重新分组，不复用生产路径的bin列。
    selected=set()
    for cell in cells:
        limits=[[None,.02,.04,.07,None],[None,1,2,None],[None,.3,.5,.7,None],[None,.1,.3,None]]
        mask=np.ones(len(actual),dtype=bool)
        for field,idx,cut in zip(['年最大净负荷平均增长率','分布式电源装机渗透率','初始转供率','电缆长度占比'],
                [cell['增长档'],cell['装机档'],cell['转供档'],cell['电缆档']],limits):
            if cut[idx] is not None:mask &= actual[field]>cut[idx]+1e-9
            if cut[idx+1] is not None:mask &= actual[field]<=cut[idx+1]+1e-9
        mask &= actual['负荷分散系数'].le(1.15+1e-9) if cell['分散档']==0 else actual['负荷分散系数'].gt(1.15+1e-9)
        sub=actual.loc[mask]
        assert len(sub)==cell['有效样本数']
        assert math.isclose(float(np.median(sub['容载比'])),cell['中位数'],abs_tol=1e-12)
        assert math.isclose(float(sorted(sub['容载比'])[len(sub)//2]),cell['代表节点容载比'],abs_tol=1e-12)
        low,high=map(float,cell['推荐窗口'].split('～'))
        assert math.isclose(high-low,.2,abs_tol=1e-12)
        assert sub['容载比'].between(low-1e-9,high+1e-9).any(),('窗口内无真实方案',cell)
        for index in [sub['容载比'].idxmin(),sub['容载比'].idxmax(),sub.sort_values('容载比',kind='stable').index[len(sub)//2]]:
            selected.add(tuple(round(float(actual.loc[index,f]),9) for f in fields))
    for _,r in sequence.iterrows():selected.add(tuple(round(float(r[f]),9) for f in fields))
    bykey={key(r):r for r in good}
    from openpyxl import load_workbook
    book=load_workbook(PUBLIC/'推广矩阵与11点敏感性分析.xlsx',read_only=True,data_only=True)
    for cable in range(3):
        for kd in range(2):
            sheet=book[f'矩阵_电缆{cable+1}_分散{kd+1}']
            assert (sheet.max_row,sheet.max_column)==(13,6)
            for row_index,(growth,eta) in enumerate((g,e) for g in range(4) for e in range(3)):
                expected=[r['推荐窗口'] for r in sorted([r for r in cells if r['电缆档']==cable and r['分散档']==kd and r['增长档']==growth and r['装机档']==eta],key=lambda r:r['转供档'])]
                assert [sheet.cell(row_index+2,j+3).value for j in range(4)]==expected
    sheet=book['负荷分散系数11点完整序列']
    assert sheet.max_row==12
    headers=list(next(sheet.iter_rows(values_only=True)))
    actual_sequence=pd.DataFrame(list(sheet.iter_rows(min_row=2,values_only=True)),columns=headers)
    pd.testing.assert_frame_equal(sequence.reset_index(drop=True),actual_sequence,check_dtype=False,check_exact=False,rtol=1e-12,atol=1e-9)
    book.close()
    if lp:
        with ProcessPoolExecutor(max_workers=workers,mp_context=get_context('spawn')) as pool:
            checks=list(pool.map(check_lp,[bykey[k] for k in sorted(selected)],chunksize=4))
        write(REVIEW/'独立LP核验.json',dict(状态='通过',原始库SHA256=sha(library),节点数=len(checks),节点=checks))
    else:
        checks=json.loads((REVIEW/'独立LP核验.json').read_text())
        assert checks['原始库SHA256']==sha(library) and checks['状态']=='通过'
    write(REVIEW/'扩展数据独立核验.json',dict(状态='通过',设计节点数=len(rows),有效节点数=len(good),
        不可行节点数=len(rows)-len(good),分散系数节点=sorted(actual['负荷分散系数'].unique().tolist()),
        查表矩阵数=6,区间单元数=len(cells),窗口宽度=.2,窗口滑动步长=.05,
        原始库SHA256=sha(library),有效样本SHA256=sha(PUBLIC/'统一有效样本.csv'),
        窗口依据SHA256=sha(PUBLIC/'窗口与原始节点依据.json'),敏感性序列SHA256=sha(PUBLIC/'负荷分散系数11点完整序列.csv'),
        独立LP记录SHA256=sha(REVIEW/'独立LP核验.json'),
        说明='窗口以最接近组内中位数的真实节点定位，偶数样本同距时取较大容载比；全体极值另列。相同代表值保留相同窗口。'))
    print('扩展数据通过：',len(rows),'组设计，',len(good),'个有效节点，288单元、11点序列及费用独立复算。')


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--lp',action='store_true')
    parser.add_argument('--workers',type=int,default=8)
    args=parser.parse_args();verify(args.lp,args.workers)
