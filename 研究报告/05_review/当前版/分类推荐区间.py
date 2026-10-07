"""从既有25200组结果生成按供电类别、负荷分散系数分类的推荐区间，不重算模型。"""
from pathlib import Path
import json,math,itertools,statistics,hashlib
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
HERE=Path(__file__).resolve().parent;R=HERE.parents[1];D=R/'01研究推荐/当前版';A=R/'02图表/当前版';O=HERE/'1007术语及分类推荐专项修订'
AREA={'A':'A类','BC':'B/C类'};KD={'较均匀':1.,'偏集中':1.15}
def read(p):return json.loads(p.read_text())
def save(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2))
def table(title,heads,rows):return title+'\n\n| '+' | '.join(heads)+' |\n| '+' | '.join(['---']*len(heads))+' |\n'+''.join('| '+' | '.join(map(str,r))+' |\n' for r in rows)+'\n'
def interval(vals):return f'{math.floor((min(vals)+1e-10)*100)/100:.2f}～{math.ceil((max(vals)-1e-10)*100)/100:.2f}'
def main():
 rows=read(D/'外推逐点结果.json');raw=read(D/'按供电类别与分布推荐矩阵.json');t=read(A/'计算表.json');s=read(A/'生成段落.json'); formal=[];blocks=[]
 for lt,area,shape in itertools.product(('架空为主','电缆为主'),AREA,KD):
  sub=[r for r in raw if r['线路类型']==lt and r['供电区域']==area and r['负荷分布']==shape]
  key=f'classified_{len(blocks)}';kd=KD[shape];tl=sub[0]['主变与线路成本比']
  title=f'表5-0 {AREA[area]}、负荷分散系数{kd:.2f}的容载比推荐区间（{lt}）'
  heads=['分布式电源装机渗透率','年最大净负荷平均增长率','95%峰值持续时间/h','主变与储能单位功率投资比（3 h）','容载比推荐区间']
  values=[]
  for r in sub:
   entry={'供电区域类别':AREA[area],'负荷分散系数':kd,'线路类型':lt,'分布式电源装机渗透率':r['分布式电源装机渗透率'],'年最大净负荷平均增长率':r['年负荷平均增长率'],'95%峰值持续时间/h':r['95%峰值持续时间'].replace(' h',''),'主变与线路单位功率投资比':tl,'主变与储能单位功率投资比（3 h）':r['主变与储能成本比_3h'],'容载比推荐区间':r['计算容载比范围'],'容载比最小值':r['容载比最小值'],'容载比最大值':r['容载比最大值'],'计算点数':r['计算点数']}
   formal.append(entry);values.append([entry[k] for k in heads])
  t[key]=table(title,heads,values)+f'注：本表主变与线路单位功率投资比为{tl}，分别计入三档线路价格。容载比按MVA/MW计，区间由对应条件下的计算结果归纳。\n'
  block=f'#### {AREA[area]}供电区域，负荷分散系数{kd:.2f}，{lt}\n\n'
  point=[r for r in rows if r['线路类型']==lt and r['供电区域']==area and r['负荷分布']==shape]
  width=[r['容载比最大值']-r['容载比最小值'] for r in sub]
  block+=f'该类条件采用{50 if area=="A" else 40} MVA初始主变规格，初始站间转供率为{50 if area=="A" else 28.29:.2f}%。负荷分散系数{kd:.2f}表示各站最大负荷之和为区域同时最大负荷的{kd:.2f}倍。分类推荐同时考虑分布式电源装机规模、负荷增长和95%峰值持续时间，表内区间宽度中位数为{statistics.median(width):.2f}。各站负荷峰值对本地设备提出容量要求，区域同时最大负荷则决定容载比分母，因此两者应在相同供电范围内对应分析。\n\n'
  block+=f'线路投资按{lt}的敷设比例计算，主变与线路单位功率投资比为{tl}。这三档价格条件下，所选主变容量保持一致，线路费用差异计入评价期费用比较。主变与储能单位功率投资比按3 h条件列示；持续时间超过3 h时，相同储能能量能够维持的充放电功率下降，实际配置仍按表中对应持续时间确定。\n\n[[TABLE:{key}]]\n\n'
  for eta in ['0.00～1.00','1.00～2.00','2.00～3.00']:
   q=[r for r in sub if r['分布式电源装机渗透率']==eta]
   by={g:interval([v for r in q if r['年负荷平均增长率']==g for v in [r['容载比最小值'],r['容载比最大值']]]) for g in dict.fromkeys(x['年负荷平均增长率'] for x in q)}
   desc='；'.join(f'年最大净负荷平均增长率{g}时为{v}' for g,v in by.items())
   block+=f'分布式电源装机渗透率为{eta}时，合并本表所列持续时间和储能价格条件，容载比范围表现为：{desc}。这些合并范围用于观察增长条件的影响，具体推荐仍按表中持续时间及主变与储能单位功率投资比定位，以保留不同储能配置对应的较窄区间。\n\n'
  blocks.append(block)
 for key in ['cost_matrix_0','cost_matrix_1']:t.pop(key,None)
 s['cost_banded_matrix']='\n'.join(blocks)
 s['matrix_results']='分类表共列432组推荐条件。与按供电类别合并的108组结果相比，区间宽度中位数由'+f'{statistics.median(r["容载比最大值"]-r["容载比最小值"] for r in read(D/"通用推荐矩阵.json")):.2f}缩小至{statistics.median(r["容载比最大值"]-r["容载比最小值"] for r in raw):.2f}。区间保留相应条件下全部已算配置的端点，差异较大的组合通过分类列示，不另行截取端点。'
 # Explicit line investment ratio analysis, including cost and capacity responses.
 line_summary=[]
 for area,lt in itertools.product(AREA,('架空为主','电缆为主')):
  for price in [.75,1.,1.25]:
   q=[r for r in rows if r['供电区域']==area and r['线路类型']==lt and r['线路价格倍数']==price]
   line_summary.append([AREA[area],lt,f'{price:.2f}',f'{q[0]["主变线路成本比"]:.2f}',interval([r['推荐容载比'] for r in q]),f'{statistics.mean(r["联络项目数"] for r in q):.2f}',f'{statistics.mean(r["全寿命费用万元"] for r in q):.2f}'])
 t['line_ratio_effects']=table('表5-0 主变与线路单位功率投资比及方案结果',['供电区域类别','线路类型','线路价格倍数','主变与线路单位功率投资比','容载比范围','新增联络项目数/项','平均费用现值/万元'],line_summary)
 # Paired explicit ratios, keep direction from low ratio to high ratio.
 groups=defaultdict(dict)
 for r in rows:groups[tuple(r[k] for k in ['供电区域','负荷分布','装机渗透率','年负荷平均增长率','95%峰值持续时间h','储能价格倍数','线路类型'])][r['线路价格倍数']]=r
 pairs=[]
 for g,x in groups.items():
  lo=x[1.25];hi=x[.75];pairs.append({'供电区域类别':AREA[g[0]],'线路类型':g[-1],'低投资比':lo['主变线路成本比'],'高投资比':hi['主变线路成本比'],'容载比变化':hi['推荐容载比']-lo['推荐容载比'],'费用现值变化万元':hi['全寿命费用万元']-lo['全寿命费用万元']})
 paired=[]
 for ar,lt in itertools.product(AREA.values(),('架空为主','电缆为主')):
  q=[x for x in pairs if x['供电区域类别']==ar and x['线路类型']==lt]; paired.append([ar,lt,f'{q[0]["低投资比"]:.2f}→{q[0]["高投资比"]:.2f}',len(q),f'{statistics.mean(x["容载比变化"] for x in q):.2f}',f'{statistics.mean(x["费用现值变化万元"] for x in q):.2f}'])
 t['line_ratio_paired']=table('表5-0 主变与线路单位功率投资比提高时的配对结果',['供电区域类别','线路类型','投资比变化','配对数','容载比平均变化','费用现值平均变化/万元'],paired)
 save(A/'计算表.json',t);save(A/'生成段落.json',s);save(D/'分类容载比推荐区间.json',formal);save(O/'线路投资比配对分析.json',pairs)
 wb=load_workbook(D/'通用容载比推荐矩阵.xlsx')
 # Existing numerical audit sheets retained, formal first sheets use intervals only.
 for name in ['分类推荐区间','线路投资比分析']:
  if name in wb:del wb[name]
 ws=wb.create_sheet('分类推荐区间',0);keys=list(formal[0]);ws.append(keys)
 for r in formal:ws.append([r[k] for k in keys])
 ws=wb.create_sheet('线路投资比分析',1);ws.append(['供电区域类别','线路类型','线路价格倍数','主变与线路单位功率投资比','容载比范围','新增联络项目数/项','平均费用现值/万元'])
 for row in line_summary:ws.append(row)
 for name in ['弹性推荐矩阵','按供电类别与负荷分散系数查表','档位分布']:wb[name].sheet_state='hidden'
 ws=wb['使用说明'];ws.delete_rows(1,ws.max_row);ws.append(['项目','内容'])
 for row in [('推荐入口','分类推荐区间：按供电区域类别、负荷分散系数、源荷、增长、持续时间及两类单位功率投资比选择'),('推荐区间','对应分类全部已算方案的容载比最小值至最大值，显示端点向外保留两位小数，不使用0.20向上归档'),('线路投资比','两类线路分别计算三档价格；本次各价格对应容量相同，费用分析单独列示'),('计算范围','仅归纳已计算条件；边界共享的条件同时查看相邻行'),('负荷分散系数','同范围各站最大负荷之和除以区域同时最大负荷，取1.00和1.15；站间负荷权重随原情景保留'),('历史审查','旧108行分档和432行原始分类作为隐藏审查表保留，正式推荐使用首表')]:ws.append(row)
 for name in ['分类推荐区间','线路投资比分析','使用说明']:
  sh=wb[name];sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
  for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='29465B');c.alignment=Alignment(wrap_text=True,vertical='center')
  for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=24
  for row in sh.iter_rows(min_row=2):
   for c in row:c.alignment=Alignment(wrap_text=True,vertical='center');c.number_format='0.00' if isinstance(c.value,float) else 'General'
 wb.active=0;wb.save(D/'通用容载比推荐矩阵.xlsx')
 plot(rows)
 save(O/'分类推荐生成核对.json',{'状态':'通过','原始点SHA256':hashlib.sha256((D/'外推逐点结果.json').read_bytes()).hexdigest(),'分类数':len(formal),'配对数':len(pairs),'最大容载比变化':max(abs(x['容载比变化']) for x in pairs),'区间宽度中位数':statistics.median(x['容载比最大值']-x['容载比最小值'] for x in formal)})
 print('分类推荐',len(formal),'线路配对',len(pairs))
def plot(rows):
 import matplotlib;matplotlib.use('Agg')
 import matplotlib.pyplot as plt
 from matplotlib import font_manager
 font_manager.fontManager.addfont('/usr/share/fonts/wps-office/FZFSK.TTF');font_manager.fontManager.addfont('/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf')
 plt.rcParams.update({'font.family':['Times New Roman',font_manager.FontProperties(fname='/usr/share/fonts/wps-office/FZFSK.TTF').get_name()],'font.size':10,'axes.unicode_minus':False})
 fig,axes=plt.subplots(2,2,figsize=(8.0,7.2),sharex=True,sharey=True)
 for ax,(area,shape) in zip(axes.flat,itertools.product(AREA,KD)):
  for price,col,off in [(.5,'#9b4b39',-.045),(1.,'#29465b',0),(1.25,'#778f68',.045)]:
   q=[r for r in rows if r['供电区域']==area and r['负荷分布']==shape and r['装机渗透率']==1.8 and r['95%峰值持续时间h']==3 and r['线路价格倍数']==1 and r['储能价格倍数']==price and r['线路类型']=='架空为主'];q.sort(key=lambda r:r['年负荷平均增长率']);k=q[0]['主变储能成本比_3h']
   ax.plot([100*r['年负荷平均增长率'] for r in q],[r['推荐容载比'] for r in q],'o-',ms=3,lw=1.2,label=f'{k:.2f}',color=col)
  from matplotlib.ticker import FormatStrFormatter
  ax.set_title(f'{AREA[area]}；负荷分散系数{KD[shape]:.2f}');ax.grid(alpha=.2);ax.set_ylim(1.98,2.62);ax.set_xticks([0,2,4,7,10]);ax.yaxis.set_major_formatter(FormatStrFormatter('%.2f'));ax.legend(title='主变与储能单位功率投资比',fontsize=8,title_fontsize=8)
 fig.supxlabel('年最大净负荷平均增长率 / %');fig.supylabel('容载比 /（MVA/MW）');fig.tight_layout();fig.savefig(A/'planning_tier_heatmap.png',dpi=300,bbox_inches='tight');plt.close(fig)
if __name__=='__main__':main()
