from pathlib import Path
import json,csv,math,itertools
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
HERE=Path(__file__).resolve().parent;R=HERE.parents[1];IT='当前版'
D=R/'01研究推荐'/IT;A=R/'02图表'/IT;T=R/'03_MD'/IT
BASE=D/'统一起点基准'
tables=json.loads((A/'计算表.json').read_text());sections=json.loads((A/'生成段落.json').read_text())
def table(title,headers,rows):
 return title+'\n\n| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(str(x) for x in r)+' |\n' for r in rows)
def dec(x):return f'{float(x):.2f}'
def readcsv(p):
 with p.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def f(r,k):return float(r[k])
def main():
 from 规划档位矩阵 import main as rebuild_matrix
 rebuild_matrix()
 global tables,sections
 tables=json.loads((A/"计算表.json").read_text());sections=json.loads((A/"生成段落.json").read_text())
 # 逐站段落使用当前已复核年度状态，既有历史版只提供分析组织思路。
 profile_meta=[]
 for region,name in [('pizhou','邳州'),('city','市区')]:
  datasets={scheme:readcsv(BASE/region/(scheme+'_stations.csv')) for scheme in ['rigid','elastic']}
  stations=sorted({x['station'] for x in datasets['elastic']})
  blocks=['### 5.4.'+('4' if region=='pizhou' else '5')+' '+name+'逐站规划方案\n\n'+('邳州各站同时考虑正向供电与反向净负荷。以下结合各站设备结构、负荷转接及储能规模说明弹性方案的配置依据。' if region=='pizhou' else '市区各站按A类供电区域进行规划，站级需求以正向供电和单台主变退出为主。以下结合重载站负荷转出和承接站容量配置说明建设安排。')]
  for station in stations:
   states={sch:sorted([x for x in ds if x['station']==station],key=lambda x:int(x['year'])) for sch,ds in datasets.items()}
   e=states['elastic'][-1];r=states['rigid'][-1];first=states['elastic'][0]
   initial=f(first,'prior_unit_1_mva')+f(first,'prior_unit_2_mva')
   blocks.append('#### '+station+'站规划分析')
   lead=f"该站基期主变配置为{dec(first['prior_unit_1_mva'])}+{dec(first['prior_unit_2_mva'])} MVA，规划第4年转接前最大正向、反向净负荷分别为{dec(e['forward_mw'])} MW和{dec(e['reverse_mw'])} MW，正向95%峰值持续时间为{dec(e['forward_duration_h'])} h。"
   events=[]
   for ix,row in enumerate(states['elastic']):
    acts=[]
    if f(row,'third_transformer_commissioned')>.5:acts.append('扩建50 MVA主变')
    purchase=f(row,'purchased_unit_mva')-(50 if f(row,'third_transformer_commissioned')>.5 else 0)
    if purchase>1e-5:acts.append(f'实施主变增容，新购容量{purchase:.2f} MVA')
    if f(row,'new_storage_power_mw')>1e-5:acts.append(f"新增储能{dec(row['new_storage_power_mw'])} MW/{dec(row['new_storage_energy_mwh'])} MWh")
    if acts:events.append('第'+str(ix+1)+'年'+ '，'.join(acts))
   if events:build='弹性方案'+ '；'.join(events)+'。'
   else:build='弹性方案沿用基期主变容量，后续各年通过站间负荷转接满足容量需求。'
   if f(e,'normal_load_transferred_out_mw')>1e-5:
    role=f"第4年转出{dec(e['normal_load_transferred_out_mw'])} MW、承接{dec(e['normal_load_transferred_in_mw'])} MW后，正向校核负荷为{dec(e['post_transfer_forward_mw'])} MW。负荷转出降低正向需求，同时将反向校核上界提高至{dec(e['post_transfer_reverse_upper_mw'])} MW，配置按两个方向共同确定。"
   elif f(e,'normal_load_transferred_in_mw')>1e-5:
    role=f"第4年承接{dec(e['normal_load_transferred_in_mw'])} MW，正向校核负荷增至{dec(e['post_transfer_forward_mw'])} MW。该站利用相应主变容量承担片区内部负荷，接收能力按承接后的负荷进行正常运行及单台主变退出校核。"
   else:role=f"第4年正向校核负荷为{dec(e['post_transfer_forward_mw'])} MW，设备配置主要承担本站供电与反向输送需求。"
   margin=.95*(f(e,'capacity_mva')-max(f(e,k) for k in ['unit_1_mva','unit_2_mva','unit_3_mva']))+f(e,'storage_energy_mwh')/max(2.15,f(e,'forward_duration_h'))-f(e,'post_transfer_forward_mw')
   comparison=f"刚性方案第4年主变总容量为{dec(r['capacity_mva'])} MVA，储能为{dec(r['storage_power_mw'])} MW；弹性方案相应为{dec(e['capacity_mva'])} MVA和{dec(e['storage_power_mw'])} MW。"
   if f(r,'storage_power_mw')-f(e,'storage_power_mw')>.001:
    comparison+=f"弹性方案储能减少{f(r,'storage_power_mw')-f(e,'storage_power_mw'):.2f} MW，退出工况的容量支撑更多由在役主变和站间负荷配置承担。"
   elif f(e,'capacity_mva')>f(r,'capacity_mva')+.001:comparison+='弹性方案增加主变容量，为本站或片区承接负荷提供支撑。'
   else:comparison+='两类方案的本站主变配置接近，差异主要通过区域其他站点的容量和负荷安排体现。'
   blocks.extend([lead+build,role+comparison])
   tk=region+'_'+station+'_profile';rows_station=[]
   for sch,label in [('rigid','刚性'),('elastic','弹性')]:
    q=states[sch][-1];units='+'.join(dec(q[k]) for k in ['unit_1_mva','unit_2_mva','unit_3_mva'] if f(q,k)>0)
    rows_station.append([label,units,dec(q['storage_power_mw'])+'/'+dec(q['storage_energy_mwh']),dec(q['normal_load_transferred_out_mw']),dec(q['normal_load_transferred_in_mw']),dec(q['post_transfer_forward_mw'])])
   tables[tk]=table('表5-0 '+station+'站第4年配置与负荷转接',['方案','主变组合/MVA','储能/MW/MWh','转出/MW','承接/MW','正向负荷/MW'],rows_station)
   blocks.append('[[TABLE:'+tk+']]')
   profile_meta.append({'区县':name,'站':station,'弹性目标容量MVA':f(e,'capacity_mva'),'刚性目标容量MVA':f(r,'capacity_mva'),'弹性最不利退出余量MW':margin,'年度状态':len(states['elastic'])})
  sections[region+'_station_profiles']='\n\n'.join(blocks)
 p=next(T.glob('05 *.md'));s=p.read_text()
 if '[[SECTION:pizhou_station_profiles]]' not in s:
  s=s.replace('## 5.5 措施作用与工程衔接','[[SECTION:pizhou_station_profiles]]\n\n[[SECTION:city_station_profiles]]\n\n## 5.5 措施作用与工程衔接');p.write_text(s)
 (HERE/'逐站分析对应记录.json').write_text(json.dumps(profile_meta,ensure_ascii=False,indent=2))
 (A/'计算表.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2));(A/'生成段落.json').write_text(json.dumps(sections,ensure_ascii=False,indent=2))
 path=D/'统一起点逐年方案与复核.xlsx';wb=load_workbook(path)
 for ws in wb:
  for row in ws:
   for cell in row:
    if isinstance(cell.value,float):cell.number_format='0.00'
    if isinstance(cell.value,str):cell.value=cell.value.replace('第三台','扩建主变')
 wb.save(path)
 print('推荐矩阵',len(json.loads((D/'通用推荐矩阵.json').read_text())),'行；逐站分析',len(profile_meta),'站')
if __name__=='__main__':main()
