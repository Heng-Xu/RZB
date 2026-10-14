"""从原始记录补齐2025年站级时序，再汇总区域绝对值峰段；不改规划模型。"""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
import hashlib, json
import numpy as np
import pandas as pd
from openpyxl import load_workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.font_manager import fontManager, FontProperties

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT/'研究报告/数据来源/当前版/区域时序'
ASSET = ROOT/'研究报告/02图表/当前版'
RAW = ROOT/'实验/研究/data/tuomin'
SOURCES = {'city': RAW/'补充数据/徐州市区脱敏.zip',
           'pizhou': RAW/'电网建模数据_Agent整合版_V1.2/邳州主变负载率.xlsx'}

def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()

def events(mask):
    edges = np.diff(np.r_[False, mask, False].astype(int))
    return list(zip(map(int, np.flatnonzero(edges == 1)), map(int, np.flatnonzero(edges == -1))))

def main():
    fontManager.addfont('/usr/share/fonts/wps-office/FZFSK.TTF')
    plt.rcParams.update({'font.family': ['Times New Roman', FontProperties(fname='/usr/share/fonts/wps-office/FZFSK.TTF').get_name()], 'axes.unicode_minus': False, 'font.size': 11})
    hours = pd.date_range('2025-01-01', periods=8760, freq='h')
    position = {h: i for i, h in enumerate(hours)}
    summary = {}
    for region, cn in [('pizhou', '邳州'), ('city', '市区')]:
        mapping = pd.read_csv(DATA/f'{region}_mapping.csv')
        source = SOURCES[region]
        if region == 'city':
            mapping = mapping[mapping.in_study_cohort.astype(str).str.lower() == 'true']
            assert (mapping.city_source_sha256 == digest(source)).all()
            ids = sorted(mapping.model_station_id.unique())
            raw = np.full((8760, len(ids)), np.nan)
            with ZipFile(source) as z:
                for _, r in mapping.iterrows():
                    w = load_workbook(BytesIO(z.read(r.archive_member)), read_only=True, data_only=True)
                    sheet = w.active; sheet.reset_dimensions()
                    col = ids.index(r.model_station_id)
                    for cells in sheet.iter_rows(values_only=True):
                        if str(cells[0]).startswith('2025-') and cells[1] is not None:
                            i = position[pd.Timestamp(cells[0])]
                            assert np.isnan(raw[i, col])
                            raw[i, col] = float(cells[1])
                    w.close()
        else:
            mapping = mapping[(mapping.resolved_voltage_kv == 110) & (mapping.station_status == 'supported_for_station_aggregation')]
            assert (mapping.hourly_source_sha256 == digest(source)).all()
            ids = sorted(mapping.resolved_station_id.unique())
            raw = np.full((8760, len(ids)), np.nan)
            cols = {sid: (mapping.loc[mapping.resolved_station_id == sid, 'source_column'].astype(int)-1).tolist() for sid in ids}
            w = load_workbook(source, read_only=True, data_only=True)
            for cells in w['Sheet3'].iter_rows(min_row=3, values_only=True):
                if not str(cells[0]).startswith('2025-'): continue
                i = position[pd.Timestamp(cells[0])]
                for j, sid in enumerate(ids):
                    values = [cells[k] for k in cols[sid]]
                    if all(v not in (None, '') for v in values):
                        assert np.isnan(raw[i,j])
                        raw[i,j] = sum(map(float, values))
            w.close()
        filled = raw.copy(); ledger = []
        for i,j in zip(*np.where(np.isnan(raw))):
            neighbors = []
            if region == 'pizhou':
                neighbors = [i-1,i+1]
                assert i > 0 and i < 8759 and np.isfinite(raw[neighbors,j]).all()
            else:
                for direction in [-1,1]:
                    for weeks in range(1,7):
                        k = i + direction*168*weeks
                        if 0 <= k < 8760 and np.isfinite(raw[k,j]):
                            neighbors.append(k); break
            assert neighbors
            filled[i,j] = raw[neighbors,j].mean()
            ledger.append({'hour': str(hours[i]), 'station': ids[j], 'filled_mw': filled[i,j], 'reference_hours': ';'.join(str(hours[k]) for k in neighbors)})
        assert np.isfinite(filled).all()
        assert np.array_equal(raw[np.isfinite(raw)], filled[np.isfinite(raw)])
        total = filled.sum(axis=1); magnitude = abs(total); peak = magnitude.max()
        mask = magnitude >= .95*peak; runs = events(mask); start,end = max(runs, key=lambda r:r[1]-r[0])
        missing = np.isnan(raw).any(axis=1)
        assert int(missing.sum()) == {'pizhou':1,'city':421}[region]
        pd.DataFrame({'hour':hours, 'net_mw':total, 'abs_mw':magnitude, 'imputed':missing.astype(int), 'near_peak':mask.astype(int)}).to_csv(DATA/f'{region}_8760.csv',index=False)
        pd.DataFrame(ledger).to_csv(DATA/f'{region}_fill_log.csv',index=False)
        np.savez_compressed(DATA/f'{region}_stations.npz', power_mw=filled, observed=np.isfinite(raw), station_ids=np.array(ids))
        item = {'region':cn, 'hours':8760, 'imputed_hours':int(missing.sum()), 'peak_mw':float(peak), 'threshold_mw':float(.95*peak), 'peak_direction':'正向' if total[magnitude.argmax()] > 0 else '反向', 'max_duration_h':end-start, 'cumulative_h':int(mask.sum()), 'longest_start':str(hours[start]), 'longest_end_exclusive':str(hours[end-1]+pd.Timedelta(hours=1)), 'events':[{'start':str(hours[a]), 'end_exclusive':str(hours[b-1]+pd.Timedelta(hours=1)), 'hours':b-a} for a,b in runs], 'diversity':float(filled.clip(min=0).max(axis=0).sum()/total.max()), 'station_peak_sum_mw':float(filled.clip(min=0).max(axis=0).sum()), 'source_sha256':digest(source), 'mapping_sha256':digest(DATA/f'{region}_mapping.csv'), 'curve_sha256':digest(DATA/f'{region}_8760.csv'), 'preserved_observations':True}
        summary[region] = item
        fig,axes = plt.subplots(2,1,figsize=(8.4,5.8),gridspec_kw={'height_ratios':[1,1]})
        for ax in axes:
            ax.axhline(.95*peak,color='#ac4b33',ls='--',lw=1.2,label=f'95%峰值 {peak*.95:.2f} MW')
            ax.spines[['top','right']].set_visible(False); ax.set_ylabel('区域净负荷绝对值/MW'); ax.grid(alpha=.15)
        axes[0].plot(hours,magnitude,color='#315c78',lw=.45,label='全年小时曲线')
        axes[0].scatter(hours[mask],magnitude[mask],s=12,color='#ac4b33',zorder=3)
        axes[0].xaxis.set_major_locator(mdates.MonthLocator(interval=2)); axes[0].xaxis.set_major_formatter(mdates.DateFormatter('%m月'))
        axes[0].set_xlim(hours[0],hours[-1]); axes[0].legend(loc='lower left',fontsize=9,ncol=2)
        lo=max(0,start-18);hi=min(8760,end+18)
        axes[1].step(hours[lo:hi],magnitude[lo:hi],where='post',color='#315c78',lw=1.2)
        axes[1].axvspan(hours[start],hours[end-1]+pd.Timedelta(hours=1),alpha=.18,color='#ac4b33')
        axes[1].set_xlim(hours[lo],hours[hi-1]);axes[1].xaxis.set_major_formatter(mdates.DateFormatter('%m-%d\n%H:%M'))
        axes[1].set_title(f'最长近峰时段：{hours[start]:%m月%d日 %H:%M}—{hours[end-1]+pd.Timedelta(hours=1):%m月%d日 %H:%M}（{end-start} h）',fontsize=11)
        axes[1].legend(fontsize=9,loc='lower left')
        fig.tight_layout();fig.savefig(ASSET/f'{region}_regional_peak.png',dpi=240,bbox_inches='tight');plt.close(fig)
    (DATA/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__ == '__main__': main()
