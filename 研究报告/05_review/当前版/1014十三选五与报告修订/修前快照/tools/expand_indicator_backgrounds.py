"""按原六参数背景域扩大补充指标样本，仅导出结果和核验资料。"""
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor
from functools import lru_cache
from collections import Counter
import argparse
import copy
import gzip
import hashlib
import itertools
import json
import math
import sys
import time
from types import FunctionType
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import extend_indicator_analysis as first
import locked_parameter_results as baseline

ROOT = first.ROOT
TRIAL = ROOT/'研究报告/试算/20261014_补充指标全背景扩展'
PUBLIC = ROOT/'研究报告/01研究推荐/当前版/补充指标全背景扩展_20261014'
SOURCE = ROOT/'研究报告/数据来源/当前版/补充指标全背景扩展_20261014'
REVIEW = ROOT/'研究报告/05_review/当前版/1014补充指标全背景扩展'
RAW = TRIAL/'原始求解记录.jsonl.gz'
BOOK = PUBLIC/'补充指标与相关性分析_全背景扩展.xlsx'
SMALL_RAW = first.RAW
SMALL_PUBLIC = first.PUBLIC
PRICE_NODES = [.7,.8,.9,1.,1.1,1.2,1.3]
CALIBRATE = first.model.locked.calibrate_initial_network
SMALL_JOBS = first.tasks()
SMALL_SPEC = first.design()


@lru_cache(maxsize=64)
def cached_network(base, pairs, target, pf):
    return CALIBRATE(np.array(base), list(pairs), target, pf)


def enable_network_cache():
    def same_network(base,pairs,target,pf):
        edges,thermal,scale,audit=cached_network(tuple(base),tuple(pairs),target,pf)
        return [dict(e) for e in edges],thermal.copy(),scale,copy.deepcopy(audit)
    first.model.locked.calibrate_initial_network=same_network


def solve_task(parameters):
    try:
        return first.solve_task(parameters)
    except AssertionError:
        # 保持原检查容差，重算由于整数/原始可行性数值误差超差的节点。
        normal=first.model._solve
        real=normal.__globals__['highspy']
        settings=dict(mip_feasibility_tolerance=1e-9,primal_feasibility_tolerance=1e-9,
                      dual_feasibility_tolerance=1e-9)
        class PrecisionHighs:
            def Highs(self):
                solver=real.Highs()
                for name,value in settings.items():solver.setOptionValue(name,value)
                return solver
            def __getattr__(self,name):return getattr(real,name)
        first.model._solve=FunctionType(normal.__code__,dict(normal.__globals__,highspy=PrecisionHighs()),
            normal.__name__,normal.__defaults__,normal.__closure__)
        try:
            row=first.solve_task(parameters)
            row['precision_retry']=settings
            return row
        finally:
            first.model._solve=normal


def backgrounds():
    return [dict(growth_rate=g,cable_fraction=cable,initial_transfer_fraction=transfer,diversity=kd)
            for g,(cable,transfer),kd in itertools.product(
                baseline.OTHER['growth_rate'],baseline.allowed_pairs(),baseline.OTHER['diversity'])]


def energy_nodes(eta):
    upper=first.model._curve(eta,1.,4)[1].sum()/first.model.LOAD_HOURS
    return [i/40 for i in range(1,27)
            if .12-1e-10<=i/40/eta<=.40+1e-10 and i/40<=upper+1e-10]


@lru_cache(maxsize=1)
def tasks():
    # 原3301节点不删减；所有重复参数组合只保留一次。
    jobs={first.key(p):dict(p) for p in SMALL_JOBS}
    def add(background,eta,energy,h,purpose,**prices):
        p=dict(background,eta=eta,energy_penetration=round(energy,10),h95_h=h,
               line_price=1.,storage_price=1.)
        p.update(prices)
        k=first.key(p)
        if k not in jobs:jobs[k]=dict(p,node_purpose=purpose)
    for b in backgrounds():
        for eta in [1.,2.,3.]:
            for energy in energy_nodes(eta):
                add(b,eta,energy,4,'全背景电量完整序列')
            # 两种供能水平分别比较八个连续时长，扩大峰段的背景覆盖。
            for energy in [.15*eta,.20*eta]:
                for h in range(1,9):add(b,eta,energy,h,'全背景双供能峰段序列')
            # 在每种装机和每个背景下独立改变两类价格。
            for driver in ['line_price','storage_price']:
                for price in PRICE_NODES:
                    add(b,eta,.20*eta,4,'全背景相对投资完整序列',**{driver:price})
    return list(jobs.values())


def design():
    jobs=tasks()
    spec=copy.deepcopy(SMALL_SPEC)
    spec.update(样本版本='full_background_extension_20261014',设计节点数=len(jobs),
        背景组合数=len(backgrounds()),原版背景节点=dict(
            年最大净负荷平均增长率=baseline.OTHER['growth_rate'],
            电缆及转供联合节点=baseline.allowed_pairs(),负荷分散系数=baseline.OTHER['diversity']),
        电量完整序列=dict(装机节点=[1.,2.,3.],电量节点={str(e):energy_nodes(e) for e in [1.,2.,3.]},峰段_h=4),
        峰段完整序列=dict(装机节点=[1.,2.,3.],电量与装机之比=[.15,.20],时间_h=list(range(1,9))),
        投资完整序列=dict(装机节点=[1.,2.,3.],电量与装机之比=.20,峰段_h=4,价格倍率=PRICE_NODES),
        前次设计节点数=len(SMALL_JOBS),前次有效节点数=3214,
        新增采样原则='交叉原版888种增长、线路、转供及分散背景；各指标保留完整受控序列；互异节点等权',
        原始库格式='gzip JSONL；包含原始可行及不可行结果、设备、费用、间隙和物理复算',
        任务SHA256=hashlib.sha256(json.dumps(jobs,sort_keys=True).encode()).hexdigest(),
        扩展入口SHA256=first.sha(Path(__file__)),
        前次原始库SHA256=first.sha(SMALL_RAW),
        性能缓存='复用同基期设备与转供率的网络标定，不改变标定算法或数值')
    return json.loads(json.dumps(spec))


def read_rows():
    with gzip.open(RAW,'rt',encoding='utf-8') as stream:return [json.loads(line) for line in stream]


def preserved_paths():
    paths=list((ROOT/'研究报告/03_MD/当前版').glob('*.md'))
    paths += [ROOT/p for p in ['研究报告/02图表/当前版/计算表.json','研究报告/02图表/当前版/生成段落.json',
        '研究报告/04_word/研究报告终稿.docx','研究报告/04_word/研究报告终稿.pdf',
        '研究报告/甲方反馈/10.11/供电公司讨论稿.docx',
        '研究报告/甲方反馈/10.11/表5-4及表5-6_新旧参数对照.xlsx',
        '实验/研究/rebuild_2026/current_economic_assumptions.py',str(first.LOCKED_FILE.relative_to(ROOT)),
        str(first.MODEL_FILE.relative_to(ROOT)),str(SMALL_RAW.relative_to(ROOT)),
        str((SMALL_PUBLIC/'补充指标与相关性分析.xlsx').relative_to(ROOT)),
        str((baseline.TRIAL/'统一原始求解记录.jsonl').relative_to(ROOT)),
        str((baseline.PUBLIC/'推广矩阵与相关性分析_最新收资统一结果.xlsx').relative_to(ROOT)),
        str((baseline.PUBLIC/'统一有效样本.csv').relative_to(ROOT)),
        str((baseline.PUBLIC/'固定其余五项后的组内相关.csv').relative_to(ROOT))]]
    return [p for p in paths if p.is_file()]


def solve(workers,limit=None):
    TRIAL.mkdir(parents=True,exist_ok=True)
    spec=design();path=TRIAL/'试算设计.json'
    if path.exists():assert json.loads(path.read_text())==spec,'设计变化须新建版本'
    else:first.dump(path,spec)
    snapshot=REVIEW/'本轮未修改文件登记.json'
    if not snapshot.exists():first.dump(snapshot,{str(p.relative_to(ROOT)):first.sha(p) for p in preserved_paths()})
    old=[json.loads(line) for line in SMALL_RAW.read_text().splitlines()]
    prior=read_rows() if RAW.exists() else []
    assert len({first.key(r) for r in prior})==len(prior)
    expected={first.key(p):p for p in tasks()}
    assert all(first.key(r) in expected for r in prior)
    existing={first.key(r) for r in prior}
    reused=[]
    for r in old:
        k=first.key(r)
        if k not in existing:
            row=dict(r,record_origin='前次同模型节点复用',sample_pool='full_background_extension_20261014')
            assert row['node_purpose']==expected[k]['node_purpose']
            reused.append(row);existing.add(k)
    with gzip.open(RAW,'at',encoding='utf-8',compresslevel=6) as stream:
        for r in reused:stream.write(json.dumps(r,ensure_ascii=False,allow_nan=False)+'\n')
    todo=[p for p in tasks() if first.key(p) not in existing]
    if limit:todo=todo[:limit]
    print('完整设计',len(expected),'复用',len(reused),'已有',len(prior),'待算',len(todo),flush=True)
    start=time.monotonic()
    with ProcessPoolExecutor(max_workers=workers,initializer=enable_network_cache) as pool, \
            gzip.open(RAW,'at',encoding='utf-8',compresslevel=6) as stream:
        for i,r in enumerate(pool.map(solve_task,todo,chunksize=16),1):
            r.update(record_origin='全背景独立求解',sample_pool='full_background_extension_20261014')
            stream.write(json.dumps(r,ensure_ascii=False,allow_nan=False)+'\n')
            if i%500==0 or i==len(todo):
                stream.flush()
                print('新增',i,'/',len(todo),'累计',len(existing)+i,'用时',round(time.monotonic()-start,1),'s',flush=True)


def frames():
    rows=read_rows()
    d=pd.DataFrame([r for r in rows if r['verified']]).sort_values(first.model.INPUTS).reset_index(drop=True)
    # 采用声明的节点值分组，消除0.2×3与0.6的浮点表示差异；原始库保留完整值。
    d[first.model.INPUTS]=d[first.model.INPUTS].round(10)
    extra=pd.DataFrame([first.controlled(d,f,driver) for f,driver in first.NEW])
    cm=d[first.METRICS].corr().rename(index=first.LABELS,columns=first.LABELS).rename_axis('指标').reset_index()
    sm=d[first.METRICS].corr(method='spearman').rename(index=first.LABELS,columns=first.LABELS).rename_axis('指标').reset_index()
    samples=d[first.METRICS+['line_price','storage_price','clr','capacity_mva','storage_mw','storage_mwh',
        'storage_effective_mw','line_count','cost','relative_gap','node_purpose']].rename(columns=first.LABELS|dict(
        clr='容载比',capacity_mva='主变容量_MVA',storage_mw='储能额定功率_MW',storage_mwh='储能额定能量_MWh',
        storage_effective_mw='储能有效功率_MW',line_count='新增联络_条',cost='费用现值_万元',
        line_price='线路价格倍率',storage_price='储能价格倍率',relative_gap='相对求解间隙',node_purpose='节点用途'))
    small=pd.read_csv(SMALL_PUBLIC/'新增指标组内相关.csv')
    compare=small[['指标','有效组数','有效节点数','组内Pearson','组内Spearman']].merge(
        extra[['指标','有效组数','有效节点数','组内Pearson','组内Spearman']],on='指标',suffixes=('_前次','_全背景'))
    compare['Pearson变化']=compare['组内Pearson_全背景']-compare['组内Pearson_前次']
    original=pd.read_csv(baseline.PUBLIC/'固定其余五项后的组内相关.csv')
    return d,dict(新增指标组内相关=extra,与前次补充结果对照=compare,原六参数组内相关=original,
        十三项指标Pearson=cm,十三项指标Spearman=sm,有效样本=samples)


def select_lp(d):
    selected={int(d[f].idxmin()) for f in first.METRICS}|{int(d[f].idxmax()) for f in first.METRICS}
    # 按各原版背景各取一个有效节点，再覆盖每个输入节点与规则间隔。
    fields=['growth_rate','cable_fraction','initial_transfer_fraction','diversity']
    for _,g in d.groupby(fields,sort=True):selected.add(int(g.index[len(g)//2]))
    for f in first.model.INPUTS:
        for _,g in d.groupby(f,sort=True):selected.add(int(g.index[len(g)//2]))
    selected.update(int(i) for i in np.linspace(0,len(d)-1,512))
    return sorted(selected)


def setup_readonly_scope():
    # 复用前次经过核验的设备、费用和全年曲线检查程序，改为读取独立扩展目录。
    first.TRIAL,first.PUBLIC,first.SOURCE,first.REVIEW,first.RAW=TRIAL,PUBLIC,SOURCE,REVIEW,RAW
    first.tasks,first.design,first.read_rows,first.frames=tasks,design,read_rows,frames


def export():
    rows=read_rows();expected={first.key(p) for p in tasks()}
    assert len(rows)==len(expected) and {first.key(r) for r in rows}==expected
    d,fs=frames();PUBLIC.mkdir(parents=True,exist_ok=True)
    overview=pd.DataFrame([dict(项目=k,值=v) for k,v in dict(
        设计组合=len(rows),有效组合=len(d),不可行组合=len(rows)-len(d),
        原版背景组合=len(backgrounds()),前次有效组合=3214,
        模型='同一静态正常/N-1费用比选',正文='本轮仅交付数据，正文修正待讨论').items()])
    fs={'结果概览':overview,**fs}
    for name,frame in fs.items():frame.to_csv(PUBLIC/(name+'.csv'),index=False,encoding='utf-8-sig')
    book=Workbook(write_only=True)
    from openpyxl.cell import WriteOnlyCell
    for name,frame in fs.items():
        sheet=book.create_sheet(name);sheet.freeze_panes='B2'
        for i,column in enumerate(frame.columns,1):
            sheet.column_dimensions[get_column_letter(i)].width=min(38,max(14,len(str(column))*1.6))
        cells=[]
        for title in frame.columns:
            c=WriteOnlyCell(sheet,value=str(title));c.font=Font(name='宋体',bold=True,color='FFFFFF')
            c.fill=PatternFill('solid',fgColor='315C78');c.alignment=Alignment(wrap_text=True,vertical='center')
            cells.append(c)
        sheet.append(cells);sheet.row_dimensions[1].height=32
        for row in frame.itertuples(index=False,name=None):
            sheet.append([None if pd.isna(v) else v for v in row])
        sheet.auto_filter.ref=f'A1:{get_column_letter(len(frame.columns))}{len(frame)+1}'
    book.save(BOOK)
    original_source=json.loads((ROOT/'研究报告/数据来源/当前版/补充指标相关性_20261014/参数范围及来源.json').read_text())
    first.dump(SOURCE/'参数范围及来源.json',dict(
        前次指标及曲线依据=original_source,扩大设计=design(),
        背景来源=dict(原版锁定采样入口='tools/locked_parameter_results.py',sha256=first.sha(ROOT/'tools/locked_parameter_results.py')),
        独立输入='源荷、增长、线路、转供、峰段、分散及两类价格；正反向功率和负载率为计算指标',
        分组='按声明节点编码固定其余输入，组内去均值，节点等权，保留零容量响应组'))
    pd.DataFrame(backgrounds()).rename(columns=first.LABELS).to_csv(SOURCE/'原版888种背景组合.csv',index=False,encoding='utf-8-sig')
    print('工作簿已导出',len(d),'有效节点',flush=True)


def verify(lp=False,workers=8):
    enable_network_cache()
    d,fs=frames();selected=select_lp(d);path=REVIEW/'全背景分层独立LP核验.json'
    if lp:
        jobs=[d.loc[i].to_dict() for i in selected]
        with ProcessPoolExecutor(max_workers=workers,initializer=enable_network_cache) as pool:
            proof=list(pool.map(first.lp_task,jobs,chunksize=4))
        first.dump(path,proof)
    proof=json.loads(path.read_text())
    assert {tuple(p['节点']) for p in proof}=={first.key(d.loc[i]) for i in selected}
    assert all(p['verified'] for p in proof)
    # 前次通用核验覆盖所有节点的曲线、容量、费用、CSV和工作簿，另保留更广的分层LP。
    old_selected={int(d[f].idxmin()) for f in first.METRICS}|{int(d[f].idxmax()) for f in first.METRICS}
    old_selected.update(int(i) for i in np.linspace(0,len(d)-1,128))
    by_key={tuple(p['节点']):p for p in proof}
    # 将通用核验要求的128个节点另算，避免依赖分层抽取恰好重合。
    missing=[d.loc[i].to_dict() for i in sorted(old_selected) if first.key(d.loc[i]) not in by_key]
    if missing:
        with ProcessPoolExecutor(max_workers=workers,initializer=enable_network_cache) as pool:
            for p in pool.map(first.lp_task,missing,chunksize=4):by_key[tuple(p['节点'])]=p
    first.dump(REVIEW/'独立LP核验.json',[by_key[first.key(d.loc[i])] for i in sorted(old_selected)])
    setup_readonly_scope()
    # 通用入口工作簿名不同，使用同目录的校验副本，文件字节与交付件一致。
    alias=PUBLIC/'补充指标与相关性分析.xlsx'
    if not alias.exists() or first.sha(alias)!=first.sha(BOOK):
        import shutil
        shutil.copy2(BOOK,alias)
    result=first.verify(lp=False,workers=workers)
    for name,digest in json.loads((REVIEW/'本轮未修改文件登记.json').read_text()).items():
        assert first.sha(ROOT/name)==digest,'本轮范围外文件发生变化:'+name
    rows=read_rows()
    assert sum(r['record_origin']=='前次同模型节点复用' for r in rows)==3301
    # 每个原版背景确实进入设计，组数来自互异组合，不靠重复节点增量。
    bg=['growth_rate','cable_fraction','initial_transfer_fraction','diversity']
    present={tuple(round(float(r[f]),10) for f in bg) for r in rows}
    assert {tuple(b[f] for f in bg) for b in backgrounds()}<=present
    result.update(原版背景组合数=len(backgrounds()),前次节点复用数=3301,
        全背景分层独立LP节点数=len(proof),全背景分层独立LP工况数=sum(p['checked_pairs'] for p in proof),
        正文及既有成果='本轮哈希未变',工作簿SHA256=first.sha(BOOK),原始库SHA256=first.sha(RAW))
    first.dump(REVIEW/'数据核验.json',result)
    protected=json.loads((REVIEW/'本轮未修改文件登记.json').read_text())
    paths=list(PUBLIC.glob('*.csv'))+[BOOK,RAW,TRIAL/'试算设计.json',SOURCE/'参数范围及来源.json',
        SOURCE/'原版888种背景组合.csv',Path(__file__)]
    first.dump(REVIEW/'结果文件登记.json',dict(状态='通过',结果文件={str(p.relative_to(ROOT)):first.sha(p) for p in paths},范围外文件=protected))
    return result


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--solve',action='store_true');ap.add_argument('--export',action='store_true')
    ap.add_argument('--verify',action='store_true');ap.add_argument('--lp',action='store_true')
    ap.add_argument('--workers',type=int,default=8);ap.add_argument('--limit',type=int)
    args=ap.parse_args()
    if args.solve:solve(args.workers,args.limit)
    if args.export:export()
    if args.verify or args.lp:print(json.dumps(verify(args.lp,args.workers),ensure_ascii=False,indent=2),flush=True)
    if not any([args.solve,args.export,args.verify,args.lp]):print(json.dumps(design(),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
