"""补充指标的声明设计、断点续算、原始库复核及报告资产；锁定库只读。"""
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor
import argparse
import hashlib
import itertools
import json
import math
import sys
import time
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from 实验.研究.rebuild_2026.indicator_extension_20261014 import model

TRIAL = ROOT/'研究报告/试算/20261014_补充指标相关性'
PUBLIC = ROOT/'研究报告/01研究推荐/当前版/补充指标相关性_20261014'
SOURCE = ROOT/'研究报告/数据来源/当前版/补充指标相关性_20261014'
REVIEW = ROOT/'研究报告/05_review/当前版/1014补充指标相关性'
RAW = TRIAL/'原始求解记录.jsonl'
MODEL_FILE = Path(model.__file__)
LOCKED_FILE = Path(model.locked.__file__)
LABELS = dict(eta='分布式电源装机渗透率', energy_penetration='分布式电源电量渗透率',
    growth_rate='年最大净负荷平均增长率', cable_fraction='电缆长度占比',
    initial_transfer_fraction='初始转供率', support_duration_h='储能支撑持续时间_h',
    h95_h='95%峰值持续时间_h', diversity='负荷分散系数',
    initial_max_loading='站级负载率最大值', forward_peak_mw='最大正向净负荷_MW',
    reverse_peak_mw='最大反向净负荷_MW', transformer_line_ratio='主变与线路单位功率投资比',
    transformer_storage_ratio='主变与储能单位功率投资比')
METRICS = ['eta', 'energy_penetration', 'initial_max_loading', 'forward_peak_mw',
           'reverse_peak_mw', 'growth_rate', 'h95_h', 'diversity',
           'initial_transfer_fraction', 'transformer_line_ratio', 'transformer_storage_ratio',
           'cable_fraction', 'support_duration_h']
SHORT = ['ηC','ηE','βmax','P+','P−','ḡ','H95','KD','r₀','RTL','RTB','α','H']
NEW = [('energy_penetration','energy_penetration'), ('h95_h','h95_h'),
       ('transformer_line_ratio','line_price'), ('transformer_storage_ratio','storage_price'),
       ('forward_peak_mw','growth_rate'), ('reverse_peak_mw','energy_penetration'),
       ('initial_max_loading','energy_penetration')]


def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def dump(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n')
def key(row): return tuple(round(float(row[f]), 10) for f in model.INPUTS)


def tasks():
    jobs = {}
    def add(p, purpose):
        p = {
            **dict(eta=2., energy_penetration=.4, growth_rate=.04, cable_fraction=.1,
            initial_transfer_fraction=.4, h95_h=4, diversity=1.1, line_price=1., storage_price=1.), **p}
        if not model.locked.matching_scenarios(p): return
        eta, energy = p['eta'], p['energy_penetration']
        if eta == 0 and energy != 0: return
        if eta > 0 and (energy <= 0 or not .12-1e-10 <= energy/eta <= .40+1e-10): return
        # 在求解前按装机和逐时出力上限判断源荷组合。
        _, upper, _ = model._curve(eta, 1., int(p['h95_h']))
        if energy > upper.sum()/model.LOAD_HOURS+1e-10: return
        k = key(p)
        if k not in jobs: jobs[k] = dict(p, node_purpose=purpose)
    # 装机、电量及峰段密集交叉；首尾包含无光伏节点。
    for eta, energy, h in itertools.product(np.arange(25)/8, np.arange(27)/40, range(1,9)):
        add(dict(eta=float(eta), energy_penetration=float(energy), h95_h=h), '源荷峰段密集交叉')
    # 在不同增长、线路、转供及分散背景下重复完整受控序列。
    for eta, g, (cable, transfer), kd in itertools.product(
        [1.,2.,3.], [0.,.04,.08], [(.1,.2),(.5,.5),(.9,.8)], [1.,1.15]):
        base = dict(eta=eta, energy_penetration=min(.20*eta,.6), growth_rate=g,
                    cable_fraction=cable, initial_transfer_fraction=transfer, diversity=kd)
        for energy in np.arange(1,27)/40:
            add(dict(base, energy_penetration=float(energy)), '多背景电量序列')
        for h in range(1,9): add(dict(base, h95_h=h), '多背景峰段序列')
        for field in ['line_price','storage_price']:
            for multiplier in [.7,.8,.9,1.,1.1,1.2,1.3]:
                add(dict(base, **{field:multiplier}), '多背景相对投资序列')
    # 每项既有输入在补充库中也具有独立受控序列。
    for field, nodes in dict(growth_rate=[0,.01,.02,.03,.04,.05,.055,.06,.07,.075,.08,.09,.1],
            cable_fraction=[.1,.3,.5,.8,.9], initial_transfer_fraction=np.arange(1,10)/10,
            diversity=[1.,1.05,1.1,1.15]).items():
        for value in nodes: add({field:float(value)}, '既有参数受控衔接')
    return list(jobs.values())


def design():
    jobs = tasks()
    return dict(版本='indicator_extension_20261014', 设计节点数=len(jobs), 输入字段=model.INPUTS,
        指标字段=LABELS, 饱和负荷密度='按用户要求不纳入',
        取值=dict(装机渗透率='0～3，步长0.125', 电量渗透率='0～65%，步长2.5个百分点',
            峰值持续时间_h=list(range(1,9)), 相对价格倍率=[.7,.8,.9,1.,1.1,1.2,1.3],
            利用水平关系='0.12≤电量渗透率/装机渗透率≤0.40，并校核逐时出力不超过装机'),
        时序='8760小时合成源荷曲线核定电量和95%绝对峰段；规划采用对应极值静态截面',
        支撑时间='补充情景取H=H95，采用既有静态能量折算',
        负载率='初始每站100 MVA，0.95功率因数，取五站最大值',
        价格='主变工程单价固定；线路与储能分别比较基准价格的70%～130%，为研究情景',
        主变投资分母='新购铭牌容量折算的有功功率；扩建工程仍按1223万元/项计费',
        原模型SHA256=sha(LOCKED_FILE), 扩展映射SHA256=sha(MODEL_FILE),
        任务SHA256=hashlib.sha256(json.dumps(jobs,sort_keys=True).encode()).hexdigest())


def solve_task(p):
    row = model.solve(dict(p, gap=1e-7, seconds=30.))
    if not row['verified'] and row['status'] != 'Infeasible':
        row = model.solve(dict(p, gap=1e-7, seconds=120.))
    assert row['verified'] or row['status'] == 'Infeasible', (key(row),row['status'])
    return row


def solve(workers, limit=None):
    TRIAL.mkdir(parents=True, exist_ok=True)
    spec = design(); declaration = TRIAL/'试算设计.json'
    if declaration.exists(): assert json.loads(declaration.read_text()) == spec, '设计变化须另建版本'
    else: dump(declaration,spec)
    previous = [json.loads(line) for line in RAW.read_text().splitlines()] if RAW.exists() else []
    assert len({key(r) for r in previous}) == len(previous)
    done = {key(r) for r in previous}
    jobs = [p for p in tasks() if key(p) not in done]
    if limit: jobs = jobs[:limit]
    print('待算',len(jobs),'已有',len(previous),'完整设计',spec['设计节点数'],flush=True)
    start = time.monotonic()
    with ProcessPoolExecutor(max_workers=workers) as pool, RAW.open('a') as stream:
        for i,r in enumerate(pool.map(solve_task,jobs,chunksize=4),1):
            stream.write(json.dumps(r,ensure_ascii=False,allow_nan=False)+'\n');stream.flush()
            if i%100 == 0 or i == len(jobs):
                print('新增',i,'/',len(jobs),'用时',round(time.monotonic()-start,1),'s',flush=True)


def read_rows():
    return [json.loads(line) for line in RAW.read_text().splitlines()]


def controlled(d, field, driver):
    fixed = [f for f in model.INPUTS if f != driver]
    n = d.groupby(fixed,dropna=False)[driver].transform('nunique')
    sub = d[n >= 2].copy()
    gx = sub.groupby(fixed)[field].transform('mean')
    gy = sub.groupby(fixed).clr.transform('mean')
    x = sub[field]-gx; y = sub.clr-gy
    den = float(np.sqrt((x@x)*(y@y)))
    pearson = float(x@y/den) if den > 1e-12 else None
    # 先组内排序再去组均值，保持与组内Pearson相同的受控范围。
    xr = sub.groupby(fixed)[field].rank(method='average')
    yr = sub.groupby(fixed).clr.rank(method='average')
    xrank = xr-xr.groupby([sub[f] for f in fixed]).transform('mean')
    yrank = yr-yr.groupby([sub[f] for f in fixed]).transform('mean')
    rankden = float(np.sqrt((xrank@xrank)*(yrank@yrank)))
    spearman = float(xrank@yrank/rankden) if rankden > 1e-12 else None
    spans = sub.groupby(fixed).capacity_mva.agg(lambda a:float(a.max()-a.min()))
    return dict(指标=LABELS[field], 受控变化项=driver, 有效组数=len(spans), 有效节点数=len(sub),
        组内Pearson=pearson, 组内Spearman=spearman,
        主变容量响应组占比=float((spans>1e-7).mean()),
        统计说明='固定其余独立输入，组内去均值、节点等权；零容量响应组保留')


def frames():
    rows=read_rows(); d=pd.DataFrame([r for r in rows if r['verified']]).sort_values(model.INPUTS).reset_index(drop=True)
    correlations=pd.DataFrame([controlled(d,f,driver) for f,driver in NEW])
    cm=d[METRICS].corr().rename(index=LABELS,columns=LABELS).rename_axis('指标').reset_index()
    sm=d[METRICS].corr(method='spearman').rename(index=LABELS,columns=LABELS).rename_axis('指标').reset_index()
    samples=d[METRICS+['line_price','storage_price','clr','capacity_mva','storage_mw','storage_mwh',
        'storage_effective_mw','line_count','cost','relative_gap','node_purpose']].rename(columns=LABELS|dict(
        clr='容载比',capacity_mva='主变容量_MVA',storage_mw='储能额定功率_MW',storage_mwh='储能额定能量_MWh',
        storage_effective_mw='储能有效功率_MW',line_count='新增联络_条',cost='费用现值_万元',
        line_price='线路价格倍率',storage_price='储能价格倍率',relative_gap='相对求解间隙',node_purpose='节点用途'))
    return d,dict(有效样本=samples,新增指标组内相关=correlations,十三项指标Pearson=cm,十三项指标Spearman=sm)


def verify(lp=False, workers=8):
    assert json.loads((TRIAL/'试算设计.json').read_text())==design()
    rows=read_rows();expected={key(p):p for p in tasks()}
    assert len(rows)==len(expected) and {key(r) for r in rows}==set(expected)
    assert len({key(r) for r in rows})==len(rows)
    checked=0;worst=0.;certificates={}
    for r in rows:
        for f in model.INPUTS:assert math.isclose(r[f],expected[key(r)][f],abs_tol=1e-10)
        inp=model.make_input(r); cert=inp['profile_certificate']
        assert cert==r['profile_certificate']
        certificates[(r['eta'],r['energy_penetration'],r['h95_h'])]=cert
        for f,v in model.indicators(inp).items():assert math.isclose(r[f],v,abs_tol=1e-9)
        assert r['support_duration_h']==r['h95_h']
        if not r['verified']:
            assert r['status']=='Infeasible';continue
        units=np.array(r['units']);modules=np.array(r['modules'])
        assert abs(units.sum()-r['capacity_mva'])<1e-8
        assert abs(units.sum()/inp['peak']-r['clr'])<1e-10
        raw=dict(transformer=float(np.where(units[:,:2]>50,units[:,:2],0).sum()*inp['transformer_cost_per_purchased_mva']
                   +units[:,2].sum()/50*inp['third_cost']),
                 storage=float(modules.sum()*inp['storage_unit_cost']*r['storage_price']),line=len(r['built'])*inp['line_cost'])
        for f,cost in raw.items():assert abs(cost*inp['factors'][f]-r['cost_components'][f])<1e-5
        assert abs(sum(r['cost_components'].values())-r['cost'])<1e-4
        assert abs(modules.sum()*.1-r['storage_mw'])<1e-9
        assert abs(modules.sum()*.215-r['storage_mwh'])<1e-9
        assert abs(modules.sum()*inp['storage_effective']-r['storage_effective_mw'])<1e-9
        assert r['physical_check']['violating_pairs']==0
        checked+=r['physical_check']['checked_pairs'];worst=max(worst,r['physical_check']['max_violation_mw'])
    d,fs=frames()
    # 各指标两端及规则抽取的节点另由LP检查，避免只检查优化器约束残差。
    selected={int(d[f].idxmin()) for f in METRICS}|{int(d[f].idxmax()) for f in METRICS}
    selected.update(int(i) for i in np.linspace(0,len(d)-1,128))
    proof_path=REVIEW/'独立LP核验.json'
    if lp:
        jobs=[d.loc[i].to_dict() for i in sorted(selected)]
        with ProcessPoolExecutor(max_workers=workers) as pool:proof=list(pool.map(lp_task,jobs))
        dump(proof_path,proof)
    proof=json.loads(proof_path.read_text())
    assert {tuple(p['节点']) for p in proof}=={key(d.loc[i]) for i in selected}
    assert all(p['verified'] for p in proof)
    for name,frame in fs.items():
        saved=pd.read_csv(PUBLIC/(name+'.csv'))
        pd.testing.assert_frame_equal(frame,saved,check_dtype=False,check_exact=False,rtol=1e-11,atol=1e-9)
    book=load_workbook(PUBLIC/'补充指标与相关性分析.xlsx',read_only=True,data_only=True)
    cell_count=0
    for name,frame in fs.items():
        sheet=book[name];iterator=sheet.iter_rows(values_only=True)
        assert list(next(iterator))==list(frame.columns)
        values=list(iterator);assert len(values)==len(frame)
        for a,b in zip(values,frame.itertuples(index=False,name=None)):
            for av,bv in zip(a,b):
                if pd.isna(bv):assert av is None
                elif isinstance(bv,(float,int,np.number)):assert math.isclose(av,bv,rel_tol=1e-11,abs_tol=1e-9)
                else:assert av==bv
                cell_count+=1
    book.close()
    result=dict(状态='通过',设计节点=len(rows),有效节点=len(d),不可行节点=len(rows)-len(d),
        独立全年源荷曲线=len(certificates),全年小时数=8760,物理复算工况=checked,
        最大物理残差_MW=worst,独立LP节点=len(proof),独立LP工况=sum(p['checked_pairs'] for p in proof),
        工作簿回读单元格=cell_count,最大相对求解间隙=float(d.relative_gap.max()),
        原模型SHA256=sha(LOCKED_FILE),扩展映射SHA256=sha(MODEL_FILE),
        饱和负荷密度='未纳入',相关分析范围='补充静态有效情景库')
    dump(REVIEW/'数据核验.json',result)
    return result


def lp_task(row):
    check=model.independent_lp_check(row,row['units'],row['modules'],row['built'])
    assert check['verified'],(key(row),check)
    return dict(节点=key(row),**check)


def export():
    rows=read_rows()
    assert len(rows)==len(tasks()) and len({key(r) for r in rows})==len(rows), '完整设计计算结束后才能导出'
    d,fs=frames();PUBLIC.mkdir(parents=True,exist_ok=True)
    for name,frame in fs.items():frame.to_csv(PUBLIC/(name+'.csv'),index=False,encoding='utf-8-sig')
    with pd.ExcelWriter(PUBLIC/'补充指标与相关性分析.xlsx',engine='openpyxl') as writer:
        for name,frame in fs.items():frame.to_excel(writer,sheet_name=name,index=False)
        for sheet in writer.book:
            sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width=min(38,max(13,len(str(col[0].value))*1.8))
    dump(SOURCE/'参数范围及来源.json',dict(实际范围依据=[
        dict(地区='邳州',年电量渗透率=1698270.73/4216392.40,用户负荷利用小时=4216392.40/956.45,最长95峰段_h=2),
        dict(地区='市区',年电量渗透率=590451.12/6791485.10,用户负荷利用小时=6791485.10/1545.20,最长95峰段_h=6)],
        研究设定=design(),原始曲线=dict(path=str(model.PV_SOURCE.relative_to(ROOT)),sha256=sha(model.PV_SOURCE)),
        定义依据='2025年4月配电网规划设计技术导则修订稿；当前报告第3.3—3.5节及第4章',
        实际数据说明='两地区电量渗透率为估算值；实际最长峰段来自补齐的全年净负荷；补充样本为合成情景'))
    print('导出补充工作簿，有效节点',len(d),flush=True)


def report_assets():
    from sync_current_report_results import table,f,pct
    d,fs=frames();t={}
    rows=[]
    for r in fs['新增指标组内相关'].to_dict('records'):
        rows.append([r['指标'].removesuffix('_MW').removesuffix('_h'),r['有效组数'],r['有效节点数'],
            f(r['组内Pearson']) if pd.notna(r['组内Pearson']) else '无变异',
            f(r['组内Spearman']) if pd.notna(r['组内Spearman']) else '无变异',pct(r['主变容量响应组占比'])])
    t['extended_controlled_correlations']=table('表3-0 补充指标与容载比的组内相关',
        ['指标','有效组数','有效节点数','Pearson','Spearman','容量响应组占比'],rows)
    cm=d[METRICS].corr()
    # 分为两个列块以适合Word正文宽度，行序和符号保持一致。
    for part,indices in enumerate([list(range(7)),list(range(7,13))],1):
        t['extended_indicator_correlations_'+str(part)]=table('表3-0 参数及运行指标的Pearson相关矩阵（'+str(part)+'）',
            ['指标']+[SHORT[i] for i in indices],[[SHORT[j]]+[f(cm.iloc[j,i]) for i in indices] for j in range(13)])
    t['extended_parameter_nodes']=table('表3-0 补充指标的取值与静态仿真衔接',['指标','取值及计算','模型作用'],[
        ['分布式电源电量渗透率','0～65%，步长2.5个百分点','按年电量平衡校准典型光伏出力'],
        ['95%峰值持续时间','1～8 h，步长1 h','按合成年曲线核定，支撑时长取相同值'],
        ['主变与线路单位功率投资比','线路价格取参考值的70%～130%，步长10个百分点','线路能力固定，改变联络工程费用'],
        ['主变与储能单位功率投资比','储能价格取参考值的70%～130%，步长10个百分点','柜型和持续时间固定，改变采购费用'],
        ['最大正向、反向净负荷','对同期汇总净负荷分别取两方向极值，MW','由源荷输入形成容量需求'],
        ['站级负载率最大值','各站净负荷绝对峰值除以初始容量和功率因数，取五站最大值','反映规划前局部容量压力']])
    return t,d,fs


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--solve',action='store_true');parser.add_argument('--export',action='store_true')
    parser.add_argument('--verify',action='store_true');parser.add_argument('--lp',action='store_true')
    parser.add_argument('--workers',type=int,default=8);parser.add_argument('--limit',type=int)
    args=parser.parse_args()
    if args.solve:solve(args.workers,args.limit)
    if args.export:export()
    if args.verify or args.lp:print(json.dumps(verify(args.lp,args.workers),ensure_ascii=False,indent=2))
    if not any([args.solve,args.export,args.verify,args.lp]):print(json.dumps(design(),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
