"""从两地区近五年原始统计表生成实际案例相关矩阵；--check复核已生成资料。"""
from pathlib import Path
import argparse
import csv
import hashlib
import io
import json
import math
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('实验/研究/data/tuomin/电网建模数据_Agent整合版_V1.2/近5年容载比.xlsx')
OUTPUT = Path('研究报告/数据来源/当前版/实际案例相关性')
LABELS = ['主变容量', '年网供最大负荷', '年度同比增长率', '容载比']
FIELDS = ['capacity_mva', 'peak_mw', 'growth', 'clr']


def table(caption, header, rows, note=''):
    text = caption + '\n\n| ' + ' | '.join(header) + ' |\n| ' + ' | '.join(['---'] * len(header)) + ' |\n'
    text += ''.join('| ' + ' | '.join(map(str, row)) + ' |\n' for row in rows)
    return text + ('\n注：' + note + '\n' if note else '')


def calculate(root):
    source = root / SOURCE
    ws = load_workbook(source, data_only=True)['Sheet1']
    samples, matrices, tables, targets = [], {}, {}, []
    for region, code, row in [('邳州', 'QX-00005', 19), ('市区', 'QX-00007', 9)]:
        assert ws.cell(row, 1).value == code and ws.cell(row, 2).value == '110千伏'
        records, previous = [], None
        for year, column in zip(range(2021, 2026), [3, 6, 9, 12, 15]):
            assert ws.cell(1, column).value == f'{year}年'
            capacity = float(ws.cell(row, column).value) * 10  # 万kVA -> MVA
            peak = float(ws.cell(row, column + 1).value) * 10  # 万kW -> MW
            clr = capacity / peak
            assert abs(clr - float(ws.cell(row, column + 2).value)) <= .0051
            record = dict(region=region, year=year, capacity_mva=capacity, peak_mw=peak,
                          growth=peak / previous - 1 if previous else None, clr=clr,
                          source_sheet='Sheet1', source_row=row,
                          source_columns=f'{ws.cell(row,column).column_letter}:{ws.cell(row,column+2).column_letter}')
            records.append(record)
            samples.append(record)
            previous = peak
        frame = pd.DataFrame(records)[FIELDS]
        pearson, spearman = frame.corr(), frame.corr(method='spearman')
        counts = frame.notna().astype(int).T @ frame.notna().astype(int)
        assert counts.loc['growth', 'clr'] == 4 and counts.loc['peak_mw', 'clr'] == 5
        # 独立中心化点积复核Pearson；排序后的点积复核Spearman。
        for a in FIELDS:
            for b in FIELDS:
                pair = frame[[a, b]].dropna()
                for method, matrix in [('pearson', pearson), ('spearman', spearman)]:
                    values = pair.rank(method='average').to_numpy() if method == 'spearman' else pair.to_numpy()
                    centered = values - values.mean(axis=0)
                    expected = np.dot(centered[:, 0], centered[:, 1]) / math.sqrt(np.sum(centered[:, 0] ** 2) * np.sum(centered[:, 1] ** 2))
                    assert abs(matrix.loc[a, b] - expected) < 1e-12
        matrices[region] = dict(字段=FIELDS, Pearson=pearson.values.tolist(), Spearman=spearman.values.tolist(), 有效样本数=counts.values.tolist())
        tables['actual_corr_' + ('pizhou' if region == '邳州' else 'city')] = table(
            f'表3-0 {region}实际年度指标的Pearson相关矩阵', ['指标'] + LABELS,
            [[label] + [f'{pearson.iloc[i,j]:.2f}' for j in range(4)] for i,label in enumerate(LABELS)])
        targets.extend([[region, label, str(int(counts.loc[field, 'clr'])), f'{pearson.loc[field,"clr"]:.2f}', f'{spearman.loc[field,"clr"]:.2f}']
                        for field,label in zip(FIELDS[:3], LABELS[:3])])
    tables['actual_corr_target'] = table('表3-0 两地区实际指标与容载比的相关性', ['地区', '指标', '样本数', 'Pearson', 'Spearman'], targets)
    return samples, matrices, tables


def maintain(root=ROOT, check=False):
    samples, matrices, generated = calculate(root)
    output = root / OUTPUT
    asset = root / '研究报告/02图表/当前版/计算表.json'
    tables = json.loads(asset.read_text())
    record = {'样本来源':str(SOURCE), '来源SHA256':hashlib.sha256((root/SOURCE).read_bytes()).hexdigest(),
              '生成器SHA256':hashlib.sha256((root/'tools/prepare_case_correlations.py').read_bytes()).hexdigest(),
              '样本单位':'同一地区110 kV年度统计；地区分别计算', '样本':samples, '相关矩阵':matrices}
    stream = io.StringIO(newline='')
    writer = csv.DictWriter(stream, fieldnames=list(samples[0]));writer.writeheader();writer.writerows(samples)
    if check:
        saved=json.loads((output/'复算.json').read_text())
        for key in record:
            if key != '相关矩阵':
                assert saved[key] == record[key], '实际案例相关性记录需复算: ' + key
        for region in matrices:
            for key in ['字段', '有效样本数']:
                assert saved['相关矩阵'][region][key] == matrices[region][key]
            for method in ['Pearson', 'Spearman']:
                np.testing.assert_allclose(saved['相关矩阵'][region][method], matrices[region][method], atol=1e-12, rtol=0)
        assert (output/'年度样本.csv').read_text() == stream.getvalue().replace('\r\n','\n')
        for key,value in generated.items():
            assert tables[key] == value, '实际案例相关表与原始统计不一致: ' + key
    else:
        output.mkdir(parents=True, exist_ok=True)
        (output/'复算.json').write_text(json.dumps(record, ensure_ascii=False, indent=2)+'\n')
        (output/'年度样本.csv').write_text(stream.getvalue().replace('\r\n','\n'))
        tables.update(generated)
        asset.write_text(json.dumps(tables,ensure_ascii=False,indent=2)+'\n')
    print('实际案例相关性通过：2地区、10条年度记录、逐对核验样本量及两种相关系数')


if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check',action='store_true')
    maintain(check=parser.parse_args().check)
