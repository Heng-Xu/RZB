"""独立公式：对称三栏承载原生OMML，由文字段落居中，消除显示数学与表格对齐叠加。"""
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import json
import re

MARKER='ReportEquation'
INDENTS=['firstLine','firstLineChars','hanging','hangingChars','left','right','start','end',
         'leftChars','rightChars','startChars','endChars']

def configure_equation_table(table,paragraph):
    """只改变布局；保留公式原生对象及全部字符、分式和矩阵结构。"""
    t=table._tbl;p=paragraph._p
    pr=t.tblPr
    for tag in ['tblStyle','tblLook','tblpPr','tblInd','tblCellMar','tblCaption']:
        for e in pr.findall(qn('w:'+tag)):pr.remove(e)
    ind=OxmlElement('w:tblInd');ind.set(qn('w:w'),'0');ind.set(qn('w:type'),'dxa');pr.insert(2,ind)
    mar=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        node=OxmlElement('w:'+side);node.set(qn('w:w'),'0');node.set(qn('w:type'),'dxa');mar.append(node)
    pr.append(mar)
    caption=OxmlElement('w:tblCaption');caption.set(qn('w:val'),MARKER);pr.append(caption)
    for cell in table.rows[0].cells:
        cp=cell._tc.get_or_add_tcPr()
        for node in cp.findall(qn('w:tcMar')):cp.remove(node)
        valign=cp.find(qn('w:vAlign'));cm=OxmlElement('w:tcMar')
        for side in ['top','left','bottom','right']:
            node=OxmlElement('w:'+side);node.set(qn('w:w'),'0');node.set(qn('w:type'),'dxa');cm.append(node)
        cp.insert(list(cp).index(valign) if valign is not None else len(cp),cm)
        for pp in cell.paragraphs:
            pi=pp._p.get_or_add_pPr()
            for node in pi.findall(qn('w:pStyle')):pi.remove(node)
            pin=pi.find(qn('w:ind'))
            if pin is None:pin=OxmlElement('w:ind');pi.append(pin)
            pin.attrib.clear()
            for name in INDENTS:pin.set(qn('w:'+name),'0')
    # 不使用oMathPara的独立数学段落对齐；独立占行的原生oMath由宿主段落居中。
    for mp in list(p.findall(qn('m:oMathPara'))):
        offset=list(p).index(mp)
        for math in mp.findall(qn('m:oMath')):
            mp.remove(math);p.insert(offset,math);offset+=1
        p.remove(mp)
    jc=p.get_or_add_pPr().find(qn('w:jc'))
    if jc is None:jc=OxmlElement('w:jc');p.get_or_add_pPr().append(jc)
    jc.set(qn('w:val'),'center')

def equation_tables(doc):
    return [t for t in doc.tables if t._tbl.xpath('./w:tblPr/w:tblCaption/@w:val')==[MARKER]]

def check_equation_layout(doc):
    tables=equation_tables(doc)
    policy=json.loads((Path(__file__).parent/'report_policy.json').read_text())['math_format']
    for i,table in enumerate(tables,1):
        label=f'独立公式{i}'
        number=table.cell(0,2).text.strip('（）')
        assert len(table.rows)==1 and len(table.columns)==3,label+'须采用对称三栏'
        t=table._tbl
        assert t.xpath('./w:tblPr/w:jc/@w:val')==['center'],label+'表格未居中'
        assert t.xpath('./w:tblPr/w:tblLayout/@w:type')==['fixed'],label+'须固定栏宽'
        assert not t.xpath('./w:tblPr/w:tblpPr'),label+'不得浮动'
        grid=list(map(int,t.xpath('./w:tblGrid/w:gridCol/@w:w')))
        assert grid[0]==grid[2],label+'左右栏须等宽'
        assert sum(grid)==int(t.xpath('./w:tblPr/w:tblW/@w:w')[0]),label+'栏宽总和须等于表宽'
        tc=t.xpath('./w:tr/w:tc')
        assert [int(x.xpath('./w:tcPr/w:tcW/@w:w')[0]) for x in tc]==grid,label+'单元格宽度与网格不符'
        for c in tc:
            for side in ['left','right']:
                assert c.xpath('./w:tcPr/w:tcMar/w:'+side+'/@w:w')==['0'],label+'单元格内边距须清零'
            for p in c.xpath('./w:p'):
                assert all(p.xpath('./w:pPr/w:ind/@w:'+name)==['0'] for name in INDENTS),label+'段落缩进须全部清零'
        p=tc[1].xpath('./w:p')[0]
        assert p.xpath('./w:pPr/w:jc/@w:val')==['center'],label+'公式宿主段落未居中'
        assert len(p.xpath('./m:oMath'))==1 and not p.xpath('./m:oMathPara'),label+'须由宿主段落承载单个原生公式'
        assert not p.xpath('./w:r/w:t'),label+'公式主体不得混入式号或空格'
        assert tc[2].xpath('./w:p/w:pPr/w:jc/@w:val')==['right'],label+'式号未右置'
        matrices=p.xpath('./m:oMath/m:m')
        if number in policy.get('single_line_equations',[]):
            assert not matrices,label+'可单行显示，不得强制分行：'+number
        for matrix in matrices:
            columns=[x.get(qn('m:val')) for x in matrix.findall('./'+ '/'.join(qn('m:'+tag) for tag in ['mPr','mcs','mc','mcPr','mcJc']))]
            if number in policy.get('continuation_equations',[]):
                assert columns==['right','left'],label+'连续推导须在关系符或运算符处对齐'
            else:
                assert columns==['center'],label+'独立多行关系须逐行居中：'+number
                assert all(len(row)==1 for row in matrix.findall(qn('m:mr'))),label+'独立多行关系不得分成左右两列'
    return tables


def pdf_formula_spans(page,mark):
    """由相邻正文限定公式区域，避免下一段重音字形的字体外框误入公式。"""
    yc=(mark[1]+mark[3])/2
    spans=[s for b in page.get_text('dict')['blocks'] for line in b.get('lines',[]) for s in line['spans']]
    prose=[s for s in spans if len(re.findall('[\u4e00-\u9fff]',s['text']))>=4]
    lo=max((s['bbox'][3] for s in prose if s['bbox'][3]<yc),default=75)
    hi=min((s['bbox'][1] for s in prose if s['bbox'][1]>yc),default=770)
    return [s for s in spans if s['bbox'][0]>90 and s['bbox'][2]<470
            and s['bbox'][1]>=lo and s['bbox'][3]<=hi and s['text'].strip()]


def check_pdf_equation_rows(doc,pdf,tolerance):
    """按关系符的实际基线划分公式行，核验各独立行的可见主体中心。"""
    policy=json.loads((Path(__file__).parent/'report_policy.json').read_text())['math_format']
    result=[]
    for table in equation_tables(doc):
        number=table.cell(0,2).text.strip('（）')
        matrices=table._tbl.xpath('./w:tr/w:tc[2]/w:p/m:oMath/m:m')
        if not matrices or number in policy.get('continuation_equations',[]):continue
        expected=len(matrices[0].findall(qn('m:mr')))
        found=[(page,w) for page in pdf for w in page.get_text('words')
               if w[0]>468 and w[4]=='（'+number+'）']
        assert len(found)==1,('多行公式式号缺失或重复',number)
        page,mark=found[0]
        formula=pdf_formula_spans(page,mark)
        anchors=[]
        for s in formula:
            if s['size']>=11 and any(c in s['text'] for c in '=≤≥'):
                y=s['origin'][1]
                if not any(abs(y-v)<1 for v in anchors):anchors.append(y)
        anchors.sort()
        assert len(anchors)==expected,('多行公式实际行数与原生结构不对应',number,len(anchors),expected)
        bands=[[] for _ in anchors]
        for s in formula:
            index=min(range(len(anchors)),key=lambda i:abs(s['origin'][1]-anchors[i]))
            bands[index].append(s)
        for index,band in enumerate(bands,1):
            x0=min(s['bbox'][0] for s in band);x1=max(s['bbox'][2] for s in band)
            center=(page.rect.width+doc.sections[0].left_margin.pt-doc.sections[0].right_margin.pt)/2
            offset=abs((x0+x1)/2-center)
            assert offset<=tolerance,('PDF独立多行公式未逐行居中',number,index,offset)
            result.append(dict(公式=number,行=index,页码=page.number+1,左边界_pt=round(x0,3),
                               右边界_pt=round(x1,3),中心偏差_pt=round(offset,4)))
    return result
