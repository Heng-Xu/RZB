from pathlib import Path
import json,itertools,math
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font,Alignment,PatternFill
ROOT=Path(__file__).resolve().parents[4];D=Path(__file__).parent;A=ROOT/'研究报告/02图表/当前版';W=ROOT/'研究报告/01研究推荐/当前版/通用容载比推荐矩阵.xlsx'
T=json.loads((A/'计算表.json').read_text());S=json.loads((A/'生成段落.json').read_text())
def table(k,title,cols,rows,note=''):
 T[k]=title+'\n\n| '+' | '.join(cols)+' |\n| '+' | '.join(['---']*len(cols))+' |\n'+''.join('| '+' | '.join(f'{v:.2f}' if isinstance(v,float) else str(v) for v in row)+' |\n' for row in rows)+(('\n注：'+note+'\n') if note else '')
paths=[ROOT/'研究报告/试算/20261010_五参数推广/main.jsonl',ROOT/'研究报告/试算/20261010_负荷分散系数扩展/extension.jsonl']
records=[json.loads(x) for p in paths for x in p.read_text().splitlines()];m=pd.DataFrame(records)
keys=['eta','growth_rate','line_cost_coefficient','initial_transfer_fraction','support_duration_h','diversity'];names=dict(zip(keys,['分布式电源装机渗透率','年最大净负荷平均增长率','线路成本系数','初始转供率','储能支撑持续时间','负荷分散系数']))
assert len(m)==12600 and m[keys].drop_duplicates().shape[0]==12600
assert m.verified.all() and m.relative_gap.max()<=.005001
assert all(r['independent_network_check']['violating_pairs']==0 for r in records)
labels=[['≤1.00','(1.00, 2.00]','>2.00'],['≤2%','(2%, 4%]','(4%, 7%]','>7%'],['≤1.65','(1.65, 6.85]','>6.85'],['≤30%','(30%, 50%]','>50%'],['≤2 h','(2 h, 4 h]','>4 h'],['≤1.10','>1.10']]
cut=[[1,2],[.02,.04,.07],[1.65,6.85],[.3,.5],[2,4],[1.1]]
m['group']=m.apply(lambda r:tuple(sum(r[k]>t for t in ts) for k,ts in zip(keys,cut)),axis=1);groups={b:q for b,q in m.groupby('group')}
sections=[];matrix=[]
for di,hi,ei in itertools.product(range(2),range(3),range(3)):
 rows=[]
 for gi,ki,ri in itertools.product(range(4),range(3),range(3)):
  b=(ei,gi,ki,ri,hi,di);q=groups[b];lo=math.floor(q.clr.min()*100)/100;up=math.ceil(q.clr.max()*100)/100
  vals=[labels[1][gi],labels[2][ki],labels[3][ri],f'{lo:.2f}～{up:.2f}'];rows.append(vals)
  matrix.append([labels[5][di],labels[0][ei],labels[4][hi],*vals,q.clr.min(),q.clr.max()])
 key=f'six_matrix_{di}_{hi}_{ei}';table(key,f'表5-0 负荷分散系数{labels[5][di]}、分布式电源装机渗透率{labels[0][ei]}、储能支撑持续时间{labels[4][hi]}的容载比计算范围',['年最大净负荷平均增长率','线路成本系数','初始转供率','容载比计算范围'],rows)
 sections.append('[[TABLE:'+key+']]')
S['new_matrix_tables']='\n\n'.join(sections)
S['new_matrix_summary']='推广计算覆盖12600组参数组合，按负荷分散系数分为两组，共形成648个参数区间单元。负荷分散系数不大于1.10组包括1.00、1.10两个计算节点，大于1.10组采用1.15节点；储能支撑持续时间采用2、4、6 h节点。全部结果通过设备及网络约束核验，费用相对求解间隙不超过千分之五。容载比计算值总体为2.01～3.28，表中各单元端点向外保留两位小数。区间形式覆盖参数定义域，数值结论限于已计算节点及所设模板。'
table('new_parameter_bins','表5-0 推广矩阵参数区间',['参数','区间','取值依据'],[[names[k],'；'.join(ls),'DL/T 5729—2023表6.3.4' if k=='growth_rate' else '研究设置；参照首尾开口分档形式'] for k,ls in zip(keys,labels)],'负荷分散系数按定义不小于1；分布式电源装机渗透率和线路成本系数非负；初始转供率在0～100%内；储能支撑持续时间为正。首尾档不以有限计算节点截断，容载比输出仍为有限范围。')
corr=[];inter=[];effects=[];pairs=[]
for key in keys:
 corr.append([names[key],m[key].corr(m.clr),m[key].corr(m.clr,method='spearman'),m[key].corr(m.cost),m[key].corr(m.cost,method='spearman')])
 others=[k for k in keys if k!=key];q=m[m[key]==m[key].min()].merge(m[m[key]==m[key].max()],on=others,suffixes=('_lo','_hi'));d=q.clr_hi-q.clr_lo
 effects.append([names[key],len(q),d.mean(),d.abs().max(),f'{(d.abs()>1e-7).mean():.0%}',(q.cost_hi-q.cost_lo).mean()])
 for _,r in q.iterrows():pairs.append({'参数':names[key],**{names[k]:r[k] for k in others},'低值':r[key+'_lo'],'高值':r[key+'_hi'],'容载比差':r.clr_hi-r.clr_lo,'费用差/万元':r.cost_hi-r.cost_lo})
for k,j in itertools.combinations(keys,2):inter.append([names[k],names[j],m[k].corr(m[j]),m[k].corr(m[j],method='spearman')])
table('new_main_effects','表5-0 核心参数最低与最高计算节点的配对变化',['参数','配对数','平均容载比差','最大绝对差','容载比变化占比','平均费用差/万元'],effects)
table('new_correlations','表5-0 推广参数与容载比、评价期费用现值的相关性',['参数','容载比Pearson','容载比Spearman','费用Pearson','费用Spearman'],corr)
table('new_input_correlations','表5-0 推广核心参数之间的相关性',['指标一','指标二','Pearson','Spearman'],inter,'六项输入独立交叉取值，相关系数接近零是计算设计的结果，不表示实际地区的参数相互独立。')
for key in ['new_aux_effects','new_aux_correlations']:
 T[key]='\n'.join(l for l in T[key].splitlines() if not l.startswith('| 负荷分散系数 |'))+'\n'
S['new_main_effects']='分布式电源装机渗透率提高后，日间反向功率需求增加；年最大净负荷平均增长率同时改变容量需求和容载比分母。负荷分散系数在区域同时最大负荷不变的条件下改变各站独立峰值，从而影响局部主变N-1及站间支援要求。线路成本系数、初始转供率和储能支撑持续时间分别作用于工程费用、存量转供预算和储能有效支撑功率，需结合设备组合解释响应。'
S['new_aux_effects']=S['new_aux_effects'].replace('负荷分散系数、','')
(A/'计算表.json').write_text(json.dumps(T,ensure_ascii=False,indent=2));(A/'生成段落.json').write_text(json.dumps(S,ensure_ascii=False,indent=2))
w=load_workbook(W)
for old in ['五参数分类矩阵','六参数分类矩阵','主网格计算值','主参数与目标相关性','主参数间相关性','主参数配对']:
 if old in w:w.remove(w[old])
def sheet(name,cols,rows):
 sh=w.create_sheet(name);sh.append(cols)
 for row in rows:sh.append([v.item() if hasattr(v,'item') else v for v in row])
 sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
 for c in sh[1]:c.font=Font(name='宋体',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='29465B');c.alignment=Alignment(wrap_text=True,vertical='center')
 for col in sh.columns:
  sh.column_dimensions[col[0].column_letter].width=24
  for c in col[1:]:c.number_format='0.00' if isinstance(c.value,float) else 'General'
 return sh
sheet('六参数分类矩阵',['负荷分散系数区间','分布式电源装机渗透率区间','储能支撑持续时间区间','年最大净负荷平均增长率区间','线路成本系数区间','初始转供率区间','容载比计算范围','容载比最小值','容载比最大值'],matrix)
cols=keys+['capacity_mva','peak_mw','clr','cost','cost_lower','relative_gap','storage_mw','storage_mwh','line_count'];name2={**names,'capacity_mva':'公用主变总容量/MVA','peak_mw':'年网供最大负荷/MW','clr':'容载比','cost':'评价期费用现值/万元','cost_lower':'费用下界/万元','relative_gap':'相对求解间隙','storage_mw':'储能额定功率/MW','storage_mwh':'储能额定能量/MWh','line_count':'新增联络项目数'}
sh=sheet('主网格计算值',[name2[k] for k in cols],m[cols].itertuples(index=False,name=None))
for k in ['growth_rate','initial_transfer_fraction','relative_gap']:
 for row in sh.iter_rows(min_row=2,min_col=cols.index(k)+1,max_col=cols.index(k)+1):row[0].number_format='0%'
sheet('主参数与目标相关性',['参数','容载比Pearson','容载比Spearman','费用Pearson','费用Spearman'],corr)
sheet('主参数间相关性',['指标一','指标二','Pearson','Spearman'],inter)
sheet('主参数配对',list(pairs[0]),[list(x.values()) for x in pairs])
sh=w['使用说明'];sh.delete_rows(1,sh.max_row);sh.append(['说明'])
for v in ['六参数分类矩阵共648个区间单元，12600组计算组合。','负荷分散系数≤1.10和>1.10分别查表；计算节点1.00、1.10、1.15。','只有年最大净负荷平均增长率分界直接采用DL/T 5729—2023表6.3.4，其他分界为研究设置。','输入区间首尾开口；容载比为有限计算范围，未计算的参数节点需另行校核。']:sh.append([v])
w.save(W)
(D/'矩阵扩展核验.json').write_text(json.dumps({'组合数':len(m),'区间单元':len(matrix),'最大相对间隙':m.relative_gap.max(),'六参数配对':effects,'相关性':corr},ensure_ascii=False,indent=2))
print('matrix saved',len(matrix))
