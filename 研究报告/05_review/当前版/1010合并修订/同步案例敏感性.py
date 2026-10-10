from pathlib import Path
import json
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import fontManager,FontProperties
ROOT=Path(__file__).resolve().parents[4];D=Path(__file__).parent;A=ROOT/'研究报告/02图表/当前版';P=ROOT/'研究报告/03_MD/当前版';O=ROOT/'研究报告/试算/20261010_案例敏感性'
rows=json.loads((O/'汇总.json').read_text());assert len(rows)==14 and all(r['verified'] for r in rows)
T=json.loads((A/'计算表.json').read_text());S=json.loads((A/'生成段落.json').read_text())
regnames={'pizhou':'邳州','city':'市区'};bases={r['region']:r for r in rows if r['parameter']=='baseline'}
for r in rows:
 assert max(r['solver']['matrix_check'].values())<1e-5
 assert all(c['violating_pairs']==0 for c in r['checks'])
 assert abs(r['summary']['cost_recomputation_error_10k'])<1e-4
for reg,b in bases.items():
 # 同一输入保留历次求解有效下界中较强者。
 old=json.loads((ROOT/f'研究报告/试算/20261010_四小时案例/{reg}/elastic_storage/summary.json').read_text());b['solver']['dual_bound']=max(old['solver']['dual_bound'],b['solver']['dual_bound'])
 b['solver']['mip_gap']=(b['solver']['objective']-b['solver']['dual_bound'])/b['solver']['objective']
def table(k,title,cols,rs,note=''):
 T[k]=title+'\n\n| '+' | '.join(cols)+' |\n| '+' | '.join(['---']*len(cols))+' |\n'+''.join('| '+' | '.join(f'{x:.2f}' if isinstance(x,float) else str(x) for x in row)+' |\n' for row in rs)+(('\n注：'+note+'\n') if note else '')
cols=['地区','参数取值','年网供最大负荷/MW','公用主变总容量/MVA','容载比','储能额定功率/MW','新增联络项目数','可行费用/万元','费用下界/万元'];outrows=[];groups={}
for key,slug,name in [('growth_delta','growth','年最大净负荷平均增长率'),('line_factor','line','线路成本系数'),('duration','duration','储能支撑持续时间')]:
 rs=[];paras=[]
 for reg,cn in regnames.items():
  group=sorted([r for r in rows if r['region']==reg and r['parameter']==key]+[bases[reg]],key=lambda r:0 if r['parameter']=='baseline' and key=='growth_delta' else (1 if r['parameter']=='baseline' and key=='line_factor' else (4 if r['parameter']=='baseline' else r['value'])))
  groups[reg,key]=group
  for r in group:
   base=r['parameter']=='baseline';value=(r['growth_baseline']+(0 if base else r['value'])) if key=='growth_delta' else ((1.65 if reg=='pizhou' else 6.85)*(1 if base else r['value']) if key=='line_factor' else (4 if base else r['value']))
   text=f'{value:.0%}' if key=='growth_delta' else (f'{value:g} h' if key=='duration' else f'{value:.2f}')
   if base:text+='（基准）'
   y=r['summary']['annual'][-1];obj=r['solver']['objective'];low=r['solver']['dual_bound'];r['display_value']=text
   rs.append([cn,text,y['peak_mw'],y['capacity_mva'],y['clr'],y['storage_nameplate_mw'],y['cumulative_lines'],obj,low])
   outrows.append([cn,name,text,y['peak_mw'],y['capacity_mva'],y['clr'],y['storage_nameplate_mw'],y['cumulative_lines'],obj,low,r['solver']['mip_gap']])
  l,b,u=group;yl,yb,yu=[r['summary']['annual'][-1] for r in group]
  if key=='growth_delta':
   paras.append(f"{cn}在较低、基准及较高年最大净负荷平均增长率下，目标年负荷分别为{yl['peak_mw']:.2f}、{yb['peak_mw']:.2f}、{yu['peak_mw']:.2f} MW，对应主变容量{yl['capacity_mva']:.2f}、{yb['capacity_mva']:.2f}、{yu['capacity_mva']:.2f} MVA，容载比为{yl['clr']:.2f}、{yb['clr']:.2f}、{yu['clr']:.2f}。上述变化说明容量随设备投运调整，不能由负荷增长倍率直接推定同倍率的容载比变化。")
  elif key=='line_factor':
   paras.append(f"{cn}线路成本系数由基准的75%增至125%时，所取得可行方案的新增联络项目数由{yl['cumulative_lines']}项变为{yu['cumulative_lines']}项，主变容量由{yl['capacity_mva']:.2f} MVA变为{yu['capacity_mva']:.2f} MVA，储能额定功率由{yl['storage_nameplate_mw']:.2f} MW变为{yu['storage_nameplate_mw']:.2f} MW。工程规模与费用共同列示，可区分相同配置的造价变化和候选措施组合变化。")
  else:
   paras.append(f"{cn}在2、4、6 h正向支撑要求下，目标年主变容量分别为{yl['capacity_mva']:.2f}、{yb['capacity_mva']:.2f}、{yu['capacity_mva']:.2f} MVA，储能额定功率分别为{yl['storage_nameplate_mw']:.2f}、{yb['storage_nameplate_mw']:.2f}、{yu['storage_nameplate_mw']:.2f} MW，容载比为{yl['clr']:.2f}、{yb['clr']:.2f}、{yu['clr']:.2f}。"+('三个候选均未配置储能，所列主变与联络组合能够在这些持续时间条件下满足静态约束。' if max(yl['storage_nameplate_mw'],yb['storage_nameplate_mw'],yu['storage_nameplate_mw'])==0 else '储能有效支撑功率随持续时间折算，额定功率变化需要连同主变与联络配置判断。'))
 table('case_sensitivity_'+slug,'表3-0 两地区'+name+'敏感性',cols,rs,'设备规模及容载比为目标年可行候选值，费用为2022—2025年新增措施在共同评价期内的费用现值。各情景采用1%的相对求解间隙目标，未收敛结果同时列出有效下界，不将费用差表述为精确最优差。')
 S['case_sensitivity_'+slug]='\n\n'.join(paras)
b=bases['city'];y=b['summary']['annual'][-1]
S['case_sensitivity_comparison']=f'''年最大净负荷平均增长率的扰动同时作用于各年功率规模及容载比分母；线路成本系数直接改变联络投资，设备组合可随经济条件调整；储能支撑持续时间改变单位储能可提供的持续功率。三项参数通过不同机制作用于容载比，不能将费用变化或储能规模变化直接当作容载比的同义指标。

市区零扰动基准取得主变容量{y['capacity_mva']:.2f} MVA、容载比{y['clr']:.2f}、费用{b['solver']['objective']:.2f}万元的可行配置。第五章所列市区规划候选为3324.00 MVA、容载比2.17、费用58346.09万元。相同输入下两组费用接近、容量组合不同，表明细小容载比差可能来自离散配置选择与求解精度；本节以零扰动候选作比较基准，地区既定规划方案另按第五章列示。'''
(A/'计算表.json').write_text(json.dumps(T,ensure_ascii=False,indent=2));(A/'生成段落.json').write_text(json.dumps(S,ensure_ascii=False,indent=2))
fontManager.addfont('/usr/share/fonts/wps-office/FZFSK.TTF');plt.rcParams.update({'font.family':['Times New Roman',FontProperties(fname='/usr/share/fonts/wps-office/FZFSK.TTF').get_name()],'axes.unicode_minus':False,'font.size':12})
fig,axs=plt.subplots(3,2,figsize=(8,9.3));keys=['growth_delta','line_factor','duration'];labels=['年最大净负荷平均增长率','线路成本系数','储能支撑持续时间']
for j,(key,label) in enumerate(zip(keys,labels)):
 for reg,color,marker in [('pizhou','#29465b','o'),('city','#9b4b39','s')]:
  g=groups[reg,key];x=np.arange(3);clr=[r['summary']['annual'][-1]['clr'] for r in g];cost=np.array([r['solver']['objective']/1e4 for r in g]);lo=np.array([r['solver']['dual_bound']/1e4 for r in g]);shift=-.035 if reg=='pizhou' else .035
  axs[j,0].plot(x,clr,marker=marker,color=color,label=regnames[reg]);axs[j,1].errorbar(x+shift,cost,yerr=[cost-lo,np.zeros(3)],marker=marker,color=color,capsize=3,label=regnames[reg])
 ticks=['基准−2个百分点','基准','基准+2个百分点'] if key=='growth_delta' else (['基准×75%','基准','基准×125%'] if key=='line_factor' else ['2 h','4 h','6 h'])
 for ax in axs[j,:]:ax.set_xticks(range(3),ticks,rotation=20,ha='right',fontsize=10);ax.grid(alpha=.18);ax.spines[['top','right']].set_visible(False)
 axs[j,0].set_title(label,fontsize=12);axs[j,0].set_ylabel('容载比');axs[j,1].set_ylabel('评价期费用现值/亿元');axs[j,0].legend(fontsize=10)
fig.tight_layout();fig.savefig(A/'case_sensitivity.png',dpi=240,bbox_inches='tight');plt.close(fig)
w=load_workbook(ROOT/'研究报告/01研究推荐/当前版/统一起点逐年方案与复核.xlsx')
if '案例三参数敏感性' in w:w.remove(w['案例三参数敏感性'])
sh=w.create_sheet('案例三参数敏感性');sh.append(['地区','变化参数','取值','年网供最大负荷/MW','公用主变总容量/MVA','容载比','储能额定功率/MW','新增联络项目数','可行费用/万元','费用下界/万元','相对求解间隙'])
for row in outrows:sh.append(row)
sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
for col in sh.columns:
 sh.column_dimensions[col[0].column_letter].width=23
 for c in col[1:]:c.number_format='0%' if c.column==11 else ('0.00' if isinstance(c.value,float) else 'General')
for c in sh[1]:c.font=Font(name='宋体',bold=True);c.alignment=Alignment(wrap_text=True)
w.save(ROOT/'研究报告/01研究推荐/当前版/统一起点逐年方案与复核.xlsx')
(D/'案例敏感性核验.json').write_text(json.dumps({'独立计算组合':len(rows),'展示行数':len(outrows),'全部可行':True,'最大求解间隙':max(r['solver']['mip_gap'] for r in rows),'基准与扰动分别核验':True},ensure_ascii=False,indent=2))
print('sensitivity synchronized',len(rows))
