"""交付文档、案例结果、通用矩阵及配套工作簿的一致性验收。"""
from pathlib import Path
import json,re,zipfile,hashlib
from collections import Counter
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
import pymupdf
HERE=Path(__file__).resolve().parent
R=HERE.parents[1]
ITER='当前版'
TEXT=R/'03_MD'/ITER
DATA=R/'01研究推荐'/ITER
ASSET=R/'02图表'/ITER
OUT=R/'04_word'

def read(p):return json.loads(p.read_text())
def norm(s):return re.sub(r'\s+','',s)
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def same_cell(actual,expected):
 if isinstance(expected,float):return isinstance(actual,(int,float)) and abs(actual-expected)<1e-10
 return actual==expected
def sequential(items,kind):
 grouped={}
 for c,n in items:grouped.setdefault(c,[]).append(int(n))
 for c,ns in grouped.items():assert ns==list(range(1,len(ns)+1)),(kind,c,ns)
 return sum(map(len,grouped.values()))

def main():
 source=(TEXT/'研究报告装配稿.md').read_text()
 for file,expected in read(HERE/'装配源文件登记.json').items():assert digest(R.parent/file)==expected,('装配后源文件发生变化',file)
 assert len(re.findall(r'^# 第[一二三四五六]章 ',source,re.M))==6
 assert not re.search(r'^#{1,4} 4\.7(?:\s|\.)',source,re.M)
 assert source.index('## 5.2 三类方案')<source.index('## 5.4 寻优结果')<source.index('## 5.6 参数拓展')<source.index('## 5.7 规划建议')
 for term in ['丰县','共同反事实','规划分母代理','建模绝对峰','最大站绝对峰','模型站数','正向/反送任务之和','分母代理','四条路径','[[TABLE:','[[SECTION:','[[FIG:','为本项目本研究','第三台','392个站年的','不另创导则指标','不将全年零功率','不能合并解释','价格口径','Sheet1 第','本报告在文字中称','最大最优性间隙']:
  assert term not in source,term
 refs=re.findall(r'^\[(\d+)\]\s+',source,re.M)
 assert list(map(int,refs))==list(range(1,len(refs)+1)),refs
 assert len(refs)==89
 body=source.split('# 参考文献')[0]
 cited=[int(n) for m in re.finditer(r'\[([0-9]+(?:,[0-9]+)*)\]',body) for n in m[1].split(',')]
 assert list(dict.fromkeys(cited))==list(range(1,len(refs)+1))
 assert set(cited)==set(map(int,refs))
 tables=sequential(re.findall(r'^表(\d+)-(\d+)\s',source,re.M),'表')
 figures=sequential(re.findall(r'^图(\d+)-(\d+)\s',source,re.M),'图')
 eqs=read(ASSET/'公式/公式清单.json')
 assert sequential([x['number'].split('-') for x in eqs],'公式')==35
 mathsource=list(re.finditer(r'\$\$([\s\S]*?)\$\$|\$([^$\n]+)\$',source))
 displays=sum(m[1] is not None for m in mathsource)
 doc=Document(OUT/'研究报告终稿.docx')
 assert [p.text for p in doc.paragraphs if p.style.name in ('Heading 1','Heading 2')]==re.findall(r'^#{1,2} (.*)$',source,re.M),'Word目录层级标题与Markdown源不一致'
 assert [p.text for p in doc.paragraphs if p.style.name.startswith('Heading ')]==re.findall(r'^#{1,4} (.*)$',source,re.M),'Word各级标题与Markdown源不一致'
 objects=doc._element.xpath('.//m:oMath')
 assert not doc._element.xpath('.//w:ins | .//w:del | .//w:commentReference')
 table_review=read(HERE/'Word表格内容复核.json')
 assert table_review['状态']=='通过' and table_review['表格']==tables
 assert table_review['WordSHA256']==digest(OUT/'研究报告终稿.docx') and table_review['装配稿SHA256']==digest(TEXT/'研究报告装配稿.md')
 assert len(objects)==len(mathsource)
 assert len(doc._element.xpath('.//m:oMathPara'))==displays==35
 assert len(doc.inline_shapes)==figures==14
 assert len(doc.tables)==tables+displays
 for tab in doc.tables:
  for row in tab.rows:
   for cell in row.cells:
    assert cell.vertical_alignment==1
    for paragraph in cell.paragraphs:
     ind=paragraph._p.xpath('./w:pPr/w:ind')
     assert ind and ind[0].get(qn('w:firstLine'))=='0' and ind[0].get(qn('w:firstLineChars'))=='0'
     assert not any(ind[0].get(qn('w:'+a)) for a in ['hanging','hangingChars'])
     if not tab._tbl.xpath('.//m:oMathPara'):assert paragraph.alignment==1
  if not tab._tbl.xpath('.//m:oMathPara'):continue
  assert len(tab.rows)==1 and len(tab.columns)==3
  assert tab.cell(0,0).width==tab.cell(0,2).width
  assert len(tab.cell(0,1)._tc.xpath('.//m:oMathPara'))==1
  para=tab.cell(0,1).paragraphs[0]
  assert para.alignment==1
  assert para._p.xpath('.//m:jc')[0].get(qn('m:val'))=='center'
  assert re.fullmatch(r'（\d+-\d+）',tab.cell(0,2).text)
 docx_paths=re.findall(r'!\[\]\(([^)]+\.png)\)',source)
 assert len(docx_paths)==14 and all((R.parent/Path(p)).is_file() for p in docx_paths)
 font=read(HERE/'图件字体说明.json')
 assert font['中文字体']=='方正仿宋_GBK'
 assert doc.styles['Normal'].element.xpath('.//w:rFonts')[0].get(qn('w:eastAsia'))==font['中文字体']
 # 目录以实际排版的正文页核定，同时检查Word缓存。
 pdf=pymupdf.open(OUT/'研究报告终稿.pdf')
 layout=read(HERE/'1004二轮审查/版面抽查记录.json')
 assert layout['状态']=='通过' and layout['PDFSHA256']==digest(OUT/'研究报告终稿.pdf'),'PDF已更新：需实际渲染抽查并更新版面记录后重新核验'
 assert all(1<=page<=len(pdf) for page in layout['抽查页'])
 pages=[norm(p.get_text()) for p in pdf]
 cache=read(HERE/'目录页码核对.json')
 assert len(cache)==len(read(HERE/'目录标题.json'))
 for item in cache:
  actual=[i+1 for i,t in enumerate(pages) if i>=4 and norm(item['标题']) in t]
  assert actual and actual[0]==item['页码'],(item,actual[:1])
 indices={x['锚点']:str(x['页码']) for x in cache}
 fields=0
 for field in doc._element.xpath('.//w:fldSimple'):
  match=re.search(r'PAGEREF\s+(\w+)',field.get(qn('w:instr'),''))
  if not match:continue
  assert ''.join(x.text or '' for x in field.iter(qn('w:t')))==indices[match[1]],(match[1],''.join(x.text or '' for x in field.iter(qn('w:t'))))
  fields+=1
 assert fields==len(cache)
 with zipfile.ZipFile(OUT/'研究报告终稿.docx') as z:
  ns={'ep':'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'}
  tree=etree.fromstring(z.read('docProps/app.xml'))
  assert int(tree.find('ep:Pages',ns).text)==len(pdf)
 assert not any('¿' in p.get_text() for p in pdf),'PDF数学解析错误符号'
 for page in pdf:
  for info in page.get_image_info():
   rect=pymupdf.Rect(info['bbox'])
   assert rect.x0>=-1 and rect.x1<=page.rect.width+1 and rect.y0>=-1 and rect.y1<=page.rect.height+1,('图形越出页面',page.number,rect)
 # 技术、费用及求解精度采用独立复算记录。
 review=read(DATA/'统一起点基准/review.json')
 base=read(HERE/'基准独立复核.json')
 external=read(HERE/'外推独立复核.json')
 stats=read(DATA/'外推核验.json')
 assert sum(x['站年'] for x in base)==392 and all(x['状态']=='通过' for x in base)
 assert external['状态']=='通过' and external['计算组合']==25200 and external['站级复核']==126000
 assert stats['全部约束核对']=='通过' and stats['最大最优性间隙']<=.005
 assert review['primary_cost_relative_gap']<=.01 and review['capacity_preference_relative_gap']<=.005
 assert review['cost_order_imposed'] is False
 assert review['same_measure_set'] and review['same_start_per_district']
 costs=review['cost_npv_10k']
 rigid=sum(x['rigid'] for x in costs.values());elastic=sum(x['elastic'] for x in costs.values())
 assert all(x['elastic']<x['rigid'] for x in costs.values())
 reduction=(1-elastic/rigid)*100
 assert f'{reduction:.2f}%' in source
 for value in (rigid,elastic):assert f'{value:.2f}' in source
 assert '推荐采用弹性方案' in source
 matrix=read(DATA/'通用推荐矩阵.json')
 assert len(matrix)==108
 allowed={'线路类型','分布式电源装机渗透率','年负荷平均增长率','95%峰值持续时间','主变与线路成本比','主变与储能成本比_3h','主选参考档位','覆盖参考档位','涉及参考档位','主选档占比','计算容载比范围','容载比最小值','容载比最大值','档位计数','计算点数'}
 assert all(set(x)==allowed for x in matrix)
 assert set(x['线路类型'] for x in matrix)=={'架空为主','电缆为主'}
 assert len(set(x['主变与线路成本比'] for x in matrix))==2
 assert len(set(x['主变与储能成本比_3h'] for x in matrix))==3
 assert sum(x['计算点数'] for x in matrix)>=25200  # 相邻源荷、持续时间档共享边界点
 for row in matrix:
  low,high=map(float,row['计算容载比范围'].split('～'))
  assert 2<=low<=high<=2.6
  assert low<=row['容载比最小值']<=row['容载比最大值']<=high+1e-10
  counts=row['档位计数'];assert sum(counts.values())==row['计算点数']
  chosen=min(counts,key=lambda t:(-counts[t],float(t)))
  assert row['主选参考档位']==chosen and abs(row['主选档占比']-counts[chosen]/row['计算点数'])<1e-10
  assert float(chosen) in (2.,2.2,2.4,2.6)
  assert row['覆盖参考档位']==max(counts,key=float)
  assert row['涉及参考档位']=='/'.join(sorted(counts,key=float))
  assert float(row['覆盖参考档位'])+1e-9>=row['容载比最大值']
 cases=read(HERE/'案例与矩阵衔接核对.json')
 for case in cases:
  row=matrix[case['矩阵行']-1]
  assert case['落入模板计算范围']==(row['容载比最小值']<=case['容载比']<=row['容载比最大值'])
  import math
  assert abs(float(case['案例上限参考档'])-math.ceil(case['容载比']/.20-1e-10)*.20)<1e-10
 workbook=load_workbook(DATA/'通用容载比推荐矩阵.xlsx',read_only=True)
 ws=workbook['弹性推荐矩阵'];assert ws.max_row==109 and ws.max_column==12
 headers=[c.value for c in ws[1]]
 assert not any(any(t in h for t in ['刚性','柜数','措施','项目数']) for h in headers)
 for row,expected in zip(ws.iter_rows(min_row=2,values_only=True),matrix):
  assert all(same_cell(value,expected[key]) for key,value in zip(headers,row)),('总矩阵工作簿不一致',expected)
 assert workbook['参数与计算点'].max_row==25201
 structured=read(DATA/'按供电类别与分布推荐矩阵.json')
 assert len(structured)==432
 structured_sheet=workbook['按供电类别与分布查表']
 assert structured_sheet.max_row==433 and structured_sheet.max_column==15
 structured_headers=[c.value for c in structured_sheet[1]]
 assert structured_headers==['总矩阵行','供电区域','负荷分布']+headers
 for row,expected in zip(structured_sheet.iter_rows(min_row=2,values_only=True),structured):
  assert all(same_cell(value,expected[key]) for key,value in zip(structured_headers,row)),('分列矩阵工作簿不一致',expected)
 for i,parent in enumerate(matrix,1):
  children=[x for x in structured if x['总矩阵行']==i]
  assert {(x['供电区域'],x['负荷分布']) for x in children}=={(a,s) for a in ('A','BC') for s in ('偏集中','较均匀')}
  assert sum(x['计算点数'] for x in children)==parent['计算点数']
  assert min(x['容载比最小值'] for x in children)==parent['容载比最小值']
  assert max(x['容载比最大值'] for x in children)==parent['容载比最大值']
  merged=Counter()
  for child in children:
   merged.update(child['档位计数'])
   assert float(child['覆盖参考档位'])+1e-9>=child['容载比最大值']
  assert dict(merged)==parent['档位计数']
 diagnostic=read(HERE/'1004二轮审查/二轮技术复核.json')
 assert len(diagnostic['下限诊断'])==16
 assert diagnostic['主选档不足覆盖单元数']==60 and diagnostic['多档单元数']==72
 assert sum(x['仅保留增长分档下限方案']['推荐容载比']<=2. for x in diagnostic['下限诊断'])==6
 profiles=read(HERE/'逐站分析对应记录.json')
 assert len(profiles)==49
 assert len(re.findall(r'^#### [A-Za-z0-9-]+站规划分析',source,re.M))==49
 for district,slug in [('邳州','pizhou'),('市区','city')]:
  import csv
  with (DATA/'统一起点基准'/slug/'elastic_stations.csv').open(encoding='utf-8-sig') as stream:
   stations={x['station']:x for x in csv.DictReader(stream) if x['year']=='2025'}
  for profile in profiles:
   if profile['区县']!=district:continue
   row=stations[profile['站']]
   assert abs(profile['弹性目标容量MVA']-float(row['capacity_mva']))<1e-5
   assert profile['弹性最不利退出余量MW']>=-1e-4
 # 文字、表格中的小数统一到两位，章节编号和文献标识保留。
 precision=re.sub(r'^#{1,4}[^\n]*','',body,flags=re.M)
 precision=re.sub(r'(?<![\d.])\d+\.\d+\.\d+(?:\.\d+)*','',precision)
 bad=[]
 for match in re.finditer(r'(?<![\d.])\d+\.\d+(?![\d.])',precision):
  if len(match[0].split('.')[1])!=2:
   assert match[0] in ('2.3','5.6'),match[0]
 assert body.count('第五章 “一片一策”弹性容载比规划建议')==1
 workbook=load_workbook(DATA/'统一起点逐年方案与复核.xlsx',read_only=True)
 assert workbook['逐年区域方案'].max_row==17 and workbook['容量费用复核'].max_row==5
 assert sum(workbook[x].max_row-1 for x in ['邳州刚性逐站','邳州弹性逐站','市区刚性逐站','市区弹性逐站'])==392
 # 源文件说明、摘要及结论与模型保持相同范围。
 notes=(R/'数据来源'/ITER/'数据来源与参数说明.md').read_text()
 assert all(term in notes for term in ('原值乘10','25200','1.15','432','期末剩余价值','SOC','2.001'))
 style=read(HERE/'语体与术语审查.json');assert style['状态']=='通过'
 result={'状态':'通过','章节':6,'参考文献':len(refs),'PDF页数':len(pdf),'真实数据表':tables,'图件':figures,'独立居中公式':displays,'行内原生公式':len(objects)-displays,'原生数学对象':len(objects),'目录页码核对':fields,
 '基准独立复核站年':392,'一般站群外推组合':25200,'外推独立复核站级记录':126000,'通用矩阵行数':108,'结构分列矩阵行数':432,'下限诊断组合':16,'案例校核点':cases,
 '刚性费用现值万元':rigid,'弹性费用现值万元':elastic,'弹性费用降低比例':1-elastic/rigid,'费用排序作为约束':False,
 '费用阶段最优性间隙':review['primary_cost_relative_gap'],'容量偏好阶段最优性间隙':review['capacity_preference_relative_gap'],'外推最大最优性间隙':stats['最大最优性间隙'],
 '公式处理':'交付Word全部为原生OMML，行内可编辑；独立公式在等宽侧栏间居中，编号单独右置。PDF由同一LaTeX源生成数学图形后在本地排版导出；排版副本不替代原生Word。',
 '版面复核':'自动核对实际分页、目录缓存、各级标题和图形边界；本轮人工抽查另见1004二轮审查/版面抽查记录.json。',
 '审查范围':'数据和模型复算、文档结构与本地版面核对；交付供电科院审核。',
 '字体':font,'语体审核':style['技能'],'文件SHA256':{p.name:digest(p) for p in [OUT/'研究报告终稿.docx',OUT/'研究报告终稿.pdf',DATA/'通用容载比推荐矩阵.xlsx',DATA/'统一起点逐年方案与复核.xlsx']}}
 (HERE/'最终核验.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
 (HERE/'重建一致性核对.json').write_text(json.dumps({'状态':'通过','适用修订':'10.03反馈二轮审查',
  '范围':f'当前源文件哈希、Word各级标题、{tables}张数据表文字、{len(objects)}处原生数学对象及{fields}项目录实际页码一致性核对',
  'WordSHA256':digest(OUT/'研究报告终稿.docx'),'装配源文件登记':'装配源文件登记.json',
  '数据表文字复核':'Word表格内容复核.json','独立公式':displays,'行内公式':len(objects)-displays,'图件':figures,'目录':fields},ensure_ascii=False,indent=2))
 print(json.dumps({k:result[k] for k in ['状态','PDF页数','真实数据表','图件','原生数学对象','目录页码核对','基准独立复核站年','外推独立复核站级记录','通用矩阵行数']},ensure_ascii=False))
if __name__=='__main__':main()
