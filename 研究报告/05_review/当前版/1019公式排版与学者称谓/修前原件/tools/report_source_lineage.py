"""原始资料增加辅助核对列时，按真实归档原件核验历史摘要。"""
from pathlib import Path
import hashlib
import json


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check_auxiliary_cells(original, current, expected_changes):
    """只接受已登记空白单元格的追加；既有数据与公式逐格保持。"""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    a=load_workbook(original,read_only=True,data_only=False)
    b=load_workbook(current,read_only=True,data_only=False)
    assert a.sheetnames==b.sheetnames, '原始资料工作表变化'
    actual=[]
    for sa,sb in zip(a,b):
        for i in range(1,max(sa.max_row,sb.max_row)+1):
            for j in range(1,max(sa.max_column,sb.max_column)+1):
                before,after=sa.cell(i,j).value,sb.cell(i,j).value
                if before==after:continue
                assert before is None, '原始资料已有单元格发生变化：'+sa.title+'!'+get_column_letter(j)+str(i)
                actual.append({'sheet':sa.title,'cell':get_column_letter(j)+str(i),'before':before,'after':after})
    a.close();b.close()
    assert actual==expected_changes, '原始资料辅助列变化未逐项对应登记'


def preserved_original(root, name, expected):
    """实际工作表或已归档旧原件均须与登记摘要一致，其他变化仍阻止交付。"""
    current=root/name
    if digest(current)==expected:return current
    policy=json.loads((root/'tools/report_policy.json').read_text())
    review=root/policy.get('plain_delivery_review','研究报告/05_review/当前版/1019七项表述与清稿')
    record=json.loads((review/'原始资料版本衔接.json').read_text())
    row=record['更新原始资料'][name]
    assert row['历史原SHA256']==expected, '原始资料历史摘要不对应：'+name
    original=root/row['历史原件']
    assert digest(original)==expected, '原始资料归档原件摘要不对应：'+name
    assert digest(current)==row['当前SHA256'], '原始资料更新后再次变化：'+name
    snapshot=root/row['当前核读原件']
    assert digest(snapshot)==row['当前SHA256'], '当前原始资料核读原件缺失：'+name
    check_auxiliary_cells(original,current,row['单元格变化'])
    return original
