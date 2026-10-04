"""将重算后的案例、费用边界及正文摘要同步到当前成果源。"""
from pathlib import Path
import csv,json,sys,math,re
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from 驱动因子分析 import table
ROOT=Path(__file__).resolve().parents[3];sys.path.insert(0,str(ROOT))
from 实验.研究.rebuild_2026.current_economic_assumptions import cost_factors,line_capex
HERE=Path(__file__).parent;R=ROOT/'研究报告';A=R/'02图表/当前版';D=R/'01研究推荐/当前版';BASE=D/'统一起点基准';T=R/'03_MD/当前版'
REG={'pizhou':'邳州','city':'市区'};SC={'rigid':'刚性','elastic':'弹性'}
def readcsv(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def f(r,k):return float(r[k])
def dec(x):return f'{float(x):.2f}'
def main():
    tables=json.loads((A/'计算表.json').read_text());sections=json.loads((A/'生成段落.json').read_text());factors=cost_factors()
    portfolio={(c,s):(readcsv(BASE/c/(s+'_years.csv')),readcsv(BASE/c/(s+'_stations.csv'))) for c in REG for s in SC}
    costs={};component_rows=[];measure_rows=[];transfer_rows=[];comparison=[];independent=[]
    old_totals={'rigid':208566.29,'elastic':35683.63,'reduction':82.89}
    for (c,s),(years,stations) in portfolio.items():
        components={kind:sum(f(y,kind+'_capex_10k')*factors[kind][int(y['year'])] for y in years) for kind in ('transformer','storage','line')}
        costs[c,s]=sum(components.values())
        summary=next(x for x in json.loads((BASE/c/'summary.json').read_text()) if x['scheme']==s)
        assert abs(costs[c,s]-summary['objective_npv_10k'])<1e-3
        component_rows.append([REG[c],SC[s],*[dec(v) for v in components.values()],dec(costs[c,s])])
        measure_rows.append([REG[c],SC[s],dec(sum(f(y,'new_transformer_purchase_mva') for y in years)),int(sum(f(y,'new_third_transformers') for y in years)),dec(years[-1]['installed_storage_power_mw'])+'/'+dec(years[-1]['installed_storage_energy_mwh']),int(years[-1]['new_lines_in_service'])])
        initial=sum(f(st,'prior_unit_1_mva')+f(st,'prior_unit_2_mva') for st in stations if st['year']=='2022');previous=initial;display=[]
        for i,y in enumerate(years,1):
            display.append([i,dec(y['net_peak_proxy_mw']),dec(y['capacity_mva']),dec(f(y,'capacity_mva')-previous),dec(y['new_transformer_purchase_mva']),dec(y['new_storage_power_mw'])+'/'+dec(y['new_storage_energy_mwh']),int(y['new_lines_commissioned']),dec(f(y,'existing_transfer_mw')+f(y,'new_line_transfer_mw')),dec(y['clr'])]);previous=f(y,'capacity_mva')
            transfer_rows.append([REG[c],SC[s],i,dec(y['existing_transfer_mw']),dec(y['new_line_transfer_mw']),int(y['new_lines_in_service']),dec(f(y,'existing_transfer_mw')+f(y,'new_line_transfer_mw'))])
        tables[c+'_'+s+'_years']=table('表5-0 '+REG[c]+SC[s]+'方案逐年措施及容载比',['仿真年','年网供最大负荷/MW','主变容量/MVA','净增/MVA','新购/MVA','新增储能MW/MWh','新建联络/项','总转接/MW','容载比'],display)
        minimum=1e9
        for st in stations:
            S=f(st,'capacity_mva');N=S-max(f(st,k) for k in ('unit_1_mva','unit_2_mva','unit_3_mva'));B=f(st,'storage_energy_mwh')/max(2.15,f(st,'forward_duration_h'));BR=f(st,'storage_energy_mwh')/max(2.15,f(st,'reverse_duration_h'))
            minimum=min(minimum,.95*N+B-f(st,'post_transfer_forward_mw'),.95*S+B-f(st,'post_transfer_forward_mw'),.95*S+BR-f(st,'post_transfer_reverse_upper_mw'))
        assert minimum>=-1e-4
        independent.append({'地区':REG[c],'方案':SC[s],'站年':len(stations),'最小容量余量MW':minimum,'费用偏差万元':costs[c,s]-summary['objective_npv_10k'],'状态':'通过'})
    totals={s:sum(costs[c,s] for c in REG) for s in SC};reduction=(1-totals['elastic']/totals['rigid'])*100
    for c,name in REG.items():
        yearsR=portfolio[c,'rigid'][0];yearsE=portfolio[c,'elastic'][0]
        sections[c+'_results']=name+'刚性方案第1—4年容载比分别为'+'、'.join(dec(y['clr']) for y in yearsR)+'，弹性方案分别为'+'、'.join(dec(y['clr']) for y in yearsE)+'。刚性新增措施全寿命费用现值为'+dec(costs[c,'rigid'])+'万元，弹性为'+dec(costs[c,'elastic'])+'万元。各年主变容量、储能和联络建设列于下表。'
        comparison.extend([[name,i,dec(r['clr']),dec(e['clr']),dec(f(e,'clr')-f(r,'clr'))] for i,(r,e) in enumerate(zip(yearsR,yearsE),1)])
    sections['cost_results']=f"两地区刚性方案新增措施全寿命费用现值合计{totals['rigid']:.2f}万元，弹性合计{totals['elastic']:.2f}万元，弹性费用降低{reduction:.2f}%。本评价期内三类设备均无寿命更新支出，费用差来自新增投资、固定维护费用及投运时点。"
    tables['cost_comparison']=table('表5-0 刚性与弹性方案全寿命费用现值比较',['地区','刚性/万元','弹性/万元','差额/万元','弹性费用降幅'],[[REG[c],dec(costs[c,'rigid']),dec(costs[c,'elastic']),dec(costs[c,'rigid']-costs[c,'elastic']),f"{(1-costs[c,'elastic']/costs[c,'rigid'])*100:.2f}%"] for c in REG])
    tables['cost_components']=table('表5-0 全寿命费用按工程类型分解',['地区','方案','主变/万元','储能/万元','线路/万元','合计/万元'],component_rows)
    tables['measure_totals']=table('表5-0 全期新购主变及第4年在役状态',['地区','方案','新购主变/MVA','扩建主变/项','储能MW/MWh','新增联络/项'],measure_rows)
    tables['transfer_results']=table('表5-0 既有与新增转供能力的实际利用',['地区','方案','仿真年','既有能力使用/MW','新增能力使用/MW','在役项目/项','总转接/MW'],transfer_rows)
    tables['rigid_elastic_clr_comparison']=table('表5-0 刚性与弹性容载比逐年对照',['地区','仿真年','刚性容载比','弹性容载比','差值'],comparison)
    storage={s:(sum(f(portfolio[c,s][0][-1],'installed_storage_power_mw') for c in REG),sum(f(portfolio[c,s][0][-1],'installed_storage_energy_mwh') for c in REG)) for s in SC}
    tables['rigid_elastic_interpretation']=table('表5-0 刚性与弹性方案的技术经济比较',['比较内容','刚性方案','弹性方案'],[['仿真起点','同区县容载比2.00','同区县容载比2.00'],['指标条件','各年不超过2.00','各年大于2.00至2.60'],['技术条件','正常、反向、全负荷主变N-1、持续在役及负荷守恒','与刚性相同'],['增容、储能及联络','三类措施共同寻优','三类措施共同寻优'],['第4年两地区在役储能',*[f'{storage[s][0]:.2f} MW/{storage[s][1]:.2f} MWh' for s in SC]],['新增措施费用现值',dec(totals['rigid'])+'万元',dec(totals['elastic'])+'万元'],['规划比较','作为比较基准',f'费用降低{reduction:.2f}%，优先比选']])
    tables['economic_assumptions']=table('表4-1 统一经济评价参数',['参数','取值','依据'],[['基准年/截止年','2021/2041','研究评价期'],['折现率','6.00%','研究实价假定'],['主变/线路寿命','40/40年','甲方反馈参数'],['储能寿命','25年','甲方反馈参数'],['变电维护费率','初始投资4.20%/年','甲方反馈参数'],['线路维护费率','初始投资2.20%/年','甲方反馈参数'],['储能维护费率','初始投资3.00%/年','研究参数']])
    tables['lifecycle_factors']=table('表4-0 单位初始投资的2021年末现值因子（式4-10）',['投运年','主变','线路','储能'],[[y,*[dec(factors[k][y]) for k in ('transformer','line','storage')]] for y in range(2022,2026)])
    tables['line_cost_components']=table('表4-0 两类等效联络线投资构成',['项目','架空为主工程量','费用/万元','电缆为主工程量','费用/万元'],[['架空','2.70 km','64.80','0.30 km','7.20'],['电缆电气','0.30 km','36.00','2.70 km','324.00'],['电缆土建','0.30 km','18.00','2.70 km','162.00'],['等效开关','2台','6.00','2台','6.00'],['合计','3 km及2台开关','124.80','3 km及2台开关','499.20']])
    tables['general_assumptions']=tables['general_assumptions'].replace('两类一般站群','两类五站模板')
    if '| 线路类型 |' not in tables['general_assumptions']:
        tables['general_assumptions']+='| 线路类型 | 架空为主/电缆为主；电缆占比10.00%/90.00% | 与供电类别交叉计算 |\n'
    (A/'计算表.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2));(A/'生成段落.json').write_text(json.dumps(sections,ensure_ascii=False,indent=2))
    (HERE/'基准独立复核.json').write_text(json.dumps(independent,ensure_ascii=False,indent=2))
    wb=Workbook();ws=wb.active;ws.title='逐年区域方案';headers=list(next(iter(portfolio.values()))[0][0]);ws.append(['地区']+headers)
    for (c,s),(ys,ss) in portfolio.items():
        for row in ys:ws.append([REG[c]]+[number(row[k]) for k in headers])
        ws2=wb.create_sheet(REG[c]+SC[s]+'逐站');heads=list(ss[0]);ws2.append(heads)
        for row in ss:ws2.append([number(row[k]) for k in heads])
    ws=wb.create_sheet('容量费用复核');ws.append(['地区','方案','费用现值万元','目标年容量MVA','目标年容载比','站年','最小容量余量MW'])
    for e in independent:
        c=next(k for k,v in REG.items() if v==e['地区']);s=next(k for k,v in SC.items() if v==e['方案']);last=portfolio[c,s][0][-1]
        ws.append([e['地区'],e['方案'],costs[c,s],f(last,'capacity_mva'),f(last,'clr'),e['站年'],e['最小容量余量MW']])
    for sh in wb:
        sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
        for cell in sh[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='29465B');cell.alignment=Alignment(wrap_text=True)
        for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=20
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,float):cell.number_format='0.00'
    wb.save(D/'统一起点逐年方案与复核.xlsx')
    last={c:{s:portfolio[c,s][0][-1] for s in SC} for c in REG}
    tiers={c:math.ceil(f(last[c]['elastic'],'clr')/.20-1e-10)*.20 for c in REG}
    sections['case_tier_results']=f"按完整精度结果判定，邳州案例对应{tiers['pizhou']:.2f}上限参考档，市区对应{tiers['city']:.2f}上限参考档。两地实际配置的容载比均显示约2.20，档位判定仍使用未取整的容量与负荷；实际建设规模按逐年方案表落实。"
    sections['case_tier_conclusion']=f"按0.20间隔分档，邳州案例采用{tiers['pizhou']:.2f}上限参考档，市区采用{tiers['city']:.2f}上限参考档。档位按未取整的容量与负荷判定，近似显示的容载比用于说明实际设备配置。"
    (A/'生成段落.json').write_text(json.dumps(sections,ensure_ascii=False,indent=2))
    matrix=json.loads((D/'通用推荐矩阵.json').read_text())
    cases=[]
    for c,eta,g,h,lt in [('pizhou',1.630652935,.080741736,4,'架空为主'),('city',.31788,.018,5,'电缆为主')]:
        eb='1.00～2.00' if eta>1 else '0.00～1.00';gb='＞7.00%～10.00%' if g>.07 else '0.00%～2.00%'
        ix=next(i for i,m in enumerate(matrix) if m['线路类型']==lt and m['分布式电源装机渗透率']==eb and m['年负荷平均增长率']==gb and m['95%峰值持续时间']=='3～6 h' and m['主变与储能成本比_3h']=='＞0.05～0.08')
        actual=f(last[c]['elastic'],'clr');row=matrix[ix]
        cases.append({'地区':REG[c],'矩阵行':ix+1,'容载比':actual,'案例上限参考档':dec(tiers[c]),'模板主选档':row['主选参考档位'],'落入模板计算范围':row['容载比最小值']<=actual<=row['容载比最大值'],'说明':'实际异容量49站案例的四年规划与等容量五站目标年模板分别比较；不补点扩大模板范围'})
    (HERE/'案例与矩阵衔接核对.json').write_text(json.dumps(cases,ensure_ascii=False,indent=2))
    for p in T.glob('0[1-6] *.md'):
        text=p.read_text().replace('208566.29',dec(totals['rigid'])).replace('35683.63',dec(totals['elastic'])).replace('82.89',dec(reduction))
        if p.name.startswith('05'):
            start=text.index('根据两地区技术经济对比，推荐采用弹性方案：');end=text.index('\n\n',start)
            text=text[:start]+f"根据两地区技术经济对比，推荐采用弹性方案：仿真第4年邳州容载比{f(last['pizhou']['elastic'],'clr'):.2f}、市区{f(last['city']['elastic'],'clr'):.2f}，对应配置通过所设技术约束，新增措施费用低于刚性。其他片区可按通用矩阵选择参考档位，以实际初态进行刚性、弹性方案比选，结合设备和线路条件确定建设规模。"+text[end:]
        if p.name.startswith('06'):
            text=re.sub(r'规划第4年，弹性方案邳州储能.*?费用优势',f"规划第4年，弹性方案邳州储能为{f(last['pizhou']['elastic'],'installed_storage_power_mw'):.2f} MW/{f(last['pizhou']['elastic'],'installed_storage_energy_mwh'):.2f} MWh，市区为{f(last['city']['elastic'],'installed_storage_power_mw'):.2f} MW/{f(last['city']['elastic'],'installed_storage_energy_mwh'):.2f} MWh。费用优势",text)
            text=re.sub(r'根据案例比较，推荐邳州和市区规划第4年采用容载比约2.21的弹性方案。',f"根据案例比较，推荐邳州和市区规划第4年分别采用容载比约{f(last['pizhou']['elastic'],'clr'):.2f}、{f(last['city']['elastic'],'clr'):.2f}的弹性方案。",text)
            text=text.replace('线路及储能各采用三档价格，形成九组交叉经济条件，共完成25200个目标年情景比选。两个成本比分别设置三档，与源荷和持续时间分档共同形成162行弹性容载比推荐矩阵。','线路类型分别取架空为主和电缆为主，各结合三档线路价格和三档储能价格，共完成25200个目标年情景比选。配对结果表明当前线路价格范围未改变容量档位，因此合并同类线路的价格条件，形成108行参考档位矩阵。')
            text=text.replace('推荐范围按这些条件归纳，以相应情景结果及两地区案例形成区间，供其他片区选择容量比选范围。','矩阵按0.20间隔列示主选参考档位、同档占比及计算范围。两地区四年案例与目标年模板分别保留，规划人员按实际供电类别和工程条件比选相邻档位。')
        p.write_text(text)
    abstract=[
        '分布式新能源接入后，110 kV主变同时承担正向供电和反向输送需求，增加容量又可能使区域容载比超过一般规划上限。本研究以徐州邳州和市区为案例，分析新能源渗透指标与传统容载比的关系。弹性容载比是本课题采用的研究概念，指满足相同技术约束、经全寿命费用比较确定、允许适度偏离一般容载比上限的取值。',
        '研究分析年度负荷、光伏规模和变电站净负荷，列出候选指标并开展同口径相关性分析，结合容量约束识别核心参数。费用采用实际主变工程及储能采购资料；主变、线路寿命均为40年，储能为25年，变电和线路维护费率分别为4.20%/年、2.20%/年。刚性、弹性方案共用本地区容载比2.00仿真起点，后续四年校核正常运行、主变N-1、反向容量和负荷转接条件。',
        f"规划第4年，邳州刚性、弹性容载比分别为{f(last['pizhou']['rigid'],'clr'):.2f}、{f(last['pizhou']['elastic'],'clr'):.2f}，市区为{f(last['city']['rigid'],'clr'):.2f}、{f(last['city']['elastic'],'clr'):.2f}。两地区新增措施全寿命费用现值分别合计{totals['rigid']:.2f}万元和{totals['elastic']:.2f}万元，弹性方案降低{reduction:.2f}%。该费用比较对应本研究的负荷、技术条件和参考价格。",
        '通用外推按导则供电区域类别构造A类与B/C类五站模板，并将架空为主、电缆为主两类线路与源荷、增长率、峰段及价格条件交叉计算，完成25200组目标年静态比选。结果形成108行参考档位矩阵，按0.20间隔列示主选档位与同档占比，并以热力图展示。规划中按片区条件选择比选起点，保留2.00上限刚性方案作费用对照，再以实际设备、线路及价格确定建设方案。',
    ]
    (HERE/'内容提要.json').write_text(json.dumps(abstract,ensure_ascii=False,indent=2))
    (HERE/'1003反馈修订/本轮费用结果.json').write_text(json.dumps({'原费用用于核对':old_totals,'本轮费用':totals,'费用降低百分数':reduction,'现值因子':factors,'目标年结果':last},ensure_ascii=False,indent=2))
    print(json.dumps({'费用':totals,'降幅':reduction,'目标年':{c:{s:float(y['clr']) for s,y in v.items()} for c,v in last.items()}},ensure_ascii=False))
def number(value):
    try:return float(value) if '.' in value or 'e' in value.lower() else int(value)
    except (ValueError,TypeError):return value
if __name__=='__main__':main()
