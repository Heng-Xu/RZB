"""依据完整计算点形成线路分类矩阵与0.20参考档位。"""
from pathlib import Path
import json,math,itertools,collections
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from 驱动因子分析 import table
HERE=Path(__file__).parent;R=HERE.parents[1];D=R/'01研究推荐/当前版';A=R/'02图表/当前版'
ETA=[(0,1),(1,2),(2,3)];G=[(0,.02),(.02,.07),(.07,.1)];H=[(1,3),(3,6)];B=[(.04,.05),(.05,.08),(.08,.12)]
def tier(r):return min(2.6,max(2.,math.ceil(r/.2-1e-10)*.2))
def in_growth(v,low,high):return low<v<=high or low==0 and v==0
def select(rows,lt,eta,g,h,b=None,base=False):
    return [r for r in rows if r['线路类型']==lt and eta[0]<=r['装机渗透率']<=eta[1] and in_growth(r['年负荷平均增长率'],*g) and h[0]<=r['95%峰值持续时间h']<=h[1] and (b is None or b[0]<r['主变储能成本比_3h']<=b[1]) and (not base or r['线路价格倍数']==1 and r['储能价格倍数']==1)]
def summarize(rows):
    values=[r['推荐容载比'] for r in rows];counts=collections.Counter(round(tier(x),2) for x in values)
    chosen=min(counts,key=lambda t:(-counts[t],t))
    return chosen,counts[chosen]/len(rows),min(values),max(values),{f'{k:.2f}':v for k,v in sorted(counts.items())}
def main():
    rows=json.loads((D/'外推逐点结果.json').read_text());matrix=[];blocks=[]
    tables=json.loads((A/'计算表.json').read_text());sections=json.loads((A/'生成段落.json').read_text())
    for lt in ('架空为主','电缆为主'):
        for e,g,h,(bi,b) in itertools.product(ETA,G,H,enumerate(B)):
            sub=select(rows,lt,e,g,h,b);assert sub
            chosen,share,low,high,counts=summarize(sub)
            matrix.append({'线路类型':lt,'分布式电源装机渗透率':f'{e[0]:.2f}～{e[1]:.2f}','年负荷平均增长率':('＞' if g[0] else '')+f'{g[0]*100:.2f}%～{g[1]*100:.2f}%','95%峰值持续时间':f'{h[0]}～{h[1]} h','主变与线路成本比':f"{min(r['主变线路成本比'] for r in sub):.2f}～{max(r['主变线路成本比'] for r in sub):.2f}",'主变与储能成本比_3h':('＞' if bi else '')+f'{b[0]:.2f}～{b[1]:.2f}','主选参考档位':f'{chosen:.2f}','主选档占比':share,'计算容载比范围':f'{math.floor((low+1e-10)*100)/100:.2f}～{math.ceil((high-1e-10)*100)/100:.2f}','容载比最小值':low,'容载比最大值':high,'档位计数':counts,'计算点数':len(sub)})
        sub=[m for m in matrix if m['线路类型']==lt]
        key='cost_matrix_'+str(len(blocks));headers=['装机渗透率','增长率','峰段/h','主变与储能成本比','上限参考档','同档占比','计算范围']
        tables[key]=table('表5-0 '+lt+'线路条件下的容载比参考档位（主变与线路成本比'+sub[0]['主变与线路成本比']+'）',headers,[[m['分布式电源装机渗透率'],m['年负荷平均增长率'],m['95%峰值持续时间'].replace(' h',''),m['主变与储能成本比_3h'],m['主选参考档位'],f"{m['主选档占比']*100:.2f}%",m['计算容载比范围']] for m in sub])
        blocks.append('[[TABLE:'+key+']]')
    assert len(matrix)==108
    sections['cost_banded_matrix']='\n\n'.join(blocks);tables.pop('cost_matrix_2',None)
    sections['matrix_results']='线路价格及类型的配对比较未使本次模板的容量取值发生变化，因此矩阵按线路类型分别列示，并合并各类型的三个线路价格条件，共108个参数单元。原始费用及项目数量仍逐点保存。主选档用于前期比选，工程方案按实际设备和线路条件复核。'
    (D/'通用推荐矩阵.json').write_text(json.dumps(matrix,ensure_ascii=False,indent=2))
    keys=['线路类型','分布式电源装机渗透率','年负荷平均增长率','95%峰值持续时间','主变与线路成本比','主变与储能成本比_3h','主选参考档位','主选档占比','计算容载比范围','计算点数']
    wb=Workbook();ws=wb.active;ws.title='弹性推荐矩阵';ws.append(keys)
    for m in matrix:ws.append([m[k] for k in keys])
    ws=wb.create_sheet('参数与计算点');ws.append(list(rows[0]));[ws.append(list(r.values())) for r in rows]
    ws=wb.create_sheet('档位分布');ws.append(['矩阵行','线路类型','2.00档点数','2.20档点数','2.40档点数','2.60档点数','最小原值','最大原值'])
    for i,m in enumerate(matrix,1):ws.append([i,m['线路类型']]+[m['档位计数'].get(f'{x:.2f}',0) for x in (2,2.2,2.4,2.6)]+[m['容载比最小值'],m['容载比最大值']])
    ws=wb.create_sheet('使用说明');ws.append(['项目','内容'])
    for row in [['档位方法','按0.20上限档向上归档；同档频次并列优先较低档'],['含义','档位作为允许上限的比选起点，不要求将实际设备容量补齐至档值；2.00上限另作刚性对照'],['主选档占比','等权参数组合频次，非真实概率；不足60%时比较相邻档'],['计算范围','仅记录本单元条件差异形成的全部计算范围，不作为统一推荐区间'],['线路价格','类型与价格分别计算；对容量无变化后合并价格档，费用保留逐点'],['案例与模板','四年案例与目标年静态模板分别求解，不将案例点补入模板范围']]:ws.append(row)
    for sh in wb:
        sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
        for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='29465B');c.alignment=Alignment(wrap_text=True)
        for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=22
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,float):cell.number_format='0.00%' if sh.cell(1,cell.column).value=='主选档占比' else '0.00'
    ws.append(['分档节点','装机渗透率1.00、2.00及峰段3 h的计算点在相邻单元共享；实际参数位于这些节点时同时核对两侧单元'])
    wb.save(D/'通用容载比推荐矩阵.xlsx')
    (A/'计算表.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2));(A/'生成段落.json').write_text(json.dumps(sections,ensure_ascii=False,indent=2))
    plot(rows)
    print('矩阵',len(matrix),'行；主选档占比范围',min(m['主选档占比'] for m in matrix),max(m['主选档占比'] for m in matrix))
def plot(rows):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager,colors
    from matplotlib.ticker import FormatStrFormatter
    font_manager.fontManager.addfont('/usr/share/fonts/wps-office/FZFSK.TTF')
    font_manager.fontManager.addfont('/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf')
    family=font_manager.FontProperties(fname='/usr/share/fonts/wps-office/FZFSK.TTF').get_name()
    plt.rcParams.update({'font.family':['Times New Roman',family],'font.size':11,'axes.unicode_minus':False})
    fig,axes=plt.subplots(2,2,figsize=(8,6.2),sharex=True,sharey=True)
    cmap=colors.ListedColormap(['#eef2f4','#bed4dc','#73a0b0','#29465b']);norm=colors.BoundaryNorm([1.9,2.1,2.3,2.5,2.7],4)
    for i,lt in enumerate(('架空为主','电缆为主')):
        for j,h in enumerate(H):
            z=np.zeros((3,3));shares=np.zeros((3,3));ax=axes[i,j]
            for gi,g in enumerate(G):
                for ei,e in enumerate(ETA):
                    sub=select(rows,lt,e,g,h,base=True);z[gi,ei],shares[gi,ei],*_=summarize(sub)
            im=ax.imshow(z,cmap=cmap,norm=norm,aspect='auto');ax.set_title(f'{lt}；峰段{h[0]}～{h[1]} h')
            for gi,ei in itertools.product(range(3),repeat=2):ax.text(ei,gi,f'{z[gi,ei]:.2f}\n{shares[gi,ei]*100:.2f}%',ha='center',va='center',color='white' if z[gi,ei]>2.4 else '#182c37')
            ax.set_xticks(range(3),['0.00～1.00','1.00～2.00','2.00～3.00']);ax.set_yticks(range(3),['0.00%～2.00%','＞2.00%～7.00%','＞7.00%～10.00%'])
            if i==1:ax.set_xlabel('装机渗透率')
            if j==0:ax.set_ylabel('年负荷平均增长率')
    fig.tight_layout(rect=(0,0,1,.95));cbar=fig.colorbar(im,ax=axes.ravel().tolist(),location='top',shrink=.65,pad=.08,ticks=[2,2.2,2.4,2.6]);cbar.ax.xaxis.set_major_formatter(FormatStrFormatter('%.2f'));cbar.set_label('主选上限参考档（线路与储能基准价）')
    fig.savefig(A/'planning_tier_heatmap.png',dpi=300,bbox_inches='tight');plt.close(fig)
if __name__=='__main__':main()
