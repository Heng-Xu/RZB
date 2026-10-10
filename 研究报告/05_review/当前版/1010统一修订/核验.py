"""本轮报告的数值、参数区间、展示、原生Word和WPS PDF一致性检查。"""
from pathlib import Path
import json,re,hashlib,sys,math,itertools
import pandas as pd
from docx import Document
from docx.oxml.ns import qn
import pymupdf
ROOT=Path(__file__).resolve().parents[4];D=Path(__file__).parent;R=D.parent;A=ROOT/'研究报告/02图表/当前版';M=ROOT/'研究报告/试算/20261010_五参数推广';C=ROOT/'研究报告/试算/20261010_四小时案例';W=ROOT/'研究报告/01研究推荐/当前版';O=ROOT/'研究报告/04_word'
def read(p):return json.loads(p.read_text())
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s):return re.sub(r'\s+','',s)
rows={k:[json.loads(l) for l in (M/(k+'.jsonl')).read_text().splitlines()] for k in ['pilot','main','sensitivity']}
keys=['eta','growth_rate','line_cost_coefficient','initial_transfer_fraction','support_duration_h'];cuts=[[1,2],[.02,.04,.07],[1.65,6.85],[.3,.5],[2,4]]
for kind,count in [('pilot',96),('main',4200),('sensitivity',363)]:
 data=rows[kind];assert len(data)==count;assert len({r['task_id'] for r in data})==count
 for r in data:
  assert r['verified'] and r['matrix_violation']<1e-5 and r['relative_gap']<=.0050001
  assert r['independent_network_check']['violating_pairs']==0
  assert r['bin']==[sum(r[k]>t for t in cut) for k,cut in zip(keys,cuts)]
  assert abs(r['peak_mw']-250*(1+r['growth_rate'])**4)<1e-7
  assert abs(r['capacity_mva']-sum(sum(u) for u in r['units']))<1e-7
  assert abs(r['clr']-r['capacity_mva']/r['peak_mw'])<1e-8
  assert abs(r['cost']-sum(r['cost_components'].values()))<1e-4
  assert abs(r['storage_mwh']-.215*r['storage_modules'])<1e-7
 assert read(M/(kind+'_summary.json'))['source_sha256']==digest(ROOT/'实验/研究/rebuild_2026/matrix_20261010/run.py')
main=pd.DataFrame(rows['main']);assert main[keys].drop_duplicates().shape[0]==4200
extra=[json.loads(l) for l in (ROOT/'研究报告/试算/20261010_负荷分散系数扩展/extension.jsonl').read_text().splitlines()]
assert len(extra)==8400 and len({r['extension_id'] for r in extra})==8400
for r in extra:
 assert r['verified'] and r['relative_gap']<=.005001 and r['matrix_violation']<1e-5
 assert r['independent_network_check']['violating_pairs']==0
 assert abs(r['capacity_mva']-sum(map(sum,r['units'])))<1e-7
 assert abs(r['clr']-r['capacity_mva']/r['peak_mw'])<1e-8
 assert abs(r['cost']-sum(r['cost_components'].values()))<1e-4
allmain=pd.DataFrame(rows['main']+extra);assert allmain[keys+['diversity']].drop_duplicates().shape[0]==12600
mat=pd.read_excel(W/'通用容载比推荐矩阵.xlsx',sheet_name='六参数分类矩阵');assert len(mat)==648 and '样本数' not in mat.columns
allmain['sixbin']=allmain.apply(lambda r:tuple(r['bin'])+(int(r['diversity']>1.1),),axis=1)
for b,g in allmain.groupby('sixbin'):
 ei,gi,ki,ri,hi,di=b;idx=((di*3+hi)*3+ei)*36+gi*9+ki*3+ri;t=mat.iloc[idx]
 assert abs(t['容载比最小值']-g.clr.min())<1e-7 and abs(t['容载比最大值']-g.clr.max())<1e-7
 lo,up=map(float,t['容载比计算范围'].split('～'));assert lo<=g.clr.min()+1e-10 and up>=g.clr.max()-1e-10
sensitivity=read(ROOT/'研究报告/试算/20261010_案例敏感性/汇总.json');assert len(sensitivity)==14
for r in sensitivity:
 assert r['verified'] and max(r['solver']['matrix_check'].values())<1e-5
 assert all(x['violating_pairs']==0 for x in r['checks'])
 assert abs(r['summary']['cost_recomputation_error_10k'])<1e-4
cases=[]
for reg in ['city','pizhou']:
 inp=read(C/reg/'inputs.json');assert all(v['forward_duration_h']==4 for v in inp['stations'].values())
 for case in ['rigid_storage','elastic_storage']:
  z=read(C/reg/case/'summary.json');assert z['input_sha256']==digest(C/reg/'inputs.json');assert max(z['solver']['matrix_check'].values())<1e-5
  assert all(x['violating_pairs']==0 for x in z['static_network_checks']);assert abs(z['cost_recomputation_error_10k'])<1e-4
  st=pd.read_csv(C/reg/case/'stations.csv');prev=inp['initial_capacity_mva']
  for y in z['annual']:
   assert abs(y['clr']-y['capacity_mva']/y['peak_mw'])<1e-8
   assert abs(st[st.year==y['year']].capacity_mva.sum()-y['capacity_mva'])<1e-7
   assert y['capacity_mva']>=prev-1e-7
   if case=='rigid_storage':assert y['capacity_mva']<=max(prev,2*y['peak_mw'])+1e-7
   else:assert y['clr']>=2.001-1e-7
   prev=y['capacity_mva']
  cases.append({'地区':reg,'方案':case,'费用':z['cost_npv_10k'],'下界':z['solver']['dual_bound'],'间隙':z['solver']['mip_gap']})
 assert read(C/reg/'elastic_storage/summary.json')['cost_npv_10k']<read(C/reg/'rigid_storage/summary.json')['solver']['dual_bound']
source=ROOT/'研究报告/03_MD/当前版/研究报告装配稿.md';body=source.read_text().split('# 参考文献')[0]
for token in ['BDZ-','SIM-CITY-','25200','加权增长','平均增长情景','供电区域按A类规划','0% 时净现值','退出','补充仿真站','推荐区间宽度与取值差异','上一轮静态网络模型']:
 assert token not in body,token
assert not re.search(r'\d+\.\d+%',body),'正文百分比精度'
assert '千分之一点六' in body and '千分之二点五' in body
assert '内容提要' not in ''.join(p.text for p in Document(O/'研究报告终稿.docx').paragraphs)
assert '储能支撑持续时间' in body and '线路成本系数' in body
assert '首尾开口' in body and '有限范围' in body
for p in [W/'通用容载比推荐矩阵.xlsx',W/'统一起点逐年方案与复核.xlsx']:
 sheets=pd.read_excel(p,sheet_name=None)
 for name,df in sheets.items():assert not re.search(r'BDZ-\d{5}|SIM-CITY-(KL|XSZ|YQ)',df.to_string()),(p,name)
refs=read(R/'引文顺序与对应.json');doc=Document(O/'研究报告终稿.docx')
assert len(doc.inline_shapes)==len(re.findall(r'^!\[\]',source.read_text(),re.M))
figure_paths=re.findall(r'^!\[\]\(([^)]+)\)',source.read_text(),re.M)
for shape,path in zip(doc.inline_shapes,figure_paths):
 blip=shape._inline.xpath('.//a:blip')[0]
 part=doc.part.related_parts[blip.get(qn('r:embed'))]
 assert hashlib.sha256(part.blob).hexdigest()==digest(ROOT/path),('Word图件未同步',path)
assert not doc._element.xpath('.//w:sectPr/w:docGrid')
style_ids={s.style_id for s in doc.styles};assert all(x.get(qn('w:val')) in style_ids for x in doc._element.xpath('.//w:pStyle'))
maths=len(doc._element.xpath('.//m:oMath'));displays=len(doc._element.xpath('.//m:oMathPara'));expected_displays=sum(len(re.findall(r'\\tag\{',p.read_text())) for p in (ROOT/'研究报告/03_MD/当前版').glob('0[1-6] *.md'));assert displays==expected_displays and maths>=displays
hashes=read(R/'装配源文件登记.json');assert all(digest(ROOT/k)==v for k,v in hashes.items())
result={'状态':'通过','主网格':12600,'试点':96,'专项及边界':363,'参数区间单元':648,'案例敏感性组合':14,'地区案例':cases,'图件':len(doc.inline_shapes),'原生数学对象':maths,'独立公式':displays,'数据表':len(doc.tables)-displays,'参考文献':len(refs),'WordSHA256':digest(O/'研究报告终稿.docx'),'原始记录SHA256':{k:digest(M/(k+'.jsonl')) for k in rows},'图件SHA256':{p.name:digest(p) for p in A.glob('*.png')}}
if '--without-pdf' not in sys.argv:
 pdf=pymupdf.open(O/'研究报告终稿.pdf');assert 'wps' in str(pdf.metadata).lower() or 'kingsoft' in str(pdf.metadata).lower()
 reg=read(R/'WPS原生导出登记.json');assert reg['WordSHA256']==result['WordSHA256'] and reg['PDFSHA256']==digest(O/'研究报告终稿.pdf')
 pages=[norm(p.get_text()) for p in pdf];toc=read(R/'目录页码核对.json');lines=[norm(l) for p in list(pdf)[:3] for l in p.get_text().splitlines()]
 for item in toc:
  title=norm(item['标题']);actual=[i+1 for i,t in enumerate(pages) if i>=3 and title in t];assert actual and actual[0]==item['页码'],item
  hits=[l for l in lines if l.startswith(title)];assert len(hits)==1,(item,hits)
  num=re.search(r'(\d+)$',hits[0][len(title):]);assert num and int(num[1])==item['页码'],(item,hits)
 for p in pdf:
  assert '¿' not in p.get_text()
  for im in p.get_image_info():
   r=pymupdf.Rect(im['bbox']);assert r.x0>=-1 and r.x1<=p.rect.width+1 and r.y0>=-1 and r.y1<=p.rect.height+1
 result.update(PDF页数=len(pdf),PDFSHA256=digest(O/'研究报告终稿.pdf'),目录核对=len(toc))
(D/('预核验.json' if '--without-pdf' in sys.argv else '最终核验.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps({k:v for k,v in result.items() if k not in ['图件SHA256','原始记录SHA256','地区案例']},ensure_ascii=False))
