import os,sys,re,json,subprocess,copy,math,hashlib
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_LINE_SPACING,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
import cairosvg
ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT/'tools'))
TEXT=ROOT/'研究报告/03_MD/当前版';REVIEW=Path(__file__).parent
ASSET=ROOT/'研究报告/02图表/当前版';MATH=ASSET/'公式';MATH.mkdir(exist_ok=True)
OUT=ROOT/'研究报告/04_word';OUT.mkdir(exist_ok=True,parents=True)
TITLE='徐州地区分布式新能源高渗透率地区110 kV电网容载比弹性指标优化研究'
TABLES=json.loads((ASSET/'计算表.json').read_text());SECTIONS=json.loads((ASSET/'生成段落.json').read_text())
def font(run,size=12,bold=False,east='方正仿宋_GBK'):
 run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold
 pr=run._element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
 rf.set(qn('w:eastAsia'),east);rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman')
def style(s,size,bold=False,east='方正仿宋_GBK'):
 s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=bold
 rf=s.element.get_or_add_rPr().find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');s.element.get_or_add_rPr().append(rf)
 rf.set(qn('w:eastAsia'),east)
def refdoc():
 source=ROOT/'研究报告/范文与语料/国网宿迁公司宿迁市“十五五”电网侧新型储能发展项目研究报告-初稿.docx'
 doc=Document(source);body=doc._element.body
 for el in list(body):
  if el.tag!=qn('w:sectPr'):body.remove(el)
 for sec in doc.sections:
  sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.54);sec.bottom_margin=Cm(2.54);sec.left_margin=Cm(3.1697);sec.right_margin=Cm(3.1697)
  for part in [sec.header,sec.footer,sec.first_page_header,sec.first_page_footer]:
   for el in list(part._element):part._element.remove(el)
 style(doc.styles['Normal'],12);fmt=doc.styles['Normal'].paragraph_format;fmt.line_spacing=1.25;fmt.space_before=Pt(0);fmt.space_after=Pt(0);fmt.first_line_indent=Pt(24)
 for n,size in [(1,18),(2,16),(3,13),(4,12)]:
  s=doc.styles['Heading '+str(n)];style(s,size,True);s.paragraph_format.line_spacing=1.5;s.paragraph_format.first_line_indent=Pt(0);s.paragraph_format.space_before=Pt(8 if n>1 else 0);s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=True
 for nm in ['Body Text','First Paragraph','Compact']:
  if nm not in doc.styles:
   s=doc.styles.add_style(nm,WD_STYLE_TYPE.PARAGRAPH);s.base_style=doc.styles['Normal']
  s=doc.styles[nm];style(s,12);s.paragraph_format.line_spacing=Pt(17.4);s.paragraph_format.first_line_indent=Pt(24);s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(0)
 doc.save(REVIEW/'排版参照.docx')
def assemble():
 sources=sorted(TEXT.glob('0[1-6] *.md'))+[ASSET/'计算表.json',ASSET/'生成段落.json']
 (REVIEW/'装配源文件登记.json').write_text(json.dumps({str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},ensure_ascii=False,indent=2))
 txt='\n\n'.join(p.read_text().strip() for p in sorted(TEXT.glob('0[1-6] *.md')))
 txt=re.sub(r'\[\[SECTION:([^\]]+)\]\]',lambda m:SECTIONS[m[1]],txt)
 txt=re.sub(r'\[\[TABLE:([^\]]+)\]\]',lambda m:TABLES[m[1]],txt)
 txt=re.sub(r'\[\[FIG:([^|]+)\|([^\]]+)\]\]',lambda m:'![]('+str((ASSET/(m[1]+'.png')).relative_to(ROOT))+'){width=5.6in}\n\n'+m[2],txt)
 registry={int(k):v for k,v in json.loads((REVIEW/'文献原编号登记.json').read_text()).items()}
 from report_citations import arrange
 txt,refs,reference_records=arrange(txt,registry)
 (TEXT/'参考文献.md').write_text(refs)
 (REVIEW/'引文顺序与对应.json').write_text(json.dumps(reference_records,ensure_ascii=False,indent=2))
 txt+='\n\n'+refs
 # 按正文顺序连续编号，公式引用同步改号。
 counters={};mapping={}
 for match in re.finditer(r'\\tag\{(\d+)-(\d+)\}',txt):
  chapter=match[1];counters[chapter]=counters.get(chapter,0)+1;mapping[match[1]+'-'+match[2]]=chapter+'-'+str(counters[chapter])
 txt=re.sub(r'\\tag\{([^}]+)\}',lambda m:'\\tag{'+mapping[m[1]]+'}',txt)
 txt=re.sub(r'式([（(]?)(\d+-\d+)',lambda m:'式'+m[1]+mapping.get(m[2],m[2]),txt)
 for kind in ['表','图']:
  counts={}
  def caption(m):
   c=m[1];counts[c]=counts.get(c,0)+1;return kind+c+'-'+str(counts[c])+' '
  txt=re.sub(r'^'+kind+r'(\d+)-[\da-z-]+\s+',caption,txt,flags=re.M)
 # 公式保留LaTeX源，经Pandoc转换为Word原生OMML。
 mathlist=[]
 def equation(m):
  tex=m[1];tags=re.findall(r'\\tag\{([^}]+)\}',tex)
  number=tags[0] if tags else str(len(mathlist)+1)
  tex=re.sub(r'\\tag\{[^}]+\}','',tex)
  mathlist.append({'number':number,'tex':tex})
  return '$$'+tex+'$$'
 txt=re.sub(r'\$\$([\s\S]*?)\$\$',equation,txt)
 assert '[[' not in txt
 (TEXT/'研究报告装配稿.md').write_text(txt)
 (MATH/'公式清单.json').write_text(json.dumps(mathlist,ensure_ascii=False,indent=2))
 result=subprocess.run(['pandoc',str(TEXT/'研究报告装配稿.md'),'--reference-doc='+str(REVIEW/'排版参照.docx'),'--from=markdown+tex_math_dollars','--to=docx','--output='+str(REVIEW/'原始装配.docx')],check=True,cwd=ROOT,capture_output=True,text=True)
 assert 'Could not convert TeX math' not in result.stderr,result.stderr
 if result.stderr:print(result.stderr,file=sys.stderr)
 return mathlist
def configure_doc(mathlist):
 doc=Document(REVIEW/'原始装配.docx');sec=doc.sections[0]
 sec.different_first_page_header_footer=True;sec.header_distance=Cm(1.27);sec.footer_distance=Cm(1.27)
 sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.54);sec.bottom_margin=Cm(2.54);sec.left_margin=Cm(3.1697);sec.right_margin=Cm(3.1697)
 # 清除模板悬挂的多级编号，正文编号已经写在标题文本。
 for n in range(1,5):
  s=doc.styles['Heading '+str(n)];ppr=s.element.get_or_add_pPr()
  for node in ppr.findall(qn('w:numPr')):ppr.remove(node)
 normal=doc.styles['Normal'];style(normal,12);normal.paragraph_format.line_spacing=Pt(17.4)
 numeq=0
 for p in doc.paragraphs:
  pf=p.paragraph_format;t=p.text.strip();level=int(p.style.name.split()[-1]) if p.style.name.startswith('Heading ') else 0
  if level:
   ppr=p._p.get_or_add_pPr()
   for nd in ppr.findall(qn('w:numPr')):ppr.remove(nd)
   np=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');np.append(ni);ppr.append(np)
   size={1:18,2:16,3:13,4:12}.get(level,12);pf.first_line_indent=Pt(0);pf.line_spacing=Pt(size*1.5);pf.keep_with_next=True;pf.space_before=Pt(8 if level>1 else 0);pf.space_after=Pt(6)
   if level==1:pf.page_break_before=True
   for r in p.runs:font(r,size,True)
  elif p._element.xpath('.//m:oMathPara'):
   it=mathlist[numeq];numeq+=1
   pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Pt(0);pf.line_spacing=1.15;pf.space_before=Pt(5);pf.space_after=Pt(5);pf.keep_with_next=True
   for mp in p._element.xpath('.//m:oMathPara'):
    pr=mp.find(qn('m:oMathParaPr'))
    if pr is None:pr=OxmlElement('m:oMathParaPr');mp.insert(0,pr)
    for existing in pr.findall(qn('m:jc')):pr.remove(existing)
    jc=OxmlElement('m:jc');jc.set(qn('m:val'),'center');pr.append(jc)
   # 编号单独置于行尾，数学对象独立居中。
   p.add_run('（'+it['number']+'）')
  elif p._element.xpath('.//w:drawing'):
   pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Pt(0);pf.line_spacing=1;pf.space_before=Pt(6);pf.space_after=Pt(0);pf.keep_with_next=True
  elif t.startswith('表') and re.match(r'表\d+-',t):
   pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Pt(0);pf.line_spacing=Pt(14);pf.space_before=Pt(6);pf.space_after=Pt(4);pf.keep_with_next=True
   for r in p.runs:font(r,10.5,False,'宋体')
  elif t.startswith('图') and re.match(r'图\d+-',t):
   pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Pt(0);pf.line_spacing=Pt(14);pf.space_before=Pt(3);pf.space_after=Pt(6)
   for r in p.runs:font(r,10.5,False,'宋体')
  else:
   pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.first_line_indent=Pt(24);pf.line_spacing=Pt(17.4);pf.space_before=Pt(0);pf.space_after=Pt(0);pf.widow_control=True
   if p._p.xpath('.//m:f | .//m:nary | .//m:m'):pf.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
   if t.startswith('式中'):
    pf.first_line_indent=Pt(0);pf.left_indent=Pt(0);pf.right_indent=Pt(0)
   if re.match(r'^\[\d+\]',t):pf.first_line_indent=Pt(0);pf.line_spacing=Pt(16);pf.space_after=Pt(4);pf.alignment=WD_ALIGN_PARAGRAPH.LEFT
   for r in p.runs:font(r,10.5 if re.match(r'^\[\d+\]',t) else 12)
 assert numeq==len(mathlist),(numeq,len(mathlist))
 # 重建为标准Word表格，消除转换器沿用模板时产生的无效表格样式和结构。
 for old in list(doc.tables):
  values=[[c.text for c in row.cells] for row in old.rows]
  new=doc.add_table(rows=len(values),cols=len(values[0]))
  for row,vs in zip(new.rows,values):
   for cell,text in zip(row.cells,vs):cell.text=text
  old._tbl.addprevious(new._tbl);old._tbl.getparent().remove(old._tbl)
 for tab in doc.tables:
  tab.autofit=False;tab.alignment=1;n=len(tab.rows[0].cells);total=14.6606
  weights=[1]*n
  head=[c.text for c in tab.rows[0].cells]
  if head[0]=='储能价格倍数' and len(tab.rows)==4:
   caption=tab._tbl.getprevious()
   caption_text=''.join(caption.itertext()) if caption is not None else ''
   if '负荷分散系数1.00的储能投资比比较' in caption_text:
    heading=caption.getprevious()
    if heading is not None and heading.tag==qn('w:p'):
     # 每类首张短表另页起排，两组比较表连贯显示，避免孤立表题。
     heading.get_or_add_pPr().append(OxmlElement('w:pageBreakBefore'))
  if any('站号' in x for x in head):weights[head.index(next(x for x in head if '站号' in x))]=1.2
  for j,x in enumerate(head):
   if 'MW/MWh' in x:weights[j]=1.45
  if head[0]=='资料':weights=[1.2,1.2,1.2,1.2]
  if head[0]=='参数' and n==3:weights=[1.1,2,2]
  if head[0]=='参数' and n==2:weights=[1,2]
  if '主变组合/MVA' in head and n==6:weights=[0.65,1.80,1.40,1,1,1.05]
  if '容载比推荐区间' in head and n==5:weights=[1.0,1.35,1.05,1.4,1.2]
  if head[0]=='参数' and '计算依据与用途' in head:weights=[1.3,1.4,2.3]
  if '主变与线路单位功率投资比' in head and n==7:weights=[.75,.85,.7,1.25,1.2,.8,1.15]
  widths=[Cm(total*w/sum(weights)) for w in weights]
  grid=tab._tbl.tblGrid
  for child in list(grid):grid.remove(child)
  for width in widths:
   gc=OxmlElement('w:gridCol');gc.set(qn('w:w'),str(width.twips));grid.append(gc)
  for col,width in zip(tab.columns,widths):col.width=width
  pr=tab._tbl.tblPr
  tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(Cm(total).twips))
  bord=pr.find(qn('w:tblBorders'))
  if bord is None:bord=OxmlElement('w:tblBorders');pr.append(bord)
  for edge in ['top','bottom','insideH','insideV','left','right']:
   el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'555555');bord.append(el)
  for idx,row in enumerate(tab.rows):
   trp=row._tr.get_or_add_trPr();csp=OxmlElement('w:cantSplit');trp.append(csp)
   if idx==0:repeat=OxmlElement('w:tblHeader');trp.append(repeat)
   for j,cell in enumerate(row.cells):
    cell.width=widths[j];cell.vertical_alignment=1
    cp=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
    for side in ['left','right']:
     e=OxmlElement('w:'+side);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');mar.append(e)
    cp.append(mar)
    for p in cell.paragraphs:
     f=p.paragraph_format;f.first_line_indent=Pt(0);f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.line_spacing=Pt(13);f.space_before=Pt(1);f.space_after=Pt(1);f.keep_with_next=idx==0
     if len(tab.rows)<=10 or (len(tab.rows)==13 and n==6 and head[:2]==['年最大净负荷平均增长率','装机渗透率']):
      f.keep_with_next=idx<len(tab.rows)-1
     for r in p.runs:font(r,9,idx==0,'宋体')
 # 公式采用等宽侧栏，数学对象居中，式号独立在右栏。
 for p in list(doc.paragraphs):
  if not p._element.xpath('.//m:oMathPara'):continue
  number=p.runs[-1].text
  p.runs[-1]._r.getparent().remove(p.runs[-1]._r)
  tab=doc.add_table(rows=1,cols=3);tab.autofit=False;tab.alignment=1
  widths=[Cm(1.3),Cm(12.0606),Cm(1.3)]
  for col,w in zip(tab.columns,widths):col.width=w
  for cell,w in zip(tab.rows[0].cells,widths):
   cell.width=w;cell.vertical_alignment=1
   margin=OxmlElement('w:tcMar')
   for side in ['left','right']:
    nd=OxmlElement('w:'+side);nd.set(qn('w:w'),'0');nd.set(qn('w:type'),'dxa');margin.append(nd)
   cell._tc.get_or_add_tcPr().append(margin)
  pr=tab._tbl.tblPr;tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(Cm(14.6606).twips))
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','bottom','insideH','insideV','left','right']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'nil');borders.append(e)
  pr.append(borders)
  rowpr=tab.rows[0]._tr.get_or_add_trPr();rowpr.append(OxmlElement('w:cantSplit'))
  mid=tab.cell(0,1);mid._tc.remove(mid.paragraphs[0]._p)
  p._p.addprevious(tab._tbl);mid._tc.append(p._p)
  for cell in [tab.cell(0,0),tab.cell(0,2)]:
   pp=cell.paragraphs[0];pp.paragraph_format.first_line_indent=Pt(0);pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.space_after=Pt(0)
  for cell in tab.rows[0].cells:
   for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
  pp=tab.cell(0,2).paragraphs[0];pp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
  font(pp.add_run(number),9)
  sys.path.insert(0,str(ROOT/'tools'))
  from report_math_layout import configure_equation_table
  configure_equation_table(tab,p)
 # 页眉页脚采用页码域与目录缓存。
 h=sec.header.add_paragraph('徐州地区110 kV电网容载比弹性指标优化研究');h.alignment=1
 for r in h.runs:font(r,9,False,'宋体')
 f=sec.footer.add_paragraph();f.alignment=1;r=f.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
 # 封面、目录由同一版面生成；不插空白填充页。
 body=doc._element.body;front=Document();frontsec=front.sections[0]
 def add(text='',size=12,bold=False,align=0):
  p=front.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(max(17.5,size*1.5));p.alignment=align
  r=p.add_run(text);font(r,size,bold);return p
 p=add(TITLE,22,True,1);p.paragraph_format.space_before=Pt(140);p.paragraph_format.space_after=Pt(50)
 add('研究报告',20,True,1);p=add('2026年10月',14,False,1);p.paragraph_format.space_before=Pt(120)
 p=add('目  录',18,True,1);p.paragraph_format.page_break_before=True;p.paragraph_format.space_after=Pt(15)
 headings=[]
 for i,p in enumerate(doc.paragraphs):
  if p.style.name in ['Heading 1','Heading 2']:
   text=p.text;level=int(p.style.name[-1]);anchor='h'+str(i)
   bs=OxmlElement('w:bookmarkStart');bs.set(qn('w:id'),str(1000+i));bs.set(qn('w:name'),anchor);be=OxmlElement('w:bookmarkEnd');be.set(qn('w:id'),str(1000+i));p._p.insert(1,bs);p._p.append(be)
   headings.append((text,level,anchor));t=add('',12,level==1);t.paragraph_format.left_indent=Pt(12 if level==2 else 0);t.paragraph_format.space_after=Pt(2);t.paragraph_format.tab_stops.add_tab_stop(Cm(14.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
   link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),anchor);link.set(qn('w:history'),'1')
   r=t.add_run(text);font(r,12,level==1);t._p.remove(r._r);link.append(r._r);t._p.append(link)
   t.add_run('\t')
   for kind in ['begin','instruction','separate','result','end']:
    r=t.add_run();font(r,12,level==1)
    if kind=='instruction':
     nd=OxmlElement('w:instrText');nd.set(qn('xml:space'),'preserve');nd.text=' PAGEREF '+anchor+' \\h ';r._r.append(nd)
    elif kind=='result':r.text='待更新'
    else:
     nd=OxmlElement('w:fldChar');nd.set(qn('w:fldCharType'),kind);r._r.append(nd)
   for r in t.runs:font(r,12,level==1)
 for el in reversed(list(front._element.body)):
  if el.tag!=qn('w:sectPr'):body.insert(0,copy.deepcopy(el))
 (REVIEW/'目录标题.json').write_text(json.dumps(headings,ensure_ascii=False,indent=2))
 # Word常见标点与单位保留，关闭修订和注释；作者为项目组。
 doc.core_properties.author='研究项目组';doc.core_properties.last_modified_by='研究项目组';doc.core_properties.title=TITLE;doc.core_properties.subject='110 kV容量规划与措施经济比选'
 settings=doc.settings.element
 for node in settings.findall(qn('w:trackRevisions')):settings.remove(node)
 # 不继承范文网格及Mac数学字体；显式行距减少不同办公软件的字体度量差异。
 for section in doc.sections:
  for node in section._sectPr.findall(qn('w:docGrid')):section._sectPr.remove(node)
 for s in doc.styles:
  if s.type==WD_STYLE_TYPE.PARAGRAPH:
   pr=s.element.get_or_add_pPr()
   for node in pr.findall(qn('w:snapToGrid')):pr.remove(node)
   node=OxmlElement('w:snapToGrid');node.set(qn('w:val'),'0');pr.append(node)
 math_font=json.loads((ROOT/'tools/report_policy.json').read_text())['math_format']['font']
 for node in settings.xpath('.//m:mathFont'):node.set(qn('m:val'),math_font)
 # 用原生OMML结构配合明确的西文字体，防止WPS将变量改映射到替代数学字形。
 for run in doc._element.xpath('.//m:r'):
  mr=run.find(qn('m:rPr'))
  if mr is None:mr=OxmlElement('m:rPr');run.insert(0,mr)
  sty=mr.find(qn('m:sty'));scr=mr.find(qn('m:scr'))
  text=''.join(node.text or '' for node in run.findall(qn('m:t')))
  italic=any(('A'<=c<='Z') or ('a'<=c<='z') or ('α'<=c<='ϕ') for c in text)
  if sty is not None and sty.get(qn('m:val')) in ['p','b'] and scr is None:italic=False
  nor=mr.find(qn('m:nor'))
  if nor is None:nor=OxmlElement('m:nor');mr.append(nor)
  nor.set(qn('m:val'),'1')
  pr=run.find(qn('w:rPr'))
  if pr is None:pr=OxmlElement('w:rPr');run.insert(list(run).index(mr)+1,pr)
  rf=pr.find(qn('w:rFonts'))
  if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
  for key in ['ascii','hAnsi','cs']:rf.set(qn('w:'+key),math_font)
  rf.set(qn('w:eastAsia'),'方正仿宋_GBK')
  for key in ['i','iCs']:
   node=pr.find(qn('w:'+key))
   if node is None:node=OxmlElement('w:'+key);pr.append(node)
   node.set(qn('w:val'),'1' if italic else '0')
 for node in settings.findall(qn('w:updateFields')):settings.remove(node)
 node=OxmlElement('w:updateFields');node.set(qn('w:val'),'true');settings.append(node)
 # 表格单元格及独立公式段落的特殊缩进显式置零，覆盖模板继承。
 for p in doc._element.xpath('.//w:tbl//w:p | .//w:p[m:oMathPara] | .//w:p[starts-with(w:r/w:t,"式中")]'):
  pr=p.find(qn('w:pPr'))
  if pr is None:pr=OxmlElement('w:pPr');p.insert(0,pr)
  ind=pr.find(qn('w:ind'))
  if ind is None:ind=OxmlElement('w:ind');pr.append(ind)
  for attr in ['firstLine','firstLineChars','hanging','hangingChars','left','right','start','end','leftChars','rightChars','startChars','endChars']:ind.set(qn('w:'+attr),'0')
 # 图片及图题必须覆盖正文样式中的字符缩进，点值首行置零不足以取消继承。
 figure_policy=json.loads((ROOT/'tools/report_policy.json').read_text())['figure_paragraph_format']
 for p in doc._element.xpath('.//w:body//w:p'):
  drawing=p.xpath('.//w:drawing | .//w:pict')
  caption=re.match(r'^图\d+-', ''.join(p.xpath('.//w:t/text()')).strip())
  if not drawing and not caption:continue
  pr=p.get_or_add_pPr();ind=pr.find(qn('w:ind'))
  if ind is None:ind=OxmlElement('w:ind');pr.append(ind)
  ind.attrib.clear()
  for attr in figure_policy['zero_indent_attributes']:ind.set(qn('w:'+attr),'0')
  jc=pr.find(qn('w:jc'))
  if jc is None:jc=OxmlElement('w:jc');pr.append(jc)
  jc.set(qn('w:val'),figure_policy['alignment'])
 doc.save(OUT/'研究报告终稿.docx')
 print('Word',OUT/'研究报告终稿.docx','表',len(doc.tables),'公式',numeq,'图',len(doc.inline_shapes))
def main():
 subprocess.run([sys.executable,str(ROOT/'tools/check_report_policy.py'),'--stage','source'],cwd=ROOT,check=True)
 refdoc();ml=assemble();configure_doc(ml)
 subprocess.run([sys.executable,str(ROOT/'tools/check_report_policy.py'),'--stage','docx'],cwd=ROOT,check=True)
if __name__=='__main__':main()
