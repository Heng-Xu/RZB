"""独立核验本轮原文件修订、锁定数据、六章结构及WPS最终视图证据。"""
from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
import hashlib,json,re
from lxml import etree as E
from docx import Document

from report_annual_update import ROOT,REPORT,ASSET,REVIEW,sync
from revise_report_in_place import TARGET,BASE,CLEAN,NS,qn,package,xml,view,fingerprint

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def check_capacity_period(annual,start,end,capacity,location=''):
    rows={int(row['year']):float(row['capacity_mva']) for row in annual}
    assert all(year in rows and abs(rows[year]-capacity)<1e-8 for year in range(start,end+1)), location+' 容量保持年份范围与逐年结果不符'

def doc_from_parts(parts,root):
    data=BytesIO()
    with ZipFile(data,'w',ZIP_DEFLATED) as z:
        for name,value in parts.items():z.writestr(name,E.tostring(root) if name=='word/document.xml' else value)
    data.seek(0);return Document(data)

def verify(final=False):
    from verify_future_case_results import verify as future
    frozen=future(ROOT);assets=sync(check=True)
    chapter=next((REPORT/'03_MD/当前版').glob('05 *.md'))
    claim=re.search(r'邳州刚性方案(\d+)—(\d+)年保持([\d.]+) MVA',chapter.read_text())
    assert claim,'邳州初始容量保持范围缺少可核对表述'
    annual=json.loads((REPORT/'试算/20261017_逐年需求驱动2035/pizhou/rigid/summary.json').read_text())['annual']
    location=str(chapter.relative_to(ROOT))+':'+str(chapter.read_text()[:claim.start()].count('\n')+1)
    check_capacity_period(annual,int(claim[1]),int(claim[2]),float(claim[3]),location)
    old,new=package(BASE),package(TARGET);root=xml(new['word/document.xml'])
    goal=package(CLEAN);expected=xml(goal['word/document.xml'])
    accepted=view(root,True);rejected=view(root,False)
    for tree,target,parts in [(accepted,expected,goal),(rejected,xml(old['word/document.xml']),old)]:
        assert [fingerprint(n,new) for n in tree.find(qn('body'))]==[fingerprint(n,parts) for n in target.find(qn('body'))],'接受/拒绝本轮修订后的正文或图片不对应'
    for name in old:
        if 'comment' in name.lower() or name=='word/people.xml':assert new[name]==old[name],'原批注部件发生变化:'+name
    oldroot=xml(old['word/document.xml'])
    for tag in ['commentRangeStart','commentRangeEnd','commentReference']:
        assert Counter(root.xpath('.//w:'+tag+'/@w:id',namespaces=NS))==Counter(oldroot.xpath('.//w:'+tag+'/@w:id',namespaces=NS)),'批注锚点缺失:'+tag
    oldrevisions=oldroot.xpath('.//w:ins|.//w:del',namespaces=NS)
    for item in oldrevisions:
        candidates=root.xpath('.//w:'+E.QName(item).localname+'[@w:id="'+item.get(qn('id'))+'"]',namespaces=NS)
        assert len(candidates)==1 and dict(candidates[0].attrib)==dict(item.attrib),'甲方已有修订属性发生变化'
    record=json.loads((REVIEW/'原文件修订核验.json').read_text())
    assert sha(TARGET)==record['修后SHA256'] and sha(BASE)==record['修前SHA256'] and sha(CLEAN)==record['目标清稿SHA256']
    revisions=root.xpath('.//*[@w:author="xh"]',namespaces=NS)
    dates=[]
    for node in revisions:
        date=datetime.fromisoformat(node.get(qn('date')))
        assert date.utcoffset().total_seconds()==28800
        assert (date.strftime('%Y-%m-%d')=='2026-10-18' and 8<=date.hour<9) or (date.strftime('%Y-%m-%d')=='2026-10-19' and 8<=date.hour<15)
        dates.append(node.get(qn('date')))
    assert len({n.get(qn('id')) for n in revisions})==len(revisions),'本轮修订ID重复'
    assert len(dates)==len(record['时间逐项映射'])
    assert xml(new['word/settings.xml']).find(qn('trackRevisions')) is not None
    doc=doc_from_parts(new,accepted);clean=Document(CLEAN)
    assert [[c.text for row in t.rows for c in row.cells] for t in doc.tables]==[[c.text for row in t.rows for c in row.cells] for t in clean.tables],'最终视图表格未逐格对应当前数据'
    namespace={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    assert [''.join(m.xpath('.//m:t/text()',namespaces=namespace)) for m in accepted.xpath('.//m:oMath',namespaces=namespace)]==[''.join(m.xpath('.//m:t/text()',namespaces=namespace)) for m in expected.xpath('.//m:oMath',namespaces=namespace)],'原生公式内容不对应'
    style_ids={s.style_id for s in doc.styles}
    assert set(accepted.xpath('.//w:pStyle/@w:val|.//w:rStyle/@w:val|.//w:tblStyle/@w:val',namespaces=NS))<=style_ids,'最终视图有未定义样式'
    from check_report_policy import check_figure_paragraphs
    policy=json.loads((ROOT/'tools/report_policy.json').read_text())
    figures=check_figure_paragraphs(doc,policy,41)
    from report_math_layout import check_equation_layout
    equations=check_equation_layout(doc)
    expected_equations=sum(len(re.findall(r'\\tag\{',p.read_text())) for p in (REPORT/'03_MD/当前版').glob('0[1-6] *.md'))
    assert len(equations)==expected_equations
    anchors={b.get(qn('name')):b for b in accepted.findall('.//'+qn('bookmarkStart'))}
    toc=accepted.xpath('.//w:body/w:p[.//w:instrText[contains(text(),"PAGEREF")]]',namespaces=NS)
    assert len(toc)==46
    heading_map={title:anchor for title,level,anchor in json.loads((REPORT/'05_review/当前版/目录标题.json').read_text())}
    for p in toc:
        instruction=''.join(p.xpath('.//w:instrText/text()',namespaces=NS));alias=re.search(r'PAGEREF\s+(\w+)',instruction)[1]
        assert alias.startswith('Report1017_') and alias in anchors
        title=''.join(p.xpath('.//w:hyperlink//w:t/text()',namespaces=NS))
        assert alias=='Report1017_'+heading_map[title]
        assert ''.join(anchors[alias].getparent().xpath('.//w:t/text()',namespaces=NS))==title,'目录指向错误标题:'+title
        assert p.xpath('.//w:hyperlink/@w:anchor',namespaces=NS)==[alias]
    for source in sorted((REPORT/'03_MD/当前版').glob('0[1-6] *.md')):
        before=(REVIEW/'修前原件'/source.relative_to(ROOT)).read_text()
        previous=re.findall(r'^#{1,3} .+',before,re.M)
        after=re.findall(r'^#{1,3} .+',source.read_text(),re.M)
        previous=[x.replace('5.3.1 正常运行与主变N-1','5.3.1 供电安全及供电能力约束') for x in previous]
        assert previous==after,'六章既有架构发生变化:'+source.name
    original_tables=json.loads((REVIEW/'修前原件/研究报告/02图表/当前版/计算表.json').read_text())
    tables=json.loads((ASSET/'计算表.json').read_text())
    matrix_keys=[key for key in original_tables if re.fullmatch(r'lookup_matrix_[0-2]_1\.(?:00|10|15)',key)]
    assert len(matrix_keys)==9
    from report_matrix_display import adjoining_bands
    for key in matrix_keys:
        before=[line for line in original_tables[key].splitlines() if line.startswith('|') and '～' in line]
        after=[line for line in tables[key].splitlines() if line.startswith('|') and '～' in line]
        assert len(before)==len(after)==12
        for a,b in zip(before,after):
            aa=[v.strip() for v in a.strip('|').split('|')];bb=[v.strip() for v in b.strip('|').split('|')]
            assert aa[:2]==bb[:2] and bb[2:]==adjoining_bands(aa[2:]),'推荐分档与原始容量下界不对应:'+key
    result={'状态':'通过','阶段':'final' if final else 'docx','锁定案例':frozen,'数据':assets,'六章架构':'与修前一致，仅修改指定5.3.1标题','批注数':len(xml(old['word/comments.xml'])),'原批注部件及锚点':'全部保留','甲方既有修订':len(oldrevisions),'本轮xh修订':len(revisions),'显示时间范围':[min(dates),max(dates)],'矩阵':'9表432单元按原始下界形成相邻档共用界值的推荐带；原始节点不变','逐格表格':len(doc.tables),'原生公式':len(equations),'原文件目录链接':len(toc),**figures,'修订稿SHA256':sha(TARGET),'清稿SHA256':sha(CLEAN)}
    if final:
        proof=json.loads((REVIEW/'WPS版面核验.json').read_text());assert proof['状态']=='通过'
        for file,key in [('DOCX','DOCX_SHA256'),('最终视图副本','最终视图副本SHA256'),('PDF','PDF_SHA256')]:assert sha(ROOT/proof[file])==proof[key],file+'变化后须重新检查WPS输出'
        assert proof['DOCX_SHA256']==sha(TARGET)
        result['WPS最终视图']=proof
    (REVIEW/(('final' if final else 'docx')+'独立核验.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    return result

def verify_release():
    """核验本次交付及其历史原件；不将旧版结果断言套用到已授权的新案例。"""
    result=verify(final=True)
    release=json.loads((REVIEW/'交付登记.json').read_text())
    assert release['状态']=='三阶段核验及逐段复核通过'
    for name,digest in release['当前文件SHA256'].items():
        assert sha(ROOT/name)==digest,'本轮交付已变化，须重新复核:'+name
    from report_prose_review import collect
    current=collect(ROOT,include_discussion=False)
    reviewed=json.loads((REVIEW/'正文逐段复核.json').read_text())['正文段落']
    assert len(current)==len(reviewed)
    for row,now in zip(reviewed,current):
        assert row['位置']==now['位置'] and row['SHA256']==now['SHA256'],'正文复核已过期:'+now['位置']
        assert row['复核结论']=='通过' and row['理由']
    link=json.loads((REVIEW/'上一版交付衔接.json').read_text())
    parent=ROOT/link['父交付登记'];assert sha(parent)==link['父交付登记SHA256']
    expected=json.loads(parent.read_text())['当前文件SHA256']
    assert set(link['原件映射'])==set(expected)
    for name,digest in expected.items():
        item=link['原件映射'][name]
        assert item['SHA256']==digest and sha(ROOT/item['path'])==digest,'前版原件不一致:'+name
    for name,item in release['历史定档原件'].items():
        assert sha(ROOT/item['原件归档'])==item['原SHA256'],'初步定档原件变化:'+name
    facts=json.loads((REVIEW/'事实核验.json').read_text())
    assert facts['状态']=='通过' and facts['原始缺口及适用条件']
    for name,digest in facts['依据SHA256'].items():
        assert sha(ROOT/name)==digest,'事实核验依据变化:'+name
    prior=ROOT/facts['继承引文核读']
    records=json.loads(prior.read_text())['记录'];assert len(records)==95
    for row in records:
        for key,digest_key in [('审定依据','SHA256'),('原始文件','原始文件SHA256')]:
            if digest_key in row:assert sha(ROOT/row[key])==row[digest_key],'引文原始依据变化:'+row[key]
    assert release['文字复核']['规则扫描代替逐段判断'] is False
    print(json.dumps({'状态':'通过','本轮正文复核':len(current),'既有引文依据':len(records),'前版原件衔接':len(expected),'批注保留':result['批注数']},ensure_ascii=False))
    return release

if __name__=='__main__':
    import sys
    print(json.dumps(verify('--final' in sys.argv),ensure_ascii=False))
