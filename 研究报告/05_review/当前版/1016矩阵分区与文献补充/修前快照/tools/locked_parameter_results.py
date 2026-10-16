"""锁定参数结果核验及复算入口；默认只读，复算/寻优必须另指定空目录。"""
from pathlib import Path
from concurrent.futures import ProcessPoolExecutor
from multiprocessing import get_context
import argparse
import hashlib
import itertools
import json
import math
import sys
import numpy as np
import pandas as pd
from scipy.optimize import linprog
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from 实验.研究.rebuild_2026.matrix_20261012 import model
TRIAL=ROOT/'研究报告/试算/20261012_参数联合情景锁定'
PUBLIC=ROOT/'研究报告/01研究推荐/当前版/参数联合情景_20261012'
SOURCE=ROOT/'研究报告/数据来源/当前版/参数联合情景_20261012'
REVIEW=ROOT/'研究报告/05_review/当前版/1012参数结果锁定'
STATION=ROOT/'实验/研究/data/tuomin/补充数据/徐州地区变电站全停全转情况分析汇总20250705.xlsx'
FIELDS=model.FIELDS
LABEL=dict(zip(FIELDS,['分布式电源装机渗透率','年最大净负荷平均增长率',
    '电缆长度占比','初始转供率','储能支撑持续时间_h','负荷分散系数']))

OTHER=dict(eta=[i/4 for i in range(13)],growth_rate=[0,.02,.03,.04,.055,.07,.08,.1],
           support_duration_h=[2,4,6],diversity=[1,1.1,1.15])
CUTS=dict(eta=[1,2],growth_rate=[.02,.04,.07],cable_fraction=[.1,.3],
          initial_transfer_fraction=[.3,.5,.7],support_duration_h=[2,4],diversity=[1.1])
TEXT=dict(eta=['≤1','(1,2]','>2'],growth_rate=['≤2%','(2%,4%]','(4%,7%]','>7%'],
          cable_fraction=['≤10%','(10%,30%]','>30%'],
          initial_transfer_fraction=['≤30%','(30%,50%]','(50%,70%]','>70%'],
          support_duration_h=['≤2 h','(2 h,4 h]','>4 h'],diversity=['≤1.10','>1.10'])

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def dump(p,value):p.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

def key(r):return tuple(round(float(r[k]),9) for k in FIELDS)

def binof(k,x):return sum(x>c+1e-9 for c in CUTS[k])

def bkey(r):return tuple(binof(k,r[k]) for k in FIELDS)

def allowed_pairs():
    return sorted({(q,r) for s in model.AREA_SCENARIOS.values()
                   for q,r in itertools.product(s['cable_nodes'],s['transfer_nodes'])})

def tasks():
    regular=[dict(zip(OTHER,v),cable_fraction=q,initial_transfer_fraction=r,node_purpose='主网格')
             for v in itertools.product(*OTHER.values()) for q,r in allowed_pairs()]
    # 既有校准研究已经识别的7%增长边界，预声明补核而不按结果挑点。
    boundary=[dict(eta=e,growth_rate=g,cable_fraction=q,initial_transfer_fraction=.3,
                   support_duration_h=h,diversity=1.15,node_purpose='7%增长边界补充')
              for e,g,q,h in itertools.product([.5,1.5,2.5],[.071,.075],[.1,.3,.5],[2,4,6])]
    return regular+boundary

def presentation_tasks():
    return tasks()

def station_evidence():
    sheet=load_workbook(STATION,read_only=True,data_only=True).worksheets[0]
    rows=[]
    for i,r in enumerate(sheet.iter_rows(min_row=3,values_only=True),3):
        if r[0] not in ['邳州','徐州市区'] or not r[2]:continue
        if not isinstance(r[9],(float,int)) or not isinstance(r[10],(float,int)):continue
        rows.append(dict(地区=r[0],电压等级=r[1],变电站=r[2],原表行号=i,
                         站外最大可供负荷_MW=r[9],损失负荷_MW=r[10],原表转供率=r[12]))
    d=pd.DataFrame(rows);summary=[]
    for region,g in d.groupby('地区'):
        for scope,sub in [('全部电压等级',g),('仅110 kV',g[g.电压等级.eq('110kV')])]:
            a=sub['站外最大可供负荷_MW'].sum();b=sub['损失负荷_MW'].sum()
            summary.append(dict(地区=region,统计范围=scope,站点数=len(sub),
                                可转供负荷合计_MW=a,损失负荷合计_MW=b,
                                负荷加权转供率=a/(a+b),
                                计算口径='站外最大可供负荷合计／（站外最大可供负荷合计＋损失负荷合计）'))
    return d,pd.DataFrame(summary)

def corr_frame(d,fields,method='pearson'):
    return d[fields].corr(method=method).rename(index=LABEL,columns=LABEL).rename_axis('参数').reset_index()

def vif_values(d,fields):
    corr=d[fields].corr().to_numpy();v=np.diag(np.linalg.inv(corr));direct=[]
    z=(d[fields]-d[fields].mean())/d[fields].std(ddof=0)
    for i in range(len(fields)):
        y=z.iloc[:,i].to_numpy();x=np.delete(z.to_numpy(),i,axis=1)
        a=np.linalg.lstsq(x,y,rcond=None)[0]
        direct.append(float(y@y/((y-x@a)@(y-x@a))))
    assert np.allclose(v,direct,rtol=1e-10,atol=1e-10)
    return pd.DataFrame({'参数':[LABEL[f] for f in fields],'VIF':v})

def correlations(d):
    result=[]
    for f in FIELDS:
        fixed=[k for k in FIELDS if k!=f]
        n=d.groupby(fixed)[f].transform('nunique');sub=d[n>=2].copy()
        x=sub[f]-sub.groupby(fixed)[f].transform('mean')
        y=sub.clr-sub.groupby(fixed).clr.transform('mean')
        xy=float(x@y/np.sqrt((x@x)*(y@y)))
        assert abs(xy-float(x.corr(y)))<1e-12
        spans=sub.groupby(fixed).capacity_mva.agg(lambda a:a.max()-a.min())
        result.append(dict(参数=LABEL[f].removesuffix('_h'),有效受控组数=len(spans),
                           有效节点数=len(sub),组内Pearson=xy,
                           主变容量响应组数=int((spans>1e-7).sum()),
                           主变容量响应组占比=float((spans>1e-7).mean()),
                           统计方式='固定其余五项；联合域内每组至少两个有效节点；组内去均值，节点等权，保留零响应组'))
    return pd.DataFrame(result)

def matrix_cells(d):
    grouping=[d[k].map(lambda x,k=k:binof(k,x)) for k in FIELDS]
    cells=[];endpoints=set()
    for bs,g in d.groupby(grouping,sort=True):
        lo=float(g.clr.min());hi=float(g.clr.max())
        endpoints.update([int(g.clr.idxmin()),int(g.clr.idxmax())])
        text=f'{math.floor((lo+1e-10)*100)/100:.2f}～{math.ceil((hi-1e-10)*100)/100:.2f}'
        cells.append(dict(zip(['bin_'+k for k in FIELDS],bs))|
                     {LABEL[k]:TEXT[k][i] for k,i in zip(FIELDS,bs)}|
                     dict(有效样本数=len(g),容载比下界=lo,容载比上界=hi,容载比区间=text,
                          本单元电缆节点='、'.join(f'{q:.0%}' for q in sorted(g.cable_fraction.unique())),
                          本单元转供节点='、'.join(f'{r:.0%}' for r in sorted(g.initial_transfer_fraction.unique())),
                          下界节点='；'.join(map(str,key(d.loc[g.clr.idxmin()]))),
                          上界节点='；'.join(map(str,key(d.loc[g.clr.idxmax()])))))
    cells=pd.DataFrame(cells)
    assert cells.有效样本数.sum()==len(d)
    return cells,endpoints

def matrix_frames(cells):
    frames={};index=[]
    fixed=['bin_cable_fraction','bin_support_duration_h','bin_diversity']
    for (q,h,kd),g in cells.groupby(fixed,sort=True):
        rs=sorted(g.bin_initial_transfer_fraction.unique())
        assert rs==[0,1,2,3],'每张矩阵必须由有效样本覆盖全部四个转供区间'
        name=f'矩阵_电缆{q+1}_时间{h+1}_分散{kd+1}'
        records=[]
        for (growth,eta),row in g.groupby(['bin_growth_rate','bin_eta'],sort=True):
            by={int(r.bin_initial_transfer_fraction):r.容载比区间 for r in row.itertuples()}
            assert set(by)==set(rs),'不得补造没有有效样本的单元'
            records.append({'年最大净负荷平均增长率':TEXT['growth_rate'][growth],
                            '分布式电源装机渗透率':TEXT['eta'][eta]}|
                           {f'初始转供率{TEXT["initial_transfer_fraction"][r]}':by[r] for r in rs})
        frames[name]=pd.DataFrame(records)
        index.append(dict(工作表=name,电缆长度占比=TEXT['cable_fraction'][q],
                          储能支撑持续时间=TEXT['support_duration_h'][h],
                          负荷分散系数=TEXT['diversity'][kd],
                          转供率列='、'.join(TEXT['initial_transfer_fraction'][r] for r in rs),
                          有效区间单元数=len(g)))
    assert len(frames)==18
    return frames,pd.DataFrame(index)

def independent_initial_check(inp):
    """独立星形LP检查初始转供指标，与解析标定程序分别组装。"""
    base=inp['base'];es=[e for e in inp['edges'] if e['existing']];flows=[]
    for s in range(5):
        arcs=[]
        for e in es:
            a,b=e['a'],e['b']
            if a//2==s:arcs.append((a,b,e['capacity']))
            elif b//2==s:arcs.append((b,a,e['capacity']))
        rows=[];rhs=[]
        for f in range(10):
            rows.append([float(a==f) for a,b,c in arcs]);rhs.append(base[f//2]/2)
            rows.append([float(b==f) for a,b,c in arcs]);rhs.append(max(inp['thermal'][f]-base[f//2]/2,0))
        for j in range(5):
            if j==s:continue
            rows.append([float(b//2==j) for a,b,c in arcs]);rhs.append(max(.95*100-base[j],0))
        rows.append([1.]*len(arcs));rhs.append(inp['existing_budget'][s])
        r=linprog(-np.ones(len(arcs)),A_ub=rows,b_ub=rhs,bounds=[(0,c) for a,b,c in arcs],method='highs')
        assert r.success,r.message
        flows.append(-r.fun)
    assert np.max(np.abs(np.array(flows)-inp['initial_transfer_audit']['station_transfer_mw']))<1e-7
    return dict(station_transfer_mw=flows,weighted_rate=sum(flows)/sum(base))

def public_frame(d):
    rows=[]
    for r in d.to_dict('records'):
        rows.append({LABEL[k]:r[k] for k in FIELDS}|dict(容载比=r['clr'],
            主变容量_MVA=r['capacity_mva'],储能额定功率_MW=r['storage_mw'],储能额定能量_MWh=r['storage_mwh'],
            储能有效支撑功率_MW=r['storage_effective_mw'],新增联络合计_条=r['line_count'],
            其中独立新出线_条=r['new_receiving_feeder_count'],费用现值_万元=r['cost'],
            主变费用现值_万元=r['cost_components']['transformer'],储能费用现值_万元=r['cost_components']['storage'],
            线路费用现值_万元=r['cost_components']['line'],适用联合情景='；'.join(r['joint_area_scenarios']),节点用途=r['node_purpose']))
    return pd.DataFrame(rows)


def display_frame(cells):
    return cells[[LABEL[k] for k in FIELDS]+['有效样本数','容载比下界','容载比上界','容载比区间',
        '本单元电缆节点','本单元转供节点','下界节点','上界节点']]


def equal_frame(expected, actual, name):
    try:
        pd.testing.assert_frame_equal(expected.reset_index(drop=True),actual.reset_index(drop=True),
            check_dtype=False,check_exact=False,rtol=1e-12,atol=1e-9)
    except AssertionError as error:
        raise AssertionError(f'{name}: {error}') from error


def verify_workbook(frames, index):
    book=load_workbook(PUBLIC/'推广矩阵与相关性分析_最新收资统一结果.xlsx',read_only=True,data_only=True)
    checked=0
    for name,frame in frames.items():
        sheet=book[name]
        assert (sheet.max_row,sheet.max_column)==(len(frame)+1,len(frame.columns)),name
        actual=sheet.iter_rows(values_only=True)
        assert list(next(actual))==list(frame.columns),name
        for arow,brow in zip(actual,frame.itertuples(index=False,name=None)):
            for a,b in zip(arow,brow):
                if isinstance(b,(int,float,np.integer,np.floating)) and not isinstance(b,(bool,np.bool_)):
                    assert a is not None and math.isclose(a,float(b),rel_tol=1e-12,abs_tol=1e-9),(name,a,b)
                else:assert a==b,(name,a,b)
                checked+=1
    book.close()
    # 直接读OOXML链接，避免载入整本大工作簿。
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
        'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    relns='{http://schemas.openxmlformats.org/package/2006/relationships}'
    with ZipFile(PUBLIC/'推广矩阵与相关性分析_最新收资统一结果.xlsx') as z:
        wb=ET.fromstring(z.read('xl/workbook.xml'))
        rid=next(s.get('{'+ns['r']+'}id') for s in wb.findall('s:sheets/s:sheet',ns) if s.get('name')=='矩阵索引')
        rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        target=next(s.get('Target') for s in rels.findall(relns+'Relationship') if s.get('Id')==rid)
        path=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        sheet=ET.fromstring(z.read(path));links=sheet.findall('s:hyperlinks/s:hyperlink',ns)
        rp=Path(path);rels=ET.fromstring(z.read(str(rp.parent/'_rels'/(rp.name+'.rels'))))
        targets={s.get('Id'):s.get('Target') for s in rels.findall(relns+'Relationship')}
        assert len(links)==18
        for link,name in zip(links,index.工作表):assert targets[link.get('{'+ns['r']+'}id')]==f"#'{name}'!A1"
    return checked


def lp_task(row):
    result=model.independent_lp_check(row,row['units'],row['modules'],row['built'])
    assert result['verified'],(key(row),result)
    return dict(节点=key(row),**result)


def frames_from_rows(rows):
    d=pd.DataFrame([r for r in rows if r['verified']]).sort_values(FIELDS).reset_index(drop=True)
    cells,endpoints=matrix_cells(d);matrices,index=matrix_frames(cells)
    keep=[f for f in FIELDS if f!='cable_fraction'];five=d[keep].drop_duplicates().reset_index(drop=True)
    frames={'参数与容载比组内相关':correlations(d),'六参数输入相关':corr_frame(d,FIELDS),
        '六参数输入秩相关':corr_frame(d,FIELDS,'spearman'),'六参数输入VIF':vif_values(d,FIELDS),
        '五参数输入相关':corr_frame(five,keep),'五参数输入VIF':vif_values(five,keep),
        '矩阵索引':index,**matrices,'有效区间汇总':display_frame(cells),'有效样本':public_frame(d)}
    return d,cells,endpoints,index,frames


def verify(lp=False,workers=8):
    manifest=json.loads((REVIEW/'锁定清单.json').read_text())
    for record in manifest['文件']:
        path=ROOT/record['path']
        assert path.is_file() and sha(path)==record['sha256'],f'锁定文件变化: {record["path"]}'
    rows=list(map(json.loads,(TRIAL/'统一原始求解记录.jsonl').read_text().splitlines()))
    assert len(rows)==34686 and len({key(r) for r in rows})==len(rows)
    expected={key(r):r for r in tasks()};assert set(expected)=={key(r) for r in rows}
    for r in rows:
        assert r['node_purpose']==expected[key(r)]['node_purpose']
        assert r['joint_area_scenarios']==model.matching_scenarios(r) and r['joint_domain_passed']
        if not r['verified']:assert r['status']=='Infeasible',(key(r),r['status'])
    assert sum(r['verified'] for r in rows)==32112
    assert sum(r['record_origin']=='上轮同模型节点复用' for r in rows)==27198
    d,cells,endpoints,index,frames=frames_from_rows(rows)
    assert len(cells)==864 and cells.有效样本数.min()==2
    assert {bkey(r) for r in tasks()}=={tuple(r['bin_'+k] for k in FIELDS) for r in cells.to_dict('records')}
    assert d.relative_gap.max()<1.1e-7
    assert max(r['physical_check']['max_violation_mw'] for r in rows if r['verified'])<1e-5
    # 从设备配置及当前冻结价格独立回算容量、储能、三项费用与容载比。
    cost_inputs={}
    for r in rows:
        if not r['verified']:continue
        cost_key=(r['cable_fraction'],r['support_duration_h'],r['storage_price'],r['line_length_km'])
        if cost_key not in cost_inputs:cost_inputs[cost_key]=model.make_input(r)
        inp=cost_inputs[cost_key];units=np.asarray(r['units']);modules=np.asarray(r['modules'])
        peak=250*(1+r['growth_rate'])**4
        assert abs(peak-r['peak_mw'])<1e-9
        capacity=float(units.sum());assert abs(capacity-r['capacity_mva'])<1e-9
        assert abs(capacity/peak-r['clr'])<1e-9
        assert abs(modules.sum()*.1-r['storage_mw'])<1e-9
        assert abs(modules.sum()*.215-r['storage_mwh'])<1e-9
        assert abs(modules.sum()*inp['storage_effective']-r['storage_effective_mw'])<1e-9
        assert r['line_count']==len(r['built'])
        raw=dict(transformer=float(np.where(units[:,:2]>50,units[:,:2],0).sum()*inp['transformer_cost_per_purchased_mva']
                     +units[:,2].sum()/50*inp['third_cost']),
                 storage=float(modules.sum()*inp['storage_unit_cost']*inp['p']['storage_price']),
                 line=len(r['built'])*inp['line_cost'])
        for k,v in raw.items():assert abs(v*inp['factors'][k]-r['cost_components'][k])<1e-5,(key(r),k)
        assert abs(sum(r['cost_components'].values())-r['cost'])<1e-4
    pairs=[('统一有效样本.csv','有效样本'),('推广矩阵有效区间.csv','有效区间汇总'),
        ('固定其余五项后的组内相关.csv','参数与容载比组内相关'),('六参数输入相关.csv','六参数输入相关'),
        ('五参数输入相关.csv','五参数输入相关')]
    for filename,name in pairs:equal_frame(frames[name],pd.read_csv(PUBLIC/filename),filename)
    detail,actual=station_evidence()
    equal_frame(detail,pd.read_csv(SOURCE/'实际转供收资站点及来源.csv'),'实际转供逐站来源')
    equal_frame(actual,pd.read_csv(SOURCE/'实际转供率统计范围核验.csv'),'实际转供统计')
    frames['实际转供统计核验']=actual
    initial=[]
    for rho in sorted(d.initial_transfer_fraction.unique()):
        row=d[d.initial_transfer_fraction.eq(rho)].iloc[0].to_dict()
        check=independent_initial_check(model.make_input(row))
        assert abs(check['weighted_rate']-rho)<1e-8
        initial.append(check)
    checked=verify_workbook(frames,index)
    selected=set(endpoints)|set(d[d.node_purpose.ne('主网格')].index)
    for _,group in d.groupby(['eta','initial_transfer_fraction','support_duration_h'],sort=True):
        selected.add(int(group.sort_values(['growth_rate','cable_fraction','diversity']).index[len(group)//2]))
    proofs=json.loads((REVIEW/'矩阵端点及分层独立LP核验.json').read_text())
    assert len(proofs)==len(selected)==2069 and all(r['verified'] for r in proofs)
    assert {tuple(r['节点']) for r in proofs}=={key(d.loc[i]) for i in selected}
    if lp:
        print('独立LP重新核验2069节点',flush=True)
        with ProcessPoolExecutor(max_workers=workers,mp_context=get_context('spawn')) as pool:
            newproof=list(pool.map(lp_task,[d.loc[i].to_dict() for i in sorted(selected)]))
        assert sum(x['checked_pairs'] for x in newproof)==60948
    result=dict(核验通过=True,锁定文件数=len(manifest['文件']),设计节点数=len(rows),有效节点数=len(d),
        不可行节点数=len(rows)-len(d),矩阵数=len(index),有效区间数=len(cells),每格最少有效样本=int(cells.有效样本数.min()),
        原始库回算容量储能及三项费用节点数=len(d),工作簿复算回读单元格数=checked,矩阵索引链接数=18,
        初始转供独立LP节点数=len(initial),独立LP原核验节点数=len(proofs),本轮独立LP重新核验节点数=len(proofs) if lp else 0,
        本轮独立LP重新核验工况数=60948 if lp else 0,核验过程不写稿件=True)
    return result,frames


def new_output(path):
    path=path.resolve()
    protected=[TRIAL,PUBLIC,SOURCE,REVIEW,ROOT/'研究报告/甲方反馈',Path(model.__file__).parent]
    assert not any(path==p or path.is_relative_to(p) or p.is_relative_to(path) for p in protected), '复算请使用独立临时目录'
    assert not path.exists() or not any(path.iterdir()),'复算目录必须为空'
    path.mkdir(parents=True,exist_ok=True)
    return path


def export(frames,path):
    path=new_output(path)
    with pd.ExcelWriter(path/'锁定数据复算.xlsx',engine='openpyxl') as writer:
        for name,frame in frames.items():frame.to_excel(writer,sheet_name=name,index=False)
    for name in ['有效样本','有效区间汇总','参数与容载比组内相关','六参数输入相关','五参数输入相关']:
        frames[name].to_csv(path/(name+'.csv'),index=False,encoding='utf-8-sig')
    print('数值复算输出:',path,flush=True)


def solve_task(parameters):
    result=model.solve(dict(parameters,gap=1e-7,seconds=30))
    assert result['verified'] or result['status']=='Infeasible',(key(result),result['status'])
    result['record_origin']='锁定模型独立复算'
    return result


def solve_to(path,workers,limit):
    path=new_output(path);jobs=tasks()
    if limit is not None:
        assert limit>0
        jobs=jobs[:limit]
    dump(path/'复算设计.json',dict(用途='独立重算；不替换已锁定数据',参数字段=FIELDS,
        其余四项节点=OTHER,联合情景=model.AREA_SCENARIOS,分档=TEXT,任务数=len(jobs),完整设计数=len(tasks()),
        模型SHA256=sha(Path(model.__file__))))
    with ProcessPoolExecutor(max_workers=workers) as pool,(path/'复算原始记录.jsonl').open('w') as f:
        for i,result in enumerate(pool.map(solve_task,jobs),1):
            f.write(json.dumps(result,ensure_ascii=False)+'\n');f.flush()
            if i%1000==0 or i==len(jobs):print('独立重算',i,'/',len(jobs),flush=True)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--lp',action='store_true',help='重新独立LP核验全部2069个端点/分层节点')
    parser.add_argument('--workers',type=int,default=8)
    parser.add_argument('--report',type=Path,help='核验记录输出位置，须在锁定文件之外')
    parser.add_argument('--export-to',type=Path,help='从原始库重建数值表，输出到另一个空目录')
    parser.add_argument('--solve-to',type=Path,help='从锁定模型重新寻优，输出到另一个空目录')
    parser.add_argument('--limit',type=int,help='独立寻优时只算前N项，供复现入口自检')
    args=parser.parse_args()
    assert args.workers>0
    if args.limit is not None and args.solve_to is None:parser.error('--limit只用于--solve-to')
    if args.solve_to:
        if args.lp or args.export_to or args.report:parser.error('--solve-to不能与核验/导出混用')
        solve_to(args.solve_to,args.workers,args.limit);return
    if args.report:
        path=args.report.resolve()
        assert path.suffix=='.json','核验记录使用JSON文件'
        manifest=json.loads((REVIEW/'锁定清单.json').read_text())
        assert path!=REVIEW/'锁定清单.json' and path not in {ROOT/r['path'] for r in manifest['文件']}
    result,frames=verify(args.lp,args.workers)
    if args.export_to:export(frames,args.export_to)
    if args.report:
        args.report.parent.mkdir(parents=True,exist_ok=True);dump(args.report,result)
    print(json.dumps(result,ensure_ascii=False,indent=2),flush=True)


if __name__=='__main__':main()
