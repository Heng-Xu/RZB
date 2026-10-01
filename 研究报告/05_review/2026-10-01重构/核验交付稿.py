from pathlib import Path
from collections import defaultdict
import json,re,hashlib,zipfile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from PIL import Image
from fontTools.ttLib import TTCollection
import pymupdf,subprocess
ROOT=Path(__file__).resolve().parents[3];R=Path(__file__).parent;O=ROOT/'研究报告/04_word/2026-10-01重构';A=ROOT/'研究报告/02图表/2026-10-01重构';M=ROOT/'研究报告/01研究推荐/2026-10-01条件矩阵';S=ROOT/'研究报告/数据来源/2026-10-01本轮来源说明'
def read(p):return json.loads(p.read_text())
def main():
 d=Document(O/'研究报告_WPS兼容稿.docx');pdf=pymupdf.open(O/'研究报告_WPS兼容稿.pdf');assert len(pdf)==150,len(pdf)
 paras=[p.text for p in d.paragraphs];body='\n'.join(paras);cells='\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells);alltext=body+'\n'+cells
 for token in ['[[TABLE:','[[FIG:','[[SECTION:','FORMULA_','{width=','TODO','待补充','ChatGPT','作为AI']:
  assert token not in alltext,token
 heads=[p.text for p in d.paragraphs if p.style.name=='Heading 1'];assert len(heads)==7,heads;assert all('附录' not in h for h in heads)
 assert heads[4].startswith('第五章') and heads[5].startswith('第六章 结论'),heads
 assert len(d.tables)==122,len(d.tables);assert len(d.inline_shapes)==54,len(d.inline_shapes)
 captions={}
 for kind in ['表','图']:
  numbers=defaultdict(list)
  for t in paras:
   m=re.match('^'+kind+r'(\d+)-(\d+)\s',t)
   if m:numbers[int(m[1])].append(int(m[2]))
  for c,ns in numbers.items():assert ns==list(range(1,len(ns)+1)),(kind,c,ns)
  captions[kind]={str(c):len(ns) for c,ns in numbers.items()}
 assert sum(captions['表'].values())==122,captions;assert sum(captions['图'].values())==15,captions
 formulas=read(A/'公式/公式清单.json');assert len(formulas)==39
 eqids={x['number'] for x in formulas}
 for ref in re.findall(r'式[（(]?(\d+-\d+)',alltext):assert ref in eqids,ref
 alt=[shape._inline.docPr.get('descr','') for shape in d.inline_shapes];assert sum(x.startswith('公式（') for x in alt)==39
 for p in formulas:
  im=Image.open(A/'公式'/(p['number']+'.png'));assert im.width/p['width_in']>450
  assert p['width_in']<=4.95+1e-6
 # 公式中文采用明确覆盖字形的字体，正文图像只含像素，不依赖阅读端字体。
 cjk=set(re.findall('[\u4e00-\u9fff]',''.join(x['tex'] for x in formulas)))
 fontpath=subprocess.check_output(['fc-match','-f','%{file}','Noto Serif CJK SC'],text=True)
 cmap=set()
 for f in TTCollection(fontpath).fonts:cmap.update(f.getBestCmap())
 assert all(ord(c) in cmap for c in cjk),cjk
 refs=[p for p in paras if re.match(r'^\[\d+\]',p)];assert [int(re.match(r'^\[(\d+)\]',p)[1]) for p in refs]==list(range(1,81))
 for cit in re.findall(r'\[([\d—,,-]+)\]',alltext):assert all(1<=int(n)<=80 for n in re.findall(r'\d+',cit)),cit
 toc=read(R/'目录页码核对.json');fields=d._element.xpath('.//w:fldSimple');assert len(fields)==len(toc)
 expected={p['锚点']:str(p['页码']) for p in toc}
 for f in fields:
  anchor=re.search(r'PAGEREF\s+(\w+)',f.get(qn('w:instr')))[1];assert list(f.iter(qn('w:t')))[0].text==expected[anchor]
 for t in d.tables:
  assert len(t._tbl.tblGrid)==len(t.columns);assert t.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
  for row in t.rows:assert row._tr.xpath('./w:trPr/w:cantSplit')
 with zipfile.ZipFile(O/'研究报告_WPS兼容稿.docx') as z:
  for name in z.namelist():
   if name.endswith('.xml'):
    tree=etree.fromstring(z.read(name));assert not tree.xpath('//*[local-name()="trackRevisions"]')
   if name.endswith('.rels'):
    tree=etree.fromstring(z.read(name));assert all(not(e.get('TargetMode')=='External' and 'image' in e.get('Type','')) for e in tree)
  assert b'<Pages>150</Pages>' in z.read('docProps/app.xml')
 for i,p in enumerate(pdf):
  txt=p.get_text();assert '�' not in txt and '{width=' not in txt
  if i>0:assert txt.strip().endswith(str(i+1)),(i+1,txt[-80:])
 audit=read(R/'矩阵独立复核.json');assert len(audit)==72 and sum(x['站年'] for x in audit)==7056;assert all(x['核查']=='通过' for x in audit)
 inp=read(R/'情景输入独立复核.json');assert inp['总核对站年']==7056
 wb=load_workbook(M/'容载比条件矩阵与全部措施.xlsx',read_only=True,data_only=True);em=list(wb['弹性指标'].values);ms=list(wb['弹性措施'].values)
 assert len(em)==37 and len(ms)==37
 for row in em[1:]:assert all(float(v)>2 for v in row[3:7]),row
 assert len(list(wb['全部站年措施'].values))==7057
 notes=Document(S/'数据来源与参数计算说明.docx');nt='\n'.join(p.text for p in notes.paragraphs)
 assert '原光伏表未写明数值单位' in nt and '2021—2022年无同来源月度记录' in nt
 files=[O/'研究报告_WPS兼容稿.docx',O/'研究报告_WPS兼容稿.pdf',S/'数据来源与参数计算说明.docx',S/'数据来源与参数计算说明.pdf',S/'原格定位与建模输入.xlsx',M/'容载比条件矩阵与全部措施.xlsx']
 report={'状态':'通过','报告页数':len(pdf),'正文结构':'六章及参考文献，正文无附录','表格':captions['表'],'图件':captions['图'],'公式数':39,'参考文献数':80,'目录项数':len(toc),'情景数':18,'路径数':72,'复核站年':7056,'弹性年度矩阵':'36个地区情景、144个年度值全部大于2.0','公式处理':'480dpi图片嵌入；中文字体覆盖完整；公式带TeX替代文字','排版核验':'LibreOffice实际Word排版导出PDF，150页；人工抽查主要图表及39组公式','WPS云端':'未实测','语体规则':'lieflat-less-ai-tone白名单，结构、数字、公式及引文保持','原始资料说明':'单列文件；光伏单位采用假定及历史回推如实说明','文件':[]}
 for p in files:report['文件'].append({'路径':str(p.relative_to(ROOT)),'字节数':p.stat().st_size,'SHA256':hashlib.sha256(p.read_bytes()).hexdigest()})
 (R/'交付稿最终核验.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print('交付稿通过：150页、122表、15图、39公式；72路径及7056站年复核通过')
if __name__=='__main__':main()
