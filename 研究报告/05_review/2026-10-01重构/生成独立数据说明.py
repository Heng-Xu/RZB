from pathlib import Path
import sys,json,csv,hashlib,shutil
from collections import defaultdict
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook,Workbook
sys.path.insert(0,str(Path(__file__).resolve().parent))
from 装配Word import font,ROOT,REVIEW
from 整理数据与图表 import rows,RAW,BASE,OUT as ASSETS
from 整理方案与核查 import CASES,REG,SCHEME
DEST=ROOT/'研究报告/数据来源/2026-10-01本轮来源说明';DEST.mkdir(parents=True,exist_ok=True)
DOC=Document(REVIEW/'排版参照.docx');MD=[]
def paragraph(text):
 p=DOC.add_paragraph(text);p.paragraph_format.first_line_indent=Pt(24);p.paragraph_format.line_spacing=1.25
 for r in p.runs:font(r,12)
 MD.append(text+'\n\n')
def heading(text,level=1):
 p=DOC.add_heading(text,level);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.5
 np=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');np.append(ni);p._p.get_or_add_pPr().append(np)
 for r in p.runs:font(r,16 if level==1 else 13,True)
 MD.append('#'*level+' '+text+'\n\n')
def table(headers,data):
 def display(v):return format(v,'.10g') if isinstance(v,float) else str(v)
 t=DOC.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,head in enumerate(headers):t.rows[0].cells[i].text=str(head)
 for row in data:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=display(v)
 for i,row in enumerate(t.rows):
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(2)
    for r in p.runs:font(r,9,i==0,'宋体')
  rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
  if i==0:rp.append(OxmlElement('w:tblHeader'))
 pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders');pr.append(b)
 for name in ['top','bottom','left','right','insideH','insideV']:
  e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');b.append(e)
 MD.append('| '+' | '.join(map(str,headers))+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(map(display,row))+' |\n' for row in data)+'\n')
def main():
 heading('研究报告数据来源与参数计算说明（2026年10月）')
 paragraph('本文件与本轮六章研究报告、容载比条件矩阵工作簿配套，集中保存原始收资定位、建模换算、估计规则及费用参数依据。报告正文以邳州、市区的实际建模输入开展运行分析与寻优，丰县只作描述性对照。以下资料说明不作为研究报告正文附录。')
 heading('1 本轮计算版本与适用范围')
 paragraph('基准采用2026年10月1日同措施、同起点的四路径联合规划，2022—2025年刚性R≤2.0、弹性R≥2.001，市区同方案每年低于邳州至少0.001，弹性研究上侧为2.6。地区内两方案共同2021年反事实起点，主变、储能和新线状态非减。仅计新增措施，费用排序不强制。')
 paragraph('正常负荷转接用途为区县内年度负荷归属调整，区域任务守恒、容载比分母固定；电源仍在原站，反送上界为原反送峰加转出峰。共同10%为研究使用上限，不是标准故障转供比例。既有全站对仿真通道覆盖新增任务，新线在本轮正价格范围被支配而选零。')
 paragraph('矩阵包括12个两价格交叉条件和6个技术条件，共18组、72条地区方案路径、7056个站年。基准和9个需要改变布局或任务的情景完整计算；线路组合使用正费用及同任务可替代证明，沿用对应储能情景布局后独立复核。重复布局不计作新增观测样本。')
 paragraph('当前原始与派生结果目录：'+str(BASE.relative_to(ROOT))+'；矩阵目录：研究报告/01研究推荐/2026-10-01条件矩阵。9月29日ordered_guide及事故转供代理旧方案不作为本轮终稿数值依据。')
 heading('2 原始文件与完整性记录')
 paths=[RAW/'近5年容载比.xlsx',RAW/'逐月分县分布式光伏.xlsx',RAW/'2025设备负载统计表.xlsx',RAW/'邳州主变负载率.xlsx',ROOT/'实验/研究/data/tuomin/补充数据/徐州市区脱敏.zip',RAW/'光伏8760小时数据.xlsx',RAW/'110（35）kv设备明细.xlsx',RAW/'江苏徐州邢楼110千伏变电站主变扩建等工程建设规模及投资汇总表(1).xlsx',ROOT/'参考政策/DL-T+5729-2023+配电网规划设计技术导则.pdf']
 manifests=[]
 for p in paths:
  assert p.exists(),p
  manifests.append({'文件':str(p.relative_to(ROOT)),'字节数':p.stat().st_size,'SHA256':hashlib.sha256(p.read_bytes()).hexdigest()})
 (DEST/'本轮原始文件登记.json').write_text(json.dumps(manifests,ensure_ascii=False,indent=2))
 for rec in manifests:paragraph(rec['文件']+'；'+str(rec['字节数'])+'字节；SHA256：'+rec['SHA256'])
 heading('3 年度容量与负荷的原格定位')
 w=load_workbook(RAW/'近5年容载比.xlsx',read_only=True,data_only=True);ws=w.active;annual=[]
 for region,row in [('邳州',19),('市区',9),('丰县',34)]:
  for year,capcol,loadcol in zip(range(2021,2026),['C','F','I','L','O'],['D','G','J','M','P']):
   ca=ws[capcol+str(row)].value;lo=ws[loadcol+str(row)].value;annual.append([region,year,ws.title+'!'+capcol+str(row),ca,ws.title+'!'+loadcol+str(row),lo,float(ca)*10,float(lo)*10])
 table(['地区','年份','容量格','原容量','负荷格','原负荷','容量MVA','负荷MW'],annual)
 paragraph('原容量万千伏安、负荷万千瓦分别乘10。原负荷列名称为降压负荷，作为年网供最大负荷代理；未取得同刻形成记录时不改称已核定标准分母。原表比值存在显示舍入，报告均由换算容量与负荷重算。2025年邳州容量2139.5 MVA、市区3701 MVA。')
 paragraph('2021年共同反事实容量为邳州1463.5 MVA、市区2462 MVA。邳州控制目标2×731.77=1463.54 MVA，再按槽位规格离散形成；市区按1.7×1449.77=2464.609 MVA控制离散初态。二者不是原表实有容量，原表为2027与3113.5 MVA。各站起点与市区标定见基准站年表和city_2021_district_calibration.csv。')
 heading('4 月度光伏与前两年回推')
 w=load_workbook(RAW/'逐月分县分布式光伏.xlsx',read_only=True,data_only=True);ws=w.active;year=None;records=[]
 for i,rr in enumerate(ws.iter_rows(values_only=True),1):
  if str(rr[1]).startswith('202') and str(rr[1]).endswith('年'):year=int(str(rr[1])[:4])
  if rr[0] in ['QX-00005','QX-00007','QX-00001'] and year in [2023,2024,2025] and all(x is not None for x in rr[1:13]):
   name={'QX-00005':'邳州','QX-00007':'市区','QX-00001':'丰县'}[rr[0]]
   for m,value in enumerate(rr[1:13],1):records.append([name,year,m,f'{ws.title}!{chr(65+m)}{i}',float(value),float(value)*10])
 table(['地区','年份','12月原格','12月原值','采用MW'],[[r[0],r[1],r[3],r[4],r[5]] for r in records if r[2]==12])
 paragraph('原光伏表未写明数值单位。本轮按项目统一建模口径将原数值乘10解释为MW，此项属于明确的单位采用假定，尚待原统计提供方核定。报告中的建模光伏规模、源荷比及电量估计均在该口径下形成，不把换算后的全部数值称为原表已经注明单位的实测。')
 paragraph('2021—2022年无同来源月度记录。对每个月m，C(y,m)=C(2023,m)×[C(2023,m)/C(2024,m)]^(2023−y)。2023—2025年月度规模用原表换算。图展示完整建模序列并将前两年标为回推。历史压力采用各站2025年春秋工作日10—15时低净负荷模板月份的同月倍率，不只使用12月倍率。')
 metrics=json.loads((ASSETS/'指标复算.json').read_text());mp=metrics['model_pv'];modeled=[]
 for rid,vals in mp.items():
  for y,monthvals in vals.items():
   for m,vv in enumerate(monthvals,1):modeled.append([{'QX-00005':'邳州','QX-00007':'市区'}[rid],int(y),m,vv,'几何回推' if int(y)<2023 else '原表按统一单位换算'])
 wb=Workbook();s=wb.active;s.title='年度原格'
 s.append(['地区','年份','容量格','原容量','负荷格','原负荷','容量MVA','负荷MW'])
 for r in annual:s.append(r)
 s=wb.create_sheet('光伏原格');s.append(['地区','年份','月份','原格','原值','采用MW'])
 for r in records:s.append(r)
 s=wb.create_sheet('五年建模光伏');s.append(['地区','年份','月份','MW','性质'])
 for r in modeled:s.append(r)
 s=wb.create_sheet('基期设备')
 st=rows(BASE/'pizhou/rigid_stations.csv')+rows(BASE/'city/rigid_stations.csv');st=[r for r in st if r['year']=='2022'];s.append(['地区','站号','类别','初态1MVA','初态2MVA','第三台预留位置'])
 for r in st:s.append(['邳州' if r in rows(BASE/'pizhou/rigid_stations.csv') else '市区',r['station'],r['area_class'],float(r['prior_unit_1_mva']),float(r['prior_unit_2_mva']),int(r['source_available_third_slots'])])
 for s in wb:s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 wb.save(DEST/'原格定位与建模输入.xlsx')
 heading('5 站级负荷、方向与持续时间')
 paragraph('邳州主变负载率.xlsx逐列时序经电压与站号映射形成110 kV序列，市区徐州市区脱敏.zip内有功净负荷文件经已核定映射去重。市区源表值(kV)表头错误，第二列按MW读取。全零序列不作为正常轻载；补充SIM-CITY站无对应设备原行，其规格与C类为研究假定。')
 paragraph('邳州缺1 h，年度积分按前后均值补值；市区代表序列421 h按同周类型均值补值。缺测记录保留，补值不称为观测。全年峰值和持续时间统计按既定时序模板，连续事件以相邻1 h判断。反向峰为零时无事件。模型储能使用E/max(2.15,H95)，零与数值下限1 h均受铭牌2.15 h限制。')
 paragraph('正向：2025年独立站峰乘地区年度尺度；邳州尺度为区县原负荷/2025样本同期峰，市区先按原样本年度比例形成场景，再标定到区县年度负荷。反向：模板用户负荷为max(模板净功率,同月夜间中位数)，模板光伏为用户负荷−净功率，历史净功率为用户负荷×年度尺度−光伏×同月装机倍率，取负部分。2025年再与独立站反向极值比较取较大值。')
 paragraph('压力结果直接见基准rigid/elastic_stations.csv的forward_mw、reverse_mw、forward_duration_h、reverse_duration_h。图中的绝对峰为max(正向,反向)；站任务之和不替代区域同期峰。设备年最大负载率来自2025设备负载统计表主变1表P列，110 kV筛选后按区县统计；其不是模型同时点负载率。')
 heading('6 电量、源荷比和关联分析')
 paragraph('光伏8760小时数据为3408 MW系统光伏资料典型曲线，原始情景并非本区县2025年分布式电源逐站实测，日期按2025非闰年映射。光伏电量=逐时系数×相邻月末平均装机×1 h求和。用户电量=覆盖系数×带符号净交换年积分＋光伏电量，忽略损耗及未单独建模其他电源。')
 paragraph('邳州覆盖系数1.005033、市区1.169435；用户峰值采用日间净功率加回光伏与夜间候选负荷较大值。2025年源荷比估计邳州1.630653、市区0.317876；电量渗透率估计40.2778%与8.6940%。丰县最大负荷代理544.55 MW，借用两地年负荷率均值0.502489，形成2397002 MWh用户电量和30.0384%渗透率，只用于描述性对照。')
 paragraph('邳州12个月相关分析固定容量2139.5 MVA，不称为标准年度容载比。与电量渗透率的Pearson描述性系数分别为容量/月正向峰0.6556、绝对有功压力比−0.6579。按12个月样本使用，不作因果推断。年度丰县与邳州仅两点，不输出有统计解释力的跨地区相关系数。')
 heading('7 工程价格、计算系数及假定')
 table(['参数','本轮值','依据或计算'],[
 ['替换工程','16.573333万元/新购MVA','Sheet1!P21+P36除以10×(E21+E36)，(1151+1335)/150'],['第三台50 MVA','1223万元/项目','Sheet1第14、16行静态投资(1206+1240)/2'],['储能柜','21.0456262万元/柜','浏阳成交2104562.62元，10柜最低1 MW/2.15 MWh'],['储能单柜规模','0.1 MW/0.215 MWh','10柜保证包线性拆分，规模外推假定'],['储能上侧对照','27.2万元/柜','苏州100 kW/215 kWh子项最高限价，非中标价'],['新线','106.8万元/条','3×(0.9×24+0.1×120)+2×3'],['新线可转任务','7.491 MW','墩振线23—24杆区段负荷重建；热限代理9.05 MW'],['主变／线路比','约1.22364','(16.573333/0.95)/(106.8/7.491)'],['主变／储能比3 h','约0.05941','(16.573333/0.95)/(210.456262/(2.15/3))'],['功率因数','0.95','共同静态筛查假定'],['经济评价','2021年末—2041年、6%','研究评价期及实价折现假定'],['寿命与运维','主变/线20年、1%；储能10年、3%','统一研究参数，不冒称导则定额'],['更新规则','投运年＋寿命严格小于2041','同实价更新；运维投运次年起至2041'],['正常转接使用限制','10%','本轮研究假定，非事故转供标准']])
 paragraph('成本档案位于研究报告/数据来源/2026-09-28_10kV线路与成本依据。华容县2020规划表7-1提供24万元/km架空、120万元/km电缆电气、3万元/台智能开关；线路参考未自行增加无依据的土建、征地或间隔费用。徐州线路迁改116.17万元项目无对应长度，不用于元/km标定。')
 paragraph('储能案例原公告：苏州项目编号TCAXCG2024156-2，2025年3月25日；浏阳中标2024年8月15日。保证规模采用最低2.15 MWh，标题2.2 MWh不交替代入。原档案文件定位及哈希参考本目录一并复制的工程依据登记，旧登记中的方案控制条件以本文件覆盖。')
 shutil.copy2(ROOT/'研究报告/数据来源/2026-09-28_10kV线路与成本依据/source_manifest.csv',DEST/'工程与规范原始档案登记.csv')
 heading('8 规范、研究指标与结果文件')
 paragraph('规范状态于2026年10月1日通过全国标准信息公共服务平台核对：DL/T 5729—2023于2024年6月28日实施，DL/T 2041—2025于2026年6月18日实施。项目2025年4月修订稿仍按研究资料使用，不替代正式版条款。')
 paragraph('正式版状态依据：https://std.samr.gov.cn/hb/search/stdHBDetailedCNF?id=2E1288291AD60971E06397BE0A0ABFD2；https://std.samr.gov.cn/hb/search/stdHBDetailedCNF?id=52C44B7747821D39E06397BE0A0A8891。')
 paragraph('DL/T 5729—2023：2.0.4网供负荷、2.0.9及6.3.2容载比、6.3.3行政区县边界、6.3.4增长分档、6.3.5特定阶段调整、6.2单元供电安全、7.1.2故障或检修转移、8.2.1主变规格。装机及电量渗透率定义按项目2025年4月修订稿研究口径，不称为正式标准条款。H95、两个费用比和四年1∶2∶3∶4权重均为本项目计算方法。')
 paragraph('完整措施工作簿：研究报告/01研究推荐/2026-10-01条件矩阵/容载比条件矩阵与全部措施.xlsx。独立核查：研究报告/05_review/2026-10-01重构/矩阵独立复核.json。表列全部年度、站年及站对转接，目录中各情景保留原精度CSV和求解状态。正文只展示与技术论证相关的图表与公式，不加入程序路径或核验日志。')
 paragraph('工作簿的弹性指标、弹性措施两页对应正文18组条件下的弹性年度路径，年度值全部超过2.0。经济推荐、推荐措施两页对应实际费用较低的路径；平滑增长市区、增长加快两地区可能优选刚性。两类推荐各自使用对应容量、费用和措施，不拼接成新方案。')
 paragraph('现金流、单位换算和容量守恒由输出另行重算。72条路径全部通过，最低余量容许数值误差1×10^−4 MW，费用差小于0.002万元。最优性证据与静态容量证据分别使用，不宣称已经验证交流潮流、供电恢复动作或全年储能荷电状态。')
 for sec in DOC.sections:sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(3.1697);sec.right_margin=Cm(3.1697)
 DOC.core_properties.author='研究项目组';DOC.core_properties.title='研究报告数据来源与参数计算说明'
 DOC.save(DEST/'数据来源与参数计算说明.docx');(DEST/'数据来源与参数计算说明.md').write_text(''.join(MD).rstrip()+'\n')
 print('独立数据说明',DEST)
if __name__=='__main__':main()
