"""按区县统一源荷统计范围；只读原始资料和已冻结优化结果。"""
from pathlib import Path
import argparse,csv,json,hashlib,zipfile,math,unicodedata
from collections import defaultdict
import numpy as np
from scipy.stats import spearmanr
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment

ROOT=Path(__file__).resolve().parents[3]
RAW=ROOT/'实验/研究/data/tuomin/电网建模数据_Agent整合版_V1.2'
ARCHIVE=ROOT/'历史归档/区县统计调整前-2026-09-27.zip'
CASE_FILE=ROOT/'实验/研究/rebuild_2026/source_audit/capacity_release_simulation/reserve_policy_v4/recommendation_matrix/annual_indicator_recommendation.csv'
REGIONS={'QX-00001':'丰县','QX-00003':'贾汪','QX-00004':'沛县','QX-00005':'邳州','QX-00007':'市区','QX-00008':'睢宁','QX-00009':'铜山','QX-00010':'新沂'}
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def write(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
def interval(a,key):return [min(float(r[key]) for r in a),max(float(r[key]) for r in a)]
def show(v,d=3):
    if v[0]==v[1]:return f'{v[0]:.{d}f}'
    k=10**d;return f'{math.floor(v[0]*k)/k:.{d}f}～{math.ceil(v[1]*k)/k:.{d}f}'

def generate(out):
    # 旧研究底稿提供原路径特征；对照原CSV，不继承旧的边界判断。
    with zipfile.ZipFile(ARCHIVE) as z:
        data=json.loads(z.read('研究报告/01研究推荐/研究推荐数据.json'))
    w=openpyxl.load_workbook(RAW/'逐月分县分布式光伏.xlsx',read_only=True,data_only=True)
    s=w['Sheet1'];annual=[]
    for y,start in [(2023,4),(2024,16),(2025,29)]:
        for line in range(start,start+8):
            annual.append({'id':f'CS{len(annual)+1:03d}','region':REGIONS[s.cell(line,1).value],'year':y,
               'installed_10k_kw':float(s.cell(line,13).value),'installed_mw':float(s.cell(line,13).value)*10,
               'installed_location':f'Sheet1!M{line}',
               'installed_source':str((RAW/'逐月分县分布式光伏.xlsx').relative_to(ROOT))})
    w.close()
    w=openpyxl.load_workbook(RAW/'近5年容载比.xlsx',read_only=True,data_only=True)
    s=w['Sheet1'];county_rows={}
    for line in range(9,47):
        if s.cell(line,2).value=='110千伏':county_rows[REGIONS[s.cell(line,1).value]]=line
    for a in annual:
        line=county_rows[a['region']];col={2023:10,2024:13,2025:16}[a['year']]
        a['county_reference_peak_mw']=float(s.cell(line,col).value)*10
        a['county_source_load_ratio']=a['installed_mw']/a['county_reference_peak_mw']
        a['load_location']=f'Sheet1!{openpyxl.utils.get_column_letter(col)}{line}'
        a['load_source']=str((RAW/'近5年容载比.xlsx').relative_to(ROOT))
        a['ratio_formula']='区县年末已并网分布式光伏装机MW / 区县110kV年度降压峰值MW'
        a['ratio_basis']='区县统计近似；降压峰未等同于用户最大用电负荷；未把110与35kV上下级峰值相加'
        a['scope']='源表区县范围；市区按源表市区统计单元，不改写为29站'
    w.close();assert len(annual)==24
    county={(r['region'],r['year']):r for r in annual}
    with CASE_FILE.open(encoding='utf-8-sig') as f:original=list(csv.DictReader(f))
    bykey={(REGIONS[r['study_region_id']],int(r['voltage_kv']),int(r['year'])):r for r in original}
    cases=data['base_cases']
    for r in cases+data['price_cases']:
        a=county.get((r['region'],r['year']))
        r['county_source_load_ratio']=a['county_source_load_ratio'] if a else None
        r['county_source_load_basis']='区县装机/区县110kV降压峰统计近似' if a else '本表无2022年实际装机记录'
        r['county_source_load_source_locations']=[a['installed_location'],a['load_location']] if a else []
        r['standard_source_load_status']='区县范围已统一；研究源荷比按降压峰近似，未冒充用户毛负荷实测值'
        if r in cases:
            raw=bykey[r['region'],r['voltage_kv'],r['year']]
            assert abs(float(raw['recommended_clr'])-r['recommended_clr'])<1e-12
            assert abs(float(raw['recommended_capacity_mva'])-r['recommended_capacity_mva'])<1e-12
    complete=[r for r in cases if r['county_source_load_ratio'] is not None]
    assert len(complete)==9
    def group(r):
        if r['voltage_kv']==110 and r['transfer_scope']=='no_station_tie_candidates':return 'C1'
        if r['voltage_kv']==110:return 'C2' if r['county_source_load_ratio']<=1 else 'C3'
        return 'C4' if r['county_source_load_ratio']<=1 else 'C5'
    grouped=defaultdict(list)
    for r in complete:grouped[group(r)].append(r)
    conditions={
      'C1':'110 kV；尖峰模板5 h；无候选站间互济；区县统计源荷比≤1；恢复比例0.45',
      'C2':'110 kV；尖峰模板4 h；有互济候选；区县统计源荷比≤1；恢复比例0.60',
      'C3':'110 kV；尖峰模板4 h；有互济候选；区县统计源荷比>1；恢复比例0.60',
      'C4':'35 kV；尖峰模板3 h；无候选站间互济；区县统计源荷比≤1；恢复比例0.60',
      'C5':'35 kV；尖峰模板3 h；无候选站间互济；区县统计源荷比>1；恢复比例0.60'}
    matrix=[]
    for gid,items in sorted(grouped.items()):
        ranges={k:interval(items,k) for k in ['county_source_load_ratio','net_peak_cagr_since_2021',
            'forward_reference_peak_mw','reverse_station_peak_max_mw','forward_d95_max_template_hours',
            'reverse_d95_max_template_hours','recommended_clr','initial_capacity_mva']}
        matrix.append({'id':gid,'engineering_conditions':conditions[gid],'ranges':ranges,
           'relative_price_vectors':[[1,1]],'sample_count':len(items),'independent_path_count':1,
           'support_cases':[r['case_id'] for r in items],
           'research_reference_clr_interval':ranges['recommended_clr'] if len(items)>1 else None,
           'status':'区县统计口径的五类指标样本支持研究范围' if len(items)>1 else '单例，仅展示点值',
           'restriction':'源荷比采用降压峰近似；各范围是现有案例包络，不是任意组合均可行的矩形域'})
    names={'county_source_load_ratio':'区县源荷比（降压峰统计近似）',
       'net_peak_cagr_since_2021':'相应优化单元净峰平均增长率',
       'forward_reference_peak_mw':'相应优化单元正向参考净峰',
       'reverse_station_peak_max_mw':'最大站级反向净峰',
       'forward_d95_max_template_hours':'正向尖峰持续时间模板',
       'reverse_d95_max_template_hours':'反向尖峰持续时间模板'}
    correlations=[]
    scopes=[('2023至2025全部完整指标案例',complete)]
    scopes +=[(f'{region}{v}kV',[r for r in complete if r['region']==region and r['voltage_kv']==v])
              for region,v in sorted({(r['region'],r['voltage_kv']) for r in complete})]
    for scope,items in scopes:
        y=np.array([r['recommended_clr'] for r in items])
        for k,label in names.items():
            x=np.array([r[k] for r in items]);constant=np.ptp(x)==0 or np.ptp(y)==0
            correlations.append({'scope':scope,'metric':label,'n':len(items),
              'pearson':None if constant else float(np.corrcoef(x,y)[0,1]),
              'spearman':None if constant else float(spearmanr(x,y).statistic),
              'interpretation':'常量不可辨识' if constant else '仅描述；共用模板、同一路径；不是因果或独立地区验证'})
    data['status']='区县统计源荷口径及2023至2025五类指标研究矩阵已形成；用户峰值分母采用现有降压峰近似'
    data['county_scope_definition']={'unit':'区县，市区按源表统计范围','installed_unit':'万千瓦，用户确认',
       'source_load_definition':'同范围已并网分布式电源装机容量 / 同范围用户最大用电负荷',
       'calculation_used':'分子使用区县年末分布式光伏装机；分母用区县110kV降压峰作为现有数据近似',
       'aggregation':'不对110/35kV上下级净峰相加；不把区县装机分摊到电压层级或29站',
       'outcome_scope':'源荷比是区县背景条件；容载比、净峰增长及尖峰按原优化单元标记；市区R仅代表29站样本',
       'threshold':'研究分档取1作为装机与参考净峰相当的界限，不是导则规定阈值',
       'missing_years':[2021,2022]}
    data['annual_county_statistics']=annual;data['conditional_research_matrix']=matrix
    data['outcome_correlations']=correlations;data['complete_indicator_cases']=9
    data['excluded_from_full_indicator_matrix']=[r['case_id'] for r in cases if r['year']==2022]
    data['source_hashes'] +=[{'file':str(ARCHIVE.relative_to(ROOT)),'sha256':sha(ARCHIVE)},
                             {'file':str(CASE_FILE.relative_to(ROOT)),'sha256':sha(CASE_FILE)},
                             {'file':str(Path(__file__).relative_to(ROOT)),'sha256':sha(Path(__file__))}]
    data['source_hashes']=list({r['file']:r for r in data['source_hashes']}.values())
    assert all(sha(ROOT/r['file'])==r['sha256'] for r in data['source_hashes'])
    for r in data['historical_context']:
        a=county[r['region'],r['year']]
        r['county_source_load_ratio']=a['county_source_load_ratio']
        r['source_load_missing_reason']='区县背景口径已统一；研究分母采用110kV降压峰近似，不冒充用户最大毛负荷实测值'
    # 旧删点记录对应旧分类，留在旧版归档，不套用到新分类。
    data.pop('internal_deletion_sensitivity',None)
    data['annual_installed_candidates']=annual
    out.mkdir(parents=True,exist_ok=True);write(out/'研究推荐数据.json',data)
    wb=openpyxl.Workbook();wb.remove(wb.active)
    def sheet(title,head,rows):
        ws=wb.create_sheet(title);ws.append(head)
        for row in rows:ws.append(row)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24445C')
        for row in ws:
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        def width(v):return sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in str(v or ''))
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(65,max(15,max(width(c.value) for c in col)+2))
        for row in ws:
            ws.row_dimensions[row[0].row].height=max(35,max(math.ceil(width(c.value)/(ws.column_dimensions[c.column_letter].width-2)) for c in row)*16+10)
    sheet('阅读说明',['项目','口径'],[
      ['区县边界','邳州两电压等级共用邳州源荷背景指标；市区使用源表市区统计范围，29站是结果样本。'],
      ['源荷比分母','采用区县110kV年度降压峰作为用户最大负荷的现有数据近似，保留这个区别，不相加上下级峰值。'],
      ['年度与样本','2023至2025：24条区县年度统计、9条完整指标优化案例；2022年3条原结果单独保留。'],
      ['研究矩阵','1的源荷分档是研究界限；区间为原最优案例R包络，单例不提供区间。'],
      ['工程边界','新地区须保留初始资产、负荷分布和互济条件再优化；不是导则推荐值或接入承诺。']])
    sheet('区县年度源荷统计',['编号','区县','年度','装机万千瓦','装机MW','区县110kV降压峰MW','区县源荷比（降压峰近似）','装机位置','负荷位置','口径说明'],[
      [a[k] for k in ['id','region','year','installed_10k_kw','installed_mw','county_reference_peak_mw','county_source_load_ratio','installed_location','load_location','ratio_basis']] for a in annual])
    table=[]
    for m in matrix:
        r=m['ranges'];n=m['sample_count']
        table.append([m['id'],m['engineering_conditions'],show(r['county_source_load_ratio']),
          show([v*100 for v in r['net_peak_cagr_since_2021']],2)+'%',show(r['forward_reference_peak_mw']),
          show(r['reverse_station_peak_max_mw']),show(r['forward_d95_max_template_hours'],0),
          show(r['reverse_d95_max_template_hours'],0),'(1,1)',
          show(r['recommended_clr']) if n>1 else f"仅一点：{r['recommended_clr'][0]:.3f}",n,
          show(r['initial_capacity_mva'],1),'；'.join(m['support_cases']),m['status']])
    sheet('五指标研究推荐矩阵',['类别','工程条件','区县源荷比（降压峰近似）','净峰平均增长率','正向参考净峰MW','最大站反向净峰MW','正向D95 h','反向D95 h','归一相对成本(主变/储能,线路/储能)','R研究范围','年度案例数','初始容量MVA','支持案例','状态'],table)
    sheet('年度优化结果',['区县','电压kV','年度','区县源荷比（降压峰近似）','净峰平均增长率','正向参考净峰MW','最大站反向净峰MW','推荐R','主变容量MVA','储能柜数','既有转供MW','新线转供MW','刚性全路径现值万元','弹性全路径现值万元','源荷状态'],[
      [r['region'],r['voltage_kv'],r['year'],r['county_source_load_ratio'],r['net_peak_cagr_since_2021'],r['forward_reference_peak_mw'],r['reverse_station_peak_max_mw'],r['recommended_clr'],r['recommended_capacity_mva'],int(r['storage_modules_in_service']),float(r['existing_tie_mw']),float(r['new_line_transfer_mw']),float(r['rigid_path_npv_10k_cny']),float(r['elastic_path_npv_10k_cny']),r['county_source_load_basis']] for r in cases])
    sheet('指标与R关联',['统计范围','指标','案例数','Pearson','Spearman','解释边界'],[[c[k] for k in ['scope','metric','n','pearson','spearman','interpretation']] for c in correlations])
    price_names={'transformer_0_7':'主变价格系数0.7','storage_1_5':'储能价格系数1.5','line_1_5':'线路价格系数1.5','base':'基准价'}
    sheet('价格条件年度结果',['区县','年度','价格情景','区县源荷比（降压峰近似）','主变系数','储能系数','线路系数','推荐R','主变容量MVA','储能柜数','刚性全路径现值万元','弹性全路径现值万元'],[
      [r['region'],r['year'],price_names[r.get('price_case','base')],r['county_source_load_ratio'],r['transformer_cost_scale'],r['storage_cost_scale'],r['line_cost_scale'],r['recommended_clr'],r['recommended_capacity_mva'],int(r['storage_modules_in_service']),float(r['rigid_path_npv_10k_cny']),float(r['elastic_path_npv_10k_cny'])]
      for r in [x for x in cases if x['region']=='邳州' and x['voltage_kv']==110]+data['price_cases']])
    sheet('成本配对效应',['价格案例','年度','主变储能相对系数','线路储能相对系数','基准R','价格情景R','R变化','容量变化MVA','储能柜变化'],[[price_names[p['price_case']]]+[p[k] for k in ['year','transformer_storage_relative','line_storage_relative','base_clr','scenario_clr','delta_clr','delta_capacity_mva','delta_storage_modules']] for p in data['paired_price_effects']])
    sheet('原出力代理辅助量',['区县','电压kV','年度','原出力毛负荷代理比','使用范围'],[
      [r['region'],r['voltage_kv'],r['year'],r['source_load_output_proxy_ratio'],'仅解释原仿真构造，不替代装机口径源荷比'] for r in cases])
    sheet('来源指纹',['文件','SHA256'],[[r['file'],r['sha256']] for r in data['source_hashes']])
    wb.save(out/'研究型推荐矩阵.xlsx')
    lines=['# 区县口径容载比研究推荐方案','','研究统计边界统一为区县，市区按原表统计单元处理。邳州110 kV与35 kV共用邳州的年度装机背景；市区29站使用市区背景指标，其经济结果仍只代表29站。没有把全县光伏分配到电压层级或29站，也没有改变原设备与成本最优路径。','','## 源荷统计口径','','源荷比的定义采用同供电范围已并网分布式电源装机容量除以该范围用户最大用电负荷。现有历史负荷表记录的是降压峰，本研究以区县110 kV年度降压峰作为用户最大负荷的统计近似，明确标为“区县源荷比（降压峰近似）”。110 kV与35 kV上下级峰值不相加。这个近似解决现有数据的研究计算，不代表已经取得用户最大毛负荷的实测值。','','年末容量取12月记录，单位万千瓦由用户确认，折算MW乘10。邳州2023至2025年装机与参考降压峰之比分别为0.741827、1.305933、1.630653；市区为0.164490、0.250060、0.321280。2022年没有实际装机记录，原结果保留，但不参加完整五类指标矩阵。','','## 从年度案例归纳研究类型','','先按电压、尖峰持续时间和互济条件保留工程类别，再按区县统计源荷比分档。研究取1作为装机容量与参考降压峰相当的分界，它不是导则规定的高渗透阈值。各类同时显示净峰增长、净负荷极值、持续时间及措施相对成本的案例范围，推荐R取原完整成本最优路径的年度R包络。单个案例只列点值。','','| 类别与工程条件 | 区县源荷比（降压峰近似） | 净峰平均增长率 | 正向参考净峰/最大站反向峰（MW） | 正向/反向尖峰模板（h） | 相对成本（主变/储能，线路/储能） | R研究范围 | 案例数 |','| --- | --- | --- | --- | --- | --- | --- | ---: |']
    for m in matrix:
        r=m['ranges'];lines.append(f"| {m['id']}：{m['engineering_conditions']} | {show(r['county_source_load_ratio'])} | {show([v*100 for v in r['net_peak_cagr_since_2021']],2)}% | {show(r['forward_reference_peak_mw'])} / {show(r['reverse_station_peak_max_mw'])} | {show(r['forward_d95_max_template_hours'],0)} / {show(r['reverse_d95_max_template_hours'],0)} | (1,1) | {show(r['recommended_clr']) if m['sample_count']>1 else '仅一点：'+show(r['recommended_clr'])} | {m['sample_count']} |")
    lines+=['','区间按三位小数向外展示，准确值保留在JSON和年度结果中。上、下限来自不同年度的最优案例，并不表示单一工况下任意R都可实施；五个指标的独立最小值、最大值也不组成已经验证的矩形可行域。初始容量和支持案例在工作簿中列明。各类型目前仅有一条独立规划路径，不能把两个年度当作两个独立地区。','','## 关联解释','','相关分析只使用2023至2025年的九条完整指标案例，价格敏感性另做同物理条件配对。源荷比是区县背景指标，净峰增长、峰值及持续时间是相应优化单元的指标，不能将市区29站结果表述为全市区最优配置。','','| 优化单元 | 区县统计源荷比与R（Pearson / Spearman） |','| --- | --- |']
    for c in correlations:
        if c['metric']==names['county_source_load_ratio'] and c['n']==3:lines.append(f"| {c['scope']} | {c['pearson']:.3f} / {c['spearman']:.3f} |")
    lines+=['','每个单元只有三个年度，以上系数只描述已有案例。源荷比增加与R降低是否同时发生，还取决于继承容量、净峰变化和工程措施。源荷比与R都使用负荷统计，存在共同分母或共同趋势影响，不能据此宣称新能源装机对R的独立因果效应。持续时间在同一单元内不变，其独立作用尚不能辨识。','','## 成本维度和应用顺序','','邳州110 kV在2025年的区县统计源荷比为1.630653，基准R为2.191437。主变价格系数0.7或储能价格系数1.5时，已有原敏感性路径的R为2.209211；线路价格系数1.5时R仍等于基准。四个已求解离散价格向量的R包络为2.191～2.210，容量与储能替代结果保留原值。不能把离散价格向量之间的任意组合当作已求解条件。','','应用时先按区县计算装机与负荷指标，再按工程类别、年度净峰和价格条件查找样本范围，最后保留站级空间分布、初始资产和互济条件重新比选新地区方案。区县指标低于1也可能存在个别站反向压力，不能据此免除站级校核。当前模型没有完成交流潮流、电压、短路和保护校核，研究矩阵不作为接入承诺。','','本次只更新指标统计与推荐展示，不生成正式报告正文或Word，不改写原逐年措施、成本和R。']
    (out/'研究型推荐方案.md').write_text('\n'.join(lines)+'\n')
    audit={'status':'PASS','county_records':24,'full_indicator_base_cases':9,'year_2022_cases_retained_only':3,
       'matrix_types':5,'multi_case_types':3,'model_results_unchanged':True,
       'unit':'万千瓦（用户确认）','county_boundary':'用户确认按区县统计；源表市区统计范围保留',
       'denominator_assumption':'区县110kV降压峰为用户最大用电负荷的研究近似，不叠加35kV峰值',
       'standard_gross_load_ratio_verified':False,'invented_2022_installed_capacity':False}
    write(out/'补充分析核验.json',audit)
    return audit

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output',type=Path,required=True)
    args=p.parse_args();print(json.dumps(generate(args.output.resolve()),ensure_ascii=False))
