"""只读使用冻结结果；输出研究区间，不求解、不改写阶段0或原始数据。"""
from pathlib import Path
import csv, json, hashlib, math, argparse, unicodedata
from collections import defaultdict
import numpy as np
from scipy.stats import spearmanr
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parents[1]
BASE = ROOT/'实验/研究/rebuild_2026/source_audit/capacity_release_simulation/reserve_policy_v4'
RAW = ROOT/'实验/研究/data/tuomin/电网建模数据_Agent整合版_V1.2'
NAMES = {'QX-00001':'丰县','QX-00002':'鼓楼','QX-00003':'贾汪','QX-00004':'沛县','QX-00005':'邳州','QX-00006':'泉山','QX-00007':'市区','QX-00008':'睢宁','QX-00009':'铜山','QX-00010':'新沂','QX-00011':'徐州直属'}
LABELS = {
 'source_load_output_proxy_ratio':'光伏出力与毛负荷代理比（辅助指标）',
 'net_peak_cagr_since_2021':'年最大正向参考净负荷平均增长率',
 'forward_reference_peak_mw':'年最大正向参考净负荷',
 'reverse_station_peak_max_mw':'研究范围内最大站级反向净负荷极值',
 'forward_d95_max_template_hours':'正向尖峰最长连续持续时间模板',
 'reverse_d95_max_template_hours':'反向尖峰最长连续持续时间模板'}

def read(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def freeze_check():
    m=json.loads((ROOT/'研究报告/00审查/证据/阶段0冻结清单.json').read_text())
    assert all((ROOT/r['path']).is_file() and sha(ROOT/r['path'])==r['sha256'] for r in m['files'])
    assert not any((ROOT/p).exists() for p in m['removed_directories'])
    return len(m['files'])

def extent(rows,key):return [min(float(r[key]) for r in rows),max(float(r[key]) for r in rows)]

def display_interval(v,digits=3):
    if v[0]==v[1]:return f'{v[0]:.{digits}f}'
    k=10**digits
    return f'{math.floor(v[0]*k)/k:.{digits}f}～{math.ceil(v[1]*k)/k:.{digits}f}'

def group_id(r):
    # 分组仅依据工程条件和增长率；不使用待推荐的容载比。
    if int(r['voltage_kv'])==35:return 'M4' if r['net_peak_cagr_since_2021']>0.07 else 'M5'
    if r['transfer_scope']=='no_station_tie_candidates':return 'M1'
    return 'M3' if r['net_peak_cagr_since_2021']>0.07 else 'M2'

def main():
    original_count=freeze_check()
    current_manifest=json.loads((ROOT/'研究报告/00审查/证据/阶段0冻结清单.json').read_text())
    registered={r['path'] for r in current_manifest['files']}
    try:output_key=str((OUT/'研究推荐数据.json').relative_to(ROOT))
    except ValueError:output_key=None
    if output_key in registered:
        raise SystemExit('此输出目录已封存。请用 --output 指定新的复算目录，不覆盖现有结果。')
    OUT.mkdir(parents=True,exist_ok=True)
    inputs=[BASE/'recommendation_matrix/annual_indicator_recommendation.csv',
            BASE/'recommendation_matrix/cost_sensitivity/annual_recommendations.csv',
            RAW/'逐月分县分布式光伏.xlsx',RAW/'近5年容载比.xlsx',
            RAW/'光伏装机.xlsx',
            ROOT/'参考文献/前沿调研/pdfs/D8_政策标准/《配电网规划设计技术导则（修订稿）》.pdf',
            ROOT/'参考政策/分布式电源接入电力系统承载力评估导则（DLT 2041-2025).pdf',
            BASE/'recommendation_matrix/cost_sensitivity/summary.csv',
            BASE/'recommendation_matrix/cost_sensitivity/annual_schemes.csv']
    # 原值保留。用户确认万千瓦；用同区域2026年MW快照交叉校核，不回填原表。
    w=openpyxl.load_workbook(inputs[2],read_only=True,data_only=True)
    s=w['Sheet1']; annual=[]
    for year,start in [(2023,4),(2024,16),(2025,29)]:
        for row in range(start,start+8):
            region=s.cell(row,1).value
            annual.append({'region':NAMES[region],'year':year,'raw_value':s.cell(row,13).value,
                           'unit':'万千瓦（2026-09-27用户确认，内部量级校核支持；原表未注明）',
                           'working_capacity_mw':float(s.cell(row,13).value)*10,
                           'source':str(inputs[2].relative_to(ROOT)),
                           'location':f'Sheet1!M{row}','role':'县域年度装机统计候选；不是分电压或29站口径'})
    assert len(annual)==24 and len({(a['region'],a['year']) for a in annual})==24
    monthly_june_pz_mw=float(s.cell(44,7).value)*10
    w.close()
    w=openpyxl.load_workbook(inputs[4],read_only=True,data_only=True)
    assets=w['主变（跨区合并）1'];snapshot_values=[]
    for lineno,cells in enumerate(assets.iter_rows(min_row=2,values_only=True),2):
        if cells[1]=='QX-00005' and cells[3] in ('110kV','35kV'):
            snapshot_values.append({'row':lineno,'voltage':cells[3],'date':str(cells[4]),'online_mw':float(cells[7])})
    snapshot_pz_mw=sum(r['online_mw'] for r in snapshot_values)
    unit_check={'working_unit':'万千瓦','conversion_to_mw':10,
        'user_record':'2026-09-27用户确认“那就是对的，可以按照这个单位做统计和计算”；不冒称原表注明单位',
        'monthly_pz_2026_june_mw':monthly_june_pz_mw,'monthly_source_location':'Sheet1!G44',
        'asset_pz_2026_110_and_35_online_mw_sum':snapshot_pz_mw,
        'relative_difference':snapshot_pz_mw/monthly_june_pz_mw-1,
        'asset_rows':snapshot_values,
        'interpretation':'差异约0.215%，支持万千瓦的量级；快照时点不同，不据此认定逐站边界一致；220kV层不再叠加以免重复统计'}
    w.close()
    installed={(a['region'],a['year']):a for a in annual}
    w=openpyxl.load_workbook(inputs[3],read_only=True,data_only=True);s=w['Sheet1']
    historical=[];region=None
    for row in range(9,47):
        if s.cell(row,1).value is not None:region=NAMES[s.cell(row,1).value]
        voltage=s.cell(row,2).value
        if voltage not in ('110千伏','35千伏'):continue
        for year,col in [(2023,9),(2024,12),(2025,15)]:
            cap=float(s.cell(row,col).value);peak=float(s.cell(row,col+1).value)
            baseline=float(s.cell(row,4).value)
            source=installed[region,year]
            historical.append({'region':region,'year':year,'voltage_kv':110 if voltage=='110千伏' else 35,
                'installed_raw':source['raw_value'],'installed_unit':source['unit'],
                'installed_working_mw':source['working_capacity_mw'],
                'installed_over_voltage_net_peak_screen':source['raw_value']/peak,
                'capacity_10k_mva':cap,'stepdown_peak_10k_mw':peak,
                'reported_historical_clr':float(s.cell(row,col+2).value),'recomputed_historical_clr':cap/peak,
                'net_peak_cagr_since_2021':(peak/baseline)**(1/(year-2021))-1,
                'standard_source_load_ratio':None,
                'source_load_missing_reason':'已采用万千瓦工作单位；分电压供电范围与用户用电峰值口径未闭合，不以降压净峰替代',
                'installed_source':source['source'],'installed_location':source['location'],
                'load_source':str(inputs[3].relative_to(ROOT)),
                'load_location':f'Sheet1!{openpyxl.utils.get_column_letter(col)}{row}:{openpyxl.utils.get_column_letter(col+2)}{row}',
                'target_role':'历史配置水平；不是优化推荐值'})
    w.close();assert len(historical)==48
    rows=read(inputs[0])
    for lineno,r in enumerate(rows,2):
        for k in list(r):
            if k in LABELS or k in ('recommended_clr','recommended_capacity_mva','contingency_service_fraction','initial_capacity_mva'):
                r[k]=float(r[k])
        r['year']=int(r['year']);r['voltage_kv']=int(r['voltage_kv'])
        r['region']=NAMES[r['study_region_id']];r['source_line']=lineno
        r['case_id']=f"{r['region']}{r['voltage_kv']}kV-{r['year']}-基准价"
        r['standard_source_load_ratio']=None
        r['standard_source_load_status']='待同边界年度装机与用户峰值确认，代理比不替代'
        assert abs(r['recommended_capacity_mva']/r['forward_reference_peak_mw']-r['recommended_clr'])<1e-6
    assert len(rows)==12
    # 三条敏感性路径共12行，剔除重复的基准路径4行。
    prices=read(inputs[1]);price_rows=[];paired=[]
    bykey={(r['study_region_id'],r['voltage_kv'],r['year']):r for r in rows}
    for lineno,p in enumerate(prices,2):
        b=bykey[p['study_region_id'],int(p['voltage_kv']),int(p['year'])]
        if p['case_id']=='base':
            assert abs(float(p['recommended_clr'])-b['recommended_clr'])<1e-8
            continue
        c=dict(b);c.update({'price_case':p['case_id'],'price_source_line':lineno,
                          'recommended_clr':float(p['recommended_clr']),
                          'recommended_capacity_mva':float(p['recommended_capacity_mva']),
                          'storage_modules_in_service':int(p['storage_modules']),
                          'transformer_cost_scale':float(p['transformer_scale']),
                          'line_cost_scale':float(p['line_scale']),
                          'storage_cost_scale':float(p['storage_scale']),
                          'rigid_path_npv_10k_cny':float(p['rigid_path_npv_10k_cny']),
                          'elastic_path_npv_10k_cny':float(p['elastic_path_npv_10k_cny'])})
        # 年度场景现金流与转供量必须来自该价格方案的独立结果，不继承基准值。
        for k in ('year_investment_lifecycle_npv_10k_cny','existing_tie_mw','new_line_transfer_mw'):
            c.pop(k,None)
        c['case_id']=f"{c['region']}{c['voltage_kv']}kV-{c['year']}-{p['case_id']}"
        price_rows.append(c)
        paired.append({'price_case':p['case_id'],'year':c['year'],
                       'transformer_scale':c['transformer_cost_scale'],
                       'storage_scale':c['storage_cost_scale'],'line_scale':c['line_cost_scale'],
                       'transformer_storage_relative':c['transformer_cost_scale']/c['storage_cost_scale'],
                       'line_storage_relative':c['line_cost_scale']/c['storage_cost_scale'],
                       'base_clr':b['recommended_clr'],'scenario_clr':c['recommended_clr'],
                       'delta_clr':c['recommended_clr']-b['recommended_clr'],
                       'delta_capacity_mva':c['recommended_capacity_mva']-b['recommended_capacity_mva'],
                       'delta_storage_modules':int(c['storage_modules_in_service'])-int(b['storage_modules_in_service']),
                       'source_line':lineno})
    assert len(price_rows)==12
    corr=[]
    scopes=[('全部优化案例（混合工程条件，仅描述）',rows)]
    scopes += [(f'{r[0]}{r[1]}kV',[x for x in rows if (x['region'],x['voltage_kv'])==r]) for r in sorted({(x['region'],x['voltage_kv']) for x in rows})]
    for scope,subset in scopes:
        y=np.array([float(r['recommended_clr']) for r in subset])
        for key,label in LABELS.items():
            x=np.array([float(r[key]) for r in subset]);constant=np.ptp(x)==0 or np.ptp(y)==0
            corr.append({'scope':scope,'metric':label,'n':len(subset),
                         'pearson':None if constant else float(np.corrcoef(x,y)[0,1]),
                         'spearman':None if constant else float(spearmanr(x,y).statistic),
                         'status':'常量，不可辨识' if constant else '描述相关；共同模板和同一路径，不给显著性或因果结论'})
    historical_corr=[]
    for v in (110,35):
        subset=[r for r in historical if r['voltage_kv']==v]
        for key,label in [('installed_raw','县域年末光伏装机工作值（万千瓦）'),('net_peak_cagr_since_2021','降压负荷峰值平均增长率')]:
            x=np.array([r[key] for r in subset]);y=np.array([r['recomputed_historical_clr'] for r in subset])
            historical_corr.append({'voltage_kv':v,'metric':label,'n':len(subset),'county_count':8,
                 'pearson':float(np.corrcoef(x,y)[0,1]),'spearman':float(spearmanr(x,y).statistic),
                 'status':'历史配置相关，不是经济最优；光伏全县范围与电压层级不同；万千瓦工作单位有内部量级校核支持'})
    groups=defaultdict(list)
    for r in rows:groups[group_id(r)].append(r)
    conditions={
      'M1':'110 kV；正向尖峰模板5 h；无站间联络候选；净峰平均增长率≤4%；恢复比例0.45',
      'M2':'110 kV；正向尖峰模板4 h；有既有及新线互济候选；净峰平均增长率≤7%；恢复比例0.60',
      'M3':'110 kV；正向尖峰模板4 h；有既有及新线互济候选；净峰平均增长率>7%；恢复比例0.60',
      'M4':'35 kV；正向尖峰模板3 h；无站间联络候选；净峰平均增长率>7%；恢复比例0.60',
      'M5':'35 kV；正向尖峰模板3 h；无站间联络候选；净峰平均增长率≤7%；恢复比例0.60'}
    matrix=[];stability=[]
    for gid,subset in sorted(groups.items()):
        lim=extent(subset,'recommended_clr');n=len(subset)
        m={'id':gid,'engineering_conditions':conditions[gid],
           'standard_source_load_interval':None,'source_load_status':'待补；当前不能按源荷比分档',
           'growth_rate_interval':extent(subset,'net_peak_cagr_since_2021'),
           'forward_peak_mw_interval':extent(subset,'forward_reference_peak_mw'),
           'maximum_station_reverse_peak_mw_interval':extent(subset,'reverse_station_peak_max_mw'),
           'forward_d95_h_interval':extent(subset,'forward_d95_max_template_hours'),
           'reverse_d95_h_interval':extent(subset,'reverse_d95_max_template_hours'),
           'relative_cost_vectors':[[1.0,1.0]],'cost_vector_order':['主变/储能归一价格比','线路/储能归一价格比'],
           'optimal_case_clr_extent':lim,
           'research_reference_clr_interval':lim if n>=2 else None,
           'display_interval':display_interval(lim) if n>=2 else f'仅一点：{lim[0]:.3f}',
           'sample_count':n,'independent_path_count':1,
           'support_cases':[r['case_id'] for r in subset],
           'status':'待源荷比完善的样本支持研究范围' if n>=2 else '单例展示，不放行区间推荐',
           'scope':'同电压及工程条件下的经济最优案例包络；不是矩形可行域、置信区间或任意可实施R范围'}
        matrix.append(m)
        if n>=3:
            for held in subset:
                train=[r for r in subset if r['case_id']!=held['case_id']]
                remaining=extent(train,'recommended_clr')
                stability.append({'group':gid,'deleted_case':held['case_id'],'original_n':n,
                  'remaining_lower':remaining[0],'remaining_upper':remaining[1],
                  'deleted_r':held['recommended_clr'],
                  'covered':remaining[0]-1e-9<=held['recommended_clr']<=remaining[1]+1e-9,
                  'role':'内部删点敏感性；不是独立验证或未来预测检验'})
    # 对数恒等式单独显示共同分母作用，避免把R与负荷的关系误解释为因果。
    decomposition=[]
    for region,v in sorted({(r['region'],r['voltage_kv']) for r in rows}):
        series=sorted([r for r in rows if r['region']==region and r['voltage_kv']==v],key=lambda r:r['year'])
        for a,b in zip(series,series[1:]):
            ds=math.log(b['recommended_capacity_mva']/a['recommended_capacity_mva'])
            dp=math.log(b['forward_reference_peak_mw']/a['forward_reference_peak_mw'])
            dr=math.log(b['recommended_clr']/a['recommended_clr'])
            assert abs(dr-(ds-dp))<1e-8
            decomposition.append({'region':region,'voltage_kv':v,'from_year':a['year'],'to_year':b['year'],
                                  'delta_log_capacity':ds,'delta_log_peak':dp,'delta_log_ratio':dr})
    definitions=[
       ['源荷比','并网分布式电源装机容量／同范围用户最大用电负荷','无量纲或%','修订稿3.1.18、6.5.4；不是DL/T 2041-2025定义','万千瓦工作单位已校核；仍待分电压、同供电范围的装机容量及用户最大用电负荷'],
       ['年最大正向参考净负荷平均增长率','(本年净峰／2021年净峰)^(1/经过年数)-1','%/年','修订稿6.2.2.1表2给出增长率分档；项目明确采用CAGR','不是全年平均负荷增长率；不同终点的累计平均率不能当同一固定窗口'],
       ['净负荷极值','正向：年度参考净峰；反向：各站年度反向极值的最大值','MW','修订稿3.1.16、6.2.2、7.4.2.2；DL/T 2041-2025第6章设备负荷口径','反向最大站峰不是同步区域反向峰；代理、观测和年度参考缩放须注明'],
       ['尖峰持续时间','站级不低于各自方向峰值95%的最长连续时段，再取研究范围站间最大值','h','项目研究定义；两份导则均未规定本项目D95公式','历史年份沿用2025模板，不能称逐年实测；D95不是95分位数'],
       ['措施相对成本','(本情景主变单价／基准主变单价)／(本情景储能单价／基准储能单价)，线路同理','无量纲','项目经济模型与原价格敏感性输入；不是导则指标','不能把万元/MVA与万元/柜直接相除；不能用优化后总成本当预测输入'],
       ['规划参考容载比','优化后在役主变总额定容量／该年度正向参考净峰','MVA/MW，工程惯例按比值','修订稿3.1.16、6.2.2；DL/T 2041-2025未给推荐R区间','当前分母固定为规划参考值，不等同措施实施后重新测得的最大净负荷']]
    data={'status':'研究方法与补充分析完成；标准源荷轴待补，完整五维推荐未放行',
          'source_hashes':[{'file':str(p.relative_to(ROOT)),'sha256':sha(p)} for p in inputs],
          'sample_counts':{'optimized_base_rows':12,'price_additional_rows':12,'unique_physical_contexts':12,
                           'independent_base_paths':3,'annual_installed_records':24,'historical_voltage_rows':48},
          'annual_installed_candidates':annual,'installed_unit_check':unit_check,'historical_context':historical,
          'base_cases':rows,'price_cases':price_rows,'outcome_correlations':corr,
          'historical_correlations':historical_corr,'conditional_research_matrix':matrix,
          'paired_price_effects':paired,'internal_deletion_sensitivity':stability,
          'ratio_decomposition':decomposition}
    (OUT/'研究推荐数据.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
    wb=openpyxl.Workbook();wb.remove(wb.active)
    def sheet(name,headers,records):
        ws=wb.create_sheet(name);ws.append(headers)
        for record in records:ws.append(record)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24445C')
        for row in ws.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
        def text_width(value):
            return sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in str(value or ''))
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(65,max(15,max(text_width(c.value) for c in col)+2))
        for row in ws.iter_rows(min_row=2):
            lines=max(math.ceil(text_width(c.value)/max(1,ws.column_dimensions[c.column_letter].width-2)) for c in row)
            ws.row_dimensions[row[0].row].height=max(35,lines*16+10)
    sheet('阅读说明',['项目','内容'],[
      ['状态',data['status']],['研究范围','3个工程单元、12个物理年度条件。价格案例不扩充物理样本数。'],
      ['区间含义','现有成本最优路径的年度R包络，不是区间内每个R都可实施，也不是统计置信区间。'],
      ['源荷比','留空。2023至2025年年度装机按万千瓦工作口径；分电压及用户最大负荷待核对。'],
      ['五指标组织','保留原五类工程指标。净负荷、持续时间和价格分别保留方向或措施分量，不强行合成加权总分。'],
      ['边界','35 kV单独列示；市区29站只代表样本。新地区和新条件必须重求完整规划路径。']])
    sheet('指标定义',['指标','公式含义','单位','依据','口径限制'],definitions)
    sheet('研究型推荐矩阵',['类型','工程条件','标准源荷比','净峰增长率范围','正向参考净峰MW','最大站反向峰MW','正向D95 h','反向D95 h','归一相对成本(主变/储能,线路/储能)','研究参考R','案例数','独立路径数','支持案例','状态'],[
      [m['id'],m['engineering_conditions'],None,
       display_interval([v*100 for v in m['growth_rate_interval']],2)+'%',
       display_interval(m['forward_peak_mw_interval']),display_interval(m['maximum_station_reverse_peak_mw_interval']),
       display_interval(m['forward_d95_h_interval'],0),display_interval(m['reverse_d95_h_interval'],0),'(1,1)',
       m['display_interval'],m['sample_count'],1,'；'.join(m['support_cases']),m['status']] for m in matrix])
    sheet('年度优化案例',['单元','电压kV','年度','标准源荷比','出力毛负荷代理比(辅助)','净峰平均增长率','正向参考净峰MW','最大站反向峰MW','正向D95 h','反向D95 h','推荐R','主变容量MVA','储能在役柜数','既有联络转供MW','新线转供MW','刚性全路径成本现值万元','弹性全路径成本现值万元','来源行'],[
      [r['region'],r['voltage_kv'],r['year'],None,r['source_load_output_proxy_ratio'],r['net_peak_cagr_since_2021'],r['forward_reference_peak_mw'],r['reverse_station_peak_max_mw'],r['forward_d95_max_template_hours'],r['reverse_d95_max_template_hours'],r['recommended_clr'],r['recommended_capacity_mva'],int(r['storage_modules_in_service']),float(r['existing_tie_mw']),float(r['new_line_transfer_mw']),float(r['rigid_path_npv_10k_cny']),float(r['elastic_path_npv_10k_cny']),r['source_line']] for r in rows])
    sheet('年度装机候选',['区域','年度','原表12月值','单位状态','折算MW工作值','来源文件','位置','使用边界'],[[a[k] for k in ['region','year','raw_value','unit','working_capacity_mw','source','location','role']] for a in annual])
    sheet('装机单位校核',['项目','值或说明'],[
      ['采用单位',unit_check['working_unit']],['用户记录',unit_check['user_record']],
      ['邳州2026年6月按万千瓦折算MW',monthly_june_pz_mw],['邳州2026年110与35kV已并网容量合计MW',snapshot_pz_mw],
      ['相对差异',unit_check['relative_difference']],['解释边界',unit_check['interpretation']]])
    sheet('历史地区对照',['区域','电压kV','年度','县域年末光伏万kW工作值','容量万MVA','降压峰值万MW','源表历史R','容量除净峰重算R','净峰平均增长率','全县装机/本电压净峰对照比(非标准源荷比)','装机位置','容量负荷位置','用途'],[
      [r[k] for k in ['region','voltage_kv','year','installed_raw','capacity_10k_mva','stepdown_peak_10k_mw','reported_historical_clr','recomputed_historical_clr','net_peak_cagr_since_2021','installed_over_voltage_net_peak_screen','installed_location','load_location','target_role']] for r in historical])
    sheet('指标与推荐R关联',['范围','指标','案例数','Pearson','Spearman','解释限制'],[[r[k] for k in ['scope','metric','n','pearson','spearman','status']] for r in corr])
    sheet('历史配置关联',['电压kV','指标','记录数','地区数','Pearson','Spearman','解释限制'],[[r[k] for k in ['voltage_kv','metric','n','county_count','pearson','spearman','status']] for r in historical_corr])
    sheet('成本配对效应',['价格案例','年度','主变系数','储能系数','线路系数','主变储能相对系数','线路储能相对系数','基准R','敏感性R','R变化','容量变化MVA','储能柜变化','来源行'],[[p[k] for k in ['price_case','year','transformer_scale','storage_scale','line_scale','transformer_storage_relative','line_storage_relative','base_clr','scenario_clr','delta_clr','delta_capacity_mva','delta_storage_modules','source_line']] for p in paired])
    price_names={'base':'基准价','transformer_0_7':'主变价格系数0.7','storage_1_5':'储能价格系数1.5','line_1_5':'线路价格系数1.5'}
    scheme_names={'rigid':'刚性','elastic':'弹性'}
    price_summary=read(inputs[7]);schemes=read(inputs[8])
    sheet('价格全路径费用',['价格情景','方案','主变价格系数','储能价格系数','线路价格系数','全路径成本现值万元'],[
      [price_names[r['case_id']],scheme_names[r['scheme']],float(r['transformer_scale']),float(r['storage_scale']),float(r['line_scale']),float(r['cost_npv_10k_cny'])] for r in price_summary])
    sheet('价格方案年度配置',['价格情景','方案','年度','容量MVA','容载比','储能在役柜数'],[
      [price_names[r['case_id']],scheme_names[r['scheme']],int(r['year']),float(r['capacity_mva']),float(r['clr']),int(r['storage_modules'])] for r in schemes])
    sheet('内部删点敏感性',['类型','删除案例','原案例数','余下R下限','余下R上限','被删除案例R','是否覆盖','用途'],[[r[k] for k in ['group','deleted_case','original_n','remaining_lower','remaining_upper','deleted_r','covered','role']] for r in stability])
    sheet('R变化分解',['区域','电压kV','起年','末年','容量对数变化','负荷对数变化','R对数变化'],[[r[k] for k in ['region','voltage_kv','from_year','to_year','delta_log_capacity','delta_log_peak','delta_log_ratio']] for r in decomposition])
    sheet('来源指纹',['来源文件','SHA256'],[[r['file'],r['sha256']] for r in data['source_hashes']])
    wb.save(OUT/'研究型推荐矩阵.xlsx')
    # 生成一张便于直接审阅的表；保留准确值在JSON和年度页。
    md=['# 容载比研究区间补充分析','','本分析从现有成本最优路径归纳研究范围，原模型、设备方案及逐年结果保持不变。源荷比必须改用装机容量口径。当前标准源荷轴缺少同范围用户最大用电负荷和分电压年度容量，不能用出力代理值填补。','','## 区间构建方法','','先按电压等级、正向尖峰持续时间模板、互济候选和故障恢复比例分层，再用年最大正向净负荷平均增长率分档。增长率7%界限参考本地2025年4月规划导则修订稿表2，城市样本集中在4%以内。分层规则不使用待推荐的R。每层汇总经济最优案例的R最小值和最大值；价格变化单独做同条件配对，避免增加虚假的独立地区数量。','','区间上下限来自不同年度案例，反映条件变化对应的配置水平。它们不是同一工况下可以任意选择的连续容量区间。原始上下限保留全部精度，表中只向外保留三位小数。单个案例只显示点值，不给区间推荐。','','## 按工程条件归纳的研究范围','','下表保留五类指标。净负荷极值、持续时间和相对成本分别含两个分量，不能任意压缩为一个加权分数。所有源荷比格暂缺，其余四类指标已计算；本表尚未形成完整五维推荐。','','| 类型与工程条件 | 源荷比 | 净峰平均增长率 | 正向参考净峰 / 最大站反向峰（MW） | 正向/反向尖峰模板（h） | 归一相对成本（主变/储能，线路/储能） | 容载比研究范围 | 案例数 |','| --- | --- | --- | --- | --- | --- | --- | ---: |']
    for m in matrix:
        md.append(f"| {m['id']}：{m['engineering_conditions']} | 待补 | {display_interval([v*100 for v in m['growth_rate_interval']],2)}% | {display_interval(m['forward_peak_mw_interval'])} / {display_interval(m['maximum_station_reverse_peak_mw_interval'])} | {display_interval(m['forward_d95_h_interval'],0)} / {display_interval(m['reverse_d95_h_interval'],0)} | (1,1) | {m['display_interval']} | {m['sample_count']} |")
    md += ['','这些类型是由既有案例工程条件归纳的研究分类。M1由市区29站支持，M2、M3由邳州110 kV支持，M4、M5由邳州35 kV支持。各类型目前均只对应一条完整规划路径，不能据此宣称已经验证其他地区。源荷比补齐后，应在类型内再按源荷比分档；若每档案例不足，保留较粗分档，不人为补齐组合。','','## 指标与推荐容载比的关联','','本次新增指标与推荐R的Pearson和Spearman描述统计。价格敏感性不进入物理指标相关统计。源荷代理比只作为辅助构造量，不能用其相关系数说明标准源荷比的作用。','','| 工程单元 | 净峰增长率与R（Pearson / Spearman） | 正向净峰与R（Pearson / Spearman） | 反向最大站峰与R（Pearson / Spearman） |','| --- | --- | --- | --- |']
    for scope,subset in scopes[1:]:
        selected=[next(c for c in corr if c['scope']==scope and c['metric']==LABELS[k]) for k in ['net_peak_cagr_since_2021','forward_reference_peak_mw','reverse_station_peak_max_mw']]
        md.append('| '+scope+' | '+' | '.join(f"{c['pearson']:.3f} / {c['spearman']:.3f}" for c in selected)+' |')
    md += ['','每个单元只有四个年度，且共用时序模板。持续时间在单元内没有变化，因此无法从这些年份辨识持续时间的独立作用。正向峰值同时出现在R的分母中，其负相关包含数学定义作用，不能全部解释为规划规律。','','邳州35 kV在2022至2023年保留251 MVA容量，正向参考净峰下降约30.48%，R由2.013升至2.895。该变化直接说明，继承容量和年度负荷下降会提高R；不能仅凭较高R归因于新能源渗透水平。容量和负荷的对数变化分解见工作簿，逐行满足Δln R=Δln S−Δln P。','','## 成本条件对照','','邳州110 kV的2025年基准结果为R=2.191437、2096 MVA和35柜储能。主变投资系数降为0.7，或储能投资系数升为1.5时，已有敏感性结果均改为R=2.209211、2113 MVA和3柜储能，即R增加0.017774、容量增加17 MVA、储能减少32柜。2022至2024年的R在三条价格敏感性路径中均未改变。线路投资系数升为1.5时，2025年的R和设备容量仍等于基准结果。','','因此，现有成本证据支持相同负荷条件下的主变与储能替代解释。2025年该物理条件下，四种已求解价格方案的R包络为2.191～2.210。它仅涵盖四个离散价格向量，不代表其间任意价格组合均已求解，也不能用作全地区的成本分档阈值。','','## 年度装机数据补查','','已找到《逐月分县分布式光伏.xlsx》中2023至2025年的12月记录，共八个统计区域、24条年度原值。邳州为56.4702、112.6578、155.9638；市区为25.3679、36.2438、49.118。记录和坐标列入工作簿。用户已确认按万千瓦进行统计和计算，内部量级校核也支持这一判断：邳州2026年6月原值折算为1800.876 MW，2026年110与35 kV已并网容量表合计1797.0071 MW，相差约0.215%。两个时点不同，该校核支持单位量级，不能替代供电范围校核。已采用万千瓦作为本次研究工作口径，原表保持不变；标准源荷比还须对齐供电范围和用户用电峰值。2021、2022年未在本表找到实际记录，不以几何回推值当作装机实测。','','同时补充八个区域、两个电压等级、三个年度的48条历史配置对照。历史容载比是实际既有配置水平，不能直接用其最小值、最大值充当经济推荐区间。县域光伏统计值重复用于两个电压层级背景说明，关联分析分电压计算，避免把48行误当作48个独立装机样本。净负荷表的降压峰值也不能未经确认直接代替所供用户最大用电负荷。','','## 稳定性与适用边界','','对至少三个案例的类别做内部删点敏感性检查。删除某个边界案例后，余下包络未必覆盖该点，这表示当前范围依赖少量边界案例。该检查只评估已有样本范围的稳定性，不能称为独立地区验证或未来预测检验。每类检查结果保存在工作簿和JSON中。','','后续使用顺序为：核对供电范围与标准源荷比，匹配工程条件，查找样本支持范围，再保留站级负荷分布和联络条件重求完整规划路径。范围之外或缺少关键指标时，给出“证据不足、需补充优化”，不自动外推。当前模型没有完成交流潮流、电压、短路与保护校核，本文区间不作为工程接入承诺。','','本次补充分析不生成正式研究报告正文。按用户授权替换原版推荐矩阵及说明，并保留原设备、成本和逐年容载比结果；替换前后的指纹另存记录。']
    (OUT/'研究型推荐方案.md').write_text('\n'.join(md)+'\n')
    # 新输出基础校验：缺失项、支持区间、统计样本数及原文件指纹。
    checks=[]
    def check(condition,name):
        assert condition,name
        checks.append(name)
    check(len(matrix)==5,'五个工程分类全部生成')
    check(sum(m['sample_count'] for m in matrix)==12,'基础案例分组无重复无遗漏')
    check(all(m['standard_source_load_interval'] is None for m in matrix),'未以代理比填补标准源荷比')
    check(all(m['research_reference_clr_interval'] is None for m in matrix if m['sample_count']==1),'单例不生成推荐区间')
    check(len(rows)+len(price_rows)==24,'价格基准四行剔除，无重复计数')
    for m in matrix:
        for r in groups[m['id']]:check(m['optimal_case_clr_extent'][0]<=r['recommended_clr']<=m['optimal_case_clr_extent'][1],r['case_id']+'范围回溯')
    wb2=openpyxl.load_workbook(OUT/'研究型推荐矩阵.xlsx',read_only=True,data_only=True)
    check(wb2['年度装机候选'].max_row==25,'工作簿年度装机24条')
    check(wb2['年度优化案例'].max_row==13,'工作簿基础优化12条')
    check(wb2['历史地区对照'].max_row==49,'工作簿历史对照48条')
    check(all(wb2['年度优化案例'].cell(i,4).value is None for i in range(2,14)),'工作簿标准源荷比保留空值')
    wb2.close()
    check(all(sha(ROOT/r['file'])==r['sha256'] for r in data['source_hashes']),'所有分析输入读取前后哈希一致')
    check(freeze_check()==original_count,'阶段0冻结588文件保持不变')
    audit={'calculation_status':'PASS','complete_five_indicator_recommendation_status':'BLOCKED_PENDING_SOURCE_LOAD_ALIGNMENT',
           'checks':checks,'original_frozen_files_unchanged':original_count,
           'case_counts':data['sample_counts'],'source_load_assumptions_used':False,
           'note':'单位采用万千瓦工作口径；待同边界用户峰值确认，没有新增虚构物理样本，没有重求或改写原模型结果'}
    (OUT/'补充分析核验.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'status':'PASS','matrix_rows':len(matrix),'original_frozen_files':original_count,
                      'historical_rows':len(historical),'checks':len(checks),'outputs':str(OUT)},ensure_ascii=False))

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=OUT,help='可选：另建复算输出目录；封存目录禁止覆盖')
    args=parser.parse_args();OUT=args.output.resolve();main()
