"""图表交付复核：来源、公式、字体、内部编码、输出尺寸及配套完整性。"""
from pathlib import Path
import json,csv,re,math,hashlib,collections,zipfile
import numpy as np
import openpyxl
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
import pymupdf
ROOT=Path(__file__).resolve().parents[3];OUT=ROOT/'研究报告/02图表'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def load(p):return json.loads(p.read_text())
def save(p,o):p.write_text(json.dumps(o,ensure_ascii=False,indent=2,allow_nan=False))
CHECKS=[]
def check(name,test,detail=''):
 if isinstance(detail,set):detail=sorted(detail)
 CHECKS.append({'item':name,'passed':bool(test),'detail':detail});assert test,(name,detail)
def close(a,b):return math.isclose(a,b,rel_tol=1e-10,abs_tol=1e-10)
plans=[p for p in load(OUT/'source_data/图表规划结构.json') if p['执行状态']=='拟生成'];idx=load(OUT/'source_data/交付索引.json');records={r['id']:r for r in load(ROOT/'研究报告/00审查/data_source.json')['records']};used=set()
check('正式索引与规划逐项一致',{p['规划编号'] for p in plans}=={p['number'] for p in idx})
q=load(OUT/'source_data/文字与布局预检.json')
for item in q:
 check(item['number']+'实际绘制文字在画布内',not item['text_outside_canvas'],item['text_outside_canvas'])
 check(item['number']+'文字词表预检',not item['humanizer_lexical_flags'])
for p in idx:
 n=p['number'];tr=load(OUT/p['source_data']);check(n+'章节编号对应',int(n[1:n.index('-')])==p['chapter'])
 for f in p['files']:check(n+'文件存在 '+f,(OUT/f).is_file())
 for r in tr['records']:
  used.add(r['id']);check(n+'登记原值 '+r['id'],r==records[r['id']])
 svg=(OUT/next(f for f in p['files'] if f.endswith('.svg'))).read_text()
 check(n+'SVG无内部编码',not re.search(r'QX-\d+|BDZ-\d+|pizhou_|reserve_policy',svg))
 check(n+'SVG字体声明','Times New Roman' in svg and 'Microsoft YaHei' in svg)
 png=Image.open(OUT/next(f for f in p['files'] if f.endswith('.png')));dpi=png.info.get('dpi',(0,0));check(n+'PNG分辨率',dpi[0]>=599)
 width=png.width/dpi[0]*2.54;check(n+'固定插入宽度',abs(width-tr['width_cm'])<.012,(width,tr['width_cm']))
 pdf=pymupdf.open(OUT/next(f for f in p['files'] if f.endswith('.pdf')));fonts=pdf[0].get_fonts();names=[v[3] for v in fonts]
 check(n+'PDF中文字体已固定',any('YaHei' in s for s in names),names);check(n+'PDF西文字体已固定',any('TimesNewRoman' in s for s in names),names)
 txt=''.join(page.get_text() for page in pdf);check(n+'PDF无内部编码',not re.search(r'QX-\d+|BDZ-\d+',txt))
 if p['kind']=='table':
  wb=openpyxl.load_workbook(OUT/f'final_tables/{n}.xlsx',rich_text=True);ws=wb['正式表格'];fontset=set()
  for row in ws:
   for cell in row:
    v=cell.value
    if isinstance(v,list):
     for block in v:
      if hasattr(block,'font'):fontset.add(block.font.rFont)
    elif v is not None:fontset.add(cell.font.name)
  check(n+'Excel文字字体',fontset<={'Microsoft YaHei','Times New Roman'} and 'Times New Roman' in fontset,fontset)
  d=Document(OUT/f'final_tables/{n}_表格片段.docx');check(n+'可编辑表格片段',len(d.tables)==1)
  actual=sum(c.width.cm for c in d.tables[0].columns);check(n+'表格网格宽度',abs(actual-tr['width_cm'])<.02,actual)
  check(n+'表格未超母版净宽',actual<=d.sections[0].page_width.cm-d.sections[0].left_margin.cm-d.sections[0].right_margin.cm+.02)
  for row in d.tables[0].rows:
   for cell in row.cells:
    for par in cell.paragraphs:
     for run in par.runs:
      f=run._element.find('.//'+qn('w:rFonts'));check(n+'Word中英文字体',f is not None and f.get(qn('w:eastAsia'))=='微软雅黑' and f.get(qn('w:ascii')) in ['Times New Roman','Microsoft YaHei'])
 # 把每图选中的原始精度单独导出，显示数据CSV与原值CSV分开。
 if tr['records']:
  columns=['编号','数据名称','数值','单位','来源文件','来源位置','计算方式','SHA256']
  with (OUT/'source_data'/f'{n}_原始数值.csv').open('w',encoding='utf-8-sig',newline='') as f:
   w=csv.writer(f);w.writerow(columns)
   for r in tr['records']:w.writerow([r['id'],r['name'],r['value'],r['unit'],r['source'],r['source_location'],r['formula'],r['source_sha256']])
# 用原始Excel单元格独立核对新增区县数据及派生比值。
update=load(ROOT/'研究报告/01研究推荐/研究推荐数据.json');books={}
for r in update['annual_county_statistics']:
 vals=[]
 for source,loc in [(r['installed_source'],r['installed_location']),(r['load_source'],r['load_location'])]:
  if source not in books:books[source]=openpyxl.load_workbook(ROOT/source,data_only=True,read_only=True)
  sheet,cell=loc.split('!');vals.append(float(books[source][sheet][cell].value)*10)
 check(r['id']+'装机原表换算',close(vals[0],r['installed_mw']))
 check(r['id']+'降压峰原表换算',close(vals[1],r['county_reference_peak_mw']))
 check(r['id']+'区县比值计算',close(vals[0]/vals[1],r['county_source_load_ratio']))
# 年度推荐容量/参考净峰重新计算，避免程序字段别名掩盖比值口径。
base='实验/研究/rebuild_2026/source_audit/capacity_release_simulation/reserve_policy_v4/'
with (ROOT/base/'annual_matrix.csv').open(encoding='utf-8-sig') as f:annual=list(csv.DictReader(f))
for r in annual:
 for s in ['rigid','elastic']:
  calculated=float(r[s+'_capacity_mva'])/float(r['reference_peak_mw']);declared=float(r[s+'_clr']);check('年度比值 '+r['study_region_id']+r['voltage_kv']+r['year']+s,abs(calculated-declared)<1e-8)
# 费用分项使用同口径总量；分项舍入尾差保留说明。
for n in ['图4-3','表4-4']:
 tr=load(OUT/f'source_data/{n}.json');rr=tr['records']
 for rid in ['D017043','D017047','D017051','D017055','D017059','D017063']:
  i=int(rid[1:]);parts=sum(records[f'D{i-k:06d}']['value'] for k in [1,2,3]);check(n+'成本分项合计 '+rid,close(parts,records[rid]['value']))
check('无未冻结数值记录',used<=set(load(ROOT/'研究报告/00审查/data_source.json')['admitted_record_ids']))
# 只核验本阶段实际引用文件，不改变阶段0材料。
for r in load(OUT/'source_data/来源核验.json')['source_checks']:check('来源哈希保持 '+r['source'],sha(ROOT/r['source'])==r['sha256'])
frozen={r['path']:r['sha256'] for r in load(ROOT/'研究报告/00审查/证据/阶段0冻结清单.json')['files']}
for source in ['研究报告/00审查/最终数据字典.xlsx','研究报告/00审查/data_source.json']:
 check('阶段0字典保持 '+source,sha(ROOT/source)==frozen[source])
result={'status':'PASS','check_count':len(CHECKS),'formal_count':len(idx),'figure_count':sum(x['kind']=='figure' for x in idx),'table_count':sum(x['kind']=='table' for x in idx),'used_admitted_records':len(used),'numeric_accuracy':'原值、派生比值与原始Excel单元格复核通过','layout_scope':'SVG/PNG/PDF按物理尺寸复核；Word表格片段完成结构与字体检查，本机无Word排版引擎，未宣称已在Word中逐页渲染','checks':CHECKS}
save(OUT/'source_data/最终核验.json',result);print({k:v for k,v in result.items() if k!='checks'})
