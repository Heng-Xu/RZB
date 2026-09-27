#!/usr/bin/env python3
"""按开题章节生成报告图表；数值只读准入记录，完整结果不重求。
运行：/home/xh/anaconda3/envs/xuzhou110kv_clr/bin/python 研究报告/02图表/scripts/generate_stage1.py
"""
from pathlib import Path
import json,csv,re,hashlib,math,os,sys
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager as fm
from matplotlib.text import Text
from matplotlib.patches import Rectangle,FancyArrowPatch
from PIL import Image,ImageDraw
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
from openpyxl.cell.rich_text import CellRichText,TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.page import PageMargins
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[3];OUT=ROOT/'研究报告/02图表'
for name in ['final_figures','final_tables','source_data','preview']: (OUT/name).mkdir(parents=True,exist_ok=True)
BASE='实验/研究/rebuild_2026/source_audit/capacity_release_simulation/reserve_policy_v4/'
CN=Path(os.environ.get('STAGE1_CN_FONT','/home/xh/.local/share/fonts/stage1/msyh.ttf'))
EN=Path('/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf')
assert CN.is_file() and EN.is_file(),'缺少实际字体文件'
for p in [CN,EN]:fm.fontManager.addfont(str(p))
assert fm.FontProperties(fname=str(CN)).get_name()=='Microsoft YaHei','中文字体身份不符'
plt.rcParams.update({'font.family':['Times New Roman','Microsoft YaHei'],'font.size':9,'axes.labelsize':9,'axes.titlesize':9.5,'xtick.labelsize':9,'ytick.labelsize':9,'legend.fontsize':8.5,'axes.unicode_minus':True,'svg.fonttype':'none','pdf.fonttype':42,'savefig.dpi':600,'axes.edgecolor':'#525252','axes.linewidth':.65,'text.color':'#242424','axes.labelcolor':'#242424','xtick.color':'#242424','ytick.color':'#242424'})
BLUE='#315D81';ORANGE='#B57630';GRAY='#737373';LIGHT='#E8EEF3';DARK='#242424'
SCHEMES=['rigid','elastic'];SNAME={'rigid':'刚性方案','elastic':'弹性方案'}
COLORS={'rigid':BLUE,'elastic':ORANGE}
UNITS=[('QX-00005',110),('QX-00007',110),('QX-00005',35)]
UNAME={('QX-00005',110):'邳州110 kV',('QX-00007',110):'市区29站110 kV',('QX-00005',35):'邳州35 kV（支撑层）'}
SHORT={('QX-00005',110):'邳州110 kV',('QX-00007',110):'市区29站110 kV',('QX-00005',35):'邳州35 kV'}
PRICE=['base','transformer_0_7','storage_1_5','line_1_5']
PNAME={'base':'基准价格','transformer_0_7':'主变价格×0.7','storage_1_5':'储能价格×1.5','line_1_5':'线路价格×1.5'}
SOURCES=json.loads((OUT/'source_data/准入记录快照.json').read_text());RECORDS={r['id']:r for r in SOURCES['records']}
INDEX=json.loads((OUT/'source_data/准入CSV字段定位.json').read_text())
UPDATE=json.loads((OUT/'source_data/区县口径展示快照.json').read_text())['data']
PLANS=[x for x in json.loads((OUT/'source_data/图表规划结构.json').read_text()) if x['执行状态']=='拟生成']
LEDGER=set();SUPP=[];META=[];QA=[];CURRENT=None

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p,obj):p.write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False))
def record(rid):LEDGER.add(rid);return RECORDS[rid]['value']
def frame(file,fields):
 source=BASE+file;rows=list(csv.DictReader((ROOT/source).open(encoding='utf-8-sig')))
 for i,r in enumerate(rows):
  for f in fields:
   e=INDEX[source][str(i)][f];r[f]=record(e['record_id'])
 return pd.DataFrame(rows)
def annual():return frame('annual_matrix.csv',['reference_peak_mw','rigid_budget_fraction','contingency_fraction','rigid_cost_npv_10k_cny','elastic_cost_npv_10k_cny','rigid_capacity_mva','rigid_clr','rigid_storage_modules','elastic_capacity_mva','elastic_clr','elastic_storage_modules','rigid_existing_tie_mw','rigid_new_line_transfer_mw','elastic_existing_tie_mw','elastic_new_line_transfer_mw'])
def bcases():
 # 冻结展示更新只为年度指标补入已登记的区县源荷背景。
 d=frame('recommendation_matrix/annual_indicator_recommendation.csv',['station_count','net_peak_cagr_since_2021','forward_reference_peak_mw','reverse_station_peak_max_mw','forward_d95_max_template_hours','reverse_d95_max_template_hours','initial_capacity_mva','contingency_service_fraction','recommended_clr','recommended_capacity_mva','storage_modules_in_service'])
 for i,r in d.iterrows():
  item=next(v for v in UPDATE['base_cases'] if v['study_region_id']==r.study_region_id and v['voltage_kv']==int(r.voltage_kv) and v['year']==int(r.year))
  d.loc[i,'county_ratio']=item.get('county_source_load_ratio') if item.get('county_source_load_ratio') is not None else np.nan
  if int(r.year)>=2023:SUPP.append({'kind':'区县口径更新','region':item['region'],'year':int(r.year),'ratio':item['county_source_load_ratio'],'source_locations':item['county_source_load_source_locations']})
 return d
def county():
 SUPP.extend(UPDATE['annual_county_statistics']);return pd.DataFrame(UPDATE['annual_county_statistics'])
def prices():return frame('recommendation_matrix/cost_sensitivity/summary.csv',['transformer_scale','storage_scale','line_scale','cost_npv_10k_cny'])
def priceyears():return frame('recommendation_matrix/cost_sensitivity/annual_schemes.csv',['transformer_scale','storage_scale','line_scale','capacity_mva','clr','storage_modules','path_npv_10k_cny'])
def subset(df,u,year=None):
 a=df[(df.study_region_id==u[0])&(df.voltage_kv.astype(int)==u[1])].copy();a['year']=a.year.astype(int)
 return a[a.year==year] if year is not None else a.sort_values('year')
def datarow(df,u,y):return subset(df,u,y).iloc[0]
def costdata():
 rows=[]
 for i,u in enumerate([('QX-00005',35),('QX-00005',110),('QX-00007',110)]):
  for k,s in enumerate(SCHEMES):
   n=17040+i*8+k*4;vals=[record('D'+str(n+j).zfill(6)) for j in range(4)]
   rows.append({'单元':UNAME[u],'方案':SNAME[s],'首次投资现值':vals[0],'更新投资现值':vals[1],'固定运维现值':vals[2],'合计现值':vals[3]})
 d=pd.DataFrame(rows)
 return pd.concat([d[d['单元']==UNAME[u]] for u in UNITS],ignore_index=True)
def clean_axes(ax):
 ax.spines[['top','right']].set_visible(False);ax.grid(axis='y',color='#DDDDDD',linewidth=.5);ax.set_axisbelow(True)
def canvas(height=8.8,rows=1,cols=1,width=14.4):
 fig,axs=plt.subplots(rows,cols,figsize=(width/2.54,height/2.54),squeeze=False)
 # 固定物理尺寸，标题、题下注与图例留专用空间。
 fig.subplots_adjust(left=.13,right=.97,bottom=.19,top=.76,hspace=.80,wspace=.42)
 return fig,axs

def header(fig,note):
 fig.suptitle(CURRENT['规划编号']+' '+CURRENT['图表名称'],x=.5,y=.975,fontsize=11)
 if note:fig.text(.5,.902,note,ha='center',va='top',fontsize=8.5)
 source='资料来源：项目统计及规划计算。'
 if CURRENT['规划编号']=='表2-1':source='资料来源：国内外研究现状分析报告。'
 elif CURRENT['规划编号'] in ['图3-1','图3-4']:source='资料来源：区县光伏装机及降压负荷统计。'
 elif CURRENT['规划编号']=='表3-2':source='资料来源：近5年容载比统计表。'
 fig.text(.02,.022,source,ha='left',fontsize=8)

def grouped(ax,labels,vals1,vals2,ylabel,dec=0):
 x=np.arange(len(labels));w=.34
 for shift,v,s in [(-w/2,vals1,'rigid'),(w/2,vals2,'elastic')]:
  bar=ax.bar(x+shift,v,w,label=SNAME[s],color=COLORS[s],edgecolor=DARK,linewidth=.45,hatch='//' if s=='elastic' else None)
  for b,value in zip(bar,v):
   ax.annotate(f'{value:,.{dec}f}',(b.get_x()+b.get_width()/2,b.get_height()),xytext=(0,3),textcoords='offset points',ha='center',va='bottom',fontsize=8)
 ax.set_xticks(x,labels);ax.set_ylabel(ylabel);ax.set_ylim(0,max(max(vals1),max(vals2),1)*1.24);clean_axes(ax)

def finish(fig,data,note,conclusion,calc='原值读取；按展示单位舍入',width=14.4,kind='figure'):
 header(fig,note);fig.canvas.draw();rend=fig.canvas.get_renderer();W,H=fig.canvas.get_width_height()
 alltexts=[];clipped=[]
 unused_ticks=set()
 for ax in fig.axes:
  for axis in [ax.xaxis,ax.yaxis]:
   drawn=axis._update_ticks() if ax.axison else []
   used={id(t.label1) for t in drawn}|{id(t.label2) for t in drawn}
   unused_ticks.update(id(t.label1) for t in axis.majorTicks+axis.minorTicks if id(t.label1) not in used)
   unused_ticks.update(id(t.label2) for t in axis.majorTicks+axis.minorTicks if id(t.label2) not in used)
 for t in fig.findobj(Text):
  if id(t) in unused_ticks or not t.get_visible() or not t.get_text():continue
  tx=t.get_text();alltexts.append(tx)
  assert not re.search(r'QX-\d+|BDZ-\d+|pizhou_|reserve_policy|D\d{6}',tx),tx
  b=t.get_window_extent(rend)
  # 轴的隐藏偏移对象以及空文本不算图内文字。
  if b.x0 < -1 or b.y0 < -1 or b.x1>W+1 or b.y1>H+1:clipped.append(tx)
 dest=OUT/('final_figures' if kind=='figure' else 'final_tables');base=CURRENT['规划编号']
 files=[]
 for ext in ['svg','png','pdf']:
  p=dest/(base+'.'+ext);fig.savefig(p,dpi=600,facecolor='white',metadata={'Creator':'研究图表'} if ext in ['svg','pdf'] else None);files.append(str(p.relative_to(OUT)))
 # 全文文字清单用于humanizer逐项语境审查，不作为自动工具通过的替代。
 data=data.replace({np.nan:None}) if isinstance(data,pd.DataFrame) else data
 if isinstance(data,pd.DataFrame):
  data.to_csv(OUT/'source_data'/(base+'.csv'),index=False,encoding='utf-8-sig');rows=data.to_dict('records')
 else:rows=data
 trace={'number':base,'title':CURRENT['图表名称'],'chapter':CURRENT['所属章节名称'],'section':CURRENT['所属小节'],'data':rows,'records':[RECORDS[i] for i in sorted(LEDGER)],'supplemental_sources':SUPP,'document_sources':CURRENT['数据来源'],'processing':CURRENT['数据处理方式'],'calculation':calc,'note':note,'display_text':alltexts,'files':files,'width_cm':width,'height_cm':round(fig.get_figheight()*2.54,3),'scripts':['scripts/prepare_sources.py','scripts/revise_plan.py','scripts/generate_stage1.py']}
 dump(OUT/'source_data'/(base+'.json'),trace)
 qa={'number':base,'text_outside_canvas':clipped,'font_families':['Microsoft YaHei','Times New Roman'],'internal_code_check':'PASS','humanizer_lexical_flags':[tx for tx in alltexts if re.search('显著提升|明显优于|赋能|闭环|综上可知|具有重要意义|革命性|一键',tx)],'width_cm':width,'height_cm':trace['height_cm'],'dpi':600}
 QA.append(qa);META.append({'plan':CURRENT,'files':files,'note':note,'conclusion':conclusion,'calculation':calc,'source_data':'source_data/'+base+'.json','texts':alltexts,'qa':qa})
 plt.close(fig);print('已生成',base,CURRENT['图表名称'],flush=True)

FP=fm.FontProperties(family=['Times New Roman','Microsoft YaHei'],size=9)
def wrap(txt,maxpt,size=9):
 # 采用实际字体度量，避免中文和英文宽度估计失真。
 fig=plt.figure(figsize=(1,1));renderer=fig.canvas.get_renderer();prop=fm.FontProperties(family=['Times New Roman','Microsoft YaHei'],size=size)
 lines=[]
 for para in str(txt).split('\n'):
  line=''
  for c in para:
   w=renderer.get_text_width_height_descent(line+c,prop,False)[0]*72/fig.dpi
   if line and w>maxpt:lines.append(line);line=c
   else:line+=c
  lines.append(line)
 plt.close(fig);return '\n'.join(lines)
def rich(v):
 if not isinstance(v,str):return v
 blocks=[]
 for s in re.findall(r'[\x00-\x7f]+|[^\x00-\x7f]+',v):blocks.append(TextBlock(InlineFont(rFont='Times New Roman' if s.isascii() else 'Microsoft YaHei',sz=9),s))
 return CellRichText(blocks)
def docrun(p,text,size=9):
 for s in re.findall(r'[\x00-\x7f]+|[^\x00-\x7f]+',str(text)):
  r=p.add_run(s);r.font.name='Times New Roman' if s.isascii() else 'Microsoft YaHei';r.font.size=Pt(size);r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
def export_editable(df,widths,note,width):
 base=CURRENT['规划编号'];dest=OUT/'final_tables';wb=Workbook();ws=wb.active;ws.title='正式表格'
 ws.append([CURRENT['规划编号']+' '+CURRENT['图表名称']]);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(df.columns))
 for row in [list(df.columns)]+df.values.tolist():ws.append(row)
 for row in ws:
  for c in row:
   c.font=Font(name='Times New Roman' if isinstance(c.value,(int,float)) else 'Microsoft YaHei',size=9,color='242424')
   if isinstance(c.value,str):c.value=rich(c.value)
   c.alignment=Alignment(horizontal='center' if c.row<=2 else 'left',vertical='center',wrap_text=True)
   if c.row==2:c.border=Border(top=Side(style='thin'),bottom=Side(style='thin'));c.fill=PatternFill('solid',fgColor='E8EEF3')
   if c.row==ws.max_row:c.border=Border(bottom=Side(style='thin'))
 for i,w in enumerate(widths,1):ws.column_dimensions[chr(64+i)].width=max(10,w*75)
 ws.freeze_panes='A3';ws.auto_filter.ref=f'A2:{chr(64+len(df.columns))}{ws.max_row}'
 ws.page_setup.orientation='landscape' if width>15 else 'portrait';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.print_title_rows='1:2';ws.page_margins=PageMargins(left=.5,right=.5,top=.5,bottom=.5)
 # 完整精度另见逐表CSV；正式工作表保存报告显示值。
 ws2=wb.create_sheet('说明');ws2.append(['统计口径',note]);ws2.append(['完整源数据','../source_data/'+base+'.json']);ws2.append(['字体','中文微软雅黑，英文和数字Times New Roman'])
 wb.save(dest/(base+'.xlsx'))
 d=Document();sec=d.sections[0]
 if width>15:sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Cm(29.7);sec.page_height=Cm(21)
 else:sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.left_margin=sec.right_margin=Cm(3.17);sec.top_margin=sec.bottom_margin=Cm(2.54)
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;docrun(p,base+' '+CURRENT['图表名称'],11)
 tab=d.add_table(rows=1,cols=len(df.columns));tab.autofit=False
 for i,col in enumerate(tab.columns):col.width=Cm(width*widths[i])
 for i,c in enumerate(tab.rows[0].cells):c.width=Cm(width*widths[i]);docrun(c.paragraphs[0],df.columns[i]);c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
 for row in df.values.tolist():
  cells=tab.add_row().cells
  for i,(c,v) in enumerate(zip(cells,row)):c.width=Cm(width*widths[i]);docrun(c.paragraphs[0],v)
 for ridx,row in enumerate(tab.rows):
  trPr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trPr.append(ns)
  if ridx==0:trPr.append(OxmlElement('w:tblHeader'))
  for cell in row.cells:
   tcPr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single' if (ridx==0 and side in ['top','bottom']) or (ridx==len(tab.rows)-1 and side=='bottom') else 'nil');e.set(qn('w:sz'),'6');b.append(e)
   tcPr.append(b)
   for p in cell.paragraphs:p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.1
 p=d.add_paragraph();docrun(p,'注：'+note,8)
 d.save(dest/(base+'_表格片段.docx'))

def table(df,note,conclusion,widths=None,width=14.4):
 widths=widths or [1/len(df.columns)]*len(df.columns);widths=np.asarray(widths,dtype=float);widths/=widths.sum()
 pts=width/2.54*72;wrapped=[];heights=[]
 for row in [list(df.columns)]+df.values.tolist():
  r=[wrap(v,pts*w-9,9) for v,w in zip(row,widths)];wrapped.append(r);heights.append(max(t.count('\n')+1 for t in r)*12+8)
 height=(sum(heights)+85)/72*2.54
 assert height<=23.8,'表过高，请拆分'
 fig=plt.figure(figsize=(width/2.54,height/2.54));ax=fig.add_axes([0,0,1,1]);ax.set_xlim(0,pts);ax.set_ylim(0,sum(heights)+85);ax.axis('off')
 top=sum(heights)+45;left=0
 for i,(row,h) in enumerate(zip(wrapped,heights)):
  if i==0:ax.add_patch(Rectangle((2,top-h),pts-4,h,facecolor=LIGHT,edgecolor='none'))
  x=2
  for k,(tx,w) in enumerate(zip(row,widths)):
   ax.text(x+w*pts/2,top-h/2,tx,ha='center',va='center',fontsize=9,linespacing=1.30)
   x+=pts*w
  if i==0:ax.plot([2,pts-2],[top,top],color=DARK,lw=.85);ax.plot([2,pts-2],[top-h,top-h],color=DARK,lw=.65)
  elif i<len(wrapped)-1:ax.plot([2,pts-2],[top-h,top-h],color='#D3D3D3',lw=.3)
  top-=h
 ax.plot([2,pts-2],[top,top],color=DARK,lw=.85)
 # 表下注释自适应换行；标题与表体分开。
 footer=wrap('注：'+note,pts-12,8)
 footerlines=footer.count('\n')+1
 # 必要时增加底部空间，保持9 pt的表格字体而不压缩。
 if footerlines>2:
  extra=(footerlines-2)*11
  height+=(extra/72*2.54);fig.set_size_inches(width/2.54,height/2.54);ax.set_ylim(-extra,sum(heights)+85)
 ax.text(5,top-8,footer,ha='left',va='top',fontsize=8,linespacing=1.25)
 # finish默认题下注会挤表，表格注在表体下方，note只进说明。
 finish(fig,df,'',conclusion,'显示值舍入；完整精度和公式见逐表源数据',width,'table')
 META[-1]['note']=note
 p=OUT/'source_data'/(CURRENT['规划编号']+'.json');tr=json.loads(p.read_text());tr['note']=note;dump(p,tr)
 export_editable(df,widths,note,width)
 META[-1]['files'] += ['final_tables/'+CURRENT['规划编号']+'.xlsx','final_tables/'+CURRENT['规划编号']+'_表格片段.docx']

def flow(labels,subtitle,height=11.5):
 fig,axs=canvas(height);ax=axs[0,0];ax.set_position([.07,.10,.86,.71]);ax.axis('off');ax.set_xlim(0,1);ax.set_ylim(0,1)
 ys=np.linspace(.91,.08,len(labels));h=min(.105,.68/len(labels))
 for i,(tx,y) in enumerate(zip(labels,ys)):
  ax.add_patch(Rectangle((.08,y-h/2),.84,h,facecolor=LIGHT,edgecolor=BLUE,lw=.8));ax.text(.5,y,tx,ha='center',va='center',fontsize=9)
  if i<len(labels)-1:ax.add_patch(FancyArrowPatch((.5,y-h/2),(.5,ys[i+1]+h/2),arrowstyle='->',mutation_scale=10,color=GRAY,lw=.8))
 return fig

def price2025(df):return df[df.year.astype(int)==2025]
def draw(item):
 global CURRENT,LEDGER,SUPP
 CURRENT=item;LEDGER=set();SUPP=[];num=item['规划编号'];old=item['原规划编号']
 if num=='图1-1':
  labels=['区域年度统计、站级运行记录与设备资料','识别正向峰值、反向压力及尖峰持续时间','建立主变、储能与局部互济措施费用模型','比较刚性与弹性规划方案的完整年度路径','归纳片区指标与推荐容量组合','典型区域案例分析与工程应用建议']
  fig=flow(labels,'研究内容按开题报告各章展开');finish(fig,{'steps':labels},'运行特征、工程费用、片区建议与案例分析','形成运行特征分析、费用比较和片区建议的连续研究路线','基于开题研究任务与现行模型的定性整理');return
 if num=='表2-1':
  df=pd.DataFrame([
   ['承载力评价','国外综述2.1；国内综述3.2','接入容量、网络条件与评价边界','结合110 kV容量配置开展规划比选'],
   ['运行特征','国外综述2.2；国内综述3.3','净负荷、反向潮流与电压问题','区分区域正向净峰与站级反向峰'],
   ['灵活资源与网架','国外综述2.3；国内综述3.4','储能、网络增强及协同配置','比较主变、储能及已具备资料的局部互济'],
   ['费用与决策','国外综述2.4；国内综述3.5至3.6','投资、更新、运维与差异化投资','以统一评价期费用比较完整规划路径'],
   ['区域案例','国内综述3.7；综述第4章','研究尺度与案例应用','分析邳州及市区29站的配置差异']],columns=['研究主题','现有综述位置','工程关注点','本课题研究内容'])
  table(df,'按现有研究综述归纳主题。','研究主题对应运行分析、工程费用、片区建议及案例应用。',[.15,.23,.29,.33]);return
 if num=='表3-1':
  d=bcases();rows=[]
  for u in UNITS:
   r=datarow(d,u,2025);rows.append([UNAME[u],int(r.station_count),f'{r.initial_capacity_mva:.1f}',f'{r.forward_reference_peak_mw:.3f}','有候选互济' if u==UNITS[0] else '无站间互济候选'])
  table(pd.DataFrame(rows,columns=['研究单元','站数\n（座）','2021年容量\n（MVA）','2025年参考净峰\n（MW）','网络条件']),'市区为29站样本；邳州35 kV为支撑层，分层计算。','三单元分别采用自身基期资产及参考净峰。',[.26,.09,.19,.23,.23]);return
 if num=='表3-2':
  rows=[];rr=[r for r in RECORDS.values() if r['source'].endswith('近5年容载比.xlsx') and '110kV 2025年' in r['name']]
  regions=list(dict.fromkeys(r['name'].split()[0] for r in rr))
  for name in regions:
   a=[r for r in rr if r['name'].startswith(name+' ')];vals={r['variable']:record(r['id']) for r in a};S=vals['S原表'];P=vals['P原表'];rows.append([name,f'{S:.1f}',f'{P:.3f}',f"{vals['R原表']:.3f}",f'{S/P:.3f}'])
  table(pd.DataFrame(rows,columns=['区域','主变容量\n（MVA）','降压峰\n（MW）','原表容载比\n（MVA/MW）','容量/降压峰\n（MVA/MW）']),'按区域原表统计；原表比值保留原始舍入，末列另行计算。','区域容量和降压峰存在差异，原表舍入比值不覆盖为计算值。',[.14,.21,.22,.21,.22]);return
 if num=='图3-1':
  d=county();d=d[d.year==2025].sort_values('county_source_load_ratio');fig,axs=canvas(9.5);ax=axs[0,0];ax.set_position([.15,.14,.77,.62])
  vals=d.county_source_load_ratio.values*100;bars=ax.barh(d.region,vals,color=BLUE,edgecolor=DARK,lw=.45)
  for b,v in zip(bars,vals):ax.text(v+3,b.get_y()+b.get_height()/2,f'{v:.1f}%',va='center',fontsize=8.5)
  ax.set_xlim(0,max(vals)*1.24);ax.set_xlabel('区县源荷比（降压峰近似，%）');ax.spines[['top','right']].set_visible(False);ax.grid(axis='x',color='#ddd',lw=.5);ax.set_axisbelow(True)
  finish(fig,d[['region','year','installed_mw','county_reference_peak_mw','county_source_load_ratio']],'年末光伏装机 / 同区县110 kV年度降压峰','八区县源荷背景存在差异，结果代表区县统计近似。','装机万千瓦×10转MW；比值×100显示百分比');return
 if num=='图3-2':
  d=annual();fig,axs=canvas(12.5,rows=3);fig.subplots_adjust(top=.79,bottom=.13,hspace=.80,left=.19)
  rows=[]
  for ax,u in zip(axs[:,0],UNITS):
   a=subset(d,u);peak=a.reference_peak_mw.astype(float).tolist();ys=[2021]+a.year.tolist()
   if u[0]=='QX-00005':base=record('D000343' if u[1]==110 else 'D000338')
   else:
    raw=[r for r in RECORDS.values() if r['source'].endswith('近5年容载比.xlsx') and r['name'].startswith('市区 110kV ') and r['variable']=='P原表']
    p={int(r['name'].split()[2][:4]):record(r['id']) for r in raw};base=peak[-1]*p[2021]/p[2025]
   v=[base]+peak;ax.plot(ys,v,color=BLUE,marker='o',ms=3,lw=1.1);ax.set_title(UNAME[u],loc='left',fontsize=9.5);ax.set_xticks(ys);ax.set_ylabel('MW');clean_axes(ax);ax.set_ylim(0,max(v)*1.2)
   for y,t in zip(ys,v):rows.append({'单元':UNAME[u],'年':y,'参考净峰MW':t})
  finish(fig,pd.DataFrame(rows),'市区为29站参考峰值；历史年按年度统计比例缩放','各研究单元年度分母变化不同，35 kV数据单独作为支撑层分析。','邳州取原表；市区2021=样本2025峰×原表2021峰/原表2025峰');return
 if num=='图3-3':
  d=bcases();fig,axs=canvas(12.2,rows=3);fig.subplots_adjust(top=.79,bottom=.13,hspace=.85,left=.19)
  rows=[]
  for ax,u in zip(axs[:,0],UNITS):
   a=subset(d,u);a=a[a.year>=2023];x=a.county_ratio.astype(float)*100;y=a.recommended_clr.astype(float)
   ax.scatter(x,y,s=30,facecolor=BLUE,edgecolor=DARK,lw=.5)
   for xx,yy,yr in zip(x,y,a.year):ax.annotate(str(yr),(xx,yy),xytext=(5,4),textcoords='offset points',fontsize=8.5)
   ax.set_title(UNAME[u],loc='left');ax.set_ylabel('比值（MVA/MW）');ax.set_xlabel('区县源荷比（降压峰近似，%）');clean_axes(ax);ax.margins(x=.25,y=.4)
   for _,r in a.iterrows():rows.append({'单元':UNAME[u],'年':int(r.year),'区县源荷比':r.county_ratio,'推荐比值':r.recommended_clr})
  finish(fig,pd.DataFrame(rows),'2023至2025年；各单元3个年度案例','已有年度样本的关系随单元而异，不能据此确定通用线性响应。','只绘散点与年度标记，不作拟合或外推');return
 if num=='图3-4':
  d=county();regions=['丰县','贾汪','沛县','邳州','市区','睢宁','铜山','新沂'];fig,axs=canvas(9.7);ax=axs[0,0];x=np.arange(8);tones=[LIGHT,BLUE,ORANGE]
  for k,y in enumerate([2023,2024,2025]):
   a=d[d.year==y].set_index('region').loc[regions];ax.bar(x+(k-1)*.25,a.installed_mw,.25,label=str(y)+'年',color=tones[k],edgecolor=DARK,lw=.45,hatch='//' if k==2 else None)
  ax.set_xticks(x,regions);ax.set_ylabel('年末装机（MW）');ax.set_ylim(0,d.installed_mw.max()*1.12);clean_axes(ax);fig.legend(ncol=3,loc='upper center',bbox_to_anchor=(.55,.865),frameon=False)
  finish(fig,d[['region','year','installed_mw']],'各区县12月记录；单位统一为MW','年末光伏装机增长幅度在区县之间存在差异。','原表万千瓦×10转换MW');return
 if num=='图3-5':
  d=bcases();rows=[]
  for u in UNITS:
   r=datarow(d,u,2025);rows.append({'单元':UNAME[u],'参考正向净峰MW':r.forward_reference_peak_mw,'最大站反向峰MW':r.reverse_station_peak_max_mw})
  df=pd.DataFrame(rows);fig,axs=canvas(11.3,rows=2);fig.subplots_adjust(left=.16,bottom=.14,top=.77,hspace=.8)
  for ax,col,title in zip(axs[:,0],['参考正向净峰MW','最大站反向峰MW'],['区域参考正向净峰','单元内最大站反向峰']):
   vals=df[col];b=ax.bar(np.arange(3),vals,color=BLUE,edgecolor=DARK,lw=.45);ax.set_xticks(np.arange(3),[SHORT[u] for u in UNITS]);ax.set_ylabel('MW');ax.set_title(title,loc='left');ax.set_ylim(0,max(vals)*1.27);clean_axes(ax)
   for v,bb in zip(vals,b):ax.annotate(f'{v:.2f}',(bb.get_x()+bb.get_width()/2,v),xytext=(0,3),textcoords='offset points',ha='center',fontsize=8.5)
  finish(fig,df,'正向为区域参考净峰；反向为单站极值','正向需求和站级反向压力分别描述不同运行场景。','不将最大站反向峰当作区域同步反向峰');return
 if num=='表3-3':
  d=bcases();rows=[]
  for u in UNITS:
   r=datarow(d,u,2025);rows.append([UNAME[u],f'{r.forward_reference_peak_mw:.3f}',f'{r.reverse_station_peak_max_mw:.3f}',f'{r.forward_d95_max_template_hours:.0f}',f'{r.reverse_d95_max_template_hours:.0f}'])
  table(pd.DataFrame(rows,columns=['研究单元','正向参考净峰\n（MW）','最大站反向峰\n（MW）','正向尖峰模板\n（h）','反向尖峰模板\n（h）']),'尖峰时长按2025年站级模板，阈值为该方向峰值的95%。','尖峰持续时间与反向压力随研究单元不同。',[.26,.21,.21,.16,.16]);return
 if num=='图4-1':
  labels=['主变新购容量、储能柜数、候选新线长度','首次投资：按工程系数与采购包价计算','更新投资：按计算寿命和投运年度安排','固定运维：按新增投资及年费率计算','2022至2041年各批措施现金流','按6%折现至2021年，汇总增量费用现值']
  record('D017020');fig=flow(labels,'单位：万元；规划投运期2022至2025年');finish(fig,{'steps':labels},'规划投运期2022至2025年；费用评价期2022至2041年','费用比较同时包含首次投资、更新和固定运维。','依据成本模型F15a至F15e定性整理');return
 if num=='表4-1':
  specs=[('主变投资折算系数','D017038','35 kV，万元/新购MVA','工程案例折算'),('主变投资折算系数','D017039','110 kV，万元/新购MVA','工程案例折算'),('储能单柜价格','D017036','万元/柜','苏州招标最高限价'),('储能十柜价格','D017037','万元/10柜','浏阳等效中标价'),('单柜额定功率','D017031','MW/柜','设备规格'),('单柜额定能量','D017032','MWh/柜','设备规格'),('候选新线长度','D017026','km','规划假设'),('新线单位投资','D017027','万元/km','规划折算'),('年折现率','D017020','%','研究假设'),('主变/线路寿命','D017021','年','研究假设'),('储能寿命','D017023','年','研究假设'),('主变/新线年运维率','D017024','%','研究假设'),('储能年运维率','D017025','%','研究假设')]
  rows=[]
  for name,rid,unit,source in specs:
   v=record(rid);v=v*100 if unit=='%' else v;rows.append([name,f'{v:.6f}'.rstrip('0').rstrip('.'),unit,source])
  table(pd.DataFrame(rows,columns=['参数','数值','单位/范围','依据性质']),'主变按新购完整容量计价；储能报价锚点来自不同地区及年份。','各项参数区分工程案例、设备规格和研究假设。',[.29,.15,.27,.29]);return
 if num=='表4-2':
  d=annual();rows=[]
  for u in UNITS:
   r=datarow(d,u,2025);rows.append([UNAME[u],'2.0','3.2',f'{r.rigid_budget_fraction*100:.0f}%','不设此项预算',f'{r.contingency_fraction*100:.0f}%'])
  table(pd.DataFrame(rows,columns=['研究单元','刚性比值\n上限','弹性比值\n上限','刚性累计\n增配预算','弹性增配\n预算','停运恢复\n比例']),'两方案采用相同输入、候选措施及价格；上限、预算与恢复比例为研究设置。','方案差异来自参考比值上限与累计增配预算的组合。',[.27,.12,.12,.16,.18,.15]);return
 if num=='图4-2':
  d=annual();fig,axs=canvas(9.5);ax=axs[0,0];v=[];rows=[]
  for u in UNITS:
   r=datarow(d,u,2025);v.append((r.rigid_cost_npv_10k_cny,r.elastic_cost_npv_10k_cny));rows.append({'单元':UNAME[u],'刚性现值万元':v[-1][0],'弹性现值万元':v[-1][1],'降幅':1-v[-1][1]/v[-1][0]})
  grouped(ax,[SHORT[u].replace('29站','29站\n') for u in UNITS],[r[0] for r in v],[r[1] for r in v],'增量费用现值（万元）',2);fig.legend(ncol=2,loc='upper center',bbox_to_anchor=(.55,.865),frameon=False)
  finish(fig,pd.DataFrame(rows),'2022至2041年评价期；折现至2021年','在相同候选措施与价格下，三个研究单元的弹性路径费用均低于对应刚性路径。','(刚性现值−弹性现值)/刚性现值计算降幅');return
 if num=='图4-3':
  df=costdata();fig,axs=canvas(10.8);ax=axs[0,0];ax.set_position([.27,.17,.68,.59]);left=np.zeros(6);names=['首次投资现值','更新投资现值','固定运维现值']
  for k,col in enumerate(names):
   ax.barh(np.arange(6),df[col],left=left,label=col.replace('现值',''),color=[BLUE,LIGHT,ORANGE][k],edgecolor=DARK,lw=.4,hatch='//' if k==2 else None);left+=df[col].to_numpy()
  labels=[row.单元+'\n'+row.方案 for row in df.itertuples()];ax.set_yticks(np.arange(6),labels);ax.invert_yaxis();ax.set_xlabel('费用现值（万元）');ax.set_xlim(0,max(left)*1.22);ax.spines[['top','right']].set_visible(False)
  for i,v in enumerate(left):ax.text(v+30,i,f'{v:.2f}',va='center',fontsize=8)
  fig.legend(loc='upper center',bbox_to_anchor=(.55,.865),ncol=3,frameon=False)
  finish(fig,df,'2022至2041年评价期；折现至2021年','费用差由首次投资、更新投资和固定运维共同构成。','分项与合计统一取独立复算原值；与目标值存在约10⁻⁵万元尾差');return
 if num=='图4-4':
  d=annual();fig,axs=canvas(11.7,rows=2);fig.subplots_adjust(top=.77,bottom=.15,hspace=.8,left=.17);rows=[]
  for ax,field,title,ylabel,dec in zip(axs[:,0],['capacity_mva','storage_modules'],['主变总容量','储能在役数量'],['MVA','柜'],[1,0]):
   vv=[[float(datarow(d,u,2025)[s+'_'+field]) for u in UNITS] for s in SCHEMES];grouped(ax,[SHORT[u] for u in UNITS],*vv,ylabel,dec);ax.set_title(title,loc='left')
   if field=='capacity_mva':fig.legend(ncol=2,loc='upper center',bbox_to_anchor=(.55,.865),frameon=False)
   for i,u in enumerate(UNITS):rows.append({'单元':UNAME[u],'指标':field,'刚性':vv[0][i],'弹性':vv[1][i]})
  finish(fig,pd.DataFrame(rows),'2025年末配置；容量与柜数分开表示','两方案的容量与储能组合不同，形成相应费用差异。');return
 if num=='图4-5':
  d=prices();fig,axs=canvas(9.6);ax=axs[0,0];v=[[float(d[(d.case_id==p)&(d.scheme==s)].iloc[0].cost_npv_10k_cny) for p in PRICE] for s in SCHEMES];grouped(ax,[PNAME[p].replace('价格','价格\n') for p in PRICE],*v,'增量费用现值（万元）',2);fig.legend(ncol=2,loc='upper center',bbox_to_anchor=(.55,.865),frameon=False)
  df=pd.DataFrame([{'情景':PNAME[p],'方案':SNAME[s],'现值万元':v[k][i]} for k,s in enumerate(SCHEMES) for i,p in enumerate(PRICE)])
  finish(fig,df,'2022至2041年评价期；每次只改变一项价格系数','四个价格向量下均完成两方案整条规划路径比较。','已求解完整路径现值，不将原方案费用简单缩放代替重求结果');return
 if num=='图4-6':
  d=price2025(priceyears());d=d[d.scheme=='elastic'].set_index('case_id').loc[PRICE];fig,axs=canvas(12.5,rows=3);fig.subplots_adjust(top=.78,bottom=.14,hspace=.8,left=.18)
  for ax,col,label in zip(axs[:,0],['clr','capacity_mva','storage_modules'],['参考比值（MVA/MW）','主变容量（MVA）','储能数量（柜）']):
   vals=d[col].astype(float);ax.scatter(np.arange(4),vals,s=28,facecolor=BLUE,edgecolor=DARK,lw=.5);ax.set_xticks(np.arange(4),[PNAME[p].replace('价格','价格\n') for p in PRICE]);ax.set_ylabel(label);clean_axes(ax);ax.set_xlim(-.4,3.7)
   if col=='storage_modules':ax.set_ylim(-2,max(vals)*1.35+2)
   else:ax.set_ylim(min(vals)*.98,max(vals)*1.03)
   for xx,v in enumerate(vals):ax.annotate(f'{v:.3f}' if col=='clr' else f'{v:.0f}',(xx,v),xytext=(0,5),textcoords='offset points',ha='center',fontsize=8)
  out=d.reset_index()[['case_id','clr','capacity_mva','storage_modules']].copy();out.case_id=out.case_id.map(PNAME)
  finish(fig,out,'2025年弹性方案；纵轴采用局部范围展示离散配置差异','主变降价或储能涨价时，2025年方案增加17 MVA、减少32柜储能；线路涨价情景保持基准配置。','只比较四个独立价格向量，不绘连续响应曲线');return
 if num=='表4-3':
  d=price2025(priceyears());s=prices();rows=[]
  for p in PRICE:
   r=d[(d.case_id==p)&(d.scheme=='elastic')].iloc[0];r0=s[(s.case_id==p)&(s.scheme=='rigid')].iloc[0];r1=s[(s.case_id==p)&(s.scheme=='elastic')].iloc[0];rows.append([PNAME[p],f'{r.transformer_scale:g}/{r.storage_scale:g}/{r.line_scale:g}',f'{r0.cost_npv_10k_cny:.2f}',f'{r1.cost_npv_10k_cny:.2f}',f'{r.clr:.3f}',int(r.storage_modules)])
  table(pd.DataFrame(rows,columns=['价格情景','主变/储能/线路\n价格系数','刚性费用\n（万元）','弹性费用\n（万元）','2025推荐\n比值','2025储能\n（柜）']),'费用为完整规划路径现值；配置为2025年弹性方案。','价格条件影响费用及2025年主变与储能替代。',[.20,.23,.15,.15,.14,.13]);return
 if num=='表4-4':
  df=costdata();shown=df.copy()
  for col in df.columns[2:]:shown[col]=df[col].map(lambda v:f'{v:.2f}')
  shown.columns=['研究单元','方案','首次投资\n现值（万元）','更新投资\n现值（万元）','固定运维\n现值（万元）','合计\n现值（万元）']
  table(shown,'2022至2041年评价期，折现至2021年；合计按未舍入分项计算，显示值可能存在尾差。','三个单元的费用构成均可由相同科目比较。',[.25,.08,.18,.16,.16,.17]);return
 if num=='图5-1':
  labels=['区县源荷背景：年末光伏装机 / 年度降压峰','负荷增长：自2021年起参考净峰年均增长率','运行压力：区域正向参考净峰、最大站反向峰','尖峰持续时间：正向与反向95%峰值模板时长','经济条件：主变、储能与线路相对价格','推荐结果：参考容载比及对应工程措施组合']
  fig,axs=canvas(11.0);ax=axs[0,0];ax.set_position([.055,.13,.89,.64]);ax.set_xlim(0,1);ax.set_ylim(0,1);ax.axis('off')
  blocks=['区县源荷背景\n年末装机 / 110 kV降压峰','净峰增长\n自2021年起年均增长率','运行压力\n参考正向净峰 / 最大站反向峰','尖峰持续时间\n正向与反向95%峰值模板时长','经济条件\n主变、储能与线路相对价格']
  ys=np.linspace(.88,.12,5)
  for y,tx in zip(ys,blocks):
   ax.add_patch(Rectangle((.02,y-.065),.58,.13,fc=LIGHT,ec=BLUE,lw=.8));ax.text(.31,y,tx,ha='center',va='center',fontsize=9)
   ax.plot([.60,.67],[y,y],color=GRAY,lw=.8)
  ax.plot([.67,.67],[ys[-1],ys[0]],color=GRAY,lw=.8)
  ax.add_patch(FancyArrowPatch((.67,.50),(.73,.50),arrowstyle='->',mutation_scale=10,color=GRAY,lw=.8))
  ax.add_patch(Rectangle((.73,.31),.25,.38,fc=LIGHT,ec=BLUE,lw=.8));ax.text(.855,.50,'片区条件归纳\n\n推荐参考容载比\n工程措施组合',ha='center',va='center',fontsize=9)
  finish(fig,{'indicators':blocks,'output':'片区条件归纳、推荐参考容载比及工程措施组合'},'区县背景与优化单元指标分别计算','五类输入指标共同描述片区条件，推荐结果由完整规划路径比选形成。','定性整理已登记五类指标与规划输出');return
 if num=='图5-2':
  d=annual();fig,axs=canvas(12.6,rows=3);fig.subplots_adjust(top=.78,bottom=.13,hspace=.72,left=.18)
  rows=[]
  for ax,u in zip(axs[:,0],UNITS):
   a=subset(d,u)
   for s in SCHEMES:
    v=a[s+'_clr'].astype(float);ax.plot(a.year,v,label=SNAME[s],color=COLORS[s],marker='o' if s=='rigid' else 's',ms=3,lw=1.1,linestyle='-' if s=='rigid' else '--')
    for y,t in zip(a.year,v):rows.append({'单元':UNAME[u],'年':y,'方案':SNAME[s],'规划参考容载比':t})
   ax.axhline(2,color=GRAY,lw=.6,ls=':');ax.set_xticks(a.year);ax.set_ylim(0,3.4);ax.set_ylabel('比值（MVA/MW）');ax.set_title(UNAME[u],loc='left');clean_axes(ax)
  fig.legend(*axs[0,0].get_legend_handles_labels(),ncol=2,loc='upper center',bbox_to_anchor=(.55,.865),frameon=False)
  finish(fig,pd.DataFrame(rows),'2022至2025年；两方案采用相同年度参考净峰','年度推荐比值随容量和参考净峰变化，弹性上限3.2不等于推荐目标值。','容量/固定参考净峰；点线来自完整路径，2.0虚线为刚性研究上限');return
 if num=='表5-1':
  d=bcases();rows=[]
  for u in UNITS:
   for _,r in subset(d,u).iterrows():rows.append([UNAME[u],int(r.year),f'{r.forward_reference_peak_mw:.3f}',f'{r.recommended_capacity_mva:.1f}',f'{r.recommended_clr:.3f}',int(r.storage_modules_in_service)])
  table(pd.DataFrame(rows,columns=['研究单元','年','参考净峰\n（MW）','推荐容量\n（MVA）','推荐比值\n（MVA/MW）','储能\n（柜）']),'推荐结果来自各单元完整规划路径；2022年不补填缺少的装机指标。','年度推荐值与其容量及固定参考净峰配套引用。',[.27,.09,.20,.18,.17,.09]);return
 if num=='表5-2':
  SUPP.extend(UPDATE['conditional_research_matrix']);rows=[]
  def ran(v,fmt):return fmt(v[0]) if v[0]==v[1] else fmt(v[0])+'至'+fmt(v[1])
  def outward(v,places=3):
   a,b=v
   if a==b:return f'{a:.{places}f}（点值）'
   k=10**places;return f'{math.floor(a*k)/k:.{places}f}至{math.ceil(b*k)/k:.{places}f}'
  def bounds(v,places=3,scale=1,multiline=False):
   a,b=v[0]*scale,v[1]*scale
   if a==b:return f'{a:.{places}f}'
   k=10**places;sep='\n至' if multiline else '至'
   return f'{math.floor(a*k)/k:.{places}f}'+sep+f'{math.ceil(b*k)/k:.{places}f}'
  names=['市区29站\n110 kV，无互济','邳州110 kV\n有互济，源荷≤1','邳州110 kV\n有互济，源荷>1','邳州35 kV\n无互济，源荷≤1','邳州35 kV\n无互济，源荷>1']
  for i,r in enumerate(UPDATE['conditional_research_matrix']):
   v=r['ranges'];rows.append([names[i],outward(v['county_source_load_ratio']),bounds(v['net_peak_cagr_since_2021'],2,100),bounds(v['forward_reference_peak_mw'],3,1,True),bounds(v['reverse_station_peak_max_mw'],3),f"{v['forward_d95_max_template_hours'][0]:g}/{v['reverse_d95_max_template_hours'][0]:g}",'1 / 1',outward(v['recommended_clr']),f"{r['sample_count']}个年度\n1条路径"])
  df=pd.DataFrame(rows,columns=['工程类型','区县源荷比\n（降压峰近似）','净峰年均\n增长率\n（%）','正向参考净峰\n（MW）','最大站反向峰\n（MW）','正/反向\n时长（h）','相对价格\n主变/储能、\n线路/储能','推荐比值\n（MVA/MW）','支持案例'])
  table(df,'2023至2025年已有案例范围；两类仅有点值，范围端点向外舍入。主变/储能与线路/储能均为基准价格系数之比。','五类工程条件对应样本参考范围；每类只有一条独立规划路径。',[.15,.13,.09,.12,.12,.08,.10,.12,.09],23.0);return
 if num=='表5-3':
  df=pd.DataFrame([
   ['110 kV，具备互济候选','先确认供受端余量及通道，再比选主变与储能','邳州110 kV','年度参考净峰、站级反向压力、区段负荷'],
   ['110 kV，无站间互济候选','以存量容量、主变调整和储能配置共同比选','市区29站','样本设备、尖峰时长与完整费用路径'],
   ['35 kV支撑层','单独计算容量、负荷与费用，支撑110 kV规划分析','邳州35 kV','上下级峰值及容量避免重复统计'],
   ['区县源荷背景变化','更新装机与降压峰，再核对站级负荷分布','八区县统计','区县比值与局部反向压力分别判断'],
   ['价格条件变化','按实际工程价格重新比较完整规划路径','邳州价格情景','主变、储能和线路价格向量']],columns=['工程条件','规划应用建议','结果依据','重点工程资料'])
  table(df,'推荐范围用于已有样本条件对照；新地区应结合实际设备与运行资料重新比选。','应用顺序由区域背景、局部条件及工程费用共同决定。',[.22,.37,.15,.26]);return
 if num=='表6-1':
  d=bcases();rows=[]
  for u in UNITS[:2]:
   r=datarow(d,u,2025);rows.append([UNAME[u],int(r.station_count),f'{r.initial_capacity_mva:.1f}',f'{r.forward_d95_max_template_hours:g}/{r.reverse_d95_max_template_hours:g}',f'{r.contingency_service_fraction*100:.0f}%', '既有及拟建候选' if u==UNITS[0] else '无站间候选'])
  table(pd.DataFrame(rows,columns=['案例单元','站数\n（座）','基期容量\n（MVA）','正/反向尖峰\n模板（h）','停运恢复\n比例','互济条件']),'基期2021年；两案例按各自同边界数据比较。','两案例分别代表有候选互济与无站间候选的条件。',[.27,.09,.18,.18,.13,.15]);return
 if num in ['图6-1','图6-3']:
  d=annual();u=UNITS[0] if num=='图6-1' else UNITS[1];a=subset(d,u);fields=['capacity_mva','storage_modules']+(['transfer'] if num=='图6-1' else [])
  fig,axs=canvas(12.7 if len(fields)==3 else 11.5,rows=len(fields));fig.subplots_adjust(top=.78,bottom=.13,hspace=.8,left=.17);rows=[]
  for ax,f in zip(axs[:,0],fields):
   v=[]
   for s in SCHEMES:
    values=(a[s+'_existing_tie_mw']+a[s+'_new_line_transfer_mw']).astype(float).tolist() if f=='transfer' else a[s+'_'+f].astype(float).tolist();v.append(values)
   label={'capacity_mva':'主变容量（MVA）','storage_modules':'储能数量（柜）','transfer':'候选转供功率（MW）'}[f];grouped(ax,a.year.tolist(),*v,label,1 if f=='capacity_mva' else 2 if f=='transfer' else 0)
   if f=='capacity_mva':fig.legend(ncol=2,loc='upper center',bbox_to_anchor=(.55,.865),frameon=False)
   for k,s in enumerate(SCHEMES):
    for y,val in zip(a.year,v[k]):rows.append({'年':y,'方案':SNAME[s],'指标':label,'数值':val})
  finish(fig,pd.DataFrame(rows),'2022至2025年；'+('转供为场景功率，非年度电量' if num=='图6-1' else '市区仅代表29站样本'),'年度工程配置解释两方案容量、储能'+('和互济' if num=='图6-1' else '')+'组合的差异。','转供为既有及拟建候选功率之和；其他指标取冻结年度原值');return
 if num=='图6-2':
  fig,axs=canvas(9.8);ax=axs[0,0];ax.set_position([.06,.15,.88,.62]);ax.set_xlim(0,1);ax.set_ylim(0,1);ax.axis('off')
  boxes=[(.02,.64,.28,.22,'墩南线\n下游区段'),(.70,.64,.28,.22,'河炮线\n受入侧'),(.02,.18,.28,.22,'墩振线\n下游区段'),(.70,.18,.28,.22,'河湾变\n拟建新出线')]
  for x,y,w,h,tx in boxes:ax.add_patch(Rectangle((x,y),w,h,fc=LIGHT,ec=BLUE,lw=.8));ax.text(x+w/2,y+h/2,tx,ha='center',va='center',fontsize=9)
  for y,tx,ls,rid in [(.75,'既有联络候选','-','D017030'),(.29,'拟建互济候选','--','D017029')]:
   v=record(rid);ax.add_patch(FancyArrowPatch((.31,y),(.69,y),arrowstyle='->',mutation_scale=12,color=BLUE,lw=1,linestyle=ls));ax.text(.5,y+.12,tx,ha='center',fontsize=9);ax.text(.5,y-.10,f'静态通道包络 {v:g} MW',ha='center',fontsize=8.5)
  finish(fig,{'boxes':[b[-1] for b in boxes],'existing_path_envelope_mw':record('D017030'),'new_line_envelope_mw':record('D017029')},'工程关系示意；箭头表示候选受入方向','两类候选分别利用既有联络与拟建出线承担局部互济。','通道包络不是实际转供结果；转供还受区段、供端及受端条件共同限制');return
 if num in ['表6-2','表6-3']:
  d=annual();u=UNITS[0] if num=='表6-2' else UNITS[1];r=datarow(d,u,2025);spec=[('规划容量（MVA）','capacity_mva',1),('规划参考容载比（MVA/MW）','clr',3),('储能在役数量（柜）','storage_modules',0)]
  if num=='表6-2':spec += [('既有候选转供功率（MW）','existing_tie_mw',3),('拟建候选转供功率（MW）','new_line_transfer_mw',3)]
  spec += [('完整路径费用现值（万元）','cost_npv_10k_cny',2)]
  rows=[[name,f'{float(r["rigid_"+col]):.{dec}f}',f'{float(r["elastic_"+col]):.{dec}f}'] for name,col,dec in spec]
  saving=1-r.elastic_cost_npv_10k_cny/r.rigid_cost_npv_10k_cny
  rows.append(['费用降幅（%）','比较基准',f'{saving*100:.2f}'])
  table(pd.DataFrame(rows,columns=['指标','刚性方案','弹性方案']),'配置为2025年；费用为2022至2041年评价期现值，折现至2021年。','两方案在各自共同起点和候选条件下形成不同容量与费用组合。',[.56,.22,.22]);return
 if num=='表7-1':
  d=annual();save=[]
  for u in UNITS:
   r=datarow(d,u,2025);save.append(f'{SHORT[u]}{(1-r.elastic_cost_npv_10k_cny/r.rigid_cost_npv_10k_cny)*100:.2f}%')
  df=pd.DataFrame([
   ['运行特征','区县装机背景、区域正向峰与站级反向压力分别分析','图3-1至图3-5；表3-3'],
   ['费用比较','弹性路径费用降幅：'+ '；'.join(save),'图4-2；表4-4'],
   ['片区建议','按工程条件形成五类年度样本参考范围','图5-2；表5-2至表5-3'],
   ['价格响应','邳州110 kV主变降价或储能涨价情景的2025配置发生替代','图4-5至图4-6；表4-3'],
   ['典型案例','邳州与市区29站的措施组合结合互济条件分别分析','图6-1至图6-3；表6-2至表6-3']],columns=['研究内容','主要结果','图表依据'])
  table(df,'结论对应已分析区域与研究设置；市区结果仅代表29站样本。','运行特征、工程费用和局部网络条件共同支撑差异化规划建议。',[.16,.58,.26]);return
 raise ValueError('尚未实现 '+num)

def descriptions():
 lines=['# 图表说明文件','\n本文件为写作配套资料，保留完整来源与工程口径；不属于研究报告正文。图表已按开题章节编号，正式图内以数据与结果为主。']
 for m in META:
  p=m['plan'];n=p['规划编号'];trace=json.loads((OUT/m['source_data']).read_text());rids=[r['id'] for r in trace['records']]
  lines+=['\n## '+n+' '+p['图表名称'],'\n| 项目 | 内容 |\n| --- | --- |']
  fields={'图/表编号':n,'图/表名称':p['图表名称'],'所属章节':str(p['所属章节'])+' '+p['所属章节名称'],'所属小节':p['所属小节'],'图表类型':p['图表类型'],'图表目的':p['图表目的'],'数据来源':p['数据来源'],'来源位置与编号':'；'.join(rids)+'；精确文件、单元格/CSV行、SHA256见'+m['source_data'],'数据处理方式':p['数据处理方式'],'计算方法':m['calculation'],'统计/计算口径':m['note'] or trace.get('note') or '定性关系依据已审查模型与开题任务','绘制脚本位置':'scripts/generate_stage1.py，按编号分支；scripts/prepare_sources.py负责核验','图表所表达的核心结论':m['conclusion'],'正文引用建议':p['正文引用建议'],'工程解释说明':p['工程解释摘要'],'字体统一检查':'已完成：实际Microsoft YaHei与Times New Roman；英文数字优先Times New Roman，中文回退至微软雅黑；可编辑表格按文字段设置字体','humanizer检查':'已逐项进行工程语境审查；文字清单见逐图JSON；不将词表扫描等同于人工语境检查','格式审查':'已完成画布边界检查；最终版式检查记录见图表审查清单','正式文件':'；'.join(m['files'])}
  for k,v in fields.items():lines.append('| '+k+' | '+str(v).replace('|','/').replace('\n','；')+' |')
 lines+=['\n## 共同工程口径','\n费用为2022至2041年增量费用现值，折现至2021年，包含首次投资、更新投资及固定运维；网损、残值、处置、差异性存量运维、可靠性损失和套利等未评估，不写成真实零费用。两方案费用差同时反映参考比值上限和增配预算差异。','\n规划参考容载比=主变容量/固定年度参考净峰，不称措施后实测物理比值。区县源荷比采用年末光伏装机/区县110 kV年度降压峰的研究近似，2022年无装机数据不回填。九个完整指标年度案例来自三条独立规划路径；五类参考范围是样本包络，不是可任意组合的矩形可行域。','\n站级反向条件与单台停运恢复是静态容量代理，不替代完整导则承载力、交流潮流、电压、短路和保护校核。储能采用静态持续时间模型，不包含荷电状态、效率或衰减。市区共同观测8339小时、缺421小时，持续时间沿用已登记的补值模板；历史年度为模板及年度统计比例缩放。原表市区值(kV)表头按已确认有功净负荷MW解释。','\n## 移入配套说明的内容','\n对象与数据来源结构、数据质量、模型框架、决策变量、约束、求解及实施核查不另设正式重复图表。变量与公式详见阶段0《公式与术语审查》，数据来源关系详见《data_source.json》。模型依次最小费用、同费用下最大累计年度参考比值、最少转供；不逐年拼接方案。','\n## 文字审查依据','\n已按本地humanizer技能及用户指定的[humanizer规则](https://github.com/blader/humanizer)检查标题、坐标轴、图例、流程框、表头和注释。删除空泛优劣判断、内部过程词和重复限定；保留会改变指标含义的区域、时间和分母说明。参考睢宁魏集容量规划报告的数据引出、结果介绍和工程解释方式。']
 (OUT/'figure_description.md').write_text('\n'.join(lines))
 dump(OUT/'source_data/交付索引.json',[{'number':m['plan']['规划编号'],'title':m['plan']['图表名称'],'chapter':m['plan']['所属章节'],'kind':'figure' if m['plan']['规划编号'].startswith('图') else 'table','files':m['files'],'source_data':m['source_data'],'conclusion':m['conclusion']} for m in META])
 dump(OUT/'source_data/文字与布局预检.json',QA)
 dump(OUT/'source_data/字体登记.json',{'Chinese':{'family':'Microsoft YaHei','path':str(CN),'sha256':sha(CN),'source':'https://github.com/qbob/MSFonts/blob/master/fonts/msyh.ttf'},'English':{'family':'Times New Roman','path':str(EN),'sha256':sha(EN)},'fallback_policy':'Times New Roman优先处理英文数字，中文由实际Microsoft YaHei绘制','template_difference':'Word母版中文为仿宋_GB2312；本阶段按任务指定微软雅黑，阶段2正文应同步统一','renderer_versions':{'python':sys.version.split()[0],'matplotlib':matplotlib.__version__,'pandas':pd.__version__}})

def previews():
 for kind in ['figure','table']:
  files=[OUT/m['files'][1] for m in META if (m['plan']['规划编号'].startswith('图'))==(kind=='figure')]
  thumbs=[]
  for p in files:
   im=Image.open(p).convert('RGB');im.thumbnail((620,600));tile=Image.new('RGB',(650,630),'white');tile.paste(im,((650-im.width)//2,15));ImageDraw.Draw(tile).text((12,610),p.stem,fill='black');thumbs.append(tile)
  for start in range(0,len(thumbs),6):
   group=thumbs[start:start+6];grid=Image.new('RGB',(1950,1260),'#dddddd')
   for i,t in enumerate(group):grid.paste(t,((i%3)*650,(i//3)*630))
   grid.save(OUT/'preview'/f'{kind}_contact_{start//6+1}.png')
if __name__=='__main__':
 for item in PLANS:draw(item)
 descriptions();previews();print('TOTAL',len(META))
