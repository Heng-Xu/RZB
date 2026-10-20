"""Package the approved five-slide talk and audit its factual/native contents.

Run with the xuzhou110kv_clr Python environment (docx, Playwright, PyMuPDF).
This script renders existing slide sources; it never authors slide SVGs.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from pathlib import Path

import pymupdf
from playwright.sync_api import sync_playwright
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw


PROJECT = Path(__file__).resolve().parent
OUT = PROJECT.parent
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    exported = max((PROJECT / "exports").glob("*.pptx"), key=lambda x: x.stat().st_mtime)
    pptx = OUT / "EPG-NSGA-II_MIND2026_Oral_Corrected.pptx"
    shutil.copy2(exported, pptx)
    sections = re.findall(r"^# (\d\d_[^\n]+)\n(.*?)(?=^# |\Z)",
                          (PROJECT / "notes/total.md").read_text(), re.M | re.S)
    scripts = [(name, body.strip().removesuffix("---").strip()) for name, body in sections]
    assert len(scripts) == 5
    titles = ["Cover", "Background and problem", "Proposed solution", "Results", "Thank you"]
    counts = [len(re.findall(r"\b[\w]+(?:[-'][\w]+)*\b", body)) for _, body in scripts]

    md = OUT / "EPG-NSGA-II_English_Script.md"
    md.write_text("# EPG-NSGA-II — English oral presentation script\n\n"
                  f"Five slides · {sum(counts)} spoken words · approximately 4–5 minutes.\n\n"
                  "The spoken text below matches the PowerPoint speaker notes.\n\n" +
                  "\n\n---\n\n".join(f"## Slide {i}: {title}\n\n{body}"
                    for i, (title, (_, body)) in enumerate(zip(titles, scripts), 1)) + "\n")
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(8)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[style_name].font.name = "Arial"
        doc.styles[style_name].font.color.rgb = RGBColor.from_string("102B46")
    doc.add_heading("EPG-NSGA-II", 0)
    doc.add_paragraph("English oral presentation script")
    doc.add_paragraph(f"5 slides | {sum(counts)} words | approximately 4–5 minutes")
    doc.add_paragraph("The spoken text matches the PowerPoint speaker notes. "
                      "Pause briefly at each paragraph and between slides.")
    for i, (title, (_, body)) in enumerate(zip(titles, scripts), 1):
        heading = doc.add_heading(f"Slide {i} · {title}", 1)
        heading.paragraph_format.keep_with_next = True
        for paragraph in body.split("\n\n"):
            doc.add_paragraph(paragraph)
    script_docx = OUT / "EPG-NSGA-II_English_Script.docx"
    doc.save(script_docx)

    # Verify recovered source bytes and rederive the complete Table I ratios.
    manifest = json.loads((PROJECT / "sources/historical_model_manifest.json").read_text())
    for item in manifest:
        assert sha(PROJECT / item["local"]) == item["sha256"], item["local"]
    last = {}
    for row in csv.DictReader((PROJECT / "sources/historical_model/results/ea_converge.csv").open()):
        key = (int(row["K"]), row["algo"], int(row["seed"]))
        if key not in last or int(row["gen"]) > int(last[key]["gen"]):
            last[key] = row
    expected = {
        10: [1.000, .998, .996, 1.004, 1.007],
        20: [1.000, .991, .991, 1.006, 1.007],
        30: [1.000, .979, .988, 1.008, 1.009],
        40: [1.000, .985, .998, 1.048, 1.048],
    }
    algorithms = ["classic", "nsga3", "warmstart_only", "repair_only", "fg"]
    ratios = {}
    for k in expected:
        means = {}
        for algo in algorithms:
            vals = [r for key, r in last.items() if key[0] == k and key[1] == algo]
            assert len(vals) == 3
            assert all(int(r["gen"]) == (800 if algo in ["classic", "nsga3"] else 400) for r in vals)
            means[algo] = sum(float(r["hv"]) for r in vals) / len(vals)
        ratios[k] = {a: means[a] / means["classic"] for a in algorithms}
        assert [round(ratios[k][a], 3) for a in algorithms] == expected[k]
    costs = list(csv.DictReader((PROJECT / "sources/historical_model/results/tables/elastic_vs_rigid_comparison.csv").open()))
    elastic = next(r for r in costs if r["scheme"] == "elastic_cdslr")
    capped = next(r for r in costs if r["selection"] == "min_EENS")
    assert float(elastic["annualized_cost_MCNY_per_yr"]) == 228.43
    assert float(capped["annualized_cost_MCNY_per_yr"]) == 236.40
    assert float(elastic["EENS_MWh_per_yr"]) == 0
    assert float(capped["EENS_MWh_per_yr"]) == 561
    reduction = 100 * (236.40 - 228.43) / 236.40
    assert round(reduction, 1) == 3.4

    # Audit the final Office package, including actual embedded notes.
    normalize = lambda t: " ".join(t.split())
    with zipfile.ZipFile(pptx) as z:
        slide_names = sorted(n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
        assert len(slide_names) == 5
        table_count = picture_count = chevron_count = process_count = math_count = 0
        texts = []
        for i, name in enumerate(slide_names, 1):
            root = ET.fromstring(z.read(name))
            table_count += len(root.findall(".//a:tbl", NS))
            picture_count += len(root.findall(".//p:pic", NS))
            chevron_count += sum(x.get("prst") == "chevron" for x in root.findall(".//a:prstGeom", NS))
            process_count += sum(x.get("prst") == "flowChartProcess" for x in root.findall(".//a:prstGeom", NS))
            math_count += len(root.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
            texts.append(" ".join(t.text or "" for t in root.findall(".//a:t", NS)))
            note = ET.fromstring(z.read(f"ppt/notesSlides/notesSlide{i}.xml"))
            bodies = []
            for shape in note.findall(".//p:sp", NS):
                ph = shape.find(".//p:ph", NS)
                if ph is not None and ph.get("type") == "body":
                    bodies.append(" ".join(t.text or "" for t in shape.findall(".//a:t", NS)))
            assert normalize(" ".join(bodies)) == normalize(scripts[i-1][1]), f"notes mismatch on slide {i}"
        assert table_count == 0 and process_count == 4 and picture_count == 0 and math_count == 2
        for value in ["236.40", "228.43", "561.00", "0.00", "+4.8%", "−3.4%", "1.048"]:
            assert value in texts[3], value

    # Chromium honors SVG script positioning and embeds text fonts in its PDF.
    # CairoSVG previously collapsed script runs; it must not render this deck.
    # Math previews are vector paths; the PPTX itself retains native Office Math.
    previews = PROJECT / "validation/preview"
    previews.mkdir(parents=True, exist_ok=True)
    pdf_path = OUT / "EPG-NSGA-II_MIND2026_Oral_Corrected.pdf"
    svgs = sorted((PROJECT / "svg_output").glob("*.svg"))
    html = ('<!doctype html><html><head><meta charset="utf-8"><style>'
            '@page{size:1280px 720px;margin:0}'
            'html,body{margin:0;padding:0;width:1280px;background:white}'
            '.page{width:1280px;height:720px;break-after:page;overflow:hidden}'
            '.page:last-child{break-after:auto}'
            'svg{display:block;width:1280px;height:720px}'
            '*{-webkit-print-color-adjust:exact;print-color-adjust:exact}'
            '</style></head><body>' +
            ''.join('<div class="page">' + svg.read_text() + '</div>' for svg in svgs) +
            '</body></html>')
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1280, "height": 720})
        page.set_content(html, wait_until="load")
        page.evaluate("document.fonts.ready")
        page.pdf(path=str(pdf_path), width="1280px", height="720px",
                 print_background=True, prefer_css_page_size=True,
                 display_header_footer=False, tagged=True)
        browser.close()
    pdf = pymupdf.open(pdf_path)
    assert len(pdf) == 5, f"PDF page count: {len(pdf)}"
    rendered = []
    embedded_fonts = {}
    for svg, pg in zip(svgs, pdf):
        assert abs(pg.rect.width / pg.rect.height - 16/9) < 1e-5
        png = previews / f"{svg.stem}.png"
        pg.get_pixmap(matrix=pymupdf.Matrix(8/3, 8/3), alpha=False).save(png)
        rendered.append(png)
        for font in pg.get_fonts(full=True):
            xref, ext, kind, base, *_ = font
            assert xref > 0 and pdf.extract_font(xref)[3], f"Unembedded font: {base}"
            embedded_fonts[base] = {"type": kind, "extension": ext, "embedded": True}
        assert '\ufffd' not in pg.get_text(), "PDF contains replacement glyph"
    pdf.close()
    sheet = Image.new("RGB", (1300, 1170), "#E6EDF0")
    draw = ImageDraw.Draw(sheet)
    for i, path in enumerate(rendered):
        image = Image.open(path).convert("RGB")
        image.thumbnail((630, 355))
        x, y = 10 + (i % 2) * 650, 25 + (i // 2) * 390
        sheet.paste(image, (x, y))
        draw.text((x, y - 17), f"Slide {i+1}", fill="#102B46")
    sheet.save(previews / "contact_sheet.png")

    report = {
        "status": "passed", "slides": 5, "language": "English", "words_by_slide": counts,
        "total_spoken_words": sum(counts), "minutes_at_130_wpm": sum(counts) / 130,
        "source_manifest_hashes_verified": len(manifest), "normalized_final_HV": ratios,
        "cost_reduction_percent_unrounded": reduction,
        "pptx_native_tables": table_count, "pptx_native_chevrons": chevron_count,
        "pptx_native_process_shapes": process_count, "pptx_native_math": math_count,
        "pptx_picture_count": picture_count, "speaker_notes_match_script": True,
        "preview_rendering": "Chromium vector PDF; all text fonts embedded; vector math previews; not an Office render",
        "PDF_fonts": embedded_fonts,
        "provenance_qualification": "The camera-ready PDF governs the reported comparison. "
          "Archived CSV values match exactly; the historical rows are labelled selected(front=1) "
          "and min_EENS. The paper states the same knee-point rule, but that selection-rule "
          "history was not independently reconstructed. No optimization rerun was performed.",
        "files": {str(f.relative_to(OUT)): {"sha256": sha(f), "bytes": f.stat().st_size}
                  for f in [pptx, md, script_docx, pdf_path]},
    }
    (PROJECT / "validation/delivery_audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: report[k] for k in ["status", "slides", "total_spoken_words", "minutes_at_130_wpm", "pptx_native_tables", "speaker_notes_match_script"]}, indent=2))
    for f in [pptx, script_docx, md, pdf_path]:
        print(f)


if __name__ == "__main__":
    main()
