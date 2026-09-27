"""Generate the client-facing Q&A companion without changing the report."""
from pathlib import Path
import json,re
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'研究报告/甲方反馈'
DATA=OUT/'甲方问询Q&A_内容.json'
DOC=OUT/'甲方问询Q&A.docx'

def runfont(run,size=11,bold=False,east='方正仿宋_GBK',color='222222'):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)
    rf=run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'),east);rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman')

def para(d,text,style=None,size=11,after=5,keep=False,color='222222'):
    p=d.add_paragraph(style=style);p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.line_spacing=1.25;p.paragraph_format.keep_with_next=keep
    p.paragraph_format.widow_control=True
    runfont(p.add_run(text),size,color=color)
    return p

def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for values in rows:
        for c,v in zip(t.add_row().cells,values):c.text=str(v)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for c,w in zip(row.cells,widths):
            c.width=Cm(w)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                p.paragraph_format.line_spacing=1.1
                for r in p.runs:runfont(r,10,bold=i==0)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDF1F4');c._tc.get_or_add_tcPr().append(shade)
    return t

def main():
    data=json.loads(DATA.read_text());d=Document()
    d.core_properties.title=data['title'];d.core_properties.subject=data['subtitle']
    d.core_properties.author='项目研究组'
    sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2)
    sec.header_distance=sec.footer_distance=Cm(.9)
    for name in ['Normal','Heading 1','Heading 2','Title']:
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor.from_string('222222')
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'方正仿宋_GBK')
        st.paragraph_format.line_spacing=1.25
    for name,size in [('Heading 1',15),('Heading 2',12)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
        st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.keep_with_next=True
    hp=sec.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    runfont(hp.add_run('容载比弹性指标优化研究 · 甲方问询'),9,color='666666')
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    runfont(fp.add_run('第 '),9,color='666666')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
    runfont(fp.add_run(' 页'),9,color='666666')
    title=para(d,data['title'],size=21,after=12);title.runs[0].bold=True
    para(d,data['subtitle'],size=12,after=10)
    para(d,data['basis'],size=10,after=8)
    para(d,'每题先给出可直接回答的要点，再列追问补充及报告依据。报告依据采用章节、图表和公式编号，便于现场查阅。',size=10,after=12)
    para(d,'主要结果速查',style='Heading 2',size=12,keep=True)
    table(d,['研究单元','刚性费用／弹性费用（万元）','降幅','2025推荐比值'],[
        ['邳州110 kV','5245.66／4410.44','15.92%','2.191'],
        ['市区29站110 kV','3211.15／1912.29','40.45%','2.280'],
        ['邳州35 kV支撑层','3737.48／1517.74','59.39%','1.815']], [4.3,6.1,2.5,4.1])
    para(d,'费用为2022至2041年增量费用现值，折现至2021年；比值采用MVA/MW口径。各研究单元分别评价。',size=9,after=12,color='555555')
    para(d,'问答导航',style='Heading 2',size=12,keep=True)
    for s in data['sections']:
        first=s['items'][0]['number'];last=s['items'][-1]['number']
        para(d,f"Q{first:02d}—Q{last:02d}  {s['title']}",size=11,after=5)
    para(d,'配套资料：七章研究报告；独立《数据来源说明》；《当前报告数据来源与计算审查》。',size=9,after=0,color='555555')
    for s in data['sections']:
        d.add_page_break();para(d,s['title'],style='Heading 1',size=15,keep=True)
        for x in s['items']:
            p=para(d,f"Q{x['number']:02d}  {x['question']}",style='Heading 2',size=12,keep=True)
            p.runs[0].bold=True
            p=para(d,x['answer'],size=11,after=5,keep=True)
            para(d,'追问时可补充：'+x['followup'],size=10,after=5,keep=True,color='444444')
            para(d,'报告依据：'+x['source'],size=9,after=12,color='666666')
    d.save(DOC)
    print(DOC)

if __name__=='__main__':main()
