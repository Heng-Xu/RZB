"""只读既有两套原始库，比较相同分档下的五参数矩阵；不寻优或修改报告。"""
from pathlib import Path
from collections import defaultdict
import gzip
import hashlib
import itertools
import json
import math

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
NAME = '五参数推广矩阵新旧对比_20261014'
PUBLIC = ROOT/'研究报告/01研究推荐/当前版'/NAME
SOURCE = ROOT/'研究报告/数据来源/当前版'/NAME
REVIEW = ROOT/'研究报告/05_review/当前版/1014五参数推广矩阵新旧对比'
OLD = ROOT/'研究报告/01研究推荐/当前版/参数联合情景_20261012'
NEW = ROOT/'研究报告/01研究推荐/当前版/补充指标全背景扩展_20261014'
OLD_RAW = ROOT/'研究报告/试算/20261012_参数联合情景锁定/统一原始求解记录.jsonl'
NEW_RAW = ROOT/'研究报告/试算/20261014_补充指标全背景扩展/原始求解记录.jsonl.gz'
FIVE = ['分布式电源装机渗透率', '年最大净负荷平均增长率', '电缆长度占比', '初始转供率', '负荷分散系数']
KEYS = ['装机档', '增长档', '电缆档', '转供档', '负荷分散系数']
H = '储能支撑持续时间_h'
ENERGY = '分布式电源电量渗透率'
PRICE = ['线路价格倍率', '储能价格倍率']
CUTS = [[1, 2], [.02, .04, .07], [.1, .3], [.3, .5, .7]]
LABELS = [['≤1', '(1,2]', '>2'], ['≤2%', '(2%,4%]', '(4%,7%]', '>7%'],
          ['≤10%', '(10%,30%]', '>30%'], ['≤30%', '(30%,50%]', '(50%,70%]', '>70%']]
KD = [1., 1.1, 1.15]
DOMAINS = [range(3), range(4), range(3), range(4), KD]
MAPPING = dict(zip(['eta', 'growth_rate', 'cable_fraction', 'initial_transfer_fraction', 'diversity'], FIVE)) | {
    'support_duration_h': H, 'clr': '容载比', 'capacity_mva': '主变容量_MVA', 'cost': '费用现值_万元'}
NEW_MAPPING = MAPPING | {'energy_penetration': ENERGY, 'line_price': PRICE[0], 'storage_price': PRICE[1]}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n', encoding='utf-8')


def interval(lo, hi):
    return f'{math.floor((lo+1e-10)*100)/100:.2f}～{math.ceil((hi-1e-10)*100)/100:.2f}'


def normalize(frame):
    out = frame.copy()
    cols = FIVE+[H]+[c for c in [ENERGY]+PRICE if c in out]
    out[cols] = out[cols].round(10)
    return out


def read_and_check_raw(raw, public, mapping):
    opener = gzip.open if raw.suffix == '.gz' else open
    rows = []
    states = defaultdict(int)
    with opener(raw, 'rt', encoding='utf-8') as stream:
        for line in stream:
            record = json.loads(line)
            states[record['status']] += 1
            if record['verified']:
                rows.append({label: record[field] for field, label in mapping.items()})
    actual = normalize(pd.DataFrame(rows))
    expected = normalize(pd.read_csv(public)[list(mapping.values())])
    sorting = FIVE+[H]+[c for c in [ENERGY]+PRICE if c in actual]
    for frame in [actual, expected]:
        assert not frame.duplicated(sorting).any()
    pd.testing.assert_frame_equal(actual.sort_values(sorting).reset_index(drop=True),
        expected.sort_values(sorting).reset_index(drop=True), check_dtype=False,
        check_exact=False, rtol=1e-12, atol=1e-9)
    return normalize(pd.read_csv(public)), dict(states)


def cell_key(row):
    return tuple([sum(float(row[f]) > cut+1e-9 for cut in cuts)
                  for f, cuts in zip(FIVE[:4], CUTS)]+[float(row[FIVE[4]])])


def endpoint(row):
    fields = FIVE+[H]+[c for c in [ENERGY]+PRICE if c in row]
    return json.dumps({f: float(row[f]) for f in fields}, ensure_ascii=False, sort_keys=True)


def matrix(frame):
    selected = frame[frame[FIVE[-1]].isin(KD)].copy()
    for i in range(4):
        selected[KEYS[i]] = selected[FIVE[i]].map(lambda v, i=i: sum(v > cut+1e-9 for cut in CUTS[i]))
    records = []
    for key, group in selected.groupby(KEYS, sort=True):
        lo, hi = float(group.容载比.min()), float(group.容载比.max())
        records.append(dict(zip(KEYS, key)) | {FIVE[i]: LABELS[i][key[i]] for i in range(4)} | {
            '有效样本数': len(group), '容载比下界': lo, '容载比上界': hi,
            '容载比区间': interval(lo, hi), '区间宽度': hi-lo, '区间中点': (hi+lo)/2,
            '下界参数': endpoint(group.loc[group.容载比.idxmin()]),
            '上界参数': endpoint(group.loc[group.容载比.idxmax()])})
    cells = pd.DataFrame(records)
    # 独立逐记录归档，复算所有单元的计数和极值。
    direct = defaultdict(lambda: [math.inf, -math.inf, 0])
    for row in selected.to_dict('records'):
        item = direct[cell_key(row)]
        item[0] = min(item[0], row['容载比'])
        item[1] = max(item[1], row['容载比'])
        item[2] += 1
    for row in cells.to_dict('records'):
        key = tuple(row[k] for k in KEYS)
        assert np.allclose(direct[key][:2], [row['容载比下界'], row['容载比上界']], atol=1e-12, rtol=1e-12)
        assert direct[key][2] == row['有效样本数']
    assert cells.有效样本数.sum() == len(selected)
    return cells


def adjacent(cells, view):
    keyed = cells.set_index(KEYS)
    records = []
    for key, left in keyed.iterrows():
        for axis, domain in enumerate(DOMAINS):
            values = list(domain)
            at = values.index(key[axis])
            if at == len(values)-1:
                continue
            other = list(key)
            other[axis] = values[at+1]
            other = tuple(other)
            if other not in keyed.index:
                continue
            right = keyed.loc[other]
            records.append({'版本': view, '变化参数': FIVE[axis], '左单元': str(key), '右单元': str(other),
                '左区间': left.容载比区间, '右区间': right.容载比区间,
                '展示区间不同': left.容载比区间 != right.容载比区间,
                '区间不重叠': bool(left.容载比上界 < right.容载比下界-1e-9 or right.容载比上界 < left.容载比下界-1e-9),
                '中点差绝对值': abs(left.区间中点-right.区间中点),
                '两单元平均宽度': (left.区间宽度+right.区间宽度)/2})
    return pd.DataFrame(records)


def metrics(cells, pairs, view, selected):
    return {'版本': view, '参与矩阵样本数': int(cells.有效样本数.sum()),
        '不同五参数节点数': len(selected[selected[FIVE[-1]].isin(KD)][FIVE].drop_duplicates()),
        '有效单元数': len(cells), '原区间要求单元数': 432, '不同展示区间数': cells.容载比区间.nunique(),
        '单元平均区间宽度': float(cells.区间宽度.mean()), '相邻可比单元对数': len(pairs),
        '相邻展示区间不同对数': int(pairs.展示区间不同.sum()),
        '相邻区间不重叠对数': int(pairs.区间不重叠.sum()),
        '相邻单元平均中点差': float(pairs.中点差绝对值.mean()),
        '中点差与平均宽度之比': float(pairs.中点差绝对值.sum()/pairs.两单元平均宽度.sum())}


def write_workbook(tables):
    book = Workbook()
    book.remove(book.active)
    for name, frame in tables.items():
        sheet = book.create_sheet(name)
        sheet.append(list(frame.columns))
        for values in frame.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(v) else v.item() if isinstance(v, np.generic) else v for v in values])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='366092')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(48, max(13, len(str(column[0].value))*2+2))
            for cell in column:
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = '0.000000'
    path = PUBLIC/'五参数推广矩阵_新旧对比.xlsx'
    book.save(path)
    # 回读全部单元，避免交付表与计算表不一致。
    checked = 0
    readback = load_workbook(path, read_only=True, data_only=True)
    assert readback.sheetnames == list(tables)
    for name, frame in tables.items():
        rows = list(readback[name].iter_rows(values_only=True))
        assert rows[0] == tuple(frame.columns)
        assert len(rows)-1 == len(frame)
        for actual, expected in zip(rows[1:], frame.itertuples(index=False, name=None)):
            for a, b in zip(actual, expected):
                if pd.isna(b):
                    assert a is None, (name, a, b)
                elif isinstance(b, (float, np.floating)):
                    assert math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-10), (name, a, b)
                else:
                    assert a == b, (name, a, b)
                checked += 1
    readback.close()
    return checked


def panels(cells, prefix):
    lookup = cells.set_index(KEYS)
    tables = {}
    for cable, kd in itertools.product(range(3), KD):
        rows = []
        for growth, eta in itertools.product(range(4), range(3)):
            row = {'年最大净负荷平均增长率': LABELS[1][growth], '分布式电源装机渗透率': LABELS[0][eta]}
            for transfer in range(4):
                key = (eta, growth, cable, transfer, kd)
                row['初始转供率'+LABELS[3][transfer]] = lookup.loc[key, '容载比区间'] if key in lookup.index else '无有效样本'
            rows.append(row)
        tables[f'{prefix}_电缆{cable+1}_分散{kd:.2f}'] = pd.DataFrame(rows)
    return tables


def main():
    for directory in [PUBLIC, SOURCE, REVIEW]:
        directory.mkdir(parents=True, exist_ok=True)
    preserved = list(OLD.glob('*'))+list(NEW.glob('*'))+[OLD_RAW, NEW_RAW]
    preserved += list((ROOT/'研究报告/03_MD/当前版').glob('*.md'))
    preserved += [ROOT/p for p in ['研究报告/02图表/当前版/计算表.json', '研究报告/02图表/当前版/生成段落.json',
        '研究报告/04_word/研究报告终稿.docx', '研究报告/04_word/研究报告终稿.pdf',
        '实验/研究/rebuild_2026/matrix_20261012/model.py', '实验/研究/rebuild_2026/indicator_extension_20261014/model.py',
        '实验/研究/rebuild_2026/current_economic_assumptions.py']]
    before = {str(p.relative_to(ROOT)): sha(p) for p in preserved if p.is_file()}
    expanded_registry = json.loads((ROOT/'研究报告/05_review/当前版/1014补充指标全背景扩展/结果文件登记.json').read_text())
    for path, expected in expanded_registry['结果文件'].items():
        assert sha(ROOT/path) == expected, path
    old, old_status = read_and_check_raw(OLD_RAW, OLD/'统一有效样本.csv', MAPPING)
    new, new_status = read_and_check_raw(NEW_RAW, NEW/'有效样本.csv', NEW_MAPPING)
    assert len(old) == 32112 and len(new) == 90420
    reference_prices = np.isclose(new[PRICE[0]], 1) & np.isclose(new[PRICE[1]], 1)
    reference_energy = np.isclose(new[ENERGY], .2*new[FIVE[0]], rtol=0, atol=1e-10)
    views = {
        '旧库_时间2_4_6h包络': old,
        '新库_全部背景包络': new,
        '新库_原价格及2_4_6h': new[reference_prices & new[H].isin([2, 4, 6])],
        '旧库_固定4h': old[old[H].eq(4)],
        '新库_原价格固定4h': new[reference_prices & new[H].eq(4)],
        '新库_原价格4h电量比0.20': new[reference_prices & new[H].eq(4) & reference_energy],
    }
    paired = views['旧库_固定4h'].merge(views['新库_原价格4h电量比0.20'], on=FIVE,
        suffixes=('_旧', '_新'), validate='one_to_one')
    paired['容载比变化'] = paired.容载比_新-paired.容载比_旧
    matched_old = paired[FIVE+['容载比_旧']].rename(columns={'容载比_旧': '容载比'})
    matched_new = paired[FIVE+['容载比_新']].rename(columns={'容载比_新': '容载比'})
    # 端点记录需要时长；配对两侧均为4h、原价格。
    matched_old[H] = 4
    matched_new[H] = 4
    views['旧库_完全相同五参数节点'] = matched_old
    views['新库_完全相同五参数节点'] = matched_new
    cells, pairs, overview = {}, {}, []
    dimension_metrics = []
    for name, selected in views.items():
        cells[name] = matrix(selected)
        pairs[name] = adjacent(cells[name], name)
        overview.append(metrics(cells[name], pairs[name], name, selected))
        for parameter, group in pairs[name].groupby('变化参数', sort=False):
            dimension_metrics.append({'版本': name, '参数': parameter, '相邻可比对数': len(group),
                '展示区间不同对数': int(group.展示区间不同.sum()),
                '区间不重叠对数': int(group.区间不重叠.sum()),
                '平均中点差': float(group.中点差绝对值.mean()),
                '中点差与宽度之比': float(group.中点差绝对值.sum()/group.两单元平均宽度.sum())})
    old_name, new_name = list(views)[:2]
    comparison = cells[old_name].merge(cells[new_name], on=KEYS, how='left', suffixes=('_旧', '_新'), indicator=True)
    comparison['新库有样本'] = comparison['_merge'].eq('both')
    missing = comparison[~comparison.新库有样本].drop(columns='_merge').copy()
    common = comparison[comparison.新库有样本].copy()
    common['下界变化'] = common.容载比下界_新-common.容载比下界_旧
    common['上界变化'] = common.容载比上界_新-common.容载比上界_旧
    common['宽度变化'] = common.区间宽度_新-common.区间宽度_旧
    assert len(cells[old_name]) == 432 and len(cells[new_name]) == 423 and len(missing) == 9
    common_metrics = []
    common_index = cells[new_name].set_index(KEYS).index
    for name in [old_name, new_name]:
        aligned_cells = cells[name].set_index(KEYS).loc[common_index].reset_index()
        aligned_pairs = adjacent(aligned_cells, name)
        assert len(aligned_cells) == 423 and len(aligned_pairs) == 1473
        common_metrics.append({'版本': name, '共同单元数': len(aligned_cells),
            '平均区间宽度': float(aligned_cells.区间宽度.mean()),
            '相同相邻组合数': len(aligned_pairs),
            '展示区间不同对数': int(aligned_pairs.展示区间不同.sum()),
            '区间不重叠对数': int(aligned_pairs.区间不重叠.sum()),
            '平均中点差': float(aligned_pairs.中点差绝对值.mean())})
    selection = pd.DataFrame([
        ['第一层', '与容载比的相关强度', '沿用固定其余输入后的组内相关；比较绝对值。|r|≥0.40为已讨论的候选工作门槛。'],
        ['第一层', '较弱关联指标退出主矩阵索引', '95%峰值持续时间0.261、储能支撑时间0.269、两项投资比绝对值0.159及0.197。'],
        ['第二层', '参数间信息重叠', '以|r|≥0.70检查高相关组；电量—装机0.920、反向净负荷—装机0.893、站级负载率—装机0.776、正向净负荷—增长率0.999。'],
        ['第二层', '同组代表', '同组优先保留基础规划输入参数；阈值本身不唯一决定装机与电量的代表选择。'],
        ['核验', '五项低共线性', '扩展库最大参数间|r|为0.291，VIF为1.000～1.118；描述有效仿真情景库的信息重叠。'],
        ['矩阵', '最终五项', '分布式电源装机渗透率、年最大净负荷平均增长率、电缆长度占比、初始转供率、负荷分散系数。'],
        ['条件', '筛选与仿真条件', '时长、源荷曲线和价格仍影响设备配置。移出索引后须固定或用声明范围形成包络。'],
        ['证据', '样本口径', '基础参数相关取原32112节点库，补充指标取扩展90420节点库；本轮比较不把两库系数伪称为单一全变量排名。'],
    ], columns=['步骤', '内容', '采用逻辑'])
    notes = pd.DataFrame([
        ['结论', '新库能生成候选矩阵，整体区分度未提高，建议沿用原数据形成五参数矩阵。'],
        ['旧库五参数组织', '原2/4/6h三个背景取单元极值包络；保留电缆三档和分散系数三个定值，共9表、432单元。'],
        ['分档', '装机≤1、(1,2]、>2；增长≤2%、(2%,4%]、(4%,7%]、>7%；电缆≤10%、(10%,30%]、>30%；转供≤30%、(30%,50%]、(50%,70%]、>70%。'],
        ['新库全部背景', '电量、时长1～8h及两类价格0.7～1.3倍按真实采样合并；与旧包络条件范围不同，另给同价格/同时间及同节点比较。'],
        ['定值分散', '按当前报告1.00、1.10、1.15分别展示；新库另有1个1.05点单独保留，不擅自混入定值表。'],
        ['覆盖缺口', '新库缺9格，均为增长>7%、转供≤30%、分散1.15，跨三装机档和三电缆档；旧库由7.1%、7.5%边界节点覆盖。'],
        ['公平比较', '配对比较固定4h、原价格、新库电量/装机=0.20，取两库完全相同五参数节点；此参考条件由原扩展设计预先给定。'],
        ['区分度', '同一参数相邻档、其余四项保持相同，比较两位小数展示区间是否不同及完整精度区间是否不重叠；各单元和单元对等权。'],
        ['宽度', '区间变宽表示同档内包含更分散的结果，须结合相邻档重叠判断，不能单独作为改善证据。'],
        ['范围', '仅导出分析及候选查表结果；未改源稿、正式Word/PDF、锁定原始库、模型或价格。未新增求解。'],
    ], columns=['项目', '说明'])
    tables = {'采用逻辑': selection, '比较说明': notes, '总体对比': pd.DataFrame(overview),
        '分参数区分度': pd.DataFrame(dimension_metrics), '共同单元区分度': pd.DataFrame(common_metrics),
        '共同423格比较': common.drop(columns='_merge'),
        '新库缺失9格': missing, '相同节点比较': paired,
        '定值以外样本': new[~new[FIVE[-1]].isin(KD)]}
    for i, name in enumerate(views, 1):
        tables[f'单元{i}_{name}'] = cells[name]
    tables.update(panels(cells[old_name], '旧数据五参'))
    tables.update(panels(cells[new_name], '新数据候选'))
    for name, frame in {'总体对比': tables['总体对比'], '分参数区分度': tables['分参数区分度'],
        '共同单元比较': tables['共同423格比较'], '新库缺失单元': missing,
        '相同节点比较': paired, '原数据五参数矩阵单元': cells[old_name],
        '新数据候选矩阵单元': cells[new_name]}.items():
        frame.to_csv(PUBLIC/(name+'.csv'), index=False, encoding='utf-8-sig')
    checked = write_workbook(tables)
    source = {'原始资料': {str(p.relative_to(ROOT)): before[str(p.relative_to(ROOT))]
                           for p in [OLD_RAW, NEW_RAW, OLD/'统一有效样本.csv', NEW/'有效样本.csv']},
        '展示口径': dict(五参数=FIVE, 分档=CUTS, 分散定值=KD, 单元总数=432),
        '比较条件': notes.to_dict('records'), '筛选表述': selection.to_dict('records'),
        '没有追加采样或寻优': True}
    dump(SOURCE/'比较口径及来源.json', source)
    for path, expected in before.items():
        assert sha(ROOT/path) == expected, path
    same = int(np.isclose(paired.容载比变化, 0, rtol=0, atol=1e-9).sum())
    summary = {'建议': '原数据五参数矩阵；扩展库用于补充指标相关性及背景条件核验',
        '旧原始状态': old_status, '新原始状态': new_status,
        '旧五参数单元': 432, '新五参数单元': 423, '新缺失单元': 9,
        '共同单元数': len(common), '旧共同单元平均宽度': float(common.区间宽度_旧.mean()),
        '新共同单元平均宽度': float(common.区间宽度_新.mean()),
        '宽度增加单元数': int((common.宽度变化 > 1e-9).sum()),
        '宽度相同单元数': int((abs(common.宽度变化) <= 1e-9).sum()),
        '宽度减少单元数': int((common.宽度变化 < -1e-9).sum()),
        '共同单元区分度': common_metrics,
        '相同五参数节点对': len(paired), '容载比相同节点': same,
        '容载比增加节点': int((paired.容载比变化 > 1e-9).sum()),
        '容载比减少节点': int((paired.容载比变化 < -1e-9).sum()),
        '配对结果解释': '同节点容载比降低为参考源荷出力变化的配置结果，不等同于分档区分度提高。',
        '核验': {'状态': '通过', '原始库与有效样本逐节点对照': True,
            '全部版本单元计数与端点独立逐记录复算': True, '工作簿回读单元格': checked,
            '既有源稿成果模型数据SHA256保持': True, '新增寻优节点': 0}}
    dump(REVIEW/'数据核验.json', summary)
    dump(REVIEW/'本轮未修改文件登记.json', before)
    markdown = ('# 五参数推广矩阵新旧对比\n\n'
        '已按用户确认的两层筛选逻辑组织分析，尚未修改研究报告。筛选逻辑及数值口径见工作簿“采用逻辑”。\n\n'
        '建议沿用原32112节点库，移出储能时间索引后，以2/4/6h三个原背景取各单元容载比包络。'
        '保留原有四项区间分档和分散系数三个定值，形成9张五参数查表、432格。\n\n'
        '新库按相同分档可生成423格候选矩阵。共同423格的平均宽度由'
        f'{summary["旧共同单元平均宽度"]:.4f}增至{summary["新共同单元平均宽度"]:.4f}；'
        '308格变宽，115格相同，没有缩窄格。固定原价格及原2/4/6h范围后，区间重叠增多的判断仍成立。'
        '扩大库主要扩展电量、时长和价格背景，不是对五项主参数均匀加密，故样本多未提高主矩阵区分度。\n\n'
        f'固定4h、原价格及预设电量/装机比0.20后，完全相同五参数节点有{len(paired)}对，'
        f'{same}对容载比相同。配对版本完整比较见工作簿，避免将少量参考节点造成的窄区间误判为改善。\n\n'
        '“旧数据五参”九个工作表可供后续五参数报告组织审阅；“新数据候选”九表仅用于对照，空格保留真实缺口。'
        '原模型、原库、既有矩阵及报告成果均未改写。所有比较由原始库和公开有效样本双向核对，无新增求解。\n')
    (PUBLIC/'分析结论.md').write_text(markdown, encoding='utf-8')
    outputs = list(PUBLIC.glob('*'))+list(SOURCE.glob('*'))+[REVIEW/'数据核验.json', REVIEW/'本轮未修改文件登记.json', Path(__file__)]
    dump(REVIEW/'结果文件登记.json', {str(p.relative_to(ROOT)): sha(p) for p in outputs if p.is_file()})
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(tables['总体对比'].to_string(index=False))


if __name__ == '__main__':
    main()
