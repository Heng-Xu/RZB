"""核对当前报告来源并生成本轮人读、机读审查文件。"""
from pathlib import Path
from functools import lru_cache
import csv, hashlib, json, math, re
import openpyxl
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / '研究报告/数据来源'
QA = ROOT / '研究报告/05_review/闭环核验'

def read(p):
    return json.loads(p.read_text())

@lru_cache(None)
def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def same(a, b):
    if a is None or b is None:
        return a is None and b is None
    try:
        return math.isclose(float(a), float(b), rel_tol=1e-12, abs_tol=1e-9)
    except (TypeError, ValueError):
        return a == b

@lru_cache(None)
def csvrows(p):
    with p.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

@lru_cache(None)
def workbook(p):
    return openpyxl.load_workbook(p, read_only=True, data_only=True)

def cell(p, loc):
    sheet, address = loc.split('!')
    return workbook(p)[sheet][address].value

def record_check(r):
    p = ROOT / r['source']
    assert sha(p) == r['source_sha256'], r['id'] + '来源哈希'
    loc = r['source_location']
    m = re.fullmatch(r'CSV物理行(\d+)；(?:字段)?(.+)', loc)
    if m:
        raw = csvrows(p)[int(m[1])-2][m[2]]
        if m[2]=='amount_yuan':
            r['historical_raw_value'] = r['raw_value']
            r['raw_value'] = raw
        else:
            assert same(raw, r['raw_value']), r['id'] + '原值'
        assert same(None if raw in ('', 'None') else float(raw)/10000 if m[2]=='amount_yuan' and r['unit'].startswith('万元') else raw, r['value']), r['id'] + 'CSV取值'
        return '原CSV行及字段逐项核对'
    m = re.fullmatch(r'JSON \$\.(\w+)\[(\d+)\]\.(\w+)', loc)
    if m:
        assert same(read(p)[m[1]][int(m[2])][m[3]], r['value']), r['id']
        return '原JSON坐标逐项核对'
    m = re.match(r'L(\d+):\s*(.*)', loc)
    if m:
        line = p.read_text().splitlines()[int(m[1])-1]
        assert line.strip() == m[2].strip(), r['id'] + '代码行'
        nums = re.findall(r'(?<![\w.])-?\d+(?:\.\d+)?(?![\w.])', line)
        assert any(same(n, r['value']) for n in nums), r['id'] + '代码常数'
        return '参数声明行及数值核对；属于研究设置'
    if loc.startswith('Sheet1!P') and 'C列工程规模' in loc:
        import sys
        sys.path.insert(0,str(ROOT/'实验/研究'))
        from rebuild_2026.incremental_cost import replacement_cost_coefficients
        voltage = 35 if 'P43' in loc else 110
        c = next(x for x in replacement_cost_coefficients() if x['voltage_kv']==voltage)
        assert same(c['base_coefficient_10k_cny_per_purchased_mva'],r['value']),r['id']
        r['calculation_inputs'] = c
        return '原工程投资及新购容量加总复算；完整分子、分母和行号附后'
    m = re.fullmatch(r'Sheet1![A-Z]+\d+', loc)
    if m:
        raw = cell(p, loc)
        assert same(raw, r['raw_value']), r['id'] + '原表'
        factor = 10 if r['unit'] in ('MW', 'MVA') else 1
        assert same(float(raw)*factor, r['value']), r['id'] + '单位换算'
        return 'Excel单元格及单位换算核对'
    raise AssertionError('尚未支持的精确定位：' + r['id'] + ' ' + loc)

def derivation(r):
    field = r.get('internal_field', '')
    if 'reference_peak' in field:
        return '年度固定正向参考峰：邳州读取同层年度降压负荷；市区以29站2025年同步峰乘全市区年度峰比例。见正文3.3、式（4-4），原值定位另列。'
    if 'clr' in field:
        return '同年度、同范围主变容量MVA除以固定正向参考峰MW；按式（4-4）计算。'
    if field=='station_count':
        return '按研究单元纳入模型的唯一站点身份计数；同层、同范围，不将全区背景站点计入局部样本。'
    if 'capacity_mva' in field:
        if 'initial' in field:return '2021年共同起点的离散主变容量加总；市区按样本与全区年度容量比例分配，非历史实测逐站台账。'
        return '同年度、同研究单元逐站两台选定主变额定容量加总，单位MVA；式（4-4）、（4-11）。'
    if 'storage_modules' in field:
        return '同年度、同研究单元的在役储能整数柜数加总；柜数非有效功率，后者按式（4-9）另算。'
    if 'tie_mw' in field or 'line_transfer_mw' in field:
        return '读取完整路径对应年度的既有或拟建区段转供功率；供受端同额计量，式（4-16）、（4-17）；单位MW，非电量。'
    if 'contingency' in field:
        return '单台主变停运恢复比例的研究设置；邳州0.60、市区0.45，式（4-14）；没有完整工程N−1验证。'
    if 'budget' in field:
        return '刚性研究单元累计正向增配预算比例；式（4-15），不是逐站或逐年独立限额。'
    if field=='reported_or_sample_scaled_2021_peak_mw':
        return '2021年同电压参考峰：邳州取原统计负荷；市区样本同步峰按全市区年度峰比例缩放；正文3.3。'
    if 'cagr' in field:
        return '年度参考峰与2021年参考峰的比值取1/(年度−2021)次幂，再减1；式（4-7）。'
    if 'reverse_station_peak' in field:
        return '先取每站反向峰，再取各站中的最大值；式（4-5），不是区域同步峰。'
    if 'd95' in field:
        return '2025年模板中达到方向峰值95%的最长连续小时段；年度显示取各站最大时长；式（4-8）。'
    if 'cost' in field or 'npv' in field:
        return '沿完整规划路径按式（4-20）至（4-25）重建首次、更新及运维现金流，并折现至2021年；原求解量及复算量分别保留。'
    if 'scale' in field:
        return '措施价格相对于同规格基准的倍率；式（4-10），改变参数后重新求解完整路径。'
    return r['formula']

def main():
    registry = read(ROOT/'研究报告/03_MD/图表引用登记.json')
    original = read(OUT/'data_source.json')
    records, supplements, materials = {}, {}, []
    for item in registry:
        p = ROOT/item['source']; d = read(p)
        for r in d.get('records', []):
            if r['id'] not in records:
                r = dict(r)
                r['verification'] = record_check(r)
                r['current_calculation_note'] = derivation(r)
                records[r['id']] = r
        for s in d.get('supplemental_sources', []):
            if 'installed_source' in s:
                supplements[s['id']] = s
        materials.append({**item, 'source_sha256':sha(p), 'title':d['title'],
            'processing':d['processing'], 'calculation':d['calculation'],
            'data':d['data'], 'record_ids':[r['id'] for r in d.get('records', [])],
            'supplemental_sources':d.get('supplemental_sources', [])})
    for s in supplements.values():
        a, b = ROOT/s['installed_source'], ROOT/s['load_source']
        s['installed_source_sha256'], s['load_source_sha256'] = sha(a), sha(b)
        assert same(float(cell(a, s['installed_location']))*10, s['installed_mw']), s['id']+'装机'
        assert same(float(cell(b, s['load_location']))*10, s['county_reference_peak_mw']), s['id']+'负荷'
        assert same(s['installed_mw']/s['county_reference_peak_mw'], s['county_source_load_ratio']), s['id']+'源荷比'
        s['verification'] = '两个原始单元格、万千瓦至MW换算及比值独立重算'
    city_path=ROOT/'实验/研究/rebuild_2026/source_audit/city_2025_scenario_sensitivity.csv'
    city_rows=csvrows(city_path)
    for rid,field,variant,name in [('X001','observed_hours','observed','市区29站共同观测小时数'),('X002','imputed_hours','weekly_mean','市区补值模板小时数')]:
        index,row=next((i,r) for i,r in enumerate(city_rows) if r['model_station_id']=='__DISTRICT__' and r['variant']==variant)
        r={'id':rid,'name':name,'variable':field,'value':int(row[field]),'raw_value':row[field],'unit':'h','source':str(city_path.relative_to(ROOT)),'source_location':f'CSV物理行{index+2}；字段{field}','source_sha256':sha(city_path),'formula':'年度共同小时计数；8339+421=8760','internal_field':field,'limitation':'观测与周类比补值分列，不能称为全年实测'}
        r['verification']=record_check(r);r['current_calculation_note']=r['formula'];records[rid]=r
    assert records['X001']['value']+records['X002']['value']==8760
    # Attach direct evidence for numerical model settings, without matching unrelated identical numbers.
    parameter_sources=[]
    for relative in ['实验/研究/rebuild_2026/joint_lifecycle_optimizer.py','实验/研究/rebuild_2026/simulation_reserve_policy.py','实验/研究/rebuild_2026/annual_no_tie_investment_submodel.py']:
        p=ROOT/relative
        terms=['MAX_PLANNING_STORAGE_MODULES','POWER_FACTOR','factor = 0.95','DESIGNED_NEW_LINE','EXISTING_PATH_CAP','def cost_factors','network_om:','line_life:','mip_rel_gap','time_limit','optimum + 1e-5','optimum + 1e-6','minimum_negative_reserve - 1e-6','contingency','budget','clr_cap']
        for i,line in enumerate(p.read_text().splitlines(),1):
            if any(t in line for t in terms):parameter_sources.append({'source':relative,'line':i,'text':line.strip(),'sha256':sha(p),'nature':'实际实现与研究设置；不作为导则参数'})
    gap = dict(next(r for r in original['records'] if r['id']=='D000346'))
    gap['verification']=record_check(gap);gap['current_calculation_note']='市区2021年离散共同起点容量减比例折算目标；单位MVA，非费用或相关系数。'
    records[gap['id']]=gap
    graph = {n['path']:n for n in original['source_graph']}
    reached = set()
    def visit(path):
        if path in reached or path not in graph:return
        reached.add(path)
        for p in graph[path]['parents']:visit(p)
    for r in records.values():visit(r['source'])
    nodes = [graph[p] for p in sorted(reached)]
    for n in nodes:
        assert sha(ROOT/n['path']) == n['sha256'], n['path']
    equations = read(ROOT/'研究报告/05_review/装配清单.json')['equations'][:29]
    eqsources = read(ROOT/'研究报告/03_MD/公式来源对应.json')
    for e in equations:e['source_note'] = eqsources[e['number']]
    narrative = read(ROOT/'研究报告/03_MD/检查证据/正文数值定位.json')
    section_materials={
      '3.1':['表3-1','表3-2'],'3.2':['图3-1','图3-2'],'3.3':['图3-3','图3-4','表3-3'],'3.5':['图3-6'],
      '5.1':['图5-1'],'5.2':['表5-1','图5-3'],'5.3':['表5-2'],'5.4':['表5-3','图5-6'],'5.5':['表5-4'],
      '6.1':['表3-1','表6-1'],'6.2':['表6-2','图6-1','图6-2','图3-2'],'6.3':['表6-3','图6-3'],'7.1':['表7-1','表5-1','表5-2','表5-3']}
    occurrence_sections=[]
    for p in sorted((ROOT/'研究报告/03_MD').glob('[0-9][0-9] *.md')):
        if p.name[:2] not in ['03','05','06','07']:continue
        section=''
        for line in p.read_text().splitlines():
            heading=re.match(r'^#{2,3} (\d+\.\d+)',line)
            if heading:section=heading[1]
            if line.startswith(('#','|')):continue
            line=re.sub(r'!\[[^\]]*\]\([^)]*\)','',line)
            for m in re.finditer(r'(?<![a-zA-Z0-9_.])-?\d+\.\d+(?![a-zA-Z0-9_.])',line):
                occurrence_sections.append((p.name,m[0],section))
    assert len(occurrence_sections)==len(narrative)
    for n,(chapter,value,section) in zip(narrative,occurrence_sections):
        assert chapter==n['chapter'] and value==n['value']
        n['section']=section
        n['legacy_numeric_matches']=n.pop('source_figures_tables')
        n['reviewed_source_materials']=section_materials.get(section,[])
        if value=='0.649':n['reviewed_source_materials']=['D000346']
        if section=='6.2':
            direct={'0.742':['CS003'],'4.75':['D017030'],'9.05':['D017029'],
                    '44.543':['D017027'],'133.629':['D017027','式（4-23）']}
            if value in direct:n['reviewed_source_materials']=direct[value]
            if value=='133.629':
                n['calculation']='3 km × 44.543 万元/km = 133.629 万元；候选线路长度为正文6.2及式（4-23）的研究设置'
                assert same(3*records['D017027']['value'],float(value))
        assert n['reviewed_source_materials'],n
        n['review_basis']='按正文小节、指标含义和工程范围定位；原候选同值匹配只保留作检查器追溯，不作为语境出处'
    scientific = read(QA/'本轮独立复算.json')
    known_docs = {'研究报告/AGENTS.md','研究报告/终稿/写作与科学表达硬约束.md'} | {n['path'] for n in read(QA/'历史文件定位.json')}
    numeric_problems = [p for p in scientific['problems'] if p.removeprefix('原输入或原结果变化：') not in known_docs]
    assert not numeric_problems, numeric_problems
    stage2 = read(ROOT/'研究报告/03_MD/检查证据/阶段2自动核验.json')
    assert stage2['passed'], stage2['failed']
    docx = ROOT/'研究报告/04_word/研究报告终稿.docx'
    pdf = docx.with_suffix('.pdf')
    data = {'schema_version':'current-report-1.0',
        'report':{'docx':str(docx.relative_to(ROOT)),'docx_sha256':sha(docx),'pdf':str(pdf.relative_to(ROOT)),'pdf_sha256':sha(pdf)},
        'scope':'当前七章报告、17图、15表、29组公式；来源记录不代表工程校核完成',
        'source_dictionary':{'path':'研究报告/数据来源/data_source.json','sha256':sha(OUT/'data_source.json')},
        'region_mapping':original['region_mapping'], 'materials':materials,
        'parameter_sources':parameter_sources, 'literature_evidence':read(ROOT/'研究报告/03_MD/检查证据/第二章文献核定.json'),
        'records':list(records.values()), 'county_statistics':list(supplements.values()),
        'source_graph':nodes, 'equations':equations, 'narrative_decimal_locations':narrative,
        'verification':{'stage2_passed':True,'stage2_checks':stage2['checks'],
            'independent_raw_status':scientific['status'], 'independent_numeric_passed':not numeric_problems,
            'independent_checks':scientific['checks'], 'documentation_differences':scientific['problems'],
            'documentation_resolution':'现行AGENTS及历史文件归档由阶段2逐文件核验；不修改原模型结果',
            'limitations':['数值候选定位需结合研究范围、年度和指标含义阅读，数值相等本身不构成来源证明',
            '未重新调用求解器；六路径由原始输入和逐站动作独立复算',
            '缺少完整潮流、标准承载力、逐时SOC、损耗和完整拆改残值费用']}}
    (OUT/'当前报告数据来源.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
    d = Document()
    for section in d.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)
    for style in d.styles:
        if not hasattr(style,"font"):continue
        style.font.name = 'Times New Roman'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '方正仿宋_GBK')
    d.styles['Normal'].font.size = Pt(10)
    d.styles['Normal'].paragraph_format.space_after = Pt(3)
    d.add_heading('当前报告数据来源与计算审查', 0)
    d.add_paragraph('本文件对应本轮修订的Word和PDF。用于从报告位置查到原始单元格、结果字段、计算规则和适用限制。原始数据字典与本轮对应文件同时保留。')
    d.add_paragraph(f'覆盖17张图、15张表、29组公式，直接引用{len(records)}条来源记录，并独立复算{len(supplements)}条区县装机、降压峰及源荷比。正文{len(narrative)}次数值小数定位见后附记录；整数和参数依据随图表及公式登记。')
    d.add_paragraph('功率统一用MW，主变容量用MVA，储能电量用MWh，时长用h，长度用km，金额用万元。报告正文的峰值和比值主要保留三位小数，容量主要保留一位，费用结果保留两位。原始精度不截断；采购锚点、工程参数和公式系数保留计算所需精度。百分比与比例通过乘100换算，不混读。')
    d.add_paragraph('图表分项独立舍入时可能出现0.01万元尾差。源码、原始记录、完整计算值和对外显示值分别保留。方正仿宋_GBK是本环境实际可用的仿宋字体；报告与图件统一使用该字体，英文和数字使用Times New Roman。')
    d.add_paragraph('标准设备承载力、县级校核和可开放容量公式已与DL/T 2041-2025扫描原件PDF第6至8页核对。80%反向系数有运行方式条件，研究筛查系数不替代标准评估。3.2上限、恢复比例、预算、寿命、运维和折现率属于研究设置。')
    d.add_heading('报告版本与核验边界', 1)
    for key,value in data['report'].items():d.add_paragraph(key+'：'+value)
    d.add_paragraph(f'阶段2完成{stage2["checks"]}项检查。本轮独立复算执行{scientific["checks"]}项检查，覆盖六路径、12条年度案例、原表、离散容量、储能有效功率、转供和现金流。原检查中的五项文件变化为写作规则与已移除旧文档，已通过现行规则及归档哈希分别核验。完整求解器复现沿用既有冻结证据，本轮没有再次求解。')
    d.add_paragraph('人类化表达按humanizer技能通读及复核，保持事实、数值和证据强度。程序扫描只检查常见词句；公式、参数名及来源路径保持可追溯写法。')
    d.add_heading('逐图逐表来源与计算', 1)
    for m in materials:
        d.add_heading(m['number']+' '+m['title'], 2)
        d.add_paragraph('正文位置：'+m['chapter']+'；原来源文件：'+m['source'])
        d.add_paragraph('处理：'+m['processing']+'。计算与显示：'+m['calculation']+'。')
        d.add_paragraph('逐项记录：'+('、'.join(m['record_ids']) or '文字示意或区县补充原表定位'))
        d.add_paragraph('来源SHA256：'+m['source_sha256'])
    d.add_heading('区县原表取值与复算', 1)
    for s in supplements.values():
        d.add_paragraph(f'{s["id"]} {s["region"]} {s["year"]}年：装机 {s["installed_10k_kw"]}万千瓦×10={s["installed_mw"]:.9g} MW；降压峰 {s["county_reference_peak_mw"]:.9g} MW；源荷近似比值={s["county_source_load_ratio"]:.9g}。')
        d.add_paragraph('装机原表：'+s['installed_source']+' '+s['installed_location']+'；负荷原表：'+s['load_source']+' '+s['load_location'])
    d.add_heading('公式与成本科目', 1)
    for e in equations:
        d.add_paragraph('式（'+e['number']+'）：'+'；'.join(e['source_note']))
    d.add_paragraph('首次投资按主变新购完整容量、储能采购规模差和新线新增状态分别核算；更新投资按原价和计算寿命发生；固定运维从投运次年计至2041年；三项统一折现至2021年。损耗、辅助工程、拆改、可靠性损失与残值另列待计算框架，未赋值，不能读作真实零成本。')
    d.add_heading('观测小时与研究参数依据', 1)
    d.add_paragraph('市区29站观测8339 h与补值421 h分别来自市区场景表的区域observed和weekly_mean行，合计8760 h。原始小时计数已核对，补值不作为观测数据。')
    for r in parameter_sources:
        d.add_paragraph(r['source']+' 第'+str(r['line'])+'行：'+r['text'])
    d.add_heading('正文小数定位', 1)
    for n in narrative:
        d.add_paragraph(n['chapter']+'；数值 '+n['value']+'；原句邻域：'+n['context']+'；来源定位：'+'、'.join(n['reviewed_source_materials'])+('；计算：'+n['calculation'] if 'calculation' in n else ''))
    d.add_heading('直接引用记录明细', 1)
    for r in records.values():
        d.add_paragraph(f'{r["id"]} {r["name"]}：{r["value"]} {r["unit"]}')
        d.add_paragraph('来源：'+r['source']+'；位置：'+r['source_location'])
        d.add_paragraph('计算：'+r['current_calculation_note']+'；限制：'+r.get('limitation','见正文范围说明'))
    d.save(OUT/'当前报告数据来源与计算审查.docx')
    print(json.dumps({'records':len(records),'county_rows':len(supplements),'graph_nodes':len(nodes),'materials':len(materials),'formulas':len(equations),'passed':True},ensure_ascii=False))

if __name__ == '__main__':main()
