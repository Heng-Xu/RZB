"""Build the independent six-section source note from its Markdown source."""
from pathlib import Path
import re, subprocess
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'研究报告/03_MD/数据来源说明.md'
OUT=ROOT/'研究报告/数据来源/数据来源说明.docx'

def font(run, size=11):
    run.font.name='Times New Roman';run.font.size=Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'方正仿宋_GBK')

def main():
    text=SRC.read_text()
    assert len(re.findall(r'^## ',text,re.M))==6
    assert '缺项及后续补充要求' not in text
    # Use Pandoc's supported math delimiters for editable native Word math.
    text=re.sub(r'\\\((.*?)\\\)',lambda m:'$'+m[1]+'$',text)
    proc=subprocess.run(['pandoc','-f','markdown+tex_math_dollars','-t','docx','-o',str(OUT)],input=text,text=True,capture_output=True,check=True)
    d=Document(OUT)
    for sec in d.sections:
        sec.top_margin=sec.bottom_margin=Cm(2)
        sec.left_margin=sec.right_margin=Cm(2)
        p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'方正仿宋_GBK')
    normal.paragraph_format.line_spacing=1.25
    normal.paragraph_format.space_after=Pt(5)
    for p in d.paragraphs:
        heading=p.style.name.startswith('Heading')
        size=18 if p.style.name=='Heading 1' else 13 if heading else 11
        if heading:p.paragraph_format.keep_with_next=True
        for run in p.runs:font(run,size)
    for ti, table in enumerate(d.tables):
        table.autofit=False
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'808080');borders.append(e)
        table._tbl.tblPr.append(borders)
        widths=[3.1,7.0,6.9] if ti==0 else [5.0,1.7,4.0,6.3]
        for col,w in zip(table.columns,widths):col.width=Cm(w)
        for ri,row in enumerate(table.rows):
            if ri==0:
                repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
            for cell,w in zip(row.cells,widths):
                cell.width=Cm(w)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.1
                    p.paragraph_format.space_after=Pt(3)
                    for run in p.runs:font(run,9)
                    if ri==0:
                        for run in p.runs:run.bold=True
    d.save(OUT)
    print(OUT)

if __name__=='__main__':main()
