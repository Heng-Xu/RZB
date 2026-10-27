"""成果总结原生公式：直接复用研究报告的OMML，检查同义源式与对称三栏布局。"""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
import re
from lxml import etree as E

W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
NS={'w':W,'m':M}
qw=lambda name:'{'+W+'}'+name
qm=lambda name:'{'+M+'}'+name
SPECS={'1':{'report_number':'4-7','source_tag':'4-10','chapter':'04','continuation':True},
       '2':{'report_number':'4-8','source_tag':'4-11','chapter':'04','continuation':False},
       '3':{'report_number':'5-8','source_tag':'5-8','chapter':'05','continuation':True}}
INDENTS=['firstLine','firstLineChars','hanging','hangingChars','left','right','start','end','leftChars','rightChars','startChars','endChars']


def node(tag,math=False,**attrs):
    q=qm if math else qw;e=E.Element(q(tag))
    for k,v in attrs.items():e.set(q(k),str(v))
    return e


def latex_key(value):
    return re.sub(r'\s+','',re.sub(r'\\tag\{[^}]+\}','',value.strip('$')))


def formula_number(block):
    match=re.search(r'\\tag\{([^}]+)\}',block)
    assert match and match[1]in SPECS,'成果总结公式编号无登记'
    return match[1]


def source_check(root,text):
    formulas=re.findall(r'\$\$([\s\S]*?)\$\$',text)
    assert [formula_number(b)for b in formulas]==list(SPECS),'成果总结须按序列出成本及供电约束公式'
    evidence=[]
    for block in formulas:
        num=formula_number(block);spec=SPECS[num]
        path=next((root/'研究报告/03_MD/当前版').glob(spec['chapter']+' *.md'))
        original=next(b for b in re.findall(r'\$\$([\s\S]*?)\$\$',path.read_text())if r'\tag{'+spec['source_tag']+'}'in b)
        assert latex_key(original)==latex_key(block),'成果总结公式须对应研究报告源式：'+num
        evidence.append({'成果总结式号':num,'研究报告源式号':spec['source_tag'],'当前Word式号':spec['report_number'],'源稿':str(path.relative_to(root))})
    return evidence


def reference_math(root,number):
    spec=SPECS[number]
    with ZipFile(root/'研究报告/04_word/研究报告终稿.docx')as z:doc=E.fromstring(z.read('word/document.xml'))
    candidates=[t for t in doc.xpath('//w:body/w:tbl[.//m:oMath]',namespaces=NS)if ''.join(t.xpath('.//w:t/text()',namespaces=NS))=='（'+spec['report_number']+'）']
    assert len(candidates)==1,'研究报告原生公式缺失：'+spec['report_number']
    maths=candidates[0].xpath('.//m:oMath',namespaces=NS);assert len(maths)==1
    return deepcopy(maths[0])


def normalize_math(math):
    for run in math.findall('.//'+qm('r')):
        mr=run.find(qm('rPr'))
        if mr is None:mr=node('rPr',math=True);run.insert(0,mr)
        text=''.join(run.xpath('./m:t/text()',namespaces=NS));style=mr.find(qm('sty'));script=mr.find(qm('scr'))
        italic=any(('A'<=c<='Z')or('a'<=c<='z')or('α'<=c<='ϕ')for c in text)
        if style is not None and style.get(qm('val'))in ['p','b']and script is None:italic=False
        nor=mr.find(qm('nor'))
        if nor is None:nor=node('nor',math=True);mr.append(nor)
        nor.set(qm('val'),'1')
        pr=run.find(qw('rPr'))
        if pr is None:pr=node('rPr');run.insert(list(run).index(mr)+1,pr)
        for tag in ['rFonts','sz','szCs','i','iCs']:
            for child in pr.findall(qw(tag)):pr.remove(child)
        pr.insert(0,node('rFonts',ascii='Times New Roman',hAnsi='Times New Roman',cs='Times New Roman',eastAsia='方正仿宋_GBK'))
        for tag,value in [('sz',28),('szCs',28),('i',int(italic)),('iCs',int(italic))]:pr.append(node(tag,val=value))
    # 明确控制属性字体，避免继承中文正文字体。
    for control in math.findall('.//'+qm('ctrlPr')):
        pr=control.find(qw('rPr'))
        if pr is None:pr=node('rPr');control.append(pr)
        for child in list(pr):
            if child.tag in [qw('rFonts'),qw('sz'),qw('szCs')]:pr.remove(child)
        pr.insert(0,node('rFonts',ascii='Times New Roman',hAnsi='Times New Roman',cs='Times New Roman',eastAsia='方正仿宋_GBK'))
        pr.append(node('sz',val=28));pr.append(node('szCs',val=28))
    return math


def equation_table(root,block):
    number=formula_number(block);t=node('tbl');pr=node('tblPr')
    pr.append(node('tblW',w=8306,type='dxa'));pr.append(node('jc',val='center'));pr.append(node('tblInd',w=0,type='dxa'));pr.append(node('tblLayout',type='fixed'));pr.append(node('tblCaption',val='ReportEquation'))
    borders=node('tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:borders.append(node(name,val='nil'))
    pr.append(borders);mar=node('tblCellMar')
    for side in ['top','bottom','left','right']:mar.append(node(side,w=0,type='dxa'))
    pr.append(mar);t.append(pr);grid=node('tblGrid');widths=[800,6706,800]
    for width in widths:grid.append(node('gridCol',w=width))
    t.append(grid);tr=node('tr');rp=node('trPr');rp.append(node('cantSplit'));tr.append(rp)
    for i,width in enumerate(widths):
        tc=node('tc');cp=node('tcPr');cp.append(node('tcW',w=width,type='dxa'));cp.append(node('vAlign',val='center'));cm=node('tcMar')
        for side in ['top','bottom','left','right']:cm.append(node(side,w=0,type='dxa'))
        cp.append(cm);tc.append(cp);p=node('p');pp=node('pPr');pp.append(node('ind',**{k:0 for k in INDENTS}));pp.append(node('jc',val='right'if i==2 else'center'));pp.append(node('spacing',before=50,after=50,line=360,lineRule='auto'));pp.append(node('snapToGrid',val=0));pp.append(node('keepNext'));p.append(pp)
        if i==1:p.append(normalize_math(reference_math(root,number)))
        elif i==2:
            r=node('r');rpr=node('rPr');rpr.append(node('rFonts',ascii='Times New Roman',hAnsi='Times New Roman',cs='Times New Roman',eastAsia='方正仿宋_GBK'));rpr.append(node('sz',val=24));r.append(rpr);txt=node('t');txt.text='（'+number+'）';r.append(txt);p.append(r)
        tc.append(p);tr.append(tc)
    t.append(tr)
    return t


def math_text(node_):return ''.join(n.text or ''for n in node_.iter(qm('t')))


INLINE_KEYS={r'C_\mathrm{plan}':'Cplan',r'C_e':'Ce',r'\mathcal E_T':'ET',
             r'\mathcal E_B':'EB',r'\mathcal E_L':'EL',**{k:k for k in 'CIyLqrTkte'}}


def inline_segments(value):
    """正文美元标记只承载已出现在成本公式中的符号。"""
    pieces=re.split(r'(\$[^$]+\$)',value)
    result=[]
    for piece in pieces:
        if not piece:continue
        if piece.startswith('$'):
            token=piece[1:-1];assert token in INLINE_KEYS,'未登记的行内符号：'+token
            result.append((True,token,INLINE_KEYS[token]))
        else:result.append((False,piece,piece))
    return result


def inline_math(root,token,size=32):
    """复用原公式的变量、下标和花体，生成行内原生OMML。"""
    key=INLINE_KEYS[token];reference=reference_math(root,'2'if key in ['Cplan','Ce','ET','EB','EL','e']else'1')
    if len(key)>1:
        choices=[n for n in reference.iter(qm('sSub'))if math_text(n)==key]
    else:choices=[n for n in reference.iter(qm('r'))if math_text(n)==key]
    assert choices,'原公式中未找到行内符号：'+token
    result=E.Element(qm('oMath'),nsmap=NS);result.append(deepcopy(choices[0]));normalize_math(result)
    for run in result.iter(qm('r')):
        pr=run.find(qw('rPr'))
        for tag in ['sz','szCs']:pr.find(qw(tag)).set(qw('val'),str(size))
    return result


def check_native(root,path):
    from docx import Document
    from report_math_layout import check_equation_layout,equation_tables
    doc=Document(path);tables=equation_tables(doc);assert len(tables)==len(SPECS)
    records=[]
    for t in tables:
        number=t.cell(0,2).text.strip('（）');assert number in SPECS
        assert math_text(t._tbl)==math_text(reference_math(root,number)),'数学符号与研究报告不一致：'+number
        for r in t._tbl.xpath('.//m:r'):
            assert r.find('./'+qm('rPr')+'/'+qm('nor')).get(qm('val'))=='1','公式须明确正斜体字形'
            assert r.find('./'+qw('rPr')+'/'+qw('rFonts')).get(qw('ascii'))=='Times New Roman'
            assert r.find('./'+qw('rPr')+'/'+qw('i')).get(qw('val'))in ['0','1']
        records.append({'式号':number,'研究报告Word式号':SPECS[number]['report_number'],'字符':math_text(t._tbl),'原生OMML':'通过','字体与正斜体':'通过'})
        # 仅在内存副本恢复来源编号，复用项目对连续推导与单行关系的布局检查。
        t.cell(0,2).paragraphs[0].runs[0].text='（'+SPECS[number]['report_number']+'）'
    check_equation_layout(doc)
    return records


def prose_bounds(spans,yc):
    # 数字较多的正文行也会只有一两个汉字，仍须排除在公式边界之外。
    prose=[s for s in spans if s['size']>=15.5 and re.search('[\u4e00-\u9fff]',s['text'])]
    low=max((s['bbox'][3]for s in prose if s['bbox'][3]<yc),default=70)
    high=min((s['bbox'][1]for s in prose if s['bbox'][1]>yc),default=775)
    return low,high


def check_pdf(root,path,pdf):
    from docx import Document
    from report_math_layout import equation_tables
    doc=Document(path);results=[]
    for t in equation_tables(doc):
        number=t.cell(0,2).text
        marks=[(page,word)for page in pdf for word in page.get_text('words')if word[0]>460 and word[4]==number]
        assert len(marks)==1,('PDF式号缺失或重复',number)
        page,mark=marks[0]
        all_spans=[s for b in page.get_text('dict')['blocks']for line in b.get('lines',[])for s in line['spans']]
        # 公式内“i为故障站”是下标条件；正文边界用16磅正文识别。
        yc=(mark[1]+mark[3])/2
        low,high=prose_bounds(all_spans,yc)
        spans=[s for s in all_spans if s['bbox'][0]>90 and s['bbox'][2]<460 and s['bbox'][1]>=low and s['bbox'][3]<=high and s['text'].strip()]
        assert spans,('PDF公式无可读主体',number)
        left=min(s['bbox'][0]for s in spans);right=max(s['bbox'][2]for s in spans)
        offset=abs((left+right)/2-page.rect.width/2)
        assert offset<=1.,('PDF公式主体偏离版心中心',number,offset)
        assert mark[0]>right and 475<=mark[2]<=506,('PDF式号位置错误',number)
        results.append({'式号':number,'PDF页':page.number+1,'主体左右边界_pt':[round(left,3),round(right,3)],'中心偏差_pt':round(offset,3),'式号右边界_pt':round(mark[2],3)})
    return results
