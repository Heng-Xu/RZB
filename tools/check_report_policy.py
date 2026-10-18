"""报告强制约束：源稿、原生Word格式、目录域和实际PDF一致性。"""
from pathlib import Path
import argparse, hashlib, json, re

ROOT=Path(__file__).resolve().parents[1]

def body_policy_errors(sources,sections,tables,patterns):
    """只检查正文实际引用的材料，并保留可编辑源的位置。"""
    errors=[]
    def scan(value,location,stack=()):
        for equation in re.finditer(r'\\begin\{aligned\}([\s\S]*?)\\end\{aligned\}',value):
            for row in re.split(r'\\\\',equation[1]):
                if len(re.findall(r'&(?:=|\\(?:leq|geq))',row))>1:
                    line_no=value[:equation.start()].count('\n')+1
                    errors.append(f'{location}:{line_no}: 多组等式或约束须分别换行，防止Word公式截断')
        for line_no,line in enumerate(value.splitlines(),1):
            from report_prose_review import LOGIC_RULES,is_explanatory_line
            for rule,pattern in patterns.items():
                if rule in LOGIC_RULES and is_explanatory_line(line):continue
                candidate = re.sub(r'\$[^$]*\$', '', line) if rule=='最终功率电量展示保留两位小数' else line
                if rule=='最终功率电量展示保留两位小数' and '$$' in line:continue
                for match in re.finditer(pattern,candidate):
                    if (rule=='项目核心增长指标采用完整名称'
                        and line[max(0,match.start()-3):match.start()]=='年负荷'
                        and re.search(r'DL/T\s*5729[—－-]2023',line)):
                        # 仅豁免明确对应2023版规范的原指标名称。
                        continue
                    errors.append(f'{location}:{line_no}: {rule}: {match[0]}')
            for kind,key in re.findall(r'\[\[(SECTION|TABLE):([^\]]+)\]\]',line):
                token=(kind,key)
                if token in stack:
                    raise ValueError(f'正文引用循环: {location}:{line_no}: {token}')
                data=sections if kind=='SECTION' else tables
                name='生成段落.json' if kind=='SECTION' else '计算表.json'
                scan(data[key],f'{location}:{line_no} → {name}[{key}]',stack+(token,))
    for location,value in sources:
        scan(value,location)
    return errors

def matrix_policy_errors(source,location,sections,tables,pattern):
    """只在推荐矩阵范围内应用时长展示规则，保留源行与嵌套资产键。"""
    lines=source.splitlines();start=next(i for i,line in enumerate(lines) if line.startswith('## 5.6 '))
    end=next(i for i,line in enumerate(lines[start+1:],start+1) if line.startswith('## 5.7 '))
    part='\n'.join(lines[start:end])
    found=body_policy_errors([(location,part)],sections,tables,{'推荐矩阵不展示已排除的时长参数':pattern})
    return [re.sub(re.escape(location)+r':(\d+)',lambda m:location+':'+str(int(m[1])+start),e,count=1) for e in found]

def check_figure_paragraphs(doc,policy,expected_count):
    from docx.oxml.ns import qn
    figure_policy=policy['figure_paragraph_format']
    figures=doc._element.xpath('.//w:body//w:p[.//w:drawing or .//w:pict]')
    captions=[p for p in doc._element.xpath('.//w:body//w:p') if re.match(r'^图\d+-',''.join(p.xpath('.//w:t/text()')).strip())]
    assert len(figures)==expected_count, '图片段落数与正文源不一致'
    assert len(captions)==len(figures),'图片与图题数量不一致'
    if figure_policy['inline_only']:
        assert not doc._element.xpath('.//w:body//wp:anchor | .//w:body//w:pict'),'图片应使用嵌入型，避免浮动定位覆盖段落居中'
    for i,p in enumerate(figures+(captions if figure_policy['include_captions'] else []),1):
        label=f'图片或图题段落{i}'
        assert p.xpath('./w:pPr/w:jc/@w:val')==[figure_policy['alignment']],label+'必须居中对齐'
        inds=p.xpath('./w:pPr/w:ind');assert len(inds)==1,label+'必须显式设置无缩进'
        for attr in figure_policy['zero_indent_attributes']:
            assert inds[0].get(qn('w:'+attr))=='0',label+'缩进未清零: '+attr
    return {'图片无缩进居中':len(figures),'图题无缩进居中':len(captions)}

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root',type=Path,default=ROOT)
    parser.add_argument('--stage',choices=['source','docx','final'],default='source')
    args=parser.parse_args();root=args.root.resolve()
    policy=json.loads((root/'tools/report_policy.json').read_text())
    report=root/'研究报告';text=report/'03_MD/当前版';asset=report/'02图表/当前版';review=report/'05_review/当前版'
    tables=json.loads((asset/'计算表.json').read_text());sections=json.loads((asset/'生成段落.json').read_text())
    from report_storage_display import storage_duration_table
    assert tables['storage_duration_cost']==storage_duration_table(), '计算表.json[storage_duration_cost]: 有效功率、费用或两位小数展示未对应完整精度计算'
    sources=sorted(text.glob('0[1-6] *.md'));assert len(sources)==6
    body='\n\n'.join(p.read_text() for p in sources)
    body=re.sub(r'\[\[SECTION:([^\]]+)\]\]',lambda m:sections[m[1]],body)
    body=re.sub(r'\[\[TABLE:([^\]]+)\]\]',lambda m:tables[m[1]],body)
    errors=body_policy_errors(
        [(str(p.relative_to(root)),p.read_text()) for p in sources],
        sections,tables,policy['forbidden_body_patterns'])
    # 参数切片必须有响应；完整原始结果仍保留在计算资料。
    for key in re.findall(r'\[\[TABLE:([^\]]+)\]\]', '\n'.join(p.read_text() for p in sources)+'\n'+'\n'.join(sections.values())):
        if not key.startswith(('restore_storage_price_', 'restore_support_duration_h_', 'restore_line_cost_coefficient_', 'restore_initial_transfer_fraction_', 'restore_eta_', 'restore_growth_rate_')):
            continue
        rows=[[c.strip() for c in ln.strip('|').split('|')] for ln in tables[key].splitlines() if ln.startswith('|')][2:]
        if len(rows)>1 and len({tuple(row[1:]) for row in rows})==1:
            errors.append('无响应参数切片不得重复展示: '+key)
    for m in re.finditer(r'由(\d+(?:\.\d+)?)(%| MW| MVA|万元)?变为\1\2(?=[，；。 ])',body):
        errors.append('相同数值不能表述为变化: '+m[0])
    for path in sources:
        for line,text_line in enumerate(path.read_text().splitlines(),1):
            if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',text_line):
                errors.append(f'{path.relative_to(root)}:{line}: 非法控制字符，检查公式转义')
    from prepare_case_correlations import maintain as check_actual_correlations
    check_actual_correlations(root,check=True)
    chapter_three=next(p.read_text() for p in sources if p.name.startswith('03 '))
    actual_section=chapter_three.split('## 3.5 核心参数依据与相关性分析')[1].split('## 3.6 ')[0]
    for key in ['current_controlled_correlations','current_input_correlations','new_parameter_bins']:
        assert '[[TABLE:'+key+']]' in actual_section, '第3.5节缺少参数或相关分析: '+key
    assert not any('[[TABLE:'+key+']]' in body for key in ['actual_corr_pizhou','actual_corr_city','actual_corr_target']), '实际年度相关表已退出正文'
    # 以已确认参数库复算展示，明确区分实际年度关联与受控仿真组内相关。
    import sync_current_report_results as current_results
    assert tables['economic_assumptions']==current_results.economic_assumptions_table(), '计算表.json[economic_assumptions]: 经济费率展示未对应锁定参数的两位小数值'
    generated,generated_sections,cells,index,data=current_results.parameter_assets()
    for key,value in generated.items():
        assert tables[key]==value,'计算表.json:'+key+' 与锁定参数库不一致'
    for key,value in generated_sections.items():
        assert sections[key]==value,'生成段落.json:'+key+' 与锁定参数库不一致'
    # 所有指标及筛选表必须实际进入第3.5节，五参数矩阵由锁定节点归纳。
    required=['extended_parameter_nodes','extended_indicator_correlations_1','extended_indicator_correlations_2',
        'parameter_redundancy_groups','parameter_selection_summary']
    for key in required:
        assert '[[TABLE:'+key+']]' in actual_section, '第3.5节缺少筛选证据:'+key
    assert len(cells)==8 and len(index)==2
    chapter_five=next(p.read_text() for p in sources if p.name.startswith('05 '))
    assert '### 5.4.8 保留2021实际容量的方案对照' in chapter_five
    referenced='\n'.join(p.read_text() for p in sources)+'\n'+'\n'.join(sections.values())
    for key in ['uncompressed_annual_summary','uncompressed_cost_summary','typical_transfer_matrix','current_transfer_sequence_second']:
        assert '[[TABLE:'+key+']]' in referenced,'第五章缺少已确认结果：'+key
    assert not re.search(r'\[\[TABLE:current_matrix_',referenced),'矩阵不得单独按线路费用分类'
    assert all('[[TABLE:lookup_matrix_'+f'{c}_{kd:.2f}'+']]' in sections['new_matrix_tables']
               for c in range(3) for kd in [1.,1.1,1.15]),'缺少电缆分区或负荷分散系数矩阵'
    assert '线路成本系数' not in sections['new_matrix_tables']
    chapter_path=next(p for p in sources if p.name.startswith('05 '))
    errors.extend(matrix_policy_errors(chapter_five,str(chapter_path.relative_to(root)),sections,tables,
                                       policy['matrix_display']['forbidden_duration_pattern']))
    from report_citations import verify as verify_citations
    citation_result=verify_citations(root,args.stage)
    assert '取0.80' in chapter_five and '故障站取1.00' in chapter_five,'反向承载系数须明确'
    assert '[[TABLE:new_correlations]]' not in body,'旧版无条件仿真相关表已退出正文'
    # 连续式中解释应合并到一段，不能用多个空行分列符号解释。
    if re.search(r'式中[^\n]*\n\s*\n式中',body):errors.append('相邻式中说明应合并')
    for equation in re.findall(r'\$\$([\s\S]*?)\$\$',body):
        if equation.count(r'\qquad')>=2 and r'\begin{aligned}' not in equation:
            errors.append('包含多组关系的长公式须显式分行')
    for name in re.findall(r'\[\[FIG:([^|]+)\|',body):
        if not (asset/(name+'.png')).is_file():errors.append('缺少图件: '+name)
    profiles=report/'数据来源/当前版/区域时序'
    summary=json.loads((profiles/'summary.json').read_text())
    import csv
    for region in policy['profile_regions']:
        p=profiles/f'{region}_8760.csv';rows=list(csv.DictReader(p.open()))
        assert len(rows)==policy['required_profile_hours']
        assert summary[region]['curve_sha256']==hashlib.sha256(p.read_bytes()).hexdigest()
        assert len({r['hour'] for r in rows})==len(rows)
        assert all(r['net_mw'] and r['abs_mw'] for r in rows)
    assert not errors,'\n'.join(sorted(set(errors)))
    result={'状态':'通过','阶段':args.stage,'规则版本':policy['version'],'源稿SHA256':{str(p.relative_to(root)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}}
    if args.stage!='source':
        from docx import Document
        from docx.oxml.ns import qn
        docpath=report/'04_word/研究报告终稿.docx';doc=Document(docpath)
        matrices=[tab for tab in doc.tables if len(tab.columns)==6 and len(tab.rows)==13
                  and [c.text for c in tab.rows[0].cells[:2]]==['年最大净负荷平均增长率','装机渗透率']]
        assert len(matrices)==9,'原生Word须包含九张完整电缆分区矩阵'
        for tab in matrices:
            for i,row in enumerate(tab.rows):
                for cell in row.cells:
                    for p in cell.paragraphs:
                        assert p.paragraph_format.keep_with_next==(i<12),'区间矩阵须按整表连续分页'
        assert doc.settings.element.xpath('.//m:mathFont/@m:val')==[policy['math_format']['font']],'数学字体须采用Times New Roman'
        for run in doc._element.xpath('.//m:r'):
            pr=run.find(qn('w:rPr'));assert pr is not None,'公式字符须显式指定字体'
            fonts=pr.find(qn('w:rFonts'));assert fonts is not None
            assert all(fonts.get(qn('w:'+a))==policy['math_format']['font'] for a in ['ascii','hAnsi','cs'])
            assert run.find(qn('m:rPr')).find(qn('m:nor')) is not None,'公式字符须保留明确的正斜体字形'
        def zero_indent(p):
            inds=p.xpath('./w:pPr/w:ind');assert len(inds)==1,'未显式设置段落缩进'
            ind=inds[0]
            for a in ['firstLine','firstLineChars','left','right']:
                assert ind.get(qn('w:'+a))=='0',(a,ind.attrib)
            assert all(ind.get(qn('w:'+a),'0')=='0' for a in ['hanging','hangingChars','leftChars','rightChars','startChars','endChars'])
        result.update(check_figure_paragraphs(doc,policy,len(re.findall(r'\[\[FIG:',body))))
        from report_math_layout import check_equation_layout
        layouts=check_equation_layout(doc)
        display=[t._tbl.xpath('./w:tr/w:tc[2]/w:p')[0] for t in layouts]
        expected=len(re.findall(r'\\tag\{',body));assert len(display)==expected
        for p in display:
            zero_indent(p);assert p.xpath('./w:pPr/w:jc/@w:val')==['center']
            row=p.getparent().getparent()
            assert all(x.xpath('./w:pPr/w:keepNext') for x in row.xpath('.//w:p')),'公式须与式中解释保持同页'
        explanations=[p for p in doc._element.xpath('.//w:p') if ''.join(p.xpath('./w:r/w:t/text()')).startswith('式中')]
        for p in explanations:zero_indent(p)
        bookmarks={b.get(qn('w:name')):b for b in doc._element.xpath('.//w:bookmarkStart')}
        headings=json.loads((review/'目录标题.json').read_text());anchors={x[2] for x in headings}
        assert anchors<=bookmarks.keys()
        fields=[];field_results={}
        for p in doc._element.xpath('.//w:p[w:r/w:instrText]'):
            m=re.search(r'PAGEREF\s+(\w+)',''.join(p.xpath('./w:r/w:instrText/text()')))
            if not m:continue
            anchor=m[1];assert anchor in anchors
            assert p.xpath('./w:hyperlink/@w:anchor')==[anchor]
            fields.append(anchor)
            active=False;cached=[]
            for r in p.findall(qn('w:r')):
                kinds=r.xpath('./w:fldChar/@w:fldCharType')
                if 'separate' in kinds:active=True;continue
                if 'end' in kinds:active=False
                if active:cached.extend(r.xpath('./w:t/text()'))
            field_results[anchor]=''.join(cached)
        assert len(fields)==len(headings) and set(fields)==anchors
        for anchor in anchors:
            p=bookmarks[anchor].getparent();assert p[0].tag==qn('w:pPr'),'书签不能置于段落属性前'
        result.update(独立公式=len(display),式中段落=len(explanations),目录链接=len(fields),WordSHA256=hashlib.sha256(docpath.read_bytes()).hexdigest())
        hashes=json.loads((review/'装配源文件登记.json').read_text())
        assert all(hashlib.sha256((root/k).read_bytes()).hexdigest()==v for k,v in hashes.items()),'源稿变化后未重新装配'
        if args.stage=='final':
            import pymupdf
            pdfpath=report/'04_word/研究报告终稿.pdf';pdf=pymupdf.open(pdfpath)
            assert any(s in str(pdf.metadata).lower() for s in ['wps','kingsoft'])
            major=[entry for entry in pdf.get_toc() if entry[0]==1]
            first=next(entry[2] for entry in major if entry[1].startswith('第一章'))
            references=next(entry[2] for entry in major if entry[1]=='参考文献')
            body_pages=references-first
            assert abs(body_pages-policy['target_body_pages'])<=policy['body_page_tolerance'],('六章正文篇幅偏离约160页',body_pages)
            result['六章正文页数']=body_pages
            matrix_pdf=[]
            for page in pdf:
                compact=re.sub(r'\s+','',page.get_text())
                if not re.search(r'表5-\d+电缆长度占比.+容载比区间矩阵',compact):continue
                for tab in page.find_tables().tables:
                    header=re.sub(r'\s+','',''.join(x or '' for x in tab.extract()[0]))
                    if header.startswith('年最大净负荷平均增长率装机渗透率'):
                        assert (tab.row_count,tab.col_count)==(13,6),('WPS矩阵跨页或不完整',page.number+1)
                        matrix_pdf.append(dict(页码=page.number+1,行数=tab.row_count,列数=tab.col_count))
            assert len(matrix_pdf)==9,('WPS完整矩阵数量不符',len(matrix_pdf))
            result['PDF矩阵整表核验']=matrix_pdf
            assert len(pdf)>=policy['minimum_pdf_pages'],f'PDF篇幅不足：{len(pdf)}页'
            # 同时验证渲染后的图片位置，防止Word结构正确但导出排版偏移。
            pdf_images=[(page,im) for page in pdf for im in page.get_image_info()]
            assert len(pdf_images)==len(doc.inline_shapes),'PDF图件数量与Word不一致'
            left=doc.sections[0].left_margin.pt;right=doc.sections[0].right_margin.pt
            offsets=[]
            for page,im in pdf_images:
                r=pymupdf.Rect(im['bbox']);offset=abs((r.x0+r.x1)/2-(page.rect.width+left-right)/2)
                assert offset<=policy['figure_paragraph_format']['pdf_center_tolerance_pt'],('PDF图片未居中',page.number+1,offset)
                offsets.append(offset)
            result['PDF图片最大居中偏差_pt']=round(max(offsets,default=0),4)
            # 单行公式另外检查渲染位置；多行公式按对称栏宽核验并逐式目视复核。
            manifest=json.loads((asset/'公式/公式清单.json').read_text())
            formula_offsets=[]
            for equation in manifest:
                if r'\begin{aligned}' in equation['tex'] or r'\sum' in equation['tex']:continue
                found=[]
                for page in pdf:
                    found.extend((page,pymupdf.Rect(word[:4])) for word in page.get_text('words') if word[0]>468 and word[4]=='（'+equation['number']+'）')
                assert len(found)==1,('独立式号缺失或重复',equation['number'])
                page,number=found[0];spans=[]
                for block in page.get_text('dict')['blocks']:
                    for line in block.get('lines',[]):
                        for span in line['spans']:
                            x0,y0,x1,y1=span['bbox']
                            if x0>125 and x1<470 and y0<number.y1+2 and y1>number.y0-15:spans.append(span)
                assert spans,('未找到公式主体',equation['number'])
                offset=abs((min(s['bbox'][0] for s in spans)+max(s['bbox'][2] for s in spans))/2-(page.rect.width+left-right)/2)
                assert offset<=policy['math_format']['pdf_center_tolerance_pt'],('PDF公式主体未居中',equation['number'],offset)
                formula_offsets.append(dict(公式=equation['number'],页码=page.number+1,中心偏差_pt=round(offset,4)))
            result['PDF单行公式中心核验']=formula_offsets
            toc=json.loads((review/'目录页码核对.json').read_text())
            norm=lambda x:re.sub(r'\s+','',x)
            pages=[norm(p.get_text()) for p in pdf]
            # 正文第一页由第一章标题和印刷页码共同确定；封面+两页目录。
            lines=[norm(l) for p in list(pdf)[:3] for l in p.get_text().splitlines()]
            for entry in toc:
                title=norm(entry['标题']);n=entry['页码']
                assert field_results[entry['锚点']]==str(n),('Word页码域缓存与实际页不一致',entry)
                assert title in pages[n-1] and n>=4,entry
                hits=[l for l in lines if l.startswith(title)]
                assert len(hits)==1 and re.search(r'(\d+)$',hits[0][len(title):]),(entry,hits)
                assert int(re.search(r'(\d+)$',hits[0][len(title):])[1])==n,(entry,hits)
            # 检查WPS实际导出的目录内部链接是否指向对应标题页。
            links=[l for p in list(pdf)[:3] for l in p.get_links() if l.get('kind')==1]
            assert len(links)>=len(toc),('PDF目录缺少内部跳转',len(links),len(toc))
            title_links=[l for l in links if l['from'].x0<300]
            assert len(title_links)==len(toc)
            assert [l['page']+1 for l in title_links]==[e['页码'] for e in toc],'目录标题跳转目标与对应页码不一致'
            for entry in toc:
                assert entry['页码']-1 in {l['page'] for l in links},entry
            result.update(PDF页数=len(pdf),PDFSHA256=hashlib.sha256(pdfpath.read_bytes()).hexdigest())
            exported=json.loads((review/'WPS原生导出登记.json').read_text())
            assert exported['WordSHA256']==result['WordSHA256'] and exported['PDFSHA256']==result['PDFSHA256'],'Word/PDF导出登记未同步'
    result['引文一致性']=citation_result
    out=review/('1017甲方批注与逐年案例更新' if (review/'1017甲方批注与逐年案例更新/修前哈希.json').exists() else '1011局部优化');out.mkdir(exist_ok=True)
    (out/f'{args.stage}约束检查.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps({k:v for k,v in result.items() if k!='源稿SHA256'},ensure_ascii=False))

if __name__=='__main__':main()
