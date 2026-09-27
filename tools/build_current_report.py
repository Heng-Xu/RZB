"""Fill the existing project template from the seven report chapters and the unified bibliography."""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import json, re, hashlib, subprocess, sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from lxml import etree

ROOT=Path.cwd()
SRC=ROOT/'研究报告/03_MD'
OUT=ROOT/'研究报告/04_word'
QA=ROOT/'研究报告/05_review'
TMP=Path('/tmp/rzb-closed-loop'); TMP.mkdir(exist_ok=True)
TEMPLATE=ROOT/'研究报告/中期研究报告初稿/03+XX项目-研究报告.docx'
TITLE='徐州地区分布式新能源高渗透率地区110kV电网容载比弹性指标优化研究'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
FILES=sorted(SRC.glob('[0-9][0-9] *.md'))+[SRC/'参考文献.md']
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def child(parent,tag,**attrs):
 e=OxmlElement(tag)
 for k,v in attrs.items():e.set(qn('w:'+k),str(v))
 parent.append(e);return e
def fonts(rpr):
 f=rpr.find(qn('w:rFonts'))
 if f is None:f=child(rpr,'w:rFonts')
 for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','方正仿宋_GBK'),('cs','Times New Roman')]:f.set(qn('w:'+k),v)
 for k in list(f.attrib):
  if 'theme' in k.lower():del f.attrib[k]
def inline_text(node):
 if isinstance(node,list):return ''.join(inline_text(x) for x in node)
 if not isinstance(node,dict):return ''
 t,c=node.get('t'),node.get('c')
 if t=='Str':return c
 if t in ['Space','SoftBreak','LineBreak']:return ' '
 if t=='Math':return c[1]
 if t=='Code':return c[1]
 if t=='Link':return inline_text(c[1])
 if t=='Image':return inline_text(c[1])
 return inline_text(c)

def wrap_math(latex):
 if latex.startswith("A_{m,y}(t)="):
  latex=latex.replace("+K_{m,y}\\sum",r" \\ {}+K_{m,y}\sum").replace("+\\mu_m",r" \\ {}+\mu_m")
  return r"\begin{aligned}"+latex+r"\end{aligned}"
 """Split only top-level horizontal separators; never split braces/delimiters."""
 depth=0; delim=0; cuts=[];i=0
 while i<len(latex):
  if latex[i]=='{' and (i==0 or latex[i-1]!='\\'):depth+=1
  elif latex[i]=='}' and (i==0 or latex[i-1]!='\\'):depth-=1
  if latex.startswith('\\left',i):delim+=1
  if latex.startswith('\\right',i):delim-=1
  m=re.match(r'\\(?:qquad|quad)\b',latex[i:])
  if m and depth==0 and delim==0:cuts.append((i,i+len(m[0])))
  i+=1
 if len(latex)>105 and cuts:
  segments=[];start=0
  for a,b in cuts:segments.append(latex[start:a].strip());start=b
  segments.append(latex[start:].strip())
  return '\\begin{aligned}'+r' \\ '.join(segments)+'\\end{aligned}'
 return latex

def walk(node, fn):
 if isinstance(node,dict):
  fn(node)
  for v in node.values():walk(v,fn)
 elif isinstance(node,list):
  for v in node:walk(v,fn)

def main():
 assert len(FILES)==8
 head=subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()
 upstream='本次以本地文件及SHA256为依据'
 OUT.mkdir(exist_ok=True);QA.mkdir(exist_ok=True)
 manifest={'commit':head,'upstream':upstream,'template':{'path':str(TEMPLATE.relative_to(ROOT)),'sha256':sha(TEMPLATE)},'sources':[{'path':str(p.relative_to(ROOT)),'sha256':sha(p)} for p in FILES]}
 asts=[]; equations=[]; heads=[]
 for file in FILES:
  ast=json.loads(subprocess.check_output(['pandoc','-f','markdown+tex_math_single_backslash-implicit_figures','-t','json',str(file)]))
  def trans(n):
   if n.get('t')=='Str':
    n['c']=re.sub(r'(?<=\d)(kV|MW(?:h)?|MVA|km|h|s)(?=$|[^A-Za-z])',r' \1',n['c'])
    n['c']=re.sub(r'(kV)(?=\d)',r'\1 ',n['c'])
   if n.get('t')=='Image':
    final=ROOT/'研究报告/02图表/final_figures'/Path(n['c'][2][0]).name
    n['c'][2][0]=str(final.resolve() if final.exists() else (file.parent/n['c'][2][0]).resolve())
   if n.get('t')=='Link':
    target=n['c'][2][0]
    if not re.match(r'^[a-z]+:|^#',target):n['c'][2][0]=str((file.parent/target).resolve())
  walk(ast,trans)
  for b in ast['blocks']:
   if b['t']=='Header':
    level=b['c'][0]; text=inline_text(b['c'][2]);heads.append({'level':level,'text':text})
   if b['t']=='Para' and len(b['c'])==1 and b['c'][0]['t']=='Math' and b['c'][0]['c'][0]['t']=='DisplayMath':
    raw=b['c'][0]['c'][1]; tag=re.search(r'\\tag\{([^}]+)\}',raw)
    assert tag,raw
    num=tag[1];clean=re.sub(r'\\tag\{[^}]+\}','',raw).strip()
    wrapped=wrap_math(clean);b['c'][0]['c'][1]=wrapped
    b['c'] += [{'t':'Space'},{'t':'Str','c':'（'+num+'）'}]
    equations.append({'source':file.name,'number':num,'latex':clean,'layout_latex':wrapped})
  asts.append(ast)
 ast=deepcopy(asts[0]);ast['blocks']=[b for a in asts for b in a['blocks']]
 (TMP/'content.json').write_text(json.dumps(ast,ensure_ascii=False))
 proc=subprocess.run(['pandoc','-f','json','-t','docx','--reference-doc',str(TEMPLATE),str(TMP/'content.json'),'-o',str(TMP/'content.docx')],capture_output=True,text=True)
 (TMP/'pandoc.log').write_text(proc.stderr)
 assert proc.returncode==0,proc.stderr
 assert 'Could not convert' not in proc.stderr,proc.stderr
 doc=Document(TEMPLATE);content=Document(TMP/'content.docx')
 body=doc._element.body
 originals=list(body)
 cover=deepcopy(doc.sections[0]._sectPr);front=deepcopy(doc.sections[1]._sectPr);normal=deepcopy(doc.sections[2]._sectPr)
 template_styles={s.name:s for s in doc.styles}
 for s in doc.styles:fonts(s.element.get_or_add_rPr())
 # Keep template font sizes, spacing and numbering definitions. The explicit
 # chapter numbers in Markdown override inherited automatic heading numbers.
 for name in ['Heading 1','Heading 2','Heading 3']:
  ppr=doc.styles[name].element.get_or_add_pPr();num=ppr.find(qn('w:numPr'))
  if num is not None:ppr.remove(num)
  child(ppr,'w:widowControl',val='1')
  if name=='Heading 1':child(ppr,'w:pageBreakBefore')
 for name in ['Normal','表']:
  child(doc.styles[name].element.get_or_add_pPr(),'w:widowControl',val='1')
 for s in content.styles:
  if s.name not in template_styles:
   # Supplemental semantic styles are based on existing template styles.
   cloned=deepcopy(s.element)
   for part in list(cloned):
    if etree.QName(part).localname in ['rPr','pPr']:cloned.remove(part)
   base=cloned.find(qn('w:basedOn'))
   if base is None:base=child(cloned,'w:basedOn',val=doc.styles['Normal'].style_id)
   else:base.set(qn('w:val'),doc.styles['Normal'].style_id)
   fonts(child(cloned,'w:rPr'));doc.styles.element.append(cloned)
 stylemap={s.style_id:doc.styles[s.name].style_id for s in content.styles if s.name in doc.styles}
 for e in list(body):body.remove(e)
 for e in originals[:21]:body.append(deepcopy(e))
 # Fill cover text while keeping the template runs and paragraph geometry.
 for p in doc.paragraphs:
  if p.text=='XX项目研究报告':
   for r in p.runs:r.text=''
   p.runs[0].text=TITLE+'\n研究报告'
  elif p.text=='XX公司':
   for r in p.runs:r.text=''
   p.runs[0].text='国网江苏电力设计咨询有限公司'
  elif p.text=='XX年XX月':
   for r in p.runs:r.text=''
   p.runs[0].text='2026年9月'
 # Front section retains the template's Roman numbering; no invented abstract.
 toc=deepcopy(originals[25]);tc=toc.find(qn('w:sdtContent'))
 proto=deepcopy(next(e for e in tc if e.tag==qn('w:p')))
 for e in list(tc):tc.remove(e)
 from docx.text.paragraph import Paragraph
 title=Paragraph(proto,doc._body);title.text='目  录';title.style=doc.styles['TOC Heading'];tc.append(proto)
 tp=OxmlElement('w:p');tc.append(tp);p=Paragraph(tp,doc._body)
 r=p.add_run();child(r._r,'w:fldChar',fldCharType='begin')
 r=p.add_run();ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=' TOC \\o "1-3" \\h \\z \\u ';r._r.append(ins)
 r=p.add_run();child(r._r,'w:fldChar',fldCharType='separate')
 p.add_run('目录更新中')
 r=p.add_run();child(r._r,'w:fldChar',fldCharType='end')
 body.append(toc)
 boundary=OxmlElement('w:p');child(boundary,'w:pPr').append(front);body.append(boundary)
 relmap={}
 for rid,rel in content.part.rels.items():
  if rel.reltype.endswith('/image'):
   newrid,_=doc.part.get_or_add_image(BytesIO(rel.target_part.blob));relmap[rid]=newrid
  elif rel.is_external:relmap[rid]=doc.part.relate_to(rel.target_ref,rel.reltype,is_external=True)
 for e in content._element.body:
  if e.tag==qn('w:sectPr'):continue
  e=deepcopy(e)
  for n in e.iter():
   for k,v in list(n.attrib.items()):
    if k.startswith('{'+NS['r']+'}') and v in relmap:n.set(k,relmap[v])
   if n.tag in [qn('w:pStyle'),qn('w:rStyle'),qn('w:tblStyle')]:
    val=n.get(qn('w:val'));n.set(qn('w:val'),stylemap.get(val,val))
  body.append(e)
 body.append(normal)
 # Pandoc writes display math and the external tag to consecutive paragraphs.
 # Join them using template-width tab stops so the number sits on the right.
 equation_index=0
 for p in list(doc.paragraphs):
  wrappers=p._p.xpath('./m:oMathPara')
  if not wrappers:continue
  nxt=p._p.getnext()
  number=''.join(nxt.xpath('.//w:t/text()')) if nxt is not None else ''
  assert re.fullmatch(r'（4-\d+）',number),number
  eq=equations[equation_index];equation_index+=1
  p.paragraph_format.first_line_indent=Pt(0)
  p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
  p.paragraph_format.line_spacing=1.2
  p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=True
  pr=p._p.get_or_add_pPr();child(pr,'w:ind',firstLine='0',firstLineChars='0')
  tabs=child(pr,'w:tabs');child(tabs,'w:tab',val='center',pos='3730');child(tabs,'w:tab',val='right',pos='8312')
  p.add_run('\t');tab=p._p[-1];p._p.remove(tab);p._p.insert(1,tab)
  wrapper=wrappers[0];pos=list(p._p).index(wrapper)
  for math in wrapper.findall(qn('m:oMath')):p._p.insert(pos,deepcopy(math));pos+=1
  p._p.remove(wrapper)
  p.add_run('\t'+number)
  nxt.getparent().remove(nxt)
  for r in p.runs:r.font.size=Pt(12)
  pprfonts=child(pr,'w:rPr');fonts(pprfonts);child(pprfonts,'w:sz',val='24')
  size=11 if len(eq['latex'])>230 else 12
  for mr in p._p.xpath('.//m:r'):
   wr=mr.find(qn('w:rPr'))
   if wr is None:
    wr=OxmlElement('w:rPr');mr.insert(1 if mr.find(qn('m:rPr')) is not None else 0,wr)
   fonts(wr);child(wr,'w:sz',val=str(size*2))
 assert equation_index==29
 # Apply typography with inherited sizes; ordinary runs follow Normal's 16 pt.
 for p in doc.paragraphs:
  for r in p.runs:fonts(r._r.get_or_add_rPr())
  if p._p.xpath('.//m:oMath') and re.fullmatch(r'（4-\d+）',p.text.strip()):
   if p._p.xpath('.//m:oMathPara'):
    p.paragraph_format.first_line_indent=Pt(0)
    pr=p._p.get_or_add_pPr();ind=pr.find(qn('w:ind'))
    if ind is not None:ind.set(qn('w:firstLineChars'),'0')
    p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=True
    p.paragraph_format.line_spacing=1.35
    # Long expressions use smaller variable text, remaining editable OMML.
    idx=next((i for i,x in enumerate(equations) if p.text.strip()=='（'+x['number']+'）'),None)
    size=12 if idx is None else (11 if len(equations[idx]['latex'])>230 else 12)
    for mr in p._p.xpath('.//m:r'):
     wr=mr.find(qn('w:rPr'))
     if wr is None:wr=OxmlElement('w:rPr');mr.insert(1 if mr.find(qn('m:rPr')) is not None else 0,wr)
     fonts(wr);child(wr,'w:sz',val=str(size*2))
  if re.match(r'^[图表]\d+-\d+\s',p.text):
   p.style=doc.styles['Caption'];p.paragraph_format.first_line_indent=Pt(0)
   pr=p._p.get_or_add_pPr();ind=pr.find(qn('w:ind'))
   if ind is not None:ind.set(qn('w:firstLineChars'),'0')
   p.paragraph_format.keep_together=True
   p.paragraph_format.keep_with_next=p.text.startswith('表')
  if p._p.xpath('.//w:drawing'):
   p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
   p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
  if p.text.startswith('资料来源：'):
   for r in p.runs:r.font.size=Pt(10)
   p.paragraph_format.line_spacing=1.2
  if p.style.name.startswith('Heading'):
   # Prevent template numbering duplication without changing the words.
   child(p._p.get_or_add_pPr(),'w:numPr');np=p._p.get_or_add_pPr().find(qn('w:numPr'));child(np,'w:numId',val='0')
 # Fix image geometry to fit inherited printable width and a caption on-page.
 width=int(doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin)
 for im in doc.inline_shapes:
  ratio=im.height/im.width;im.width=width;im.height=int(width*ratio)
  maxheight=int(Cm(17.2))
  if im.height>maxheight:im.width=int(maxheight/ratio);im.height=maxheight
 table_audit=[]
 for index,t in enumerate(doc.tables):
  n=len(t._tbl.tr_lst[0].tc_lst)
  if len(t.columns)==0:
   for _ in range(n):child(t._tbl.tblGrid,'w:gridCol',w='1385')
  t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  caption=''
  prev=t._tbl.getprevious()
  if prev is not None:caption=''.join(prev.xpath('.//w:t/text()'))
  wide=n>=8; avail=11906-2880 if wide else 11906-3594
  # Landscape applies only to the wide source table; other sections are copied.
  if wide:
   before=OxmlElement('w:p');sp=deepcopy(normal)
   # First body section must retain the template's page-number restart.
   child(before,'w:pPr').append(sp)
   # Move section break before table title rather than between title and table.
   if prev is not None:prev.addprevious(before)
   else:t._tbl.addprevious(before)
   after=OxmlElement('w:p');land=deepcopy(normal)
   sz=land.find(qn('w:pgSz'));sz.set(qn('w:w'),'16838');sz.set(qn('w:h'),'11906');sz.set(qn('w:orient'),'landscape')
   for node in land.findall(qn('w:pgNumType')):land.remove(node)
   child(after,'w:pPr').append(land);t._tbl.addnext(after);avail=16838-3594
  sizes=[1/n]*n
  if n==6:sizes=[.22,.12,.165,.165,.165,.16];sizes=[s/sum(sizes) for s in sizes]
  if n==3:sizes=[.19,.42,.39]
  if n==4:sizes=[.22,.12,.27,.39] if 'SHA256' in t.cell(0,3).text else [.22,.28,.22,.28]
  for col,portion in zip(t.columns,sizes):col.width=Twips(int(avail*portion))
  for rowidx,row in enumerate(t.rows):
   trpr=row._tr.get_or_add_trPr();child(trpr,'w:cantSplit')
   if rowidx==0:child(trpr,'w:tblHeader')
   for cell,portion in zip(row.cells,sizes):
    cell.width=Twips(int(avail*portion));cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
     p.style=doc.styles['表'];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.15
     p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
     p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
     p.paragraph_format.keep_with_next=False
     for r in p.runs:fonts(r._r.get_or_add_rPr());r.font.bold=rowidx==0
     # Long paths and SHA256 may wrap within a word; the characters stay intact.
     child(p._p.get_or_add_pPr(),'w:wordWrap',val='0')
  # Template contains no table object; use its existing table paragraph style
  # and preserve Pandoc's simple grid, with explicit printable dimensions.
  pr=t._tbl.tblPr;tw=pr.find(qn('w:tblW'))
  if tw is None:tw=child(pr,'w:tblW')
  tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(avail))
  borders=pr.find(qn('w:tblBorders'))
  if borders is None:borders=child(pr,'w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   child(borders,'w:'+edge,val='single',sz='4',color='000000')
  table_audit.append({'index':index,'title':caption,'columns':n,'rows':len(t.rows),'landscape':wide,'width_twips':avail})
 if any(x['landscape'] for x in table_audit):
  endnum=doc._element.body.sectPr.find(qn('w:pgNumType'))
  if endnum is not None:endnum.attrib.pop(qn('w:start'),None)
 # Fill inherited header; the full title wraps poorly, so retain the project
 # title at a size that fits two lines within the inherited header band.
 for s in doc.sections:
  for p in s.header.paragraphs:
   if 'XX项目' in p.text:
    for r in p.runs:r.text=''
    p.runs[0].text=TITLE
   for r in p.runs:fonts(r._r.get_or_add_rPr());r.font.size=Pt(9)
  # The template footer is empty: add PAGE in its existing centered tab.
  for p in s.footer.paragraphs:
   if not p._p.xpath('.//w:instrText'):
    p.style=doc.styles['Footer'];p.add_run('\t')
    r=p.add_run();child(r._r,'w:fldChar',fldCharType='begin')
    r=p.add_run();ins=OxmlElement('w:instrText');ins.text=' PAGE ';r._r.append(ins)
    r=p.add_run();child(r._r,'w:fldChar',fldCharType='end')
   for r in p.runs:fonts(r._r.get_or_add_rPr())
 setting=doc.settings.element
 upd=setting.find(qn('w:updateFields'))
 if upd is None:upd=child(setting,'w:updateFields',val='true')
 # Cambria Math remains the mathematical glyph provider for Word's stretchy
 # operators; variable and Latin runs have explicit Times New Roman fonts.
 for e in doc._element.xpath('.//m:r'):
  rpr=e.find(qn('w:rPr'))
  if rpr is None:rpr=OxmlElement('w:rPr');e.insert(1 if e.find(qn('m:rPr')) is not None else 0,rpr)
  fonts(rpr)
 for p in doc.paragraphs:
  if p._p.xpath('.//m:oMath') and not re.fullmatch(r'（4-\d+）',p.text.strip()):
   p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
   p.paragraph_format.first_line_indent=None
   p.paragraph_format.line_spacing=None
   pr=p._p.get_or_add_pPr()
   for e in list(pr):
    if e.tag in [qn('w:ind'),qn('w:rPr')]:pr.remove(e)
   for mr in p._p.xpath('.//m:r'):
    wr=mr.find(qn('w:rPr'))
    if wr is None:wr=OxmlElement('w:rPr');mr.append(wr)
    for e in list(wr):
     if e.tag==qn('w:sz'):wr.remove(e)
    child(wr,'w:sz',val='32')
 # Ensure inherited fonts in cover, TOC, header/footer and defaults are unified.
 for e in doc.styles.element.xpath('.//w:rFonts'):fonts(e.getparent())
 filename=OUT/'研究报告终稿.docx';doc.save(filename)
 manifest['images']=[{'path':n['c'][2][0],'sha256':sha(Path(n['c'][2][0]))} for a in asts for b in a['blocks'] for n in (b.get('c',[]) if b['t']=='Para' else []) if isinstance(n,dict) and n.get('t')=='Image']
 manifest['equations']=equations;manifest['headings']=heads;manifest['tables']=table_audit
 (QA/'装配清单.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps({'docx':str(filename),'headings':len(heads),'equations':len(equations),'tables':len(doc.tables),'images':len(doc.inline_shapes),'pandoc_warnings':proc.stderr},ensure_ascii=False))
if __name__=='__main__':main()
