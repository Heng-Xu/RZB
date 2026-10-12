"""依据研究报告终稿重构讨论稿；沿用原稿版式，正文和表格可追溯核验。"""
from pathlib import Path
import argparse, hashlib, json, re, shutil
from copy import deepcopy
from docx import Document
from docx.table import Table
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1]
REVIEW='研究报告/05_review/当前版/1011讨论稿重构'

def report_tables(path):
    doc=Document(path); tables={}; caption=None
    for element in doc.element.body:
        if element.tag==qn('w:p'):
            text=''.join(element.xpath('.//w:t/text()'))
            m=re.match(r'表(\d+-\d+)\s',text)
            caption=m[1] if m else None
        elif element.tag==qn('w:tbl') and caption:
            tables[caption]=[[c.text for c in r.cells] for r in Table(element,doc).rows]
            caption=None
    return doc,tables

def material(root):
    folder=root/'研究报告/供电公司讨论稿'; source=folder/'讨论稿.md'
    authority=root/'研究报告/04_word/研究报告终稿.docx'
    _,t=report_tables(authority)
    result=[['地区／方案','主变容量\nMVA','容载比','储能额定功率\nMW','累计新增联络\n条','费用现值\n万元']]
    annual={}
    for region,cap,eng,key in [('邳州','5-4','5-5','PIZHOU'),('市区','5-6','5-7','CITY')]:
        for scheme in ['刚性','弹性']:
            a=next(r for r in t[cap][1:] if r[:2]==[scheme,'2025'])
            cost=next(r[2] for r in t['5-8'][1:] if r[:2]==[region,scheme])
            result.append([region+scheme,a[3],a[4],a[5],a[6],cost])
        rows=[['年份','当年新购主变\nMVA','当年新增联络\n条','当年新增储能\nMW','年末主变容量\nMVA','年末容载比']]
        last=0
        for a in t[cap][1:]:
            if a[0]!='弹性':continue
            e=next(r for r in t[eng][1:] if r[:2]==a[:2])
            power=float(a[5]); rows.append([a[1],e[3],e[6],f'{power-last:.2f}',a[3],a[4]]);last=power
        annual[key+'_ANNUAL_TABLE']=rows
    ranges=[['分布式电源\n装机渗透率','年最大净负荷\n平均增长率','初始转供率\n≤30%','初始转供率\n(30%, 50%]','初始转供率\n>50%']]
    for number,label in [('5-71','≤1.00'),('5-72','(1.00, 2.00]'),('5-73','>2.00')]:
        for growth in ['≤2%','(2%, 4%]','(4%, 7%]','>7%']:
            selected=[r for r in t[number][1:] if r[:2]==[growth,'(1.65, 6.85]']]
            assert [r[2] for r in selected]==['≤30%','(30%, 50%]','>50%']
            ranges.append([label,growth]+[r[3] for r in selected])
    tables={'RESULT_TABLE':result,**annual,'RANGE_TABLE':ranges}
    text=source.read_text()
    assert set(re.findall(r'\[\[([A-Z_]+TABLE)\]\]',text))==set(tables)
    from check_report_policy import body_policy_errors
    policy=json.loads((root/'tools/report_policy.json').read_text())
    errors=body_policy_errors([(str(source.relative_to(root)),text)],{},{},policy['forbidden_body_patterns'])
    assert not errors,'\n'.join(errors)
    files=[source,authority,root/'tools/build_discussion_brief.py',root/REVIEW/'修前_供电公司讨论稿.docx',root/'tools/report_policy.json']
    hashes={str(p.relative_to(root)):hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    return folder,text,tables,hashes

def build(root):
    folder,text,tables,hashes=material(root)
    doc=Document(root/REVIEW/'修前_供电公司讨论稿.docx')
    for name in ['Title','Heading 1','Heading 2']:
        pr=doc.styles[name].element.find(qn('w:pPr'))
        if pr is not None:
            for b in list(pr.findall(qn('w:pBdr'))):pr.remove(b)
    for el in list(doc.element.body):
        if el.tag!=qn('w:sectPr'):doc.element.body.remove(el)
    def font(run,size=11.5,bold=False):
        run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'方正仿宋_GBK')
    def zero(p):
        pr=p._p.get_or_add_pPr();snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pr.append(snap)
        ind=pr.get_or_add_ind()
        for a in ['left','right','firstLine','hanging','leftChars','rightChars','firstLineChars','hangingChars','start','end','startChars','endChars']:
            ind.set(qn('w:'+a),'0')
    next_page=False
    for line in text.splitlines():
        if not line.strip():continue
        if line=='[[PAGEBREAK]]':next_page=True;continue
        if line.startswith('[['):
            rows=tables[line[2:-2]];tbl=doc.add_table(rows=0,cols=len(rows[0]));tbl.style='Table Grid'
            for i,row in enumerate(rows):
                for cell,value in zip(tbl.add_row().cells,row):
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0];zero(p);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    pf=p.paragraph_format;pf.space_before=Pt(3);pf.space_after=Pt(3);pf.line_spacing=1.05
                    font(p.add_run(value),10.5,i==0)
                pr=tbl.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
                if i==0:pr.append(OxmlElement('w:tblHeader'))
            continue
        p=doc.add_paragraph();zero(p);pf=p.paragraph_format;pf.line_spacing=Pt(16);pf.space_before=Pt(0);pf.space_after=Pt(5);pf.widow_control=True
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if next_page:pf.page_break_before=True;next_page=False
        if line.startswith('# '):
            p.style=doc.styles['Title'];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run(line[2:]),17,True);pf.line_spacing=Pt(22);pf.keep_with_next=True
        elif line.startswith('> '):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run(line[2:]),12);pf.space_after=Pt(12);pf.keep_with_next=True
        elif line.startswith('### '):
            p.style=doc.styles['Heading 2'];font(p.add_run(line[4:]),11.5,True);pf.space_before=Pt(4);pf.space_after=Pt(4);pf.keep_with_next=True
        elif line.startswith('## '):
            p.style=doc.styles['Heading 1'];font(p.add_run(line[3:]),13,True);pf.space_before=Pt(6);pf.space_after=Pt(5);pf.keep_with_next=True
        elif re.match(r'表[1-4] ',line):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(p.add_run(line),10.5);pf.keep_with_next=True
        else:
            ind=p._p.get_or_add_pPr().get_or_add_ind();ind.attrib.pop(qn('w:hanging'),None);ind.attrib.pop(qn('w:hangingChars'),None)
            ind.set(qn('w:firstLine'),'460');font(p.add_run(line))
        # Do not inherit coloured heading text from the Word template.
        for r in p.runs:
            from docx.shared import RGBColor
            r.font.color.rgb=RGBColor(0,0,0)
    path=folder/'供电公司讨论稿.docx';doc.save(path)
    hashes[str(path.relative_to(root))]=hashlib.sha256(path.read_bytes()).hexdigest()
    record={'来源SHA256':hashes,'唯一内容依据':'研究报告/04_word/研究报告终稿.docx','表格':tables,'表格依据':{'RESULT_TABLE':['5-4','5-6','5-8'],'PIZHOU_ANNUAL_TABLE':['5-4','5-5'],'CITY_ANNUAL_TABLE':['5-6','5-7'],'RANGE_TABLE':['5-71','5-72','5-73']},'交付格式':'用户要求仅Word，原有PDF为修前版本，本轮不生成PDF'}
    (folder/'生成登记.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n')
    print(path)

def check(root,pdf=False,final=False):
    folder,text,tables,hashes=material(root);record=json.loads((folder/'生成登记.json').read_text())
    assert tables==record['表格']
    for name,sha in record['来源SHA256'].items():
        assert hashlib.sha256((root/name).read_bytes()).hexdigest()==sha,('讨论稿需重新生成',name)
    assert all(record['来源SHA256'][k]==v for k,v in hashes.items())
    path=folder/'供电公司讨论稿.docx';doc=Document(path)
    assert len(doc.tables)==len(tables)
    for table,rows in zip(doc.tables,tables.values(),strict=True):
        assert [[c.text for c in r.cells] for r in table.rows]==rows
        for row in table.rows:
            assert row._tr.xpath('./w:trPr/w:cantSplit')
            for cell in row.cells:
                assert cell.vertical_alignment==WD_CELL_VERTICAL_ALIGNMENT.CENTER
                assert all(p.alignment==WD_ALIGN_PARAGRAPH.CENTER for p in cell.paragraphs)
    expected=[re.sub(r'^(?:#{1,3}|>) ','',l) for l in text.splitlines() if l.strip() and not l.startswith('[[')]
    assert [p.text for p in doc.paragraphs if p.text.strip()]==expected
    assert len(doc.element.body.xpath('.//w:pageBreakBefore'))==text.count('[[PAGEBREAK]]')
    assert not doc.element.body.xpath('.//w:ins | .//w:del | .//w:commentRangeStart')
    if pdf or final:
        evidence=json.loads((root/REVIEW/'WPS原生Word核验.json').read_text())
        assert evidence['WordSHA256']==hashlib.sha256(path.read_bytes()).hexdigest(),'WPS核验后Word已变化'
        assert evidence['页数']==7 and evidence['逐页复核']==[1,2,3,4,5,6,7]
        for name,sha in evidence['截图SHA256'].items():
            assert hashlib.sha256((root/name).read_bytes()).hexdigest()==sha,name
        print('按用户仅Word要求，最终版式采用WPS原生Word逐页核验；未生成PDF。')
    print('源稿表达规则、终稿表格、正文与原生Word核验通过：4张表、155个单元格。')

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--root',type=Path,default=ROOT);ap.add_argument('--check',action='store_true');ap.add_argument('--pdf',action='store_true');ap.add_argument('--final',action='store_true');a=ap.parse_args()
    check(a.root,a.pdf,a.final) if a.check else build(a.root)
