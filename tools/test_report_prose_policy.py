"""验证排除式说明检查能拦截复发、保留实质约束并定位生成文本。"""
from pathlib import Path
import json
import unittest
from check_report_policy import body_policy_errors

POLICY=json.loads((Path(__file__).with_name('report_policy.json')).read_text())

class ProsePolicyTests(unittest.TestCase):
    def test_transfer_power_limit_and_diversity_sampling_remain_current(self):
        for value in ['新增联络的工程计算限值取8.00 MVA，按0.95功率因数折算有功上限为7.60 MW。',
                      '| 负荷分散系数 | 1.00、1.05、1.10、1.15 | 4 |']:
            errors=self.check('[[TABLE:sample]]',tables={'sample':value})
            self.assertTrue(any('计算表.json[sample]:1' in e for e in errors))
        for value in ['新增联络采用折算后的有功上限8.00 MW。',
                      '| 负荷分散系数 | 1.05～1.25，步长0.02 | 11 |',
                      '功率因数取0.95，站级负载率按主变额定容量计算。']:
            self.assertEqual(self.check(value),[])

    def test_extended_matrix_terms_locate_stale_scope_and_keep_real_definition(self):
        for value in ['历史回算与未来预测均保留所采用的补值，年度规划峰值继续使用原统计值。',
                      '其中1.15为本研究情景范围的上端。',
                      '五参数矩阵共九张表432个单元。']:
            self.assertTrue(any('生成段落.json[matrix]:1' in e for e in
                                self.check('[[SECTION:matrix]]',sections={'matrix':value})))
        for value in ['负荷分散系数在1.05—1.25内按0.02间隔取11个节点，以1.15划分两档。',
                      '矩阵由六张表288个单元组成，窗口宽0.20。',
                      '系数1.00对应各站正向峰值完全同时发生的边界情景。']:
            self.assertEqual(self.check(value),[])
    def test_transfer_terms_use_engineering_names(self):
        for value in ['既有边/MW', '新增边/MW', '年度正常及事故追加区段预案', '其中，学者Zhu等']:
            self.assertTrue(any('计算表.json[transfer]:1' in e for e in
                                self.check('[[TABLE:transfer]]', tables={'transfer':value})))
        for value in ['既有联络/MW', '新增联络/MW', '各年度正常运行和主变N-1时的负荷转供方案',
                      '其中，Zhenyu Zhu等对线路进行寿命费用评价。']:
            self.assertEqual(self.check(value), [])

    def test_verified_values_and_purpose_statements_locate_sources(self):
        for value in ['地区案例与通用模板的用途不同，分别保留其负荷来源和计算条件。',
                      '全市装机原截图未标单位，仅用于计算相对倍数。',
                      '| 年份 | 全市装机数值 |',
                      '原截图未标单位，本研究仅使用其相对2025年的增长倍数。']:
            errors=self.check('[[TABLE:forecast]]', tables={'forecast':value})
            self.assertTrue(any('计算表.json[forecast]:1' in e for e in errors))
        for value in ['全市光伏装机容量预测为7090.00 MW。',
                      '额定容量以MVA计，有功功率以MW计。',
                      '现有资料未提供2021—2022年逐月光伏装机记录。',
                      '推广仿真采用与地区案例相同的成本模型和寻优方法。']:
            self.assertEqual(self.check(value), [])

    def test_annual_case_model_wording(self):
        for value in ['以该比例乘当年最大正向净负荷形成年度预算。']:
            errors=self.check('[[SECTION:annual]]',{'annual':value})
            self.assertTrue(any('生成段落.json[annual]:1' in e for e in errors))
        for value in ['既有能力采用收资MW值，所有年份固定封顶。',
                      '源荷区段内负荷和光伏共同切换。',
                      '由于电源仍归属原站，反向压力按电源原归属保守处理。',
                      '若仅将用电负荷转出而电源保持接入原站，本地反送需求可能增加。']:
            self.assertEqual(self.check(value),[])
    def test_three_measure_wording_locates_unwanted_implementation_details(self):
        for value in ['在配套费用参考值的20%范围内优先减少联络。',
                      '在该配套费用参考值的120%以内优先减少联络。',
                      '源荷区段切换数为整数决策。']:
            errors=self.check('[[SECTION:annual]]', {'annual':value})
            self.assertTrue(any('生成段落.json[annual]:1' in e for e in errors))
        for value in ['主变增容、储能配置和10 kV联络三种措施联合寻优。',
                      '弹性方案不限制容载比上限。', '储能柜数采用整数变量。']:
            self.assertEqual(self.check(value), [])
    def test_body_omits_result_lock_maintenance_process(self):
        for value in ['推广矩阵沿用已确认参数结果，各单元为有限范围。',
                      '地区案例保留原锁定配置，在8760小时序列上校核。']:
            self.assertTrue(self.check(value))
        for value in ['推广矩阵各单元为有效节点形成的有限范围。',
                      '地区案例在各年8760小时源荷上分别校核正常和主变N-1。']:
            self.assertEqual(self.check(value),[])
    def test_equipment_terms_and_storage_precision(self):
        for value in ['基期主变按离散规格配置。', '离散设备投运改变容量。',
                      '市区3座模拟站参与地区计算。', '有效功率为0.08170 MW。']:
            errors=self.check('[[SECTION:equipment]]',{'equipment':value})
            self.assertTrue(any('生成段落.json[equipment]:1' in e for e in errors))
        for value in ['主变按可选额定容量逐台配置。', '文献采用离散粒子群算法。',
                      '采用相依离散卷积刻画不确定性。', '有效功率为0.82 MW。',
                      r'计算过程为$E=0.1634\ \mathrm{MWh}$。']:
            self.assertEqual(self.check(value), [])

    def test_storage_duration_power_and_unrounded_cost(self):
        from decimal import Decimal
        from report_storage_display import effective_power,storage_duration_table
        self.assertEqual(effective_power(1), Decimal('1.00'))
        self.assertEqual(effective_power(2), Decimal('.8170'))
        self.assertEqual(effective_power(4), Decimal('.4085'))
        table=storage_duration_table()
        self.assertIn('| 1.00 | 2.15 | 0.76 | 1.00 | 210.46 |', table)
        self.assertIn('| 2.00 | 2.15 | 0.76 | 0.82 | 257.60 |', table)
        with self.assertRaises(ValueError):effective_power(0)

    def test_economic_table_preserves_actual_rate_precision(self):
        from sync_current_report_results import economic_assumptions_table
        value=economic_assumptions_table()
        for rate in ['4.20%/年', '2.20%/年', '3.00%/年', '6.00%']:
            self.assertIn(rate,value)
        self.assertNotIn('初始投资4%/年',value)

    def test_matrix_time_exclusion_is_scoped_and_located(self):
        from check_report_policy import matrix_policy_errors
        source='地区案例支撑持续时间为2 h。\n## 5.6 五参数推广矩阵\n[[SECTION:matrix]]\n## 5.7 规划建议'
        pattern=POLICY['matrix_display']['forbidden_duration_pattern']
        for value in ['矩阵归纳2、4、6 h支撑条件。','储能支撑持续时间为4 h。']:
            errors=matrix_policy_errors(source,'第五章.md',{'matrix':value},{},pattern)
            self.assertTrue(any('第五章.md:3 → 生成段落.json[matrix]:1' in e for e in errors))
        self.assertEqual(matrix_policy_errors(source,'第五章.md',{'matrix':'按电缆分区列示五参数矩阵。'},{},pattern),[])

    def test_saved_discussion_uses_its_verified_report_policy(self):
        from unittest.mock import patch
        import revise_discussion_report_section as discussion
        prior = discussion.ROOT/'研究报告/05_review/当前版/1016研究报告结果同步'
        preserved = prior/'修前快照/研究报告/04_word/研究报告终稿.docx'
        with patch.object(discussion.r, 'REPORT', preserved):
            historical = discussion.policy_for_report()
        self.assertEqual(historical['version'], '2026-10-14-preliminary-baseline')
        value = '复合平均增长率为4%。'
        self.assertEqual(body_policy_errors([('已审定讨论稿', value)], {}, {},
                                           historical['forbidden_body_patterns']), [])
        self.assertTrue(self.check(value))
        with patch.object(discussion.r, 'REPORT', discussion.ROOT/'研究报告/04_word/研究报告终稿.docx'):
            self.assertEqual(discussion.policy_for_report(), POLICY)

    def test_capacity_growth_and_diagram_terminology(self):
        for value in ['保留2021原表容量。','平均增长率决定行位置。','复合平均增长率为4%。',
                      '本项目年负荷平均增长率为4%。',
                      '图1 六参数推广矩阵及三项核心参数独立扰动。']:
            self.assertTrue(self.check('[[TABLE:terms]]', tables={'terms':value}))
        for value in ['保留2021实际容量。','年最大净负荷平均增长率决定行位置。',
                      'DL/T 5729—2023的年负荷平均增长率分档。',
                      '文献采用年负荷增长率描述当地发展水平。']:
            self.assertEqual(self.check(value), [])

    def test_typical_transfer_response_and_matrix_conditions(self):
        for value in ['推荐矩阵按线路系数分表。','九个细节点均严格下降。','所有转供率节点严格下降。']:
            self.assertTrue(self.check('[[SECTION:matrix]]',{'matrix':value}))
        for value in ['四档推荐区间上下界均下降，九个细节点保留平台。',
                      '线路成本系数用于内部计价，电缆长度占比50%。',
                      '80%节点主变增加，新增联络减少，方案费用继续下降。']:
            self.assertEqual(self.check(value),[])
    def test_statistics_require_displayed_table_evidence(self):
        for value in ['Spearman按组内排序检验关联方向[12][13]。',
                      'Spearman最大绝对值为0.26，说明排序关联也较弱。']:
            errors=self.check('[[SECTION:correlations]]',{'correlations':value})
            self.assertTrue(any('生成段落.json[correlations]:1' in e for e in errors))
        self.assertEqual(self.check('五项参数中，电缆长度占比与初始转供率的Pearson系数为0.29。'),[])
        self.assertEqual(self.check('正文保留Pearson统计。',{'archive':'Spearman原始核验保留。'}),[])
    def check(self,text,sections=None,tables=None):
        return body_policy_errors([('第三章.md',text)],sections or {},tables or {},POLICY['forbidden_body_patterns'])

    def test_user_examples_are_detected(self):
        examples=[
            '不能将当前无响应解释为设备寿命没有经济意义。',
            '不能脱离设备组合单独确定。',
            '不能直接等同于设备新增过多。',
            '该比较保持线路技术能力不变，避免把造价变化解释成转供能力变化。',
            '二者的来源和统计范围不同，不将其作为两地区已核定的实际初始转供率直接比较。',
            '上述差异用于说明案例基础，不能通过两个地区的横向差值认定某一参数对容载比的独立影响。',
            '分布式电源电量渗透率下降并不表示当月光伏发电量减少。',
        ]
        for text in examples:
            with self.subTest(text=text):
                self.assertTrue(self.check(text))

    def test_supporting_file_narration_is_blocked(self):
        for value in ['具体参数组合与设备结果见配套工作簿。','求解精度及经济比较核验记录见配套资料。','逐点设备配置保存在配套工作簿中。']:
            self.assertTrue(self.check(value))
        for value in ['主变更换规格及费用见研究报告第5.4节。','导则建议范围见表4[1]。']:
            self.assertEqual(self.check(value),[])

    def test_paragraph_logic_and_explanation_exemptions(self):
        for value in ['上述逐年结果由初始设备和费用目标形成；是否需要超出上限，还需比较其他方案。',
                      '本次投运设备未达到更新年限。费用范围未计残值。所列费用用于比较。']:
            self.assertTrue(self.check(value))
        for value in ['方案由既定条件形成。若需论证超出上限的必要性，应检验建议范围内方案。',
                      '设备寿命超过评价期，因此本期费用由建设和运维构成。',
                      '式中，S为容量；因此处使用MVA单位，按功率因数折算。',
                      '注：输入按已计算节点归类；是否补充节点需结合新资料。',
                      '| 注 | 统一范围；因此采用同一单位 |']:
            self.assertEqual(self.check(value),[])

    def test_substantive_limits_are_preserved(self):
        examples=[
            '建议范围内不能满足供电安全准则时，可结合实际情况适当提高。',
            '当前候选集合内单靠联络无法弥补相关缺口。',
            '当前研究未开展逐时荷电状态仿真。',
            '现有资料缺少同范围用户年电量记录。',
            '既有联络不重复计建设费用，新增联络按投运年计费。',
            '储能额定功率保持0.50 MW。',
            '高渗透率情景下未取得可行方案。',
        ]
        for text in examples:
            with self.subTest(text=text):
                self.assertEqual(self.check(text),[])

    def test_known_input_and_template_misstatements_are_detected(self):
        for text in [
            '年网供最大负荷统一标定到容载比2.00。',
            '将地区名称和规模转换为五站仿真模板结构。',
            '两类方案本站总容量相同，仍需结合主变组合和邻站转接安排解释储能配置。',
        ]:
            with self.subTest(text=text):
                self.assertTrue(self.check(text))
        self.assertEqual(self.check('两类方案共用实际年度负荷和初始设备，推广计算采用独立五站模板。'),[])

    def test_nested_generated_text_is_located(self):
        errors=self.check('标题\n[[SECTION:analysis]]',{'analysis':'结果\n[[TABLE:comparison]]'},
                          {'comparison':'表题\n不能直接等同于设备新增过多。'})
        self.assertTrue(any('第三章.md:2 → 生成段落.json[analysis]:2 → 计算表.json[comparison]:2' in e for e in errors))

    def test_unreferenced_records_are_not_body(self):
        self.assertEqual(self.check('正文',{'archived':'不能直接等同于设备新增过多。'}),[])

    def test_actual_policy_version_name_is_preserved(self):
        self.assertEqual(self.check('2025年4月《配电网规划设计技术导则（修订稿）》第6.2.2.1条[69]。'),[])
        self.assertTrue(self.check('研究报告修订稿采用新结果。'))

    def test_multiple_constraints_need_explicit_math_break(self):
        self.assertTrue(self.check(r'\begin{aligned} A&\leq B, C&\leq D.\end{aligned}'))
        self.assertEqual(self.check(r'\begin{gathered} A\leq B,\\ C\leq D.\end{gathered}'),[])

    def test_independent_relations_centered_and_continuations_aligned(self):
        self.assertTrue(self.check(r'\begin{aligned} A&=B,\\ C&=D.\end{aligned}'))
        self.assertEqual(self.check(r'\begin{gathered} A=B,\\ C=D.\end{gathered}'),[])
        self.assertEqual(self.check(r'\begin{aligned} A&=B\\&+C.\end{aligned}'),[])

    def test_reference_cycle_fails(self):
        with self.assertRaisesRegex(ValueError,'正文引用循环'):
            self.check('[[SECTION:a]]',{'a':'[[SECTION:a]]'})

    def test_feedback_1013_prose_and_fixed_values(self):
        for text in ['采用收资《近5年容载比》中的记录。',
                     '储能价格另作专项敏感性分析，不增加主矩阵查表维度。',
                     '| 储能支撑持续时间 | 2、4、6 h | ≤2 h；(2 h,4 h]；>4 h |',
                     '| 负荷分散系数 | 1.00、1.10、1.15 | ≤1.10；>1.10 |']:
            with self.subTest(text=text):self.assertTrue(self.check(text))
        for text in ['根据项目提供的2021—2025年容量及负荷记录计算。',
                     '| 储能支撑持续时间 | 2、4、6 h | 2 h；4 h；6 h |',
                     '六项参数分别表征源荷、网络和措施条件。']:
            with self.subTest(text=text):self.assertEqual(self.check(text),[])

    def test_local_review_examples_are_located(self):
        examples=['甲方后续选取查表参数时，可结合本表判断，本版保留六项完整结果。',
                  '六项参数在本版均列入推广矩阵。',
                  '《配电网规划设计技术导则》修订资料第6.2.2条。',
                  '采用《徐州地区变电站全停全转情况分析汇总20250705》。',
                  '| 电量性质 | 典型场景及样本外推 |',
                  '注：7%增长分界附近另有7.1%、7.5%的54个补充组合。',
                  '刚性方案的被动超限由保留既有设备形成。']
        for text in examples:
            with self.subTest(text=text):
                errors=self.check('[[TABLE:review]]',tables={'review':text})
                self.assertTrue(any('计算表.json[review]:1' in e for e in errors))
        self.assertEqual(self.check('采用变电站全停条件下的负荷转移统计数据。'),[])
        self.assertEqual(self.check('主变容量保持不变，当年负荷减少，容载比随之升高。'),[])

    def test_selection_narrative_and_deleted_feedback_phrase(self):
        for value in ['扩展相关性分析用于支撑参数筛选，推广矩阵沿用原结果库。',
                      '矩阵沿用原数据，只筛掉时间指标。',
                      '两项分母分别采用典型场景和样本外推估计。']:
            errors=self.check('[[SECTION:selection]]',{'selection':value})
            self.assertTrue(any('生成段落.json[selection]:1' in e for e in errors))
        for value in ['十三项候选指标经相关强度及信息重叠分析确定五项推广参数。',
                      '矩阵归纳2、4、6 h支撑条件下的容量范围。',
                      '源荷与价格受控情景分别列明有效组数和节点数。']:
            self.assertEqual(self.check(value),[])

    def test_sampling_and_chinese_parameter_names(self):
        for value in ['源荷、网络和储能条件采用五站模板联合采样。',
                      '电量、峰段和价格条件采用97548个互异组合。',
                      '表3-17 五项推广参数的共线性核验',
                      '方差膨胀因子按VIFj=1/(1−Rj2)计算。',
                      'ηC、ḡ、α、r₀依次表示四项参数。']:
            errors=self.check('[[TABLE:sampling]]',tables={'sampling':value})
            self.assertTrue(any('计算表.json[sampling]:1' in e for e in errors))
        for value in ['仿真采用统一五站模板，各指标的采样节点与点数集中列于表3-11。',
                      '分布式电源装机渗透率与初始转供率的相关系数为0.00。',
                      '五项推广参数间的最大Pearson绝对值为0.29。']:
            self.assertEqual(self.check(value),[])

class EquationLayoutTests(unittest.TestCase):
    def matrix(self,t,alignments):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        m=t._tbl.xpath('.//m:oMath')[0];matrix=OxmlElement('m:m');pr=OxmlElement('m:mPr');cols=OxmlElement('m:mcs')
        for alignment in alignments:
            col=OxmlElement('m:mc');cp=OxmlElement('m:mcPr');jc=OxmlElement('m:mcJc');jc.set(qn('m:val'),alignment);cp.append(jc);col.append(cp);cols.append(col)
        pr.append(cols);matrix.append(pr)
        for _ in range(2):
            row=OxmlElement('m:mr')
            for _ in alignments:row.append(OxmlElement('m:e'))
            matrix.append(row)
        m.append(matrix)

    def sample(self):
        from docx import Document
        from docx.shared import Cm
        from docx.oxml import OxmlElement
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from report_math_layout import configure_equation_table
        d=Document();t=d.add_table(rows=1,cols=3);t.autofit=False;t.alignment=1
        widths=[Cm(1.3),Cm(12.0606),Cm(1.3)]
        for col,width in zip(t.columns,widths):col.width=width
        for c,width in zip(t.rows[0].cells,widths):c.width=width
        pr=t._tbl.tblPr;pr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblW').set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w',str(sum(x.twips for x in widths)))
        p=t.cell(0,1).paragraphs[0];mp=OxmlElement('m:oMathPara');mp.append(OxmlElement('m:oMath'));p._p.append(mp)
        t.cell(0,2).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
        configure_equation_table(t,p)
        return d,t

    def test_native_formula_is_centered_by_host_paragraph(self):
        from report_math_layout import check_equation_layout
        d,t=self.sample();self.assertEqual(len(check_equation_layout(d)),1)

    def test_independent_rows_reject_relation_aligned_columns(self):
        from report_math_layout import check_equation_layout
        d,t=self.sample();t.cell(0,2).paragraphs[0].add_run('（3-2）');self.matrix(t,['right','left'])
        with self.assertRaisesRegex(AssertionError,'逐行居中'):check_equation_layout(d)

    def test_centered_rows_and_aligned_continuation_are_accepted(self):
        from report_math_layout import check_equation_layout
        for number,columns in [('3-2',['center']),('4-7',['right','left']),('5-8',['right','left'])]:
            d,t=self.sample();t.cell(0,2).paragraphs[0].add_run('（'+number+'）');self.matrix(t,columns)
            self.assertEqual(len(check_equation_layout(d)),1)

    def test_short_objective_rejects_forced_line_break(self):
        from report_math_layout import check_equation_layout
        d,t=self.sample();t.cell(0,2).paragraphs[0].add_run('（5-9）');self.matrix(t,['center'])
        with self.assertRaisesRegex(AssertionError,'不得强制分行'):check_equation_layout(d)

    def pdf_rows(self,second_offset=0):
        import pymupdf
        d,t=self.sample();t.cell(0,2).paragraphs[0].add_run('（3-2）');self.matrix(t,['center'])
        pdf=pymupdf.open();page=pdf.new_page()
        page.insert_text((475,215),'（3-2）',fontname='china-s',fontsize=9)
        for y,value,offset in [(200,'A = B',0),(230,'C = D',second_offset)]:
            width=pymupdf.get_text_length(value,fontname='helv',fontsize=12)
            page.insert_text(((page.rect.width-width)/2+offset,y),value,fontsize=12)
        return d,pdf

    def test_pdf_rows_use_each_visible_row_center(self):
        from report_math_layout import check_pdf_equation_rows
        d,pdf=self.pdf_rows();rows=check_pdf_equation_rows(d,pdf,1)
        self.assertEqual(len(rows),2)
        self.assertTrue(all(r['中心偏差_pt']<1 for r in rows))

    def test_pdf_second_row_shift_is_rejected(self):
        from report_math_layout import check_pdf_equation_rows
        d,pdf=self.pdf_rows(second_offset=-18)
        with self.assertRaisesRegex(AssertionError,'未逐行居中'):check_pdf_equation_rows(d,pdf,1)

    def test_next_paragraph_accent_bbox_does_not_enter_formula(self):
        from report_math_layout import pdf_formula_spans
        spans=[dict(text='前文计算如下',bbox=(90,180,410,196)),
               dict(text='A = B',bbox=(200,210,400,225)),
               dict(text='式中变量说明',bbox=(90,240,410,255)),
               dict(text='accent',bbox=(126,220,126,260))]
        class Page:
            def get_text(self,_):return {'blocks':[{'lines':[{'spans':spans}]}]}
        self.assertEqual([s['text'] for s in pdf_formula_spans(Page(),(475,220,505,230))],['A = B'])

    def test_local_assembly_detects_layout_but_ignores_unused_namespace(self):
        from lxml import etree as E
        from docx.oxml.ns import qn
        from report_plain_revision import block_key
        d,t=self.sample();self.matrix(t,['center']);left=E.fromstring(E.tostring(t._tbl))
        right=E.fromstring(E.tostring(t._tbl).replace(b'<w:tbl ',b'<w:tbl xmlns:unused="urn:unused" ',1))
        self.assertEqual(block_key(left,{}),block_key(right,{}))
        right.find('.//'+qn('m:mcJc')).set(qn('m:val'),'left')
        self.assertNotEqual(block_key(left,{}),block_key(right,{}))

    def test_asymmetric_widths_are_rejected(self):
        from report_math_layout import check_equation_layout
        d,t=self.sample();t._tbl.tblGrid[0].set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w','800')
        with self.assertRaisesRegex(AssertionError,'左右栏须等宽'):check_equation_layout(d)

    def test_inherited_character_indent_is_rejected(self):
        from report_math_layout import check_equation_layout
        d,t=self.sample();t._tbl.xpath('./w:tr/w:tc[2]/w:p/w:pPr/w:ind')[0].set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}leftChars','200')
        with self.assertRaisesRegex(AssertionError,'缩进须全部清零'):check_equation_layout(d)

if __name__=='__main__':
    unittest.main()
