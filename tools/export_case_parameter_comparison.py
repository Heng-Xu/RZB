"""按终稿表5-4、表5-6导出旧、新参数的四张表，并逐格核验。"""
import argparse
import csv
import hashlib
import json
from math import isclose
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / '研究报告'
DOCX = REPORT / '04_word/研究报告终稿.docx'
TABLE_SOURCE = REPORT / '02图表/当前版/计算表.json'
OLD = REPORT / '试算/20261010_四小时案例'
NEW = REPORT / '试算/20261012_既有转供约束诊断'
REVIEW = REPORT / '05_review/当前版/1013案例参数对照'
LOCK = REPORT / '05_review/当前版/1012既有转供约束诊断/报告采用版本锁定清单.json'
OUTPUT = REPORT / '甲方反馈/10.11/表5-4及表5-6_新旧参数对照.xlsx'
CHECK = REVIEW / 'Excel导出核验.json'
HEADERS = ['方案', '年份', '年网供最大负荷/MW', '主变容量/MVA', '容载比', '储能额定功率/MW', '累计新线/条']
FIELDS = ['peak_mw', 'capacity_mva', 'clr', 'storage_nameplate_mw', 'cumulative_lines']
REGIONS = [('pizhou', '邳州', '5-4', '约67%'), ('city', '市区', '5-6', '约91%')]
CASES = [('rigid_storage', '刚性'), ('elastic_no_floor', '弹性')]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as handle:
        return list(csv.DictReader(handle))


def markdown_table(text):
    lines = [line for line in text.splitlines() if line.startswith('|')]
    return [[cell.strip() for cell in line.strip('|').split('|')] for line in lines if '---' not in line]


def displayed(rows):
    return [[row[0], str(row[1])] + [f'{value:.2f}' for value in row[2:6]] + [str(row[6])] for row in rows]


def source_tables():
    """只读取锁定的主方案；旧表同时与Word和当前计算表逐格比对。"""
    lock = json.loads(LOCK.read_text())
    assert lock['version'] == 'case_station_caps_normal_adjust_20261013'
    for entry in lock['model_files'] + lock['data_files']:
        assert sha(ROOT / entry['path']) == entry['sha256'], entry['path']
    main = {(r['region'], r['case']): r for r in lock['main']}
    assert len(main) == 4 and all(r['policy'] == 'station_caps' and r['verified'] for r in main.values())
    doc = Document(DOCX)
    word_tables = {}
    caption = ''
    for element in doc.element.body:
        if element.tag == qn('w:p'):
            caption = Paragraph(element, doc).text
        elif element.tag == qn('w:tbl') and caption.startswith(('表5-4 ', '表5-6 ')):
            word_tables[caption] = [[cell.text for cell in row.cells] for row in Table(element, doc).rows]
    table_json = json.loads(TABLE_SOURCE.read_text())
    sources = {DOCX, TABLE_SOURCE, LOCK}
    tables = []
    cases = []
    for region, name, number, rate in REGIONS:
        caption = f'表{number} {name}逐年容量及措施'
        for version in ['旧参数', '新参数']:
            rows = []
            for case, label in CASES:
                if version == '旧参数':
                    folder = OLD / region / ('elastic_storage' if case == 'elastic_no_floor' else case)
                else:
                    selected = main[region, case]
                    folder = NEW / selected['result_folder']
                    assert sha(folder / 'chosen_solution.npz') == selected['solution_sha256']
                    assert sha(folder.parent / 'inputs.json') == selected['input_sha256']
                    sources.update([folder / 'chosen_solution.npz', folder.parent / 'inputs.json'])
                years = read_csv(folder / 'years.csv')
                stations = read_csv(folder / 'stations.csv')
                projects = read_csv(folder / 'projects.csv')
                assert [int(a['year']) for a in years] == [2022, 2023, 2024, 2025]
                for annual in years:
                    year = int(annual['year'])
                    station_year = [s for s in stations if int(s['year']) == year]
                    capacity = sum(float(s['capacity_mva']) for s in station_year)
                    power = sum(int(s['storage_modules']) for s in station_year) * .1
                    cumulative_lines = sum(int(p['year']) <= year for p in projects)
                    assert isclose(capacity, float(annual['capacity_mva']), abs_tol=1e-8)
                    assert isclose(power, float(annual['storage_nameplate_mw']), abs_tol=1e-8)
                    assert cumulative_lines == int(annual['cumulative_lines'])
                    assert isclose(capacity / float(annual['peak_mw']), float(annual['clr']), abs_tol=1e-10)
                    rows.append([label, year] + [float(annual[field]) for field in FIELDS[:-1]] + [cumulative_lines])
                sources.update(folder / filename for filename in ['years.csv', 'stations.csv', 'projects.csv'])
                cases.append(dict(region=region, version=version, case=folder.name,
                                  source=str(folder.relative_to(ROOT)), years=4))
            if version == '旧参数':
                assert word_tables[caption] == [HEADERS] + displayed(rows), caption
                assert markdown_table(table_json[f'new_{region}_years']) == word_tables[caption]
            else:
                old_rows = tables[-1]['rows']
                assert [r[:3] for r in rows] == [r[:3] for r in old_rows], '新旧负荷及行序不一致'
            tables.append(dict(sheet=f'{name}_{version}', caption=caption,
                               note=version if version == '旧参数' else f'新参数（转供率{rate}）', rows=rows))
    return tables, sources, cases


def build(tables):
    book = Workbook()
    book.remove(book.active)
    book.properties.title = '表5-4及表5-6 新旧参数对照'
    book.properties.creator = 'xh'
    for table in tables:
        sheet = book.create_sheet(table['sheet'])
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_properties.tabColor = '4472C4' if '旧参数' in sheet.title else '548235'
        for row, value in [(1, table['caption']), (2, table['note'])]:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            cell = sheet.cell(row, 1, value)
            cell.font = Font(name='方正仿宋_GBK', size=14 if row == 1 else 11, bold=row == 1)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.row_dimensions[1].height = 28
        sheet.row_dimensions[2].height = 23
        sheet.row_dimensions[3].height = 9
        for row_number, values in enumerate([HEADERS] + table['rows'], start=4):
            sheet.row_dimensions[row_number].height = 36 if row_number == 4 else 25
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_number, column, value)
                cell.font = Font(name='方正仿宋_GBK' if row_number == 4 or column == 1 else 'Times New Roman',
                                 size=11, bold=row_number == 4)
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                if row_number > 4 and column > 1:
                    cell.number_format = '0' if column in (2, 7) else '0.00'
                if row_number == 4:
                    cell.border = Border(top=Side(style='medium'), bottom=Side(style='thin'))
                elif row_number == 12:
                    cell.border = Border(bottom=Side(style='medium'))
        for column, width in zip('ABCDEFG', [10, 10, 23, 19, 12, 23, 18]):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = 'C5'
        sheet.print_area = 'A1:G12'
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.print_options.horizontalCentered = True
        sheet.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4, header=.15, footer=.15)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUTPUT)


def verify(tables):
    book = load_workbook(OUTPUT, data_only=False)
    assert book.sheetnames == [t['sheet'] for t in tables] and len(book.sheetnames) == 4
    count = 0
    for table in tables:
        sheet = book[table['sheet']]
        assert sheet.max_row == 12 and sheet.max_column == 7
        assert str(sheet.merged_cells) == 'A1:G1 A2:G2'
        assert sheet['A1'].value == table['caption'] and sheet['A2'].value == table['note']
        for r, row in enumerate([HEADERS] + table['rows'], start=4):
            for c, expected in enumerate(row, start=1):
                cell = sheet.cell(r, c)
                actual = cell.value
                if isinstance(expected, (float, int)):
                    assert cell.data_type == 'n' and isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-10), (sheet.title, cell.coordinate)
                    assert cell.number_format == ('0' if c in (2, 7) else '0.00')
                else:
                    assert actual == expected, (sheet.title, cell.coordinate)
                count += 1
        assert all(sheet.cell(12, c).border.bottom.style == 'medium' for c in range(1, 8))
    book.close()
    return count


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check-only', action='store_true')
    args = parser.parse_args()
    tables, sources, cases = source_tables()
    sources.add(Path(__file__).resolve())
    if not args.check_only:
        build(tables)
    cells = verify(tables)
    if args.check_only:
        record = json.loads(CHECK.read_text())
        assert record['workbook']['sha256'] == sha(OUTPUT)
        for entry in record['sources']:
            assert sha(ROOT / entry['path']) == entry['sha256'], entry['path']
    else:
        record = dict(verified=True, workbook=dict(path=str(OUTPUT.relative_to(ROOT)), sha256=sha(OUTPUT)),
                      sheets=[t['sheet'] for t in tables], checked_table_cells=cells, cases=cases,
                      report_version=json.loads(LOCK.read_text())['version'],
                      notes='旧参数逐格匹配当前原生Word及计算表；新参数取已锁定正常允许调整供区的刚性和弹性无下限。原报告为4 h，本版沿用锁定方案的市区6 h、邳州2 h及逐站既有转供上限，属于前后版本对照。储能列为额定功率MW；新线列为累计新增数量。工作簿仅标注新转供率，按用户要求输出至甲方反馈/10.11。',
                      sources=[dict(path=str(p.relative_to(ROOT)), sha256=sha(p)) for p in sorted(sources)])
        REVIEW.mkdir(parents=True, exist_ok=True)
        CHECK.write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n')
    print(f'VERIFIED_EXCEL {len(tables)} tables / {cells} table cells: {OUTPUT.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
