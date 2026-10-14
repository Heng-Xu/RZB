"""讨论稿内容补充及研究报告格式复用；不修改锁定结果。"""
from copy import deepcopy
from pathlib import Path
from lxml import etree as E
from docx import Document
from docx.shared import Inches
import re
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
q=lambda n:'{'+W+'}'+n
NS={'w':W}
MATRIX_KEYS=['current_matrix_2_4_1.10','current_matrix_0_4_1.10','current_matrix_1_4_1.10','current_matrix_2_2_1.10','current_matrix_2_6_1.10','current_matrix_2_4_1.00','current_matrix_2_4_1.15']
FIGURES=['planning_tier_heatmap','discussion_cable_transfer','discussion_hours_eta','discussion_dispersion_transfer']
def plain(e):return ''.join(e.xpath('.//w:t/text()',namespaces=NS))
def report_table(report,prefix):
 d=Document(report)
 p=next(p for p in d.paragraphs if p.text.startswith(prefix))
 return E.fromstring(E.tostring(p._p.getnext()))
def rows(t):return [[plain(c) for c in r.findall(q('tc'))] for r in t.findall(q('tr'))]
def amend(changes,tables,additions,base,report,t,md_rows):
 changes.update({3:'1.1 双向供电任务与站级容量配置',4:'分布式光伏增加使主变同时承担正向供电与日间反向输送任务。容载比反映区域总体容量储备，具体配置还需结合主变规格、最大单台主变退出运行（N-1）后的剩余容量及站间转供能力确定。',5:'',6:'',7:'',8:'',9:'1.2 导则建议范围与实际容量需求',12:'1.3 主变、联络与储能联合比选'})
 # 甲方已修改段落保留有效措辞，仅更新必要事实和条件。
 changes[20]=plain_visible(base[20]).replace('年度负荷、新能源回算','年度负荷和新能源规模回算').replace('起点容载比均为2.0','起点容载比均约为2.0')
 changes[21]=plain_visible(base[21]).replace('刚性方案按容载比2.0控制','刚性方案新增容量按当年负荷的2.0倍控制；负荷回落时保留上一年度在役容量').replace('10kV','10 kV')
 changes[24]=plain_visible(base[24]).replace('转入、转出负荷保持恒定','转入、转出负荷合计相等').replace('邳州的负荷转供能力结合邳州实际区域电网转供能力进行测算，市区的负荷转供能力结合规程假设为40%。','两地区逐站转供能力均依据变电站全停条件下的统计测算，每站转出之和取原统计站外可供负荷与既有网络预算的较小值。地区全部电压等级负荷加权转供率为邳州约67%、市区约92%，110 kV案例采用匹配站点的MW上限。')
 changes[37]=plain_visible(base[37]).replace('储能退出','储能减少').replace('80616.00万元降至37138.45万元','2657.80万元降至0万元').replace('3862.44万元增至21083.57万元','0万元增至1446.64万元').replace('电缆占比高','电缆长度占比高').replace('（终稿第4.3节、表5-10）','（研究报告第4.3节、第5.4.1节）')
 changes.pop(41,None)
 # 统一可见正文中的报告名称与节引用，保留未涉及段落。
 for i,p in enumerate(base):
  if p.tag!=q('p'):continue
  value=changes.get(i,plain_visible(p));new=normalize(value)
  if new!=plain_visible(p):changes[i]=new
  elif i in changes:changes.pop(i)
 guide=rows(report_table(report,'表5-8 '))
 additions[53].insert(0,('table',guide));additions[53].insert(0,('caption','表4 高压分区或行政县（市、区）110～35 kV电网容载比选择范围[1]'))
 additions[53].append(('p','导则第6.2.2.1条a项适用于建议范围内无法满足供电安全准则的情形，c项要求加强中压联络、增加转供能力[1]。上述逐年结果由既定初始设备、网络约束和费用目标共同形成；是否需要超出分档上限，还需比较建议范围内可选设备与转供方案。'))
 # 表号4~8顺延；矩阵新增六张，对应三项背景各自其余两个取值。
 for index,items in additions.items():
  if index==53:continue
  additions[index]=[(kind,re.sub(r'^表([4-8])',lambda m:'表'+str(int(m[1])+1),value) if kind=='caption' else value) for kind,value in items]
 changes[74]='表8～表14分别比较电缆长度占比、储能支撑持续时间和负荷分散系数的条件变化。以表8为基准，表9、表10改变电缆长度占比，表11、表12改变储能支撑持续时间，表13、表14改变负荷分散系数；各表行按平均增长率和装机渗透率分组，列为初始转供率四档。'
 changes[75]='表8 电缆长度占比>30%、储能支撑4 h、负荷分散系数1.10时的容载比区间'
 changes[78]=changes[78].replace('表7','表8');changes[79]=changes[79].replace('表8','表15')
 additions[79]=[(k,v.replace('表9 ','表15 ') if k=='caption' else v) for k,v in additions[79]]
 titles=['电缆长度占比≤10%、储能支撑4 h、负荷分散系数1.10','电缆长度占比(10%,30%]、储能支撑4 h、负荷分散系数1.10','电缆长度占比>30%、储能支撑2 h、负荷分散系数1.10','电缆长度占比>30%、储能支撑6 h、负荷分散系数1.10','电缆长度占比>30%、储能支撑4 h、负荷分散系数1.00','电缆长度占比>30%、储能支撑4 h、负荷分散系数1.15']
 additions[76]=sum(([('caption',f'表{n} {title}时的容载比区间'),('table',md_rows(t[key]))] for n,title,key in zip(range(9,15),titles,MATRIX_KEYS[1:])),[])
 additions[76].append(('p','条件表比较同一分类下有效节点的容载比范围；下列图件固定其余参数，进一步识别设备配置发生变化的具体节点。'))
 captions=['图1 装机渗透率与年最大净负荷平均增长率对应的容载比','图2 电缆长度占比与初始转供率对应的容载比','图3 装机渗透率与储能支撑持续时间对应的容载比','图4 初始转供率与负荷分散系数对应的容载比']
 descriptions=['固定电缆长度占比10%、初始转供率40%、储能支撑持续时间4 h及负荷分散系数1.10。装机增长提高反向承载需求，负荷增长同时改变容量需求和容载比分母，图中分段变化对应离散设备配置。','固定装机渗透率2.50、平均增长率4%、储能支撑持续时间4 h及负荷分散系数1.10，比较初始转供率50%～90%的共同采样节点。电缆长度占比影响联络造价，初始转供率决定既有网络支援，两者共同影响新增主变与联络的经济选择。初始转供率为60%或70%时，电缆长度占比90%对应容载比2.43，其余电缆占比节点为2.26，体现较高联络造价下主变配置的替代作用。','固定平均增长率4%、电缆长度占比10%、初始转供率40%及负荷分散系数1.10。装机渗透率2.00时，支撑持续时间由2 h延长至4 h，容载比由2.05升至2.14；延长至6 h时保持2.14。设备配置的变化集中于部分装机节点，反映储能有效支撑功率与主变整台规格的共同作用。','固定装机渗透率1.50、平均增长率4%、电缆长度占比10%及储能支撑持续时间4 h，比较初始转供率20%～90%的共同有效节点。负荷分散程度改变各站局部峰值任务，转供能力决定邻站余量的可利用程度，初始转供率20%时，分散系数1.00、1.10、1.15分别对应容载比1.88、2.05、2.22；转供率达到40%后，三个分散系数下均为1.88，说明本组节点中网络支援改善能够缓解局部峰值对新增容量的要求。']
 for f,c,p in zip(FIGURES,captions,descriptions):additions[76]+=[('image',f),('figcaption',c),('p',p)]
 return changes,tables,additions

def plain_visible(p):return ''.join(p.xpath('.//w:t[not(ancestor::w:del)]/text()',namespaces=NS))
def normalize(s):
 s=s.replace('终稿','研究报告')
 s=re.sub(r'（研究报告第([\d.]+)((?:、[\d.]+)+)节）',lambda m:'（研究报告第'+m[1]+'节'+''.join('、第'+x+'节' for x in m[2].split('、')[1:])+'）',s)
 return s

def table_like_report(values,report):
 d=Document(report);tables=[E.fromstring(E.tostring(x._tbl)) for x in d.tables if not x._tbl.xpath('.//m:oMath')]
 normalized=[[v.replace('\n','') for v in row] for row in values]
 exact=next((x for x in tables if rows(x)==normalized),None)
 if exact is not None:return keep_table_together(deepcopy(exact))
 template=next(x for x in reversed(tables) if len(x.findall(q('tr'))[0].findall(q('tc')))==len(values[0]))
 tbl=deepcopy(template)
 for tr in list(tbl.findall(q('tr'))):tbl.remove(tr)
 for i,row in enumerate(values):
  tr=deepcopy(template.findall(q('tr'))[0 if i==0 else 1])
  for cell,value in zip(tr.findall(q('tc')),row):
   p=cell.find(q('p'));rpr=deepcopy(p.find('.//'+q('rPr')))
   for extra in list(cell.findall(q('p')))[1:]:cell.remove(extra)
   for child in list(p):
    if child.tag!=q('pPr'):p.remove(child)
   run=E.SubElement(p,q('r'));run.append(rpr)
   for j,part in enumerate(value.split('\n')):
    if j:E.SubElement(run,q('br'))
    E.SubElement(run,q('t')).text=part
  tbl.append(tr)
 return keep_table_together(tbl)

def keep_table_together(tbl):
 trs=tbl.findall(q('tr'))
 for i,tr in enumerate(trs):
  for p in tr.findall('.//'+q('p')):
   pr=p.find(q('pPr'));kn=pr.find(q('keepNext'))
   if kn is None:kn=E.SubElement(pr,q('keepNext'))
   kn.set(q('val'),'1' if i<len(trs)-1 else '0')
 return tbl

def image_paragraph(path,revision,parts):
 # python-docx仅用于生成嵌入图形片段，不保存或重建用户文档。
 d=Document();p=d.add_paragraph();p.add_run().add_picture(str(path),width=Inches(5.65));el=deepcopy(p._p)
 relns='http://schemas.openxmlformats.org/package/2006/relationships';rns='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
 rels=E.fromstring(parts['word/_rels/document.xml.rels']);rid='rIdDiscussion'+path.stem
 E.SubElement(rels,'{'+relns+'}Relationship',Id=rid,Type=rns+'/image',Target='media/'+path.name)
 parts['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
 parts['word/media/'+path.name]=path.read_bytes()
 for b in el.xpath('.//*[local-name()="blip"]'):b.set('{'+rns+'}embed',rid)
 for dp in el.xpath('.//*[local-name()="docPr"]'):dp.set('id',str(10000+revision.next));dp.set('name',path.stem)
 pr=E.Element(q('pPr'));el.insert(0,pr);E.SubElement(pr,q('jc')).set(q('val'),'center');E.SubElement(pr,q('spacing')).set(q('lineRule'),'auto')
 ind=E.SubElement(pr,q('ind'))
 for k in ['firstLine','hanging','left','right','start','end','firstLineChars','hangingChars','leftChars','rightChars','startChars','endChars']:ind.set(q(k),'0')
 E.SubElement(pr,q('keepNext'));E.SubElement(pr,q('rPr')).append(revision.mark('ins'))
 for r in list(el.findall(q('r'))):el.remove(r);ins=revision.mark('ins');ins.append(r);el.append(ins)
 return el


def verify_figures(root,audit,package):
 import pandas as pd,numpy as np
 raw=pd.read_csv(root/'研究报告/01研究推荐/当前版/参数联合情景_20261012/统一有效样本.csv')
 keys=list(raw.columns[:6]);indexed=raw.set_index(keys)
 images=[data for name,data in package.items() if name.startswith('word/media/') and name.endswith('.png')]
 for name in FIGURES:
  assert (root/'研究报告/02图表/当前版'/(name+'.png')).read_bytes() in images, '讨论稿图件与已核验绘图不一致:'+name
  if name=='planning_tier_heatmap':continue
  shown=pd.read_csv(audit/(name+'.csv')).set_index(keys)
  assert not shown.index.duplicated().any()
  expected=indexed.loc[shown.index]
  assert np.allclose(shown.select_dtypes('number'),expected.select_dtypes('number'),atol=1e-8)
 h=pd.read_csv(audit/'discussion_hours_eta.csv')
 assert [f'{x:.2f}' for x in h.loc[np.isclose(h['分布式电源装机渗透率'],2),'容载比']]==['2.05','2.14','2.14']
